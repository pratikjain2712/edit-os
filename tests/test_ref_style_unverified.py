"""Entries PubMed / CrossRef cannot verify are still styled, from the author's own text (refstyle.py): APA,
Harvard and EndNote shapes to Vancouver, NLM abbreviations from the full name, complete names for non-indexed
journals, missing fields queried, and nothing the author did not write. Synthetic references; HTTP is patched
(conftest: no network, Claude off)."""

import pytest
from docx import Document

from edit_os import biblio, docmodel, journals, queries, refengine, refs_other, refstyle, rules, style


class _Dir:
    def __init__(self, known):
        self.known = {journals.normalise(k): v for k, v in known.items()}

    def lookup(self, name):
        return self.known.get(journals.normalise(name))


def _journal(title, abbr):
    return journals.Journal(nlm_id=abbr, title=title, med_abbr=abbr, iso_abbr=abbr, issn_print="", issn_online="")


@pytest.fixture
def jehp(monkeypatch):
    """Nothing is found in PubMed or CrossRef; NLM lists one journal."""
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: [])
    monkeypatch.setattr(biblio, "pubmed_best_match", lambda term, retmax=10: [])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [])
    monkeypatch.setattr(biblio, "crossref_search", lambda text, rows=5: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    monkeypatch.setattr(biblio, "fallback_candidates", lambda *a, **k: [])
    monkeypatch.setattr(journals, "directory", lambda: _Dir({"Journal of Computers in Education": _journal(
        "Journal of computers in education", "J Comput Educ")}))
    monkeypatch.setattr(refengine, "_dictionary", (set(), set()))
    st = style.load("jehp")
    st.settings.setdefault("guide_replace", {})["smart_quotes"] = False
    return st


def test_apa_entry_not_found_is_styled_in_vancouver(jehp):
    text = ("Mu, T., Wang, F., Sang, L., & Xu, Y. (2020). Towards Effective Adaptive Learning: The Role of Learning "
            "Analytics. Journal of Computers in Education, 7(1), 1–26. https://doi.org/10.9999/jce.1")
    res = refengine.resolve(text, jehp)
    assert res.status == refengine.QUERY and res.record is None and res.styled
    assert res.formatted == ("Mu T, Wang F, Sang L, Xu Y. Towards effective adaptive learning: The role of learning "
                             "analytics. J Comput Educ 2020;7:1-26.")  # NLM's abbreviation of the full name


def test_harvard_entry(jehp):
    text = ("Okoro, A. and Lindqvist, B. B. (2019) Peer tutoring in rural nursing schools. Journal of Rural Nursing "
            "Practice, 12(3), pp. 140-148.")
    res = refengine.resolve(text, jehp)
    assert res.formatted == ("Okoro A, Lindqvist BB. Peer tutoring in rural nursing schools. Journal of Rural Nursing "
                             "Practice 2019;12:140-8.")  # not in NLM: the complete journal name


def test_endnote_entry_with_et_al(jehp):
    text = ("Sadrollahi, A., et al., The prevalence of abuse types in the elderly. Journal of Research and Health, "
            "2020. 10(1): p. 59-66.")
    res = refengine.resolve(text, jehp)
    assert res.formatted == ("Sadrollahi A, et al. The prevalence of abuse types in the elderly. Journal of Research "
                             "and Health 2020;10:59-66.")


def test_vancouver_entry_more_than_six_authors_dates_and_doi(jehp):
    text = ("Park H-H, Han S-H, Lee J. K., Kim A, Choi B, Jung C, Yoon D. A Study on the Knowledge of Nursing Students "
            "in India. Journal of the Korean Applied Science and Technology. 2022 Jun 30;39(3):386–396. "
            "https://doi.org/10.12925/JKOCS.2022.39.3.386")
    res = refengine.resolve(text, jehp)
    assert res.formatted == ("Park H-H, Han S-H, Lee JK, Kim A, Choi B, Jung C, et al. A study on the knowledge of "
                             "nursing students in India. Journal of the Korean Applied Science and Technology "
                             "2022;39:386-96.")


def test_supplement_kept(jehp):
    text = "Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Nurs Rural Pract. 2019;12(Suppl 2):S140-S148."
    assert refengine.resolve(text, jehp).formatted == ("Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. "
                                                      "Nurs Rural Pract 2019;12(Suppl 2):S140-8.")


@pytest.mark.parametrize("text, styled, query", [
    ("Rahayu, W. R., & Okoro, A. (2025). Improving regional clustering of tuberculosis cases. JAIEA, 4(2).",
     "Rahayu WR, Okoro A. Improving regional clustering of tuberculosis cases. JAIEA 2025;4.", "references.no_pages"),
    ("Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Journal of Rural Nursing Practice. 2019.",
     "Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Journal of Rural Nursing Practice 2019.",
     "references.no_volume"),
])
def test_missing_field_is_queried_not_guessed(jehp, text, styled, query):
    fields = {"kind": "journal_article", "authors": [], "title": "", "year": "", "doi": "", "pmid": ""}
    res = refengine.resolve(text, jehp, fields)  # (Claude's kind: an entry without a volume reads as a journal article)
    assert res.formatted == styled and (query, {}) in res.queries


def test_unparseable_authors_kept_and_queried(jehp):
    fields = {"kind": "journal_article", "authors": [], "title": "Effektivitas edukasi berbasis video", "year": "2026",
              "doi": "", "pmid": ""}
    text = "Idhayanti Ribkha Itha, Masini. Effektivitas Edukasi Berbasis Video. J Updat Keperawatan. 2026;6(1):329–35."
    res = refengine.resolve(text, jehp, fields)
    assert res.formatted == ("Idhayanti Ribkha Itha, Masini. Effektivitas edukasi berbasis video. J Updat Keperawatan "
                             "2026;6:329-35.")
    assert ("references.author_style", {}) in res.queries


def test_uncertain_split_left_to_the_cleanup(jehp):
    for text in ["M, K., H.R. H, and M. S, The effectiveness of therapy in older adults. Aging Psychology, 2025. 11(3): p. 331.",
                 "Smith J, Jones K. lung cancer in younger patients and older ones. 2019;5:1-4."]:
        res = refengine.resolve(text, jehp)
        assert res.status == refengine.QUERY and not res.styled and not res.formatted
        assert any(c.reason != "ref_styled" for c in _cleanup(text, jehp))  # the light cleanup still runs


def _cleanup(text, st):
    import tempfile, os
    doc = Document()
    for t in ["References", "1. " + text]:
        doc.add_paragraph(t)
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.docx")
        doc.save(path)
        st.settings["references"]["ai_parse"] = False
        return rules.run_all(docmodel.load(path), st, only=["ref_lookup", "references"]).changes


def test_no_invention_guard(jehp, monkeypatch):
    text = "Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Nurs Rural Pract. 2019;12:140-8."
    assert refstyle.invented("Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Nurs Rural Pract 2019;12:140-8.", text) == []
    assert refstyle.invented("Okoro A, Lindqvist B, Haddad C. Peer tutoring. Nurs Rural Pract 2019;12:140-8.", text) == ["Haddad", "C"]
    assert refstyle.invented("Okoro A. Peer tutoring. Nurs Rural Pract 2018;12:140-8.", text) == ["2018"]
    assert refstyle.invented("Lee JK, Park AC. X", "Lee J. K., Park Anna Clara. X") == []  # initials as written
    # a styled entry that would carry anything the author did not write is not used
    real = refstyle.format_article
    monkeypatch.setattr(refstyle, "format_article", lambda a, st, t: (real(a, st, t)[0].replace("Peer", "Pair"), [], ""))
    res = refengine.resolve(text, jehp)
    assert res.status == refengine.QUERY and not res.styled and not res.formatted


def test_apa_book_and_chapter_in_the_guide_order(jehp):
    st = jehp
    st.settings["references"]["other"].update(check_links=False, verify_books=False)
    book = "Tomlinson, C. A. (2001). How to differentiate instruction in mixed-ability classrooms (2nd ed.). ASCD."
    res = refs_other.resolve(book, st, refengine.parse(book))
    assert res.formatted == "Tomlinson CA. How to Differentiate Instruction in Mixed-Ability Classrooms. 2nd ed. ASCD; 2001."
    assert res.styled and ("references.publisher_location", {}) in res.queries
    chapter = ("Koedinger, K. R., & Corbett, A. (2006). Cognitive tutors: Technology bringing learning sciences to the "
               "classroom. In R. K. Sawyer (Ed.), The Cambridge Handbook of the Learning Sciences (pp. 61–77). "
               "Cambridge University Press.")
    res = refs_other.resolve(chapter, st, refengine.parse(chapter))
    assert res.formatted == ("Koedinger KR, Corbett A. Cognitive tutors: Technology bringing learning sciences to the "
                             "classroom. In: Sawyer RK, editor. The Cambridge Handbook of the Learning Sciences. "
                             "Cambridge University Press; 2006. p. 61-77.")
    endnote = "Tatara, T., Understanding elder abuse in minority populations. 2021: Routledge."
    p = refengine.parse(endnote)
    p.kind = "book"  # Claude's kind: nothing in the wording says book
    assert refs_other.resolve(endnote, st, p).formatted == (
        "Tatara T. Understanding Elder Abuse in Minority Populations. Routledge; 2021.")


def test_rules_keep_the_query_and_mark_provenance(jehp, tmp_path):
    jehp.settings["references"]["ai_parse"] = False
    doc = Document()
    for t in ["Discussion", "Peer tutoring works [1].", "References",
              "1. Okoro, A., & Lindqvist, B. (2019). Peer tutoring in rural nursing schools. Journal of Rural Nursing "
              "Practice, 12(3), 140–148."]:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    cs = rules.run_all(ms, jehp, only=["ref_lookup", "references"])
    mine = [c for c in cs.changes if c.reason == "ref_styled"]
    assert mine and all(c.reason == "ref_styled" for c in cs.changes)  # the light cleanup does not run on it too
    assert mine[0].provenance["label"] == refstyle.PROVENANCE_LABEL
    text = ms.paragraphs[-1].text
    for c in sorted(cs.changes, key=lambda c: c.start, reverse=True):
        text = text[:c.start] + c.new + text[c.end:]
    assert text == ("1. Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Journal of Rural Nursing Practice "
                    "2019;12:140-8.")
    qs = queries.collect(ms, jehp)
    assert any(q.check == "references.not_found" for q in qs.queries)  # "not verified" stays


def test_switch(jehp):
    jehp.settings["references"]["style_unverified"] = False
    text = "Okoro A, Lindqvist B. Peer tutoring in rural nursing schools. Nurs Rural Pract. 2019;12(3):140-148."
    assert not refengine.resolve(text, jehp).formatted
