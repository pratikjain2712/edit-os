"""The top of the file (header.py): sign-off lines, renumbering queries with old -> new tables, tables and figures
renumbered in order of first citation (rules_figures.py) and put in that order, the running head first."""

import zipfile
from datetime import datetime

from docx import Document

from edit_os import docmodel, frontmatter as fm, header, pipeline, qc, queries, rules, rules_figures, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

WHEN = datetime(2026, 10, 22, 11, 0)


def _style(code="ijo"):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    st.settings["front_matter"]["enabled"] = False
    return st


def _doc(tmp_path, paras, tables=(), name="t.docx"):
    """paras: text, or ("TABLE", n) for a small table after its caption."""
    doc = Document()
    for t in paras:
        if isinstance(t, tuple):
            tbl = doc.add_table(rows=2, cols=2)
            tbl.cell(0, 0).text, tbl.cell(0, 1).text = f"Item {t[1]}", "Value"
            tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Row", str(t[1])
        else:
            doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


# --- sign-off --------------------------------------------------------------------------------

def test_signoff_lines_follow_the_guide():
    st = _style()
    lines = header.lines(st, {}, ChangeSet(), (Status.ACCEPTED,), stages=("style", "copyedit"), editor="Vikram S.",
                         when=WHEN)
    assert lines == [("p", "Technical Editing done on 22-October-2026:- Vikram S."),
                     ("p", "Copyediting done on 22-10-26")]
    only_style = header.lines(st, {}, ChangeSet(), (Status.ACCEPTED,), stages=("style",), editor="", when=WHEN)
    assert only_style == [("p", "Technical Editing done on 22-October-2026:- ???")]  # no copy edit; no name yet


def test_placeholder_blocks_qc_and_the_name_clears_it(tmp_path):
    st = _style()
    for editor, flagged in (("", True), ("Vikram S.", False)):
        ms = docmodel.load(_doc(tmp_path, ["Introduction", "Text."]))
        header.render(ms, header.lines(st, {}, ChangeSet(), (Status.ACCEPTED,), stages=("style",), editor=editor,
                                       when=WHEN), tracked=True, author=editor or "Edit OS", date="2026-10-22T00:00:00Z")
        ms.save(tmp_path / "o.docx")
        out = docmodel.load(str(tmp_path / "o.docx"))
        assert out.paragraphs[0].text.startswith("Technical Editing done on")
        rep = qc.check(out, st)
        assert ("signoff.name" in {i.check for i in rep.issues}) is flagged
        xml = zipfile.ZipFile(tmp_path / "o.docx").read("word/document.xml").decode()
        assert "<w:pPr><w:rPr><w:ins " in xml  # a tracked inserted paragraph


# --- renumbering queries and tables at the top --------------------------------------------------

def test_renumbering_query_and_table_in_the_header(tmp_path):
    st = _style("mjdrdypu")
    info = {"renumbering": [{"kind": "Table", "mapping": [[1, 1], [2, 2], [3, 4], [4, 3]], "blocks": {}}],
            "after": "p0001"}
    ch = Change("x", "p0002", 0, 1, "4", "3", Layer.TABLES_FIGURES, Source.RULE, "fig_renumber",
                provenance={"kind": "renumber", "label": "Table 4 → 3", "figure_kind": "Table"})
    lines = header.lines(st, info, ChangeSet([ch]), (Status.SUGGESTED, Status.ACCEPTED), stages=("style",),
                         editor="Sophiya", when=WHEN)
    assert lines[1] == ("q", "Query: Table citations are not in sequential order. Therefore, the Tables and citations "
                             "have been renumbered. Please check and confirm.")
    assert lines[2] == ("table", [["Table", "Renumber"], ["1", "1"], ["2", "2"], ["4", "3"], ["3", "4"]])
    # Not rendered while the renumbering is rejected (or, in the clean copy, still pending).
    assert len(header.lines(st, info, ChangeSet([ch]), (Status.ACCEPTED,), stages=("style",), editor="S",
                            when=WHEN)) == 1

    # Written as a real Word table of tracked rows, below the portal's copyright notice query.
    path = _doc(tmp_path, ["AQ: The author(s) Dr X has/have not agreed to the copyright terms.", "Title", "Text."])
    ms = docmodel.load(path)
    header.render(ms, lines, tracked=True, author="Sophiya", date="2026-10-22T00:00:00Z", after="p0001")
    ms.save(tmp_path / "h.docx")
    doc = Document(str(tmp_path / "h.docx"))
    body = [el.tag.rpartition("}")[2] for el in doc.element.body][:6]
    assert body == ["p", "p", "p", "tbl", "p", "p"]
    paras = docmodel.load(str(tmp_path / "h.docx")).paragraphs
    texts = [p.text for p in paras if not p.in_table]
    assert texts[0].startswith("Technical Editing") and texts[1].startswith("AQ: The author(s)")
    assert texts[2].startswith("Query: Table citations") and texts[3:5] == ["", "Title"]
    cells = [p.text for p in paras if p.in_table]
    assert [cells[i:i + 2] for i in range(0, len(cells), 2)] == lines[2][1]
    xml = zipfile.ZipFile(tmp_path / "h.docx").read("word/document.xml").decode()
    assert xml.count("<w:trPr><w:ins ") == 5


# --- tables and figures renumbered by first citation ---------------------------------------------

TABLES_OUT_OF_ORDER = [
    "Introduction", "Results",
    "Age is given in (Table 1) and sex in Table 2. Outcomes are in (Table 4); see also Tables 3 and 4.",
    "Accuracy is summarised in Table 3.",
    "Table 1: Ages", ("TABLE", 1),
    "Table 2: Sex", ("TABLE", 2),
    "Table 3: Accuracy", ("TABLE", 3),
    "Table 4: Outcomes", ("TABLE", 4), "Abbreviations: CI, confidence interval",
    "References", "1. Author A. A study. J Test 2001;1:1-2.",
]


def test_table_renumbering_in_captions_and_citations(tmp_path):
    st = _style()
    path = _doc(tmp_path, TABLES_OUT_OF_ORDER)
    ms = docmodel.load(path)
    plan = rules_figures.renumber_plan(ms, st)
    assert set(plan) == {"Table"} and plan["Table"].mapping == {1: 1, 2: 2, 4: 3, 3: 4}
    cs = rules.run_all(ms, st)
    assert {c.reason for c in cs.changes if c.new in ("3", "4")} == {"fig_renumber"}
    qs = queries.collect(docmodel.load(path), st)
    assert "figures.order" not in {q.check for q in qs.queries}  # renumbered: the header carries the query

    out = pipeline.edit(path, st, str(tmp_path / "out"), stages=("style",), accept_all=True, editor="Vikram S.")
    clean = Document(out.path("clean"))
    texts = [p.text for p in clean.paragraphs if p.text]
    assert texts[0].startswith("Technical Editing done on") and texts[0].endswith(":- Vikram S.")
    assert texts[1] == ("AQ: Table citations are not in sequential order. Therefore, the Tables and citations have "
                        "been renumbered. Please check and confirm.")
    assert [[c.text for c in r.cells] for r in clean.tables[0].rows] == \
        [["Old", "New"], ["1", "1"], ["2", "2"], ["4", "3"], ["3", "4"]]
    body = next(t for t in texts if t.startswith("Age is given"))
    assert body == "Age is given in [Table 1] and sex in Table 2. Outcomes are in [Table 3]; see also Tables 3 and 4."
    assert next(t for t in texts if t.startswith("Accuracy is")) == "Accuracy is summarised in Table 4."
    # The tables are put in the new order, each with its caption and notes.
    caps = [t for t in texts if t.startswith("Table ") and ":" in t]
    assert caps == ["Table 1: Ages.", "Table 2: Sex.", "Table 3: Outcomes.", "Table 4: Accuracy."]
    assert [t.rows[1].cells[1].text for t in clean.tables[1:]] == ["1", "2", "4", "3"]
    k = texts.index("Table 3: Outcomes.")
    assert texts[k + 1] == "Abbreviations: CI, confidence interval"  # the note moves with its table

    # Tracked copy: the moved table is deleted where it was and inserted, edited, where it belongs.
    trk = docmodel.load(out.path("tracked"))
    texts = [p.text for p in trk.paragraphs if not p.in_table and p.text]
    caps = [t for t in texts if t.startswith("Table ") and ":" in t]
    assert caps == ["Table 1: Ages.", "Table 2: Sex.", "Table 3: Outcomes.", "Table 4: Accuracy."]  # deleted text hidden
    xml = zipfile.ZipFile(out.path("tracked")).read("word/document.xml").decode()
    assert "<w:delText>Outcomes</w:delText>" in xml or "Outcomes</w:delText>" in xml
    assert len(Document(out.path("tracked")).tables) == 6  # header table + four + the deleted original of the moved one


def test_no_renumbering_when_a_table_is_missing(tmp_path):
    st = _style()
    paras = [p for p in TABLES_OUT_OF_ORDER if p not in ("Table 3: Accuracy", ("TABLE", 3))]
    path = _doc(tmp_path, paras)
    assert rules_figures.renumber_plan(docmodel.load(path), st) == {}
    assert not [c for c in rules.run_all(docmodel.load(path), st).changes if c.reason == "fig_renumber"]
    checks = {q.check for q in queries.collect(docmodel.load(path), st).queries}
    assert "figures.order" in checks and "figures.no_table" in checks


def test_figure_renumbering_and_kjo_queries_instead(tmp_path):
    paras = ["Introduction", "Results", "The scan [Figure 2] and the photograph [Figures 1a and 1b] differ.",
             "Figure 1: Photograph (a) left (b) right", "Figure 2: Scan", "References", "1. Author A. J 2001;1:1."]
    path = _doc(tmp_path, paras)
    st = _style()
    assert rules_figures.renumber_plan(docmodel.load(path), st)["Figure"].mapping == {2: 1, 1: 2}
    cs = rules.run_all(docmodel.load(path), st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    tracked.render(docmodel.load(path), cs, tracked=False).save(tmp_path / "c.docx")
    texts = [p.text for p in docmodel.load(str(tmp_path / "c.docx")).paragraphs]
    assert "The scan [Fig. 1] and the photograph [Fig. 2a and 2b] differ." in texts  # IJO (S300 "Fig. 1")
    k = next(i for i, t in enumerate(texts) if t.startswith("Figure 2: Photograph"))
    assert texts[k + 1] == "Figure 1: Scan"  # renumbered in place (the legends move when the file is rendered)

    kjo = _style("kjo")  # the KJO sheet has no chronological-order line: the guide's query instead
    assert rules_figures.renumber_plan(docmodel.load(path), kjo) == {}
    assert "figures.order" in {q.check for q in queries.collect(docmodel.load(path), kjo).queries}


# --- running head first ----------------------------------------------------------------------------

def test_running_head_is_the_first_line(tmp_path, monkeypatch):
    """The title stands above an "Article type:" line in the manuscript: the running head and the type go above the
    title (the type line is deleted), as the editors place them."""
    st = style.load("mjdrdypu")
    st.settings["references"]["lookup"] = False
    paras = ["Effect of Valsalva Maneuver on Intraocular Pressure", "Article type: Original Article",
             "Asha Kumar1, Ravi Rao1", "1Department of Ophthalmology, City Eye Hospital, Chennai, Tamil Nadu, India",
             "Abstract", "Purpose: To test."]
    raw = {
        "article_type": "Original Article", "title": paras[0], "running_title": "Valsalva and IOP",
        "authors": [{"given_names": "Asha", "surname": "Kumar", "prefix": "", "affiliations": [1], "corresponding": True},
                    {"given_names": "Ravi", "surname": "Rao", "prefix": "", "affiliations": [1], "corresponding": False}],
        "affiliations": [{"id": 1, "department": "Department of Ophthalmology", "institution": "City Eye Hospital",
                          "city": "Chennai", "state": "Tamil Nadu", "country": "India"}],
        "correspondence": {}, "proper_nouns": ["Valsalva"],
        "paragraphs": [{"id": f"M{i}", "role": r} for i, r in
                       enumerate(["title", "article_type", "authors", "affiliation", "body", "body"], 1)],
    }
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = _doc(tmp_path, paras)
    ms = docmodel.load(path)
    ms.front_sources = fm.Sources(first_page=["Running title: Valsalva and IOP"])
    pl = fm.plan(ms, st)
    cs = rules.run_all(ms, st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    for n in pl.inserts:
        n.status = Status.ACCEPTED
    out = docmodel.load(path)
    tracked.render(out, cs, tracked=False)
    fm.render_inserts(out, pl.inserts, tracked=False)
    out.save(tmp_path / "clean.docx")
    got = [p.text for p in docmodel.load(str(tmp_path / "clean.docx")).paragraphs if p.text]
    assert got[0] == "<RH>Kumar and Rao: Valsalva and IOP"
    assert got[1] == "Original Article" and got[2].startswith("Effect of Valsalva maneuver")
    assert "Article type" not in " ".join(got)
