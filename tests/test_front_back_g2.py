"""Front and back matter (IJO, JEHP): legacy .doc title pages, affiliation linking and clean-up, statements given
only on the title page, contributions detection, sections stated twice, the acknowledgement's place, consent
headings, names with their initial after them. Owner decisions of 2026-10-27.

All text is invented."""

import io
import struct
import zipfile

from docx import Document

from edit_os import docmodel, frontmatter as fm, queries, rules, rules_sections, style, tracked
from edit_os.changes import Status

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item.")
REFS = ["References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4.",
        "2. Lee K, Rao S. Neck pain in adolescents. Spine 1998;23:1380-4."]


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def load(path, first_page=(), meta=None):
    ms = docmodel.load(path)
    ms.front_sources = fm.Sources(first_page=list(first_page), meta=meta or fm.Meta())
    return ms


def run(path, code="jehp", first_page=(), meta=None, only=("heading_tags", "sections_back_matter")):
    st = style.load(code)
    cs = rules.run_all(load(path, first_page, meta), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = path.replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs if p.text]


def section_queries(path, code="jehp", first_page=()):
    return [q.text for q in queries.collect(load(path, first_page), style.load(code)).queries
            if q.check.startswith("sections.")]


# --- 1. legacy .doc title page ---------------------------------------------------------------------------

def _word97(text: str) -> tuple[bytes, bytes]:
    """WordDocument and 1Table streams of a minimal Word 97 file holding `text` as one 8-bit piece."""
    body = text.encode("cp1252")
    at = 0x800
    wd = bytearray(at + len(body) + 16)
    struct.pack_into("<HHHHHH", wd, 0, 0xA5EC, 0xC1, 0, 0, 0, 0x0200)  # 1Table
    struct.pack_into("<i", wd, 0x4C, len(body))
    wd[at:at + len(body)] = body
    pcd = struct.pack("<HIH", 0, (at * 2) | 0x40000000, 0)
    plc = struct.pack("<ii", 0, len(body)) + pcd
    clx = b"\x02" + struct.pack("<I", len(plc)) + plc
    table = b"\x00" * 16 + clx
    struct.pack_into("<II", wd, 0x1A2, 16, len(clx))
    return bytes(wd), table


class _FakeOle:
    def __init__(self, streams):
        self.streams = streams

    def openstream(self, name):
        return io.BytesIO(self.streams[name])

    def exists(self, name):
        return name in self.streams

    def close(self):
        pass


def test_legacy_doc_title_page_read(monkeypatch):
    import olefile

    wd, table = _word97("Title page\rRunning title : Sleep in nurses\rCorresponding Author:\x0bDr. A. Rahimi\r"
                        "E-mail address : a.rahimi@example.org\x07\x13 HYPERLINK x \x14link\x15\r")
    monkeypatch.setattr(olefile, "OleFileIO", lambda f: _FakeOle({"WordDocument": wd, "1Table": table}))
    assert fm.doc_paragraphs(b"ole") == ["Title page", "Running title : Sleep in nurses",
                                         "Corresponding Author:\nDr. A. Rahimi", "E-mail address : a.rahimi@example.org",
                                         "link"]


def test_legacy_doc_title_page_taken_from_the_zip(tmp_path, monkeypatch):
    path = tmp_path / "original.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("FirstPage.doc", b"legacy")
        z.writestr("Manuscript.docx", b"x")
    monkeypatch.setattr(fm, "doc_paragraphs", lambda data: ["Title page", "Running title: Sleep in nurses"])
    src = fm.sources_from_zip(str(path))
    assert src.first_page_name == "FirstPage.doc" and src.first_page[1] == "Running title: Sleep in nurses"


# --- 2. affiliations ---------------------------------------------------------------------------------------

def test_postal_codes_and_designations_dropped():
    assert fm.drop_postal_codes("Sidi Ali University, Fez 30000") == "Sidi Ali University, Fez"
    assert fm.drop_postal_codes("75006 Paris") == "Paris"
    assert fm.drop_postal_codes("Building 2, Sector 12") == "Building 2, Sector 12"
    assert fm.drop_designations("Department of Nursing, Delta University, Zahedan, Iran, Associate Professor, "
                                "Community Health Center, Delta University") == \
        "Department of Nursing, Delta University, Zahedan, Iran, Community Health Center, Delta University"
    assert fm.drop_designations("Head and Neck Surgery") == "Head and Neck Surgery"


def _fields(authors, affs):
    f = fm.Fields(authors=[fm.Author(g, s, list(a)) for g, s, a in authors])
    f.affiliations = {i: dict(dict(department="", institution="", city="", state="", country=""), **a)
                      for i, a in affs.items()}
    return f


def test_unnumbered_affiliation_lines_follow_the_marks():
    # IJO_2733: the split numbered the affiliations from the cover letter's signature (swapped)
    lines = ["Authors:", "Dr. Ruhi Dev Jain DNB1", "Dr. Kiran Kale MS1", "Dr. Omar Shah MS2", "Department and affiliation",
             "Orbit and Oculoplasty Services", "Lotus Eye Foundation, Chennai, India",
             "Orbit and Oculoplasty Services, Lotus Eye Hospital (A Unit of Lotus Eye Foundation), Kolkata, India"]
    f = _fields([("Ruhi Dev", "Jain", [1]), ("Kiran", "Kale", [1]), ("Omar", "Shah", [2])],
                {1: dict(department="Orbit and Oculoplasty Services",
                         institution="Lotus Eye Hospital (A Unit of Lotus Eye Foundation)", city="Kolkata", country="India"),
                 2: dict(institution="Lotus Eye Foundation", city="Chennai", country="India")})
    fm._listed_affiliation_order(f, fm.Sources(first_page=lines), [])
    assert [a.affs for a in f.authors] == [[2], [2], [1]]
    assert f.affiliations[2]["department"] == "Orbit and Oculoplasty Services"  # the line above it


def test_per_author_list_linked_in_author_order():
    # IJO_505: one numbered affiliation per author, the marks after the names number the institutions
    eye = dict(department="Eye Centre", institution="City Medical College", city="New Delhi")
    path = dict(department="Department of Pathology", institution="City Medical College", city="New Delhi")
    f = _fields([("Ruchi", "Das", [1]), ("Meena", "Wadia", [1]), ("Ravi", "Saran", [2]), ("Kriti", "Kaur", [2])],
                {1: eye, 2: eye, 3: path, 4: path})
    fm._per_author_list(f)
    assert [a.affs for a in f.authors] == [[1], [2], [3], [4]]


def test_xml_field_with_two_affiliations_split():
    # jehp_1402: the portal's institute field runs two affiliations together
    affs = fm._xml_affiliations({"department": "Psychiatric", "institution": "Psychiatric University Center Riverton, "
                                 "Morocco.Health Sciences Research Laboratory, Faculty of Medicine, Hill University, "
                                 "Riverton, Morocco.", "city": "RIVERTON"})
    assert len(affs) == 2
    assert affs[0]["department"] == "Department of Psychiatric" and affs[0]["institution"] == \
        "Psychiatric University Center" and affs[0]["country"] == "Morocco"
    assert affs[1]["institution"].startswith("Health Sciences Research Laboratory")
    # the department field holding a whole affiliation; the institution field only repeats part of it
    one = fm._xml_affiliations({"department": "Department of Epidemiology, Faculty of medicine, Hill University,Riverton,"
                                              "Morocco.", "institution": "Faculty of medecine Riverton"})
    assert one == [{"department": "Department of Epidemiology", "institution": "Faculty of medicine, Hill University",
                    "city": "Riverton", "country": "Morocco"}]


def test_spelling_variants_grouped_and_department_of_for_a_bare_specialty():
    f = _fields([("Ela", "Varma", [1]), ("Nita", "Jha", [2]), ("Asha", "Rao", [2]), ("Vinod", "Kumar", [3])],
                {1: dict(department="Department of Obstetrics and Gynaecology", institution="Coastal Institute",
                         city="Puducherry", country="India"),
                 2: dict(department="Department of Obstetrics and Gynecology", institution="Coastal Institute",
                         city="Puducherry", country="India"),
                 3: dict(department="Anatomy", institution="Coastal Institute", city="Puducherry", country="India")})
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    ms = docmodel.Manuscript(path="", document=None, paragraphs=[])
    blocks, _ = fm.build(ms, f, fm.Meta(), st)
    text = {b.role: b.text for b in blocks}
    assert text["authors"] == "Ela Varma1, Nita Jha1, Asha Rao1, Vinod Kumar2"
    assert text["affiliations"] == ("1Department of Obstetrics and Gynecology, Coastal Institute, Puducherry, India, "
                                    "2Department of Anatomy, Coastal Institute, Puducherry, India")
    assert any("spelling" in w for w in f.warnings)


def test_country_spelling_and_skipped_district():
    assert fm._same_country("Turkey", "Türkiye") and not fm._same_country("Turkey", "Greece")
    corr = {"address_parts": ["Department of Ophthalmology", "Mira City Hospital", "Lale Mahallesi, 12. Cadde No: 4"],
            "city": "Ankara", "postal_code": "06800", "country": "Türkiye"}
    src = ("Postal Address: Department of Ophthalmology, Mira City Hospital, Lale Mahallesi, 12. Cadde No: 4, Çankaya, "
           "06800 Ankara, Türkiye.")
    assert fm._skipped_address_pieces(src, corr) == ["Çankaya"]


def _plan(tmp_path, monkeypatch, raw, meta, code="jehp", lead=("Abstract", "Background: A study.")):
    path = make(tmp_path, list(lead), "fm.docx")
    monkeypatch.setattr(fm, "split", lambda text: raw)
    ms = docmodel.load(path)
    ms.front_sources = fm.Sources(first_page=raw.pop("_lines"), meta=meta)
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return fm.plan(ms, st)


def _raw(lines, authors, affs):
    return {"_lines": lines, "article_type": "", "title": "Sleep in nurses", "running_title": "Sleep",
            "authors": [dict(given_names=g, surname=s, prefix="", affiliations=a, corresponding=False)
                        for g, s, a in authors],
            "affiliations": [dict(dict(id=i, department="", institution="", city="", state="", country=""), **x)
                             for i, x in affs.items()],
            "correspondence": dict(prefix="", given_names="", surname="", address_parts=[], city="", state="",
                                   postal_code="", country="", email=""),
            "proper_nouns": [], "paragraphs": []}


def test_xml_only_author_queried_not_added(tmp_path, monkeypatch):
    lines = ["Sleep in nurses", "Asha Rao1, Ravi Kumar Das2", "1City Hospital, Pune, India", "2Lake College, Pune, India"]
    raw = _raw(lines, [("Asha", "Rao", [1]), ("Ravi Kumar", "Das", [2])],
               {1: dict(institution="City Hospital", city="Pune", country="India"),
                2: dict(institution="Lake College", city="Pune", country="India")})
    meta = fm.Meta(authors=[dict(given="asha", surname="rao"), dict(given="ravi", surname="das"),
                            dict(given="mira", surname="sen", prefix="Prof")])
    pl = _plan(tmp_path, monkeypatch, raw, meta)
    text = {b.role: b.text for b in pl.blocks}
    assert text["authors"] == "Asha Rao1, Ravi K. Das2"  # the title page's list, middle name as an initial (G50)
    q = [q for q in pl.queries if q.check == "front.xml_only_author"]
    assert len(q) == 1 and "Prof mira sen" in q[0].text


def test_name_with_its_initial_after_it_kept_as_written(tmp_path, monkeypatch):
    # jehp_1683: "Natarajan S" is not expanded from the XML / copyright notice; the guide puts the initial first
    lines = ["Sleep in nurses", "Dr. Kumaran S, Priya Ragu", "City Hospital, Chennai, India"]
    raw = _raw(lines, [("Kumaran", "S", [1]), ("Priya", "Ragu", [1])],
               {1: dict(institution="City Hospital", city="Chennai", country="India")})
    meta = fm.Meta(authors=[dict(given="kumaran", surname="selvaraj"), dict(given="Priya", surname="Ragu")])
    text = {b.role: b.text for b in _plan(tmp_path, monkeypatch, raw, meta).blocks}
    assert text["running_head"].startswith("<RH>Kumaran and Ragu:")
    assert text["authors"] == "S. Kumaran, Priya Ragu"
    assert "Kumaran S, Ragu P." in text["cite"]


# --- 3. back matter from the title page ---------------------------------------------------------------------

def test_title_page_statements_used_never_the_default_nil(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Acknowledgement",
                           "We thank the participating nurses."] + REFS)
    fp = ["Title page", "Ethical Approval", "The study was approved by the ethics committee of Lake University (ID 42).",
          "Funding", "Funding was provided by the research deputy of Lake University.", "Conflicts of interest",
          "The authors declare that they have no competing interests.", "Use of AI tools",
          "ChatGPT was used only for language editing.", "Authors' contributions",
          "A.R. designed the study and wrote the draft. B.K. analysed the data and reviewed the manuscript.",
          "Covering letter", "Dear Editor,"]
    t = run(path, first_page=fp)
    i = t.index("<H1>Acknowledgement")
    assert t[i + 2:i + 12] == [
        "<H2>Author contributions",
        "A.R. designed the study and wrote the draft. B.K. analysed the data and reviewed the manuscript.",
        "<H2>Ethical approval", "The study was approved by the ethics committee of Lake University (ID 42).",
        "<H2>Financial support and sponsorship", "Funding was provided by the research deputy of Lake University.",
        "<H2>Conflicts of interest", "The authors declare that they have no competing interests.",
        "<H2>Use of AI tools", "ChatGPT was used only for language editing."]
    assert "Nil." not in t


def test_funding_sentence_inside_the_conflicts_statement(tmp_path):
    # jehp_1209
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Conflict of interest",
                           "The authors declare no conflict of interest. The study received no commercial funding."]
                + REFS)
    t = run(path)
    i = t.index("<H2>Financial support and sponsorship")
    assert t[i:i + 4] == ["<H2>Financial support and sponsorship", "The study received no commercial funding.",
                          "<H2>Conflict of interest", "The authors declare no conflict of interest."]


# --- 4. contributions ----------------------------------------------------------------------------------------

def test_contribution_statements_from_the_title_page():
    # "Contributorship statement" (IJO_1432), initials in brackets (1745), "(CRediT Statement)" (2904), names in
    # lower case (505), consent lines after the list (1347)
    for fp, want in [
        (["Contributorship statement: Patient care (AK, BR); literature search (CS)."],
         ["Patient care (AK, BR); literature search (CS)."]),
        (["Author’s contributions Conceptualization: [AK,BR]; methodology: [AK]; writing: [BR]."],
         ["Conceptualization: [AK,BR]; methodology: [AK]; writing: [BR]."]),
        (["Author Contributions (CRediT Statement)", "Conceptualization: Asha Rao", "Supervision: Ravi Das",
          "Funding: None"], ["Conceptualization: Asha Rao", "Supervision: Ravi Das"]),
        (["Contribution", "Asha Rao- review of literature", "ravi kumar das – histopathological diagnosis",
          "Declaration"], ["Asha Rao- review of literature", "ravi kumar das – histopathological diagnosis"]),
        (["Author Contributions", "Concept: Asha Rao developed the idea.", "Design: Ravi Das planned the report.",
          "Written informed consent was obtained from the patient.", "Patient Consent for Publication:"],
         ["Concept: Asha Rao developed the idea.", "Design: Ravi Das planned the report."]),
    ]:
        assert rules_sections._contribution_items(fp) == want, fp


def test_unlabelled_contributions_only_for_author_names(tmp_path):
    # jehp_1683: a "Future Research: 1] ..." list in the Discussion is not the authors' contributions
    meta = fm.Meta(authors=[dict(given="Asha", surname="Rao"), dict(given="Ravi", surname="Das")])
    path = make(tmp_path, ["Introduction", BODY, "Discussion", BODY,
                           "Future Research: 1] Multi-centric study design with data collection in five centres."]
                + REFS)
    assert "<H2>Author contributions" not in run(path, meta=meta)
    path = make(tmp_path, ["Introduction", BODY, "Discussion", BODY,
                           "Asha Rao: study design, data collection and drafting of the manuscript."] + REFS, "b.docx")
    assert "<H2>Author contributions" in run(path, meta=meta)


def test_ai_disclosure_is_not_contributions_or_conflicts(tmp_path):
    # jehp_1402
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Conflicts of interest",
                           "There are no conflicts of interest.",
                           "Disclosure: ChatGPT was used during revision for language editing and a draft figure."]
                + REFS)
    t = run(path)
    assert "<H2>Disclosure" in t and "There are no conflicts of interest." in t
    assert t.count("<H2>Conflicts of interest") == 1 and "<H2>Author contributions" not in t


# --- 5. a section stated twice --------------------------------------------------------------------------------

def test_funding_stated_twice_keeps_one_heading_and_queries(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Financial Support: Nil.",
                           "Funding: Nil. No financial support was received for this study."] + REFS)
    t = run(path)
    assert t.count("<H2>Financial support and sponsorship") == 1 and "<H2>Funding" in t
    assert any("given twice" in q for q in section_queries(path))


def test_nil_statement_beside_a_real_one_goes(tmp_path):
    # jehp_924: "Financial relationship and sponsership: None" and "Source of funding- ..."
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Acknowledge- None",
                           "Financial relationship and sponsership: None",
                           "Source of funding- This research was supported by an intramural grant of the institute."]
                + REFS)
    t = run(path)
    assert "Acknowledge- None" not in t and not any("sponsership" in x for x in t)
    i = t.index("<H2>Financial support and sponsorship")
    assert t[i + 1] == "This research was supported by an intramural grant of the institute."


# --- 6. the acknowledgement first (G148) -----------------------------------------------------------------------

def test_acknowledgement_after_the_statements_moves_before_them(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Declarations",
                           "Conflict of interest: The authors declared no conflict of interest",
                           "Funding: Self-funding", "Acknowledgements: We thank all the patients who took part."]
                + REFS)
    t = run(path)
    i = t.index("<H1>Acknowledgement")
    assert t[i - 1] == BODY and t[i + 1] == "We thank all the patients who took part."
    assert t.index("<H2>Conflicts of interest") > i and t.count("We thank all the patients who took part.") == 1


# --- 7. labels typed inline -------------------------------------------------------------------------------------

def test_label_typed_before_its_statement_with_spaces(tmp_path):
    # IJO_1745
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY,
                           "Ethical approval    All procedures were in accordance with the ethical standards."] + REFS)
    t = run(path, "ijo")
    i = t.index("<H2>Ethical approval")
    assert t[i + 1] == "All procedures were in accordance with the ethical standards."
    assert t[i + 2] == "<H2>Financial support and sponsorship"  # the added sections after it


def test_unlabelled_nil_statements_under_declarations_go(tmp_path):
    # IJO_2904
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Declarations",
                           "There were no specific funds allocated for this work.",
                           "No new datasets were generated or analysed, so there are no data available."] + REFS)
    t = run(path, "ijo")
    assert not any(x.startswith(("There were no", "No new datasets", "Declarations")) for x in t)
    assert t[t.index("<H2>Financial support and sponsorship") - 1] == BODY


def test_ethical_approve_typo_heading():
    assert rules_sections._sentence_label("Ethical approve") == "Ethical approval"


# --- 8. consent heading ---------------------------------------------------------------------------------------

def test_authors_own_consent_heading_takes_the_sheets(tmp_path):
    statement = "The patient gave written informed consent for the publication of the clinical images."
    for head in ("Author’s Declaration statement-", "Authors’ Declaration"):
        path = make(tmp_path, ["Case Report", BODY, "Discussion", BODY, head, statement] + REFS)
        for code in ("ijo", "jehp"):
            t = run(path, code)
            assert t[t.index(statement) - 1] == "<H2>Declaration of patient consent", (head, code)
