"""Books, links, conference papers and reports (refs_other). All lookups are faked; nothing goes online."""

import json

import pytest
from docx import Document

from edit_os import biblio, docmodel, queries, refengine, refs_other, rules, style, tracked
from edit_os.changes import Status


@pytest.fixture
def st():
    return style.load("ijo")


@pytest.fixture
def tmp_cache(monkeypatch, tmp_path):
    """Link checks are cached in biblio's SQLite file; keep test entries out of the real one."""
    monkeypatch.setattr(biblio, "_http", biblio._Http(str(tmp_path / "cache.sqlite")))


@pytest.fixture
def links(monkeypatch, tmp_cache):
    pages = {}
    monkeypatch.setattr(refs_other, "_fetch_link", lambda url: pages.get(url))
    return pages


@pytest.fixture
def lookups(monkeypatch, tmp_cache):
    """Fake HTTP bodies by URL substring (CrossRef, Open Library) and PubMed calls."""
    bodies = {}

    def get(self, url, host_gap):
        return next((b for k, b in bodies.items() if k in url), None)

    monkeypatch.setattr(biblio._Http, "get", get)
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: [])
    return bodies


# --- dates ---------------------------------------------------------------------------------

@pytest.mark.parametrize("raw, out", [
    ("2025 Oct 4", "2025 Oct 04"), ("2025 Sept 22", "2025 Sep 22"), ("25th March 2026", "2026 Mar 25"),
    ("Apr. 20, 2025", "2025 Apr 20"), ("1 May 2018", "2018 May 01"), ("2026-03-25", "2026 Mar 25"),
    ("25/03/2026", "2026 Mar 25"), ("2021Dec10", "2021 Dec 10"),
    ("03/04/2025", None), ("March 2026", None), ("yesterday", None),  # ambiguous or partial: never guessed
])
def test_access_date(raw, out):
    assert refs_other.access_date(raw) == out


# --- template fixes ------------------------------------------------------------------------

def test_url_template_moves_cited_date():
    f = refs_other.template("Google Scholar [Internet]. 2025 [cited 2025 May 20]. Available from: https://scholar.example.org",
                            "website")
    assert f.text == "Google Scholar. 2025. Available from: https://scholar.example.org. [Last accessed on 2025 May 20]."
    assert f.accessed == "2025 May 20"


def test_url_template_accessed_phrases():
    f = refs_other.template("Massey Unit. Keratoprosthesis. Available at https://eye.example.org/kpro Accessed on 25th March 2026.",
                            "website")
    assert f.text == "Massey Unit. Keratoprosthesis. Available from: https://eye.example.org/kpro. [Last accessed on 2026 Mar 25]."
    f = refs_other.template("Data Office. Population. Available from: https://data.example.org/566 Accessed May 13, 2025.", "")
    assert f.text.endswith("Available from: https://data.example.org/566. [Last accessed on 2025 May 13].")


def test_url_template_never_invents_access_date():
    f = refs_other.template("https://www.example.com/products/lens-a", "website")
    assert f.text == "Available from: https://www.example.com/products/lens-a." and f.accessed is None
    f = refs_other.template("Foo Society. Bar page. Available at https://x.example.org/a (accessed 03/04/2025).", "website")
    assert "Last accessed" not in f.text and "accessed 03/04/2025" in f.text and f.access_raw


def test_doi_link_is_not_a_web_reference():
    f = refs_other.template("Vera K. Coca leaf. Revista Psicologia 2019; doi:https://example.org/x", "website")
    assert "Available from" not in f.text


def test_book_template():
    f = refs_other.template("Gorlin RJ, Cohen MM, Hennekam RC. Syndromes of the head and neck. 4th edition. "
                            "New York : Oxford University Press, 2001, pp 100-109", "book")
    assert f.text == ("Gorlin RJ, Cohen MM, Hennekam RC. Syndromes of the Head and Neck. 4th ed. "
                      "New York: Oxford University Press; 2001. p. 100-9.")


def test_book_edition_words_and_title_inside():
    assert refs_other.fix_edition("Title. Second edition. London: X; 2001") == "Title. 2nd ed. London: X; 2001"
    keep = "Update from the 5th Edition of the WHO classification. Head Neck Pathol 2022"
    assert refs_other.fix_edition(keep) == keep  # an edition named inside a title is not an edition statement


def test_chapter_template():
    f = refs_other.template("Aster J, Kumar V. White cells and lymph nodes. In: Cotran RS, Kumar V, Collins T (eds). "
                            "Robbins pathologic basis of disease. 6th Edition. Oxford: W B Saunders, 1999, p. 657-658.",
                            "book_chapter")
    assert f.text == ("Aster J, Kumar V. White cells and lymph nodes. In: Cotran RS, Kumar V, Collins T, editors. "
                      "Robbins Pathologic Basis of Disease. 6th ed. Oxford: W B Saunders; 1999. p. 657-8.")


def test_classify_heuristics():
    c = refs_other.classify
    assert c("Rahman S. Breast cancer detection [thesis]. Dhaka: Brac University; 2022.") == "thesis"
    assert c("Pena IG. Tuberculosis. A paper presented in the 8th Asian Conference on Clinical Pharmacy, 2008.") == "conference"
    assert c("Morin CM. Insomnia. New York: Guilford Press; 1993.") == "book"
    assert c("Lee K. Back pain. Spine 1997;22:1380-4.") == "other"  # journal citations are not touched here
    assert c("Anything", "website") == "website"  # Claude's kind wins


# --- links -----------------------------------------------------------------------------------

def lc(**kw):
    return refs_other.LinkCheck(kw.pop("url", "https://www.example.org/guide/page-1"), **kw)


def test_link_verdicts():
    assert lc(status=200, final_url="https://example.org/guide/page-1/", title="Guide").verdict == "ok"
    assert lc(status=404).verdict == "dead"
    assert lc(status=200, final_url="https://www.example.org/", title="Home").verdict == "redirected"  # soft 404
    assert lc(status=200, title="Page Not Found | Example").verdict == "dead"
    assert lc(status=200, title="Just a moment...").verdict == "unknown"  # bot check: cannot tell
    assert lc(status=403).verdict == "unknown"
    assert lc(error="Tunnel connection failed: 502 Bad Gateway").verdict == "dead"
    assert lc(error="[SSL: CERTIFICATE_VERIFY_FAILED]").verdict == "unknown"
    assert lc(status=200, final_url="http://example.org/guide/page-1").verdict == "ok"  # www / scheme only


def test_dead_link_and_missing_date_are_queried(st, links):
    url = "https://www.example.com/products/lens-a"
    links[url] = {"url": url, "status": 404, "final_url": url}
    res = refengine.resolve(f"Lens maker. Lens A. {url}", st, {"kind": "website", "title": "Lens A", "authors": []})
    assert res.status == refs_other.OTHER_FIXED
    assert res.formatted == f"Lens maker. Lens A. Available from: {url}."
    assert [q for q, _ in res.queries] == ["references.link_dead", "references.access_date"]


def test_live_link_with_matching_title_is_verified(st, links):
    url = "https://www.example.org/fact-sheets/family-planning"
    links[url] = {"url": url, "status": 200, "final_url": url, "title": "Family planning/contraception methods",
                  "content_type": "text/html"}
    text = (f"World Health Organization. Family planning/contraception methods. 2025. Available from: {url}. "
            "[Last accessed on 2025 Sep 22].")
    res = refengine.resolve(text, st, {"kind": "website", "title": "Family planning/contraception methods"})
    assert res.status == refs_other.OTHER_VERIFIED and not res.formatted and not res.queries


def test_redirect_elsewhere_is_queried(st, links):
    url = "https://old.example.org/professionals/keratoprosthesis"
    links[url] = {"url": url, "status": 200, "final_url": "https://new.example.com/en/professionals",
                  "title": "Professionals | New Example", "content_type": "text/html"}
    res = refengine.resolve(f"Eye Unit. Keratoprosthesis service. Available from: {url}. [Last accessed on 2026 Mar 25].",
                            st, {"kind": "website", "title": "Keratoprosthesis service"})
    assert [q for q, _ in res.queries] == ["references.link_redirect"]
    assert res.queries[0][1]["url"] == "https://new.example.com/en/professionals"


def test_link_checks_are_cached(links):
    url = "https://www.example.org/a"
    links[url] = {"url": url, "status": 404, "final_url": url}
    assert refs_other.check_link(url).verdict == "dead"
    links.clear()
    assert refs_other.check_link(url).verdict == "dead"  # from biblio's cache


# --- books -----------------------------------------------------------------------------------

BOOK = "Ringsven MK, Bond D. Gerontology and Leadership Skills for Nurses. 2nd ed. Albany, NY: Delmar Publishers; 1996."


def test_book_verified_in_open_library(st, lookups):
    lookups["openlibrary.org"] = json.dumps({"docs": [
        {"title": "Gerontology and leadership skills for nurses", "author_name": ["Mary K. Ringsven", "Donna Bond"],
         "publish_year": [1991, 1996], "publisher": ["Delmar Publishers"], "publish_place": ["Albany, N.Y"]}]})
    res = refengine.resolve(BOOK, st)
    assert res.status == refs_other.OTHER_VERIFIED and not res.queries
    assert "openlibrary" in res.message


def test_book_verified_in_crossref(st, lookups):
    lookups["api.crossref.org"] = json.dumps({"message": {"items": [
        {"type": "book", "title": ["Gerontology and Leadership Skills for Nurses"],
         "author": [{"family": "Ringsven", "given": "M K"}], "issued": {"date-parts": [[1996]]}, "publisher": "Delmar"}]}})
    res = refengine.resolve(BOOK, st)
    assert res.status == refs_other.OTHER_VERIFIED and "crossref" in res.message


def test_wrong_book_is_not_accepted(st, lookups):
    lookups["openlibrary.org"] = json.dumps({"docs": [
        {"title": "Leadership for nurses", "author_name": ["Someone Else"], "publish_year": [1996]}]})
    res = refengine.resolve(BOOK, st)
    assert res.status == refengine.NOT_JOURNAL
    assert [q for q, _ in res.queries] == ["references.book_unverified"]


def test_book_without_city_is_queried(st, lookups):
    res = refengine.resolve("Thakur A. Charcoal in Dentistry. Scrivener Publishing; 2020.", st,
                            {"kind": "book", "title": "Charcoal in Dentistry", "authors": ["Thakur A"]})
    assert "references.publisher_location" in [q for q, _ in res.queries]


def test_statpearls_chapter_verified_in_pubmed_bookshelf(st, monkeypatch, tmp_cache):
    xml = ('<PubmedArticleSet><PubmedBookArticle><BookDocument><PMID>1</PMID><ArticleIdList>'
           '<ArticleId IdType="bookaccession">NBK000001</ArticleId></ArticleIdList><Book><Publisher><PublisherName>'
           'StatPearls Publishing</PublisherName></Publisher><BookTitle>StatPearls</BookTitle><PubDate><Year>2025</Year>'
           '</PubDate></Book><ArticleTitle>Ocular Cellulitis</ArticleTitle><AuthorList><Author><LastName>Mejia</LastName>'
           '</Author></AuthorList></BookDocument></PubmedBookArticle></PubmedArticleSet>')
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: ["1"] if "NBK000001" in term else [])
    monkeypatch.setattr(biblio, "_ncbi", lambda endpoint, **kw: xml)
    text = ("Mejia E, Vohra V. Ocular Cellulitis [Internet]. In: StatPearls. Treasure Island (FL): StatPearls Publishing; "
            "2025 [cited 2025 Oct 4]. Available from: http://www.ncbi.nlm.nih.gov/books/NBK000001/")
    res = refengine.resolve(text, st, {"kind": "book_chapter", "title": "Ocular Cellulitis", "authors": ["Mejia E", "Vohra V"]})
    assert res.status == refs_other.OTHER_VERIFIED and not res.queries
    assert res.formatted == ("Mejia E, Vohra V. Ocular Cellulitis. In: StatPearls. Treasure Island (FL): StatPearls "
                             "Publishing; 2025. Available from: http://www.ncbi.nlm.nih.gov/books/NBK000001/. "
                             "[Last accessed on 2025 Oct 04].")


def test_conference_gets_unverified_query(st):
    res = refengine.resolve("Pena IG. Tuberculosis treatment in Caloocan City. A paper presented in the 8th Asian "
                            "Conference on Clinical Pharmacy, Surabaya, Indonesia, July 01 - 04, 2008.", st)
    assert [q for q, _ in res.queries] == ["references.unverified"]


def test_not_a_reference_is_left_alone(st):
    res = refengine.resolve("appear in a journal article, or to be used for thesis or presentation.", st)
    assert res.status == refengine.NOT_JOURNAL and not res.queries and not res.formatted


# --- pipeline --------------------------------------------------------------------------------

def test_rules_and_queries_end_to_end(tmp_path, monkeypatch, tmp_cache):
    url = "https://www.example.org/report/vision"
    monkeypatch.setattr(refs_other, "_fetch_link", lambda u: {"url": u, "status": 404, "final_url": u})
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph(f"1. World Health Organization. World report on vision [Internet]. Geneva: WHO; 2019 "
                      f"[cited 2026 Oct 4]. Available at: {url}")
    doc.add_paragraph("2. Smith A, Jones B, et al. Ophthalmic Surgery. 2nd edition. London: Mosby; 2010.")
    path = tmp_path / "m.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    st = style.load("ijo")
    cs = rules.run_all(ms, st)
    assert any(c.reason == "ref_other" for c in cs.changes)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "u.docx"
    tracked.render(docmodel.load(str(path)), cs, tracked=False).save(out)
    texts = [p.text for p in docmodel.load(str(out)).paragraphs]
    assert texts[1] == (f"1. World Health Organization. World Report on Vision. Geneva: WHO; 2019. Available from: {url}. "
                        "[Last accessed on 2026 Oct 04].")
    assert texts[2] == "2. Smith A, Jones B, et al. Ophthalmic Surgery. 2nd ed. London: Mosby; 2010."
    qs = [q.text for q in queries.collect(ms, st).queries]
    assert queries.TEXT["references.link_dead"] in qs
    assert queries.TEXT["references.book_et_al"] in qs  # et al. with fewer than 6 authors in a book
