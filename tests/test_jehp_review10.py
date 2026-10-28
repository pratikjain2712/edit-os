"""JEHP review fixes, round 10 (synthetic text): a heading bold only through the "Strong" character style gets its
own bold, transliterated author names are not queried, a run-in label repeating the heading above is left alone,
table-footnote queries sit on the footnote (and travel with moved text), and the reference author-name query says
what differs."""

from docx import Document

from edit_os import biblio, docmodel, queries, refengine, rules

from test_jehp_citations import _refs, jehp

BODY = ("The participants completed the measures at baseline and after the programme, and the scores were compared "
        "between the two groups with the appropriate statistical tests for each outcome measure in the study.")


def _save(doc, tmp_path, name="m.docx"):
    path = str(tmp_path / name)
    doc.save(path)
    return path


# --- 1. a tagged heading bold only through a character style --------------------------------------------------------

def test_heading_bold_only_through_a_character_style_gets_its_own_bold(tmp_path):
    doc = Document()
    for t in ["Introduction", BODY, "Methods", BODY, "Results", BODY, "Discussion", BODY]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run("Limitations", style="Strong")
    for t in [BODY, "Conclusion", BODY] + _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(t)
    path = _save(doc, tmp_path)
    cs = rules.run_all(docmodel.load(path), jehp())
    lim = next(p for p in docmodel.load(path).paragraphs if p.text == "Limitations")
    assert any(c.para_id == lim.id and c.new.startswith("<H") for c in cs.changes)
    bold = [c for c in cs.changes if c.para_id == lim.id and c.reason == "fmt_heading_bold"]
    assert bold and bold[0].format == {"bold": True} and (bold[0].start, bold[0].end) == (0, len("Limitations"))


# --- 2. transliterated author names ----------------------------------------------------------------------------------

def test_transliterated_author_name_is_not_queried(tmp_path):
    doc = Document()
    for t in ["Introduction", "As Ra’ofat et al. [1] showed, and Moore et al. [2] found, the scores rose.",
              "Methods", BODY] + _refs("Raoffat F, Lane B. One study. J T 2020;1:1-2.",
                                       "Bae JS, Kim C. Two study. J T 2020;2:1-2."):
        doc.add_paragraph(t)
    path = _save(doc, tmp_path)
    found = rules.author_name_mismatches(docmodel.load(path), jehp())
    assert [(name, n) for _, _, _, name, n, _ in found] == [("Moore", 2)]


def test_same_name_folds_transliterations():
    assert rules._same_name("Ba'ezat", "Baezzat")
    assert rules._same_name("Müller-Hahn", "Muller Hahn")
    assert not rules._same_name("Wang", "Wong") and not rules._same_name("Moore", "Bae")


# --- 3. a run-in label repeating the heading above ---------------------------------------------------------------------

def test_run_in_label_repeating_its_heading_is_not_queried(tmp_path):
    doc = Document()
    for t in ["Introduction", BODY, "Methods"]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run("Coping Habits Scale (CHS)").bold = True
    doc.add_paragraph("Coping Habits Scale (CHS): " + BODY)
    doc.add_paragraph().add_run("Ethical Considerations").bold = True
    doc.add_paragraph("Ethical considerations: " + BODY)
    for t in ["Results", "The CHS score rose. " + BODY, "Discussion", BODY] + _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(t)
    path = _save(doc, tmp_path)
    qs = queries.collect(docmodel.load(path), jehp())
    checks = [q.check for q in qs.queries]
    assert "headings.duplicate" not in checks
    assert "abbrev.defined_twice" not in checks


# --- 4. table-footnote queries on the footnote; comments travel with moved text --------------------------------------

def _table(doc, caption, foot=None):
    doc.add_paragraph(caption)
    rows = [["Group", "Mean", "SD", "95% CI"], ["A", "1.0", "0.2", "1, 2"], ["B", "1.1", "0.3", "1, 3"]]
    t = doc.add_table(rows=3, cols=4)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.rows[r].cells[c].paragraphs[0].add_run(val)
    if foot:
        doc.add_paragraph(foot)
    doc.add_paragraph("")


def test_footnote_abbreviation_query_sits_on_the_footnote_or_the_caption(tmp_path):
    doc = Document()
    for t in ["Introduction", "Results are in Tables 1 and 2 [Tables 1 and 2]."]:
        doc.add_paragraph(t)
    _table(doc, "Table 1: Scores at baseline", foot="*Chi-square test, CI = confidence interval")
    _table(doc, "Table 2: Scores at follow-up")
    for t in _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(t)
    path = _save(doc, tmp_path)
    ms = docmodel.load(path)
    st = jehp()
    st.settings["figures"]["table_footnote_abbreviation_queries"] = True  # off for JEHP since 2026-10-28 (owner)
    qs = queries.collect(ms, st)
    on = [(ms.by_id(q.para_id).text, q.text) for q in qs.queries if q.check.startswith("figures.footnote_abbreviations")]
    assert not [q for q in queries.collect(ms, jehp()).queries if q.check.startswith("figures.footnote_abbreviations")]
    # SD is missing under both tables: asked once, at the first table's footnote, naming both
    assert on == [("*Chi-square test, CI = confidence interval",
                   "Kindly provide the expansion of SD in the footnotes of Tables 1 and 2."),
                  ("Table 2: Scores at follow-up", "Kindly provide the expansion of CI in the table footnote.")]


def test_moved_copy_takes_the_comments_with_it(tmp_path):
    from edit_os import header
    from edit_os.queries import QuerySet, write_comments
    from edit_os.docmodel import qn

    doc = Document()
    doc.add_paragraph("Table 1: A caption")
    path = _save(doc, tmp_path)
    ms = docmodel.load(path)
    qs = QuerySet()
    qs.add(ms.paragraphs[0].id, "text.brackets", 0, 5)
    write_comments(ms, qs)
    el = ms.paragraphs[0].element
    copy_el = header.moved_copy(el)
    tags = ("w:commentRangeStart", "w:commentRangeEnd", "w:commentReference")
    assert all(len(list(copy_el.iter(qn(t)))) == 1 for t in tags)
    assert all(not list(el.iter(qn(t))) for t in tags)


# --- 5. the reference author-name query says what differs --------------------------------------------------------------

def _rec(authors):
    return biblio.Record(source="crossref", title="A title", authors=list(authors), journal="J T", year="2020",
                         volume="1", pages="1-2")


def _parsed(names):
    return refengine.Parsed(text="", authors=names)


AFF = ("Department of Nursing, School of Health, Example University, Town", "")


def test_author_name_query_names_the_difference():
    q = refengine.author_name_query(_parsed(["Farhadi KM", "Lane B"]),
                                    _rec([AFF, ("Farhadi Kashani", "M"), AFF, ("Lane", "B")]))
    assert q == ("references.author_names", {"pairs": "Farhadi KM vs Farhadi Kashani M"})
    assert queries.TEXT[q[0]].format(**q[1]) == ("The author names differ from the published article (Farhadi KM vs "
                                                 "Farhadi Kashani M). Kindly check.")
    q = refengine.author_name_query(_parsed(["Rostami Z", "Lane B", "Kian H"]),
                                    _rec([("Rostami", "Z"), AFF, ("Kiyan", "H"), ("Lane", "B")]))
    assert q == ("references.author_order", {"names": "Rostami Z, Kiyan H, Lane B"})


def test_author_name_query_none_when_the_names_agree():
    # capitals, a record's transposed letters, short surnames read as initials, affiliations in the list
    assert refengine.author_name_query(_parsed(["MORAN Y", "HE N", "YUE Y"]),
                                       _rec([("MORNA", "Y"), AFF, ("HE", "N"), ("YUE", "Y")])) is None
