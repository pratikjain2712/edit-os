"""Structure (synthetic text only): reference-list heading variants, legends and table captions over several
paragraphs, legends one line per part, abbreviations in brackets, template leftovers, tables / figures / supplementary
material sent as other Word files of the submission."""

import io
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches

from edit_os import docmodel, leftovers, legends, pipeline, references, rules, rules_figures, style, supplied, tracked
from edit_os.changes import Status

from test_figures_insert import picture
from test_legends import REFS, W, make

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail.")


def journal(code, **figures):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    st.settings["figures"].update(figures)
    return st


def texts(path):
    return [p.text for p in docmodel.load(str(path)).paragraphs if p.text.strip()]


def apply(path, st, tmp_path, only=None):
    cs = rules.run_all(docmodel.load(path), st, only=list(only) if only else None)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return texts(out), cs


def _accepted_run(path, st, tmp_path, name="out"):
    """The whole pipeline with every suggestion accepted: (clean copy texts, outputs)."""
    out = pipeline.edit(path, st, str(tmp_path / name), stages=("style",), accept_all=True)
    return texts(out.path("clean")), out


# --- 1. the reference list's heading -------------------------------------------------------------

@pytest.mark.parametrize("text", ["References", "REFERENCES:", " REFERNCES", "References-", "References –", "Refrences",
                                  "REFRENCES.", "Bibliograpy", "7. References", "<H1>References", "Reference"])
def test_reference_heading_variants(text):
    assert references.is_heading(text)


@pytest.mark.parametrize("text", ["Preferences", "Inferences", "Referees", "Recurrences", "Reverence", "References were",
                                  "Referral"])
def test_not_reference_headings(text):
    assert not references.is_heading(text)


def test_misspelt_heading_finds_the_list_and_is_renamed(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " [1]", "Conclusion", BODY, " REFERNCES",
                           "Lee K, Park J. Back pain in adolescents. Spine 1997;22:1380-4.", "Legends-",
                           "Figure 1: Study flow"])
    ms = docmodel.load(path)
    assert [r.text[:5] for r in references.extract(ms)] == ["Lee K"]  # the legends heading ends the list
    assert rules_figures.LEGENDS_HEADING.match("Legends-")
    got, _ = apply(path, journal("ijo"), tmp_path, only=["heading_tags", "sections_headings"])
    assert "<H1>References" in got and " REFERNCES" not in got


def test_dashed_heading_loses_its_dash(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " [1]", "References-",
                           "1. Lee K, Park J. Back pain in adolescents. Spine 1997;22:1380-4."])
    got, _ = apply(path, journal("jehp"), tmp_path, only=["heading_tags", "sections_headings"])
    assert "<H1>References" in got


# --- 3. captions and legends over several paragraphs ----------------------------------------------

def test_table_label_and_title_on_two_lines_are_one_caption(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.", "Results", "Table 1. ",
                           "Baseline Characteristics of Participants (n = 110)",
                           ("table", [["Age", "Sex"], ["1", "2"]]), "Values are n (%).", BODY, *REFS])
    ms = docmodel.load(path)
    assert [(a.text, b.text[:8]) for a, b in rules_figures.split_table_captions(ms)] == [("Table 1. ", "Baseline")]
    got, _ = _accepted_run(path, journal("jehp"), tmp_path)
    i = got.index("Tables")
    assert got[i + 1] == "Table 1: Baseline characteristics of participants (n = 110)."  # in the table's first row
    assert any(t.startswith("Values are") for t in got[i:])
    assert not any(t.startswith("Baseline") for t in got[:i])  # the title is not left in the body


def test_legend_lines_and_note_move_with_the_legend(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " (Figure 1) [1].", "Figure 1. Alignment of the learning stages and ",
                           "CBE competency domains.", "Note. CBE = competency-based education.", "", "Methods", BODY,
                           *REFS])
    got, _ = _accepted_run(path, journal("jehp"), tmp_path)
    i = next(k for k, t in enumerate(got) if t.startswith("Lee K") or "Lee K" in t)
    assert got[i + 1:] == ["Figure Legend", "Figure 1: Alignment of the learning stages and CBE competency domains.",
                           "Note. CBE = competency-based education."]
    assert not any("competency domains" in t for t in got[:i])


def test_legend_text_paragraphs_after_the_references_move_with_it(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " [Figure 1] [1].", "References",
                           "1. Lee K. Back pain in adolescents. Spine 1997;22:1380-4.",
                           "Table 1: Baseline data", ("table", [["Age", "Sex"], ["1", "2"]]),
                           "Figure Legends", "Figure 1: The distribution of the subjects.",
                           "The bar chart depicts the number of patients per condition.",
                           "Figure 2: Sites sampled."])
    ms = docmodel.load(path)
    legs = rules_figures.legends(ms)
    assert [k for _, k in legs[0].conts] == ["text"]
    moves, _ = legends.plan(ms, journal("ijo"))
    blocks = next(m for m in moves if m.kind == "legends").blocks
    assert [ms.by_id(i).text[:14] for i in blocks[1]] == ["Figure 1: The ", "The bar chart "]


def test_table_note_without_a_footnote_pattern_moves_with_its_table(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.", "Table 1: Correlations",
                           ("table", [["A", "B"], ["1", "2"]]), "Cells show Spearman rho (p-value).", "Discussion", BODY,
                           *REFS])
    moves, _ = legends.plan(docmodel.load(path), journal("jehp"))
    ms = docmodel.load(path)
    blk = next(m for m in moves if m.kind == "tables").blocks[0]
    assert ms.by_id(blk[-1]).text.startswith("Cells show")


# --- 4. legends one line per part (IJO) -----------------------------------------------------------

def test_legend_parts_with_repeated_labels_run_on(tmp_path):
    path = make(tmp_path, ["Results", "Seen in Figures 1 and 2 [1].", *REFS, "Legends",
                           "Figure 1 (a & b ): Photograph of the eyelid ( yellow arrow).",
                           "Figure 1(c ): Photograph of the tarsal plate.",
                           "Figure 1 ( d):  CT scan of the orbit.", "",
                           "Figure 2(a): H and E stain ( 20 x magnification) with cellular area (Antoni A).",
                           "Figure 2 (b): Stain (40x magnification)."])
    got, _ = apply(path, journal("ijo"), tmp_path, only=["fig_captions"])
    assert got[-2:] == [
        "Figure 1: (a and b): Photograph of the eyelid (yellow arrow). (c): Photograph of the tarsal plate. (d):  CT "
        "scan of the orbit.",
        "Figure 2: (a): H and E stain (20 x magnification) with cellular area (Antoni A). (b): Stain (40x "
        "magnification)."]


# --- 5. abbreviations in brackets are no part labels ----------------------------------------------

def test_abbreviation_in_brackets_is_not_a_part_label(tmp_path):
    path = make(tmp_path, ["Results", "Seen in Figure 1 [1].", *REFS, "Figure Legends",
                           "Figure 1. (A) Slit-lamp photograph. (B) Haematoxylin and eosin (H&E) staining. (C) Alcian "
                           "blue staining."])
    got, _ = apply(path, journal("ijo"), tmp_path, only=["fig_captions"])
    assert got[-1] == "Figure 1: (a) Slit-lamp photograph. (b) Haematoxylin and eosin (H&E) staining. (c) Alcian " \
                      "blue staining."


# --- 6. a figure drawn with shapes ----------------------------------------------------------------

_TEXTBOX = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="width:200pt;height:40pt">'
            '<v:textbox><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox>'
            '</v:shape></w:pict></w:r>')


def _with_boxes(tmp_path, paras, boxes, name="m.docx"):
    """make(), then text boxes added: {index of the paragraph that anchors it: box text}."""
    path = make(tmp_path, paras, name)
    doc = Document(path)
    ps = doc.paragraphs
    for i, t in boxes.items():
        ps[i]._p.append(parse_xml(_TEXTBOX.format(text=t)))
    doc.save(path)
    return path


def test_drawn_figure_is_queried_and_its_caption_left_alone(tmp_path):
    path = _with_boxes(tmp_path, ["Introduction", BODY, "", "Figure 1 illustrates the framework [1].", "Methods", BODY,
                                  *REFS], {2: "Figure 1. Conceptual Framework"})
    ms = docmodel.load(path)
    assert [(p.text[:8], n) for p, n in rules_figures.drawn_figures(ms)] == [("Figure 1", 1)]
    assert rules_figures.legends(ms) == []
    cs = rules.run_all(ms, journal("jehp"), only=["fig_captions"])
    box = next(p for p in ms.paragraphs if rules_figures.in_textbox(p))
    assert not [c for c in cs.changes if c.para_id == box.id]


# --- 7. template leftovers --------------------------------------------------------------------------

FRONT = ["Main article", "Abstract Page", "A Study of Things in Schools: A Cross-Sectional Survey", "Abstract",
         "Background: " + BODY, "Keywords: schools", "A Study of Things in Schools: A Cross-Sectional Survey",
         "Introduction", BODY + " [1]", ".", "Methods", BODY]


def test_template_leftovers_are_deleted(tmp_path):
    # (the table in the body has its caption; its title listed again after the legends is a leftover)
    path = make(tmp_path, FRONT + ["Table 1: Baseline data", ("table", [["A", "B"], ["1", "2"]]), BODY,
                                   "Ethical approve ", "Using AI:", "", *REFS, "Bottom of Form", "Figure Legends",
                                   "Figure 1: Flow", "Table 1: Baseline data"])
    ms = docmodel.load(path)
    gone = [ms.by_id(c.para_id).text for c in leftovers.rule_leftovers(ms, journal("jehp"))]
    assert set(gone) == {"Main article", "Abstract Page", "A Study of Things in Schools: A Cross-Sectional Survey", ".",
                         "Ethical approve ", "Using AI:", "Bottom of Form", "Table 1: Baseline data"}
    assert gone.count("A Study of Things in Schools: A Cross-Sectional Survey") == 1  # the first copy stays
    assert gone.count("Table 1: Baseline data") == 1  # the caption over the table stays


def test_leftover_deletions_win_over_other_edits(tmp_path):
    path = make(tmp_path, FRONT + [*REFS])
    cs = rules.run_all(docmodel.load(path), journal("jehp"))
    ms = docmodel.load(path)
    gone = {c.para_id for c in cs.changes if c.reason == leftovers.REASON}
    assert gone and all(c.reason == leftovers.REASON for c in cs.changes if c.para_id in gone)


def test_repeated_heading_block_around_a_drawn_figure(tmp_path):
    path = _with_boxes(tmp_path, ["Introduction", BODY, "", "Materials and Methods", "Study Design and Setting", "",
                                  "Figure 1 illustrates the framework [1].", "Materials and Methods",
                                  "Study Design and Setting", BODY, *REFS], {2: "Knowledge", 5: "Figure 1. Framework"})
    ms = docmodel.load(path)
    gone = [c.para_id for c in leftovers.rule_leftovers(ms, journal("jehp"))]
    ids = [p.id for p in ms.paragraphs if p.text in ("Materials and Methods", "Study Design and Setting")]
    assert gone == ids[:2]


def test_template_text_boxes_and_repeated_pictures(tmp_path):
    path = _with_boxes(tmp_path, ["Introduction", "", BODY + " [1]", ("picture", None), *REFS, ("picture", None)],
                       {1: "Text"})
    # the picture after the references: the same image as the one in the body
    doc = Document(path)
    pics = [p for p in doc.paragraphs if p._p.xpath(".//a:blip")]
    rid = pics[0]._p.xpath(".//a:blip")[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    for b in pics[1]._p.xpath(".//a:blip"):
        b.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    doc.add_paragraph("End.")
    doc.save(path)
    ms = docmodel.load(path)
    recs = leftovers.plan(ms, journal("jehp"))
    assert [(r.boxes, len(r.pictures)) for r in recs] == [([next(p.id for p in ms.paragraphs if p.text == "Text")], 0),
                                                          ([], 1)]
    d = recs[0].to_dict()
    assert pipeline.insert_from_dict(d).to_dict() == d
    for tracked_copy in (True, False):
        for r in recs:
            r.status = Status.ACCEPTED
        ms = docmodel.load(path)
        leftovers.render(ms, recs, (Status.ACCEPTED,), tracked=tracked_copy, author="Edit OS",
                         date="2026-10-27T00:00:00Z")
        ms.save(tmp_path / "o.docx")
        body = Document(tmp_path / "o.docx").element.body
        live_pics = [b for b in body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
                     if not any(a.tag == W + "del" for a in b.iterancestors())]
        live_boxes = [x for x in body.iter(W + "txbxContent") if not any(a.tag == W + "del" for a in x.iterancestors())]
        assert len(live_pics) == 1 and not live_boxes


# --- 2. tables, figures and supplementary material in other Word files ------------------------------

def _docx(paras, boxes=None) -> bytes:
    doc = Document()
    for item in paras:
        if isinstance(item, str):
            doc.add_paragraph(item)
        elif item[0] == "table":
            rows = item[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
        elif item[0] == "picture":
            doc.add_paragraph().add_run().add_picture(io.BytesIO(picture(seed=item[1])), width=Inches(1))
    for i, t in (boxes or {}).items():
        doc.paragraphs[i]._p.append(parse_xml(_TEXTBOX.format(text=t)))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


MANUSCRIPT = ["Original Article", "Introduction", BODY + " See Table 1, Table 2 and Figure 1.", "Acknowledgement",
              "We thank the schools.", "Financial support and sponsorship", "Nil.", *REFS]


def _zip(tmp_path, files, paras=MANUSCRIPT):
    ms = make(tmp_path, paras, "Manuscript.docx")
    z = tmp_path / "sub.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(ms, "ART/Manuscript.docx")
        zf.writestr("ART/ART.xml", "<article/>")
        for name, data in files.items():
            zf.writestr(f"ART/{name}", data)
    return str(z)


TABLES = _docx(["Tables and Figure:", "Table 1: Socio-demographic Characteristics of Participants [n=298]",
                ("table", [["Variable", "Percent"], ["Age", "79.9"]]), "SD = standard deviation", "",
                "Figure 1: Conceptual framework of the model", ("picture", 3), "",
                ("table", [["Variable", "N"], ["Sex", "12"]]), "Table 2 detailed description of the cases"])


def test_tables_and_figures_from_a_word_file_are_merged(tmp_path):
    z = _zip(tmp_path, {"Revised Figure,Table .docx": TABLES})
    got, out = _accepted_run(z, journal("jehp"), tmp_path)
    i = got.index("Tables")
    assert got[i:] == ["Tables", "Table 1: Socio-demographic characteristics of participants [n=298].", "Variable",
                       "Percent", "Age", "79.9", "SD = standard deviation", "Table 2: Detailed description of the cases.",
                       "Variable", "N", "Sex", "12"]
    j = got.index("Figure Legend")
    assert got[j + 1] == "Figure 1: Conceptual framework of the model."
    assert len(Document(out.path("clean")).inline_shapes) == 1  # the picture from the file, above its legend
    man, _, inserts, _ = pipeline.load_records(str(tmp_path / "out"))
    imp = next(n for n in inserts if supplied.is_import(n))
    assert imp.name == "Revised Figure,Table .docx" and pipeline.insert_from_dict(imp.to_dict()).to_dict() == imp.to_dict()


def test_imported_tables_are_tracked_insertions_and_pending_ones_stay_out_of_the_clean_copy(tmp_path):
    z = _zip(tmp_path, {"Table 1 and 2.docx": TABLES})
    out = pipeline.edit(z, journal("jehp"), str(tmp_path / "out"), stages=("style",))
    assert "Tables" not in texts(out.path("clean"))
    body = Document(out.path("tracked")).element.body
    tbls = body.findall(W + "tbl")
    assert len(tbls) == 2 and all(r.find(f"{W}trPr/{W}ins") is not None for t in tbls for r in t.findall(W + "tr"))


def test_a_table_the_manuscript_has_is_not_imported_again(tmp_path):
    paras = MANUSCRIPT[:3] + ["Table 1: Baseline", ("table", [["A", "B"], ["1", "2"]])] + MANUSCRIPT[3:]
    z = _zip(tmp_path, {"Tables.docx": TABLES}, paras)
    _, out = _accepted_run(z, journal("jehp"), tmp_path)
    man, _, inserts, _ = pipeline.load_records(str(tmp_path / "out"))
    imp = next(n for n in inserts if supplied.is_import(n))
    assert "Table 2" in imp.text and "Table 1" not in imp.text


def test_a_questionnaire_file_is_supplementary_material(tmp_path):
    q = _docx(["QUESTIONNAIRE", "How old are you?", "Where do you live?"])
    z = _zip(tmp_path, {"Study Questionnaire.docx": q})
    got, out = _accepted_run(z, journal("jehp"), tmp_path)
    i = got.index("<H2>Supplementary Material")
    assert got[i - 1] == "We thank the schools." and got[i + 1:i + 4] == ["QUESTIONNAIRE", "How old are you?",
                                                                       "Where do you live?"]
    qs = [d for d in pipeline.load_records(str(tmp_path / "out"))[3] if d["check"].startswith("supplied")]
    assert len(qs) == 1 and "Study Questionnaire.docx" in qs[0]["text"] and "has been included" in qs[0]["text"]


def test_a_questionnaire_that_cannot_be_copied_gets_the_heading_and_a_query(tmp_path):
    q = _docx(["QUESTIONNAIRE", "", "How old are you?"], boxes={1: "Tick one"})
    z = _zip(tmp_path, {"INSTRUMENT_Questionnaire.docx": q})
    st = journal("jehp")
    assert "supplied.supplementary_file" in st.get("queries.drop")  # JEHP editors never ask it (owner, 2026-10-30)
    st.settings["queries"]["drop"] = []
    got, out = _accepted_run(z, st, tmp_path)
    i = got.index("<H2>Supplementary Material")
    assert "QUESTIONNAIRE" not in got
    qs = [d for d in pipeline.load_records(str(tmp_path / "out"))[3] if d["check"].startswith("supplied")]
    assert len(qs) == 1 and "INSTRUMENT_Questionnaire.docx" in qs[0]["text"] and "Questionnaire" in qs[0]["text"]
    assert got[i - 1] == "We thank the schools."


def test_correspondence_files_are_not_merged(tmp_path):
    z = _zip(tmp_path, {"Reply to reviewers.docx": TABLES, "Author contribution form.docx": TABLES})
    _, out = _accepted_run(z, journal("jehp"), tmp_path)
    assert not [n for n in pipeline.load_records(str(tmp_path / "out"))[2] if supplied.is_import(n)]


def test_a_questionnaire_the_manuscript_already_has_is_not_merged_again(tmp_path):
    lines = ["How old are you at your last birthday?", "Where do you live with your family?"]
    q = _docx(["QUESTIONNAIRE", *lines])
    z = _zip(tmp_path, {"Questionnaire.docx": q}, MANUSCRIPT + ["Appendix", *lines])
    _, out = _accepted_run(z, journal("jehp"), tmp_path)
    assert not [n for n in pipeline.load_records(str(tmp_path / "out"))[2] if supplied.is_import(n)]
    assert any("already" in w["message"] for w in out.summary["editor_warnings"])


def test_a_repeated_picture_with_its_legend_after_the_references_stays(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " [1]", ("picture", None), "Figure 1: Flow", *REFS,
                           ("picture", None), "Figure 1. PRISMA flow diagram", "End."])
    doc = Document(path)
    pics = [p for p in doc.paragraphs if p._p.xpath(".//a:blip")]
    emb = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    pics[1]._p.xpath(".//a:blip")[0].set(emb, pics[0]._p.xpath(".//a:blip")[0].get(emb))
    doc.save(path)
    assert leftovers.plan(docmodel.load(path), journal("jehp")) == []


def _member(tmp_path, names, xml=None):
    z = tmp_path / "m.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for n, size in names.items():
            zf.writestr(n, b"x" * size)
        if xml is not None:
            zf.writestr("A.xml", xml)
    return pipeline.manuscript_member(zipfile.ZipFile(z))


def test_the_clean_manuscript_is_chosen_over_a_track_changes_copy(tmp_path):
    assert _member(tmp_path, {"Manuscript-track changes.docx": 900, "Manuscript.docx": 500}) == "Manuscript.docx"
    assert _member(tmp_path, {"Paper highlighted.docx": 900, "Paper.docx": 500}) == "Paper.docx"
    xml = "<article><articlefile>Revised paper.docx</articlefile></article>"
    assert _member(tmp_path, {"Manuscript-track changes.docx": 900, "Revised paper.docx": 500}, xml) == \
        "Revised paper.docx"


def test_an_appendix_file_is_merged_as_supplementary_material(tmp_path):
    # owner, 2026-10-27: Appendix / Supplementary Word files are merged like a questionnaire (S372)
    a = _docx(["Appendix 1.", "Item one", "Item two"])
    z = _zip(tmp_path, {"Appendix 1.docx": a})
    got, out = _accepted_run(z, journal("jehp"), tmp_path)
    i = got.index("<H2>Supplementary Material")
    assert got[i + 1:i + 4] == ["Appendix 1.", "Item one", "Item two"]


def test_several_supplementary_files_go_under_one_heading_in_file_order(tmp_path):
    # Three Supplementary Word files: one "Supplementary Material" heading (S372), the files under it in order --
    # not a heading per file, nor the last file first.
    files = {"05_Supplementary_File1_Checklist.docx": _docx(["CHECKLIST", "Item one is listed here"]),
             "06_Supplementary_File2_Search.docx": _docx(["SEARCH TERMS", "Kettle OR teapot"]),
             "07_Supplementary_File3_Details.docx": _docx(["DETAILS", "The kettle was boiled twice"])}
    z = _zip(tmp_path, files)
    got, out = _accepted_run(z, journal("jehp"), tmp_path)
    assert got.count("<H2>Supplementary Material") == 1
    i = got.index("<H2>Supplementary Material")
    assert got[i - 1] == "We thank the schools."
    assert got[i + 1:i + 7] == ["CHECKLIST", "Item one is listed here", "SEARCH TERMS", "Kettle OR teapot", "DETAILS",
                                "The kettle was boiled twice"]
    tr = texts(out.path("tracked"))
    assert tr.count("<H2>Supplementary Material") == 1 and \
        tr.index("CHECKLIST") < tr.index("SEARCH TERMS") < tr.index("DETAILS")
