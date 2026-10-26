"""JEHP citation fixes (synthetic text): "surname, year" citations (S320), "&" between cited authors (G293),
figure / table citations with words inside the brackets (S260), malformed superscript citations (S310, S316),
powers in formulas, punctuation never superscript, guide replacements in table cells."""

import zipfile

from docx import Document

from edit_os import docmodel, rules, rules_guide, style, tracked
from edit_os.changes import Status

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def jehp():
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    return st


def _refs(*entries):
    return ["References"] + [f"{i}. {e}" for i, e in enumerate(entries, 1)]


def make(tmp_path, paras, name="m.docx", table=None):
    """Paragraphs: a string, or a list of (text, {"superscript": True, ...}) runs. `table`: cell texts of a one-row
    table placed after the first paragraph."""
    doc = Document()
    for i, item in enumerate(paras):
        p = doc.add_paragraph()
        for text, fmt in ([(item, {})] if isinstance(item, str) else item):
            run = p.add_run(text)
            for k, v in fmt.items():
                setattr(run.font, k, v)
        if i == 0 and table:
            t = doc.add_table(rows=1, cols=len(table))
            for cell, item in zip(t.rows[0].cells, table):
                cp = cell.paragraphs[0]
                for text, fmt in ([(item, {})] if isinstance(item, str) else item):
                    run = cp.add_run(text)
                    for k, v in fmt.items():
                        setattr(run.font, k, v)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def run(path, st=None, only=None):
    st = st or jehp()
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(path).replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(str(path).replace(".docx", "_tr.docx"))
    return cs, out, [p.text for p in docmodel.load(out).paragraphs]


INTRO = "Introduction"
REFS = _refs("Nagarajappa R, Batra M, Sharda AJ. Antimicrobial effect. J Dent 2015;3:1-2.",
             "Sorte SR, Rawekar AT. Future-ready medicine. J Med Educ 2025;4:1-2.",
             "Choi A, Lee B, Kim C. First study. J Test 2024;5:1-2.",
             "Choi A, Park D. Second study. J Test 2024;6:1-2.",
             "Wagner E, Hall F. Peer support. J Test 2016;7:1-2.",
             "Castillo G, Ross H. Group care. J Test 2021;8:1-2.",
             "Chaturvedi AP, Tripathi YB. Leaves. J Test 2011;9:1-2.")
CITE = ["author_year", "citations", "guide_replace"]


def test_year_beside_the_number_goes(tmp_path):
    body = [INTRO,
            "In vitro evidence: Nagarajappa et al. (2015) [1] used the agar method in their laboratory work.",
            "Sorte et al. (2025) reported strong interest in incorporating AI into the curriculum.[2]",
            "Studies by Wagner et al. (2016) and Castillo et al. (2021) [5,6] demonstrate peer interaction effects."]
    _, _, t = run(make(tmp_path, body + REFS), only=CITE)
    assert t[1] == "In vitro evidence: Nagarajappa et al.[1] used the agar method in their laboratory work."
    assert t[2] == "Sorte et al. reported strong interest in incorporating AI into the curriculum.[2]"
    assert t[3] == "Studies by Wagner et al. and Castillo et al.[5,6] demonstrate peer interaction effects."


def test_surname_year_becomes_the_number_only_when_unique(tmp_path):
    body = [INTRO,
            "As shown by Sorte et al. (2025), most respondents lacked any prior training in the area.",
            "Choi et al. (2024) found that the effect was larger in the older group of students.",  # two Choi 2024
            "Smith et al. (2019) found that the effect was larger in the younger group of students.",  # not listed
            "Earlier work by Sorte et al. (2019) was smaller in scope and in its sample size."]  # wrong year
    _, _, t = run(make(tmp_path, body + REFS), only=CITE)
    assert t[1] == "As shown by Sorte et al.[2], most respondents lacked any prior training in the area."
    assert t[2] == body[2]
    assert t[3] == body[3]
    assert t[4] == body[4]


def test_table_cell_author_year(tmp_path):
    cells = ["Sorte et al. (2025)", "Nagarajappa et al. (2015) [1]", "Sorte et al., 2025",
             "Nagarajappa et al., 2015.[1]", "Kosasih et al. (2024)"]
    path = make(tmp_path, [INTRO, "Text cites the two studies here.[1,2]"] + REFS, table=cells)
    _, out, _ = run(path, only=CITE)
    doc = Document(out)
    got = [c.text for c in doc.tables[0].rows[0].cells]
    # a study column's "Surname et al., year.[n]" keeps its year (the editors); a study not in the list stays
    assert got == ["Sorte et al.[2]", "Nagarajappa et al.[1]", "Sorte et al.[2]", "Nagarajappa et al., 2015.[1]",
                   "Kosasih et al. (2024)"]


def test_author_year_off_by_default(tmp_path):
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    body = [INTRO, "Sorte et al. (2025) reported strong interest in incorporating AI into the curriculum."]
    _, _, t = run(make(tmp_path, body + REFS), st=st, only=["author_year"])
    assert t[1] == body[1]


def test_ampersand_between_cited_authors():
    f = rules_guide.find_ampersand
    assert f("Chaturvedi & Tripathi [8] extended it", {}) == [(10, 13, " and ")]
    assert f("Chaturvedi AP & Tripathi YB (2011)", {}) == [(13, 16, " and ")]
    assert f("Sun & Kao (2024) [28]", {}) == [(3, 6, " and ")]
    assert f("Goss & Creutzfeldt¹²", {}) == [(4, 7, " and ")]
    for keep in ("Jammu & Kashmir", "Johnson & Johnson products", "Systematic Review & Meta-Analysis", "CPP & RCI"):
        assert f(keep, {}) == [], keep


def test_ampersand_in_text_not_in_references(tmp_path):
    body = [INTRO, "Chaturvedi & Tripathi [7] extended the in vitro observations to animal models."]
    refs = REFS[:-1] + ["7. Chaturvedi AP & Tripathi YB. Leaves. J Test 2011;9:1-2."]
    _, _, t = run(make(tmp_path, body + refs), only=CITE)
    assert t[1] == "Chaturvedi and Tripathi[7] extended the in vitro observations to animal models."
    assert "&" in t[-1]


def test_figure_citation_words_inside_brackets(tmp_path):
    body = [INTRO, "Records were screened in two stages (see Figure 1). The analysis of the 106 documents "
                   "(shown in figure 2) found growth. Values ranged widely (Supplementary Table S2). "
                   "The work was cited (Supplimentary Table 1), as noted.",
            "Chemical profiling used TLC, and LC-MS/MS analyses. (Table 1)",
            "Table 1: Results", "Figure 1: Flow", "Figure 2: Growth"]
    _, _, t = run(make(tmp_path, body), only=["fig_citations"])
    assert t[1] == ("Records were screened in two stages [Figure 1]. The analysis of the 106 documents [Figure 2] "
                    "found growth. Values ranged widely [Supplementary Table S2]. The work was cited "
                    "[Supplementary Table 1], as noted.")
    assert t[2] == "Chemical profiling used TLC, and LC-MS/MS analyses [Table 1]."


def test_figure_lead_words_left_by_default(tmp_path):
    st = style.load("ijo")
    body = [INTRO, "Records were screened in two stages (see Fig. 1).", "Fig. 1: Flow"]
    _, _, t = run(make(tmp_path, body), st=st, only=["fig_citations"])
    assert t[1] == body[1]


def _sup(text):
    return (text, {"superscript": True})


def _plain(text):
    return (text, {})


def test_malformed_superscript_citations(tmp_path):
    body = [INTRO,
            [_plain("Training on normal delivery procedures."), _sup(" 3,5, "), _plain("Given the high rates.")],
            [_plain("This led to safer maternity services in the region."), _sup(" 3, ,9")],
            [_plain("Teamwork and communication skills."), _sup("5,6.7")],
            [_plain("Especially when baseline knowledge is already high"), _sup("17."), ],
            [_plain("Prior reports agree on the plateau in this indicator"), _sup("18,19,20"), _plain(".")],
            [_plain("This can compromise patient outcomes. "), _sup("3"), _plain("Simulation-based training helps.")]]
    _, _, t = run(make(tmp_path, body + _refs(*[f"Author{i} A. Title. J 2001;1:1-2." for i in range(1, 21)])),
                  only=["citations"])
    assert t[1] == "Training on normal delivery procedures.[3,5] Given the high rates."
    assert t[2] == "This led to safer maternity services in the region.[3,9]"
    assert t[3] == "Teamwork and communication skills.[5-7]"
    assert t[4] == "Especially when baseline knowledge is already high.[17]"
    assert t[5] == "Prior reports agree on the plateau in this indicator.[18-20]"
    assert t[6] == "This can compromise patient outcomes.[3] Simulation-based training helps."


def test_value_with_decimal_in_superscript_is_not_a_citation(tmp_path):
    (p,) = [q for q in docmodel.load(make(tmp_path, [[_plain("The mean was"), _sup("2.1")]])).paragraphs if q.text]
    assert rules.superscript_citations(p) == []


def test_power_after_bracketed_letter_in_formula(tmp_path):
    (p,) = [q for q in docmodel.load(make(tmp_path, ["using the Yamane formula n = N / (1 + N(e)²)."])).paragraphs
            if q.text]
    assert rules.superscript_citations(p) == []


def _runs(path, idx):
    """(text, vertAlign) of each run of paragraph `idx` of a rendered file."""
    root = zipfile.ZipFile(path).read("word/document.xml")
    from lxml import etree
    paras = [p for p in etree.fromstring(root).iter(W + "p")]
    out = []
    for r in paras[idx].iter(W + "r"):
        va = r.find(f"{W}rPr/{W}vertAlign")
        out.append(("".join(t.text or "" for t in r.iter(W + "t")), None if va is None else va.get(W + "val")))
    return out


def test_full_stop_before_superscript_citation_is_not_superscript(tmp_path):
    body = [INTRO, [_plain("Stunting determinants remain lacking"), _sup("[13]"), _plain(". Next one.")],
            [_plain("Home-based rehabilitation in CR"), _sup(".24-30")]]
    path = make(tmp_path, body + _refs(*[f"Author{i} A. Title. J 2001;1:1-2." for i in range(1, 31)]))
    _, out, t = run(path, only=["citations"])
    assert t[1] == "Stunting determinants remain lacking.[13] Next one."
    assert t[2] == "Home-based rehabilitation in CR.[24-30]"
    for idx in (1, 2):
        runs = _runs(out, idx)
        assert all(va != "superscript" for text, va in runs if "." in text), runs
        assert any(va == "superscript" and "[" in text for text, va in runs), runs


def test_citation_after_possessive_name(tmp_path):
    st = jehp()
    st.settings["citations"]["next_to_author"] = True
    body = [INTRO, [_plain("Roy et al.’s case study rated the answers in scenarios of ethics."), _sup("[1]"),
                    _plain(" Others agree.")]]
    path = make(tmp_path, body + _refs("Roy AD, Das D. Efficacy. J Educ 2024;13:22."))
    _, _, t = run(path, st=st, only=["author_citations", "citations"])
    assert t[1].startswith("Roy et al.’s[1] case study")


def test_jehp_keeps_citation_where_the_author_put_it(tmp_path):
    body = [INTRO, "Roy et al. rated the answers in scenarios of ethics and communication.[1]"]
    _, _, t = run(make(tmp_path, body + _refs("Roy AD, Das D. Efficacy. J Educ 2024;13:22.")),
                  only=["author_citations", "citations"])
    assert t[1] == body[1]


def test_guide_replacements_in_table_cells(tmp_path):
    cells = ["p", "Vigorous activity >=30 min/day", "Aged 18-59 years", "-5.47", "Care & support"]
    path = make(tmp_path, [INTRO, "Body text."], table=cells)
    _, out, _ = run(path, only=["guide_replace"])
    got = [c.text for c in Document(out).tables[0].rows[0].cells]
    assert got == ["P", "Vigorous activity ≥30 min/day", "Aged 18–59 years", "−5.47", "Care and support"]
