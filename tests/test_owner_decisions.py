"""The owner's decisions of 2026-10-26 (IJO, JEHP, MJDRDYPU) that no other test file covers."""

import zipfile

from docx import Document
from docx.oxml.ns import qn

from edit_os import docmodel, frontmatter, refengine, rules, rules_format, style, tracked
from edit_os.changes import Status


def _render(path, cs, out, track=False):
    for c in cs.changes:
        c.status = Status.ACCEPTED
    tracked.render(docmodel.load(path), cs, tracked=track).save(out)
    return out


def _save(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = str(tmp_path / name)
    doc.save(path)
    return path


SYM_TEXT = "Mean −2.1 ± 0.5 (α ≤ 0.05, β ≥ 2) at 5 μg, 3 × 4 mm, µm and ©."


def test_symbol_font_for_greek_plus_minus_and_comparisons(tmp_path):
    path = _save(tmp_path, ["Introduction", SYM_TEXT])
    for code in ("ijo", "jehp"):
        cs = rules.run_all(docmodel.load(path), style.load(code), only=["format_symbols"])
        assert sorted(c.old for c in cs.changes) == sorted("−±α≤β≥")  # not ×, µ, ©, nor the unit's μ ("μg")
        for track in (False, True):
            out = _render(path, cs, str(tmp_path / f"{code}{track}.docx"), track)
            xml = zipfile.ZipFile(out).read("word/document.xml").decode()
            for code_point in ("F02D", "F0B1", "F061", "F0A3", "F062", "F0B3"):
                assert f'<w:sym w:font="Symbol" w:char="{code_point}"/>' in xml
            if not track:  # the text view (and so QC and the rules) still sees the characters
                assert docmodel.load(out).paragraphs[1].text == SYM_TEXT
    assert rules.run_all(docmodel.load(path), style.load("mjdrdypu"), only=["format_symbols"]).changes == []


def test_symbols_already_inserted_are_read_and_left(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Mean 2.1 ")
    r = p.add_run()._r
    sym = r.makeelement(qn("w:sym"), {qn("w:font"): "Symbol", qn("w:char"): "F0B1"})
    r.append(sym)
    p.add_run(" 0.5")
    path = str(tmp_path / "s.docx")
    doc.save(path)
    ms = docmodel.load(path)
    assert ms.paragraphs[0].text == "Mean 2.1 ± 0.5"
    assert rules_format.rule_symbols(ms, style.load("ijo")) == []
    other = r.makeelement(qn("w:sym"), {qn("w:font"): "Wingdings", qn("w:char"): "F0B1"})
    assert docmodel.sym_char(other) == ""


def test_reference_initials_cut_to_two_for_jehp_only():
    assert refengine.cut_initials("Chandrasena PPCM", 2) == "Chandrasena PP"
    assert refengine.cut_initials("Smith J-P", 2) == "Smith J-P"
    assert refengine.cut_initials("Cunningham ETAB Jr", 2) == "Cunningham ET Jr"
    assert refengine.cut_initials("WHO", 2) == "WHO"
    text = "Chandrasena PPCM, Ten Ham-Baloyi WAB, Lee K. A title. J Test 2001;1:1-2."
    got = rules._apply_edits(text, rules._reference_edits(text, style.load("jehp")))
    assert got.startswith("Chandrasena PP, Ten Ham-Baloyi WA, Lee K.")
    got = rules._apply_edits(text, rules._reference_edits(text, style.load("ijo")))  # IJO editors keep them
    assert got.startswith("Chandrasena PPCM, Ten Ham-Baloyi WAB")


def test_ijo_third_level_heading_italic_not_bold(tmp_path):
    st = style.load("ijo")
    st.settings["formatting"]["heading_italic_levels"] = [2]  # (IJO: H3; H2 here for a short test document)
    st.settings["formatting"]["heading_unbold_levels"] = [2]
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Some text about the study that is long enough to be a paragraph of running text. " * 3)
    doc.add_paragraph().add_run("Methods").bold = True
    doc.add_paragraph("Some text about the methods that is long enough to be a paragraph of running text. " * 3)
    doc.add_paragraph("Ethics approval", style="Heading 1").runs[0].bold = True
    doc.add_paragraph("The study was approved by the institutional ethics committee of the hospital. " * 3)
    path = str(tmp_path / "h.docx")
    doc.save(path)
    ms = docmodel.load(path)
    got = {(ms.by_id(c.para_id).text, c.reason): c.format for c in rules_format.rule_cleanup(ms, st)}
    assert got.get(("Ethics approval", "fmt_heading_unbold")) == {"bold": False}
    assert ("Ethics approval", "fmt_heading_bold") not in got  # not bolded again for its Word style
    assert ("Methods", "fmt_heading_unbold") not in got  # an H1 stays bold


def test_table_note_abbreviations_take_the_journals_form(tmp_path):
    def note(code, text):
        doc = Document()
        doc.add_paragraph("Introduction")
        doc.add_paragraph("See Table 1.")
        doc.add_paragraph("Table 1: Results")
        t = doc.add_table(rows=2, cols=2)
        for c, v in zip(t.rows[0].cells + t.rows[1].cells, ["CI", "OR", "1", "2"]):
            c.text = v
        doc.add_paragraph(text)
        path = str(tmp_path / f"{code}.docx")
        doc.save(path)
        cs = rules.run_all(docmodel.load(path), style.load(code), only=["fig_captions"])
        out = _render(path, cs, str(tmp_path / f"{code}_out.docx"))
        return docmodel.load(out).paragraphs[-1].text

    assert note("jehp", "CI: confidence interval; OR: odds ratio") == "CI = confidence interval, OR = odds ratio"
    assert note("ijo", "CI = confidence interval; OR = odds ratio") == "CI: confidence interval, OR: odds ratio"


def test_affiliation_postal_code_kept():
    src = "1Department of Ophthalmology, Sankara Nethralaya, Chennai - 600006, Tamil Nadu, India\n" \
          "2Faculty of Nursing, Hamadan University of Medical Sciences, Hamadan 6517838736, Iran\n" \
          "Address for correspondence: Dr. A, Sankara Nethralaya, Chennai 600 001"
    aff = {"department": "Department of Ophthalmology", "institution": "Sankara Nethralaya", "city": "Chennai",
           "state": "Tamil Nadu", "country": "India"}
    assert frontmatter.affiliation_postal_code(src, aff) == "600006"  # not the correspondence's
    aff2 = {"department": "Faculty of Nursing", "institution": "Hamadan University of Medical Sciences",
            "city": "Hamadan", "country": "Iran"}
    assert frontmatter.affiliation_postal_code(src, aff2) == "6517838736"
    f = frontmatter.Fields(affiliations={1: {**aff, "postal_code": "600006"}, 2: {**aff2, "postal_code": "6517838736"}})
    assert frontmatter.affiliation_text(f, 1, postal=True) == \
        "Department of Ophthalmology, Sankara Nethralaya, Chennai - 600 006, Tamil Nadu, India"
    assert frontmatter.affiliation_text(f, 2, postal=True).startswith("Faculty of Nursing, Hamadan University")
    assert "6517838736" in frontmatter.affiliation_text(f, 2, postal=True)
    assert "600" not in frontmatter.affiliation_text(f, 1)  # journals without the setting: as before
    for code in ("ijo", "jehp"):
        assert style.load(code).get("front_matter.affiliation_postal_code") is True
