from docx import Document

from edit_os import docmodel, queries, rules, rules_figures, style, tracked
from edit_os.changes import Status


def _doc(tmp_path, paras, name="f.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def _style(code):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


def _apply(path, st, tmp_path, only=("fig_citations", "fig_captions")):
    cs = rules.run_all(docmodel.load(path), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
        assert c.layer.value == "tables_figures"
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")  # renders cleanly too
    return [p.text for p in docmodel.load(str(out)).paragraphs]


BODY = [
    "Introduction",
    "Lesions were seen (Figure 1A). Figure 2 shows the scan (Table 2, Figure 1).",
    "Both eyes (Figure 1 and Figure 2) and parts (Figures 3a-b); see (Fig. 3a & 3b).",
    "Montage in Figure 4a to 4d; Picture 5 shows it {Table II}, as in Tables 1, 2 and 3 for 71](Figure 1).",
    "Table 1 and 25 patients were seen; a figure 1a and a 3-mm cut [Fig. 1). Table-1 depicts it.",
    "Results were in (Supplementary Table 1). Delayed (Figure 1,2,3,).",
]


def test_ijo_figure_citations(tmp_path):
    t = _apply(_doc(tmp_path, BODY), _style("ijo"), tmp_path)
    assert t[1] == "Lesions were seen [Fig. 1a]. Fig. 2 shows the scan [Table 2, Fig. 1]."
    assert t[2] == "Both eyes [Figs. 1 and 2] and parts [Fig. 3a-b]; see [Fig. 3a and 3b]."
    assert t[3] == ("Montage in Fig. 4a-d; Fig. 5 shows it [Table 2], as in Tables 1–3 for 71] [Fig. 1].")
    assert t[4] == "Table 1 and 25 patients were seen; a figure 1a and a 3-mm cut [Fig. 1]. Table 1 depicts it."
    assert t[5] == "Results were in [Supplementary Table 1]. Delayed [Figs. 1–3]."


def test_base_figure_citations(tmp_path):
    t = _apply(_doc(tmp_path, BODY), _style("base"), tmp_path)
    assert t[1] == "Lesions were seen [Figure 1a]. Figure 2 shows the scan [Table 2, Figure 1]."
    assert t[2] == "Both eyes [Figures 1 and 2] and parts [Figure 3a-b]; see [Figure 3a and 3b]."
    assert t[3] == "Montage in Figure 4a-d; Figure 5 shows it [Table 2], as in Tables 1, 2 and 3 for 71] [Figure 1]."


def test_parse_groups_ranges_and_roman():
    st = _style("base")
    render = lambda s: [rules_figures.render_group(g, st) for g in rules_figures.parse_groups(s)]  # noqa: E731
    assert render("Figure 1a to 1d") == ["Figure 1a-d"]
    assert render("Figures 1 to 3") == ["Figures 1–3"]
    assert render("Table IV and Table V") == ["Tables 4 and 5"]
    assert render("Tables 1A and 1B") == ["Tables 1A and 1B"]  # lettered tables are separate tables
    assert render("Figure 2A and B") == ["Figure 2a and b"]
    assert render("Figure 1(a) and (b)") == ["Figure 1a and b"]
    assert render("Table of contents") == []
    assert rules_figures.roman_to_int("XIV") == 14 and rules_figures.roman_to_int("IIV") is None


CAPTIONS = [
    "Results",
    "Table 3 shows the results.",  # a sentence, not a caption
    "References",
    "1. Author A. Title. J Test 2001;1:1-2.",
    "Figure legends",
    "Figure 1A. Slit-lamp photograph of the right eye.",
    "Figure 1B. B-scan, et al.",
    "Fig. 2 - Fundus photograph.",
    "Picture 3: Montage",
    "Table II Baseline characteristics.",
    "Table 4. Table 4. Outcomes",
    "TABLE 5: Costs",
    "Figure 6 A) axial and B) sagittal views",
]


def test_captions(tmp_path):
    t = _apply(_doc(tmp_path, CAPTIONS), _style("ijo"), tmp_path)
    assert t[1] == "Table 3 shows the results."
    assert t[5] == "Figure 1: (a) Slit-lamp photograph of the right eye"
    assert t[6] == "(b) B-scan, et al."  # run-on part of Figure 1; "et al." keeps its period
    assert t[7] == "Figure 2: Fundus photograph"  # captions keep "Figure", even for IJO
    assert t[8] == "Figure 3: Montage"
    assert t[9] == "Table 2: Baseline characteristics"
    assert t[10] == "Table 4: Outcomes"
    assert t[11] == "Table 5: Costs"
    assert t[12] == "Figure 6 A) axial and B) sagittal views"


def test_caption_end_period_kept_when_configured(tmp_path):
    st = _style("base")
    st.settings["figures"]["caption_end_period"] = True
    t = _apply(_doc(tmp_path, ["Table 1. Baseline data."]), st, tmp_path)
    assert t[0] == "Table 1: Baseline data."


def test_figure_queries(tmp_path):
    path = _doc(tmp_path, [
        "Introduction",
        "First the scan [Figure 2], then the photo [Figure 1] and Table 1; also Figure 3 and Table 4.",
        "References",
        "1. Author A. Title. J Test 2001;1:1-2.",
        "Figure 1: Photo",
        "Figure 2: Scan",
        "Table 1: Data",
        "Table 2: Unused",
    ])
    ms = docmodel.load(path)
    st = _style("ijo")
    qs = queries.collect(ms, st)
    texts = [q.text for q in qs.queries]
    assert "Figure citation not in order." in texts
    assert "Table citation not in order." not in texts
    assert "Kindly provide the caption for Figure 3." in texts
    assert "Table 4 is cited in the text but not provided. Kindly provide." in texts
    assert "Table 2 has not been cited in the text part. Kindly cite." in texts
    assert not any("Figure 1 has not" in t or "Figure 2 has not" in t for t in texts)
    order = next(q for q in qs.queries if q.check == "figures.order")
    assert ms.by_id(order.para_id).text[order.start:order.end] == "Figure 1"
    uncited = next(q for q in qs.queries if q.check == "figures.uncited")
    assert ms.by_id(uncited.para_id).text[uncited.start:uncited.end] == "Table 2: "


def test_no_caption_query_without_any_captions(tmp_path):
    """Figures often come as separate files: no captions at all is not a query per citation."""
    path = _doc(tmp_path, ["Introduction", "See [Figure 1] and [Table 1].", "References",
                           "1. Author A. Title. J Test 2001;1:1-2."])
    qs = queries.collect(docmodel.load(path), _style("ijo"))
    assert not [q for q in qs.queries if q.check.startswith("figures.")]
