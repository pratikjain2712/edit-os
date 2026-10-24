"""Whole paragraphs and tables deleted by changes: the tracked copy deletes the paragraph mark (or the
table rows), the clean copy drops them, except where Word needs the paragraph. Synthetic text only."""

import base64
import io

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from edit_os import docmodel, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status


def _delete_all(ms, ids):
    cs = []
    for p in ms.paragraphs:
        if p.id in ids and p.text:
            cs.append(Change(f"d{p.id}", p.id, 0, len(p.text), p.text, "", Layer.STRUCTURE, Source.RULE,
                             status=Status.ACCEPTED))
    return ChangeSet(cs)


def _render(path, pick, tmp_path, tracked_mode):
    ms = docmodel.load(path)
    ids = {p.id for p in ms.paragraphs if pick(p)}
    out = tmp_path / ("t.docx" if tracked_mode else "c.docx")
    tracked.render(docmodel.load(path), _delete_all(ms, ids), tracked=tracked_mode).save(out)
    return str(out)


def _mark_deleted(p_el) -> bool:
    return p_el.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None


def _png() -> io.BytesIO:  # a 1 x 1 white PNG
    return io.BytesIO(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC"))


def test_deleted_paragraphs_lose_their_mark(tmp_path):
    doc = Document()
    doc.add_paragraph("Dear Editor, please consider our manuscript.")
    doc.add_paragraph("")
    doc.add_paragraph("Yours sincerely, the authors.")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Glaucoma is common.")
    doc.add_paragraph("Funding: None.")
    doc.add_paragraph("The last paragraph stays.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    gone = ("Dear", "Yours", "Introduction", "Funding")
    t = _render(path, lambda p: p.text.startswith(gone), tmp_path, True)
    paras = list(Document(t).element.body.iter(qn("w:p")))
    assert [_mark_deleted(p) for p in paras] == [True, True, True, True, False, True, False]  # blank line too
    assert docmodel.load(t).paragraphs[0].text == ""  # the text itself is a tracked deletion
    c = _render(path, lambda p: p.text.startswith(gone), tmp_path, False)
    assert [p.text for p in docmodel.load(c).paragraphs] == ["Glaucoma is common.", "The last paragraph stays."]


def test_kept_paragraphs_images_section_breaks_last(tmp_path):
    doc = Document()
    doc.add_paragraph("Body text before.")
    pic = doc.add_paragraph("Figure placeholder ")
    pic.add_run().add_picture(_png(), width=Inches(0.2))
    doc.add_paragraph("End of section one.")
    doc.add_section()  # adds a paragraph holding the section break (w:sectPr)
    doc.paragraphs[-1].add_run("Section break line.")
    doc.add_paragraph("Final words.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    ms = docmodel.load(path)
    assert ms.paragraphs[3].element.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    t = _render(path, lambda p: p.text != "Body text before.", tmp_path, True)
    paras = list(Document(t).element.body.iter(qn("w:p")))
    # image, section break and last paragraph keep their marks (their text is still deleted)
    assert [_mark_deleted(p) for p in paras] == [False, False, True, False, False]
    c = _render(path, lambda p: p.text != "Body text before.", tmp_path, False)
    assert len(docmodel.load(c).paragraphs) == len(ms.paragraphs) - 1


def test_deleted_table_rows_and_cells(tmp_path):
    doc = Document()
    doc.add_paragraph("Before the tables.")
    t1 = doc.add_table(rows=2, cols=2)
    for i, cell in enumerate(c for row in t1.rows for c in row.cells):
        cell.text = f"Reviewer comment {i}"
    doc.add_paragraph("Between the tables.")
    t2 = doc.add_table(rows=1, cols=1)
    cell = t2.rows[0].cells[0]
    cell.text = "Keep this cell text."
    cell.add_paragraph("Delete this second line.")
    doc.add_paragraph("After the tables.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    pick = lambda p: p.text.startswith(("Reviewer", "Delete"))  # noqa: E731
    t = Document(_render(path, pick, tmp_path, True))
    tbls = t.element.body.findall(qn("w:tbl"))
    assert all(tr.find(f"{qn('w:trPr')}/{qn('w:del')}") is not None for tr in tbls[0].iter(qn("w:tr")))
    assert tbls[1].find(f".//{qn('w:trPr')}") is None
    cell_paras = list(tbls[1].iter(qn("w:p")))
    assert [_mark_deleted(p) for p in cell_paras] == [False, False]  # the cell's last paragraph keeps its mark
    c = docmodel.load(_render(path, pick, tmp_path, False))
    assert len(c.document.element.body.findall(qn("w:tbl"))) == 1
    assert [p.text for p in c.paragraphs] == ["Before the tables.", "Between the tables.", "Keep this cell text.",
                                              "", "After the tables."]


def test_paragraph_between_two_tables_stays(tmp_path):
    doc = Document()
    doc.add_paragraph("Intro.")
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "A"
    doc.add_paragraph("Separator to delete.")
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "B"
    doc.add_paragraph("End.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    c = docmodel.load(_render(path, lambda p: p.text.startswith("Separator"), tmp_path, False))
    assert [p.text for p in c.paragraphs] == ["Intro.", "A", "", "B", "End."]  # tables would merge


def test_partial_deletion_keeps_paragraph(tmp_path):
    doc = Document()
    doc.add_paragraph("Keep me. Drop me.")
    doc.add_paragraph("End.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    ms = docmodel.load(path)
    p = ms.paragraphs[0]
    cs = ChangeSet([Change("x", p.id, 8, len(p.text), p.text[8:], "", Layer.LANGUAGE, Source.RULE,
                           status=Status.ACCEPTED)])
    out = tmp_path / "t.docx"
    tracked.render(docmodel.load(path), cs, tracked=True).save(out)
    assert not _mark_deleted(next(Document(str(out)).element.body.iter(qn("w:p"))))


def test_insert_anchored_to_deleted_paragraph(tmp_path):
    from edit_os import frontmatter

    doc = Document()
    doc.add_paragraph("Old running head")
    doc.add_paragraph("Title of the study")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    ms = docmodel.load(path)
    cs = _delete_all(ms, {ms.paragraphs[0].id})
    new = frontmatter.NewParagraph("n1", ms.paragraphs[0].id, "after", [("New running head", {})], "running_head",
                                   status=Status.ACCEPTED)
    out = docmodel.load(path)
    tracked.render(out, cs, tracked=False)
    frontmatter.render_inserts(out, [new], tracked=False)
    out.save(tmp_path / "c.docx")
    assert [p.text for p in docmodel.load(str(tmp_path / "c.docx")).paragraphs] == ["New running head",
                                                                                   "Title of the study"]


# --- a paragraph the changes empty: no stray heading tag, no empty line ------------------------------

def _doc(tmp_path, *texts):
    doc = Document()
    for t in texts:
        doc.add_paragraph(t)
    path = str(tmp_path / "m.docx")
    doc.save(path)
    return path


def _ch(cid, p, start, end, new, reason, status=Status.ACCEPTED):
    return Change(cid, p.id, start, end, p.text[start:end], new, Layer.STRUCTURE, Source.RULE, reason=reason,
                  status=status)


def _texts(path):
    return [p.text for p in docmodel.load(path).paragraphs]


def test_heading_deleted_and_tagged_leaves_no_tag(tmp_path):
    # heading_tags tags the author's heading; the back-matter rule deletes it (IJO "Conflict of interest").
    path = _doc(tmp_path, "Conflict of interest", "None declared.", "End.")
    ms = docmodel.load(path)
    head = ms.paragraphs[0]
    cs = ChangeSet([_ch("tag", head, 0, 0, "<H2>", "heading_tags"),
                    _ch("del", head, 0, len(head.text), "", "sections_back_matter")])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["None declared.", "End."]
    t = tmp_path / "t.docx"
    tracked.render(docmodel.load(path), cs, tracked=True).save(t)
    first = list(Document(str(t)).element.body.iter(qn("w:p")))[0]
    assert _mark_deleted(first)
    assert "<H2>" not in "".join(x.text or "" for x in first.iter(qn("w:t")))  # no tag inserted on it


def test_heading_tag_kept_when_deletion_not_accepted(tmp_path):
    path = _doc(tmp_path, "Conflict of interest", "None declared.")
    head = docmodel.load(path).paragraphs[0]
    cs = ChangeSet([_ch("tag", head, 0, 0, "<H2>", "heading_tags"),
                    _ch("del", head, 0, len(head.text), "", "sections_back_matter", Status.REJECTED)])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["<H2>Conflict of interest", "None declared."]


def test_heading_tag_over_leading_spaces_dropped_with_heading(tmp_path):
    path = _doc(tmp_path, "  Conflicts of Interest ", "None.")
    head = docmodel.load(path).paragraphs[0]
    cs = ChangeSet([_ch("tag", head, 0, 2, "<H2>", "heading_tags"),
                    _ch("del", head, 2, len(head.text) - 1, "", "sections_back_matter")])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["None."]


def test_text_replaced_by_new_paragraphs_leaves_no_empty_line(tmp_path):
    # The statement goes and the journal's section is inserted after it, opening new paragraphs.
    path = _doc(tmp_path, "The authors declare no conflict.", "End.")
    p = docmodel.load(path).paragraphs[0]
    n = len(p.text)
    cs = ChangeSet([_ch("del", p, 0, n, "", "sections_back_matter"),
                    _ch("ins", p, n, n, " <H2>Authors’ contributions A did it.", "sections_back_matter")])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["<H2>Authors’ contributions", "A did it.", "End."]
    t = tmp_path / "t.docx"
    tracked.render(docmodel.load(path), cs, tracked=True).save(t)
    assert _mark_deleted(list(Document(str(t)).element.body.iter(qn("w:p")))[0])


def test_heading_replaced_by_tagged_heading_is_kept(tmp_path):
    path = _doc(tmp_path, "Conflict of interest", "None.")
    head = docmodel.load(path).paragraphs[0]
    cs = ChangeSet([_ch("rep", head, 0, len(head.text), "<H2>Conflicts of interest", "sections_back_matter")])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["<H2>Conflicts of interest", "None."]


def test_spaces_tidied_in_blank_paragraph_kept(tmp_path):
    path = _doc(tmp_path, "Text.", "   ", "End.")
    blank = docmodel.load(path).paragraphs[1]
    cs = ChangeSet([_ch("sp", blank, 0, 3, " ", "spaces")])
    c = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(c)
    assert _texts(c) == ["Text.", " ", "End."]
