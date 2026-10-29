"""QC agrees with the citation rule and renumbering on what is a citation; back matter the journal keeps
as written passes QC; junk lines under the reference list are not entries. All text is synthetic."""

import zipfile

from docx import Document

from edit_os import docmodel, pipeline, qc, queries, references, renumber, rules, style, tracked
from edit_os.changes import Status


def _style(code="jehp"):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


def _refs(n):
    return ["References"] + [f"{i}. Author{chr(64 + i)} A, Writer B. A study {i}. J Test 2001;{i}:1-9." for i in range(1, n + 1)]


def _doc(tmp_path, body, n_refs=5, name="m.docx"):
    doc = Document()
    for t in ["Introduction"] + body + _refs(n_refs):
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def _issues(ms, st, check):
    return [i.message for i in qc.check(ms, st).issues if i.check == check]


def _converted(ms, st):
    return [c.new.strip(".") for c in rules.rule_citations(ms, st) if "[" in c.new]


def test_counts_lists_and_statistics_are_not_round_bracket_citations(tmp_path):
    st = _style()
    path = _doc(tmp_path, [
        "Studies came from Germany (178), Australia (165) and Italy (133), as reported before.[1]",
        "Inclusion parameters: (1) Peer-reviewed studies; (2) Specifically adult units; (3) Adults. "
        "Exclusion parameters: (1) Neonatal units; (2) \"Surrogate\" as an analogy; (3) Grey literature.",
        "The scores rose [t(48) = -11.4, P < 0.001] and differed [F(3, 44) = 0.9] and r (198) = 0.41.[2]",
        "Seven behavior items (0-2) were scored, with P = 4.57 × 10⁻⁸ for the change.[3]",
    ])
    ms = docmodel.load(path)
    assert _converted(ms, st) == []
    assert rules.round_citations(ms, st) == []
    assert _issues(ms, st, "citations.brackets") == []


def test_qc_flags_what_the_rule_would_convert(tmp_path):
    st = _style()
    path = _doc(tmp_path, ["The disease is common (1). It is treated with surgery (2, 3) and drops (4)."], 4)
    ms = docmodel.load(path)
    assert _issues(ms, st, "citations.brackets") == ["Citation in round brackets: (1).",
                                                     "Citation in round brackets: (2, 3)",
                                                     "Citation in round brackets: (4)."]


def test_round_citations_without_punctuation_count_for_renumbering(tmp_path):
    # JEHP sets citations in superscript: "(1)." becomes a deletion plus an insertion, "(2, 3)" one change;
    # both are citations for renumbering and for the queries.
    st = _style()
    path = _doc(tmp_path, ["Common (1). Treated (3, 4)  Then drops (2). Rare (5)."], 5)
    ms = docmodel.load(path)
    assert [o.numbers for o in renumber.occurrences(ms, st, 5)] == [[1], [3, 4], [2], [5]]
    pl = renumber.plan(ms, st)
    assert pl.apply and pl.mapping == {1: 1, 3: 2, 4: 3, 2: 4, 5: 5}
    assert not [q for q in queries.collect(ms, st).queries if q.check == "references.uncited"]


def test_round_numbers_out_of_order_next_to_square_citations_are_counts(tmp_path):
    st = _style()
    path = _doc(tmp_path, ["Alpha [1]. Beta [2]. Klebsiella (3) grew. Gamma [3]. Then (2) eyes. Delta (4) [5]."], 5)
    ms = docmodel.load(path)
    assert [m.group(2) for _, m in rules.round_citations(ms, st)] == ["4"]


def test_number_after_a_decimal_is_a_citation_only_when_next_in_sequence(tmp_path):
    st = _style()
    path = _doc(tmp_path, ["Reliability was shown before (1). Cronbach's alpha was 0.88 (2). Mean score was 3.2 (4). "
                           "Later work agreed (3). Others too (4, 5)."], 5)
    ms = docmodel.load(path)
    found = sorted(rules.round_citations(ms, st), key=lambda x: x[1].start())
    assert [m.group(2) for _, m in found] == ["1", "2", "3", "4, 5"]


def test_superscript_exponent_and_zero_are_not_citations(tmp_path):
    st = _style()
    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph("The change was significant (P = 4.57 × 10⁻")
    p.add_run("8").font.superscript = True
    p.add_run(") in all groups.")
    p2 = doc.add_paragraph("The disease is common worldwide.")
    p2.add_run("1").font.superscript = True
    for t in _refs(2):
        doc.add_paragraph(t)
    doc.save(tmp_path / "s.docx")
    ms = docmodel.load(str(tmp_path / "s.docx"))
    assert _converted(ms, st) == ["[1]"]
    assert rules.normalise_citation("0-2", st) is None and rules.normalise_citation("0", st) is None


def test_statistics_in_square_brackets_and_citations_beyond_the_list(tmp_path):
    st = _style()
    path = _doc(tmp_path, ["Groups differed (F[4,164] = 4.7; P = 0.001) as known.[1] Others agree.[2] More.[6]"], 5)
    ms = docmodel.load(path)
    assert [o.numbers for o in renumber.occurrences(ms, st, 5)] == [[1], [2], [6]]
    assert _issues(ms, st, "references.exists") == ["Citation [6] has no entry in the reference list"]
    assert not renumber.plan(ms, st).apply  # a citation with no entry: queried, nothing renumbered
    assert [q.text for q in queries.collect(ms, st).queries if q.check == "references.missing_entry"] == [
        "Citation [6] has no entry in the reference list. Kindly provide the reference."]


def test_uncited_query_ignores_issue_numbers_in_the_reference_list(tmp_path):
    st = _style()
    doc = Document()
    for t in ["Introduction", "The disease is common.[1]", "References",
              "1. Smith A. A study. J Test 2001;12(2):1-9.", "2. Jones B. Another. J Test 2002;13(4):1-9."]:
        doc.add_paragraph(t)
    doc.save(tmp_path / "u.docx")
    ms = docmodel.load(str(tmp_path / "u.docx"))
    assert [q.text for q in queries.collect(ms, st).queries if q.check == "references.uncited"] == [
        "Reference no. 2 has not been cited. Kindly cite."]
    assert _issues(ms, st, "references.cited") == ["Reference 2 is not cited (author query needed)"]


def test_jehp_conflict_heading_variant_passes_and_missing_sections_are_added(tmp_path):
    st = _style()
    body = ["Discussion", "The disease is common in older patients treated in this clinic.[1]"]
    kept = _doc(tmp_path, body + ["Financial support and sponsorship", "Nil.", "Conflict of interest",
                                  "The authors declare no conflict of interest."], 1, "kept.docx")
    assert _issues(docmodel.load(kept), st, "sections.required") == []
    bare = _doc(tmp_path, body, 1, "bare.docx")
    assert _issues(docmodel.load(bare), st, "sections.required") == [
        "Missing section: Financial support and sponsorship", "Missing section: Conflicts of interest"]
    cs = rules.run_all(docmodel.load(bare), st, only=["heading_tags", "sections_back_matter"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(tmp_path / "bare_out.docx")
    tracked.render(docmodel.load(bare), cs, tracked=False).save(out)
    assert _issues(docmodel.load(out), st, "sections.required") == []


def test_form_markers_and_rule_lines_are_not_reference_entries(tmp_path):
    doc = Document()
    for t in ["Introduction", "Text.[1]", "References", "1. Smith A. A study. J Test 2001;12:1-9.",
              "2. Jones B. Another. J Test 2002;13:1-9.", "Bottom of Form", "_______________"]:
        doc.add_paragraph(t)
    entry = doc.add_paragraph("3. Brown C. Third. J Test 2003;14:1-9.")
    entry.style = doc.styles["Heading 1"]  # an entry the author styled as a heading is still an entry
    doc.save(tmp_path / "r.docx")
    refs = references.extract(docmodel.load(str(tmp_path / "r.docx")))
    assert [r.number for r in refs] == [1, 2, 3]


def test_manuscript_preferred_over_figure_file_named_after_it(tmp_path):
    zp = tmp_path / "s.zip"
    with zipfile.ZipFile(zp, "w") as z:
        z.writestr("Figure of manuscript file.docx", b"x" * 100)
        z.writestr("Manuscript.docx", b"x")
        z.writestr("FirstPage.docx", b"x")
    assert pipeline.manuscript_member(zipfile.ZipFile(zp)) == "Manuscript.docx"


def test_statistics_spacing_not_checked_in_supplementary_material(tmp_path):
    # S372 / G151: the supplementary material is retained as it is; the body is still checked
    body = ["Patients were grouped (n=12) by block type.", "<H2>Supplementary Material", "ANNEXURE",
            "Appendix Table 1: Block cases (n=7)", "<H1>Acknowledgement", "Staff counted (n=3) cases."]
    ms = docmodel.load(_doc(tmp_path, body))
    got = _issues(ms, _style("ijo"), "statistics.spacing")
    assert len(got) == 2 and not any("n=7" in m for m in got)
