"""Citation next to the author's name, merged citations, the year-superscript guard, legend part
labels and sub-heading tags. Synthetic text only."""

from docx import Document

from edit_os import docmodel, rules, rules_sections, style, tracked
from edit_os.changes import Status


def _style(code="ijo"):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


def _refs(*first_authors):
    return ["References"] + [f"{i}. {a}. A title {i}. J Test 2001;{i}:1-2." for i, a in enumerate(first_authors, 1)]


def _save(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        if isinstance(t, tuple):  # (text, {"bold": True}) -> one formatted run
            run = doc.add_paragraph().add_run(t[0])
            for k, v in t[1].items():
                setattr(run.font, k, v)
        else:
            doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def _run(path, st, tmp_path, only):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")
    return cs, [p.text for p in docmodel.load(str(out)).paragraphs]


CITE_RULES = ["author_citations", "renumber", "citations", "formatting", "guide_replace"]


def test_citation_moves_next_to_matching_author(tmp_path):
    body = [
        "Introduction",
        "Ames et al. showed the drug works in rats. They found more in mice [1].",
        "In a study by Brook et al., values fell (P = 0.02) [2]. Later work is new.",
        "The report by Dunn et al in 2013 was rare.[3] Evans and Frost described two cases.[4]",
        "Gray et al. found it [5].",                             # reference 5 is by Hall: stays
        "Ames et al. and Brook et al. both agree on this point [1].",  # two names: stays
        "Hall et al. [5] reported the same in a later cohort.",  # already at the name: stays
        "Themes were coded as described by Evans and Frost.[4] NVivo was used.",  # only "." between: stays
    ]
    path = _save(tmp_path, body + _refs("Ames JB, Kay T", "Brook A, Lee B", "Dunn C", "Evans D, Frost E", "Hall F"))
    _, t = _run(path, _style(), tmp_path, CITE_RULES)
    assert t[7] == "Themes were coded as described by Evans and Frost.[4] NVivo was used."
    assert t[1] == "Ames et al.[1] showed the drug works in rats. They found more in mice."
    assert t[2] == "In a study by Brook et al.,[2] values fell (P = 0.02). Later work is new."
    assert t[3] == "The report by Dunn et al.[3] in 2013 was rare. Evans and Frost[4] described two cases."
    assert t[4] == "Gray et al. found it.[5]"
    assert t[5] == "Ames et al. and Brook et al. both agree on this point.[1]"
    assert t[6] == "Hall et al.[5] reported the same in a later cohort."


def test_moved_citation_is_superscript_and_renumbered(tmp_path):
    body = [
        "Introduction",
        "Glaucoma is common [2]. Özer et al. reported a rise in density in the macula [1].",
    ]
    path = _save(tmp_path, body + _refs("Özer Ö, Baysal Z", "Weinreb RN"))
    cs, t = _run(path, _style(), tmp_path, CITE_RULES)
    # Renumbered by first citation: [2] -> [1], Özer's [1] -> [2], placed after the name.
    assert t[1] == "Glaucoma is common.[1] Özer et al.[2] reported a rise in density in the macula."
    moved = [c for c in cs.changes if c.reason == "author_citations" and c.new == "[2]"]
    assert moved and moved[0].format == {"superscript": True}


def test_author_citations_off_and_sentence_limit(tmp_path):
    body = ["Introduction", "Ames et al. said A. B was seen. C was seen. D was seen. E was seen [1]."]
    path = _save(tmp_path, body + _refs("Ames JB"))
    _, t = _run(path, _style(), tmp_path, CITE_RULES)
    assert t[1] == "Ames et al. said A. B was seen. C was seen. D was seen. E was seen.[1]"  # too far
    st = _style()
    st.settings["citations"]["next_to_author"] = False
    path = _save(tmp_path, ["Introduction", "Ames et al. said it [1]."] + _refs("Ames JB"), "b.docx")
    _, t = _run(path, st, tmp_path, CITE_RULES)
    assert t[1] == "Ames et al. said it.[1]"


def test_adjacent_citations_merge(tmp_path):
    body = ["Introduction", "It needs work. [2], [3], [5] Later [1][4] and (6), (7) here."]
    path = _save(tmp_path, body + _refs(*[f"Auth{c} A" for c in "ABCDEFG"]))
    _, t = _run(path, _style(), tmp_path, ["citations"])
    assert t[1] == "It needs work.[2,3,5] Later[1,4] and[6,7] here."


def test_year_before_superscript_is_a_citation(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    p = cell.paragraphs[0]
    p.add_run("Choudhary et al., Vision 2025")
    p.add_run("2").font.superscript = True
    q = doc.add_paragraph("Area of 3 mm")
    q.add_run("2").font.superscript = True
    path = tmp_path / "y.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    cellp = next(x for x in ms.paragraphs if x.in_table)
    assert rules.superscript_citations(cellp) == [(len("Choudhary et al., Vision 2025"), len(cellp.text), "2")]
    body = next(x for x in ms.paragraphs if x.text.startswith("Area"))
    assert rules.superscript_citations(body) == []  # mm² stays a power


def test_legend_part_labels(tmp_path):
    paras = [
        "Figure 1: Photograph showing the chamber (1A). B-scan shows echoes (1B).",
        "Figure 2: Fundus. (A, B) Discs in both eyes. (C) Imaging.",
        "Figure 3: Curves for a) diameter b) height c) index",
        "Figure 4: Cross section (2A) of the (R) eye",
    ]
    _, t = _run(_save(tmp_path, paras), _style(), tmp_path, ["fig_captions"])
    assert t[0] == "Figure 1: Photograph showing the chamber (a). B-scan shows echoes (b)"
    assert t[1] == "Figure 2: Fundus. (a, b) Discs in both eyes. (c) Imaging"
    assert t[2] == "Figure 3: Curves for (a) diameter (b) height (c) index"
    assert t[3] == "Figure 4: Cross section (2A) of the (R) eye"  # another figure's number; not a part label
    st = _style()
    st.settings["figures"]["legend_part_labels"] = False
    _, t = _run(_save(tmp_path, paras, "b.docx"), st, tmp_path, ["fig_captions"])
    assert t[1] == "Figure 2: Fundus. (A, B) Discs in both eyes. (C) Imaging"  # off: as typed


LONG = "The patients were examined in the clinic and their records were reviewed in full detail. " * 3


def test_subheadings_tagged(tmp_path):
    paras = [
        "Introduction", LONG,
        "Methods",
        ("Study design", {"bold": True}), LONG,
        ("Outcome measures", {"italic": True}), LONG,
        "Statistical analysis", LONG,                  # plain, own line, running text after it
        ("Inclusion criteria:", {"bold": True}), "Adults over 18 years", "No prior surgery",  # a list follows
        ("Funding", {"bold": True}), LONG,             # a statement, not a sub-heading
        "OCT: optical coherence tomography", LONG,     # an abbreviation line
        "Results", LONG,
        ("Data collection:", {"bold": True}), LONG,
    ] + _refs("Ames JB")
    path = _save(tmp_path, paras)
    _, t = _run(path, _style(), tmp_path, ["subheadings"])
    assert t[3] == "<H2>Study design"
    assert t[5] == "<H3>Outcome measures"  # italic under bold sub-headings
    assert t[7] == "<H2>Statistical analysis"
    assert t[9] == "Inclusion criteria:"
    assert t[12] == "Funding"
    assert t[14] == "OCT: optical coherence tomography"
    assert t[18] == "<H2>Data collection:"  # G137: punctuation at the end of a heading is retained


def test_subheadings_skip_run_in_section_and_abstract(tmp_path):
    paras = [
        "Abstract", ("Background", {"bold": True}), LONG,
        "Case Report",
        "Ocular examination: The right eye had no perception of light and the left was normal. " * 2,
        "Histopathology: The cut section showed areas of necrosis and many small round cells. " * 2,
        ("Staging and treatment:", {"bold": True}), LONG,
    ] + _refs("Ames JB")
    path = _save(tmp_path, paras)
    ms = docmodel.load(path)
    assert rules_sections.rule_subheadings(ms, _style()) == []


def test_citation_at_paragraph_end_takes_trailing_space(tmp_path):
    body = ["Introduction", "It is the main target in glaucoma [1]. ", "Ames et al. showed it in rats.[2]  ",
            "Caption-free text with a citation [1] in the middle. "]
    path = _save(tmp_path, body + _refs("Kay T", "Ames JB"))
    _, t = _run(path, _style(), tmp_path, CITE_RULES)
    assert t[1] == "It is the main target in glaucoma.[1]"   # not "glaucoma.[1] "
    assert t[2] == "Ames et al.[2] showed it in rats."
    assert t[3] == "Caption-free text with a citation[1] in the middle."  # the guide's paragraph-end spaces go (G143)


def test_caption_end_period_takes_trailing_space(tmp_path):
    _, t = _run(_save(tmp_path, ["Figure 1: Fundus photograph showing folds. "]), _style(), tmp_path, ["fig_captions"])
    assert t[0] == "Figure 1: Fundus photograph showing folds"


def test_heading_tags_replace_leading_spaces(tmp_path):
    paras = [
        "Introduction", LONG,
        " Methods:  ",
        (" Study design", {"bold": True}), LONG,
        "  Results", LONG,
    ] + _refs("Ames JB")
    path = _save(tmp_path, paras)
    _, t = _run(path, _style(), tmp_path, ["heading_tags", "sections_headings", "subheadings"])
    # IJO drops "Introduction" (editors); H1 in Title case (S: "H1 | Caps", owner reading)
    assert t[:5] == [LONG, "<H1>Methods:", "<H2>Study design", LONG, "<H1>Results"]  # G137 colon kept, G143 spaces go


def test_caption_end_period_kept_when_set_to_keep(tmp_path):
    st = _style()
    st.settings["figures"]["caption_end_period_keep"] = True
    _, t = _run(_save(tmp_path, ["Table 1. Characteristics of participants."]), st, tmp_path, ["fig_captions"])
    assert t[0] == "Table 1: Characteristics of participants."


def _table_ms(tmp_path, body, cells, notes, name="t.docx"):
    """A manuscript: body paragraphs, one table (one cell per entry: (text, superscript digits)), notes
    under it and a reference list of five."""
    doc = Document()
    doc.add_paragraph("Introduction")
    for t in body:
        if isinstance(t, tuple):
            q = doc.add_paragraph(t[0])
            q.add_run(t[1]).font.superscript = True
        else:
            doc.add_paragraph(t)
    table = doc.add_table(rows=len(cells), cols=1)
    for row, (text, sup) in zip(table.rows, cells):
        p = row.cells[0].paragraphs[0]
        p.add_run(text)
        if sup:
            p.add_run(sup).font.superscript = True
    for t in notes:
        doc.add_paragraph(t)
    for t in _refs(*[f"Auth{c} A" for c in "ABCDE"]):
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return docmodel.load(str(path))


def _table_conversions(ms):
    ids = {p.id for p in ms.paragraphs if p.in_table}
    return sorted(ms.by_id(c.para_id).text[:c.start] + c.new for c in rules.rule_citations(ms, _style())
                  if c.para_id in ids and "[" in c.new)


def test_table_footnote_markers_are_not_citations(tmp_path):
    # gtsummary-style table: the notes under it explain ¹ ² ³; "Garg et al., 2018⁴" cites.
    ms = _table_ms(tmp_path, ["Lymphocytes matter [1]. More [2,3]."],
                   [("Characteristic", ""), ("N = 4391 (100%)", "2"), ("N = 968 (19%)", "2"), ("p value", "3"),
                    ("Diabetic retinopathy", "2"), ("As in Garg et al., 2018", "4")],
                   ["¹ N not missing (unweighted) ² n (%) ³ Pearson's Chi-squared test"])
    assert _table_conversions(ms) == ["As in Garg et al., 2018[4]"]


def test_table_footnote_markers_explained_in_a_cell(tmp_path):
    # Notes typed in the table's last cell, with plain digits: "1: DSM-5 diagnostic criteria, 2: in months."
    ms = _table_ms(tmp_path, ["Gambling matters [1]."],
                   [("Tobacco dependence", "1"), ("Tobacco use duration", "2"), ("Prior work", "5"),
                    ("1: DSM-5 diagnostic criteria, 2: in months.", "")], [])
    assert _table_conversions(ms) == ["Prior work[5]"]


def test_table_marker_repeated_across_cells(tmp_path):
    cells = [("Not reported", "3"), ("12", ""), ("Not reported", "3"), ("Not reported", "3"), ("Smith", "4")]
    # The text cites in square brackets: "Not reported³" down a column is a footnote marker.
    ms = _table_ms(tmp_path, ["Seizures matter [1]. More [2]."], cells, [])
    assert _table_conversions(ms) == ["Smith[4]"]
    # The text cites in superscript: the repeated number may be a citation, so it converts.
    ms = _table_ms(tmp_path, [("Seizures matter.", "1"), ("More.", "2")], cells, [], name="u.docx")
    assert _table_conversions(ms) == ["Not reported[3]"] * 3 + ["Smith[4]"]


def test_text_after_a_table_still_cites(tmp_path):
    # Running text right under a table is no note: its "risk ² was" does not explain the cell's "Smith²".
    ms = _table_ms(tmp_path, ["Risk matters [1]."], [("Smith", "2")], ["The risk ² was higher in men than women."])
    assert _table_conversions(ms) == ["Smith[2]"]


def _sup_paragraphs(tmp_path, items):
    """One paragraph per item: a list of (text, superscript?) runs."""
    doc = Document()
    for runs in items:
        p = doc.add_paragraph()
        for text, sup in runs:
            p.add_run(text).font.superscript = sup
    path = tmp_path / "s.docx"
    doc.save(path)
    return [p for p in docmodel.load(str(path)).paragraphs if p.text]


def test_statistics_units_and_formulas_are_not_superscript_citations(tmp_path):
    # Statistic symbols and names, units and formulas take an index or a power, not a citation.
    cases = [
        [("ΔΧ²(10)", False)], [("Chi", False), ("2", True)], [("X", False), ("2", True), (", df", False)],
        [("X", False), ("2", True), ("(p-value)", False)], [("Cohen's f² effect sizes", False)],
        [("f²", False)], [("using Q²predict, RMSE", False)], [("the 𝝌2‑test and I2 statistics", False)],
        [("PSAD in ng/mL", False), ("2", True), (", median", False)], [("P within", False), ("1", True)],
        [("BMI >35 kg m-²; any", False)], [("cutoff >46 kΩ⁻¹, yielding", False)], [("a BSA of 1.73 m", False), ("2", True)],
        [("n= (1.96)\xa0", False), ("2", True), (" p q/ d", False), ("2", True)], [("n = N/(1 + Ne²), where", False)],
        [("an eta", False), ("2", True), (" of 0.1", False)], [("150 × 10³ cells", False)],
    ]
    for p in _sup_paragraphs(tmp_path, cases):
        assert rules.superscript_citations(p) == [], p.text


def test_real_superscript_citations_next_to_numbers_and_units_still_convert(tmp_path):
    cases = {
        "Smith et al.²": [("Smith et al.", False), ("2", True)],
        "irradiation at 365 nm.³": [("irradiation at 365 nm.³", False)],
        "Li et al. (2024)⁴": [("Li et al. (2024)", False), ("4", True)],
        "mean 35 (1.69) 12,28": [("mean 35 (1.69) ", False), ("12,28", True)],
        "reported worldwide²": [("reported worldwide²", False)],
        "threshold of 4 ng/mL.3,4": [("threshold of 4 ng/mL.", False), ("3,4", True)],
    }
    paras = _sup_paragraphs(tmp_path, cases.values())
    for (label, _), p in zip(cases.items(), paras):
        assert len(rules.superscript_citations(p)) == 1, label
