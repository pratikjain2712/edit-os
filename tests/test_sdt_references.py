"""A reference list inside a body-level content control (w:body/w:sdt/w:sdtContent/w:p, as Zotero / Mendeley write
their bibliography block): read, labelled and rendered as a plain list is, and read by the evaluation tools in
document order. Invented text throughout."""

import copy
import io
import os
import sys
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))

import eval_final as E  # noqa: E402
from evaluate import tracked_pairs  # noqa: E402

from edit_os import biblio, docmodel, references, refengine, rules, style, tracked  # noqa: E402
from edit_os.changes import Status  # noqa: E402

REFS = ["1.\tLee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4.",
        "2.Kim J, Park H. Neck pain in office workers. Spine 1998;23:10-4."]


def _doc(tmp_path, sdt: bool) -> str:
    doc = Document()
    for t in ["Introduction", "Back and neck pain are common [1,2].", "References"] + REFS + ["Closing note."]:
        p = doc.add_paragraph()
        parts = t.split("\t")
        for k, part in enumerate(parts):
            r = p.add_run(part)
            if k < len(parts) - 1:
                r.add_tab()
    if sdt:  # move the reference entries into one content control, where Word's citation managers put them
        body = doc.element.body
        ps = [p for p in body.findall(qn("w:p")) if p.xpath("string(.)").startswith(("1.", "2."))]
        box = etree.SubElement(body, qn("w:sdt"))
        pr = etree.SubElement(box, qn("w:sdtPr"))
        etree.SubElement(pr, qn("w:docPartObj"))
        content = etree.SubElement(box, qn("w:sdtContent"))
        ps[0].addprevious(box)
        for p in ps:
            content.append(p)
    path = tmp_path / ("sdt.docx" if sdt else "plain.docx")
    doc.save(path)
    return str(path)


def _ijo():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    return st


@pytest.mark.parametrize("tracked_mode", [True, False])
def test_labels_in_a_content_control_as_in_a_plain_list(tmp_path, tracked_mode):
    got = {}
    for sdt in (False, True):
        path = _doc(tmp_path, sdt)
        ms = docmodel.load(path)
        if sdt:
            assert [p.text for p in ms.paragraphs if p.text[:1].isdigit()] == REFS
            assert [r.number for r in references.extract(ms)] == [1, 2]
        cs = rules.run_all(ms, _ijo(), only=["references"])
        labels = [(c.old, c.new) for c in cs.changes if c.reason == "ref_label"]
        for c in cs.changes:
            c.status = Status.ACCEPTED
        out = str(tmp_path / f"out_{sdt}_{tracked_mode}.docx")
        tracked.render(docmodel.load(path), cs, tracked=tracked_mode).save(out)
        texts = [a for _, a in tracked_pairs(open(out, "rb").read()) if a[:1].isdigit()]
        got[sdt] = (labels, texts)
    assert got[True] == got[False]
    labels, texts = got[True]
    assert labels == [("", " ")]  # "2.Kim J" -> "2. Kim J"; the tab after "1." is the list's layout and stays
    assert texts[0].startswith("1.\tLee K") and texts[1].startswith("2. Kim J")


def test_tools_read_content_control_paragraphs_in_order(tmp_path):
    path = _doc(tmp_path, sdt=True)
    raw = open(path, "rb").read()
    assert E.final_text(raw)[2:] == ["References", E.norm(REFS[0]), E.norm(REFS[1]), "Closing note."]
    assert E.plain_paras(path) == E.final_text(raw)
    assert E.split_refs(E.plain_paras(path))["references"][1:3] == [E.norm(REFS[0]), E.norm(REFS[1])]


def test_split_refs_takes_the_list_not_a_reviewer_reply():
    article = ["Intro text.", "References", "1. Lee K, Rao S. Pain. Spine 1997;22:1-4.",
               "2. Kim J, Park H. Neck. Spine 1998;23:10-4.", "3. Cho M. Back. Spine 1999;24:5-9.",
               "Reply to the reviewers", "We added two references.", "References:",
               "Lee K, Rao S. Pain. Spine 1997;22:1-4.", "Thank you."]
    sp = E.split_refs(article)
    assert sp["references"][0] == "References" and len(sp["references"]) == len(article) - 1
    one = ["Text.", "References", "1. Lee K. Pain. Spine 1997;22:1-4.", "Figure 1: A chart"]
    assert E.split_refs(one)["references"] == one[1:3]


# --- footnote markers in a record's title ------------------------------------------------------------------

def test_record_title_drops_trailing_footnote_markers():
    r = biblio.Record(source="pubmed", title="Training residents on an animal eye with an induced lens opacity(2)(2).")
    assert r.title == "Training residents on an animal eye with an induced lens opacity."
    assert biblio.Record(source="crossref", title="A model of lens opacity[3]").title == "A model of lens opacity"
    # subscripts the databases write the same way stay, as does a marker in the author's own fields
    for t in ["Serum levels of vitamin D(3).", "Retention of CO(2)", "Outcomes, part (2).", "A study(2) of eyes."]:
        assert biblio.Record(source="pubmed", title=t).title == t
    assert biblio.Record(source="author", title="An eye model(2).").title == "An eye model(2)."


def test_styled_entry_has_no_doubled_marker():
    r = biblio.Record(source="pubmed", pmid="9", title="Training residents on an animal eye with an induced lens "
                      "opacity(2)(2).", authors=[("Lee", "K"), ("Rao", "S")], journal="Eye Train J", year="2001",
                      volume="27", issue="6", pages="814-6")
    out = refengine.format_record(copy.deepcopy(r), _ijo())
    assert "opacity. Eye Train J" in out and "(2)" not in out


def test_docx_bytes_roundtrip_keeps_sdt(tmp_path):
    # the content control itself survives a render (nothing flattens it)
    path = _doc(tmp_path, sdt=True)
    cs = rules.run_all(docmodel.load(path), _ijo(), only=["references"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    buf = io.BytesIO()
    tracked.render(docmodel.load(path), cs, tracked=False).document.save(buf)
    xml = zipfile.ZipFile(buf).read("word/document.xml")
    assert b"<w:sdtContent>" in xml
