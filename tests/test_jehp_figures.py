"""JEHP figures and tables (synthetic text): tables moved to the end after the figure legends, column heads and
legends in sentence case, table-note abbreviations "term = definition", a note typed into a caption split out,
caption label typos, the figure / table queries."""

from docx import Document

from edit_os import docmodel, legends, queries, rules_figures, style
from edit_os.changes import Status

from test_legends import REFS, W, apply, make

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail.")


def jehp(**figures):
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    st.settings["figures"].update(figures)
    return st


def _render(path, st, tmp_path, tracked):
    ms = docmodel.load(path)
    moves, notes = legends.plan(ms, st)
    for m in moves:
        m.status = Status.ACCEPTED
    ms = docmodel.load(path)
    legends.render(ms, moves, (Status.ACCEPTED,), tracked=tracked, author="Edit OS", date="2026-10-26T00:00:00Z")
    ms.save(tmp_path / "t.docx")
    return moves, notes, Document(tmp_path / "t.docx").element.body


def _blocks(body):
    """Body text as it reads with the tracked changes accepted: paragraphs and "TABLE: <first row>"."""
    out = []
    for el in body:
        if el.tag == W + "p":
            if el.find(f"{W}pPr/{W}rPr/{W}del") is not None:
                continue
            t = "".join(x.text or "" for x in el.iter(W + "t"))
            if t.strip():
                out.append(t)
        elif el.tag == W + "tbl":
            rows = [r for r in el.findall(W + "tr") if r.find(f"{W}trPr/{W}del") is None]
            if rows:
                out.append("TABLE: " + "".join(x.text or "" for x in rows[0].iter(W + "t")))
    return out


def test_tables_move_to_the_end_after_the_legends(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " See Table 1 and Figure 1.",
                           "Table 1: Baseline data", ("table", [["Age", "Sex"], ["1", "2"]]), "SD = standard deviation",
                           BODY, ("table", [["No", "caption"], ["3", "4"]]), "Figure 1: Study flow", BODY, *REFS])
    for tracked in (False, True):
        moves, notes, body = _render(path, jehp(), tmp_path, tracked)
        assert {m.kind for m in moves} == {"legends", "table_caption", "tables"}
        got = _blocks(body)
        i = next(k for k, t in enumerate(got) if "Lee K" in t)
        assert got[i + 1:] == ["Figure Legend", "Figure 1: Study flow", "Tables",
                               "TABLE: Table 1: Baseline data", "SD = standard deviation"]
        assert "TABLE: Nocaption" in got[:i]  # a table with no caption stays where it is
    d = next(m for m in moves if m.kind == "tables").to_dict()
    assert legends.BlockMove.from_dict(d).to_dict() == d


def test_table_caption_typed_over_the_text_that_discusses_it_moves_with_the_table(tmp_path):
    # "Table 1. ..." / two paragraphs on the results / the table: the caption goes with the table, the text stays
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.", "Table 1. Scores by school and year group",
                           "Table 1 presents the scores by school. " + BODY, BODY,
                           ("table", [["School", "Score"], ["A", "2"]]), "Discussion", BODY, *REFS])
    _, _, body = _render(path, jehp(), tmp_path, False)
    got = _blocks(body)
    i = next(k for k, t in enumerate(got) if "Lee K" in t)
    assert got[i + 1:] == ["Tables", "Table 1. Scores by school and year group", "TABLE: SchoolScore"]
    assert got[2].startswith("Table 1 presents the scores") and got[3] == BODY


def test_table_captioned_twice_keeps_the_caption_that_moves(tmp_path):
    from edit_os import rules
    from edit_os.changes import Status as St

    path = make(tmp_path, ["Introduction", BODY + " See Tables 1 and 2.", "Table 1: Scores by School",
                           ("table", [["School", "Score"], ["A", "2"]]),
                           "Table 1. Scores by school and year group, with the mean of each class.",
                           BODY, "Table 2. Ages of the pupils", ("table", [["Age", "N"], ["9", "4"]]),
                           "Table 2: Ages of the pupils", BODY, *REFS])
    st = jehp()
    ms = docmodel.load(path)
    dups = legends.duplicate_table_captions(ms, st)
    assert [(d.text[:9], k.text[:9]) for d, k, _ in dups] == [("Table 1. ", "Table 1: "), ("Table 2: ", "Table 2. ")]
    cs = rules.run_all(docmodel.load(path), st, only=["table_caption_duplicate", "fig_captions"])
    gone = [c for c in cs.changes if c.reason == legends.DUPLICATE_CAPTION]
    assert [c.old for c in gone] == [d.text for d, _, _ in dups] and all(c.new == "" for c in gone)
    assert not [c for c in cs.changes if c.reason != legends.DUPLICATE_CAPTION and c.para_id in {g.para_id for g in gone}]
    qs = queries.QuerySet()
    legends.add_queries(docmodel.load(path), st, qs)
    # only the one that said more than the caption kept is queried, on the caption kept
    assert len(qs.queries) == 1 and "Table 1 had a second caption" in qs.queries[0].text
    assert "year group, with the mean of each class”" in qs.queries[0].text
    assert legends.duplicate_table_captions(docmodel.load(path), style.load("base")) == []  # tables not moved


def test_tables_heading_not_doubled(tmp_path):
    path = make(tmp_path, [BODY + " See Table 1 and Table 2.", "Table 1: Baseline data",
                           ("table", [["Age", "Sex"], ["1", "2"]]), BODY, *REFS, "Tables", "Table 2: Outcomes",
                           ("table", [["A", "B"], ["1", "2"]])])
    _, _, body = _render(path, jehp(), tmp_path, False)
    got = _blocks(body)
    assert got.count("Tables") == 1
    assert got[-2:] == ["TABLE: Table 1: Baseline data", "TABLE: Table 2: Outcomes"]


def test_column_heads_in_sentence_case(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " Patients in Iran were seen. See Table 1.", "Table 1: Data",
                           ("table", [["Study Design", "Control Group (n = 30)", "Iran Cohort", "BMI Value"],
                                      ["Cohort", "12", "3", "4"]]),
                           ("table", [["STUDY DESIGN", "P VALUE"], ["x", "y"]])])
    out, _ = apply(path, jehp(), tmp_path)
    assert "Study design" in out and "Control group (n = 30)" in out and "Iran cohort" in out
    assert "BMI value" in out
    assert "Study Design" in out  # ALL CAPS -> Title Case (S: "In case the content is in All Caps ...")


def test_figure_legend_sentence_case_and_first_word(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " Schools in Iran took part. See Figure 1 and Figure 2.",
                           "Figure 1: Study Flow Of Participants In Iran", "Figure 2: depicts the trend. Counts rose"])
    out, _ = apply(path, jehp(), tmp_path)
    assert "Figure 1: Study flow of participants in Iran." in out  # S: captions end with a period (owner, 2026-10-26)
    assert any(t.startswith("Figure 2: Depicts the trend. Counts rose") for t in out)


def test_table_note_abbreviations(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.", "Table 1: Data",
                           ("table", [["Age", "Sex"], ["1", "2"]]),
                           "Values are mean ± SD. P < 0.05. df, degrees of freedom; SD, standard deviation.",
                           BODY])
    out, _ = apply(path, jehp(), tmp_path)
    assert "Values are mean ± SD. P < 0.05. df = degrees of freedom, SD = standard deviation" in out


def test_note_in_a_caption_goes_under_the_table(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.",
                           "Table 1. Baseline characteristics (N=100). IQR, interquartile range; SD, standard deviation.",
                           ("table", [["Age", "Sex"], ["1", "2"]]), "", BODY])
    out, _ = apply(path, jehp(), tmp_path)
    assert "Table 1: Baseline characteristics (N=100)." in out
    assert "IQR = interquartile range, SD = standard deviation" in out
    ms = docmodel.load(str(tmp_path / "out.docx"))
    i = [p.text for p in ms.paragraphs].index("IQR = interquartile range, SD = standard deviation")
    assert ms.paragraphs[i - 1].in_table  # under the table
    # Caption below its table: the note follows the caption.
    path = make(tmp_path, ["Introduction", BODY + " See Table 1.", ("table", [["Age", "Sex"], ["1", "2"]]),
                           "Table 1. Paired changes. Wilcoxon tests were used.", BODY], "b.docx")
    out, _ = apply(path, jehp(), tmp_path)
    assert "Table 1: Paired changes." in out and "Wilcoxon tests were used" in out


def test_caption_label_typos_and_queries(tmp_path):
    path = make(tmp_path, ["Introduction",
                           BODY + " See Table 1, Supplimentary Table 2 and Figures 1 and 2.",
                           "Tabel 1 Baseline data", ("table", [["Age", "Sex"], ["1", "2"]]), BODY])
    ms = docmodel.load(path)
    assert [(c.kind, c.number) for c in rules_figures.captions(ms)] == [("Table", 1)]
    assert [c.numbers for c in rules_figures.citations(ms) if c.kind == "Table"] == [[1]]
    qs = queries.QuerySet()
    rules_figures.add_queries(ms, jehp(), qs)
    texts = [q.text for q in qs.queries]
    assert "Kindly provide the captions for Figures 1–2." in texts
    assert not any("Table 2" in t for t in texts)
    out, _ = apply(path, jehp(), tmp_path)
    assert "Table 1: Baseline data." in out


def test_caption_on_a_later_line_or_in_a_heading_style(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph(BODY + " See Table 1 and Table 2.")
    doc.add_paragraph("Table 1. Baseline data", style="Heading 3")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Age"
    p = doc.add_paragraph(BODY)
    p.add_run().add_break()
    p.add_run("Table 2: Outcomes by group")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Score"
    doc.save(tmp_path / "h.docx")
    ms = docmodel.load(str(tmp_path / "h.docx"))
    assert [(c.kind, c.number) for c in rules_figures.captions(ms)] == [("Table", 1)]
    qs = queries.QuerySet()
    rules_figures.add_queries(ms, jehp(), qs)
    assert not any("not provided" in q.text for q in qs.queries)
