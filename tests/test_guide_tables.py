"""The pre-editing guide's TABLE rules for IJO, JEHP and MJDRDYPU (owner, 2026-10-27: "Follow style guide for p value,
n value and table formatting"): cell text 12 pt and left aligned, whole caption bold, Roman labels to Arabic, column
heads, stub indentation as typed spaces, notes "term = definition" with the abbreviations first, empty cell paragraphs
deleted, queries for what Word cannot track. Synthetic text only."""


from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from edit_os import docmodel, header, legends, queries, rules, rules_tables, style, tracked
from edit_os.changes import Status

JOURNALS = ("ijo", "jehp", "mjdrdypu")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _st(code):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


def _doc(tmp_path, build, name="m.docx"):
    doc = Document()
    build(doc)
    path = str(tmp_path / name)
    doc.save(path)
    return path


def _table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    return t


def _render(path, cs, out, track=False):
    for c in cs.changes:
        c.status = Status.ACCEPTED
    tracked.render(docmodel.load(path), cs, tracked=track).save(out)
    return out


def _texts(path):
    return [p.text for p in docmodel.load(path).paragraphs]


# --- G193 / G194: 12 pt, aligned left -------------------------------------------------------------

def test_cells_take_12_pt_and_left_alignment(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        t = _table(doc, [["Group", "Mean"], ["A", "1.2"]])
        for cell in (t.cell(0, 1), t.cell(1, 1)):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].runs[0].font.size = Pt(9)
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        cs = rules.run_all(docmodel.load(path), _st(code), only=["format_fonts"])
        by_text = {}
        for c in cs.changes:
            by_text.setdefault(c.old, {}).update(c.format)
        assert by_text["Mean"].get("align") == "left" and by_text["1.2"].get("align") == "left"
        assert by_text["Mean"]["size"] == 12 and by_text["1.2"]["size"] == 12
        out = _render(path, cs, str(tmp_path / f"{code}.docx"), track=True)
        tbl = Document(out).element.body.find(f"{W}tbl")
        ps = list(tbl.iter(f"{W}p"))
        jcs = [p.find(f"{W}pPr/{W}jc") for p in ps]
        assert all(j is None or j.get(f"{W}val") in ("left", "start") for j in jcs)
        changed = [p for p in ps if p.find(f"{W}pPr/{W}pPrChange") is not None]
        assert len(changed) == 2  # tracked paragraph property changes on the two centred cells
        assert all(r.find(f"{W}rPr/{W}sz").get(f"{W}val") == "24" for r in tbl.iter(f"{W}r"))
    # a journal not switched over keeps the cells' own alignment (base: table_align null)
    cs = rules.run_all(docmodel.load(path), _st("ijo"), only=["format_fonts"])
    st = _st("ijo")
    st.settings["formatting"]["fonts"]["table_align"] = None
    cs = rules.run_all(docmodel.load(path), st, only=["format_fonts"])
    assert not any("align" in (c.format or {}) for c in cs.changes if c.old in ("Mean", "1.2"))


# --- G199 / G201: caption bold, Roman to Arabic ---------------------------------------------------

def test_whole_table_caption_bold_and_roman_labels(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Scores are shown (Table IV) and in Table I.")
        doc.add_paragraph("Table I. Baseline scores")
        _table(doc, [["Group", "Mean"], ["A", "1.2"]])
        doc.add_paragraph("Figure 1: Fundus photograph")
        doc.add_paragraph("Table IV: Follow up")
        _table(doc, [["Group", "Mean"], ["B", "3.4"]])
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        st = _st(code)
        cs = rules.run_all(docmodel.load(path), st, only=["fig_citations", "fig_captions", "format_cleanup"])
        bold = [c for c in cs.changes if c.reason == "fmt_table_caption"]
        assert sorted(c.old for c in bold) == ["Table I. Baseline scores", "Table IV: Follow up"]
        fig = [c for c in cs.changes if c.old.startswith("Figure 1") and (c.format or {}).get("bold")]
        assert fig and all(c.old == "Figure 1:" for c in fig)  # a figure legend: its label only
        out = _render(path, cs, str(tmp_path / f"{code}.docx"))
        texts = _texts(out)
        assert "Table 1: Baseline scores." in texts and "Table 4: Follow up." in texts  # S: caption ends with "."
        body = next(t for t in texts if t.startswith("Scores"))
        assert "[Table 4]" in body and "Table 1." in body
        cap = next(p for p in Document(out).paragraphs if p.text.startswith("Table 1:"))
        assert all(r.bold for r in cap.runs if r.text.strip())


def test_caption_in_a_tables_first_row_is_styled_too(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("See Table 2.")
        t = doc.add_table(rows=3, cols=2)
        top = t.cell(0, 0).merge(t.cell(0, 1))
        top.text = "Table II. Visual Acuity Outcomes"
        for (r, c), v in zip([(1, 0), (1, 1), (2, 0), (2, 1)], ["Eye", "VA", "OD", "0.2"]):
            t.cell(r, c).text = v
    path = _doc(tmp_path, build)
    cs = rules.run_all(docmodel.load(path), _st("ijo"), only=["fig_captions", "format_cleanup"])
    out = _render(path, cs, str(tmp_path / "o.docx"))
    assert _texts(out)[2] == "Table 2: Visual acuity outcomes."
    assert any(c.reason == "fmt_table_caption" for c in cs.changes)
    assert any(c.reason == "fmt_table_head" and c.old == "Eye" for c in cs.changes)  # row 2 is the head row


# --- G203-G205: column heads ----------------------------------------------------------------------

def test_column_heads_sentence_case_and_all_caps_title_case(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("The BCVA was recorded in Table 1.")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Study Design", "MEAN VALUE", "BCVA"], ["Cohort", "1.2", "0.3"]])
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        cs = rules.run_all(docmodel.load(path), _st(code), only=["fig_captions", "format_cleanup"])
        texts = _texts(_render(path, cs, str(tmp_path / f"{code}.docx")))
        assert texts[3:6] == ["Study design", "Mean Value", "BCVA"], code
        assert {c.old for c in cs.changes if c.reason == "fmt_table_head"} >= {"Study Design", "MEAN VALUE", "BCVA"}


# --- G294 / G296 inside the cells -----------------------------------------------------------------

def test_p_and_n_italic_inside_cells(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Group (n = 30)", "p value"], ["A", "p = 0.03"]])
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        cs = rules.run_all(docmodel.load(path), _st(code), only=["formatting"])
        got = {(c.old, c.new, (c.format or {}).get("italic")) for c in cs.changes}
        assert ("n", "n", True) in got
        assert sum(1 for c in cs.changes if c.old == "p" and c.new == "P" and c.format.get("italic")) == 2, code


# --- G208 / G209: indentation -----------------------------------------------------------------------

def _indented_doc(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        t = _table(doc, [["Variable", "n", "%"], ["Age", "", ""], ["\t18-30", "12", "40"], ["\t31-50", "18", "60"],
                         ["\t\tUnder 40", "9", "30"], ["Sex", "", ""], ["Male", "10", "33"], [" Female", "5", "17"]])
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        t.cell(6, 0).paragraphs[0]._p.get_or_add_pPr().append(ind)
    return _doc(tmp_path, build, "ind.docx")


def test_stub_indentation_becomes_five_spaces_per_level(tmp_path):
    path = _indented_doc(tmp_path)
    for code in JOURNALS:
        cs = rules.run_all(docmodel.load(path), _st(code), only=["table_indent"])
        out = _render(path, cs, str(tmp_path / f"{code}.docx"))
        stub = [t.rows[r].cells[0].text for t in Document(out).tables for r in range(1, 8)]
        # a tab: level 1 (5 spaces); two tabs: level 2 (10); a half-inch paragraph indent reads as a tab (cleared);
        # one typed space is no level of its own (padding)
        assert stub == ["Age", "     18-30", "     31-50", "          Under 40", "Sex", "     Male", "Female"], stub
        tr = _render(path, cs, str(tmp_path / f"{code}_t.docx"), track=True)
        tbl = Document(tr).element.body.find(f"{W}tbl")
        ind = [p.find(f"{W}pPr/{W}ind") for p in tbl.iter(f"{W}p") if p.find(f"{W}pPr/{W}ind") is not None]
        assert ind and ind[0].get(f"{W}left") == "0"
        assert ind[0].getparent().find(f"{W}pPrChange") is not None  # tracked


def test_rows_under_a_subheading_row_are_indented(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Variable", "n", "%"], ["Age (years)", "30", "100"], ["Sex", "", ""], ["Male", "10", "33"],
                     ["Female", "20", "67"], ["Total", "30", "100"]])
    path = _doc(tmp_path, build)
    cs = rules.run_all(docmodel.load(path), _st("jehp"), only=["table_indent"])
    out = _render(path, cs, str(tmp_path / "o.docx"))
    stub = [r.cells[0].text for r in Document(out).tables[0].rows]
    assert stub == ["Variable", "Age (years)", "Sex", "     Male", "     Female", "Total"]
    st = _st("jehp")
    st.settings["tables"]["indent"]["enabled"] = False
    assert rules.run_all(docmodel.load(path), st, only=["table_indent"]).changes == []


# --- G211-G215: table notes ---------------------------------------------------------------------------

def test_notes_abbreviations_first_term_equals_definition_no_end_periods(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Group", "Mean"], ["A", "1.2*"]])
        p = doc.add_paragraph()
        p.add_run("*").font.superscript = True
        p.add_run("P < 0.05 versus baseline.")
        q = doc.add_paragraph("SD: standard deviation; CI, confidence interval.")
        q.runs[0].italic = False
        doc.add_paragraph("Discussion")
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        cs = rules.run_all(docmodel.load(path), _st(code), only=["fig_captions"])
        for track in (False, True):
            out = _render(path, cs, str(tmp_path / f"{code}{track}.docx"), track)
            if track:
                body = Document(out).element.body
                assert body.find(f".//{W}pPr/{W}rPr/{W}del") is not None  # the list deleted where it was
                continue
            texts = [t for t in _texts(out) if t.strip()]
            i = texts.index("Mean") + 3
            assert texts[i:i + 3] == ["SD = standard deviation, CI = confidence interval", "*P < 0.05 versus baseline",
                                      "Discussion"], (code, texts)
            first = Document(out).paragraphs[2]  # the list's runs are not superscript like the "*" they precede
            assert not any(r.font.superscript for r in first.runs)


# --- G196: empty paragraphs inside cells ---------------------------------------------------------------

def _blank_cells(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        t = _table(doc, [["Group", "Mean"], ["A", "1.2"], ["B", "3.4"]])
        t.cell(1, 0).add_paragraph("")  # an empty last paragraph
        c = t.cell(2, 1)
        c.paragraphs[0].insert_paragraph_before("")  # an empty first paragraph
        c.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        t.cell(0, 1).add_paragraph("")
        t.cell(0, 1).add_paragraph("")  # two at the end
    return _doc(tmp_path, build, "blank.docx")


def test_empty_paragraphs_in_cells_deleted(tmp_path):
    path = _blank_cells(tmp_path)
    ms = docmodel.load(path)
    moves, _ = legends.plan(ms, _st("ijo"))
    blanks = [m for m in moves if m.kind == "cell_blanks"]
    assert len(blanks) == 1 and len(blanks[0].blocks[0]) == 4
    blanks[0].status = Status.ACCEPTED
    for is_tracked in (False, True):
        ms = docmodel.load(path)
        notes = legends.render(ms, blanks, (Status.ACCEPTED,), tracked=is_tracked, author="Edit OS",
                               date="2026-10-27T00:00:00Z")
        assert notes == []
        out = str(tmp_path / f"b{is_tracked}.docx")
        ms.save(out)
        tbl = Document(out).element.body.find(f"{W}tbl")
        if not is_tracked:
            per_cell = [[("".join(t.text or "" for t in p.iter(f"{W}t"))) for p in tc.findall(f"{W}p")]
                        for tc in tbl.iter(f"{W}tc")]
            assert per_cell == [["Group"], ["Mean"], ["A"], ["1.2"], ["B"], ["3.4"]]
        else:
            # accepting every revision (header._accept) gives the clean table, the text keeping its paragraph's
            # properties
            import copy
            acc = copy.deepcopy(tbl)
            header._accept(acc)
            per_cell = [[("".join(t.text or "" for t in p.iter(f"{W}t"))) for p in tc.findall(f"{W}p")]
                        for tc in acc.iter(f"{W}tc")]
            assert per_cell == [["Group"], ["Mean"], ["A"], ["1.2"], ["B"], ["3.4"]]
            last = list(acc.iter(f"{W}tc"))[5].find(f"{W}p")
            assert last.find(f"{W}pPr/{W}jc").get(f"{W}val") == "right"
    st = _st("ijo")
    st.settings["tables"]["delete_empty_cell_paragraphs"] = False
    assert not [m for m in legends.plan(docmodel.load(path), st)[0] if m.kind == "cell_blanks"]


# --- G206 / G196 queries ------------------------------------------------------------------------------

def test_queries_for_empty_head_cells_and_entries_on_separate_lines(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        t = _table(doc, [["Group", "", "", "Total"], ["A", "1", "2", "3"], ["B", "4", "5", "9"]])
        for c in range(4):
            t.cell(2, c).add_paragraph(str(c))
        doc.add_paragraph("Table 2: More")
        _table(doc, [["Group", "Value"], ["A", "1.2"]])
    path = _doc(tmp_path, build)
    for code in JOURNALS:
        qs = queries.QuerySet()
        rules_tables.add_queries(docmodel.load(path), _st(code), qs)
        assert sorted(q.check for q in qs.queries) == ["tables.cell_rows", "tables.empty_head_cells"]
    assert queries.QuerySet().queries == []
    qs = queries.QuerySet()
    rules_tables.add_queries(docmodel.load(path), _st("aian"), qs)  # off in base
    assert qs.queries == []


def test_settings_for_the_three_journals():
    for code in JOURNALS:
        st = style.load(code)
        assert st.get("formatting.fonts.table_align") == "left"
        assert st.get("formatting.table_caption_bold") is True
        assert st.get("figures.tables_to_end") is True
        assert st.get("figures.table_footnote_abbreviation") == " = "
        assert st.get("figures.table_footnote_delimiter") == ","
        assert st.get("figures.table_footnote_end_period") == "delete"
        assert st.get("figures.table_footnote_order") == "abbreviations_first"
        assert st.get("figures.table_heads_case") == "sentence" and st.get("figures.table_heads_title_case")
        assert st.get("tables.indent.enabled") and st.get("tables.delete_empty_cell_paragraphs")
        # the caption's final period: the stylesheets' "The caption should end with a period" (decision 8) stays
        assert st.get("figures.table_caption_end_period") == "add"
        assert st.get("formatting.p_values") and st.get("formatting.p_capital") and st.get("formatting.n_variable_italic")


def test_comma_joined_pairs_take_the_spaced_equals_sign(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Eye", "DS"], ["OD", "-1.0"]])
        doc.add_paragraph("Abbreviations: DS= Diopter Sphere, D= Diopter, Kf=Flat Keratometry.")
    path = _doc(tmp_path, build)
    cs = rules.run_all(docmodel.load(path), _st("ijo"), only=["fig_captions"])
    texts = _texts(_render(path, cs, str(tmp_path / "o.docx")))
    assert texts[-1] == "Abbreviations: DS = Diopter Sphere, D = Diopter, Kf = Flat Keratometry"


def test_blank_lines_that_keep_values_level_stay(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        t = _table(doc, [["Measure", "Value"], ["BCVA", ""], ["IOP", "14"]])
        for line in ("At admission", "At last visit"):  # rows typed as lines in the cells, blanks keep them level
            t.cell(1, 0).add_paragraph(line)
        for line in ("0.3", "0.5"):
            t.cell(1, 1).add_paragraph(line)
        t.cell(2, 0).add_paragraph("")  # an ordinary row: its blank goes
    path = _doc(tmp_path, build)
    ms = docmodel.load(path)
    gone = rules_tables.cell_blanks(ms)
    assert [ms.by_id(i).element.getparent() is ms.document.tables[0].cell(2, 0)._tc for i in gone] == [True]
    qs = queries.QuerySet()
    rules_tables.add_queries(ms, _st("jehp"), qs)
    assert [q.check for q in qs.queries] == ["tables.cell_rows"]


def test_bare_table_label_moves_with_its_table(tmp_path):
    from test_legends import REFS

    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Results are shown [Table 1].")
        doc.add_paragraph("Table 1")
        _table(doc, [["Eye", "DS"], ["OD", "-1.0"]])
        doc.add_paragraph("Discussion")
        for r in REFS:
            doc.add_paragraph(r)
    path = _doc(tmp_path, build)
    moves, _ = legends.plan(docmodel.load(path), _st("mjdrdypu"))
    tables = [m for m in moves if m.kind == "tables"]
    assert len(tables) == 1
    for m in tables:
        m.status = Status.ACCEPTED
    ms = docmodel.load(path)
    legends.render(ms, tables, (Status.ACCEPTED,), tracked=False, author="Edit OS", date="2026-10-27T00:00:00Z")
    ms.save(str(tmp_path / "o.docx"))
    body = [("TABLE" if el.tag == W + "tbl" else "".join(t.text or "" for t in el.iter(W + "t")))
            for el in Document(str(tmp_path / "o.docx")).element.body if el.tag in (W + "tbl", W + "p")]
    body = [t for t in body if t]
    assert body[-2:] == ["Table 1", "TABLE"] and body.index("Discussion") < body.index("Table 1")


def test_repeated_head_rows_are_not_indented(tmp_path):
    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Table 1: Results")
        _table(doc, [["Term", "Estimate"], ["Intercept", "1.0"], ["Group A", ""], ["Term", "Estimate"],
                     ["Intercept", "2.0"], ["Slope", "0.1"]])
    path = _doc(tmp_path, build)
    cs = rules.run_all(docmodel.load(path), _st("mjdrdypu"), only=["table_indent"])
    stub = [r.cells[0].text for r in Document(_render(path, cs, str(tmp_path / "o.docx"))).tables[0].rows]
    assert stub == ["Term", "Intercept", "Group A", "Term", "     Intercept", "     Slope"]
