"""Reference lists in shapes found on the shadow run (2026-10-29): the list laid out as a table (a "1." cell and the
entry beside it), labels typed right to left (".13 Adamson KA"), an entry broken after the author's name, a table's
note left under the list, and a year before a square-bracket citation. Invented text only."""

import pytest
from docx import Document

from edit_os import docmodel, first_author, frontmatter as fm, legends, qc, references, renumber, rules, style, tracked
from edit_os.changes import Status


def _style(journal="ijo"):
    st = style.load(journal)
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["back_matter"]["add_missing"] = False
    st.settings["sections"]["headings"]["drop"] = []
    return st


def _doc(tmp_path, before, rows=None, after=(), name="m.docx"):
    doc = Document()
    for t in before:
        doc.add_paragraph(t)
    if rows:
        table = doc.add_table(rows=len(rows), cols=2)
        for i, (a, b) in enumerate(rows):
            table.cell(i, 0).text = a
            table.cell(i, 1).text = b
    for t in after:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


ROWS = [
    ("1.", "Alder PQ, Brook RS. Suture abscess after glaucoma surgery. J Eye Res. 2001;12:10-4."),
    ("2.", "Cole TU, Dunn VW. Bleb infection in older adults. Eye Sci. 2010;3:20-5."),
    ("3.", "Ennis XY. Late keratitis around a loose suture. Ocul Rep. 2015;7:30-2."),
]
BODY = ["Original Article", "Introduction", "Sutures can loosen [1]. Blebs leak [3]. Infection follows [2].",
        "Conclusion", "Watch the sutures.", "References"]


def test_reference_list_laid_out_as_a_table_is_read(tmp_path):
    ms = docmodel.load(_doc(tmp_path, BODY, ROWS, ["Figure Legends", "Figure 1: A loose suture"]))
    refs = references.extract(ms)
    assert [r.number for r in refs] == [1, 2, 3]
    assert [r.label for r in refs] == ["", "", ""]  # the number is the cell beside the entry
    assert refs[0].text.startswith("Alder PQ") and refs[2].text.startswith("Ennis XY")
    assert all(ms.by_id(r.para_id).in_table for r in refs)


def test_a_data_table_after_the_list_is_not_read_as_references(tmp_path):
    ms = docmodel.load(_doc(tmp_path, BODY + ["1. Alder PQ. Suture abscess. J Eye Res. 2001;12:10-4."],
                            [("Group", "Eyes"), ("A", "12")]))
    assert [r.number for r in references.extract(ms)] == [1]


def test_first_author_block_goes_after_a_reference_table_not_under_its_heading(tmp_path):
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    ms = docmodel.load(_doc(tmp_path, BODY, ROWS, ["Figure Legends", "Figure 1: A loose suture"]))
    last_cell = references.extract(ms)[-1].para_id
    assert legends.refs_end(ms) == next(i for i, p in enumerate(ms.paragraphs) if p.id == last_cell)
    pl = first_author.plan(ms, st, fm.plan(ms, st))
    head = next(p for p in ms.paragraphs if p.text == "Figure Legends")
    assert pl.inserts and all(n.anchor == head.id and n.where == "before" for n in pl.inserts)


def test_reference_table_cited_out_of_order_is_renumbered(tmp_path):
    st = _style()
    ms = docmodel.load(_doc(tmp_path, BODY, ROWS))
    pl = renumber.plan(ms, st)
    assert pl.mapping == {1: 1, 2: 3, 3: 2} and pl.apply


def test_uncited_entry_in_a_reference_table_blocks_renumbering(tmp_path):
    st = _style()
    body = BODY[:2] + ["Sutures can loosen [1]. Infection follows [3]."] + BODY[3:]
    ms = docmodel.load(_doc(tmp_path, body, ROWS))
    pl = renumber.plan(ms, st)
    assert pl.uncited == [2] and not pl.mapping


def test_right_to_left_labels_are_labels(tmp_path):
    paras = BODY + [
        "1.\tAlder PQ, Brook RS. Suture abscess after glaucoma surgery. J Eye Res. 2001;12:10-4.",
        ".2\tCole TU, Dunn VW. Bleb infection in older adults. Eye Sci. 2010;3:20-5.",
        ".3 Ennis XY. Late keratitis around a loose suture. Ocul Rep. 2015;7:30-2.",
    ]
    ms = docmodel.load(_doc(tmp_path, paras))
    refs = references.extract(ms)
    assert [(r.number, r.label) for r in refs] == [(1, "1.\t"), (2, ".2\t"), (3, ".3 ")]
    assert refs[2].text.startswith("Ennis XY")


@pytest.mark.parametrize("journal,expected", [
    ("ijo", "3. Ennis XY."),  # label_period: the period after the number
    ("jehp", "Ennis XY."),  # label_auto_number: the typed label goes, Word numbers the entry
])
def test_right_to_left_label_is_rewritten_like_any_label(tmp_path, journal, expected):
    st = _style(journal)
    paras = BODY + ["1. Alder PQ. Suture abscess. J Eye Res. 2001;12:10-4.",
                    "2. Cole TU. Bleb infection. Eye Sci. 2010;3:20-5.",
                    ".3 Ennis XY. Late keratitis around a loose suture. Ocul Rep. 2015;7:30-2."]
    ms = docmodel.load(_doc(tmp_path, paras))
    ref = references.extract(ms)[2]
    p = ms.by_id(ref.para_id)
    cs = rules._label_changes(p, ref, st)
    text = p.text
    for c in sorted(cs, key=lambda c: c.start, reverse=True):
        text = text[:c.start] + c.new + text[c.end:]
    assert text.startswith(expected)


def test_entry_broken_after_the_authors_name_is_one_entry(tmp_path):
    paras = BODY + [
        "1. Alder PQ. Suture abscess after glaucoma surgery. J Eye Res. 2001;12:10-4.",
        "2. Tomas Fielding Grey.",
        "Bleb infection in older adults in a coastal town. Eye Sci. 2010. doi:10.9999/eye.2010.1.",
        "3. Ennis XY. Late keratitis around a loose suture. Ocul Rep. 2015;7:30-2.",
    ]
    ms = docmodel.load(_doc(tmp_path, paras))
    refs = references.extract(ms)
    assert [r.number for r in refs] == [1, 2, 3]
    assert refs[1].full_text.startswith("Tomas Fielding Grey. Bleb infection")


def test_table_note_left_under_a_numbered_list_ends_it(tmp_path):
    paras = BODY + [
        "1. Alder PQ. Suture abscess after glaucoma surgery. J Eye Res. 2001;12:10-4.",
        "2. Cole TU. Bleb infection in older adults. Eye Sci. 2010;3:20-5.",
        "* Cases reported by Cole (2010). BI = bleb infection; NR = not reported.",
        "First author:",
        "Pia Alder",
    ]
    ms = docmodel.load(_doc(tmp_path, paras))
    assert [r.number for r in references.extract(ms)] == [1, 2]


def test_first_author_lines_end_an_unlabelled_list(tmp_path):
    paras = BODY + ["Alder PQ. Suture abscess after glaucoma surgery. J Eye Res. 2001;12:10-4.",
                    "First author:", "Pia Alder", "Eye Unit, City Hospital, Pune, India"]
    ms = docmodel.load(_doc(tmp_path, paras))
    assert len(references.extract(ms)) == 1


def test_year_before_a_square_bracket_citation_cites(tmp_path):
    st = _style("jehp")
    paras = ["Original Article", "Introduction",
             "Arthritis is common [1, 2]. Cases will double by 2030 [3, 4]. Costs rise [5].", "References"] + [
        f"{i}. Alder PQ. Topic {i}. J Eye Res. 2001;12:{i}0-4." for i in range(1, 6)]
    path = _doc(tmp_path, paras)
    cs = rules.run_all(docmodel.load(path), st, only=["citations"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    ms = docmodel.load(str(out))
    body = next(p.text for p in ms.paragraphs if p.text.startswith("Arthritis"))
    assert "by 2030.[3,4] " in body and "common.[1,2]" in body  # JEHP: after the stop, no space inside
    assert not [i for i in qc.check(ms, st).issues if i.check == "citations.format"]


@pytest.mark.parametrize("after", [["Figure Legends", "Figure 1: A loose suture"], []])
def test_first_author_block_is_never_put_inside_the_reference_table(tmp_path, after):
    # the author's table has empty rows under the last entry; with or without legends after it
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    ms = docmodel.load(_doc(tmp_path, BODY, ROWS + [("", ""), ("", "")], after))
    anchor, where = legends.back_slot(ms)
    p = ms.by_id(anchor)
    assert not p.in_table
    if after:
        assert (p.text, where) == ("Figure Legends", "before")
    pl = first_author.plan(ms, st, fm.plan(ms, st))
    assert pl.inserts and all(not ms.by_id(n.anchor).in_table for n in pl.inserts)
