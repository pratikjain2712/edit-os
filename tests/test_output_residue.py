"""Defects the output scanner found in the delivered files: the author's own tracked changes and comments, the size
of punctuation inserted next to a superscript, false QC errors, a reply table after the back matter, double spaces
left by other changes, a paragraph break at the start of a hyperlink. All text is invented."""

import copy
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from edit_os import docmodel, pipeline, qc, revisions, rules_guide, style, submission, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status
from edit_os.copyedit import langtool

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _rev(tag, run, author="Author A"):
    w = OxmlElement(tag)
    w.set(qn("w:id"), "1")
    w.set(qn("w:author"), author)
    w.set(qn("w:date"), "2026-01-01T00:00:00Z")
    run.addprevious(w)
    w.append(run)
    return w


def _author_tracked(path):
    """A manuscript with the author's revisions pending: an insertion, a deletion, a format change, a deleted
    paragraph mark and a comment."""
    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph()
    p.add_run("The lens was ")
    gone = p.add_run("very ")
    p.add_run("clear")
    added = p.add_run(" and stable")
    p.add_run(".")
    dl = _rev("w:del", gone._r)
    for t in dl.iter(qn("w:t")):
        t.tag = qn("w:delText")
    _rev("w:ins", added._r)
    bold = p.add_run(" Follow-up was long.")
    rpr = bold._r.get_or_add_rPr()
    rpr.append(OxmlElement("w:b"))
    ch = OxmlElement("w:rPrChange")
    ch.set(qn("w:id"), "3")
    ch.set(qn("w:author"), "Author A")
    ch.append(OxmlElement("w:rPr"))
    rpr.append(ch)
    # "First half" with its paragraph mark deleted joins "second half."
    q = doc.add_paragraph("First half ")
    mark_rpr = OxmlElement("w:rPr")
    mark = OxmlElement("w:del")
    mark.set(qn("w:id"), "4")
    mark.set(qn("w:author"), "Author A")
    mark_rpr.append(mark)
    q._p.get_or_add_pPr().append(mark_rpr)
    doc.add_paragraph("second half.")
    c = doc.add_paragraph("Ethics statement provided.")
    doc.add_comment(c.runs, text="Provided", author="Office")
    doc.save(path)
    return str(path)


def test_accept_revisions_normalises_the_source(tmp_path):
    path = _author_tracked(tmp_path / "ms.docx")
    counts = revisions.accept_revisions(path)
    assert counts == {"insertions": 1, "deletions": 2, "formatting": 1}
    xml = zipfile.ZipFile(path).read("word/document.xml")
    for tag in (b"<w:ins ", b"<w:del ", b"<w:rPrChange", b"w:delText"):
        assert tag not in xml
    texts = [p.text for p in docmodel.load(path).paragraphs]
    assert "The lens was clear and stable. Follow-up was long." in texts
    assert "First half second half." in texts
    assert revisions.accept_revisions(path) == {"insertions": 0, "deletions": 0, "formatting": 0}  # nothing left


@pytest.fixture
def offline(monkeypatch):
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    return st


def test_pipeline_accepts_author_revisions_and_clean_copy_has_no_comments(tmp_path, offline):
    src = _author_tracked(tmp_path / "ms.docx")
    out = pipeline.edit(src, offline, str(tmp_path / "out"), stages=("style",), accept_all=True)
    man = pipeline.load_manifest(str(tmp_path / "out"))
    assert man["source_revisions_accepted"]["insertions"] == 1
    assert any("tracked changes in the manuscript were accepted" in n for n in out.summary["notes"])
    clean = zipfile.ZipFile(out.path("clean"))
    xml = clean.read("word/document.xml")
    assert b"<w:ins " not in xml or b'w:author="Author A"' not in xml
    assert b"commentReference" not in xml and b"commentRangeStart" not in xml
    assert "word/comments.xml" not in clean.namelist()
    tracked_ = zipfile.ZipFile(out.path("tracked"))
    assert b"Office" in tracked_.read("word/comments.xml")  # the editors still see it
    assert b'w:author="Author A"' not in tracked_.read("word/document.xml")


def _sized(p, text, size=None, sup=False):
    r = p.add_run(text)
    if size:
        r.font.size = size
    if sup:
        r.font.superscript = True
    return r


def test_punctuation_moved_before_a_citation_takes_the_sentence_size(tmp_path):
    from docx.shared import Pt

    doc = Document()
    p = doc.add_paragraph()
    _sized(p, "The rate rose", Pt(12))
    fld = p.add_run()  # an empty field-code run between the words and the citation, without the size
    fld._r.append(OxmlElement("w:fldChar"))
    fld._r[-1].set(qn("w:fldCharType"), "end")
    _sized(p, "1", Pt(12), sup=True)
    _sized(p, ".", Pt(12))
    doc.save(tmp_path / "m.docx")
    ms = docmodel.load(str(tmp_path / "m.docx"))
    para = ms.paragraphs[0]
    assert para.text == "The rate rose1."
    cs = ChangeSet([Change("c1", para.id, 13, 13, "", ".", Layer.REFERENCES, Source.RULE, "citations",
                           Status.ACCEPTED),
                    Change("c2", para.id, 14, 15, ".", "", Layer.REFERENCES, Source.RULE, "citations",
                           Status.ACCEPTED)])
    tracked.render(ms, cs, tracked=False, statuses=(Status.ACCEPTED,))
    dot = [r for r in docmodel.text_runs(para.element) if docmodel.run_text(r) == "."][0]
    assert dot.find(f"{W}rPr/{W}sz").get(qn("w:val")) == "24"
    assert dot.find(f"{W}rPr/{W}vertAlign") is None
    assert docmodel.paragraph_text(para.element) == "The rate rose.1"


def test_new_text_with_a_paragraph_break_at_the_start_of_a_hyperlink(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), "rId99")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Consent was obtained."
    r.append(t)
    link.append(r)
    p._p.append(link)
    doc.save(tmp_path / "m.docx")
    for is_tracked in (True, False):
        ms = docmodel.load(str(tmp_path / "m.docx"))
        pid = ms.paragraphs[0].id
        cs = ChangeSet([Change("c1", pid, 0, 0, "", "Declaration of consent" + tracked.PARAGRAPH_BREAK,
                               Layer.STRUCTURE, Source.RULE, "sections_consent", Status.ACCEPTED)])
        errors = []
        tracked.render(ms, cs, tracked=is_tracked, statuses=(Status.ACCEPTED,),
                       on_conflict=lambda pid_, err: errors.append(err))
        assert not errors
        texts = [docmodel.paragraph_text(el) for el in ms.document.element.body.iter(qn("w:p"))]
        assert texts[:2] == ["Declaration of consent", "Consent was obtained."]


def _kw_ms(tmp_path, line):
    doc = Document()
    for t in ["Abstract", line, "Introduction", "Body text."]:
        doc.add_paragraph(t)
    doc.save(tmp_path / "k.docx")
    return docmodel.load(str(tmp_path / "k.docx"))


def test_qc_keywords_semicolon_list_keeps_commas_inside_terms(tmp_path):
    st = style.load("jehp")
    ms = _kw_ms(tmp_path, "Keywords: Clinical services; education, professional; midwifery; validation study")
    assert not [i for i in qc.check(ms, st).issues if i.check == "keywords.order"]
    ms = _kw_ms(tmp_path, "Keywords: Midwifery; education, professional; validation study")
    assert [i for i in qc.check(ms, st).issues if i.check == "keywords.order"]


def test_qc_or_is_a_conjunction_not_an_odds_ratio():
    import re

    text = "Values were significant at P < 0.05 or <0.01 as shown."
    m = re.search(r"(?<=[A-Za-z0-9)]) [=<>≤≥](?=\d)", text)
    assert qc._unary_after_word(text, m)
    text = "The risk rose with OR >2 in the group."
    m = re.search(r"(?<=[A-Za-z0-9)]) [=<>≤≥](?=\d)", text)
    assert not qc._unary_after_word(text, m)


def test_qc_author_count_group_authors_and_apa_order():
    assert qc._author_count("Rao JA, Iyer W, Bose RJ, Kale MP, Das JA; Ethics, Law, and Humanities Committee, a "
                            "joint committee of the Society of Neurology") == 6
    assert qc._author_count("Rao, C, Iyer, J, Bose, J, & Das, A") == 4
    assert qc._author_count("Rao A, Iyer B, Bose C, Kale D, Das E, Sen F, Roy G") == 7


def test_qc_reference_problems_the_author_is_asked_about_are_warnings(tmp_path):
    doc = Document()
    for t in ["Introduction", "Cataract is common [2]. It recurs [3].", "References",
              "1. Rao A, Iyer B. A study one. J Test 2001;1:1-9.", "2. Rao A, Iyer B. A study two. J Test 2001;2:1-9.",
              "3. Rao A, Iyer B. A study three. J Test 2001;3:1-9."]:
        doc.add_paragraph(t)
    doc.save(tmp_path / "r.docx")
    ms = docmodel.load(str(tmp_path / "r.docx"))
    st = style.load("ijo")
    plain = {i.check: i.severity for i in qc.check(ms, st).issues if i.check in ("citations.order", "references.cited")}
    assert plain == {"citations.order": qc.ERROR, "references.cited": qc.ERROR}
    q = [{"check": "references.uncited", "text": "Reference no. 1 has not been cited. Kindly cite."}]
    asked = {i.check: i.severity for i in qc.check(ms, st, queries=q).issues
             if i.check in ("citations.order", "references.cited")}
    assert asked == {"citations.order": qc.WARNING, "references.cited": qc.WARNING}


BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "hospital over the years of the study.")


def test_reply_table_after_the_back_matter_is_deleted(tmp_path):
    doc = Document()
    for t in ["Abstract", BODY, "Discussion", BODY, "References", "1. Rao A, Iyer B. A study. J Test 2001;1:1-9.",
              "Financial support and sponsorship: Nil", "Conflicts of interest: Nil", ""]:
        doc.add_paragraph(t)
    rows = [("No", "Reviewers comment", "Authors Response", "Line no / Page no"),
            ("1.", "Please add the consent statement.", "Statement provided", "157-158")]
    t = doc.add_table(rows=2, cols=4)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.cell(i, j).text = cell
    doc.save(tmp_path / "m.docx")
    ms = docmodel.load(str(tmp_path / "m.docx"))
    gone = submission.deleted_ids(ms, style.load("ijo"))
    assert {p.text for p in ms.paragraphs if p.id in gone} >= {"Reviewers comment", "Statement provided"}
    kept = {p.text for p in ms.paragraphs if p.id not in gone}
    assert {"References", BODY, "Financial support and sponsorship: Nil"} <= kept


def test_reply_header_table_in_the_body_is_kept(tmp_path):
    doc = Document()
    for t in ["Abstract", BODY, "Methods"]:
        doc.add_paragraph(t)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Reviewers comment", "Response"
    t.cell(1, 0).text, t.cell(1, 1).text = "Clear", "Yes"
    for x in ["Discussion", BODY, "References", "1. Rao A, Iyer B. A study. J Test 2001;1:1-9."]:
        doc.add_paragraph(x)
    doc.save(tmp_path / "b.docx")
    ms = docmodel.load(str(tmp_path / "b.docx"))
    assert not submission.deleted_ids(ms, style.load("ijo"))


def _para(tmp_path, text):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(tmp_path / "d.docx")
    return docmodel.load(str(tmp_path / "d.docx"))


def _apply(text, cs):
    for c in sorted(cs, key=lambda c: (c.start, c.end), reverse=True):
        text = text[:c.start] + c.new + text[c.end:]
    return text


def test_double_spaces_left_by_other_changes_are_closed(tmp_path):
    st = style.load("jehp")
    text = "It is common among women (1).  It causes pain,  mild or severe."
    ms = _para(tmp_path, text)
    pid = ms.paragraphs[0].id
    i = text.index(" (1).")
    j = text.index(" ")
    others = [Change("a", pid, i, i + 5, " (1).", ".", Layer.REFERENCES, Source.RULE, "citations"),
              Change("b", pid, i + 5, i + 5, "", "[1]", Layer.REFERENCES, Source.RULE, "citations"),
              Change("c", pid, j, j + 1, " ", " ", Layer.HOUSE_STYLE, Source.RULE, "guide.nbsp")]
    new = rules_guide.collapse_double_spaces(ms, st, others)
    assert _apply(text, others + new) == "It is common among women.[1] It causes pain, mild or severe."
    # a deletion between two spaces takes one of them with it: rejected, both come back
    text = "The rate [x] rose."
    ms = _para(tmp_path, text)
    d = Change("d", ms.paragraphs[0].id, 9, 12, "[x]", "", Layer.HOUSE_STYLE, Source.RULE, "other")
    assert rules_guide.collapse_double_spaces(ms, st, [d]) == []
    assert _apply(text, [d]) == "The rate rose." and d.old == "[x] "
