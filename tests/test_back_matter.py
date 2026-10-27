"""Back matter (rules_sections.rule_back_matter): consent, contributions, nil statements, order.

All text is synthetic."""

from docx import Document

from edit_os import docmodel, frontmatter, queries, rules, rules_sections, style, tracked
from edit_os.changes import Status

BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic.")
REFS = ["References", "1. Smith A, Jones B. A study. Ophthalmology 2001;12:1-9.",
        "2. Smith A, Jones B. Another study. Ophthalmology 2002;13:1-9."]
STD = ["<H2>Financial support and sponsorship", "Nil.", "<H2>Conflicts of interest",
       "There are no conflicts of interest."]
CHECKS = ("sections.consent_added", "sections.patient_consent", "sections.contributions_missing")
CONSENT_TEXT = style.load("ijo").get("sections.back_matter.consent.text")


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def load(path, article_type="", first_page=(), forms=()):
    ms = docmodel.load(path)
    ms.front_sources = frontmatter.Sources(first_page=list(first_page), meta=frontmatter.Meta(article_type=article_type),
                                           contribution_forms=[list(f) for f in forms])
    return ms


def run(path, code="ijo", only=("heading_tags", "sections_back_matter"), **src):
    st = style.load(code) if isinstance(code, str) else code
    cs = rules.run_all(load(path, **src), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = path.replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs if p.text]


def section_queries(path, code="ijo", **src):
    st = style.load(code)
    return [q.text for q in queries.collect(load(path, **src), st).queries if q.check in CHECKS]


# --- patient consent ---------------------------------------------------------------------

def test_unlabelled_consent_statement_gets_heading_and_declarations_go(tmp_path):
    statement = ("Written informed consent was obtained from the patient for the publication of this case report "
                 "and the accompanying images.")
    path = make(tmp_path, ["Case Report", BODY, "Conclusion", BODY, "Declarations",
                           "Ethics: Not required for single case report.", statement,
                           "Conflict of Interest: None.", "Funding: None."] + REFS)
    t = run(path, article_type="Case Report")
    i = t.index("<H2>Declaration of patient consent")  # the IJO sheet's heading (not the editors' wording)
    assert t[i + 1] == statement
    assert t[i + 2:i + 6] == STD and t[i + 6] == "<H1>References"
    assert "Declarations" not in t and not any(x.startswith("Ethics") for x in t)
    assert section_queries(path, article_type="Case Report") == []  # consent is there: no query


def test_inline_consent_label_split_and_renamed_for_aian(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Patient  Consent- Written informed consent for publication of "
                           "clinical details and images was obtained from the patient."] + STD[1:2] + REFS)
    t = run(path, "aian")
    i = t.index("<H2>Declaration of patient consent")  # AIAN: the journal's heading
    assert t[i + 1].startswith("Written informed consent for publication")
    t = run(path, "ijo")  # IJO too: S388 (owner, 2026-10-27; it kept the author's words before)
    assert "<H2>Declaration of patient consent" in t and "<H2>Patient consent" not in t


def test_consent_added_for_clinical_photographs_with_query(tmp_path):
    path = make(tmp_path, ["Case Report", BODY, "Discussion", BODY] + REFS +
                ["Figure 1: Clinical photograph showing a lesion of the right upper eyelid",
                 "Figure 2: Histopathology of the excised lesion"])
    t = run(path, "aian", article_type="Case Report")
    i = t.index("<H2>Declaration of patient consent")
    assert t[i + 1] == CONSENT_TEXT
    assert t[i + 2:i + 6] == STD  # placed before "Financial support and sponsorship"
    assert section_queries(path, "aian", article_type="Case Report") == [
        "Please ensure that patient consent form has been obtained."]


def test_consent_not_added_without_patient_images(tmp_path):
    path = make(tmp_path, ["Case Report", BODY, "Discussion", BODY] + REFS +
                ["Figure 1: Fundus showing a pale disc", "Figure 2: Visual fields of both eyes"])
    t = run(path, article_type="Case Report")
    assert not any("consent" in x.lower() for x in t)
    # the pre-editing query for a case report with no consent statement
    assert section_queries(path, article_type="Case Report") == ["Kindly confirm whether consent form obtained or not?"]
    # a review with clinical photographs is not a patient's case: nothing added
    path = make(tmp_path, ["Discussion", BODY] + REFS + ["Figure 1: Clinical photograph of a limbal dermoid"], "r.docx")
    assert not any("consent" in x.lower() for x in run(path, article_type="Review Article"))


def test_image_article_gets_consent(tmp_path):
    path = make(tmp_path, ["A child presented with gait difficulty. " + BODY] + REFS +
                ["Figure 1: Axial MRI showing a fourth ventricular mass"])
    t = run(path, "aian", article_type="Images in Neurology")
    assert "<H2>Declaration of patient consent" in t
    assert rules_sections.identifiable_patient_images(load(path, article_type="Images in Neurology"))
    # JOS: no consent rule in its stylesheet
    assert "consent" not in str(run(path, "jos", article_type="Images in Neurology")).lower()


# --- authors' contributions ----------------------------------------------------------------

FIRST_PAGE = ["Title of the case", "Author Contributions", "Dr Asha Rao", "Department of Ophthalmology",
              "Author Contributions",
              "Asha Rao (AR): Concept, design, literature search and manuscript preparation.",
              "Ravi Kumar (RK): Data acquisition, manuscript editing and manuscript review.",
              "All authors read and approved the final manuscript.",
              "Corresponding author: Dr Asha Rao"]


def test_contributions_from_title_page_in_order(tmp_path):
    path = make(tmp_path, ["Case Report", BODY, "Conclusion", BODY, "Declaration of patient consent",
                           "The authors certify that they have obtained all appropriate patient consent forms.",
                           "Financial support and sponsorship", "Nil.", "Conflicts of interest",
                           "There are no conflicts of interest."] + REFS)
    t = run(path, first_page=FIRST_PAGE)
    i = t.index("<H2>Declaration of patient consent")
    assert t[i + 2:i + 6] == ["<H2>Authors’ contributions", FIRST_PAGE[5], FIRST_PAGE[6], FIRST_PAGE[7]]
    assert t[i + 6:i + 10] == STD
    # ABR editors do not add them; a journal without from_sources adds nothing
    assert not any("contributions" in x for x in run(path, "abr", first_page=FIRST_PAGE))
    assert not any("contributions" in x for x in run(path, "njcp", first_page=FIRST_PAGE))


def test_contributions_from_form_and_cover_letter(tmp_path):
    form = ["Author Contributions", "A.B.: Conceptualization, Methodology, Writing – Original Draft.",
            "C.D.: Supervision, Writing – Review & Editing."]
    path = make(tmp_path, ["Discussion", BODY, "Funding: None"] + REFS)
    t = run(path, forms=[form])
    i = t.index("<H2>Authors’ contributions")
    assert t[i + 1:i + 3] == form[1:] and t[i + 3:i + 7] == STD
    # a cover letter inside the manuscript
    letter = ["Cover Letter", "6. Author Contributions:", "Dr A Rao – Concept, Design, Manuscript Review.",
              "Dr B Pillai – Literature Search, Data Acquisition.", "Thank you for considering our work."]
    path = make(tmp_path, letter + ["Discussion", BODY] + REFS, "c.docx")
    t = run(path)
    i = t.index("<H2>Authors’ contributions")
    assert t[i + 1:i + 3] == letter[2:4]


def test_contributions_not_invented(tmp_path):
    template = ["AUTHOR CONTRIBUTIONS", "Author name", "Author's contributions: Role (Concepts, Design)",
                "[Add for all authors separately here]", "Contributors :", "Dr A Rao", "Dr B Pillai"]
    path = make(tmp_path, ["Discussion", BODY] + REFS)
    t = run(path, first_page=template)
    assert not any("contributions" in x.lower() for x in t)
    st = style.load("ijo")
    st.settings["sections"]["back_matter"]["contributions"]["query_missing"] = True
    qs = [q.text for q in queries.collect(load(path, first_page=template), st).queries if q.check in CHECKS]
    assert qs == ["Please provide authors’ contributions."]


def test_unlabelled_contribution_lines_get_heading_and_inline_label_splits(tmp_path):
    lines = ["Prof. Dr. Asha Rao – Concept, Definition of intellectual content, Manuscript review.",
             "Dr Ravi Kumar – Literature search, Data acquisition, Manuscript preparation."]
    path = make(tmp_path, ["Conclusion", BODY, "Funding: None."] + lines + REFS)
    t = run(path)
    i = t.index("<H2>Authors’ contributions")
    assert t[i + 1:i + 3] == lines and t[i + 3:i + 7] == STD
    path = make(tmp_path, ["Discussion", BODY, "Author contributions: US: data collection; analysis; writing."] + REFS,
                "i.docx")
    t = run(path, "aian")
    assert t[t.index("<H2>Author contributions") + 1] == "US: data collection; analysis; writing."
    # "Contribution:" under each author's name is not the section label
    path = make(tmp_path, ["Discussion", BODY, "Dr A", "Contribution: This author helped with study design."] + REFS,
                "j.docx")
    assert "Contribution: This author helped with study design." in run(path, "jehp")


TOP = ["Orbital lesions with enophthalmos: A case series", "",
       "Asha Rao: Concept, design, data acquisition, manuscript preparation.",
       "Ravi Kumar: Data acquisition, data analysis, manuscript review.",
       "All authors approved the final version of the manuscript."]


def test_contributions_on_the_manuscript_title_page_move_to_the_back(tmp_path):
    path = make(tmp_path, TOP + ["Introduction", BODY, "Conclusion", BODY, "Declaration of patient consent",
                                 "The authors certify that they have obtained all appropriate patient consent forms.",
                                 "Financial support and sponsorship", "Nil.", "Conflicts of interest",
                                 "There are no conflicts of interest."] + REFS)
    cs = rules.run_all(load(path), style.load("ijo"), only=["heading_tags", "sections_back_matter"])
    gone = [c for c in cs.changes if c.reason == "sections_contributions" and not c.new]
    assert sorted(c.old for c in gone) == sorted(TOP[2:])  # tracked deletes carrying the text they replace
    t = run(path)
    assert t[0] == TOP[0] and not any(x in t[:4] for x in TOP[2:])
    i = t.index("<H2>Declaration of patient consent")
    assert t[i + 2:i + 6] == ["<H2>Authors’ contributions"] + TOP[2:]
    assert t[i + 6:i + 10] == STD
    # a journal that does not move them (move_from_top off) leaves the title page alone
    t = run(path, "aian")
    assert t[1:4] == TOP[2:] and not any("contributions" in x.lower() for x in t)


def test_labelled_top_contributions_move_without_their_label(tmp_path):
    path = make(tmp_path, ["A title", "Authors’ contributions:", "Asha Rao: Concept, design, literature search.",
                           "Ravi Kumar: Data acquisition, manuscript review.", "Abstract", BODY, "Discussion", BODY]
                + REFS)
    t = run(path)
    assert t[:2] == ["A title", "Abstract"]
    i = t.index("<H2>Authors’ contributions")
    assert t[i + 1:i + 3] == ["Asha Rao: Concept, design, literature search.",
                              "Ravi Kumar: Data acquisition, manuscript review."]
    assert t[i + 3:i + 7] == STD


def test_top_contributions_already_at_the_back(tmp_path):
    back = ["Authors’ contributions"] + TOP[2:4]
    path = make(tmp_path, TOP + ["Introduction", BODY, "Conclusion", BODY] + back + REFS)
    t = run(path)
    assert t[0] == TOP[0] and t.count(TOP[2]) == 1 and t.count(TOP[3]) == 1  # the duplicate on top is deleted
    assert sum("contributions" in x.lower() for x in t) == 1
    assert TOP[4] not in t  # the closing line went with the top copy (the back is the authors' own section)
    # the back says something else: the top copy is left to the editor
    other = ["Authors’ contributions", "Asha Rao: Concept and design.", "Ravi Kumar: Data analysis and review."]
    path = make(tmp_path, TOP + ["Introduction", BODY, "Conclusion", BODY] + other + REFS, "o.docx")
    t = run(path)
    assert t[1:4] == TOP[2:] and sum("contributions" in x.lower() for x in t) == 1


# --- nil statements ------------------------------------------------------------------------

def test_nil_statement_variants_deleted(tmp_path):
    nil = ["Conflicts of interests: None", "Conflict(s) of Interest and Disclosure(s): None",
           "Source(s) of support in the form of grants, equipment, drugs, or all of these: Nil.",
           "Support / Funding – Nil", "Acknowledge- None", "Competing interests",
           "The authors declare that they have no competing interests.", "No conflict of interest.",
           "The authors have nothing to disclose.", "Funding:",
           "This systematic review has no specific grant from any funding agency."]
    path = make(tmp_path, ["Discussion", BODY] + nil + REFS)
    t = run(path, "njcp")
    assert t[2:6] == STD and t[6] == "<H1>References"


def test_real_statements_kept_and_nil_after_reference_list(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Funding: This work was supported by the Eye Foundation (grant 12)."]
                + REFS + ["Conflicts of interests: None"])
    t = run(path)
    i = t.index("<H2>Financial support and sponsorship")
    assert t[i + 1] == "This work was supported by the Eye Foundation (grant 12)."
    assert t[i + 2:i + 4] == STD[2:]
    assert "Conflicts of interests: None" not in t  # read as a reference entry before; it is back matter


def test_nil_statement_becomes_section_in_place_only_in_order(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Conflicts of interests: None", "Acknowledgement: None",
                           "Funding: None"] + REFS)
    st = style.load("ijo")
    cs = rules.run_all(load(path), st, only=["sections_back_matter"])
    funding = [c for c in cs.changes if c.old.startswith("Funding")]
    assert funding and funding[0].new == "<H2>Financial support and sponsorship"  # converted in place
    assert any(c.old == "Conflicts of interests: None" and c.new == "" for c in cs.changes)  # before funding: deleted
    assert run(path)[2:7] == STD + ["<H1>References"]


def test_nil_ethics_inline_only_and_standard_wording(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Consent for publication: Not applicable.", "Consent to participate",
                           "Not applicable", "Conflicts of interest", "The authors declare no conflict of interest."]
                + REFS)
    t = run(path, "abr")
    assert "Consent for publication: Not applicable." not in t
    assert "Not applicable" in t  # a heading with "Not applicable" on its own line stays
    i = t.index("<H2>Conflicts of interest")
    assert t[i + 1] == "There are no conflicts of interest."


# --- acknowledgement, supplementary material ------------------------------------------------

def test_acknowledgement_singular_and_supplementary_heading_from_the_stylesheet(tmp_path):
    paras = ["Discussion", BODY, "Acknowledgements", "We thank the staff of the eye clinic for their help.",
             "Supplementary Table 1: Search strategy", "Supplementary Figure 1: Forest plot"] + REFS
    path = make(tmp_path, paras)
    # IJO sheet: 'If heading "Acknowledgements" is present, then please delete "s"' and 'please add the heading
    # “Supplementary Material” after Acknowledgement'
    t = run(path)
    assert "<H1>Acknowledgement" in t  # IJO: G146 "Acknowledgement", first level (owner, 2026-10-26)
    assert t[t.index("<H2>Supplementary Material") + 1] == "Supplementary Table 1: Search strategy"
    # KJO follows the guide: "Acknowledgements" keeps its "s", supporting information as it is.
    t = run(path, "kjo")
    assert "<H2>Acknowledgements" in t and "Supplementary Material" not in str(t)
    # "singular" drops the s only: the author's US spelling stays.
    path2 = make(tmp_path, ["Discussion", BODY, "Acknowledgments", "We thank the staff."] + REFS, name="m2.docx")
    assert "<H2>Acknowledgment" in run(path2, "ijdr")
    assert "<H1>Acknowledgement" in run(path2)  # IJO: the guide's (UK) spelling
    st = style.load("ijo")
    st.settings["sections"]["back_matter"]["acknowledgement_heading"] = None
    assert "<H1>Acknowledgments" in run(path2, st)  # IJO: first level either way


def test_added_consent_has_its_own_reason(tmp_path):
    path = make(tmp_path, ["Discussion", BODY] + REFS + ["Figure 1: Clinical photograph of the face"])
    cs = rules.run_all(load(path, article_type="Case Report"), style.load("ijo"), only=["sections_back_matter"])
    added = [c for c in cs.changes if CONSENT_TEXT in c.new]
    assert added and added[0].reason == "sections_back_matter.consent_added"


def test_jehp_keeps_author_wording_and_conflict_heading(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Funding", "This research received no specific grant.",
                           "Conflict of interest", "The authors declare no conflict of interest."] + REFS)
    t = run(path, "jehp")
    assert t[2:6] == ["<H2>Financial support and sponsorship", "This research received no specific grant.",
                      "<H2>Conflict of interest", "The authors declare no conflict of interest."]
    # the base keeps renaming and rewording
    assert run(path, "njcp")[2:6] == STD


def test_jehp_inline_statement_split_not_deleted(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Conflicts of interest: There are no conflicts of interest."] + REFS)
    t = run(path, "jehp")
    assert t[2:6] == STD  # the conflicts statement split into heading + statement; funding added before it
