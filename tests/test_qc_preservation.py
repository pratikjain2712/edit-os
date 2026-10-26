"""Content-preservation QC (preservation.py): the output compared with the original manuscript."""

from docx import Document

from edit_os import docmodel, pipeline, preservation, qc, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status
from edit_os.frontmatter import NewParagraph
from edit_os.tracked import PARAGRAPH_BREAK

PROSE = ("Health education programmes for adolescents were assessed across several schools in the district, and the "
         "teachers who delivered them were interviewed about the barriers they met, the time each session needed, the "
         "materials they used and the response of the students to the sessions")

ARTICLE = [
    "Dear Editor,",
    "Please consider our manuscript for publication in your journal. It has not been submitted elsewhere.",
    "Yours sincerely, the authors",
    "Original Article",
    "Teaching health in schools",
    "A. Author, B. Author",
    "Introduction",
    f"{PROSE} [1].",
    "Material and Methods",
    f"{PROSE}.",
    "Result",
    f"{PROSE}.",
    "Figure 1: Flow of participants through the study",
    "Figure 2: Scores before and after the programme",
    "Figure 3: Scores by school",
    "Reviewer comment: please explain the sampling.",
    "Response: we added the sampling frame to the Methods.",
    "Discussion",
    f"{PROSE} [2].",
    f"{PROSE} [3].",
    "Limitations",
    f"{PROSE}.",
    "Conclusions",
    f"{PROSE}.",
    "Acknowledgement",
    "Nil.",
    "Funding",
    "The study received no funding from any agency.",
    "References",
    "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit. 2002;21:295-9.",
    "2. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4.",
    "3. Rao P. School health in India. Indian J Pediatr. 2010;77:101-5.",
]


def _doc(tmp_path, paras=ARTICLE):
    d = Document()
    for t in paras:
        d.add_paragraph(t)
    path = tmp_path / "ms.docx"
    d.save(path)
    return docmodel.load(str(path))


def _para(ms, text):
    return next(p for p in ms.paragraphs if p.text == text)


def _delete(ms, text, reason, status=Status.ACCEPTED):
    p = text if isinstance(text, docmodel.Paragraph) else _para(ms, text)
    return Change(id=f"d_{p.id}_{reason}", para_id=p.id, start=0, end=len(p.text), old=p.text, new="",
                  layer=Layer.STRUCTURE, source=Source.RULE, reason=reason, status=status)


def _replace(ms, text, new, reason, status=Status.ACCEPTED):
    p = _para(ms, text)
    return Change(id=f"r_{p.id}", para_id=p.id, start=0, end=len(p.text), old=p.text, new=new,
                  layer=Layer.STRUCTURE, source=Source.RULE, reason=reason, status=status)


def _run(tmp_path, changes, inserts=(), st=None):
    """Render the changes (clean) and QC the result against the original."""
    st = st or style.load("jehp")
    src = docmodel.load(str(tmp_path / "ms.docx"))
    out = docmodel.load(str(tmp_path / "ms.docx"))
    tracked.render(out, ChangeSet(list(changes)), tracked=False, statuses=(Status.ACCEPTED,))
    out.save(str(tmp_path / "out.docx"))
    out = docmodel.load(str(tmp_path / "out.docx"))
    return qc.check(out, st, source=src, changes=ChangeSet(list(changes)), inserts=list(inserts),
                    statuses=(Status.ACCEPTED,))


def _content(report):
    return [i for i in report.issues if i.check.startswith("content.")]


def test_reply_that_takes_the_discussion_is_blocked(tmp_path):
    ms = _doc(tmp_path)
    start = next(i for i, p in enumerate(ms.paragraphs) if p.text.startswith("Reviewer comment"))
    end = next(i for i, p in enumerate(ms.paragraphs) if p.text == "References")
    changes = [_delete(ms, p, "submission_reviewer_reply") for p in ms.paragraphs[start:end] if p.text.strip()]
    rep = _run(tmp_path, changes)
    issues = _content(rep)
    sections = [i.message for i in issues if i.check == "content.sections"]
    for name in ("Discussion", "Limitations", "Conclusion"):
        assert any(f"'{name}'" in m and "submission_reviewer_reply" in m for m in sections), (name, sections)
    deleted = [i for i in issues if i.check == "content.deleted"]
    assert len(deleted) == 1 and deleted[0].severity == qc.ERROR
    assert "submission_reviewer_reply" in deleted[0].message and "'Discussion'" in deleted[0].message
    assert deleted[0].para_id == _para(ms, "Discussion").id
    refs = [i for i in issues if i.check == "content.references"]
    assert refs and "2, 3" in refs[0].message
    assert not rep.passed


def test_the_reply_alone_deleted_is_fine(tmp_path):
    ms = _doc(tmp_path)
    changes = [_delete(ms, "Reviewer comment: please explain the sampling.", "submission_reviewer_reply"),
               _delete(ms, "Response: we added the sampling frame to the Methods.", "submission_reviewer_reply")]
    assert _content(_run(tmp_path, changes)) == []


def test_a_normal_run_raises_nothing(tmp_path):
    ms = _doc(tmp_path)
    legends = ["Figure 1: Flow of participants through the study", "Figure 2: Scores before and after the programme",
               "Figure 3: Scores by school"]
    last_ref = "3. Rao P. School health in India. Indian J Pediatr. 2010;77:101-5."
    changes = [
        # cover letter removed, title page rebuilt by the front matter
        *[_delete(ms, t, "submission_cover_letter") for t in ARTICLE[:3]],
        _delete(ms, "A. Author, B. Author", "front_matter"),
        # heading renames
        _replace(ms, "Introduction", "<H1>Introduction", "heading_tags"),
        _replace(ms, "Material and Methods", "<H1>Materials and Methods", "sections_headings"),
        _replace(ms, "Result", "<H1>Results", "sections_headings"),
        _replace(ms, "Conclusions", "<H1>Conclusion", "sections_headings"),
        # legends to the end: deleted here, put back after the references
        *[_delete(ms, t, "move_legends") for t in legends],
        Change(id="leg", para_id=_para(ms, last_ref).id, start=len(last_ref), end=len(last_ref), old="",
               new="".join(PARAGRAPH_BREAK + t for t in legends), layer=Layer.TABLES_FIGURES, source=Source.RULE,
               reason="move_legends", status=Status.ACCEPTED),
        # nil acknowledgement restyled by the back-matter rule
        _delete(ms, "Acknowledgement", "sections_back_matter"),
        _delete(ms, "Nil.", "sections_back_matter"),
        # references restyled
        _replace(ms, "2. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4.",
                 "2. Lee K. Back pain in adolescents. Spine (Phila Pa 1976) 1997;22:1380-4.", "ref_lookup"),
    ]
    rep = _run(tmp_path, changes)
    assert _content(rep) == [], [i.message for i in _content(rep)]


def test_moved_block_as_a_front_matter_insert_is_not_a_loss(tmp_path):
    ms = _doc(tmp_path)
    legends = ["Figure 1: Flow of participants through the study", "Figure 2: Scores before and after the programme",
               "Figure 3: Scores by school"]
    changes = [_delete(ms, t, "fig_captions") for t in legends]
    ins = [NewParagraph(id=f"n{i}", anchor=_para(ms, "References").id, where="before", runs=[(t, None)], role="legend",
                        status=Status.ACCEPTED) for i, t in enumerate(legends)]
    assert _content(_run(tmp_path, changes, ins)) == []
    # not put back: three short caption lines are under MIN_RUN_WORDS, not a lost stretch of prose
    assert _content(_run(tmp_path, changes)) == []


def test_another_rule_deleting_body_prose_is_named(tmp_path):
    ms = _doc(tmp_path)
    changes = [_delete(ms, f"{PROSE} [2].", "guide.wordy"), _delete(ms, f"{PROSE} [3].", "guide.wordy"),
               _delete(ms, "Limitations", "guide.wordy")]
    issues = _content(_run(tmp_path, changes))
    deleted = [i for i in issues if i.check == "content.deleted"]
    assert deleted and all("guide.wordy" in i.message for i in deleted)
    assert any(i.check == "content.sections" and "'Limitations'" in i.message for i in issues)
    assert any(i.check == "content.references" for i in issues)


def test_large_deletions_spread_out_count_by_rule(tmp_path):
    ms = _doc(tmp_path)
    targets = [f"{PROSE} [1].", f"{PROSE}.", f"{PROSE} [3]."]
    seen, changes = set(), []
    for p in ms.paragraphs:
        if p.text in targets and p.id not in seen and len(changes) < 4:
            seen.add(p.id)
            changes.append(Change(id=f"x{p.id}", para_id=p.id, start=0, end=len(PROSE), old=PROSE, new="",
                                  layer=Layer.LANGUAGE, source=Source.AI, reason="ai.concise", status=Status.ACCEPTED))
    issues = [i for i in _content(_run(tmp_path, changes)) if i.check == "content.deleted"]
    assert len(issues) == 1 and "ai.concise" in issues[0].message


def test_journal_drops_a_heading_by_its_sheet(tmp_path):
    ms = _doc(tmp_path)
    rep = _run(tmp_path, [_delete(ms, "Introduction", "sections_headings")], st=style.load("ijo"))
    assert _content(rep) == []
    # the same heading deleted by anything else is a loss (and submission material that took a section heading)
    rep = _run(tmp_path, [_delete(ms, "Introduction", "submission_cover_letter")], st=style.load("ijo"))
    assert sorted(i.check for i in _content(rep)) == ["content.deleted", "content.sections"]


def test_rejected_changes_are_not_losses(tmp_path):
    ms = _doc(tmp_path)
    start = next(i for i, p in enumerate(ms.paragraphs) if p.text == "Discussion")
    changes = [_delete(ms, p, "submission_reviewer_reply", Status.REJECTED) for p in ms.paragraphs[start:start + 5]]
    assert _content(_run(tmp_path, changes)) == []


def test_section_keys():
    k = preservation.heading_keys
    assert k("<H1>2. Materials and Methods:") == {"methods"} == k("Material and Methods") == k("METHODS")
    assert k("Results and Discussion") == {"results", "discussion"}
    assert k("Conclusions") == k("Conclusion") == {"conclusion"}
    assert k("Case Presentation") == {"case"}
    assert k("Financial support and sponsorship") == {"funding"}
    assert k("Conflicts of Interest") == {"conflicts"}
    assert k("The results of the trial were good.") == set()
    assert preservation.section_keys("Funding: Nil.") == {"funding"}


def test_check_without_the_original_is_unchanged(tmp_path):
    ms = _doc(tmp_path)
    assert not [i for i in qc.check(ms, style.load("jehp")).issues if i.check.startswith("content.")]
    assert set(preservation.CHECKS) <= set(qc.CHECKS)


def test_pipeline_reports_the_content_checks(tmp_path):
    _doc(tmp_path)
    out = pipeline.edit(str(tmp_path / "ms.docx"), "jehp", str(tmp_path / "out"), stages=("style",), lt_client=None)
    checks = {c["check"]: c for c in out.summary["qc"]["checks"]}
    for name in preservation.CHECKS:
        assert checks[name]["passed"], out.summary["qc"]["issues"]
