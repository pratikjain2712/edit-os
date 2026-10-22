"""Switchable behaviour added while tuning MJDRDYPU (all text synthetic): Greek-letter powers, unary signs,
doubled stops before citations, table caption periods, figure citations before the full stop, sentence-case
sub-headings, numeric range dashes, the data availability statement, and front-matter fixes."""

from docx import Document

from edit_os import docmodel, frontmatter, queries, rules, rules_figures, rules_guide, rules_sections, style, tracked
from edit_os.changes import Status

BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic over a period of two years.")
REFS = ["References", "1. Smith A, Jones B. A study. Ophthalmology 2001;12:1-9.",
        "2. Smith A, Jones B. Another study. Ophthalmology 2002;13:1-9.",
        "3. Smith A, Jones B. A third study. Ophthalmology 2003;14:1-9."]


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def load(path, first_page=()):
    ms = docmodel.load(path)
    ms.front_sources = frontmatter.Sources(first_page=list(first_page))
    return ms


def styled(code="mjdrdypu", **over):
    st = style.load(code)
    st.settings = style._merge(st.settings, over)
    return st


def run(path, st, only, first_page=()):
    cs = rules.run_all(load(path, first_page), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = path.replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs if p.text]


# --- citations ------------------------------------------------------------------------------

def test_greek_letter_power_is_not_a_citation(tmp_path):
    path = make(tmp_path, ["Introduction", "Effect sizes were large (η²H 0.46) and the χ² test was used.", BODY] + REFS)
    p = load(path).paragraphs[1]
    assert rules.superscript_citations(p) == []


def test_stop_already_before_the_citation_is_not_doubled(tmp_path):
    path = make(tmp_path, ["Introduction", "The risk rises as a person gets older. (1). It is common.",
                           "These are pathological features.[2]. On a global scale it varies."] + REFS)
    t = run(path, style.load("mjdrdypu"), ["citations"])
    assert "The risk rises as a person gets older.[1] It is common." in t
    assert "These are pathological features.[2] On a global scale it varies." in t


# --- statistics -----------------------------------------------------------------------------

def test_keep_unary_leaves_signs_after_words_closed(tmp_path):
    text = "The target was ≥400 children aged <108 months (P<0.001, n=5)."
    path = make(tmp_path, ["Methods", text] + REFS)
    t = run(path, styled(statistics={"unary_after_word": False}), ["statistics"])
    assert "The target was ≥400 children aged <108 months (P < 0.001, n = 5)." in t
    t = run(path, styled(statistics={"unary_after_word": True}), ["statistics"])
    assert "The target was ≥ 400 children aged < 108 months (P < 0.001, n = 5)." in t


# --- captions and figure citations -----------------------------------------------------------

def test_table_caption_gets_a_period_figure_legend_loses_it(tmp_path):
    assert rules_figures.caption_match(type("P", (), {"style": "", "text": "Figure 2..   "})()) is None
    path = make(tmp_path, ["Results", "See [Table 1] and [Figure 1].", BODY,
                           "Table 1. Baseline characteristics (n = 20)", "Figure 1. Axial scan of the orbit."] + REFS)
    t = run(path, styled(figures={"table_caption_end_period": "add"}), ["fig_captions"])
    assert "Table 1: Baseline characteristics (n = 20)." in t
    assert "Figure 1: Axial scan of the orbit" in t
    t = run(path, style.load("mjdrdypu"), ["fig_captions"])  # S: "The caption should end with a period"
    assert "Table 1: Baseline characteristics (n = 20)." in t
    t = run(path, style.load("abr"), ["fig_captions"])  # no such line in the ABR sheet: G200 deletes the period
    assert "Table 1: Baseline characteristics (n = 20)" in t


def test_figure_citation_moves_before_the_full_stop(tmp_path):
    path = make(tmp_path, ["Results", "All children were fully immunized. (Table 1) Most were boys.",
                           "It was not significant (P = 0.2). (Table 2)", "As Smith et al. (Figure 1) showed, it rose.",
                           "Table 1: A", "Table 2: B", "Figure 1: C"] + REFS)
    t = run(path, style.load("mjdrdypu"), ["fig_citations"])
    assert "All children were fully immunized [Table 1]. Most were boys." in t
    assert "It was not significant (P = 0.2) [Table 2]." in t
    assert "As Smith et al. [Figure 1] showed, it rose." in t
    t = run(path, styled(figures={"before_full_stop": False}), ["fig_citations"])
    assert "All children were fully immunized. [Table 1] Most were boys." in t


# --- sub-headings ------------------------------------------------------------------------------

def test_sentence_case_keeps_an_expansion_before_its_abbreviation(tmp_path):
    path = make(tmp_path, ["Borderline Personality Inventory (BPI) Scores", "Sex Differences in BPI Scores"])
    ms = load(path)
    def apply(p):
        t = p.text
        for c in sorted(rules_sections.sentence_case_changes(ms, p, "subheadings"), key=lambda c: -c.start):
            t = t[:c.start] + c.new + t[c.end:]
        return t
    assert apply(ms.paragraphs[0]) == "Borderline Personality Inventory (BPI) scores"
    assert apply(ms.paragraphs[1]) == "Sex differences in BPI scores"


def test_subheading_sentence_case_is_switchable(tmp_path):
    long = BODY + " " + BODY
    path = make(tmp_path, ["Introduction", long, "Materials and Methods", "Statistical Analysis", long,
                           "Results", long] + REFS)
    t = run(path, style.load("mjdrdypu"), ["subheadings"])
    assert "<H2>Statistical analysis" in t
    t = run(path, styled(headings={"subheading_case": "keep"}), ["subheadings"])
    assert "<H2>Statistical Analysis" in t


# --- numeric ranges ----------------------------------------------------------------------------

def test_range_dash_finder():
    def apply(t):
        for s, e, n in reversed(rules_guide.find_range_dash(t, {})):
            t = t[:s] + n + t[e:]
        return t
    assert apply("aged 18-30 years (0.68-0.91), 4 -6 days") == "aged 18–30 years (0.68–0.91), 4–6 days"
    for keep in ("COVID-19 cases", "IL-6 levels", "cited [6-9]", "on 12-05-2024", "5 - 3 = 2", "T2-weighted",
                 "a 1-year-old"):
        assert apply(keep) == keep


def test_range_dash_runs_in_tables_after_the_references(tmp_path):
    doc = Document()
    for text in ["Results", "Ages were 18-30 years.", BODY] + REFS:
        doc.add_paragraph(text)
    doc.add_paragraph("Table 1: Age groups")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "31-40"
    path = str(tmp_path / "t.docx")
    doc.save(path)
    st = style.load("mjdrdypu")
    ms = load(path)
    new = {(ms.by_id(c.para_id).in_table, c.new) for c in rules_guide.rule_guide(ms, st) if "range" in c.reason}
    assert new == {(False, "–"), (True, "–")}
    ms = load(path)
    assert not [c for c in rules_guide.rule_guide(ms, style.load("ijo")) if "range" in c.reason]


# --- back matter -------------------------------------------------------------------------------

STD = ["<H2>Financial support and sponsorship", "Nil.", "<H2>Conflicts of interest",
       "There are no conflicts of interest."]


def test_data_availability_and_contributions_from_the_title_page(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY] + STD + REFS)
    fp = ["Title page", "Data availability statement – not applicable.",
          "Author Contribution: AB was responsible for the concept, design and drafting of the manuscript."]
    t = run(path, style.load("mjdrdypu"), ["heading_tags", "sections_back_matter"], fp)
    i = t.index("<H2>Data availability statement")
    assert t[i:i + 6] == ["<H2>Data availability statement", "Not applicable.", "<H2>Author contribution",
                          "AB was responsible for the concept, design and drafting of the manuscript."] + STD[:2]
    # other journals: no data availability section
    t = run(path, style.load("ijo"), ["heading_tags", "sections_back_matter"], fp)
    assert "<H2>Data availability statement" not in t


def test_data_availability_existing_kept_missing_queried(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, "Data availability: Data are available on request."]
                + STD + REFS)
    t = run(path, style.load("mjdrdypu"), ["heading_tags", "sections_back_matter"],
            ["Data availability statement: The data are with the author."])
    assert "Data availability: Data are available on request." in t and "<H2>Data availability statement" not in t
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY] + STD + REFS, name="n.docx")
    qs = [q.text for q in queries.collect(load(path), style.load("mjdrdypu")).queries]
    assert "Please provide the missing Data availability statement, if available." in qs
    assert "Please provide the missing Data availability statement, if available." not in [
        q.text for q in queries.collect(load(path), style.load("ijo")).queries]


def test_data_availability_form_prompt_is_not_copied():
    ms = type("M", (), {"front_sources": frontmatter.Sources(first_page=[
        "Data Availability Statement:", "Please state where the data can be found, if applicable."])})()
    assert rules_sections.data_availability_from_sources(ms) == ""


# --- front matter --------------------------------------------------------------------------------

def test_department_form_normalises_the_label():
    assert frontmatter.department_form("Dept. of Psychiatry") == "Department of Psychiatry"
    assert frontmatter.department_form("Department Of Forensic Medicine") == "Department of Forensic Medicine"
    assert frontmatter.department_form("Department of Department of Pathology") == "Department of Pathology"
    assert frontmatter.department_form("Deptt. Anatomy") == "Department of Anatomy"
    assert frontmatter.department_form("Departmental Unit") == "Departmental Unit"
