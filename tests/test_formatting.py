"""Tracked formatting changes (Change.format) and the formatting rules (rules_format)."""

import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from edit_os import docmodel, queries, rules, rules_format, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

W_ID, W_AUTHOR, W_DATE = qn("w:id"), qn("w:author"), qn("w:date")


@pytest.fixture
def sample(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Smith ")
    p.add_run("et al").bold = True  # bold already: italic must be added, bold kept
    p.add_run(" found p<0.05 in vitro and colour changes in vivo [3].")
    doc.add_paragraph("Nothing to do here.")
    path = tmp_path / "f.docx"
    doc.save(path)
    return str(path)


def span(text, old, nth=0):
    s = -1
    for _ in range(nth + 1):
        s = text.index(old, s + 1)
    return s, s + len(old)


def fmt_change(cid, p, old, fmt, new=None):
    s, e = span(p.text, old)
    return Change(cid, p.id, s, e, old, old if new is None else new, Layer.HOUSE_STYLE, Source.RULE,
                  status=Status.ACCEPTED, format=fmt)


def text_change(cid, p, old, new):
    s, e = span(p.text, old)
    return Change(cid, p.id, s, e, old, new, Layer.LANGUAGE, Source.AI, status=Status.ACCEPTED)


@pytest.fixture
def changeset(sample):
    p = docmodel.load(sample).paragraphs[0]
    return ChangeSet([
        fmt_change("f1", p, "et al", {"italic": True}),                      # across a run boundary below
        fmt_change("f2", p, "p", {"italic": True}, new="P"),                  # text change + format
        fmt_change("f3", p, "in vitro", {"italic": True}),
        fmt_change("f4", p, "[3]", {"superscript": True}),
        fmt_change("f5", p, "colour changes in vivo", {"small_caps": True}),  # overlaps text change t1
        text_change("t1", p, "colour", "color"),
        Change("t2", p.id, len(p.text), len(p.text), "", " Done.", Layer.LANGUAGE, Source.AI,
               status=Status.ACCEPTED),
        fmt_change("f6", p, "Smith", {"bold": True}),
    ])


def runs_with(path, para=0):
    """(text, rPr element, inside w:ins, inside w:del) for each run of a paragraph."""
    p = list(Document(path).element.body.iter(qn("w:p")))[para]
    out = []
    for r in p.iter(qn("w:r")):
        anc = {a.tag for a in r.iterancestors()}
        txt = "".join(t.text or "" for t in r.iter(qn("w:t"), qn("w:delText")))
        out.append((txt, r.find(qn("w:rPr")), qn("w:ins") in anc, qn("w:del") in anc))
    return out


def has(rpr, tag, val=None):
    if rpr is None:
        return False
    el = rpr.find(qn(tag))
    return el is not None and (val is None or el.get(qn("w:val")) == val)


def formatted_text(path, tag, val=None, para=0):
    return "".join(t for t, rpr, ins, dele in runs_with(path, para) if not dele and has(rpr, tag, val))


def test_tracked_formatting(sample, changeset, tmp_path):
    out = tmp_path / "tracked.docx"
    tracked.render(docmodel.load(sample), changeset, tracked=True).save(out)
    assert docmodel.load(str(out)).paragraphs[0].text == \
        "Smith et al found P<0.05 in vitro and color changes in vivo [3]. Done."
    assert formatted_text(out, "w:i") == "et alPin vitro"
    assert formatted_text(out, "w:vertAlign", "superscript") == "[3]"
    assert formatted_text(out, "w:smallCaps") == "color changes in vivo"
    # Every formatted original run records its previous properties, with revision attributes.
    changed = [(t, rpr) for t, rpr, ins, dele in runs_with(out) if has(rpr, "w:rPrChange")]
    assert "et al" in [t for t, _ in changed] and "[3]" in [t for t, _ in changed]
    for t, rpr in changed:
        ch = rpr.find(qn("w:rPrChange"))
        assert ch.get(W_ID) and ch.get(W_AUTHOR) == "Edit OS" and ch.get(W_DATE)
        assert rpr[-1] is ch  # rPrChange is the last child of rPr
        old = ch.find(qn("w:rPr"))
        assert old is not None and old.find(qn("w:i")) is None and old.find(qn("w:smallCaps")) is None
        if t == "et al":
            assert has(rpr, "w:b") and has(old, "w:b")  # bold kept and recorded as before
    # Inserted text carries the new format but no format revision; deleted text keeps its run.
    ins = [(t, rpr) for t, rpr, i, d in runs_with(out) if i]
    assert ("P" in [t for t, _ in ins]) and all(not has(rpr, "w:rPrChange") for _, rpr in ins)
    assert has(dict(ins)["P"], "w:i")
    assert [t for t, _, _, d in runs_with(out) if d] == ["p", "colour"]
    # Schema order inside rPr: b before i.
    for _, rpr, _, _ in runs_with(out):
        if rpr is not None and has(rpr, "w:b") and has(rpr, "w:i"):
            tags = [c.tag for c in rpr]
            assert tags.index(qn("w:b")) < tags.index(qn("w:i"))


def test_clean_formatting(sample, changeset, tmp_path):
    out = tmp_path / "clean.docx"
    tracked.render(docmodel.load(sample), changeset, tracked=False).save(out)
    assert docmodel.load(str(out)).paragraphs[0].text == \
        "Smith et al found P<0.05 in vitro and color changes in vivo [3]. Done."
    body = Document(out).element.body
    for tag in ("w:ins", "w:del", "w:rPrChange"):
        assert not list(body.iter(qn(tag)))
    assert formatted_text(out, "w:i") == "et alPin vitro"
    assert formatted_text(out, "w:vertAlign", "superscript") == "[3]"
    assert formatted_text(out, "w:b") == "Smithet al"


def test_clearing_and_noop(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Bold italic").bold = True
    p.runs[0].italic = True
    p.add_run(" plain")
    path = tmp_path / "c.docx"
    doc.save(path)
    para = docmodel.load(str(path)).paragraphs[0]
    cs = ChangeSet([fmt_change("c1", para, "Bold italic", {"bold": False, "italic": False}),
                    fmt_change("c2", para, "plain", {"superscript": False})])  # nothing to clear
    out = tmp_path / "o.docx"
    tracked.render(docmodel.load(str(path)), cs, tracked=True).save(out)
    runs = runs_with(out)
    assert not has(runs[0][1], "w:b") and not has(runs[0][1], "w:i")
    old = runs[0][1].find(qn("w:rPrChange")).find(qn("w:rPr"))
    assert has(old, "w:b") and has(old, "w:i")
    assert len(list(Document(out).element.body.iter(qn("w:rPrChange")))) == 1  # no-op left untracked


def test_overlapping_format_changes_refused(sample):
    p = docmodel.load(sample).paragraphs[0]
    cs = ChangeSet([fmt_change("a", p, "found p", {"italic": True}), fmt_change("b", p, "p<0", {"bold": True})])
    with pytest.raises(tracked.ChangeConflict):
        tracked.render(docmodel.load(sample), cs)


def test_unknown_property_rejected():
    with pytest.raises(ValueError):
        Change("x", "p0001", 0, 1, "a", "a", Layer.HOUSE_STYLE, Source.RULE, format={"underline": True})


def test_json_round_trip(changeset, tmp_path):
    path = tmp_path / "changes.json"
    changeset.dump(str(path))
    loaded = ChangeSet.load(str(path))
    assert [c.to_dict() for c in loaded.changes] == [c.to_dict() for c in changeset.changes]
    assert loaded.changes[0].format == {"italic": True} and loaded.changes[0].format_only
    assert loaded.changes[5].format is None and not loaded.changes[5].format_only
    old = {k: v for k, v in changeset.changes[5].to_dict().items() if k != "format"}  # files from before
    assert Change.from_dict(old).format is None


def test_with_comments(sample, changeset, tmp_path):
    ms = docmodel.load(sample)
    p = ms.paragraphs[0]
    qs = queries.QuerySet()
    s, e = span(p.text, "in vitro and colour")  # overlaps a format change and a text change
    qs.queries.append(queries.Query(p.id, s, e, "Kindly check.", "test", "AQ1"))
    assert queries.write_comments(ms, qs) == 1
    out = tmp_path / "commented.docx"
    tracked.render(ms, changeset, tracked=True).save(out)
    assert docmodel.load(str(out)).paragraphs[0].text == \
        "Smith et al found P<0.05 in vitro and color changes in vivo [3]. Done."
    assert formatted_text(out, "w:i") == "et alPin vitro"
    with zipfile.ZipFile(out) as z:
        assert "Kindly check." in z.read("word/comments.xml").decode()
        doc = z.read("word/document.xml").decode()
    assert "commentRangeStart" in doc and "rPrChange" in doc


# --- rules ---------------------------------------------------------------------------

def make(tmp_path, paragraphs):
    doc = Document()
    for t in paragraphs:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return str(path)


REFS = ["References",
        "1. Lee K, Rao S, et al. Back pain in adolescents in vitro. Spine 1997;22:1380-4.",
        "2. Rao S. Eye disease, p. 12. J Eye 2001;3:4-5.",
        "3. Rao S. Lens. J Eye 2002;4:5-6."]


def base():
    st = style.load("base")
    st.settings["references"]["lookup"] = False
    return st


def proposals(path, st):
    ms = docmodel.load(path)
    return [(c.reason, c.old, c.new, c.format) for c in rules_format.rule_formatting(ms, st)], ms


def test_rule_et_al_and_latin(tmp_path):
    path = make(tmp_path, ["Smith et al. reported in vitro and In vivo results; ex-vivo too.",
                           "Jones et al found the in situ grafts stable (see https://x.org/in vitro).",
                           "Retail sales, the vitro-retinal interface and invitro are not Latin."] + REFS)
    got, _ = proposals(path, base())
    italic = {"italic": True}
    assert got == [("fmt_et_al", "et al.", "et al.", italic), ("fmt_latin", "in vitro", "in vitro", italic),
                   ("fmt_latin", "In vivo", "In vivo", italic), ("fmt_latin", "ex-vivo", "ex-vivo", italic),
                   ("fmt_et_al", "et al", "et al.", italic), ("fmt_latin", "in situ", "in situ", italic)]


def test_rule_p_values(tmp_path):
    path = make(tmp_path, ["Rates differed (p<0.05; P = .03; p-value; p value) but p. 12, pp. 3-4, "
                           "p53, 5p, Hb-p=4 and the p wave did not.", "Table note: *P < 0.001"] + REFS)
    st = base()
    got, _ = proposals(path, st)
    assert [(o, n) for r, o, n, f in got] == [("p", "p"), ("P", "P"), ("p", "p"), ("p", "p"), ("P", "P")]
    st.settings["formatting"]["p_capital"] = True
    got, _ = proposals(path, st)
    assert [(o, n) for r, o, n, f in got] == [("p", "P"), ("P", "P"), ("p", "P"), ("p", "P"), ("P", "P")]
    assert {r for r, *_ in got} == {"fmt_p_value"} and all(f == {"italic": True} for *_, f in got)
    assert not proposals(path, style.load("ijo"))[0]  # IJO: the editor leaves p alone


def test_rule_n_variable(tmp_path):
    path = make(tmp_path, ["Cases (n = 34) and controls (N=12); an=3, Fig. n, n-3 and n =x are not."] + REFS)
    got, _ = proposals(path, base())
    assert [(r, o) for r, o, n, f in got] == [("fmt_n", "n"), ("fmt_n", "N")]


def test_rule_superscript_citations(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Known before.[1] Seen again [2,3] and in [9] and [2015].")
    doc.add_paragraph("[1] A list label, not a citation.")
    p = doc.add_paragraph("Already done.")
    p.add_run("[3]").font.superscript = True
    for t in REFS:
        doc.add_paragraph(t)
    path = tmp_path / "c.docx"
    doc.save(path)
    got, _ = proposals(str(path), base())
    assert [(r, o, f) for r, o, n, f in got] == [("fmt_citations", "[1]", {"superscript": True}),
                                                ("fmt_citations", "[2,3]", {"superscript": True})]


def test_references_untouched_and_already_italic(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Smith ")
    p.add_run("et al.").italic = True
    p.add_run(" agree.")
    for t in REFS:
        doc.add_paragraph(t)
    doc.add_paragraph("Lin YJ, Wang Q, et al. A stray reference outside the list. p = 0.4")
    path = tmp_path / "r.docx"
    doc.save(path)
    assert proposals(str(path), base())[0] == []


def test_rules_switch_off(tmp_path):
    path = make(tmp_path, ["Smith et al found p<0.05 in vitro (n = 4) [1]."] + REFS)
    st = base()
    st.settings["formatting"]["enabled"] = False
    assert proposals(path, st)[0] == []


def test_citation_rule_inserts_superscript_brackets(tmp_path):
    path = make(tmp_path, ["Dermoids are rare (1). Others agree [2, 3]. Some grow (3)"] + REFS)
    cs = rules.run_all(docmodel.load(path), base(), only=["citations", "formatting"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    for is_tracked in (True, False):
        out = tmp_path / f"o{is_tracked}.docx"
        tracked.render(docmodel.load(path), cs, tracked=is_tracked).save(out)
        assert docmodel.load(str(out)).paragraphs[0].text == "Dermoids are rare.[1] Others agree.[2,3] Some grow[3]"
        assert formatted_text(out, "w:vertAlign", "superscript") == "[1][2,3][3]"


def test_style_run_end_to_end(tmp_path):
    path = make(tmp_path, ["Introduction", "Smith et al found p<0.05 in vitro [1].", "Conclusion"] + REFS)
    cs = rules.run_all(docmodel.load(path), base())
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "t.docx"
    ms = docmodel.load(path)
    queries.write_comments(ms, queries.collect(ms, base()))
    tracked.render(ms, cs, tracked=True).save(out)
    i = [p.text for p in docmodel.load(str(out)).paragraphs].index("Smith et al. found p < 0.05 in vitro.[1]")
    assert formatted_text(out, "w:i", para=i) == "et al.pin vitro"
    assert formatted_text(out, "w:vertAlign", "superscript", para=i) == "[1]"


def test_citation_move_is_one_suggestion(tmp_path):
    """A citation moved after the full stop carries its superscript itself: no separate
    formatting suggestion on the text it replaces."""
    path = make(tmp_path, ["Rates were low [1]. Others agree [2]."] + REFS)
    cs = rules.run_all(docmodel.load(path), base())
    para = [c for c in cs.changes if c.para_id == "p0001"]
    assert para and not any(c.format_only for c in para)
    assert all((c.format or {}).get("superscript") for c in para if "[" in c.new)
