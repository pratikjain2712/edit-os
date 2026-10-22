"""Rebuilt reference entries: what the record lacks or garbles must not reach the output (synthetic records;
no HTTP)."""

import pytest

from edit_os import biblio, journals, refengine, rules, style


@pytest.fixture
def jehp():
    return style.load("jehp")


@pytest.fixture
def no_smart_quotes():
    st = style.load("jehp")
    st.settings.setdefault("guide_replace", {})["smart_quotes"] = False
    return st


@pytest.fixture
def small_dictionary(monkeypatch):
    monkeypatch.setattr(refengine, "_dictionary", ({"plant", "species", "management", "recurrent", "after"}, set()))


def _rec(**kw):
    base = dict(source="pubmed", title="", authors=[("Onnink", "MLA"), ("Teunissen", "LBA")], journal="Disabil Rehabil",
                year="2022", volume="44", issue="3", pages="101-9")
    base.update(kw)
    return biblio.Record(**base)


# 1. Curly quotes: the record's straight apostrophe never replaces the author's curly one.

def test_curly_apostrophe_kept_from_author(no_smart_quotes):
    text = ("Onnink MLA, Teunissen LBA. Experts’ perspectives on the impact of visual impairment. "
            "Disabil Rehabil. 2022;44(3):101-9.")
    p = refengine.parse(text)
    rec = _rec(title="Experts' perspectives on the impact of visual impairment.")
    assert "Experts’ perspectives" in refengine.format_record(rec, no_smart_quotes, p)


def test_journal_quote_convention_applied_to_record_title(jehp):
    # JEHP makes straight quotes curly (guide_replace.smart_quotes): a rebuilt title follows it.
    text = "Onnink MLA, Teunissen LBA. A different title. Disabil Rehabil. 2022;44(3):101-9."
    rec = _rec(title='On the theme of "Patient Safety": mothers\' views.')
    out = refengine.format_record(rec, jehp, refengine.parse(text))
    assert "“Patient Safety”: Mothers’ views" in out or "“Patient Safety”: mothers’ views" in out
    assert "'" not in out and '"' not in out


def test_straight_quotes_stay_where_journal_and_author_use_them(no_smart_quotes):
    text = "Onnink MLA, Teunissen LBA. Another title. Disabil Rehabil. 2022;44(3):101-9."
    rec = _rec(title="Experts' perspectives on impairment.")
    assert "Experts' perspectives" in refengine.format_record(rec, no_smart_quotes, refengine.parse(text))


def test_author_dashes_and_letters_kept_when_same_words(jehp):
    text = ("Asiksoy G, Teunissen LBA. Nurses’ assessment of artificial intelligence chatbots in years—Autism and "
            "Developmental Disabilities. Disabil Rehabil. 2022;44(3):101-9.")
    rec = _rec(title="Nurses' Assessment of Artificial ıntelligence Chatbots in Years - Autism and Developmental "
                     "Disabilities.", authors=[("Asiksoy", "G"), ("Teunissen", "LBA")])
    out = refengine.format_record(rec, jehp, refengine.parse(text))
    # (the author's capitals in "Autism and Developmental Disabilities", a name, stay)
    assert "Nurses’ assessment of artificial intelligence chatbots in years—Autism and Developmental" in out


def test_author_sentence_case_wins_over_record_title_case(jehp):
    text = ("Martin SS, Aday AW. 2024 heart disease and stroke statistics: a report from the American Heart "
            "Association. Disabil Rehabil. 2022;44(3):101-9.")
    rec = _rec(title="2024 Heart Disease and Stroke Statistics: A Report From the American Heart Association.",
               authors=[("Martin", "SS"), ("Aday", "AW")])
    out = refengine.format_record(rec, jehp, refengine.parse(text))
    assert "2024 heart disease and stroke statistics: A report from the American Heart Association." in out


def test_crossref_markup_stripped():
    rec = biblio._crossref_record({"title": ["Adapting <scp>parent‐focused</scp> interventions"],
                                   "container-title": ["Journal of Policy and Practice"]})
    assert rec.title == "Adapting parent‐focused interventions"


# 2. Genus names keep their capital in sentence case.

@pytest.mark.parametrize("title, expect", [
    ("Phytopharmacological Profile of Jasminum grandiflorum Linn. (Oleaceae).", "Jasminum grandiflorum Linn."),
    ("Potential of Jasminum Species: A Herbal Approach.", "Potential of Jasminum species"),
    ("Comparative Efficacy of Jasminum Grandiflorum in Management of Recurrent Aphthous Stomatitis.",
     "efficacy of Jasminum grandiflorum in management of recurrent aphthous stomatitis"),
    ("Effect of Jasminum sambac Flowers on Lactation.", "Effect of Jasminum sambac flowers on lactation"),
    ("Plant Species Richness after Fire.", "Plant species richness"),
])
def test_genus_capital_kept(jehp, small_dictionary, title, expect):
    rec = _rec(title=title)
    assert expect in refengine.format_record(rec, jehp)


def test_hyphen_parts_lower_cased(jehp):
    for title, expect in [("Evidence-Based Implementation of Care.", "Evidence-based implementation"),
                          ("Feeding of 6- to 23-Month-Old Children.", "6- to 23-month-old children")]:
        assert expect in refengine.format_record(_rec(title=title), jehp)


# 3. Journal names of journals NLM does not list.

class _Dir:
    def __init__(self, known):
        self.known = known

    def lookup(self, name):
        return self.known.get(name.lower())


def test_full_journal_name_drops_publisher_acronym(jehp, monkeypatch):
    monkeypatch.setattr(journals, "directory", lambda: _Dir({}))
    rec = _rec(source="crossref", journal="IJPHS", journal_full="International Journal of Public Health Science (IJPHS)")
    assert refengine.journal_name(rec, jehp) == "International Journal of Public Health Science"
    rec = _rec(source="crossref", journal="Comput. Educ. Artif. Intell.",
               journal_full="Computers and Education: Artificial Intelligence")
    assert refengine.journal_name(rec, jehp) == "Computers and Education: Artificial Intelligence"


def test_short_title_match_must_be_the_same_journal(jehp, monkeypatch):
    la_pie = journals.Journal(nlm_id="1", title="La Pie", med_abbr="Pie", iso_abbr="Pie", issn_print="", issn_online="")
    monkeypatch.setattr(journals, "directory", lambda: _Dir({"pie": la_pie}))
    rec = _rec(source="crossref", journal="pie", journal_full="Perspectives in Education")
    assert refengine.journal_name(rec, jehp) == "Perspectives in Education"
    rec = _rec(source="crossref", journal="pie", journal_full="La Pie")
    assert refengine.journal_name(rec, jehp) == "Pie"


def test_one_word_author_is_no_group(jehp):
    rec = _rec(authors=[("Mustafa", "SR"), ("Nurdiansyah", "")], title="A study.")
    assert refengine.format_record(rec, jehp).startswith("Mustafa SR, Nurdiansyah. ")


# 4. / 5. Pages and issues the record lacks.

def _p(cite):
    return refengine.parse(f"Simak VF, Rekawaty E. The effectiveness of programs. Int J Public Health Sci. {cite}.")


def test_record_without_pages_keeps_author_pages(jehp):
    rec = _rec(authors=[("Simak", "VF"), ("Rekawaty", "E")], title="The effectiveness of programs.",
               volume="6", issue="3", pages="", year="2018")
    assert refengine.format_record(rec, jehp, _p("2018;6(3):162-168")).endswith(" 2018;6:162-8.")


def test_record_first_page_only_keeps_author_range(jehp):
    rec = _rec(authors=[("Simak", "VF"), ("Rekawaty", "E")], title="The effectiveness of programs.",
               volume="11", pages="54", year="2022")
    assert refengine.format_record(rec, jehp, _p("2022;11(1):54-60")).endswith(" 2022;11:54-60.")
    # one page written as a range stays one page
    rec.pages = "e68918"
    assert refengine.format_record(rec, jehp, _p("2022;11:e68918-e68918")).endswith(" 2022;11:e68918.")


def test_author_pages_of_another_volume_not_used_and_queried(jehp, monkeypatch):
    rec = _rec(authors=[("Simak", "VF"), ("Rekawaty", "E")], title="The effectiveness of programs.",
               volume="11", pages="", year="2022", source="crossref", journal="Int J Public Health Sci")
    p = _p("2022;12(1):1-6")
    assert refengine.format_record(rec, jehp, p).endswith(" 2022;11.")
    monkeypatch.setattr(refengine, "find", lambda parsed: (rec, ""))
    res = refengine.resolve(p.text, jehp)
    assert ("references.no_pages", {}) in res.queries


@pytest.mark.parametrize("issue, volume, expect", [
    ("Supplement_1", "49", "(Suppl 1)"), ("2 Suppl", "21", "(2 Suppl)"), ("s2", "13", "(s2)"),
    ("", "45 Suppl 1", "(Suppl 1)"), ("3", "6", ""),
])
def test_supplements_kept_as_written(jehp, issue, volume, expect):
    rec = _rec(issue=issue, volume=volume, title="A study.")
    vol = volume.split()[0]
    assert f" 2022;{vol}{expect}:101-9." in refengine.format_record(rec, jehp)


@pytest.mark.parametrize("pages, expect", [("e123-e130", "e123-30"), ("123-123", "123"), ("e5-e5", "e5"),
                                           ("1441-1446", "1441-6"), ("220-235", "220-35"), ("54", "54")])
def test_format_pages(jehp, pages, expect):
    assert refengine.format_pages(pages, jehp) == expect


def test_page_rule_on_unmatched_entries(jehp, tmp_path):
    from docx import Document

    from edit_os import docmodel

    d = Document()
    d.add_paragraph("References")
    d.add_paragraph("1. Smith J. A title here. Some J. 2020;12:e123-e130.")
    d.add_paragraph("2. Smith J. A title here. Some J. 2020;12:123-123.")
    path = tmp_path / "m.docx"
    d.save(path)
    ms = docmodel.load(str(path))
    news = sorted(c.new for c in rules.rule_references(ms, jehp) if c.reason == "ref_pages")
    assert news == [":123", ":e123-30"]


def test_prime_apostrophe_and_group_star(jehp):
    text = ("Gbratto-Dobe SAW, Segnon HB. Is mother’s education essential in Côte d’Ivoire? "
            "Disabil Rehabil. 2022;44(3):101-9.")
    rec = _rec(title="Is mother's education essential in Côte d′Ivoire?",
               authors=[("Gbratto-Dobe", "SAW"), ("Segnon", "HB")])
    assert "Is mother’s education essential in Côte d’Ivoire? " in refengine.format_record(rec, jehp,
                                                                                            refengine.parse(text))
    rec = _rec(title="Standards of care.", authors=[("American Diabetes Association Professional Practice Committee*", "")])
    assert refengine.format_record(rec, jehp).startswith("American Diabetes Association Professional Practice Committee. ")
