"""Submission-only parts inside a manuscript (edit_os/submission.py). All text is synthetic."""

from docx import Document

from edit_os import docmodel, frontmatter, pipeline, rules, style, submission, tracked
from edit_os.changes import Status

BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic.")
REFS = ["References", "1. Smith A, Jones B. A study. Ophthalmology 2001;12:1-9."]
NOTICE = ("The author(s) Dr Asha K Ramaswamy, Dr Ravi Pillai has/have not agreed to the copyright terms and "
          "conditions which were sent on the author's email address.")
LETTER = [
    "Cover Letter",
    "To,\nThe Editor-in-Chief,\nJournal of Eye Studies\n\nSubject: Submission of Case Report",
    "We are pleased to submit our manuscript titled ‘Lens dislocation in a child’ for consideration as a "
    "Case Report in your journal.",
    "1. Data Access and Author Responsibility",
    "As the corresponding author I affirm that I have full access to all the data.",
    "2. Author Contributions:",
    "Dr Asha K Ramaswamy – Concept, Design, Manuscript Review.",
    "Dr Ravi Pillai – Literature Search, Data Acquisition, Manuscript Preparation.",
    "Thank you for considering our work.",
    "Sincerely,\nDr. Ravi Pillai\nCity Eye Hospital, Pune, India\nEmail: ravi@example.org",
]
TITLE_PAGE = [
    "A, et al.: Lens Dislocation in a Child: A Rare Case",
    "Case Report",
    "Running Title: Lens dislocation in a child",
    "Authors: Dr. Asha K A¹, Dr. Ravi Pillai²",
    "¹Department of Ophthalmology, City Eye Hospital, Pune, India",
    "Corresponding Author:\nDr. Ravi Pillai",
    " | Phone: +91-9000000000",
    "Number of Figures: 2\nNumber of Tables: 1",
    "Key Learning Points",
    "• Lens dislocation is rare in children.",
]
REPLY_HEAD = ["Reply to the reviewers’ comments", "We thank the reviewer for the careful reading of the manuscript."]
REPLY_ROWS = [("Reviewer Number", "Original comments of the reviewer", "Reply by the author(s)"),
              ("1", "Please add the patient consent statement.", "Added on page 9.")]
ARTICLE = ["Title – Lens dislocation in a child: a rare case", "Abstract", BODY, "Introduction", BODY,
           "Discussion", BODY, "Conclusion", BODY]


def make(tmp_path, blocks, name="m.docx"):
    """blocks: strings are paragraphs, lists of tuples are tables."""
    doc = Document()
    for b in blocks:
        if isinstance(b, list):
            t = doc.add_table(rows=len(b), cols=len(b[0]))
            for i, row in enumerate(b):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
        else:
            doc.add_paragraph(b)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def merged(tmp_path, name="m.docx", contributions=True):
    back = ["Dr Asha K Ramaswamy – Concept, Design, Manuscript Review.",
            "Dr Ravi Pillai – Literature Search, Data Acquisition, Manuscript Preparation."] if contributions else []
    return make(tmp_path, [NOTICE] + LETTER + TITLE_PAGE + REPLY_HEAD + [REPLY_ROWS] + ARTICLE + back + REFS, name)


def kinds(ms, code="ijo"):
    found = submission.detect(ms, style.load(code))
    return {(p.kind, p.action) for p in found.parts}, found


def texts(part):
    return [p.text for p in part.paras if p.text.strip()]


def test_merged_submission_parts_found(tmp_path):
    ms = docmodel.load(merged(tmp_path))
    got, found = kinds(ms)
    assert got == {("copyright_notice", "aq"), ("cover_letter", "delete"), ("reviewer_reply", "delete"),
                   ("phone", "delete"), ("counts", "delete"), ("title_repeat", "delete")}
    by = {p.kind: p for p in found.parts}
    assert texts(by["cover_letter"]) == LETTER  # heading to the signature, nothing of the title page
    assert texts(by["reviewer_reply"]) == REPLY_HEAD + [c for row in REPLY_ROWS for c in row]
    assert texts(by["title_repeat"]) == [ARTICLE[0]]
    assert not found.warnings


def test_rule_deletes_parts_and_queries_copyright_notice(tmp_path):
    path = merged(tmp_path)
    cs = rules.run_all(docmodel.load(path), style.load("ijo"), only=["submission"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    left = [p.text for p in docmodel.load(str(out)).paragraphs if p.text.strip()]
    assert left[0] == "AQ: " + NOTICE
    gone = LETTER[:6] + LETTER[8:] + REPLY_HEAD + ["Reviewer Number", ARTICLE[0], TITLE_PAGE[6].strip(), TITLE_PAGE[7]]
    assert not any(t in left for t in gone)
    assert left.count(LETTER[6]) == 1  # the back matter's copy of the contribution stays
    # the article, the title page and the key points stay
    for t in TITLE_PAGE[:6] + TITLE_PAGE[8:] + ARTICLE[1:]:
        assert t.strip() in [x.strip() for x in left], t


def test_front_matter_reads_title_page_after_cover_letter(tmp_path):
    ms = docmodel.load(merged(tmp_path))
    rules.run_all(ms, style.load("ijo"), only=["submission"])
    lead = [p.text for p in frontmatter.leading_paragraphs(ms)]
    assert lead[0] == NOTICE and lead[1] == TITLE_PAGE[0]  # the cover letter is skipped
    assert "Key Learning Points" not in lead and TITLE_PAGE[6] not in lead and LETTER[2] not in lead
    assert submission.notice_names(ms) == [("Asha K", "Ramaswamy"), ("Ravi", "Pillai")]


def test_running_head_line_and_label_with_content_roles():
    lead = [docmodel.Paragraph(f"p{i}", None, "", t) for i, t in
            enumerate(["Rao, et al.: Lens dislocation", "Corresponding Author:\nDr. Ravi Pillai", "Authors:"])]
    f = frontmatter.Fields(roles={"p0": "title", "p1": "label", "p2": "label"})
    frontmatter._fix_roles(f, lead)
    assert f.roles == {"p0": "running_title", "p1": "correspondence", "p2": "label"}


def test_pipeline_contributions_from_cover_letter_and_nothing_else_in_deleted_paragraphs(tmp_path, monkeypatch):
    from edit_os.copyedit import langtool
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    path = merged(tmp_path, contributions=False)
    out = pipeline.edit(path, st, str(tmp_path / "out"), stages=("style",), accept_all=True)
    man, changes, _, _ = pipeline.load_records(str(tmp_path / "out"))
    ms = docmodel.load(path)
    letter_ids = {p.id for p in ms.paragraphs if p.text in LETTER}
    on_letter = [c for c in changes.changes if c.para_id in letter_ids]
    assert on_letter and all(c.reason == "submission_cover_letter" for c in on_letter)  # no "<H2>" on "2. Author ..."
    clean = [p.text for p in docmodel.load(out.path("clean")).paragraphs if p.text.strip()]
    assert not any(t in clean for t in LETTER[:6] + LETTER[8:])
    i = clean.index("<H2>Authors’ contributions")  # the only copy, read from the deleted letter
    assert clean[i + 1:i + 3] == LETTER[6:8]
    assert not out.summary["render_conflicts"]


def test_letter_to_the_editor_is_not_a_cover_letter(tmp_path):
    path = make(tmp_path, ["Letter to the Editor", "Dear Editor,", "We read with interest the article submitted by Rao "
                           "et al. on lens dislocation.", BODY, "Yours sincerely,", "Dr. Ravi Pillai"] + REFS)
    got, found = kinds(docmodel.load(path))
    assert not got and not found.warnings
    # addressed like a cover letter but nothing follows it: only a warning
    path = make(tmp_path, ["To,\nThe Editor,\nJournal of Eye Studies", "We submit this letter for publication.", BODY,
                           "Yours sincerely,", "Dr. Ravi Pillai"] + REFS, "l2.docx")
    got, found = kinds(docmodel.load(path))
    assert not got and [w.kind for w in found.warnings] == ["cover_letter"]


def test_unfinished_cover_letter_and_reply_are_warnings(tmp_path):
    path = make(tmp_path, ["Cover Letter", "We submit our manuscript for consideration.", "Abstract", BODY] + REFS)
    got, found = kinds(docmodel.load(path))
    assert not got and [w.kind for w in found.warnings] == ["cover_letter"]
    path = make(tmp_path, ["Response to reviewers", "Reviewer 1: the methods are unclear.", "We clarified them.",
                           BODY, BODY] + REFS[1:], "r.docx")
    got, found = kinds(docmodel.load(path))
    assert not got and [w.kind for w in found.warnings] == ["reviewer_reply"]


def test_reply_after_references_and_data_tables_kept(tmp_path):
    data = [("Group", "n"), ("Cases", "12")]
    path = make(tmp_path, ["Abstract", BODY, "Discussion", BODY] + REFS + ["Table 1: Patients", data]
                + REPLY_HEAD + [REPLY_ROWS] + ["Reviewer 2", REPLY_ROWS])
    ms = docmodel.load(path)
    got, found = kinds(ms)
    assert got == {("reviewer_reply", "delete")}
    gone = submission.deleted_ids(ms, style.load("ijo"))
    assert {p.text for p in ms.paragraphs if p.id in gone} >= {"Reviewer 2", "Reviewer Number", REPLY_HEAD[0]}
    assert not any(p.id in gone for p in ms.paragraphs if p.text in ("Group", "Cases", "Table 1: Patients", BODY))
    # a reply in running text up to the title page
    path = make(tmp_path, ["Reviewer Comments: Reviewer #1:", "1. Please clarify the dose.", "We clarified the dose.",
                           "Title: Lens dislocation in a child", "Abstract", BODY] + REFS, "t.docx")
    got, found = kinds(docmodel.load(path))
    assert got == {("reviewer_reply", "delete")}
    assert texts(found.parts[0])[-1] == "We clarified the dose."


def test_title_repeat_needs_a_title_above_and_counts_need_a_count(tmp_path):
    path = make(tmp_path, ["Title: Lens dislocation in a child: a rare case", "Abstract", BODY,
                           "Word count: 2173 (body, excluding abstract and references)",
                           "Figures: 6 (Figures 1–6).", "Introduction",
                           "Word count can be reduced in this section, the reviewer said.", BODY] + REFS)
    ms = docmodel.load(path)
    got, found = kinds(ms)
    assert got == {("counts", "delete")}
    assert [t for part in found.parts for t in texts(part)] == [
        "Word count: 2173 (body, excluding abstract and references)", "Figures: 6 (Figures 1–6)."]


def test_phone_lines_per_journal(tmp_path):
    path = make(tmp_path, ["Lens dislocation in a child", "Dr Ravi Pillai", "Phone: +91 90000 00000",
                           "ORCID: 0000-0002-1825-0097", "Abstract", BODY] + REFS)
    assert kinds(docmodel.load(path), "ijo")[0] == {("phone", "delete"), ("orcid", "delete")}
    assert kinds(docmodel.load(path), "njcp")[0] == {("phone", "delete")}
    assert kinds(docmodel.load(path), "neurol")[0] == set()


def test_title_page_file_without_cover_letter():
    first = ["Title: Lens dislocation in a child", "Corresponding author:", "Dr. Ravi Pillai",
             "Covering letter", "To", "The Editor-in-Chief", "Dear Sir/Madam,",
             "We are pleased to submit our manuscript entitled “Lens dislocation”.", "Sincerely,\nDr. Ravi Pillai"]
    assert submission.texts_without_parts(first) == first[:3]


def test_paragraph_with_an_image_is_never_deleted(tmp_path):
    import base64
    png = tmp_path / "dot.png"
    png.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="))
    path = merged(tmp_path)
    doc = Document(path)
    target = next(p for p in doc.paragraphs if p.text == LETTER[2])
    target.add_run().add_picture(str(png))
    doc.save(path)
    ms = docmodel.load(path)
    cs = rules.run_all(ms, style.load("ijo"), only=["submission"])
    pid = next(p.id for p in ms.paragraphs if p.text.startswith(LETTER[2][:30]))
    assert pid not in {c.para_id for c in cs.changes}
    assert any(w["para_id"] == pid and "image" in w["message"] for w in submission.warnings(ms, style.load("ijo")))


def test_copyright_notice_prefix_per_journal(tmp_path):
    doc = Document()
    for t in [NOTICE, "Abstract", BODY, "Introduction", BODY] + REFS:
        doc.add_paragraph(t)
    path = tmp_path / "n.docx"
    doc.save(path)
    new = [c.new for c in rules.run_all(docmodel.load(str(path)), style.load("jehp"), only=["submission"]).changes]
    assert new == ["Query: "]
    new = [c.new for c in rules.run_all(docmodel.load(str(path)), style.load("ijo"), only=["submission"]).changes]
    assert new == ["AQ: "]


def test_reply_typed_inside_the_article_stops_at_the_next_section(tmp_path):
    # "Reviewer Comment: … Response: …" left between the Methods and the Discussion: only the exchange goes,
    # never the Discussion, Conclusion and declarations after it (a JEHP file lost them all)
    exchange = ["Reviewer Comment:", "Please add the CONSORT diagram at the end of the intervention section.",
                "Response:", "Thank you. The CONSORT diagram has been added at the end of the Intervention section."]
    path = make(tmp_path, ["Abstract", BODY, "Methods", BODY] + exchange
                + ["Discussion", BODY, "Limitations", BODY, "Conclusion", BODY, "Declaration", BODY] + REFS)
    ms = docmodel.load(path)
    gone = submission.deleted_ids(ms, style.load("jehp"))
    lost = [p.text for p in ms.paragraphs if p.id in gone]
    assert "Discussion" not in lost and "Conclusion" not in lost and "Declaration" not in lost
    assert set(lost) <= set(exchange)
