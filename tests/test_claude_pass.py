"""Claude copy-edit pass (B4), offline: synthetic text and a fake Claude."""

import json

import pytest
from docx import Document

from edit_os import copyedit, docmodel, style
from edit_os.changes import Change, Layer, Source
from edit_os.copyedit import claude_pass


@pytest.fixture(autouse=True)
def private_cache(tmp_path, monkeypatch):
    monkeypatch.setattr(claude_pass, "_cache", claude_pass._Cache(str(tmp_path / "ce_cache.sqlite")))
    monkeypatch.delenv("EDITOS_CE_EXAMPLES", raising=False)
    monkeypatch.delenv("EDITOS_CE_MODEL", raising=False)
    monkeypatch.delenv("EDITOS_MODEL", raising=False)
    claude_pass.DROPS.clear()


def make(tmp_path, paras):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return docmodel.load(str(path))


S1 = "The patients was followed up for six months after the surgery in all of the centres."
S2 = "However the difference between the two groups was small and it was seen in 12 patients [3]."


def edit(pid, old, new, kind="agreement", confidence="certain", reason="why"):
    return dict(paragraph_id=pid, old=old, new=new, kind=kind, reason=reason, confidence=confidence)


class FakeClaude:
    def __init__(self, answer):
        self.answer, self.calls = answer, []

    def __call__(self, params):
        self.calls.append(params)
        return self.answer


def para(tmp_path, text):
    return make(tmp_path, [text]).paragraphs[0]


CFG = {"max_words": 6, "max_chars": 60}


# --- structured output -------------------------------------------------------------------------

def test_parse_keeps_wellformed_items_only():
    text = json.dumps({"edits": [edit("P1", "a", "b"), {"paragraph_id": "P1", "old": 3}, "junk"],
                       "queries": [{"paragraph_id": "P2", "anchor": "x", "text": "Please check."}, {"text": 1}]})
    got = claude_pass.parse(text)
    assert [e["old"] for e in got["edits"]] == ["a"]
    assert len(got["queries"]) == 1
    assert claude_pass.parse("not json") == {"edits": [], "queries": []}
    assert claude_pass.parse("[1, 2]") == {"edits": [], "queries": []}


def test_request_puts_stable_prompt_first_and_caches_it():
    req = claude_pass.request("INSTR", "EXAMPLES", "[P1] text", "claude-haiku-5-5", "low")
    assert [b["text"] for b in req["system"]] == ["INSTR", "EXAMPLES"]
    assert req["system"][-1]["cache_control"] == {"type": "ephemeral"}
    assert req["messages"] == [{"role": "user", "content": "[P1] text"}]
    assert req["output_config"]["format"]["type"] == "json_schema"
    assert req["output_config"]["format"]["schema"] is claude_pass.SCHEMA


# --- verification ------------------------------------------------------------------------------

def test_trim_to_changed_words():
    t = claude_pass.trim(S1, 4, "patients was followed", "patients were followed")
    assert (S1[t.start:t.end], t.old, t.new) == ("was", "was", "were")
    s = "However the difference was small."
    t = claude_pass.trim(s, 0, "However the", "However, the")
    assert (t.start, t.end, t.old, t.new) == (7, 7, "", ",")


def test_verify_makes_minimal_ai_change(tmp_path):
    p = para(tmp_path, S1)
    c, why = claude_pass.verify(p, edit("P1", "patients was followed", "patients were followed"), CFG)
    assert why == ""
    assert (c.old, c.new, p.text[c.start:c.end]) == ("was", "were", "was")
    assert c.layer == Layer.LANGUAGE and c.source == Source.AI
    assert c.reason == "claude.agreement (why)"


@pytest.mark.parametrize("old,new,why", [
    ("patients were", "patients was", "not_found"),       # old not in the paragraph
    ("the", "a", "not_unique"),                            # occurs twice
    ("six months", "6 months", "digits"),                  # numbers untouched
    ("for six months after the surgery in all of the", "for half a year after surgery in all the", "too_long"),
    ("centres", "centers", "spelling_variant"),            # the spelling checker owns the convention
    ("patients was", "patients was", "no_change"),
])
def test_verify_drops_unsafe(tmp_path, old, new, why):
    p = para(tmp_path, S1)
    c, got = claude_pass.verify(p, edit("P1", old, new), CFG)
    assert c is None and got == why


@pytest.mark.parametrize("text,old,new,why", [
    (S2, "patients [3]", "patients", "digits"),                                    # citation
    ("Values are mean ± SD unless stated otherwise here.", "mean ± SD", "mean SD", "symbols"),
    ("Levels of CRP and ESR were high in most of them.", "CRP and ESR", "C-reactive protein and ESR", "acronym"),
    ("The scale was developed by Kumar and colleagues in Delhi.", "by Kumar and", "by the", "name"),
    ("See https://example.org/a_b for the form we used.", "org/a_b for", "org/ab for", "protected"),
    ("The diagnoses were confirmed in each case by biopsy.", "nfirmed", "mfirmed", "partial_word"),
])
def test_verify_drops_protected_content(tmp_path, text, old, new, why):
    p = para(tmp_path, text)
    c, got = claude_pass.verify(p, edit("P1", old, new), CFG)
    assert c is None and got == why


def test_verify_confidence_and_skipped_kinds(tmp_path):
    p = para(tmp_path, S1)
    e = edit("P1", "patients was", "patients were", confidence="likely")
    assert claude_pass.verify(p, e, CFG) == (None, "not_certain")
    assert claude_pass.verify(p, e, {**CFG, "min_confidence": "likely"})[0] is not None
    e = edit("P1", "patients was", "patients were", kind="word_choice")
    assert claude_pass.verify(p, e, {**CFG, "skip_kinds": ["word_choice"]}) == (None, "skipped_kind")


# --- the pass ----------------------------------------------------------------------------------

def test_collect_maps_local_ids_and_anchors_queries(tmp_path):
    ms = make(tmp_path, [S1, S2])
    fake = FakeClaude({
        "edits": [edit("P1", "patients was", "patients were"),
                  edit("P2", "However the", "However, the", kind="punctuation"),
                  edit("P9", "x", "y")],
        "queries": [{"paragraph_id": "P2", "anchor": "it was seen", "text": "Kindly clarify what 'it' refers to."},
                    {"paragraph_id": "P1", "anchor": "x", "text": "Second query, over the limit."}],
    })
    changes, queries = claude_pass.collect(ms, style.load("ijo"), call=fake)
    assert len(fake.calls) == 1
    assert "[P1] " + S1 in fake.calls[0]["messages"][0]["content"]
    assert [(c.para_id, c.old, c.new) for c in changes] == [(ms.paragraphs[0].id, "was", "were"),
                                                          (ms.paragraphs[1].id, "", ",")]
    assert claude_pass.DROPS["bad_paragraph"] == 1
    assert len(queries) == 1  # max_queries_per_chunk: 1
    q = queries[0]
    assert ms.paragraphs[1].text[q.start:q.end] == "it was seen" and q.check == "claude.query"


def test_run_all_rules_win_overlaps(tmp_path):
    ms = make(tmp_path, [S1])
    p = ms.paragraphs[0]
    s = p.text.index("was")
    rule = Change(id="r1", para_id=p.id, start=s, end=s + 3, old="was", new="is",
                  layer=Layer.LANGUAGE, source=Source.RULE)
    fake = FakeClaude({"edits": [edit("P1", "patients was", "patients were"),
                                 edit("P1", "after the surgery", "after surgery", kind="article")],
                       "queries": []})
    res = copyedit.run_all(ms, style.load("ijo"), only=["claude"], existing=[rule], claude_call=fake)
    assert [(c.old, c.new) for c in res.changes.changes] == [("the ", "")]


def test_off_without_key(tmp_path):
    # conftest makes ai_refs.available() False: the pass is skipped and Claude is never called.
    ms = make(tmp_path, [S1])
    assert not claude_pass.enabled(style.load("ijo"))
    assert claude_pass.collect(ms, style.load("ijo")) == ([], [])


# --- caching -----------------------------------------------------------------------------------

def test_cache_key_is_chunk_text_not_paragraph_ids(tmp_path):
    k = claude_pass._Cache.key
    assert k("m", "low", "sys", "[P1] a") == k("m", "low", "sys", "[P1] a")
    assert len({k("m", "low", "sys", "[P1] a"), k("m2", "low", "sys", "[P1] a"),
                k("m", "high", "sys", "[P1] a"), k("m", "low", "sys2", "[P1] a"),
                k("m", "low", "sys", "[P1] b")}) == 5
    # The same text later in another manuscript (other paragraph ids) is answered from the cache.
    fake = FakeClaude({"edits": [edit("P1", "patients was", "patients were")], "queries": []})
    st = style.load("ijo")
    first, _ = claude_pass.collect(make(tmp_path, [S1]), st, call=fake)
    other = make(tmp_path, ["Title of the article", S1])  # the heading is not sent
    second, _ = claude_pass.collect(other, st, call=fake)
    assert len(fake.calls) == 1
    assert second[0].para_id == other.paragraphs[1].id and second[0].new == "were"


def test_batch_prefetch_fills_cache(tmp_path):
    ms = make(tmp_path, [S1])
    st = style.load("ijo")
    answer = json.dumps({"edits": [edit("P1", "patients was", "patients were")], "queries": []})

    class Obj:
        def __init__(self, **kw):
            self.__dict__.update(kw)

    class Batches:
        def create(self, requests):
            self.requests = requests
            return Obj(id="b1")

        def retrieve(self, bid):
            return Obj(processing_status="ended")

        def results(self, bid):
            usage = Obj(input_tokens=10, output_tokens=5, cache_creation_input_tokens=0, cache_read_input_tokens=0)
            msg = Obj(stop_reason="end_turn", usage=usage, content=[Obj(type="text", text=answer)])
            return [Obj(custom_id=r["custom_id"], result=Obj(type="succeeded", message=msg)) for r in self.requests]

    batches = Batches()
    client = Obj(messages=Obj(batches=batches))
    assert claude_pass.batch_prefetch([(ms, st)], client=client, poll=0) == "b1"
    assert len(batches.requests) == 1 and len(batches.requests[0]["custom_id"]) == 64
    fake = FakeClaude({"edits": [], "queries": []})
    changes, _ = claude_pass.collect(ms, st, call=fake)
    assert not fake.calls and [c.new for c in changes] == ["were"]
    assert claude_pass.batch_prefetch([(ms, st)], client=client) is None  # nothing left to send


# --- few-shot examples -------------------------------------------------------------------------

def test_examples_come_from_training_articles_only(tmp_path):
    train = next(f"xyz_{i}_26" for i in range(100) if not claude_pass.is_test_article(f"xyz_{i}_26"))
    test = next(f"xyz_{i}_26" for i in range(100) if claude_pass.is_test_article(f"xyz_{i}_26"))
    other = next(f"abc_{i}_26" for i in range(100) if not claude_pass.is_test_article(f"abc_{i}_26"))
    assert claude_pass.is_test_article(test) and claude_pass.is_test_article(test)  # stable

    def row(article, before, old, new, after, stage="copy_edit"):
        return json.dumps(dict(article=article, stage=stage, para=1, old=old, new=new, before=before,
                               after=after, authors=["CE"], in_refs=False))

    lines = [
        row(train, "the results of this small trial were ", "", "the ", "majority of cases"),      # article
        row(test, "the results of this small trial were ", "", "a ", "held-out example here"),      # test: never
        row(train, "the results of this small trial were ", "", ",", " and the outcome was"),      # punctuation
        row(train, "the results of this small trial were cent", "re", "er", "d on the outcome"),   # spelling: no
        row(train, "the results of this small trial were in ", "12", "twelve", " patients only"),  # digits: no
        row(train, "the results of this small trial were ", "", "x", "y", stage="style_edit"),     # stage: no
        row(other, "the results of this other trial were ", "", "the ", "other journal case"),
    ]
    path = tmp_path / "changes.jsonl"
    path.write_text("\n".join(lines))
    ex = claude_pass.load_examples(str(path), "xyz", n=10)
    assert {(e["old"], e["new"]) for e in ex} == {("", "the "), ("", ",")}
    assert all("held-out" not in e["after"] for e in ex)
    ex = claude_pass.load_examples(str(path), "xyz", n=10)
    assert ex == claude_pass.load_examples(str(path), "xyz", n=10)  # deterministic: the prompt caches
    assert any(e["after"].startswith("other journal") for e in claude_pass.load_examples(str(path), "zzz", n=10))
    assert "[ -> the ]" not in claude_pass.examples_text([])  # built-in synthetic examples without a file
    assert "majority" in claude_pass.examples_text(ex)


def test_prefix_closing_follows_journal_setting():
    ijo, ipj = style.load("ijo"), style.load("ipj")
    assert "postoperative" in claude_pass.prompt(ijo)[0]
    assert "as the author wrote them" in claude_pass.prompt(ipj)[0]
