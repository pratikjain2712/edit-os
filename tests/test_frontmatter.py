import zipfile

import pytest
from docx import Document

from edit_os import docmodel, frontmatter as fm, queries, rules, style, tracked
from edit_os.changes import ChangeSet, Status

PORTAL_XML = b"""<?xml version="1.0" encoding="UTF-8"?><manuscript>
<title>Effect of Valsalva maneuver on intraocular pressure</title>
<articletype>Original Article</articletype>
<submission_date>04-Jan-26</submission_date><revision_date> </revision_date><Accepted>01-Oct-26</Accepted>
<authors>
<author><title>Dr</title><fname>Asha</fname><mname> </mname><lname>Kumar</lname><email>asha@example.org</email>
<department>Ophthalmology</department><institute>City Eye Hospital</institute><city>Chennai</city>
<state>Tamil Nadu</state><country> </country><is_corresponding>Y</is_corresponding></author>
<author><title>Mr</title><fname>ravi</fname><mname>shankar</mname><lname>rao</lname><email>ravi@example.org</email>
<department>Ophthalmology</department><institute>City Eye Hospital</institute><city>Chennai</city>
<state>Tamil Nadu</state><country> </country><is_corresponding>N</is_corresponding></author>
</authors></manuscript>"""

JATS_XML = b"""<?xml version="1.0" encoding="utf-8"?>
<article article-type="Case Report"><front><article-meta>
<article-categories><subj-group subj-group-type="Article Type"><subject>Case Report</subject></subj-group></article-categories>
<title-group><article-title>A rare case</article-title><alt-title alt-title-type="running-head">Rare case</alt-title></title-group>
<contrib-group><contrib contrib-type="author" corresp="yes"><name><surname>Nair.</surname><given-names>Suresh</given-names>
<prefix>Dr</prefix></name><email>s@example.org</email><xref ref-type="aff" rid="aff1"/></contrib></contrib-group>
<aff id="aff1"><institution>SCTIMST, Trivandrum</institution><country>India</country></aff>
<history><date date-type="received"><day>09</day><month>10</month><year>2025</year></date>
<date date-type="rev-recd"><day>01</day><month>02</month><year>2026</year></date>
<date date-type="rev-recd"><day>01</day><month>05</month><year>2026</year></date>
<date date-type="accepted"><day>09</day><month>10</month><year>2026</year></date></history>
</article-meta></front></article>"""

TITLE_PAGE = [
    "ORIGINAL ARTICLE",
    "Title: Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.",
    "Running title: Valsalva and IOP.",
    "Dr. Asha Kumar MS1, RAVI SHANKAR RAO MD1 and Jan van der Berg PhD2",
    "1Senior Resident, Department of ophthalmology, City Eye Hospital, Chennai, Tamil Nadu. Email: asha@example.org",
    "2Department of Physiology, Erasmus University, Rotterdam, Netherlands",
    "Corresponding author: Dr. Asha Kumar, Department of Ophthalmology, City Eye Hospital, 12 Main Road, "
    "Chennai 600006, Tamil Nadu. Phone +91 99999 99999. Email: asha@example.org",
]


def claude_fields(para_ids=("M1", "M2", "M3", "M4", "M5", "M6", "M7")):
    """What Claude would return for TITLE_PAGE (plus one invented author and affiliation part)."""
    roles = ["article_type", "title", "running_title", "authors", "affiliation", "affiliation", "correspondence"]
    return {
        "article_type": "ORIGINAL ARTICLE",
        "title": "Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.",
        "running_title": "Valsalva and IOP.",
        "authors": [
            {"given_names": "Asha", "surname": "Kumar", "prefix": "Dr", "affiliations": [1], "corresponding": True},
            {"given_names": "RAVI SHANKAR", "surname": "RAO", "prefix": "", "affiliations": [1], "corresponding": False},
            {"given_names": "Jan", "surname": "van der Berg", "prefix": "", "affiliations": [2], "corresponding": False},
            {"given_names": "Invented", "surname": "Person", "prefix": "", "affiliations": [2], "corresponding": False},
        ],
        "affiliations": [
            {"id": 1, "department": "Senior Resident, Department of ophthalmology", "institution": "City Eye Hospital",
             "city": "Chennai", "state": "Tamil Nadu", "country": ""},
            {"id": 2, "department": "Department of Physiology", "institution": "Erasmus University",
             "city": "Rotterdam", "state": "Zuid-Holland", "country": "Netherlands"},
        ],
        "correspondence": {"prefix": "Dr", "given_names": "Asha", "surname": "Kumar",
                           "address_parts": ["Department of Ophthalmology", "City Eye Hospital", "12 Main Road",
                                             "Phone +91 99999 99999"],
                           "city": "Chennai", "state": "Tamil Nadu", "postal_code": "600006", "country": "",
                           "email": "asha@example.org"},
        "proper_nouns": ["Valsalva", "OCT", "IOP"],
        "paragraphs": [{"id": i, "role": r} for i, r in zip(para_ids, roles)],
    }


def make(tmp_path, paras):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return str(path)


@pytest.fixture
def ijo():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["keywords"]["add_missing"] = False  # front matter only; key words are A4's
    return st


def test_recase():
    keep = ["Valsalva", "OCT"]
    assert fm.recase("Effect of Valsalva Maneuver on IOP: a Study Using OCT", "sentence", keep) == \
        "Effect of Valsalva maneuver on IOP: A study using OCT"
    assert fm.recase("EFFECT OF VALSALVA MANEUVER IN OCT", "sentence", keep) == "Effect of Valsalva maneuver in OCT"
    assert fm.recase("stimulus-sensitive myoclonus in the young – a case report", "title", []) == \
        "Stimulus-Sensitive Myoclonus in the Young – A Case Report"
    assert fm.clean_title("Title:  “Fern like deposits in dystrophy .”") == "Fern like deposits in dystrophy"


def test_names():
    cfg = {"initial_period": True}
    assert fm.format_name(fm.Author("AN", "Swamy"), cfg) == ("A. N. Swamy", False)
    assert fm.format_name(fm.Author("A.N.", "Swamy"), {"initial_period": False}) == ("A N Swamy", False)
    assert fm.format_name(fm.Author("Appulingam Narayana", "Swamy"), cfg, abbreviate=True)[0] == "Appulingam N. Swamy"
    assert fm.format_name(fm.Author("Anjum Syed Khan", "Joad"), cfg, abbreviate=True)[0] == "Anjum S. K. Joad"
    assert fm.format_name(fm.Author("MEERA", "IYER"), cfg)[0] == "Meera Iyer"
    assert fm.format_name(fm.Author("jan", "van der berg"), cfg)[0] == "Jan van der Berg"
    assert fm.cite_name(fm.Author("Asha M. E.", "Kumar")) == "Kumar AME"
    assert fm.strip_designation("Senior resident , Department of ophthalmology") == "Department of ophthalmology"
    assert fm.strip_designation("Consultant cornea and general ophthalmology") == "cornea and general ophthalmology"
    assert fm.strip_designation("Head and Neck Surgery") == "Head and Neck Surgery"


def test_xml_formats():
    m = fm.parse_xml(PORTAL_XML)
    assert m.article_type == "Original Article" and m.dates["received"].year == 2026 and "revised" not in m.dates
    assert m.authors[1]["given"] == "ravi shankar" and m.authors[0]["corresponding"]
    j = fm.parse_xml(JATS_XML)
    assert (j.article_type, j.running_title, j.authors[0]["surname"]) == ("Case Report", "Rare case", "Nair")
    assert j.dates["revised"].month == 5  # the latest revision
    assert j.authors[0]["aff"]["institution"] == "SCTIMST, Trivandrum"


def test_verify_drops_what_is_not_in_the_text():
    f = fm.verify(claude_fields(), "\n".join(TITLE_PAGE), {f"M{i}": f"p{i:04d}" for i in range(1, 8)})
    assert [a.surname for a in f.authors] == ["Kumar", "RAO", "van der Berg"]
    assert any("Invented" in n for n in f.notes)
    assert f.affiliations[1]["department"] == "Department of ophthalmology"  # designation removed by code
    assert f.affiliations[2]["state"] == ""  # "Zuid-Holland" is not in the text
    assert "Phone +91 99999 99999" not in f.correspondence["address_parts"]
    assert f.roles["p0002"] == "title"


def test_front_matter_in_the_manuscript(tmp_path, monkeypatch, ijo):
    path = make(tmp_path, TITLE_PAGE + ["Abstract", "Purpose: To test."])
    monkeypatch.setattr(fm, "split", lambda text: claude_fields())
    ms = docmodel.load(path)
    ms.front_sources = fm.Sources(meta=fm.parse_xml(PORTAL_XML))
    pl = fm.plan(ms, ijo)
    text = {b.role: b.text for b in pl.blocks}
    assert text["running_head"] == "<RH>Kumar, et al.: Valsalva and IOP"
    assert text["article_type"] == "Original Article"
    assert text["title"] == "Effect of Valsalva maneuver on intraocular pressure in OCT study"
    # IJO: middle initial, no period; the first affiliation numbered from 1 too (owner, 2026-10-27)
    assert text["authors"] == "Asha Kumar1, Ravi S Rao1, Jan van der Berg2"
    assert text["affiliations"] == ("1Department of Ophthalmology, City Eye Hospital, Chennai, Tamil Nadu, India, "
                                    "2Department of Physiology, Erasmus University, Rotterdam, Netherlands")
    corr = [b.text for b in pl.blocks if b.role == "correspondence"]
    # the PIN goes (S "delete pin/zip code"; owner, 2026-10-27)
    assert corr == ["Correspondence to: Dr. Asha Kumar,",
                    "Department of Ophthalmology, City Eye Hospital, 12 Main Road, Chennai, Tamil Nadu, India.",
                    "E-mail: asha@example.org"]
    assert text["dates"] == "Submitted: 04-Jan-2026\tRevised: ???\tAccepted: 01-Oct-2026"
    assert text["cite"] == ("Cite this article as: Kumar A, Rao RS, van der Berg J. Effect of Valsalva maneuver on "
                            "intraocular pressure in OCT study. Indian J Ophthalmol XX;XX:XX.")
    assert not pl.queries  # running head, type and full names all present

    # Through the rules: the existing paragraphs are rewritten in place, the running-title line removed.
    cs = rules.run_all(ms, ijo)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    for n in pl.inserts:
        n.status = Status.ACCEPTED
    out = docmodel.load(path)
    tracked.render(out, cs, tracked=False)
    fm.render_inserts(out, pl.inserts, tracked=False)
    out.save(tmp_path / "clean.docx")
    got = [p.text for p in docmodel.load(str(tmp_path / "clean.docx")).paragraphs if p.text]
    assert got[0] == "<RH>Kumar, et al.: Valsalva and IOP"
    assert got[1:5] == [text["article_type"], text["title"], text["authors"], text["affiliations"]]
    assert got[5:8] == corr and got[8] == text["dates"].replace("\t", "\t") and got[9] == text["cite"]
    assert got[10:] == ["Abstract", "Purpose: To test."]
    assert "Running title" not in " ".join(got)

    # Tracked output: inserted paragraphs are marked, authors are bold with superscript links.
    ms2 = docmodel.load(path)
    tracked.render(ms2, cs, tracked=True)
    fm.render_inserts(ms2, pl.inserts, tracked=True)
    ms2.save(tmp_path / "tracked.docx")
    xml = zipfile.ZipFile(tmp_path / "tracked.docx").read("word/document.xml").decode()
    assert xml.count("<w:pPr><w:rPr><w:ins ") == len(pl.inserts)
    assert '<w:vertAlign w:val="superscript"/>' in xml and "<w:tab/>" in xml


def test_missing_running_head_and_type_are_queried(tmp_path, monkeypatch, ijo):
    raw = claude_fields()
    raw.update(running_title="", article_type="")
    monkeypatch.setattr(fm, "split", lambda text: raw)
    paras = [t for t in TITLE_PAGE if not t.startswith(("Running", "ORIGINAL"))]
    path = make(tmp_path, paras + ["Introduction", "Materials and Methods", "Results", "Discussion", "Conclusion"])
    ms = docmodel.load(path)
    for item in raw["paragraphs"]:
        item["role"] = {"M1": "title", "M2": "authors", "M3": "affiliation", "M4": "affiliation",
                        "M5": "correspondence"}.get(item["id"], "body")
    pl = fm.plan(ms, ijo, fm.Sources())
    assert [q.check for q in pl.queries] == ["front.confirm_type", "front.running_head"]
    assert pl.blocks[0].text == "<RH>Kumar, et al.: ???"
    assert pl.blocks[1].text == "Original Article"
    qs = pl.add_queries(queries.QuerySet(), ms)
    assert [q.text for q in qs.queries] == ["Kindly confirm article type.", "Kindly provide running head."]


def test_xml_only_without_claude(tmp_path, ijo):
    path = make(tmp_path, ["Abstract", "Body text."])
    ms = docmodel.load(path)
    pl = fm.plan(ms, style.load("neurol"), fm.Sources(meta=fm.parse_xml(JATS_XML)))
    text = [b.text for b in pl.blocks]
    assert text[:4] == ["<RH>Nair: Rare case", "Case Report", "A Rare Case", "Suresh Nair"]
    assert "Address for correspondence: Dr. Suresh Nair," in text
    assert text[-4:-1] == ["Received: 09-10-2025", "Revised: 01-05-2026", "Accepted: 09-10-2026"]
    assert text[-1] == "How to cite this article: Nair S. A rare case. Neurol India XX;XX:XX."
    assert not pl.changes and all(n.anchor == "p0001" and n.where == "before" for n in pl.inserts)


def test_no_sources_no_proposals(tmp_path, ijo):
    ms = docmodel.load(make(tmp_path, ["Some title", "Abstract", "Body."]))
    assert fm.plan(ms, ijo).changes == [] and fm.plan(ms, ijo).inserts == []
    assert isinstance(rules.run_all(ms, ijo), ChangeSet)


def test_clean_article_type():
    assert fm.clean_article_type("Case Report (for IJO Case Reports)") == "Case Report"
    assert fm.clean_article_type("ORIGINAL ARTICLE") == "Original Article"
    assert fm.clean_article_type("Letter to the editor") == "Letter to the Editor"
    assert fm.clean_article_type("Case reports") == "Case Report"


def test_jehp_complete_affiliations_and_run_on_correspondence(tmp_path, monkeypatch):
    fields = claude_fields()
    fields["affiliations"][1].update(institution="City Eye Hospital", city="Chennai", state="Tamil Nadu",
                                     country="India")
    fields["affiliations"][0]["country"] = "India"
    path = make(tmp_path, TITLE_PAGE + ["Abstract", "Purpose: To test."])
    monkeypatch.setattr(fm, "split", lambda text: fields)
    ms = docmodel.load(path)
    ms.front_sources = fm.Sources(meta=fm.parse_xml(PORTAL_XML))
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    st.settings["front_matter"]["correspondence_one_line"] = True
    pl = fm.plan(ms, st)
    text = {b.role: b.text for b in pl.blocks}
    # each affiliation complete, departments of one place not grouped, country after each
    assert text["affiliations"] == ("1Department of Ophthalmology, City Eye Hospital, Chennai, Tamil Nadu, India, "
                                    "2Department of Physiology, City Eye Hospital, Chennai, Tamil Nadu, India")
    corr = [b.text for b in pl.blocks if b.role == "correspondence"]
    # JEHP editors drop the street when the address names the institution ("12 Main Road" goes); the PIN goes
    # (S "delete pin/zip code"; owner, 2026-10-27)
    assert corr[0] == ("Address for correspondence: Dr. Asha Kumar, Department of Ophthalmology, City Eye Hospital, "
                       "Chennai, Tamil Nadu, India.")
    assert corr[1] == "E-mail: asha@example.org"


def test_initials_written_after_the_name_go_before_it():
    # guide: "Swamy AN" -> "A. N. Swamy"; the name words stay together
    cases = [("Suganthi", "K", "K. Suganthi"), ("Helen Shaji J", "C", "J. C. Helen Shaji"),
             ("Mathangi", "D.C", "D. C. Mathangi"), ("Swamy", "AN", "A. N. Swamy")]
    for given, surname, want in cases:
        a = fm.Author(given, surname)
        assert fm.initials_first(a)
        assert fm.format_name(a, {"abbreviate_middle_names": True})[0] == want
    for given, surname in [("Ajay K.", "Jha"), ("John", "Smith III"), ("A. Fachruddin", "B")]:
        assert not fm.initials_first(fm.Author(given, surname))
    assert fm.format_name(fm.Author("John", "Smith III"), {})[0] == "John Smith III"
    # a middle name stays in full when the first name is an initial: it is the one name written out
    assert fm.format_name(fm.Author("M. Banyameen", "Iqbal"), {"abbreviate_middle_names": True})[0] == "M. Banyameen Iqbal"


def test_name_with_initials_expanded_from_the_xml():
    xml = [fm.Author("Andi Fachruddin", "Benyamin"), fm.Author("Geethu G G", "Nair")]
    full = fm._expand_from(fm.Author("A. Fachruddin", "B"), xml)
    assert (full.given, full.surname) == ("Andi Fachruddin", "Benyamin")
    assert fm._expand_from(fm.Author("Geethu G", "Nair"), xml) is None  # not a fuller form, a repeated initial
    assert fm._expand_from(fm.Author("Bo", "Nair"), xml) is None


def test_two_word_name_without_surname_is_kept():
    raw = {"authors": [{"given_names": "Retno Setyowati", "surname": "", "prefix": "", "affiliations": [1],
                        "corresponding": False},
                       {"given_names": "Nasriyah", "surname": "", "prefix": "", "affiliations": [1],
                        "corresponding": False}]}
    f = fm.verify(raw, "Nasriyah1, Retno Setyowati1", {})
    assert [(a.given, a.surname) for a in f.authors] == [("Retno", "Setyowati"), ("", "Nasriyah")]


def _affiliations(tmp_path, monkeypatch, affs, authors=None, code="jehp"):
    fields = claude_fields()
    fields["authors"] = authors or [
        {"given_names": "Asha", "surname": "Kumar", "prefix": "", "affiliations": [1], "corresponding": False},
        {"given_names": "Ravi", "surname": "Rao", "prefix": "", "affiliations": [2], "corresponding": False}]
    fields["affiliations"] = affs
    text = " ".join(str(v) for a in affs for v in a.values()) + " Asha Kumar Ravi Rao Shankar"
    path = make(tmp_path, TITLE_PAGE + [text, "Abstract", "Purpose: To test."])
    monkeypatch.setattr(fm, "split", lambda text: fields)
    ms = docmodel.load(path)
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return {b.role: b.text for b in fm.plan(ms, st).blocks}


def test_affiliations_complete(tmp_path, monkeypatch):
    blank = dict(department="", institution="", city="", state="", country="")
    # a department written alone above the one with the institution shares its institution and place
    text = _affiliations(tmp_path, monkeypatch, [
        dict(blank, id=1, department="Dept. of Dermatology"),
        dict(blank, id=2, department="Dept of Psychiatry", institution="Patil Medical College", city="Pimpri, Pune")])
    assert text["affiliations"] == ("1Department of Dermatology, Patil Medical College, Pimpri, Pune, Maharashtra, "
                                    "India, 2Department of Psychiatry, Patil Medical College, Pimpri, Pune, Maharashtra, India")


def test_affiliation_city_named_in_the_institution_stays(tmp_path, monkeypatch):
    blank = dict(department="", institution="", city="", state="", country="")
    text = _affiliations(tmp_path, monkeypatch, [
        dict(blank, id=1, department="School of Medical Sciences", institution="University of Phayao", city="Phayao",
             country="Thailand"),
        dict(blank, id=2, department="Department of Surgery", institution="AIIMS, New Delhi", city="New Delhi",
             country="India")])
    assert text["affiliations"] == ("1School of Medical Sciences, University of Phayao, Phayao, Thailand, "
                                    "2Department of Surgery, AIIMS, New Delhi, India")


def test_duplicate_affiliations_by_printed_form_and_sorted_links(tmp_path, monkeypatch):
    blank = dict(department="", institution="", city="", state="", country="")
    authors = [
        {"given_names": "Asha", "surname": "Kumar", "prefix": "", "affiliations": [1], "corresponding": False},
        {"given_names": "Ravi", "surname": "Rao", "prefix": "", "affiliations": [3, 2], "corresponding": False},
        {"given_names": "Shankar", "surname": "Rao", "prefix": "", "affiliations": [2], "corresponding": False}]
    text = _affiliations(tmp_path, monkeypatch, [
        dict(blank, id=1, department="Cardiology Department", institution="Zagazig University", country="Egypt"),
        dict(blank, id=2, department="Department of Cardiology", institution="Zagazig University", country="Egypt"),
        dict(blank, id=3, department="Department of Anatomy", institution="Zagazig University", country="Egypt")],
        authors=authors)
    assert text["affiliations"] == ("1Department of Cardiology, Zagazig University, Egypt, "
                                    "2Department of Anatomy, Zagazig University, Egypt")
    assert text["authors"] == "Asha Kumar1, Ravi Rao1,2, Shankar Rao1"


def test_ijo_keeps_the_xml_type_with_its_qualifier():
    from edit_os import style as style_mod
    ms = None  # the type comes from the XML: the text is not read
    meta = fm.Meta(article_type="Case Report (for IJO Case Reports)")
    for code, want in (("ijo", "Case Report (for IJO Case Reports)"), ("jehp", "Case Report")):
        blocks, _ = fm.build(ms, fm.Fields(), meta, style_mod.load(code))
        assert want in [getattr(b, "text", "") for b in blocks]


def test_contribution_lines_are_recognised():
    yes = ["Fatma Savur: Data acquisition, data analysis, manuscript review.",
           "Fatma Selin Kaya: Literature search, design",
           "Dr. K. Sivakumar: Literature search, manuscript editing, and manuscript review.",
           "A.B. – Conception and design",
           "Ravi Kumar (RK): Writing – original draft, supervision",
           "YJ: Data collection, manuscript writing.\nAG: Review and editing"]
    no = ["Title: Data analysis of glaucoma, a review",
          "Article Type: Review Article",
          "Running title: Data and design",
          "Ethical Approval: Not applicable as this study is a systematic review of published literature.",
          "Sub: Submission of revised manuscript for publication",
          "Corresponding author: Dr. Asha Kumar, Department of Ophthalmology",
          "Microcystic Pigmented Chromophobe Renal Cell Carcinoma: A Rare Case Report with Review of the Literature",
          "Occult Receptor–Positive Carcinoma: Case report and review of literature on data",
          "Fatma Savur1, Fatma Selin Kaya2",
          "Department of Ophthalmology, City Eye Hospital, Chennai"]
    assert [t for t in yes if not fm.is_contribution_line(t)] == []
    assert [t for t in no if fm.is_contribution_line(t)] == []


def test_contribution_paragraphs_after_a_heading():
    texts = ["Effect of X on Y", "Asha Kumar1, Ravi Rao2", "Authors' contributions",
             "AK conceived the study and drafted the manuscript; RR collected data.",
             "All authors approved the final version of the manuscript.",
             "Corresponding author: Dr. Asha Kumar"]
    lead = [docmodel.Paragraph(f"p{i}", None, "", t) for i, t in enumerate(texts)]
    assert fm.contribution_paragraphs(lead) == {"p2", "p3", "p4"}


def test_contribution_lines_in_the_title_block_are_kept(tmp_path, monkeypatch, ijo):
    # IJO_702: Claude read the "Name: role" lines under the title as author lines; the rebuild then rewrote
    # the first into the author group and deleted the others. They are contributions, not front matter.
    contrib = ["Asha Kumar: Concept, design, data acquisition, manuscript preparation",
               "Ravi Shankar Rao: Data acquisition, data analysis, manuscript review.",
               "Jan van der Berg: Literature search, design"]
    raw = claude_fields(para_ids=("M1", "M2", "M3", "M4"))
    raw["paragraphs"] = [{"id": "M1", "role": "title"}] + [{"id": f"M{i}", "role": "authors"} for i in (2, 3, 4)]
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, [TITLE_PAGE[1]] + contrib + ["Abstract", "Purpose: To test."])
    ms = docmodel.load(path)
    pl = fm.plan(ms, ijo, fm.Sources(meta=fm.parse_xml(PORTAL_XML)))
    ids = {p.id for p in ms.paragraphs if p.text in contrib}
    assert len(ids) == 3
    assert not [c for c in pl.changes if c.para_id in ids]
    assert any(b.role == "authors" for b in pl.blocks)  # the author group is still built (and inserted)
    assert any(n.role == "authors" for n in pl.inserts)
    assert any("contribution" in n for n in pl.notes)


ORDER_XML = b"""<?xml version="1.0" encoding="UTF-8"?><manuscript><title>T</title><authors>
<author><title>Mrs</title><fname>Asha</fname><lname>Rao</lname><email>a@x.org</email><institute>Ilam University of Medical
sciences, Ilam, Iran</institute><city>ilam</city><state>ilam</state><is_corresponding>Y</is_corresponding>
<author_order>Primary</author_order></author>
<author><title>Dr</title><fname>Ben</fname><lname>Lee</lname><author_order>0</author_order></author>
<author><title>Dr</title><fname>Cy</fname><lname>Kim</lname><author_order>1</author_order></author>
</authors></manuscript>"""


def test_portal_author_order_and_metadata_json():
    # The submitting author ("Primary") takes the slot the numbers leave free: 0, 1 -> third; 1, 2 -> first.
    assert [a["surname"] for a in fm.parse_xml(ORDER_XML).authors] == ["Lee", "Kim", "Rao"]
    one_based = ORDER_XML.replace(b"<author_order>0<", b"<author_order>2<")
    assert [a["surname"] for a in fm.parse_xml(one_based).authors] == ["Rao", "Kim", "Lee"]
    gap = ORDER_XML.replace(b"<author_order>1<", b"<author_order>2<")  # orders 0, 2: Primary second
    assert [a["surname"] for a in fm.parse_xml(gap).authors] == ["Lee", "Rao", "Kim"]
    import json
    m = fm.parse_metadata_json(json.dumps({"authors": [
        {"name": "Dr Reza Valizadeh", "email": "r@x.org", "author_order": "0", "corresponding_author": False},
        {"name": "Dr sajad sohrabnejad", "email": "s@x.org", "author_order": "Primary", "corresponding_author": True},
        {"name": "Mrs Kia Kazemi", "author_order": "1"}]}).encode())
    assert [(a["prefix"], a["given"], a["surname"]) for a in m.authors] == [
        ("Dr", "Reza", "Valizadeh"), ("Mrs", "Kia", "Kazemi"), ("Dr", "sajad", "sohrabnejad")]
    assert m.authors[2]["corresponding"]


def test_xml_affiliation_place_in_the_institution():
    f = fm.fields_from_meta(fm.parse_xml(ORDER_XML))
    rao = next(a for a in f.authors if a.surname == "Rao")
    assert f.affiliations[rao.affs[0]] == {"department": "", "institution": "Ilam University of Medical sciences",
                                           "city": "Ilam", "state": "", "country": "Iran"}


def test_cover_letter_signature_gives_way_to_the_xml_authors(tmp_path, monkeypatch, ijo):
    # JEHP 1764: the title page is a cover letter signed by the corresponding author, linked to an affiliation;
    # the submission has three authors.
    raw = claude_fields(para_ids=("M1",))
    raw["authors"] = [{"given_names": "Asha", "surname": "Rao", "prefix": "", "affiliations": [1],
                       "corresponding": True}]
    raw["paragraphs"] = [{"id": "M1", "role": "title"}]
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, ["Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.", "Abstract", "x"])
    ms = docmodel.load(path)
    src = fm.Sources(first_page=["Dear Editor,", "Asha Rao", "Department of Ophthalmology, City Eye Hospital"],
                     meta=fm.parse_xml(ORDER_XML))
    pl = fm.plan(ms, ijo, src)
    text = {b.role: b.text for b in pl.blocks}
    assert text["running_head"].startswith("<RH>Lee, et al.:")
    assert text["authors"].replace("1", "") == "Ben Lee, Cy Kim, Asha Rao"
    assert any("submission XML" in n for n in pl.notes)


def test_inline_numbered_affiliations_link_in_author_order():
    f = fm.Fields(authors=[fm.Author("Andrew", "Milewski", [1]), fm.Author("M.", "Hlayel"), fm.Author("W.", "Alomaim")],
                  affiliations={1: {"department": "General Education Unit", "institution": "Khalifa University, Abu "
                                    "Dhabi; 2 Radiography Department, Doha University", "city": "", "state": "",
                                    "country": ""}})
    line = ("Affiliations: 1 General Education Unit, Khalifa University, Abu Dhabi, United Arab Emirates; "
            "2 Radiography Department, University of Suffolk, Ipswich, United Kingdom; "
            "3 Department of Allied Health, University of Cumbria, Carlisle, United Kingdom")
    src = fm.Sources(first_page=["Authors: Andrew Milewski; M. Hlayel; W. Alomaim", line])
    fm._inline_numbered_affiliations(f, src, [])
    assert [a.affs for a in f.authors] == [[1], [2], [3]]
    assert f.affiliations[2] == {"department": "Radiography Department", "institution": "University of Suffolk",
                                 "city": "Ipswich", "state": "", "country": "United Kingdom"}


def test_running_title_line_after_the_key_words(tmp_path, monkeypatch, ijo):
    raw = claude_fields(para_ids=("M1",))
    raw["running_title"] = ""
    raw["paragraphs"] = [{"id": "M1", "role": "title"}]
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, ["Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.", "ABSTRACT",
                           "Background: x.", "Keywords: a; b", "Running Title: Valsalva and IOP", "INTRODUCTION",
                           "Text."])
    ms = docmodel.load(path)
    pl = fm.plan(ms, ijo, fm.Sources(first_page=TITLE_PAGE[3:], meta=fm.parse_xml(PORTAL_XML)))
    rh = next(b.text for b in pl.blocks if b.role == "running_head")
    assert rh.endswith(": Valsalva and IOP")
    line = next(p for p in ms.paragraphs if p.text.startswith("Running Title"))
    assert any(c.para_id == line.id and c.new == "" for c in pl.changes)
    assert not any(q.check == "front.running_head" for q in pl.queries)


def test_proper_noun_backstop():
    body = ("This study was carried out in Timor-Leste among mothers in rural villages of the country.\n"
            "Data were collected in Timor-Leste by trained enumerators who visited every household in turn.\n"
            "Many oral lesions and oral ulcers were seen in the patients that the dentists examined each week.")
    title = "Changes in Knowledge in Timor-Leste and Indonesia: Oral Disorders (2004-January 2026)"
    keep = fm.proper_keep([title], ["Oral"], body)
    out = fm.recase(title, "sentence", keep)
    assert out == "Changes in knowledge in Timor-Leste and Indonesia: Oral disorders (2004-January 2026)"
    assert "Oral" not in keep  # the text writes "oral" in lower case
    assert fm.recase("in Oral Disorders", "sentence", keep) == "In oral disorders"
    assert fm.recase("Supportive Interventions—A Systematic Review", "sentence", []) == \
        "Supportive interventions—A systematic review"
    assert fm.recase("in Timor-Leste", "sentence", ["Timor-Leste"]) == "In Timor-Leste"


def test_cite_name_order_follows_the_co_authors():
    f = fm.Fields(authors=[fm.Author("Nguyen Thi", "Diep"), fm.Author("", "Tran Hoa"), fm.Author("", "Al Amin")])
    fm._consistent_name_order(f)
    assert [fm.cite_name(a) for a in f.authors] == ["Diep NT", "Hoa T", "Al Amin"]


def test_designation_retired():
    assert fm.strip_designation("Retired Scientist, CSIR-NISTADS") == "CSIR-NISTADS"


def test_jehp_correspondence_prefix_from_the_xml(tmp_path):
    path = make(tmp_path, ["A title", "Abstract", "x"])
    ms = docmodel.load(path)
    st = style.load("jehp")
    st.settings["front_matter"]["ai_split"] = False
    pl = fm.plan(ms, st, fm.Sources(meta=fm.parse_xml(ORDER_XML)))
    corr = [b.text for b in pl.blocks if b.role == "correspondence"]
    assert corr[0].startswith("Address for correspondence: Mrs. Asha Rao")  # G91-92: the XML's <title>


def test_foreign_connectors_in_an_institution_stay_lowercase():
    # Turkish "ve" (and): the title page's "Başakşehir Çam ve Sakura Şehir Hastanesi" is not recased to "Ve"
    assert fm._clean_part("Başakşehir Çam ve Sakura Şehir Hastanesi") == "Başakşehir Çam ve Sakura Şehir Hastanesi"
    assert fm._clean_part("Başakşehir Çam Ve Sakura Şehir Hastanesi") == "Başakşehir Çam ve Sakura Şehir Hastanesi"
    assert fm._clean_part("universidade de são paulo") == "Universidade de São Paulo"
    assert fm._clean_part("Universidade Do Porto") == "Universidade do Porto"
    # place names, initials and a connector at the start are left alone
    assert fm._clean_part("Hospital Universitario La Paz") == "Hospital Universitario La Paz"
    assert fm._clean_part("University of Da Nang") == "University of Da Nang"
    assert fm._clean_part("De Montfort University") == "De Montfort University"
    assert fm._clean_part("Dr Y S R University") == "Dr Y S R University"
    assert fm._clean_part("city eye hospital") == "City Eye Hospital"


def test_short_title_line_below_the_title_is_deleted(tmp_path, monkeypatch):
    # JEHP: "Running Title" on one line, "Marital Enrichment and Childbearing" on the next; both go (the running
    # head carries the short title), as editors delete them
    raw = claude_fields(para_ids=("M1", "M2", "M3", "M4"))
    raw["running_title"] = "Valsalva and IOP"
    raw["paragraphs"] = [{"id": "M1", "role": "title"}, {"id": "M2", "role": "label"},
                         {"id": "M3", "role": "other_front_matter"}, {"id": "M4", "role": "authors"}]
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, ["Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.", "Running Title",
                           "Valsalva and IOP", TITLE_PAGE[3], "Abstract", "Background: x.", "Introduction", "Text."])
    ms = docmodel.load(path)
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    pl = fm.plan(ms, st, fm.Sources(meta=fm.parse_xml(PORTAL_XML)))
    gone = {c.para_id for c in pl.changes if c.new == "" and c.start == 0 and c.end == len(c.old)}
    for t in ("Running Title", "Valsalva and IOP"):
        assert next(p for p in ms.paragraphs if p.text == t).id in gone
    assert next(b.text for b in pl.blocks if b.role == "running_head").endswith(": Valsalva and IOP")


def test_bare_running_title_label_after_the_title_page(tmp_path, monkeypatch, ijo):
    raw = claude_fields(para_ids=("M1",))
    raw["running_title"] = ""
    raw["paragraphs"] = [{"id": "M1", "role": "title"}]
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, ["Effect of Valsalva Maneuver on Intraocular Pressure in OCT Study.", "Abstract Page",
                           "Running title", "Valsalva and IOP", "Abstract", "Background: x.", "Introduction", "Text."])
    ms = docmodel.load(path)
    pl = fm.plan(ms, ijo, fm.Sources(first_page=TITLE_PAGE[3:], meta=fm.parse_xml(PORTAL_XML)))
    assert next(b.text for b in pl.blocks if b.role == "running_head").endswith(": Valsalva and IOP")
    gone = {c.para_id for c in pl.changes if c.new == ""}
    for t in ("Running title", "Valsalva and IOP"):
        assert next(p for p in ms.paragraphs if p.text == t).id in gone
