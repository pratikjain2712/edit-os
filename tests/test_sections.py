"""Section rules (rules_sections): headings, abstract, key words, back matter; and paragraph
breaks in inserted text (tracked.PARAGRAPH_BREAK)."""

import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from edit_os import docmodel, queries, rules, rules_sections, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

BR = tracked.PARAGRAPH_BREAK

BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic.")


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        if isinstance(text, tuple):
            doc.add_paragraph().add_run(text[0]).bold = True
        else:
            doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def run(path, st, only=None, tracked_mode=False, tmp_path=None):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    ms = tracked.render(docmodel.load(path), cs, tracked=tracked_mode)
    out = str(path).replace(".docx", "_out.docx")
    ms.save(out)
    return [p.text for p in docmodel.load(out).paragraphs], out, cs


ARTICLE = [
    "A study of eyes",
    "Abstract:",
    "Purpose:To compare two treatments in patients with glaucoma who were seen in the clinic.",
    "Materials and Methods",
    "Patients were randomised to two arms and followed for twelve months in the clinic (2).",
    "Conclusions: Both treatments lower pressure.",
    "Key words: Glaucoma; Intraocular Pressure; Goldenhar syndrome; latanoprost; glaucoma.",
    "1. Introduction",
    BODY + " One had Goldenhar syndrome. The intraocular pressure and latanoprost doses were recorded.",
    "Materials And Methods",
    BODY,
    "Results",
    BODY,
    "Discussion",
    BODY,
    "Conflict of Interest: None",
    "Funding: Supported by the Eye Research Foundation grant 12.",
    "Acknowledgement: Nil",
    "References",
    "1. Smith A, Jones B. A study. Ophthalmology 2001;12:1-9.",
    "2. Smith A, Jones B. Another study. Ophthalmology 2002;13:1-9.",
]


@pytest.fixture
def article(tmp_path):
    return make(tmp_path, ARTICLE)


def test_layout_finds_abstract_and_keywords(article):
    lay = rules_sections.layout(docmodel.load(article))
    assert lay.abstract_head.text == "Abstract:"
    assert [p.text[:12] for p in lay.abstract] == ["Purpose:To c", "Materials an", "Patients wer", "Conclusions:"]
    assert lay.keywords.text.startswith("Key words")


def test_abstract_rules(article):
    st = style.load("base")
    st.settings["sections"]["abstract"]["rename"] = {"materials and methods": "Methods"}
    t, _, _ = run(article, st, only=["heading_tags", "sections_abstract"])
    assert t[1] == "<H1>Abstract"  # base tags it; the colon goes either way
    assert t[2].startswith("Purpose: To compare")
    assert t[3] == "Methods: "  # sub-heading on its own line: renamed, colon added, not tagged <H1>
    assert t[5].startswith("Conclusions: ")


def test_abstract_heading_added_when_missing(tmp_path):
    path = make(tmp_path, ["Title of the paper", "Background: " + BODY, "Results: " + BODY,
                           "Keywords: Cataract, glaucoma, retina", "Introduction", BODY])
    t, _, _ = run(path, style.load("ijo"), only=["sections_abstract"])
    assert t[:3] == ["Title of the paper", "Abstract", "Background: " + BODY]


def test_keywords(article):
    t, _, _ = run(article, style.load("base"), only=["sections_keywords"])
    # alphabetical, comma delimited, duplicate and final period gone, proper noun kept
    assert t[6] == "Keywords: Glaucoma, Goldenhar syndrome, intraocular pressure, latanoprost"


def test_keywords_missing_added_and_queried(tmp_path):
    path = make(tmp_path, ["Abstract", BODY, "Introduction", BODY, "References", "1. Smith A. X. Y 2001;1:1-2."])
    t, _, _ = run(path, style.load("base"), only=["sections_keywords"])
    assert t[1:3] == [BODY, "Keywords: "]
    assert "keywords.missing" in {q.check for q in queries.collect(docmodel.load(path), style.load("base")).queries}


def test_headings_ijo(article):
    t, _, _ = run(article, style.load("ijo"), only=["heading_tags", "sections_headings"])
    assert "Introduction" not in t  # IJO prints no Introduction heading: the paragraph goes
    assert t[8] == "<H1>Methods" and t[10] == "<H1>Results"  # S: "H1 | Caps" (Title case, owner)


def test_headings_base_numbering_and_minor_words(article):
    t, _, _ = run(article, style.load("base"), only=["heading_tags", "sections_headings"])
    assert t[7] == "<H1>Introduction"
    assert t[9] == "<H1>Materials and Methods"


def test_heading_case_setting(article):
    st = style.load("base")
    st.settings["sections"]["headings"]["case"] = "upper"
    t, _, _ = run(article, st, only=["heading_tags", "sections_headings"])
    assert t[11] == "<H1>RESULTS"


def test_back_matter(article):
    t, out, _ = run(article, style.load("ijo"), only=["heading_tags", "sections_back_matter"])
    i = t.index("<H2>Conflicts of interest")
    assert t[i + 1] == "There are no conflicts of interest."
    # "Conflict of Interest: None" and "Acknowledgement: Nil" go; the real funding statement stays
    assert "Conflict of Interest: None" not in t and "Acknowledgement: Nil" not in t
    j = t.index("<H2>Financial support and sponsorship")
    assert t[j + 1] == "Supported by the Eye Research Foundation grant 12."
    assert t.count("<H2>Financial support and sponsorship") == 1
    assert t[i + 2] == "<H1>References"


def test_back_matter_standard_heading_kept(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Financial support and sponsorship", "None",
                           "Conflicts of interest", "There are no conflicts of interest.", "References",
                           "1. Smith A. X. Y 2001;1:1-2."])
    t, _, _ = run(path, style.load("ijo"), only=["heading_tags", "sections_back_matter"])
    assert t[2:7] == ["<H2>Financial support and sponsorship", "Nil", "<H2>Conflicts of interest",
                      "There are no conflicts of interest.", "<H1>References"]


def test_queries(tmp_path):
    path = make(tmp_path, ["Abstract", "Purpose: " + BODY + " As shown before [3].", "Results: " + BODY,
                           "Keywords: Cataract, glaucoma, retina", "Case Report", BODY + " [1-3]", "References",
                           "1. Smith A. X. Y 2001;1:1-2.", "2. Smith A. X. Y 2001;1:1-2.", "3. Smith A. X. Y 2001;1:1-2."])
    st = style.load("ijo")
    st.settings["sections"]["abstract"]["word_limit_query"] = True
    st.settings["sections"]["abstract"]["word_limit"]["structured"] = 20
    checks = [q.check for q in queries.collect(docmodel.load(path), st).queries]
    assert {"abstract.citations", "abstract.too_long", "abstract.structured_case_report"} <= set(checks)


def test_paragraph_break_tracked_and_clean(tmp_path):
    path = make(tmp_path, [("Funding: Eye Foundation.",), "Last."])
    ms = docmodel.load(path)
    p = ms.paragraphs[0]
    cs = ChangeSet([
        Change("c1", p.id, 0, 9, "Funding: ", "<H2>Financial support" + BR, Layer.STRUCTURE, Source.RULE,
               status=Status.ACCEPTED),
        Change("c2", p.id, len(p.text), len(p.text), "", BR + "New one" + BR + "New two", Layer.STRUCTURE,
               Source.RULE, status=Status.ACCEPTED, format={"bold": False}),
    ])
    clean = tracked.render(docmodel.load(path), cs, tracked=False)
    clean.save(tmp_path / "clean.docx")
    texts = [q.text for q in docmodel.load(str(tmp_path / "clean.docx")).paragraphs]
    assert texts == ["<H2>Financial support", "Eye Foundation.", "New one", "New two", "Last."]
    tr = tracked.render(docmodel.load(path), cs, tracked=True)
    tr.save(tmp_path / "tracked.docx")
    xml = zipfile.ZipFile(tmp_path / "tracked.docx").read("word/document.xml").decode()
    assert xml.count("<w:p>") + xml.count("<w:p ") == 5
    marks = tr.document.element.body.findall(f".//{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}")
    assert len(marks) == 3  # every new paragraph mark is tracked; the original mark ends the last part
    # Rejecting everything gives the original text back in the editor's view.
    from tools.evaluate import tracked_pairs
    before = [b for b, _ in tracked_pairs(open(tmp_path / "tracked.docx", "rb").read())]
    assert "".join(before[:4]) == "Funding: Eye Foundation."


def test_evaluate_apply_text_picks_original_part():
    from tools.evaluate import apply_text

    c = Change("c", "p1", 0, 9, "Funding: ", "<H2>Financial support" + BR, Layer.STRUCTURE, Source.RULE)
    assert apply_text("Funding: Eye Foundation.", [c]) == "Eye Foundation."


def test_style_cli_renders(article, tmp_path):
    from edit_os.__main__ import main

    out = tmp_path / "out"
    assert main(["style", article, "ijo", str(out)]) == 0  # in-process: the network stays patched off
    t = [p.text for p in docmodel.load(str(out / "m_clean.docx")).paragraphs]
    assert "<H2>Conflicts of interest" in t and "Keywords: Glaucoma, Goldenhar syndrome, intraocular pressure, latanoprost" in t


def test_keywords_list_on_next_line_and_mixed_delimiters(tmp_path):
    path = make(tmp_path, ["Abstract", BODY, "Key words", "fluoxetine; Child, pediatrics, child",
                           "Introduction", BODY + " The child was treated."])
    t, _, _ = run(path, style.load("base"), only=["sections_keywords"])
    assert t[2:4] == ["Keywords:", "Child, fluoxetine, pediatrics"]
    assert "keywords.missing" not in {q.check for q in queries.collect(docmodel.load(path), style.load("base")).queries}


def test_keywords_label_with_a_qualifier(tmp_path):
    # "Keywords (MeSH) :" over the list: the journal's label replaces it all (not "Keywords: (MeSH) :")
    path = make(tmp_path, ["Abstract", BODY, "Keywords (MeSH) : ", "Fluoxetine; Child health; Pediatrics",
                           "Introduction", BODY])
    t, _, _ = run(path, style.load("jehp"), only=["sections_keywords"])
    assert t[2].startswith("Keywords:") and "MeSH" not in t[2] and t[2].count(":") == 1
    assert t[2:4] == ["Keywords:", "Child health, fluoxetine, pediatrics"]
    assert rules_sections._KW.match("Keywords (e.g., digital tools) were searched").group(3) == " "


def test_jehp_abstract_and_headings(tmp_path):
    path = make(tmp_path, ["Title of the paper", "Abstract:", "Background: " + BODY, "Methods and Material: " + BODY,
                           "Conclusions: " + BODY, "Keywords: Cataract, glaucoma, retina", "Introduction", BODY,
                           "Ethical consideration", BODY])
    t, _, _ = run(path, style.load("jehp"), only=["heading_tags", "sections_abstract"])
    assert t[1] == "Abstract"  # not tagged; colon gone
    # S: 'The sub-headings “Materials and Methods” or “Methodology” should be changed to “Methods”'
    assert t[3].startswith("Methods: ") and t[4].startswith("Conclusion: ")
    assert t[6] == "<H1>Introduction" and t[8] == "<H2>Ethical consideration"


def test_subheadings_in_sentence_case(tmp_path):
    long = ("The analysis was planned in advance and the participants from Iran were followed for a year; "
            "the analysis of participants used standard methods and the results were checked twice.")
    path = make(tmp_path, ["Introduction", long, "Materials and Methods", ("Statistical Analysis",), long,
                           ("Participants From Iran",), long, ("Analysis: Beyond Participants",), long,
                           "Ethical Considerations", long, "Discussion", long])
    st = style.load("jehp")
    t, _, _ = run(path, st, only=["heading_tags", "subheadings"])
    assert t[3] == "<H2>Statistical analysis"
    assert t[5] == "<H2>Participants from Iran"  # "Iran" is a proper noun in the text
    assert t[7] == "Analysis: Beyond Participants"  # a run-in label, not a heading: untouched
    assert t[9] == "<H2>Ethical considerations"
    st.settings["headings"]["subheading_case"] = "keep"
    assert run(path, st, only=["heading_tags", "subheadings"])[0][3] == "<H2>Statistical Analysis"


def test_letter_salutation_from_the_stylesheet(tmp_path):
    # IJO sheet: 'Every LTE article should start with “Dear Editor” ... replace with ‘Dear Editor’'
    path = make(tmp_path, ["Letter to the Editor", "Sir,", BODY, "Sir Ronald Fisher showed it."])
    t, _, _ = run(path, style.load("ijo"), only=["sections_salutation"])
    assert t[1] == "Dear Editor," and t[3] == "Sir Ronald Fisher showed it."
    path = make(tmp_path, ["Dear Sir/Madam, We read the article with interest.", BODY], name="m2.docx")
    t, _, _ = run(path, style.load("ijo"), only=["sections_salutation"])
    assert t[0] == "Dear Editor, We read the article with interest."
    t, _, _ = run(path, style.load("abr"), only=["sections_salutation"])  # no such rule in the ABR sheet
    assert t[0] == "Dear Sir/Madam, We read the article with interest."


def test_sentence_case_leaves_words_in_brackets(tmp_path):
    long = ("The examination was planned in advance and the participants were followed for a year; the indirect "
            "examination used standard methods and the results were checked twice by two observers.")
    path = make(tmp_path, ["Introduction", long, "Materials and Methods", ("Indirect Examination (Heine, Germany)",),
                           long, "Discussion", long])
    t, _, _ = run(path, style.load("ijo"), only=["heading_tags", "subheadings"])
    assert t[3] == "<H2>Indirect examination (Heine, Germany)"


def _case_report_run(path, st):
    from types import SimpleNamespace

    ms = docmodel.load(path)
    ms.front_sources = SimpleNamespace(meta=SimpleNamespace(article_type="Case Report (for IJO Case Reports)"))
    cs = rules.run_all(ms, st, only=["heading_tags", "subheadings"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(path).replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs]


def test_case_report_methods_head_is_first_level(tmp_path):
    # IJO_702: the case report's "Methods" holds the ethics statement; "Literature Search Strategy" after it is a
    # section of its own (<H1>, Title Caps), the cases and the summary after it its <H2>s
    long = ("The review was planned in advance and the cases were followed for a year; the search of the "
            "databases used standard methods and the results were checked twice by two observers. Each report was "
            "read in full.")
    paras = ["Introduction", long, "Methods", long, "Literature Search Strategy", long, "Case 1", long, "Case 2",
             long, "Summary of Previously Reported Cases", long, "Discussion", long]
    path = make(tmp_path, paras)
    t = _case_report_run(path, style.load("ijo"))
    assert t[2] == "<H1>Methods"
    assert t[4] == "<H1>Literature Search Strategy"
    assert t[6] == "<H2>Case 1" and t[8] == "<H2>Case 2"
    assert t[10] == "<H2>Summary of previously reported cases"
    # not a case report: a sub-heading of the Methods
    t, _, _ = run(path, style.load("ijo"), only=["heading_tags", "subheadings"])
    assert t[4] == "<H2>Literature search strategy"
    # the cases straight under Methods stay second level
    path = make(tmp_path, ["Introduction", long, "Methods", long, "Case 1", long, "Case 2", long, "Discussion", long],
                name="m2.docx")
    t = _case_report_run(path, style.load("ijo"))
    assert t[4] == "<H2>Case 1" and t[6] == "<H2>Case 2"


def test_topic_and_question_head_is_a_subheading(tmp_path):
    # IJO_702: "Mechanism: Why can OVM cause enophthalmos?" on its own line (plain) over the Discussion's text
    long = ("Enophthalmos reflects a mismatch between orbital container volume and its contents; the "
            "malformation may cause fat atrophy through repeated distension and the pressure it exerts over years. "
            "The orbital fat is lost and the globe sinks back into the orbit.")
    path = make(tmp_path, ["Introduction", long, "Discussion", long, "Mechanism: Why can OVM cause enophthalmos?",
                           long, "Comparator (C): Non-smokers", long, "Discussion of this: Two words", long])
    t, _, _ = run(path, style.load("ijo"), only=["heading_tags", "subheadings"])
    assert t[4] == "<H2>Mechanism: Why can OVM cause enophthalmos?"
    assert t[6] == "Comparator (C): Non-smokers"  # no question: a label, not a heading
    assert t[8] == "Discussion of this: Two words"
