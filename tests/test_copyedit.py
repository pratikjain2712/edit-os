"""Copy-edit checkers (B1-B3), offline: synthetic text, a VarCon excerpt, a mocked LanguageTool."""

import pytest
from docx import Document

from edit_os import copyedit, docmodel, style, tracked
from edit_os.changes import Status
from edit_os.copyedit import abbrev, langtool, spelling

VARCON = """\
# analyze <verified> (level 20)
A Z: analyze / B: analyse
A Z: analyzed / B: analysed
A C: analyzes / B Cv: analyses
A B: analyses | -- pl: analysis

# among <verified> (level 10)
A B. Z: among / AV B Zv: amongst

# anemia <verified> (level 35)
A C Dv: anemia / B D: anaemia

# color <verified> (level 10)
A Cv: color / B C Z: colour

# organization <verified> (level 10)
A Z: organization / B: organisation

# randomized <verified> (level 50)
A Z: randomized / B: randomised

# Composite (level 70)
A: Composite / B: Compositae

# colorectal (level 80)
A: colorectal / B: colourectal

# et (level 70)
A: et / B: aet
"""


@pytest.fixture
def variants(monkeypatch):
    v = spelling.parse(VARCON.splitlines())
    monkeypatch.setattr(spelling, "load", lambda *a, **k: v)
    return v


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return docmodel.load(str(path))


def journal(code="ijo", **copyedit_settings):
    st = style.load(code)
    for section, values in copyedit_settings.items():
        st.settings["copyedit"][section].update(values)
    return st


def apply(ms, changes):
    out = {}
    for p in ms.paragraphs:
        text = p.text
        for c in sorted((c for c in changes if c.para_id == p.id), key=lambda c: -c.start):
            assert text[c.start:c.end] == c.old
            text = text[:c.start] + c.new + text[c.end:]
        out[p.id] = text
    return [out[p.id] for p in ms.paragraphs]


# --- B2 spelling ---------------------------------------------------------------------

def test_varcon_parse_keeps_only_safe_conversions(variants):
    assert variants.to_american["anaemia"] == "anemia"
    assert variants.to_american["analysed"] == "analyzed"
    assert variants.to_british["randomized"] == "randomised"
    assert variants.to_oxford["organisation"] == "organization"  # British with -ize
    assert variants.to_oxford["color"] == "colour"
    assert "analyses" not in variants.to_american  # plural of analysis is American too
    assert "among" not in variants.to_british       # "B." marks an equal variant
    assert "composite" not in variants.to_british   # proper-noun cluster
    assert "colorectal" not in variants.to_british  # level 80: too rare to trust
    assert "et" not in variants.to_british          # too short


def test_american_conversion(tmp_path, variants):
    ms = make(tmp_path, ["Introduction",
                         "Anaemia was analysed in randomised patients. The analyses showed colour changes.",
                         "References",
                         "1. Smith A. Anaemia in colour. J Haem 2001;1:1-2."])
    changes = spelling.check(ms, journal("ijo", spelling={"convention": "american", "journals": {}}), variants)
    assert apply(ms, changes)[1] == "Anemia was analyzed in randomized patients. The analyses showed color changes."
    assert apply(ms, changes)[3] == ms.paragraphs[3].text  # references untouched
    assert all(c.layer.value == "language" and c.source.value == "rule" for c in changes)


def test_british_conversion_by_journal_and_names_kept(tmp_path, variants):
    ms = make(tmp_path, ["Results",
                         "The color of the lesion was analyzed by the World Health Organization criteria.",
                         "Organization of care was randomized."])
    st = journal("ijdr")  # IJDR follows British spelling
    out = apply(ms, spelling.check(ms, st, variants))
    assert out[1] == "The colour of the lesion was analysed by the World Health Organization criteria."
    assert out[2] == "Organisation of care was randomised."


def test_follow_author_fixes_the_minority(tmp_path, variants):
    ms = make(tmp_path, ["Discussion",
                         "Colour and organisation were analysed. Anaemia was randomised.",
                         "One color sample."])
    st = journal("xyz", spelling={"convention": "follow_author", "journals": {}})
    assert spelling.convention(st, ms, variants) == "british"
    assert apply(ms, spelling.check(ms, st, variants))[2] == "One colour sample."


def test_skip_list_and_urls(tmp_path, variants):
    ms = make(tmp_path, ["Methods", "See https://example.org/colour for colour charts."])
    st = journal("ijo", spelling={"convention": "american", "journals": {}, "skip": []})
    out = apply(ms, spelling.check(ms, st, variants))
    assert out[1] == "See https://example.org/colour for color charts."


# --- B3 abbreviations ----------------------------------------------------------------

def test_schwartz_hearst():
    assert abbrev.best_long_form("OCT", "using optical coherence tomography") == "optical coherence tomography"
    assert abbrev.best_long_form("HbA1c", "levels of glycated haemoglobin A1c") is not None
    assert abbrev.best_long_form("XYZ", "nothing here") is None
    p = docmodel.Paragraph("p1", None, "", text="Intraocular pressure (IOP) and best-corrected visual acuity (BCVA) were noted (n = 12).")
    defs = abbrev.definitions(p)
    assert [(d.short, d.long) for d in defs] == [("IOP", "Intraocular pressure"),
                                                ("BCVA", "best-corrected visual acuity")]


def test_abbreviation_queries(tmp_path):
    ms = make(tmp_path, [
        "Abstract",
        "We measured intraocular pressure (IOP) in 20 eyes; IOP was high.",
        "Keywords: Glaucoma, pressure, eye",
        "Introduction",
        "The IOP was high and the RNFL was thin; mean SD was 2 mm.",
        "Retinal nerve fiber layer (RNFL) was imaged with optical coherence tomography (OCT).",
        "Later, optical coherence tomography (OCT) was repeated in the clinic of Armonk, NY.",
        "OCT showed thinning.",
        "References",
        "1. Lee K. The PQRS study. Eye 2001;1:1-2.",
    ])
    changes, qs = abbrev.collect(ms, journal(abbreviations={"allow": []}))  # not IJO's ophthalmic allowlist
    got = sorted((q.check, q.text) for q in qs)
    assert got == [
        ("abbrev.before_definition", "RNFL: Kindly expand the abbreviation at first use."),
        ("abbrev.defined_twice", "OCT: The abbreviation has already been expanded. Kindly check."),
        ("abbrev.undefined", "IOP: Kindly expand the abbreviation."),  # defined in the abstract only
    ]
    q = next(q for q in qs if q.check == "abbrev.undefined")
    assert ms.by_id(q.para_id).text[q.start:q.end] == "IOP" and q.para_id == ms.paragraphs[4].id
    assert changes == []


def test_abbreviation_queries_skip_what_editors_do_not_query(tmp_path):
    ms = make(tmp_path, [
        "Introduction",
        "The XQV was raised and the ZMT score fell in the DELTA study.",                   # real: queried
        "Patients with DBS (deep brain stimulation) were followed.",                       # expanded after it
        "Imaging used the KRESCAN 700 system (Fabrikam Meditec, Jena, Germany).",          # model number
        "The approval code was IR.TUMS.REC.1402.17 and the STUDY DESIGN AND SETTING were listed.",
        "Authors' contributions: Conceptualisation: AB, CD; writing: AB, EF.",
        "Rao KL, Moss T, Dunn FA. A synthetic title on eyes. J Test 2014;311:1901-11.",  # a stray reference
        "Searches combined glaucoma AND pressure; the XQV rose again.",
        "References",
        "1. Lee K. A study. Eye 2001;1:1-2.",
        "Table 1: Results",
        "ZMT, zonal motion test",               # a table note: copy editors still expand ZMT in the text
    ])
    _, qs = abbrev.collect(ms, journal(abbreviations={"allow": []}))
    assert sorted(q.text for q in qs if q.check == "abbrev.undefined") == [
        "DELTA: Kindly expand the abbreviation.", "XQV: Kindly expand the abbreviation.",
        "ZMT: Kindly expand the abbreviation."]
    assert sum(q.text.startswith("XQV") for q in qs) == 1  # one query per abbreviation


def test_abbreviation_defined_in_square_brackets(tmp_path):
    ms = make(tmp_path, ["Introduction", "Rhegmatogenous retinal detachment [RRD] is common.", "RRD recurs."])
    assert abbrev.collect(ms, journal(abbreviations={"allow": []}))[1] == []


def test_abbreviation_changes_when_configured(tmp_path):
    ms = make(tmp_path, [
        "Introduction",
        "The RNFL was thin.",
        "Retinal nerve fiber layer (RNFL) was imaged with optical coherence tomography (OCT).",
        "Later, optical coherence tomography (OCT) was repeated.",
        "OCT showed thinning.",
    ])
    st = journal(abbreviations={"redefinition": "change", "move_definition": True, "allow": []})
    changes, qs = abbrev.collect(ms, st)
    assert apply(ms, changes)[1:] == [
        "The retinal nerve fiber layer (RNFL) was thin.",
        "RNFL was imaged with optical coherence tomography (OCT).",
        "Later, OCT was repeated.",
        "OCT showed thinning.",
    ]
    assert qs == []


def test_abbreviation_not_used_again_keeps_full_form(tmp_path):
    # Instructions to all journals: "If a term is not used again, an abbreviation is not needed."
    ms = make(tmp_path, [
        "Introduction",
        "Records came from the outpatient department (OPD) and the electronic medical records (EMRs).",
        "Bandage contact lens (BCL) was placed.",
        "The BCL stayed.",
        "<H1>Methods",
        "Patients were seen at the outpatient department.",
    ])
    assert abbrev.collect(ms, journal()) == ([], [])  # off by default: copy editors keep definitions
    changes, qs = abbrev.collect(ms, journal(abbreviations={"unused_definition": "change"}))
    assert apply(ms, changes)[1:4] == [
        "Records came from the outpatient department and the electronic medical records.",
        "Bandage contact lens (BCL) was placed.",
        "The BCL stayed.",
    ]
    assert qs == []
    mode = journal(abbreviations={"unused_definition": "query"})
    assert {q.check for q in abbrev.collect(ms, mode)[1]} == {"abbrev.unused"}


def test_abbreviation_used_in_a_table_is_used(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Ultrasound biomicroscopy (UBM) was done.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "UBM findings"
    doc.save(tmp_path / "t.docx")
    ms = docmodel.load(str(tmp_path / "t.docx"))
    assert abbrev.collect(ms, journal()) == ([], [])


def test_full_form_after_definition_becomes_abbreviation(tmp_path):
    # Stylesheets: "only abbreviation should be used thereafter". Headings keep the full form.
    ms = make(tmp_path, [
        "Introduction",
        "Intraocular pressure was noted before; intraocular pressure (IOP) rose.",
        "<H1>Intraocular pressure",
        "Mean intraocular pressure was 12 mmHg; IOP fell. Intraocular pressures were stable.",
        "The IOP was lower at 1 week; IOP and IOP spikes were rare.",
    ])
    on = journal(abbreviations={"full_form_after_definition": "change"})
    changes, qs = abbrev.collect(ms, on)
    assert apply(ms, changes)[1:4] == [
        "Intraocular pressure was noted before; intraocular pressure (IOP) rose.",
        "<H1>Intraocular pressure",
        "Mean IOP was 12 mmHg; IOP fell. IOPs were stable.",
    ]
    assert {c.reason for c in changes} == {"use abbreviation after its definition"}
    assert abbrev.collect(ms, journal())  # IJO: on (S: "only abbreviation should be used thereafter")
    assert abbrev.collect(ms, journal(abbreviations={"full_form_after_definition": "off"}))[0] == []


def test_full_form_kept_when_the_author_mostly_writes_it_out(tmp_path):
    # "retinoblastoma (RB)" in a paper about retinoblastoma: the editors leave the term written out.
    ms = make(tmp_path, [
        "Abstract",
        "Adult retinoblastoma (RB) is rare. Retinoblastoma was suspected.",
        "Keywords: Retinoblastoma, adult retinoblastoma, enucleation",
        "Introduction",
        "Retinoblastoma (RB) is the commonest tumour. Adult retinoblastoma is rare; RB in adults is atypical.",
        "Retinoblastoma was confirmed. Adult retinoblastoma needs follow-up.",
    ])
    on = journal("ija", abbreviations={"full_form_after_definition": "change"})
    assert [c for c in abbrev.collect(ms, on)[0] if c.reason == "use abbreviation after its definition"] == []


def test_full_form_not_shortened_in_abstract_or_key_words(tmp_path):
    ms = make(tmp_path, [
        "Abstract",
        "Intraocular pressure (IOP) was measured. IOP rose, IOP fell, intraocular pressure was stable.",
        "Keywords: Glaucoma, intraocular pressure, IOP",
        "Introduction",
        "We studied it.",
    ])
    on = journal("ija", abbreviations={"full_form_after_definition": "change"})
    assert [c for c in abbrev.collect(ms, on)[0] if c.reason == "use abbreviation after its definition"] == []


def test_allowlist_and_codes_not_queried(tmp_path):
    ms = make(tmp_path, ["Introduction",
                         "DNA and HIV were tested; P < 0.05; CI 95%; C6-T4 and GDS-15 scores; WHO grade; "
                         "ethics number IEC/0078/2022; analysed with IBM SPSS; post-COVID cases."])
    assert abbrev.collect(ms, journal())[1] == []


# --- B1 LanguageTool -----------------------------------------------------------------

def _match(text, wrong, rep, rule, cat="TYPOGRAPHY", reps=None):
    return {"offset": text.index(wrong), "length": len(wrong), "message": f"{rule} message",
            "replacements": [{"value": r} for r in (reps or [rep])],
            "rule": {"id": rule, "category": {"id": cat}}}


def test_languagetool_client_caches(tmp_path):
    calls = []

    def post(url, data):
        calls.append((url, data["language"]))
        return {"matches": [{"offset": 0}]}

    c = langtool.Client("http://127.0.0.1:9", cache_path=str(tmp_path / "lt.sqlite"), post=post)
    assert c.check("Some text.", "en-US") == [{"offset": 0}]
    assert c.check("Some text.", "en-US") == [{"offset": 0}]
    assert calls == [("http://127.0.0.1:9/v2/check", "en-US")]


def test_languagetool_keeps_only_safe_single_suggestions(tmp_path):
    body = "Patients aged 18-81 years were included , and the follow up was long. Their is a effect [1-2]."
    ms = make(tmp_path, ["Introduction", body, "References", "1. Lee K. Eye 2001;1:1-2."])
    sent = []

    def post(url, data):
        sent.append(data)
        t = data["text"]
        if t != body:
            return {"matches": []}
        return {"matches": [
            _match(t, "18-81", "18–81", "HYPHEN_TO_EN", "PUNCTUATION"),
            _match(t, " ,", ",", "COMMA_PARENTHESIS_WHITESPACE"),
            _match(t, "follow up", "follow-up", "CA_FOLLOW_UP", "GRAMMAR"),
            _match(t, "Their", "There", "THEIR_IS", "CONFUSED_WORDS"),  # not in the allowlist
            _match(t, "a effect", "an effect", "EN_A_VS_AN", "MISC", reps=["an effect", "the effect"]),
            _match(t, "1-2", "1–2", "HYPHEN_TO_EN", "PUNCTUATION"),  # inside a citation: protected
        ]}

    client = langtool.Client(cache_path=str(tmp_path / "lt.sqlite"), post=post)
    st = journal(languagetool={"warn_categories": ["CONFUSED_WORDS"]})
    changes, warnings = langtool.collect(ms, st, client)
    assert apply(ms, changes)[1] == \
        "Patients aged 18–81 years were included, and the follow-up was long. Their is a effect [1-2]."
    assert [w.rule for w in warnings] == ["THEIR_IS"]
    assert all(d["language"] == "en-US" for d in sent)  # IJO: American spelling
    assert not any("Eye 2001" in d["text"] for d in sent)  # references not sent


def test_languagetool_language_follows_spelling(tmp_path):
    ms = make(tmp_path, ["Introduction", "Text."])
    assert langtool.language(ms, journal("ijdr")) == "en-GB"
    assert langtool.language(ms, journal("ijo")) == "en-US"


def test_languagetool_skipped_without_server(tmp_path, monkeypatch):
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)
    ms = make(tmp_path, ["Introduction", "Text."])
    assert langtool.collect(ms, journal()) == ([], [])


# --- run_all + rendering --------------------------------------------------------------

def test_run_all_renders_tracked(tmp_path, variants, monkeypatch):
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)
    ms = make(tmp_path, ["Introduction", "Anaemia was analysed in 20 patients."])
    res = copyedit.run_all(ms, journal("ijo", spelling={"journals": {}, "convention": "american"}))
    assert [c.new for c in res.changes.changes] == ["Anemia", "analyzed"]
    for c in res.changes.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "clean.docx"
    tracked.render(docmodel.load(str(tmp_path / "m.docx")), res.changes, tracked=False).save(out)
    assert docmodel.load(str(out)).paragraphs[1].text == "Anemia was analyzed in 20 patients."


def test_scopes_structured_abstract_and_box(tmp_path):
    ms = make(tmp_path, ["Title", "Abstract", "Background", "Text one.", "Methods", "Text two.",
                         "Keywords: a, b, c", "Highlights", "Box text.", "<H1>Introduction", "Main text."])
    got = [[p.text for p in s] for s in abbrev.scopes(ms)]
    assert got == [["Abstract", "Background", "Text one.", "Methods", "Text two.", "Keywords: a, b, c"],
                   ["Highlights", "Box text."], ["<H1>Introduction", "Main text."]]


def test_scopes_introduction_without_heading(tmp_path):
    # IJO prints no "Introduction" heading: the text after the key words is the main text, not a box.
    ms = make(tmp_path, ["Title", "Abstract", "Purpose: Text one.", "Keywords: a, b, c",
                         "Intraocular pressure (IOP) was high.", "<H1>Methods", "IOP was measured."])
    got = [[p.text for p in s] for s in abbrev.scopes(ms)]
    assert got == [["Abstract", "Purpose: Text one.", "Keywords: a, b, c"],
                   ["Intraocular pressure (IOP) was high.", "<H1>Methods", "IOP was measured."]]
    _, queries = abbrev.collect(ms, journal("ijo"))
    assert not [q for q in queries if q.text.startswith("IOP")]


def test_formatting_only_change_does_not_block_copy_edits(monkeypatch):
    from edit_os import copyedit
    from edit_os.changes import Change, Layer, Source
    ms = docmodel.Manuscript.__new__(docmodel.Manuscript)
    fmt = Change(id="f1", para_id="p0001", start=0, end=40, old="x" * 40, new="x" * 40, layer=Layer.HOUSE_STYLE,
                 source=Source.RULE, reason="fmt_color", format={"color": False})
    edit = Change(id="s1", para_id="p0001", start=5, end=15, old="paediatric", new="pediatric", layer=Layer.LANGUAGE,
                  source=Source.RULE, reason="spelling (american)")
    monkeypatch.setattr(copyedit, "CHECKERS", {"spelling": lambda *a, **k: ([edit], [], [])})
    res = copyedit.run_all(ms, None, existing=[fmt])
    assert [c.id for c in res.changes.changes] == ["s1"]


def test_claude_cache_shared_across_threads(tmp_path):
    import threading
    from edit_os.copyedit import claude_pass
    cache = claude_pass._Cache(str(tmp_path / "c.sqlite"))
    errors = []

    def work(i):
        try:
            for j in range(20):
                cache.put(f"k{i}-{j}", {"edits": [], "queries": []})
                assert cache.get(f"k{i}-{j}") == {"edits": [], "queries": []}
        except Exception as e:  # pragma: no cover - the failure being guarded against
            errors.append(e)

    threads = [threading.Thread(target=work, args=(i,)) for i in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors


def test_abbreviation_false_queries_cut(tmp_path):
    """Lists and table notes, key words, titles, back matter, names in mixed case, degrees, institutions and
    tools are not queried; real undefined acronyms still are."""
    ms = make(tmp_path, [
        "Abstract",
        "Background: The quality of care (QoC) of wards was rated.",
        "Keywords: Quality of care (QoC); Nursing",
        "Abbreviations: KAP: Knowledge, Attitude and Practice; QVF, quorum validity factor; MINI, Mini Ward Interview",
        "Introduction",
        "Care Pathways for QZR Wards in Rural Hospitals",
        "The QoC score, the KAP questionnaire and the Arabic MINI modules were used with the QZR scale.",
        "Analyses ran in SmartPLS and the t-SNE plot; ethics approval came from the WRX Medical College board.",
        "Fit was tested with the QVF and the TREND statement was followed; the XDQ fell.",
        "Correspondence: Ravi Kumar, MD, DM",
        "Authors’ Contributions",
        "HD: Conceptualization. SM contributed to the analysis.",
        "Use of AI: ChatGPT (OpenAI) helped with the language.",
    ])
    st = journal("jehp", abbreviations={})
    got = sorted(q.text.split(":")[0] for q in abbrev.collect(ms, st)[1] if q.check == "abbrev.undefined")
    assert got == ["QZR", "XDQ"]
    # Without `defined_in_article` (other journals) the abstract's and the list's definitions and institutions count.
    st = journal("ijo", abbreviations={"allow": []})
    got = sorted(q.text.split(":")[0] for q in abbrev.collect(ms, st)[1] if q.check == "abbrev.undefined")
    assert got == ["KAP", "MINI", "QVF", "QZR", "QoC", "WRX", "XDQ"]


def test_abbreviation_list_is_not_a_second_definition(tmp_path):
    ms = make(tmp_path, [
        "Introduction",
        "Reproductive tract infections (RTI) are common.",
        "Abbreviations: Reproductive Tract Infections (RTI), AOR (Adjusted Odds Ratio)",
        "RTI was frequent; the AOR was 2.",
    ])
    qs = abbrev.collect(ms, journal("jehp", abbreviations={}))[1]
    assert not [q for q in qs if q.check in ("abbrev.defined_twice", "abbrev.before_definition")]
