"""figures_insert: figure files from the submission zip proposed above their legends (synthetic only)."""

import io
import json
import random
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from edit_os import docmodel, docview, figures_insert, pipeline, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TEXT = [
    "Original Article",
    "Introduction",
    "The lesion is shown in Figure 1 and the scan in Figure 2 (Figure 3).",
    "Discussion",
    "It grew fast.",
    "References",
    "1. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4.",
    "Figure Legends",
    "Figure 1: Clinical photograph of the lesion",
    "Figure 2: Axial scan of the orbit",
    "Figure 3: Flow of the patients",
]


def picture(fmt="PNG", size=(60, 40), seed=1) -> bytes:
    """A small noise image (noise keeps the file above the 2 KB icon threshold)."""
    rnd = random.Random(seed)
    img = Image.new("RGB", size)
    img.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(size[0] * size[1])])
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue()


def make_docx(path, texts=TEXT, picture_before=None):
    doc = Document()
    for t in texts:
        if picture_before and t.startswith(picture_before):
            doc.add_paragraph().add_run().add_picture(io.BytesIO(picture()), width=Inches(2))
        doc.add_paragraph(t)
    doc.save(path)
    return str(path)


def make_zip(tmp_path, files, texts=TEXT, picture_before=None):
    ms = make_docx(tmp_path / "Manuscript.docx", texts, picture_before)
    z = tmp_path / "sub.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(ms, "ART_1/Manuscript.docx")
        zf.writestr("ART_1/ART_1.xml", "<article/>")
        zf.writestr("ART_1/ART_1.png", picture(size=(60, 40)))  # the platform's id image beside the XML
        for name, data in files.items():
            zf.writestr(f"ART_1/{name}", data)
    return str(z)


def journal(**cfg):
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    st.settings["figures"]["insert"] = {**st.settings["figures"]["insert"], "enabled": True, **cfg}
    return st


def run(tmp_path, files, st=None, **kw):
    z = make_zip(tmp_path, files, **kw)
    return pipeline.edit(z, st or journal(), str(tmp_path / "out"), stages=("style",))


def body(path):
    return list(Document(path).element.body.iter(W + "p"))


def text(p):
    return "".join(t.text or "" for t in p.iter(W + "t"))


def pictures(path) -> list[int]:
    """Indexes of the body paragraphs holding a picture."""
    return [i for i, p in enumerate(body(path)) if docmodel.has_picture(p)]


FILES = {
    "Figure 1.jpg": picture("JPEG", (80, 40)),
    "Fig2.tif": picture("TIFF", (40, 80), seed=2),
    "Figure 3.pptx": b"PK" + b"x" * 4000,
    "First author photograph.jpg": picture("JPEG", seed=3),
    "Consent form.pdf": b"%PDF" + b"x" * 4000,
}


# --- file names --------------------------------------------------------------------------------

@pytest.mark.parametrize("name,num,part", [
    ("Figure 1.jpg", 1, ""), ("Fig2.tif", 2, ""), ("FIG2B.jpg", 2, "b"), ("figure-4 edited.png", 4, ""),
    ("1a.jpg", 1, "a"), ("F3.png", 3, ""), ("Picture 2.jpg", 2, ""), ("Figure 1pp.png", 1, ""),
    ("IJO_2845_25_Figure 1AandB.jpeg", 1, ""), ("Figure 3 - HE x40 1.jpg", 3, ""), ("Figure 1 (3).jpg", 1, ""),
])
def test_figure_number_from_file_name(name, num, part):
    assert figures_insert.parse_name(name)[:2] == (num, part)


@pytest.mark.parametrize("name", ["Flowchart.jpg", "Figures 1A-B and 2A-B.jpg", "Supplementary Figure 1.png",
                                  "Figure S2.jpg", "Image.png"])
def test_unclear_file_names_have_no_number(name):
    num, _, problem = figures_insert.parse_name(name)
    assert num is None and problem


def test_supplied_files_skip_non_figures(tmp_path):
    files = figures_insert.supplied_files(make_zip(tmp_path, FILES))
    assert {f.base: (f.number, bool(f.problem)) for f in files} == {
        "Figure 1.jpg": (1, False), "Fig2.tif": (2, False), "Figure 3.pptx": (3, True)}


# --- the pipeline ---------------------------------------------------------------------------

def test_figures_are_tracked_insertions_above_their_legends(tmp_path):
    out = run(tmp_path, FILES)
    tracked_path, clean_path = out.path("tracked"), out.path("clean")
    paras = body(tracked_path)
    pics = pictures(tracked_path)
    assert [text(paras[i + 1]) for i in pics] == ["Figure 1: Clinical photograph of the lesion",
                                                  "Figure 2: Axial scan of the orbit"]
    for i in pics:  # an inserted paragraph mark and an inserted run holding the drawing
        p = paras[i]
        assert p.find(f"{W}pPr/{W}rPr/{W}ins") is not None
        assert all(any(a.tag == W + "ins" for a in d.iterancestors()) for d in p.iter(W + "drawing"))
    # the image parts and their relationships exist; the TIFF went in as PNG
    z = zipfile.ZipFile(tracked_path)
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media) == 2 and any(n.endswith(".png") for n in media) and any(n.endswith(".jpg") for n in media)
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert rels.count("relationships/image") == 2
    # pending suggestions never reach the clean copy
    assert pictures(clean_path) == []
    # the PowerPoint figure is left to the editor, by name
    msgs = [w["message"] for w in out.summary["editor_warnings"] if w["rule"] == "figures.insert"]
    assert any("Figure 3.pptx" in m for m in msgs)
    assert out.summary["figure_inserts"]["total"] == 2
    assert out.summary["stage_info"]["style"]["figures"] == 2


def test_accepted_figures_reach_the_clean_copy_and_rejected_ones_do_not(tmp_path):
    run(tmp_path, FILES)
    outdir = str(tmp_path / "out")
    man, changes, inserts, qdicts = pipeline.load_records(outdir)
    figs = [n for n in inserts if figures_insert.is_figure_insert(n)]
    assert [n.figure for n in figs] == [1, 2] and all(n.status == Status.SUGGESTED for n in figs)
    figs[0].status, figs[1].status = Status.ACCEPTED, Status.REJECTED
    pipeline.save_records(outdir, inserts=inserts)
    out = pipeline.render_outputs(outdir)
    clean = out.path("clean")
    assert [text(body(clean)[i + 1]) for i in pictures(clean)] == ["Figure 1: Clinical photograph of the lesion"]
    assert Document(clean).inline_shapes[0].width > 0  # python-docx reads the picture back
    # the picture is as wide as the text area and keeps the file's 2:1 aspect ratio
    sec = Document(clean).sections[-1]
    shape = Document(clean).inline_shapes[0]
    assert shape.width == sec.page_width - sec.left_margin - sec.right_margin
    assert abs(shape.width / shape.height - 2) < 0.01
    assert len(pictures(out.path("tracked"))) == 1  # rejected: gone from the tracked copy too


def test_a_figure_the_manuscript_already_shows_is_not_touched(tmp_path):
    out = run(tmp_path, {"Figure 1.jpg": picture("JPEG"), "Figure 2.png": picture(seed=4)},
              picture_before="Figure 1:")
    figs = [n for n in pipeline.load_records(str(tmp_path / "out"))[2] if figures_insert.is_figure_insert(n)]
    assert [n.figure for n in figs] == [2]
    assert len(pictures(out.path("clean"))) == 1  # the author's own picture stays


def test_a_picture_no_legend_accounts_for_stops_insertion(tmp_path):
    out = run(tmp_path, {"Figure 1.jpg": picture("JPEG")}, picture_before="It grew fast.")
    assert out.summary["figure_inserts"]["total"] == 0
    assert any("no figure legend accounts for" in w["message"] for w in out.summary["editor_warnings"])


def test_several_candidates_and_missing_legends_are_warnings(tmp_path):
    out = run(tmp_path, {"Figure 1.jpg": picture("JPEG"), "Fig 1.png": picture(seed=5),
                         "Figure 7.png": picture(seed=6)})
    assert out.summary["figure_inserts"]["total"] == 1  # Figure 7: no legend -> picture + "Figure 7: ???"
    msgs = " | ".join(w["message"] for w in out.summary["editor_warnings"])
    assert "several files match" in msgs and "no legend for Figure 7" not in msgs


def test_parts_of_one_figure_go_in_one_paragraph(tmp_path):
    out = run(tmp_path, {"1a.jpg": picture("JPEG"), "1b.jpg": picture("JPEG", seed=7)})
    pics = pictures(out.path("tracked"))
    assert len(pics) == 1
    p = body(out.path("tracked"))[pics[0]]
    assert len(list(p.iter(W + "drawing"))) == 2
    assert text(body(out.path("tracked"))[pics[0] + 1]).startswith("Figure 1:")


def test_unnumbered_file_is_named_with_the_legend_it_may_be(tmp_path):
    out = run(tmp_path, {"Figure 1.jpg": picture("JPEG"), "Figure 2.jpg": picture("JPEG", seed=8),
                         "Flowchart.jpg": picture("JPEG", seed=9)})
    assert out.summary["figure_inserts"]["total"] == 2
    assert any("Flowchart.jpg" in w["message"] and "Figure 3 has no picture" in w["message"]
               for w in out.summary["editor_warnings"])


def test_position_below_puts_the_picture_under_its_legend(tmp_path):
    out = run(tmp_path, {"Figure 1.jpg": picture("JPEG")}, st=journal(position="below"))
    paras = body(out.path("tracked"))
    (i,) = pictures(out.path("tracked"))
    assert text(paras[i - 1]).startswith("Figure 1:")


def test_on_by_default_per_stylesheet_and_nothing_for_plain_docx_input(tmp_path):
    # S (41 of 42 sheets) and G266: "Image should be present first and below it the respective caption".
    for code in ("ijo", "mjdrdypu", "kjo", "idoj"):
        assert style.load(code).get("figures.insert.enabled"), code
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    src = make_docx(tmp_path / "ms.docx", ["Intro text here [Figure 1].", "Figure 1: Legend text."])
    out = pipeline.edit(src, st, str(tmp_path / "out"), stages=("style",))  # a .docx: no figure files to place
    assert out.summary["figure_inserts"]["total"] == 0 and pictures(out.path("tracked")) == []


# --- pictures already in the manuscript survive every renderer -------------------------------------

def test_existing_pictures_survive_tracked_and_clean_renders(tmp_path):
    src = make_docx(tmp_path / "ms.docx", ["Intro text here.", "Figure 1: Legend text."], picture_before="Figure 1:")
    ms = docmodel.load(src)
    pic = next(p for p in ms.paragraphs if docmodel.has_picture(p.element))
    pic.element.append(Document().add_paragraph("caption next to it").runs[0]._r)  # text in the picture paragraph
    ms.save(src)
    ms = docmodel.load(src)
    pic = next(p for p in ms.paragraphs if docmodel.has_picture(p.element))
    assert pic.text == "caption next to it"
    cs = ChangeSet([Change("c1", pic.id, 0, 7, "caption", "Caption", Layer.HOUSE_STYLE, Source.RULE, "test"),
                    Change("c2", pic.id, 7, 18, " next to it", "", Layer.HOUSE_STYLE, Source.RULE, "test")])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    for is_tracked in (True, False):
        out = tmp_path / f"out_{is_tracked}.docx"
        tracked.render(docmodel.load(src), cs, tracked=is_tracked).save(out)
        assert len(pictures(str(out))) == 1
        assert len(Document(str(out)).inline_shapes) == 1
        assert any(n.startswith("word/media/") for n in zipfile.ZipFile(out).namelist())


def test_docview_marks_picture_paragraphs(tmp_path):
    src = make_docx(tmp_path / "ms.docx", ["Intro text here.", "Figure 1: Legend text."], picture_before="Figure 1:")
    paras = docview.document(src)["paragraphs"]
    img = [p for p in paras if p["image"]]
    assert len(img) == 1 and img[0]["kind"] == "image"
    assert not any(p["image"] for p in paras if p["text"])


def test_fit_keeps_aspect_and_page_height():
    assert figures_insert.fit(200, 100, 6000, 9000) == (6000, 3000)
    assert figures_insert.fit(100, 400, 6000, 9000) == (2250, 9000)


def test_insert_records_round_trip():
    n = figures_insert.FigureInsert(id="fig_ins1", anchor="p0009", where="before", runs=[], role="figure",
                                    reason=figures_insert.REASON, figure=1,
                                    images=[{"file": "source/figures/a.jpg", "name": "Z/Figure 1.jpg",
                                             "width": 10, "height": 5}])
    d = json.loads(json.dumps(n.to_dict()))
    back = pipeline.insert_from_dict(d)
    assert isinstance(back, figures_insert.FigureInsert) and back.to_dict() == n.to_dict()
    assert back.text == "[Figure 1 image: Figure 1.jpg]"


def test_figure_without_legend_goes_in_with_placeholder_and_query(tmp_path):
    st = style.load("mjdrdypu")
    st.settings["references"]["lookup"] = False
    z = make_zip(tmp_path, {"Figure-1.jpeg": picture("JPEG"), "Figure-2.jpg": picture("JPEG", seed=3)},
                 texts=["Introduction", "Results are shown in [Figure 1] and [Figure 2].", "Discussion text."])
    out = pipeline.edit(z, st, str(tmp_path / "out"), stages=("style",))
    assert out.summary["figure_inserts"]["total"] == 2
    ms = docmodel.load(out.path("tracked"))
    pics = [i for i, p in enumerate(ms.paragraphs) if docmodel.has_picture(p.element)]
    after = [next(q.text for q in ms.paragraphs[i + 1:] if q.text.strip()) for i in pics]
    assert after == ["Figure 1: ???", "Figure 2: ???"]
    assert out.summary["queries"]["by_check"].get("figures.legend_missing") == 2
