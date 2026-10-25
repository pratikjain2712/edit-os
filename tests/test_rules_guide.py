"""The pre-editing guide's search-and-replace table (edit_os/rules_guide.py). Synthetic text only."""

import pytest
from docx import Document

from edit_os import docmodel, rules, rules_guide, style, tracked
from edit_os.changes import Status

ALL = {name: True for name in rules_guide.FINDERS}


def fix(text, rule, **cfg):
    """Apply one guide rule to a string."""
    out = text
    for s, e, new, _ in sorted(rules_guide.propose(text, {rule: True, **cfg}), reverse=True):
        out = out[:s] + new + out[e:]
    return out


@pytest.mark.parametrize("rule, before, after", [
    ("et_al_period", "Smith et al reported", "Smith et al. reported"),
    ("et_al_period", "as Rao et al[4] showed", "as Rao et al.[4] showed"),
    ("et_al_period", "Smith et al. reported", "Smith et al. reported"),
    ("ampersand", "anxiety & depression were common", "anxiety and depression were common"),
    ("ampersand", "Institute of Medical Sciences & Research", "Institute of Medical Sciences & Research"),
    ("ampersand", "stained with H&E and H & E", "stained with H&E and H & E"),
    ("ampersand", "Writing - review & editing.", "Writing - review & editing."),
    ("operator_spacing", "mean age 45.2±11.3 years", "mean age 45.2 ± 11.3 years"),
    ("operator_spacing", "a total of 2+3 lesions", "a total of 2 + 3 lesions"),
    ("operator_spacing", "CD4+8 cells and ER+ tumours", "CD4+8 cells and ER+ tumours"),
    ("operator_spacing", "at 32+3 weeks of gestation", "at 32+3 weeks of gestation"),
    ("operator_spacing", "aged 42,3±13,5 years", "aged 42,3±13,5 years"),
    ("nbsp", "from\u00a0June 2019", "from June 2019"),
    ("percent_space", "in 50 % of eyes", "in 50% of eyes"),
    ("percent_space", "in 50\u00a0% of eyes", "in 50% of eyes"),
    ("unary_minus", "a difference of -1.53 (CI -2.1 to -0.9)", "a difference of −1.53 (CI −2.1 to −0.9)"),
    ("unary_minus", "rho = -0.17", "rho = −0.17"),
    ("unary_minus", "ranges 5% -8% and (12) -3", "ranges 5% -8% and (12) -3"),
    ("unary_minus", "aged 10 -20 years, an anti-VEGF agent, IL-6", "aged 10 -20 years, an anti-VEGF agent, IL-6"),
    ("unary_comparison", "large holes (> 400 μm) and 20 to < 30 ng/mL", "large holes (>400 μm) and 20 to <30 ng/mL"),
    ("unary_comparison", "significant (P < 0.05), BMI > 30, scores < 7", "significant (P < 0.05), BMI > 30, scores < 7"),
    # the word before is the left operand: binary, spaces kept; "aged" is not an operand
    ("unary_comparison", "A P value < 0.05, a p-value < 0.01, (age ≥ 6 years)", "A P value < 0.05, a p-value < 0.01, (age ≥ 6 years)"),
    ("unary_comparison", "adults aged ≥ 60 years", "adults aged ≥60 years"),
    ("bracket_inner_space", "mean ( 80 ) and", "mean (80) and"),
    ("bracket_inner_space", 'query: TITLE ( digital ) OR TITLE ( "cardiac rehab" ) AND PUBYEAR > 2015',
     'query: TITLE ( digital ) OR TITLE ( "cardiac rehab" ) AND PUBYEAR > 2015'),
    ("value_unit_space", "10mm ports, 20mg daily and 5µg", "10 mm ports, 20 mg daily and 5 µg"),
    ("value_unit_space", "2.5cm lesion at 3000rpm, 532nm laser", "2.5 cm lesion at 3000 rpm, 532 nm laser"),
    ("value_unit_space", "10 mm, 5mmol, T2mm, 10-mm, mg/dL", "10 mm, 5mmol, T2mm, 10-mm, mg/dL"),
    ("value_unit_space", "IOP 35mmHg, 20nJ spots, 1gm/day", "IOP 35 mmHg, 20 nJ spots, 1 gm/day"),
    ("value_bracket_space", "males 20(80%) and 21(84)", "males 20 (80%) and 21 (84)"),
    ("value_bracket_space", "t(9;22) and IL6(+) and 2020(3)a", "t(9;22) and IL6(+) and 2020(3)a"),
    ("bracket_inner_space", "males ( 80 ) and (Table 1 )", "males (80) and (Table 1)"),
    ("double_period", "in 80.. Then", "in 80. Then"),
    ("double_period", "and so on... and …..", "and so on... and ….."),
    ("degree_space", "stored at 29 °C and 4° C", "stored at 29°C and 4°C"),
    ("degree_space", "an angle of 45 ° and 29°C", "an angle of 45 ° and 29°C"),
    ("multiplication", "measuring 10x20 mm and 4.5 X 3 cm", "measuring 10 × 20 mm and 4.5 × 3 cm"),
    ("multiplication", "a 2x increase, x-ray, 150×10⁹/L", "a 2x increase, x-ray, 150×10⁹/L"),
    ("multiplication", "measuring 3.5×3.4 mm and 2 × 3 mm", "measuring 3.5 × 3.4 mm and 2 × 3 mm"),
    ("multiplication", "WBC 4×109/L and 2 ×10⁹, P < 5×10-8", "WBC 4×109/L and 2 ×10⁹, P < 5×10-8"),
    ("technetium", "scan with 99 mTc", "scan with 99mTc"),
    ("p_value_capital", "(p = 0.03) and p<0.05 and a p-value", "(P = 0.03) and P<0.05 and a P-value"),
    ("p_value_capital", "group p and type p.", "group p and type p."),
])
def test_finder(rule, before, after):
    assert fix(before, rule) == after


def test_urls_dois_and_emails_are_protected():
    text = "See https://example.org/a&b et al and doi:10.1000/x..y or a.b@mail.com &"
    assert rules_guide.propose(text, ALL) == [(text.index("et al") + 3, text.index("et al") + 5, "al.",
                                               "et_al_period")]


def test_formula_paragraph_skips_arithmetic_rules():
    text = "Y = 2x3 + -1/n"
    names = {r for *_, r in rules_guide.propose(text, ALL)}
    assert not names & {"operator_spacing", "unary_minus", "unary_comparison", "multiplication"}


def test_reference_like_paragraph_is_skipped():
    assert rules_guide.propose("Rao A, Das B & Roy C. A trial. J Test 2013;52(1):53-61.", ALL) == []


def test_disabled_rules_propose_nothing():
    assert rules_guide.propose("Smith et al & 50 % of 10mm", {}) == []


@pytest.fixture
def manuscript(tmp_path):
    doc = Document()
    for text in ["Introduction",
                 "Smith et al found anxiety & depression in 50 % of 120 patients aged 45.2±11.3 years.",
                 "Lesions measured 10x20mm ( 80 ).. Then p = 0.03.",
                 "References",
                 "1. Smith A & Rao B. Anxiety & depression. J Test 2001;3:1-9."]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "50 % & 10mm"
    path = tmp_path / "g.docx"
    doc.save(path)
    return str(path)


def _apply(path, st, tmp_path):
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["back_matter"]["add_missing"] = False  # keeps paragraph positions fixed
    cs = rules.run_all(docmodel.load(path), st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(str(out)).paragraphs]


def test_defaults_base(manuscript, tmp_path):
    t = _apply(manuscript, style.load("base"), tmp_path)
    assert t[1] == "Smith et al. found anxiety and depression in 50% of 120 patients aged 45.2 ± 11.3 years."
    assert t[2] == "Lesions measured 10 × 20 mm (80). Then P = 0.03."  # G294: P capital
    assert "&" in t[4]   # never in references
    assert t[5] == "50 % & 10mm"  # tables are off by default


def test_operators_unspaced_when_statistics_spacing_is_off(manuscript, tmp_path):
    """A journal that turns statistics.space_around_operators off: the guide rule must not bring it back."""
    st = style.load("ijo")
    st.settings["statistics"]["space_around_operators"] = False
    t = _apply(manuscript, st, tmp_path)
    assert "45.2±11.3" in t[0] and "Smith et al. found anxiety and depression in 50% " in t[0]  # "Introduction" gone
    assert "45.2 ± 11.3" in _apply(manuscript, style.load("ijo"), tmp_path)[0]  # IJO: G297 spaces operators
    st.settings["guide_replace"]["operator_spacing"] = "always"
    assert "45.2 ± 11.3" in _apply(manuscript, st, tmp_path)[0]


def test_changes_carry_old_text_and_reason(manuscript):
    ms = docmodel.load(manuscript)
    changes = rules_guide.rule_guide(ms, style.load("base"))
    assert changes and all(ms.by_id(c.para_id).text[c.start:c.end] == c.old for c in changes)
    assert all(c.reason.startswith("guide.") for c in changes)


@pytest.mark.parametrize("rule, before, after", [
    ("trailing_space", "The study ended.  ", "The study ended."),
    ("trailing_space", "   ", "   "),
    ("leading_space", "  The study ended.", "The study ended."),
    ("leading_space", "\tIndented by a tab.", "\tIndented by a tab."),
    ("double_space", "the  control group and   others", "the control group and others"),
    ("smart_quotes", "the caregivers' knowledge and the study's aim", "the caregivers’ knowledge and the study’s aim"),
    ("smart_quotes", 'no "safe place" and a \'term\' here', "no “safe place” and a ‘term’ here"),
    ("smart_quotes", "a height of 5'10\" was rare", "a height of 5'10\" was rare"),
    ("smart_quotes", "see https://example.org/it's", "see https://example.org/it's"),
    ("abbreviation_colon", "Abbreviations: CI, confidence interval; aOR, adjusted odds ratio; HDL-C, HDL cholesterol",
     "Abbreviations: CI: confidence interval; aOR: adjusted odds ratio; HDL-C: HDL cholesterol"),
    ("abbreviation_colon", "Abbreviations: CI, confidence interval", "Abbreviations: CI, confidence interval"),
    ("abbreviation_colon", "We used CI, confidence intervals; OR, odds ratios", "We used CI, confidence intervals; OR, odds ratios"),
])
def test_spacing_and_quote_finders(rule, before, after):
    assert fix(before, rule) == after


def test_spacing_finders_follow_the_guide():
    cfg = rules_guide.settings(style.load("ijo"))
    assert cfg.get("trailing_space") and cfg.get("leading_space")  # G143: spaces at the start and end go
    assert not any(cfg.get(k) for k in ("double_space", "smart_quotes", "abbreviation_colon"))  # not in the guide
    assert all(rules_guide.settings(style.load("jehp")).get(k)
               for k in ("trailing_space", "leading_space", "double_space", "smart_quotes", "p_value_capital"))


def test_p_value_only_the_probability():
    from edit_os import rules_format, rules_guide
    hits = lambda t: [t[s:e] for s, e, _ in rules_guide.find_p_value(t, {})]
    for t in ["(p = 0.03)", "p<0.001", "the p-value", "p values", "(p = .04)", "p=NS"]:
        assert hits(t) == ["p"] and rules_format._P_VALUE.search(t), t
    for t in ["p53 mutation", "pH 7.4", "see p. 465-78", "pp. 12-15", "p-ANCA", "p.o.", "the p wave", "5p deletion",
              "IL-6p", "group p and q"]:
        assert hits(t) == [] and not rules_format._P_VALUE.search(t), t
