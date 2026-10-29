"""API service with FastAPI's TestClient: upload, background job, decisions, re-render, downloads, auth."""

import io
import os
import re
import time
import zipfile

import pytest
from fastapi.testclient import TestClient

from edit_os import docmodel, style
from server.app import create_app
from server.config import Settings
from server.storage import S3Storage

from test_pipeline import TEXT, make_docx, offline_copyedit, offline_style  # noqa: F401  (fixture)

KEY = {"X-API-Key": "test-key"}


@pytest.fixture
def client(tmp_path, monkeypatch):
    orig = style.load
    st = offline_style()
    monkeypatch.setattr(style, "load", lambda code: st if code.lower() == "ijo" else orig(code))
    settings = Settings(data_dir=str(tmp_path / "data"), api_keys=["test-key"], workers=1,
                        public_url="https://editos.example")
    with TestClient(create_app(settings)) as c:
        yield c


def upload(client, tmp_path, name="ms.docx", journal="ijo"):
    path = make_docx(tmp_path / name)
    with open(path, "rb") as f:
        r = client.post("/api/jobs", headers=KEY, data={"journal": journal, "external_id": "IJO_1"},
                        files={"file": (name, f, "application/octet-stream")})
    return r


def wait(client, job_id, headers=KEY, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        job = client.get(f"/api/jobs/{job_id}", headers=headers).json()
        if job["status"] in ("done", "failed"):
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_auth_required(client, tmp_path):
    assert client.get("/api/health").status_code == 200
    assert client.get("/api/jobs").status_code == 401
    assert client.get("/api/jobs", headers={"X-API-Key": "wrong"}).status_code == 401
    assert upload(client, tmp_path).status_code == 202
    r = client.post("/api/jobs", data={"journal": "ijo"}, files={"file": ("a.docx", b"x")})
    assert r.status_code == 401


def test_no_keys_configured_refuses(tmp_path):
    with TestClient(create_app(Settings(data_dir=str(tmp_path / "d")))) as c:
        assert c.get("/api/jobs", headers=KEY).status_code == 503


def test_rejects_bad_input(client):
    r = client.post("/api/jobs", headers=KEY, data={"journal": "ijo"}, files={"file": ("a.pdf", b"x")})
    assert r.status_code == 422
    r = client.post("/api/jobs", headers=KEY, data={"journal": "nope"}, files={"file": ("a.docx", b"x")})
    assert r.status_code == 422


def test_job_flow(client, tmp_path):
    r = upload(client, tmp_path)
    assert r.status_code == 202, r.text
    created = r.json()
    job_id = created["id"]
    assert created["review_url"].startswith(f"https://editos.example/review/{job_id}?token=")
    job = wait(client, job_id)
    assert job["status"] == "done", job["error"]
    assert job["summary"]["changes"]["total"] > 0 and job["external_id"] == "IJO_1"
    assert client.get("/api/jobs?external_id=IJO_1", headers=KEY).json()[0]["id"] == job_id

    # Suggestions with paragraph context, grouped by layer.
    ch = client.get(f"/api/jobs/{job_id}/changes", headers=KEY).json()
    assert "language" in ch["layers"]
    colour = next(i for i in ch["items"] if i["old"] == "colour")
    assert colour["status"] == "suggested" and colour["context"]["before"].endswith("The ")
    assert "lesion" in colour["context"]["after"]
    rand = next(i for i in ch["items"] if i["old"] == "randomised")
    qs = client.get(f"/api/jobs/{job_id}/queries", headers=KEY).json()["items"]
    missing = next(q for q in qs if q["check"] == "references.missing_entry")
    assert missing["anchor"] == "[3]"

    # Decisions persist.
    r = client.post(f"/api/jobs/{job_id}/decisions", headers=KEY,
                    json={"ids": [colour["id"]], "status": "accepted", "editor": "te1"})
    assert r.json()["updated"] == 1
    client.post(f"/api/jobs/{job_id}/decisions", headers=KEY, json={"ids": [rand["id"]], "status": "rejected"})
    client.post(f"/api/jobs/{job_id}/decisions", headers=KEY, json={"ids": [missing["key"]], "status": "rejected"})
    assert client.post(f"/api/jobs/{job_id}/decisions", headers=KEY,
                       json={"ids": ["nope"], "status": "accepted"}).status_code == 422
    ch = client.get(f"/api/jobs/{job_id}/changes", headers=KEY).json()
    status = {i["id"]: i["status"] for i in ch["items"]}
    assert status[colour["id"]] == "accepted" and status[rand["id"]] == "rejected"

    # Bulk accept by layer only touches pending ones.
    r = client.post(f"/api/jobs/{job_id}/decisions", headers=KEY, json={"layer": "language", "status": "accepted"})
    assert rand["id"] not in r.json()["ids"]
    ch = client.get(f"/api/jobs/{job_id}/changes?layer=language", headers=KEY).json()
    assert {i["status"] for i in ch["items"]} == {"accepted", "rejected"}

    # Re-render reflects the decisions.
    rep = client.post(f"/api/jobs/{job_id}/render", headers={**KEY, "X-Editor": "Asha R."}).json()
    assert rep["queries"]["rejected"] == 1
    r = client.get(f"/api/jobs/{job_id}/files/clean", headers=KEY)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    p = tmp_path / "clean.docx"
    p.write_bytes(r.content)
    first = docmodel.load(str(p)).paragraphs[0].text
    assert first.startswith("Technical Editing done on ") and first.endswith(":- Asha R.")  # the editor signs
    assert "signoff.name" not in {c["check"] for c in rep["qc"]["checks"] if not c["passed"]}
    text = _body(p)[2]
    assert "color of the lesion" in text and "randomised" in text
    comments = zipfile.ZipFile(io.BytesIO(client.get(f"/api/jobs/{job_id}/files/tracked", headers=KEY).content)) \
        .read("word/comments.xml").decode()
    assert "has no entry" not in comments
    for kind in ("queries", "qc_report", "qc_json", "changes"):
        assert client.get(f"/api/jobs/{job_id}/files/{kind}", headers=KEY).status_code == 200, kind
    assert client.get(f"/api/jobs/{job_id}/qc", headers=KEY).json()["qc"]["checks"]


def test_review_token_scoped_to_its_job(client, tmp_path):
    a = upload(client, tmp_path, "a.docx").json()
    b = upload(client, tmp_path, "b.docx").json()
    tok = {"X-Review-Token": a["review_token"]}
    assert wait(client, a["id"], headers=tok)["status"] == "done"
    assert "review_token" not in client.get(f"/api/jobs/{a['id']}", headers=tok).json()
    assert client.get(f"/api/jobs/{a['id']}/changes", headers=tok).status_code == 200
    assert client.get(f"/api/jobs/{a['id']}/files/clean?token={a['review_token']}").status_code == 200
    assert client.get(f"/api/jobs/{b['id']}", headers=tok).status_code == 401
    assert client.get("/api/jobs", headers=tok).status_code == 401
    wait(client, b["id"])


def test_failed_job_reports_error(client, tmp_path):
    r = client.post("/api/jobs", headers=KEY, data={"journal": "ijo"}, files={"file": ("bad.docx", b"not a docx")})
    job = wait(client, r.json()["id"])
    assert job["status"] == "failed" and job["error"]
    assert client.get(f"/api/jobs/{job['id']}/changes", headers=KEY).status_code == 409


class FakeS3:
    def __init__(self):
        self.objects = {}

    def upload_file(self, path, bucket, key, ExtraArgs=None):
        self.objects[(bucket, key)] = open(path, "rb").read()

    def download_file(self, bucket, key, dest):
        with open(dest, "wb") as f:
            f.write(self.objects[(bucket, key)])

    def get_paginator(self, name):
        objs = self.objects

        class P:
            def paginate(self, Bucket, Prefix):
                yield {"Contents": [{"Key": k} for b, k in objs if b == Bucket and k.startswith(Prefix)]}
        return P()


def test_s3_storage_round_trip(tmp_path):
    fake = FakeS3()
    s = S3Storage(str(tmp_path / "a"), "bucket", "eo/", client=fake)
    os.makedirs(s.job_dir("j1") + "/out")
    open(s.job_dir("j1") + "/out/x.json", "w").write("{}")
    s.save("j1")
    assert ("bucket", "eo/j1/out/x.json") in fake.objects
    s2 = S3Storage(str(tmp_path / "b"), "bucket", "eo/", client=fake)
    assert s2.restore("j1") and os.path.exists(s2.job_dir("j1") + "/out/x.json")
    assert not s2.restore("missing")


def test_editor_rewords_a_suggestion(client, tmp_path):
    job = wait(client, upload(client, tmp_path).json()["id"])
    items = client.get(f"/api/jobs/{job['id']}/changes", headers=KEY).json()["items"]
    target = next(i for i in items if i["kind"] == "change" and not (i["format"] and i["old"] == i["new"]))
    r = client.post(f"/api/jobs/{job['id']}/edit", headers=KEY, json={"id": target["id"], "new": "EDITOR TEXT"})
    assert r.status_code == 200
    after = {i["id"]: i for i in client.get(f"/api/jobs/{job['id']}/changes", headers=KEY).json()["items"]}
    assert after[target["id"]]["new"] == "EDITOR TEXT" and after[target["id"]]["status"] == "accepted"
    assert after[target["id"]]["edited"] and after[target["id"]]["source"] == "editor"
    assert client.post(f"/api/jobs/{job['id']}/render?final=true", headers=KEY).status_code == 200
    from docx import Document
    import io
    clean = client.get(f"/api/jobs/{job['id']}/files/clean", headers=KEY).content
    assert "EDITOR TEXT" in "\n".join(p.text for p in Document(io.BytesIO(clean)).paragraphs)
    bad = client.post(f"/api/jobs/{job['id']}/edit", headers=KEY, json={"id": "nope", "new": "x"})
    assert bad.status_code == 422
    two = client.post(f"/api/jobs/{job['id']}/edit", headers=KEY, json={"id": target["id"], "new": "a\nb"})
    assert two.status_code == 422


def test_document_view(client, tmp_path):
    created = upload(client, tmp_path).json()
    job = wait(client, created["id"])
    r = client.get(f"/api/jobs/{job['id']}/document", headers=KEY)
    assert r.status_code == 200
    doc = r.json()
    paras = doc["paragraphs"]
    assert [p["text"] for p in paras] == TEXT and doc["journal"].lower() == "ijo"
    kinds = {p["text"]: (p["kind"], p["level"]) for p in paras}
    assert kinds["Introduction"] == ("heading", 1) and kinds["References"] == ("heading", 1)
    assert kinds[TEXT[2]][0] == "body" and kinds[TEXT[-1]][0] == "reference"
    assert all("".join(r["text"] for r in p["runs"]) == p["text"] for p in paras)
    # Change offsets refer to the same paragraph texts.
    items = client.get(f"/api/jobs/{job['id']}/changes", headers=KEY).json()["items"]
    texts = {p["id"]: p["text"] for p in paras}
    assert all(texts[i["para_id"]][i["start"]:i["end"]] == i["old"] for i in items if i["kind"] == "change")
    assert client.get(f"/api/jobs/{job['id']}/document").status_code == 401
    assert client.get(f"/api/jobs/{job['id']}/document?token={created['review_token']}").status_code == 200


def test_editor_direct_edits(client, tmp_path):
    created = upload(client, tmp_path).json()
    job = wait(client, created["id"])
    jid = job["id"]
    paras = client.get(f"/api/jobs/{jid}/document", headers=KEY).json()["paragraphs"]
    target = next(p for p in paras if p["text"] == TEXT[2])
    items = client.get(f"/api/jobs/{jid}/changes", headers=KEY).json()["items"]
    colour = next(i for i in items if i["old"] == "colour")
    client.post(f"/api/jobs/{jid}/decisions", headers=KEY, json={"ids": [colour["id"]], "status": "accepted"})
    base = TEXT[2].replace("colour", colour["new"])
    url = f"/api/jobs/{jid}/changes"

    # Auth: no key -> 401; the job's review token works.
    body = {"para_id": target["id"], "base": base, "text": base.replace("40", "41")}
    assert client.post(url, json=body).status_code == 401
    r = client.post(url, json=body, headers={"X-Review-Token": created["review_token"], "X-Editor": "te1"})
    assert r.status_code == 201, r.text
    [made] = r.json()["created"]
    assert (made["old"], made["new"], made["source"], made["status"]) == ("40", "41", "editor", "accepted")
    assert made["author"] == "te1"
    listed = {i["id"]: i for i in client.get(url, headers=KEY).json()["items"]}
    assert listed[made["id"]]["direct"] and listed[made["id"]]["status"] == "accepted"

    # Touching the accepted suggestion -> 409 with a clear message; nothing stored.
    base2 = base.replace("40", "41")
    bad = client.post(url, headers=KEY, json={"para_id": target["id"], "base": base2,
                                              "text": base2.replace(colour["new"], "hue")})
    assert bad.status_code == 409 and "accepted suggestion" in bad.json()["detail"]
    # Stale base -> 409; unknown paragraph / paragraph break -> 422.
    assert client.post(url, headers=KEY, json={"para_id": target["id"], "base": TEXT[2], "text": "x"}).status_code == 409
    assert client.post(url, headers=KEY, json={"para_id": "p9999", "base": "", "text": "x"}).status_code == 422
    assert client.post(url, headers=KEY, json={"para_id": target["id"], "base": base2,
                                               "text": base2 + "\nmore"}).status_code == 422

    # Editing the same paragraph again replaces the earlier direct edit; overlapping pending suggestions are set aside.
    rand = next(i for i in items if i["old"] == "randomised")
    r = client.post(url, headers=KEY, json={"para_id": target["id"], "base": base2,
                                            "text": base2.replace("41", "42").replace("randomised", "enrolled")})
    assert r.status_code == 201, r.text
    assert r.json()["replaced"] == [made["id"]] and rand["id"] in r.json()["superseded"]
    status = {i["id"]: i["status"] for i in client.get(url, headers=KEY).json()["items"]}
    assert status[made["id"]] == "rejected" and status[rand["id"]] == "rejected"

    # Rendered into the clean copy, and the tracked copy renders without conflicts.
    rep = client.post(f"/api/jobs/{jid}/render", headers=KEY).json()
    assert rep["render_conflicts"] == []
    p = tmp_path / "clean.docx"
    p.write_bytes(client.get(f"/api/jobs/{jid}/files/clean", headers=KEY).content)
    text = _body(p)[2]
    assert "color of the lesion was noted in 42 enrolled patients" in text
    # After the render saved the merged records, the direct edits are still there exactly once.
    again = [i for i in client.get(url, headers=KEY).json()["items"] if i.get("direct")]
    assert len({i["id"] for i in again}) == len(again) and any(i["status"] == "accepted" for i in again)


def _body(path):
    """Paragraph texts of an output after the sign-off lines at the top (header.py)."""
    texts = [p.text for p in docmodel.load(str(path)).paragraphs]
    while texts and re.match(r"(?:Technical Editing|Copyediting) done on", texts[0]):
        texts.pop(0)
    return texts


def test_editor_header_decoded():
    from server.app import _ed
    assert _ed("Ragavi%20Ravi") == "Ragavi Ravi"
    assert _ed("S%C3%B6phiya") == "Söphiya"
    assert _ed(None) == ""


def test_sign_in_with_a_user(tmp_path):
    settings = Settings(data_dir=str(tmp_path / "data"), api_keys=["test-key"], workers=1,
                        users={"reviewer": "pw-123"})
    with TestClient(create_app(settings)) as c:
        assert c.get("/api/me").status_code == 401
        assert c.post("/api/login", json={"username": "reviewer", "password": "nope"}).status_code == 401
        r = c.post("/api/login", json={"username": "Reviewer", "password": "pw-123"})
        assert r.status_code == 200 and r.json()["user"] == "reviewer"
        tok = {"X-API-Key": r.json()["token"]}
        assert c.get("/api/me", headers=tok).json() == {"user": "reviewer", "kind": "session"}
        assert c.get("/api/journals", headers=tok).status_code == 200
        forged = tok["X-API-Key"][:-1] + ("0" if tok["X-API-Key"][-1] != "0" else "1")
        assert c.get("/api/journals", headers={"X-API-Key": forged}).status_code == 401


def test_session_tokens_expire():
    from server import sessions
    t = sessions.issue("secret", "a", hours=1, now=1000)
    assert sessions.verify("secret", t, now=1000) == "a"
    assert sessions.verify("secret", t, now=1000 + 3601) is None
    assert sessions.verify("other", t, now=1000) is None
    assert sessions.verify("secret", "s1.garbage", now=1000) is None


def test_journals_list_the_file_name_prefixes(client):
    js = {j["code"]: j for j in client.get("/api/journals", headers=KEY).json()}
    assert js["ijo"]["prefixes"] == ["ijo"]
    assert "neurol" in js["ni"]["prefixes"] and js["ni"]["prefixes"][0] == "ni"


def test_download_signs_with_the_editor_name(client, tmp_path):
    # a file rendered with nobody's name ("Technical Editing done on <date>:- ???") is signed by whoever downloads
    # it, as it was rendered (final or not); a client file never carries "???"
    job_id = upload(client, tmp_path).json()["id"]
    wait(client, job_id)

    def first_line(headers):
        doc = client.get(f"/api/jobs/{job_id}/files/clean", headers=headers).content
        xml = zipfile.ZipFile(io.BytesIO(doc)).read("word/document.xml").decode()
        return re.sub(r"<[^>]+>", "", xml.split("</w:p>")[0])

    assert first_line(KEY).endswith("???")
    assert first_line({**KEY, "X-Editor": "Asha%20Rao"}).endswith("Asha Rao")
    assert first_line(KEY).endswith("Asha Rao")  # stays signed


def test_download_named_after_the_upload():
    from server.app import download_name
    assert download_name("JEHP_73_26.zip", "Manuscript", "/x/Manuscript_tracked.docx") == "JEHP_73_26_tracked.docx"
    assert download_name("IJO 1415_26 (rev).docx", "IJO 1415_26 (rev)", "/x/IJO 1415_26 (rev)_queries.xlsx") \
        == "IJO 1415_26 _rev_queries.xlsx"
    assert download_name(None, "Manuscript", "/x/Manuscript_clean.docx") == "Manuscript_clean.docx"
