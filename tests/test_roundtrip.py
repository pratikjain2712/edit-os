import pytest
from docx import Document
from docx.oxml.ns import qn

from edit_os import docmodel, references, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status


@pytest.fixture
def sample(tmp_path):
    doc = Document()
    doc.add_heading("Early mobilisation after hip fracture", level=1)
    p = doc.add_paragraph("Patients ")
    p.add_run("mobilized within 24 hrs").bold = True
    p.add_run(" of surgery had a shorter stay (9 Vs. 12 days; p<.001).")
    doc.add_paragraph("Their was no signifcant difference in mortality [3, 4].")
    doc.add_heading("References", level=1)
    doc.add_paragraph("1. Okafor N, Lindqvist P., Hale R. Timing of first mobilisation. "
                      "J. Orthop. Rehabil., 2017, vol. 12, pp 101-9.")
    doc.add_paragraph("[2] Hale R, Okafor N. Discharge planning after hip surgery. "
                      "Geriatr Surg Rev 2019;7:12-18.")
    doc.add_paragraph("")
    doc.add_paragraph("Table 1. Baseline characteristics")
    path = tmp_path / "sample.docx"
    doc.save(path)
    return str(path)


def change(cid, para, text, old, new, layer=Layer.LANGUAGE, nth=0, status=Status.ACCEPTED):
    start = -1
    for _ in range(nth + 1):
        start = text.index(old, start + 1)
    return Change(cid, para, start, start + len(old), old, new, layer, Source.AI, status=status)


def views(path):
    """(accepted text, rejected text) per paragraph, read straight from the XML."""
    out = []
    for p in Document(path).element.body.iter(qn("w:p")):
        acc, rej = [], []
        for el in p.iter(qn("w:t"), qn("w:delText")):
            in_del = any(a.tag == qn("w:del") for a in el.iterancestors())
            in_ins = any(a.tag == qn("w:ins") for a in el.iterancestors())
            if el.tag == qn("w:t") and not in_del:
                acc.append(el.text or "")
            if (el.tag == qn("w:delText") or not in_ins):
                rej.append(el.text or "")
        out.append(("".join(acc), "".join(rej)))
    return out


@pytest.fixture
def changeset(sample):
    ms = docmodel.load(sample)
    p2, p3 = ms.paragraphs[1], ms.paragraphs[2]
    return ChangeSet([
        change("c1", p2.id, p2.text, "mobilized", "mobilised", Layer.HOUSE_STYLE),
        change("c2", p2.id, p2.text, "24 hrs", "24 h", Layer.NUMBERS),  # spans bold run end
        change("c3", p2.id, p2.text, "Vs.", "vs", Layer.HOUSE_STYLE),
        change("c4", p2.id, p2.text, "p<.001", "P < 0.001", Layer.NUMBERS),
        change("c5", p3.id, p3.text, "Their", "There"),
        change("c6", p3.id, p3.text, "signifcant", "significant"),
        change("c7", p3.id, p3.text, "[3, 4]", "[3,4]", Layer.REFERENCES),
        Change("c8", p3.id, 0, 0, "", "Overall, ", Layer.LANGUAGE, Source.EDITOR, status=Status.ACCEPTED),
        change("c9", p3.id, p3.text, "no ", "no", status=Status.REJECTED),  # must be ignored
    ])


def test_tracked_accept_and_reject_views(sample, changeset, tmp_path):
    before = [p.text for p in docmodel.load(sample).paragraphs]
    out = tmp_path / "tracked.docx"
    tracked.render(docmodel.load(sample), changeset, tracked=True).save(out)
    v = views(out)
    assert v[1][0] == "Patients mobilised within 24 h of surgery had a shorter stay (9 vs 12 days; P < 0.001)."
    assert v[2][0] == "Overall, There was no significant difference in mortality [3,4]."
    assert [r for _, r in v] == before  # rejecting everything restores the original


def test_inserted_text_keeps_formatting(sample, changeset, tmp_path):
    out = tmp_path / "tracked.docx"
    tracked.render(docmodel.load(sample), changeset, tracked=True).save(out)
    p = list(Document(out).element.body.iter(qn("w:p")))[1]
    bold_inserted = [
        t.text for ins in p.iter(qn("w:ins")) for r in ins.iter(qn("w:r"))
        if r.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None for t in r.iter(qn("w:t"))
    ]
    assert "mobilised" in bold_inserted and "24 h" in bold_inserted


def test_clean_copy(sample, changeset, tmp_path):
    out = tmp_path / "clean.docx"
    tracked.render(docmodel.load(sample), changeset, tracked=False).save(out)
    ms = docmodel.load(str(out))
    assert ms.paragraphs[2].text == "Overall, There was no significant difference in mortality [3,4]."
    body = Document(out).element.body
    assert not list(body.iter(qn("w:ins"))) and not list(body.iter(qn("w:del")))


def test_stale_change_is_rejected(sample):
    ms = docmodel.load(sample)
    p = ms.paragraphs[2]
    bad = ChangeSet([Change("x", p.id, 0, 5, "Thier", "There", Layer.LANGUAGE, Source.AI, status=Status.ACCEPTED)])
    with pytest.raises(tracked.ChangeConflict):
        tracked.render(ms, bad)


def test_reference_extraction(sample):
    refs = references.extract(docmodel.load(sample))
    assert [r.number for r in refs] == [1, 2]
    assert refs[0].text.startswith("Okafor N") and refs[0].label == "1. "
    assert refs[1].text.startswith("Hale R") and refs[1].label == "[2] "


def test_change_on_reference_entry(sample, tmp_path):
    ms = docmodel.load(sample)
    ref = references.extract(ms)[0]
    para = ms.by_id(ref.para_id)
    old = "J. Orthop. Rehabil., 2017, vol. 12, pp 101-9."
    s = para.text.index(old)
    cs = ChangeSet([Change("r1", ref.para_id, s, s + len(old), old, "J Orthop Rehabil 2017;12:101-109.",
                           Layer.REFERENCES, Source.RULE, status=Status.ACCEPTED)])
    out = tmp_path / "refs.docx"
    tracked.render(ms, cs, tracked=False).save(out)
    assert docmodel.load(str(out)).by_id(ref.para_id).text.endswith("J Orthop Rehabil 2017;12:101-109.")
