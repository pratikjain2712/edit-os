"""JEHP sections, headings, abstract / key words and back matter (synthetic text): headings sharing a paragraph with
text, misspelt main headings, nested sub-headings, the held "Methods" renaming apart from the case fix, key words on
an abstract line and MeSH terms, the funding statement beside the journal's heading, run-in statements."""

from docx import Document
from lxml import etree

from edit_os import docmodel, queries, rules, rules_sections, style, tracked
from edit_os.changes import Status

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail, in vitro and in vivo.")
REFS = ["References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."]


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        if isinstance(text, tuple):  # bold
            doc.add_paragraph().add_run(text[0]).bold = True
        else:
            doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def jehp():
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    return st


def changes(path, st, only=None):
    return rules.run_all(docmodel.load(path), st, only=only)


def run(path, st=None, only=None):
    st = st or jehp()
    cs = changes(path, st, only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    ms = tracked.render(docmodel.load(path), cs, tracked=False)
    out = str(path).replace(".docx", "_out.docx")
    ms.save(out)
    return [p.text for p in docmodel.load(out).paragraphs]


def test_heading_sharing_a_paragraph_with_text(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "DISCUSSION \n" + BODY, BODY + " [18].\nConclusion", BODY] + REFS)
    out = run(path, only=["heading_tags", "sections_headings"])
    assert out[2:7] == ["<H1>Discussion", BODY, BODY + " [18].", "<H1>Conclusion", BODY]
    qs = queries.collect(docmodel.load(path), jehp())
    assert not any(q.check == "sections.conclusion" for q in qs.queries)  # the "Conclusion" line is there


def test_misspelt_main_headings(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, ("Material and Method: ",), BODY, ("Result ",), BODY,
                           "Discussion", BODY] + REFS)
    out = run(path, only=["heading_tags", "sections_headings", "subheadings"])
    assert "<H1>Materials and Methods:" in out[2]  # G137: the colon stays
    assert out[4].startswith("<H1>Results")


def test_nested_subheadings(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Results", ("Study Selection",), BODY,
                           ("In Vitro Evidence",), ("Anti-Candida Activity (In Vitro)",), BODY,
                           ("Wound Healing (Heine Model)",), BODY,
                           ("Animal and Preclinical Studies",), ("Oral Mucosal Healing",), BODY,
                           ("Annual Production Growth",), "Figure 3: Annual production", BODY,
                           "Discussion", ("Antimicrobial Activity: Synthesis and Comparison",), BODY] + REFS)
    out = run(path, only=["heading_tags", "subheadings"])
    assert "<H2>Study selection" in out
    assert "<H2>In vitro evidence" in out and "<H3>Anti-Candida activity (in vitro)" in out  # a genus keeps its capital
    assert "<H3>Wound healing (Heine Model)" in out  # "Heine" is never lower case: the bracket keeps its capitals
    assert "<H2>Animal and preclinical studies" in out and "<H3>Oral mucosal healing" in out
    assert "<H3>Annual production growth" in out  # over a figure caption
    assert "<H2>Antimicrobial activity: Synthesis and comparison" in out


def test_run_in_methods_label_split(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Materials and Methods"])
    doc = Document(path)
    p = doc.add_paragraph()
    p.add_run("Study Design:").bold = True
    p.add_run(" " + BODY)
    for t in [("Data Sources",), BODY, "Results", BODY] + REFS:
        if isinstance(t, tuple):
            doc.add_paragraph().add_run(t[0]).bold = True
        else:
            doc.add_paragraph(t)
    doc.save(path)
    out = run(path, only=["heading_tags", "subheadings"])
    i = out.index("<H2>Study design")
    assert out[i + 1] == BODY and "<H2>Data sources" in out


def test_textbox_paragraph_is_no_heading():
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    box = etree.fromstring(f'<w:txbxContent xmlns:w="{w}"><w:p><w:r><w:t>Identification</w:t></w:r></w:p></w:txbxContent>')
    para = docmodel.Paragraph("p1", box[0], "", "Identification")
    assert rules_sections.in_textbox(para)


def test_methods_rename_in_the_abstract(tmp_path):
    path = make(tmp_path, [("ABSTRACT",), "BACKGROUND: " + BODY, "MATERIALS AND METHODS: " + BODY, "RESULTS: " + BODY,
                           "CONCLUSION: " + BODY, "Keywords: Health; Nursing; Students", "Introduction", BODY] + REFS)
    st = jehp()
    cs = changes(path, st, ["sections_abstract"])
    mm = [c for c in cs.changes if c.old == "MATERIALS AND METHODS"]
    # S (struck through, still live): "Methods" -- held for the editor again (owner, 2026-10-30: the final text has it
    # 13% of the time): the capitals are fixed and the renaming is a note on that change
    assert [c.new for c in mm] == ["Materials and Methods"] and mm[0].reason == "sections_abstract"
    assert "Methods" in str(mm[0].provenance)
    assert any(c.old == "ABSTRACT" and c.new == "Abstract" for c in cs.changes)


def test_keywords_on_an_abstract_line_and_mesh_terms(tmp_path):
    path = make(tmp_path, [("Abstract",), "Background: " + BODY, "Results: " + BODY,
                           "Conclusion: " + BODY + "\nKeywords: Emotion Regulation, Craving, Psychotherapy",
                           ("Abbreviations",), "ERT: Emotion regulation training", "Introduction", BODY] + REFS)
    out = run(path, only=["sections_keywords"])
    assert out[3] == "Conclusion: " + BODY and out[4].startswith("Keywords: Craving, emotion regulation")
    assert "Keywords: " not in out[5:]  # no second, empty line
    qs = queries.collect(docmodel.load(path), jehp())
    assert not any(q.check.startswith("keywords") for q in qs.queries)

    path = make(tmp_path, [("Abstract",), "Background: " + BODY, "Results: " + BODY,
                           "Keywords: Diabetes Mellitus, Type 2; Behavior, Addictive; Adult, Partnership Care",
                           "Introduction", BODY] + REFS, "k.docx")
    assert rules_sections.listed_keywords(docmodel.load(path)) == [
        "Diabetes Mellitus, Type 2", "Behavior, Addictive", "Adult, Partnership Care"]
    qs = queries.collect(docmodel.load(path), jehp())
    assert not any(q.check == "keywords.too_many" for q in qs.queries)
    out = run(path, only=["sections_keywords"])
    assert out[3].startswith("Keywords: ") and "; " in out[3]  # terms holding a comma keep the semicolons


def test_funding_beside_the_journal_heading_is_not_doubled(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY, ("Funding:",), "This study received no external "
                           "funding.", ("Finance support and sponsorship:",), "Nil", ("Conflicts of interest",),
                           "There are no conflicts of interest."] + REFS)
    out = run(path, only=["heading_tags", "sections_back_matter"])
    assert out.count("<H2>Financial support and sponsorship") == 1 and "<H2>Funding" in out


def test_run_in_statements_split_and_bare_nil_ethics_go(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Conclusion", BODY,
                           "Ethics approval and consent to participate: Approved by the board (No. 51). All gave consent.",
                           "Trial registration: Not prospectively registered. Financial support and sponsorship: Nil. "
                           "Conflicts of interest: There are no conflicts of interest. Consent for publication: Not "
                           "applicable. Data availability: Data are available on request.",
                           "Acknowledgments: The authors thank the staff.",
                           "Consent to Participate: Not applicable. This study did not involve human participants."] + REFS)
    out = run(path, only=["heading_tags", "sections_back_matter"])
    i = out.index("<H2>Trial registration")
    assert out[i:i + 8] == ["<H2>Trial registration", "Not prospectively registered.",
                            "<H2>Financial support and sponsorship", "Nil.",
                            "<H2>Conflicts of interest", "There are no conflicts of interest.",
                            "<H2>Data availability statement", "Data are available on request."]
    assert "<H1>Acknowledgement" in out and "The authors thank the staff." in out  # G146 spelling (owner, 2026-10-26)
    assert "<H2>Consent to participate" in out  # with a reason: kept as its own section
    assert not any("Consent for publication" in t for t in out)  # a bare "Not applicable." goes
    assert out.count("<H2>Financial support and sponsorship") == 1


def test_case_report_table_cell_is_not_the_article_type(tmp_path):
    doc = Document()
    for t in ["Abstract", "Background: " + BODY, "Results: " + BODY, "Introduction", BODY]:
        doc.add_paragraph(t)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Case report"
    path = tmp_path / "t.docx"
    doc.save(path)
    assert not rules_sections._is_case_report(docmodel.load(str(path)))


def test_quoted_keywords_lose_their_quotes_before_case_and_order(tmp_path):
    # The author's quotation marks (and the space inside the last one) go; the sheet's case and order then apply:
    # S 'First word of first keyword should be title case; rest of keywords should be in lower case'.
    path = make(tmp_path, [("Abstract",), "Background: " + BODY, "Results: " + BODY, "Conclusion: " + BODY,
                           "KEYWORDS: “Rural schools”, “Sleep quality”, “Screen Time”, “Class Attendance”, “Asthma ”",
                           "Introduction", BODY] + REFS)
    out = run(path, only=["sections_keywords"])
    assert out[4] == "Keywords: Asthma, class attendance, rural schools, screen time, sleep quality"


def test_acknowledgements_heading_on_its_own_line_drops_the_s(tmp_path):
    # S: 'If heading "Acknowledgements" is present, then please delete "s"' (G146: "Acknowledgement"), also when the
    # heading rule recases an all-caps heading
    path = make(tmp_path, ["Introduction", BODY, "Discussion", BODY, "ACKNOWLEDGEMENTS",
                           "We thank the teachers of the three schools for their help.", "ACKNOWLEDGMENTS:"] + REFS)
    out = run(path, only=["heading_tags", "sections_headings", "sections_back_matter"])
    assert "<H1>Acknowledgement" in out and not any("Acknowledgements" in t or "ACKNOWLEDG" in t for t in out)
    assert not any(t.endswith("Acknowledgments") for t in out)


def test_en_dash_compound_in_a_sentence_case_heading(tmp_path):
    # S "H2 | Sentence Case": the part after an en dash joining a compound is lowered as after a hyphen (the author's
    # dash kept); a name joined so ("Kaplan–Meier") keeps its capitals
    text = BODY + " The health of the pupils and their health-related quality of life were measured twice."
    path = make(tmp_path, ["Introduction", text, "Methods", BODY, ("Health–Related Quality of Life",), BODY,
                           ("Kaplan–Meier Survival Estimates",), BODY, "Discussion", BODY] + REFS)
    out = run(path, only=["heading_tags", "sections_headings", "subheadings"])
    assert "<H2>Health–related quality of life" in out
    assert "<H2>Kaplan–Meier survival estimates" in out
