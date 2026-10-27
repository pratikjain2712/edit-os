"""JEHP reference-list gaps (synthetic entries; PubMed, CrossRef and the NLM directory are patched): the author's
particle case kept over CrossRef's, the title-only search with its strict terms, incomplete entries completed from
their first author's paper, typed labels turned into Word numbering, full and ALL-CAPS author names, report and
book titles in Title Case with a DOI, and a spelling kept inside a capitalised name."""

import pytest
from docx import Document

from edit_os import biblio, docmodel, journals, refengine, refs_other, refstyle, rules, style, tracked
from edit_os.changes import Status
from edit_os.copyedit import spelling


def rec(title, authors=(("Okoro", "A"),), journal="J Midwif Educ", year="2020", volume="12", pages="40-6", **kw):
    return biblio.Record(source=kw.pop("source", "pubmed"), title=title, authors=list(authors), journal=journal,
                         year=year, volume=volume, pages=pages, pmid=kw.pop("pmid", "1"), **kw)


@pytest.fixture
def db(monkeypatch):
    """PubMed returns `records` for every search; CrossRef finds nothing; no NLM directory."""
    records = []
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: [r.pmid for r in records])
    monkeypatch.setattr(biblio, "pubmed_best_match", lambda term, retmax=10: [])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [r for r in records if r.pmid in ids])
    monkeypatch.setattr(biblio, "crossref_search", lambda text, rows=5: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    monkeypatch.setattr(journals, "directory", lambda: None)
    return records


def jehp():
    st = style.load("jehp")
    st.settings["references"]["ai_parse"] = False
    return st


# --- author names ------------------------------------------------------------------------------------------

def test_author_particle_case_kept_over_crossref_not_over_pubmed():
    text = "van den Berg G, Mudau PK. A title here. Perspect Educ 2022;40:112-28."
    p = refengine.parse(text)
    cr = rec("A title here", authors=[("Van den Berg", "G"), ("Mudau", "PK")], source="crossref")
    assert refengine._usable_authors(cr, p) == ["van den Berg G", "Mudau PK"]
    pm = rec("A title here", authors=[("Van den Berg", "G"), ("Mudau", "PK")])
    assert refengine._usable_authors(pm, p) == ["Van den Berg G", "Mudau PK"]  # NLM's form


def test_other_name_forms_are_the_records():
    # another order of a compound surname, a slip inside a name, capitals: the record's
    assert not refengine.same_name_other_case("Amin Embi", "M", "Embi MA", "Embi MA, X Y.")
    assert not refengine.same_name_other_case("Alvarez-Iglesias", "A", "Alvarez-iglesias A", "Alvarez-iglesias A.")
    assert not refengine.same_name_other_case("De Leon", "C", "DE LEON C", "DE LEON C.")
    assert refengine.same_name_other_case("Van Rooyen", "DR", "van Rooyen DR", "Nyarko MJ, van Rooyen DR.")


def test_full_names_put_as_surname_and_initials_and_queried():
    text = ("Atoosa Bonyani, Leila Safaeian, Mojtaba Chehrazi. A high school-based education concerning drug abuse "
            "prevention. Journal of Education and Health Promotion. 2018;8(7):1-6")
    names, et_al, _ = refstyle.full_name_authors(text)
    assert names == ["Bonyani A", "Safaeian L", "Chehrazi M"] and not et_al
    res = refstyle.style_unverified(refengine.Resolution(refengine.QUERY, refengine.parse(text)), text, jehp())
    assert res.formatted.startswith("Bonyani A, Safaeian L, Chehrazi M. A high school-based")
    assert ("references.full_names", {}) in res.queries
    assert refstyle.full_name_authors("World Health Organization, Ministry Health. Title") is None


def test_all_caps_surname_written_as_a_name():
    edits = rules._reference_edits("ALIZADEH GJ, Imani S, Nejati V, FATHABADI J. A study of things. 2019", jehp())
    assert [(e[2], e[3]) for e in edits if e[3] == "ref_author_case"] == [("Alizadeh", "ref_author_case"),
                                                                        ("Fathabadi", "ref_author_case")]
    assert not [e for e in rules._reference_edits("WHO. Global report. Geneva: WHO; 2020.", jehp())
                if e[3] == "ref_author_case"]


# --- title-only search ---------------------------------------------------------------------------------------

TITLE = "Predictive model for assessment of delirium among patients in critical care units: A prospective study"
ENTRY = f"Sheilini M, et al. {TITLE}. Journal of Education and Health Promotion, 2025. 14(1)."


def test_wrong_first_author_found_by_title_when_listed_and_queried(db):
    db.append(rec(TITLE, authors=[("Naik", "A"), ("Sheilini", "M"), ("Chaudhuri", "S")],
                  journal="J Educ Health Promot", year="2025", volume="14", pages="243", pmid="5"))
    res = refengine.resolve(ENTRY, jehp())
    assert res.status == refengine.FORMAT_FIXED and res.record.pmid == "5"
    assert res.formatted.startswith("Naik A, Sheilini M, Chaudhuri S. Predictive model")
    assert ("references.record_differs", {"what": "first author"}) in res.queries


def test_title_only_needs_the_author_in_the_list_or_a_long_title_with_the_citation(db):
    db.append(rec(TITLE, authors=[("Naik", "A"), ("Rao", "S")], journal="J Educ Health Promot", year="2025",
                  volume="14", pages="243", pmid="5"))
    res = refengine.resolve(ENTRY, jehp())  # the entry has no first page: nothing ties it to the record
    assert res.record is None and res.status == refengine.QUERY
    assert refengine.DIFFERS in res.message and "first author" in res.message  # told, not changed
    entry = f"Sheilini M, Bose K. {TITLE}. J Educ Health Promot 2025;14:243."
    res = refengine.resolve(entry, jehp())  # long title, its volume and page: the paper
    assert res.status == refengine.FORMAT_FIXED and res.record.pmid == "5"


def test_two_papers_with_the_title_are_not_chosen_between(db):
    db.append(rec(TITLE, authors=[("Naik", "A"), ("Sheilini", "M")], journal="J Educ Health Promot", year="2025",
                  volume="14", pages="243", pmid="5"))
    db.append(rec(TITLE, authors=[("Naik", "A"), ("Sheilini", "M")], journal="Nurs Open", year="2025",
                  volume="12", pages="e70001", pmid="6"))
    res = refengine.resolve(ENTRY, jehp())
    assert res.record is None


def test_title_search_off(db):
    db.append(rec(TITLE, authors=[("Naik", "A"), ("Sheilini", "M")], journal="J Educ Health Promot", year="2025",
                  volume="14", pages="243", pmid="5"))
    st = jehp()
    st.settings["references"]["title_search"] = False
    assert refengine.resolve(ENTRY, st).record is None


def test_incomplete_entry_completed_from_its_first_authors_paper(db):
    text = ("Zali A, Mousavi P, Izadi F. The effect of cognitive behavioural therapy on loneliness among elderly "
            "women in nursing homes. 2023")
    assert refengine.incomplete_article(text, refengine.parse(text))
    db.append(rec("The effect of cognitive behavioural therapy on loneliness among elderly women in nursing homes",
                  authors=[("Zali", "A"), ("Mousavi", "P"), ("Izadi", "F")], journal="Hayat", year="2023",
                  volume="29", pages="155-70", pmid="9"))
    res = refengine.resolve(text, jehp())
    assert res.status == refengine.FORMAT_FIXED
    assert res.formatted.endswith("nursing homes. Hayat 2023;29:155-70.")


def test_incomplete_entry_by_another_first_author_is_left(db):
    text = ("Zali A, Mousavi P, Izadi F. The effect of cognitive behavioural therapy on loneliness among elderly "
            "women in nursing homes. 2023")
    db.append(rec("The effect of cognitive behavioural therapy on loneliness among elderly women in nursing homes",
                  authors=[("Rahimi", "B"), ("Zali", "A")], journal="Hayat", year="2023", volume="29",
                  pages="155-70", pmid="9"))
    assert refengine.resolve(text, jehp()).status == refengine.NOT_JOURNAL


# --- typed labels ---------------------------------------------------------------------------------------------

def _doc(tmp_path, refs):
    doc = Document()
    for t in ["Introduction", "Some text [1,2].", "References"] + refs:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return str(path)


@pytest.mark.parametrize("tracked_mode", [True, False])
def test_typed_labels_become_word_numbering(tmp_path, tracked_mode):
    path = _doc(tmp_path, ["1.\tLee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4.",
                           "2. Kim J. Neck pain in adults. Spine 1998;23:10-4."])
    st = jehp()
    st.settings["references"]["lookup"] = False
    cs = rules.run_all(docmodel.load(path), st, only=["references"])
    labels = [c for c in cs.changes if c.reason == "ref_label"]
    assert [(c.old, c.new, c.format) for c in labels] == [("1.\t", "", {"list_number": True}),
                                                         ("2. ", "", {"list_number": True})]
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(tmp_path / "o.docx")
    tracked.render(docmodel.load(path), cs, tracked=tracked_mode).save(out)
    ms = docmodel.load(out)
    refs = [p for p in ms.paragraphs if "pain" in p.text]
    assert all(p.numbered for p in refs)
    if not tracked_mode:
        assert [p.text for p in refs][0].startswith("Lee K")
    num_ids = {p.element.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId").get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") for p in refs}
    assert len(num_ids) == 1
    numbering = ms.document.part.numbering_part.element
    assert numbering.xpath(f"./w:num[@w:numId='{num_ids.pop()}']")
    xml = refs[0].element.xml
    assert ("w:pPrChange" in xml) == tracked_mode


def test_ijo_keeps_typed_labels(tmp_path):
    path = _doc(tmp_path, ["1.Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    labels = [c for c in rules.run_all(docmodel.load(path), st, only=["references"]).changes if c.reason == "ref_label"]
    assert [(c.old, c.new, c.format) for c in labels] == [("", " ", None)]  # "1. Lee K" (label_period)


# --- books and reports ----------------------------------------------------------------------------------------

def test_report_and_book_titles_in_title_case_with_a_doi():
    who = ("World Health Organization. Global oral health status report: towards universal oral health coverage by "
           "2030. Geneva: WHO; 2022. DOI: https://www.who.int/publications/i/item/9789240061484")
    out = refs_other.template(who, "report").text
    assert "Global Oral Health Status Report: Towards Universal Oral Health Coverage by 2030. Geneva" in out
    dsm = ("American Psychiatric Association. Diagnostic and statistical manual of mental disorders. 5th ed. Text rev. "
           "Washington, DC: American Psychiatric Association Publishing; 2022. doi:10.1176/appi.books.9780890425787.")
    assert "Diagnostic and Statistical Manual of Mental Disorders. 5th ed." in refs_other.template(dsm, "book").text
    mmwr = ("Workowski KA. Sexually transmitted infections treatment guidelines, 2021. Geneva: CDC; 2021. "
            "MMWR Recomm Rep 2021;70:1-187.")
    assert "Sexually transmitted infections treatment guidelines" in refs_other.book_title_case(mmwr)


# --- spelling inside a name ----------------------------------------------------------------------------------

def test_spelling_kept_inside_a_capitalised_compound_name(tmp_path):
    doc = Document()
    doc.add_paragraph("Risk of bias was assessed with the Risk Of Bias In Non-randomised Studies of Interventions "
                      "(ROBINS-I) tool, and the non-randomised trials were analysed separately.")
    path = tmp_path / "s.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    v = spelling.Variants(to_american={"randomised": "randomized", "analysed": "analyzed"},
                          american={"randomized", "analyzed"}, british={"randomised", "analysed"})
    st = style.load("jehp")
    out = [(c.start, c.new) for c in spelling.check(ms, st, v)]
    text = ms.paragraphs[0].text
    assert (text.index("non-randomised trials") + 4, "randomized") in out
    assert all(s != text.index("Non-randomised") + 4 for s, _ in out)
