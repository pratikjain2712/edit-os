"""JEHP review fixes (synthetic text): scale values are not citations, a year before a citation, table captions
typed "Table 3. summarizes ...", proper nouns in key words and references, bold headings, operator spacing,
figure / table citations inside brackets with other content, queries for missing captions, equations as images,
unexpanded abbreviations and table footnotes."""

from edit_os import queries, rules

from test_jehp_citations import _refs, make, run

REFS10 = _refs(*[f"Author{i} A, Other B. Study number {i}. J Test 2020;{i}:1-2." for i in range(1, 11)])
CITE = ["citations"]


def test_likert_values_are_not_citations(tmp_path):
    body = ["Introduction",
            "Sleep matters (1). Exercise fell (2). It is low by 2100 (3).",
            "Responses are recorded on a five-point Likert scale from strongly agree (5) to strongly disagree (1).",
            "Answers were never (0), sometimes (1), often (2) and always (3).",
            "Each item had a score (2) and three (3) options; the weights were poor (1), fair (2), good (3) on the "
            "rating scale. Women were studied before (4)."]
    _, _, t = run(make(tmp_path, body + REFS10), only=CITE)
    assert t[1] == "Sleep matters.[1] Exercise fell.[2] It is low by 2100.[3]"
    assert t[2] == body[2]
    assert t[3] == body[3]
    assert t[4].startswith("Each item had a score (2) and three (3) options; the weights were poor (1), fair (2), good (3)")
    assert t[4].endswith("studied before.[4]")


def test_year_then_citation_counts_for_the_uncited_query(tmp_path):
    from edit_os import docmodel

    body = ["Introduction", "Activity fell (1) and will fall to 1.9 by 2100 (2). Iran too (3)."]
    path = make(tmp_path, body + _refs("A A. One. J T 2020;1:1-2.", "B B. Two. J T 2020;2:1-2.",
                                       "C C. Three. J T 2020;3:1-2."))
    st = jehp_style()
    qs = queries.collect(docmodel.load(path), st)
    assert not any(q.check == "references.uncited" for q in qs.queries)


def jehp_style():
    from test_jehp_citations import jehp
    return jehp()


# --- 2. "Table 3. summarizes ..." above its table -------------------------------------------------------------------

def _table_doc(tmp_path, caption, rows, foot=None, before=None, name="t.docx"):
    from docx import Document

    doc = Document()
    for t in before or ["Introduction", "Results are in Table 1 [Table 1]."]:
        doc.add_paragraph(t)
    doc.add_paragraph(caption)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.rows[r].cells[c].paragraphs[0].add_run(val)
    if foot:
        doc.add_paragraph(foot)
    for x in _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(x)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def test_lower_case_table_caption_beside_its_table(tmp_path):
    from edit_os import docmodel, rules_figures

    path = _table_doc(tmp_path, "Table 1. summarizes the mean scores in the groups",
                      [["Group", "Mean", "SD"], ["A", "1.0", "0.2"]])
    ms = docmodel.load(path)
    cap = next(p for p in ms.paragraphs if p.text.startswith("Table 1."))
    assert rules_figures.caption_match(cap) is not None
    _, _, t = run(path, only=["fig_captions"])
    assert "Table 1: Mean scores in the groups." in t
    # away from a table, "Table 2. was ..." is a broken sentence, not a caption
    body = next(p for p in ms.paragraphs if p.text.startswith("Results"))
    assert rules_figures.caption_match(docmodel.Paragraph("x", body.element, "", "Table 2. was not found here.")) is None


# --- 3. proper nouns in key words and references ----------------------------------------------------------------------

def test_proper_nouns_keep_their_capitals():
    from edit_os import refengine, rules_sections

    body = "Sixty women in Tabriz, Iran, took part. They were followed for a month in the city of Tabriz."
    assert rules_sections._keyword_case("Iran", body) == "Iran"
    assert rules_sections._keyword_case("Family Therapy", body) == "family therapy"
    title = "Waiting Times in Clinics: A Qualitative Study in Yazd City"
    assert refengine.sentence_case(title, refengine.dictionary_proper_nouns(title)) == \
        "Waiting times in clinics: A qualitative study in Yazd City"
    title = "Effectiveness of Bowen’s Family Program and Women’s Relationship Skills"
    assert refengine.sentence_case(title, refengine.dictionary_proper_nouns(title)) == \
        "Effectiveness of Bowen’s family program and women’s relationship skills"
    assert refengine.capital_surname("kumari") == "Kumari"
    assert refengine.capital_surname("van den Berg") == "van den Berg"


def test_reference_cleanup_names_and_covid():
    edits = rules._reference_edits("Patel B, kumari M. Intention to exercise in covid 19 pandemic. "
                                   "J Test Med Sci 2023;31:1-9.", jehp_style())
    news = {e[2] for e in edits}
    assert "Kumari" in news and "COVID-19" in news


def test_theory_name_in_a_reference_title_is_sentence_case():
    from edit_os import biblio, refengine

    st = jehp_style()
    r = biblio.Record(source="crossref", title="Predicting exercise intentions with the Theory of Planned "
                      "Behaviour: A meta-analysis", authors=[("Liang", "Y"), ("kumari", "M")], journal="Adv Psychol Sci",
                      year="2025", volume="33", pages="1926")
    out = refengine.format_record(r, st)
    assert "with the theory of planned behaviour: A meta-analysis" in out
    assert out.startswith("Liang Y, Kumari M.")


# --- 4. bold headings -------------------------------------------------------------------------------------------------

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail, in vitro and in vivo.")


def test_headings_the_tool_writes_are_bold(tmp_path):
    from docx import Document
    from edit_os import docmodel

    doc = Document()
    doc.add_paragraph().add_run("Abstract").bold = True
    for label in ("Background", "Materials and Methods", "Results", "Conclusion"):
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(": " + BODY)
    doc.add_paragraph("Keywords: Health, nursing, students")
    for t in ["Introduction", BODY, "Materials and Methods",
              "Demographic Questionnaire: " + BODY, "Statistical  Analysis:", BODY, "Results", BODY, "Discussion", BODY]:
        doc.add_paragraph(t)
    for x in _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(x)
    path = str(tmp_path / "b.docx")
    doc.save(path)
    cs = rules.run_all(docmodel.load(path), jehp_style())
    methods = [c for c in cs.changes if c.old == "Materials and Methods" and c.new == "Methods"]
    assert methods and methods[0].format == {"bold": True}
    ms = docmodel.load(path)
    demo = next(p for p in ms.paragraphs if p.text.startswith("Demographic"))
    bold = [c for c in cs.changes if c.para_id == demo.id and c.reason == "fmt_heading_bold"]
    assert bold and bold[0].start == 0 and bold[0].end == len("Demographic Questionnaire")
    stat = next(p for p in ms.paragraphs if p.text.startswith("Statistical"))
    assert any(c.para_id == stat.id and c.new == "<H2>" for c in cs.changes)
    assert any(c.para_id == stat.id and c.reason == "fmt_heading_bold" for c in cs.changes)


# --- 5. replacement text in a math font ------------------------------------------------------------------------------

def test_plain_characters_in_a_math_font_take_the_journal_font():
    from edit_os import rules_format

    assert not rules_format._keeps_font("Cambria Math", "‐")
    assert not rules_format._keeps_font("Cambria Math", "a")
    assert rules_format._keeps_font("Cambria Math", "∑")
    assert rules_format._keeps_font("Symbol", "a")  # a Greek alpha in the Symbol font
    assert not rules_format._keeps_font("Times New Roman", "a")


# --- 6. figure / table citations in brackets with other content --------------------------------------------------------

def test_table_citation_in_a_bracket_with_other_content(tmp_path):
    from docx import Document

    doc = Document()
    for t in ["Results", "The difference was not significant (P = 0.21; Table 1). The rest was similar (Table 1).",
              "The groups differed by 2.5 points.", " (Table 1).", "Table 1: Scores", "Discussion", BODY]:
        doc.add_paragraph(t)
    for x in _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(x)
    path = str(tmp_path / "f.docx")
    doc.save(path)
    _, _, t = run(path, only=["fig_citations"])
    # the JEHP editors leave a bracket with other content as typed (1874); split on request
    assert t[1] == "The difference was not significant (P = 0.21; Table 1). The rest was similar [Table 1]."
    st = jehp_style()
    st.settings["figures"]["split_mixed_brackets"] = True
    _, _, t2 = run(path, st=st, only=["fig_citations"])
    assert t2[1] == "The difference was not significant (P = 0.21) [Table 1]. The rest was similar [Table 1]."
    # the stray line keeps its place (the editors, 1646) with square brackets, and is asked about
    assert t[2] == "The groups differed by 2.5 points." and t[3].strip() == "[Table 1]."
    from edit_os import docmodel
    qs = queries.collect(docmodel.load(path), jehp_style())
    assert [q.check for q in qs.queries if q.check == "figures.stray_citation"] == ["figures.stray_citation"]
    st = jehp_style()
    st.settings["figures"]["stray_citations"] = "join"
    _, _, t = run(path, st=st, only=["fig_citations"])
    assert t[2] == "The groups differed by 2.5 points [Table 1]." and "[Table 1]." not in t


# --- 7. headings: hyphen prefixes, duplicates ------------------------------------------------------------------------

def test_hyphen_prefix_heading_in_sentence_case_and_duplicate_heading_query(tmp_path):
    from docx import Document
    from edit_os import docmodel

    doc = Document()
    for t in ["Introduction", BODY, "Materials and Methods"]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run("Sleep Habits Questionnaire:").bold = True
    for t in ["Sleep habits Questionnaire: " + BODY, "Results"]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run("Inter-Group Post-Intervention Evaluations").bold = True
    for t in [BODY, "Discussion", BODY] + _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(t)
    path = str(tmp_path / "h.docx")
    doc.save(path)
    _, _, t = run(path, only=["heading_tags", "subheadings"])
    assert "<H2>Inter-group post-intervention evaluations" in t
    qs = queries.collect(docmodel.load(path), jehp_style())
    dup = [q for q in qs.queries if q.check == "headings.duplicate"]
    assert len(dup) == 1 and "Sleep Habits Questionnaire" in dup[0].text


# --- 8. pictures: no caption, an equation -----------------------------------------------------------------------------

def test_picture_queries(tmp_path):
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    from edit_os import docmodel

    img = tmp_path / "x.png"
    Image.new("RGB", (40, 40), "white").save(img)
    doc = Document()
    for t in ["Introduction", BODY, "Materials and Methods", BODY]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run().add_picture(str(img), width=Inches(2.4), height=Inches(0.75))  # an equation
    for t in ["Results", BODY]:
        doc.add_paragraph(t)
    doc.add_paragraph().add_run().add_picture(str(img), width=Inches(5), height=Inches(6))  # a figure, no caption
    for t in ["Discussion", BODY] + _refs("A A. One. J T 2020;1:1-2."):
        doc.add_paragraph(t)
    path = str(tmp_path / "i.docx")
    doc.save(path)
    ms = docmodel.load(path)
    qs = queries.collect(ms, jehp_style())
    checks = [q.check for q in qs.queries]
    assert "figures.equation_image" in checks and "figures.image_no_caption" in checks
    assert queries.write_comments(ms, qs) == len(qs.queries)  # anchored on the pictures


# --- 9. operator spacing in cells, stray brackets ---------------------------------------------------------------------

def test_operator_spacing_in_cells_and_stray_brackets(tmp_path):
    path = _table_doc(tmp_path, "Table 1: Scores", [["Group", "Mean ± SD", "95% CI"],
                                                    ["A", "5.20± 1.30", "−1.20,3.40))"], ["B", "2.17±1.97", "1, 2"]],
                      foot="SD = standard deviation, CI = confidence interval")
    _, _, t = run(path, only=["guide_replace"])
    assert "5.20 ± 1.30" in t and "2.17 ± 1.97" in t
    assert queries.unmatched_brackets("−1.20,3.40))") == (10, 11)
    assert queries.unmatched_brackets("(a) one, b) two") is None
    assert queries.unmatched_brackets("mean (SD) was 2 (1.2)") is None


# --- 10. references: a book, the label's spaces ------------------------------------------------------------------------

def test_book_with_dotted_author_names():
    from edit_os import refstyle

    out = refstyle.restyle_other("Smith, J. Rao, K. lee, M. Family counselling skills, London: "
                                 "Example Press; 2015.", jehp_style(), "")
    assert out == "Smith J, Rao K, Lee M. Family counselling skills. London: Example Press; 2015."


def test_no_break_space_before_a_reference_label_goes(tmp_path):
    body = ["Introduction", "Marriage matters (1)."]
    path = make(tmp_path, body + ["References", " 1.\tLee M, Park J. The decision. J Test 2013;12:505-15."])
    _, _, t = run(path, only=["references"])
    assert t[-1].startswith("Lee M")


# --- 11. abbreviations the common list lets pass -----------------------------------------------------------------------

def test_cvi_and_cvr_are_queried(tmp_path):
    from edit_os import docmodel
    from edit_os.copyedit import abbrev

    path = make(tmp_path, ["Introduction", "Content validity was examined by seven experts and the resulting CVI and "
                           "CVR were both 1.0 for the questionnaire, with a left LV mass in Class III patients."]
                + _refs("A A. One. J T 2020;1:1-2."))
    _, qs = abbrev.collect(docmodel.load(path), jehp_style())
    asked = {q.text.split(":")[0] for q in qs if q.check == "abbrev.undefined"}
    assert {"CVI", "CVR", "LV"} <= asked and "III" not in asked


# --- 12. table footnotes ------------------------------------------------------------------------------------------------

def test_table_footnote_queries(tmp_path):
    from edit_os import docmodel

    path = _table_doc(tmp_path, "Table 1: Scores", [["Group", "Mean", "SD", "95% CI"],
                                                    ["A", "1.0", "**0.2", "***1, 2"], ["B", "1.1", "0.3", "1, 3"]],
                      foot="*Chi-square test, **Fisher’s exact test, **Mann–Whitney U test")
    qs = queries.collect(docmodel.load(path), jehp_style())
    by = {q.check: q.text for q in qs.queries}
    assert by["figures.footnote_abbreviations"] == "Kindly provide the expansion of SD and CI in the table footnote."
    assert "“**”" in by["figures.footnote_symbol_twice"]
    assert "“***”" in by["figures.footnote_symbol_missing"]


def test_caption_keeps_country_names(tmp_path):
    path = _table_doc(tmp_path, "Table 1. Participant Characteristics, United Arab Emirates",
                      [["Group", "Mean", "SD"], ["A", "1.0", "0.2"]])
    _, _, t = run(path, only=["fig_captions"])
    assert "Table 1: Participant characteristics, United Arab Emirates." in t
