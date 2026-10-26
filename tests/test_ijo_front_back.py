"""IJO front and back matter: bold section headings, the consent section's place, the abstract and key words,
correspondence and affiliations, equal contribution, title-page statements after the abstract.

All text is synthetic."""

import zipfile

from docx import Document
from lxml import etree

from edit_os import docmodel, frontmatter, queries, rules, rules_sections, style, submission, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

BR = tracked.PARAGRAPH_BREAK
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic.")
REFS = ["References", "1. Smith A, Jones B. A study. Ophthalmology 2001;12:1-9.",
        "2. Smith A, Jones B. Another study. Ophthalmology 2002;13:1-9."]
CONSENT = ("The authors certify that they have obtained all appropriate patient consent forms for the publication "
           "of the clinical images in the journal.")


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        if isinstance(text, tuple):
            doc.add_paragraph().add_run(text[0]).bold = True
        else:
            doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def load(path, first_page=(), meta=None):
    ms = docmodel.load(path)
    ms.front_sources = frontmatter.Sources(first_page=list(first_page), meta=meta or frontmatter.Meta())
    return ms


def run(path, only, code="ijo", **src):
    cs = rules.run_all(load(path, **src), style.load(code), only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = path.replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs if p.text], out, cs


def bold_texts(path):
    """Paragraph texts whose every run is bold."""
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    out = []
    for p in root.iter(W + "p"):
        runs = [r for r in p.iter(W + "r") if "".join(t.text or "" for t in r.iter(W + "t")).strip()]
        if runs and all(r.find(f"{W}rPr/{W}b") is not None for r in runs):
            out.append("".join(t.text or "" for t in p.iter(W + "t")))
    return out


BACK = ("heading_tags", "sections_back_matter")


# --- inserted paragraphs: formatting per segment ------------------------------------------------

def test_segment_format_bolds_the_heading_only(tmp_path):
    path = make(tmp_path, ["First paragraph."])
    ms = docmodel.load(path)
    p = ms.paragraphs[0]
    c = Change(id="c1", para_id=p.id, start=len(p.text), end=len(p.text), old="", new=BR + "<H2>Heading" + BR + "Nil.",
               layer=Layer.STRUCTURE, source=Source.RULE, status=Status.ACCEPTED, format={"bold": False},
               segment_format=[None, {"bold": True}, None])
    out = str(tmp_path / "o.docx")
    tracked.render(ms, ChangeSet([c]), tracked=True).save(out)
    assert bold_texts(out) == ["<H2>Heading"]
    assert Change.from_dict(c.to_dict()).segment_format == [None, {"bold": True}, None]


def test_segment_pieces_bold_a_label_inside_a_new_paragraph(tmp_path):
    path = make(tmp_path, ["First paragraph."])
    ms = docmodel.load(path)
    p = ms.paragraphs[0]
    c = Change(id="c1", para_id=p.id, start=len(p.text), end=len(p.text), old="", new=BR + "Keywords: a, b",
               layer=Layer.STRUCTURE, source=Source.RULE, status=Status.ACCEPTED, format={"bold": False},
               segment_format=[None, [[9, {"bold": True}]]])
    out = str(tmp_path / "o.docx")
    tracked.render(ms, ChangeSet([c]), tracked=False).save(out)
    root = etree.fromstring(zipfile.ZipFile(out).read("word/document.xml"))
    runs = [("".join(t.text for t in r.iter(W + "t")), r.find(f"{W}rPr/{W}b") is not None)
            for r in root.iter(W + "r") if "".join(t.text or "" for t in r.iter(W + "t"))]
    assert ("Keywords:", True) in runs and (" a, b", False) in runs


# --- back matter ------------------------------------------------------------------------------------

def test_added_sections_have_bold_headings(tmp_path):
    path = make(tmp_path, ["Discussion", BODY] + REFS)
    t, out, _ = run(path, BACK)
    assert "<H2>Financial support and sponsorship" in t and "Nil." in t
    bold = bold_texts(out)
    assert "<H2>Financial support and sponsorship" in bold and "<H2>Conflicts of interest" in bold
    assert "Nil." not in bold


def test_consent_after_references_moves_before_financial_support(tmp_path):
    path = make(tmp_path, ["Discussion", BODY] + REFS + ["Declaration of Patient Consent", CONSENT])
    t, _, _ = run(path, BACK)
    i = t.index("<H2>Declaration of patient consent")
    assert t[i + 1] == CONSENT and t[i + 2] == "<H2>Financial support and sponsorship"
    assert t.index("<H1>References") > i and t.count(CONSENT) == 1


def test_consent_after_the_funding_statements_moves_before_them(tmp_path):
    path = make(tmp_path, ["Discussion", BODY, "Conflicts of interest", "Nil", "Sources of funding", "Nil",
                           "Patient Consent", CONSENT] + REFS)
    t, _, _ = run(path, BACK)
    assert t.index(CONSENT) < t.index("<H2>Financial support and sponsorship") < t.index("<H2>Conflicts of interest")
    assert t[t.index(CONSENT) - 1] == "<H2>Patient consent"  # the author's heading, sentence case


def test_bare_consent_heading_renamed(tmp_path):
    statement = "Written informed consent was obtained from all patients for the publication of their images."
    path = make(tmp_path, ["Conclusion", BODY, "Consent", statement, "Conflict of interest",
                           "The authors declare that they have no conflict of interest."] + REFS)
    t, _, _ = run(path, BACK)
    assert t[t.index(statement) - 1] == "<H2>Declaration of patient consent"


def test_acknowledgement_from_the_title_page(tmp_path):
    path = make(tmp_path, ["Discussion", BODY] + REFS)
    fp = ["Title: A study", "Acknowledgements:", "We thank the participants of this study for their cooperation."]
    t, _, _ = run(path, BACK, first_page=fp)
    i = t.index("<H1>Acknowledgement")  # G146 "Acknowledgement", first level (owner, 2026-10-26)
    assert t[i + 1] == "We thank the participants of this study for their cooperation."
    assert t[i + 2] == "<H2>Financial support and sponsorship"


def test_front_statements_after_the_abstract_feed_the_back_matter(tmp_path):
    path = make(tmp_path, ["Abstract", "Aims: " + BODY, "Words- 194", "CONFLICT OF INTEREST- None",
                           "FUNDING/SUPPORT- Eye Research Foundation, Hyderabad (grant 12)",
                           "Running head- Fungal keratitis", "Introduction:", BODY, "Discussion", BODY] + REFS)
    t, _, _ = run(path, ("submission",) + BACK)
    assert not any(x.startswith(("Words-", "CONFLICT", "FUNDING", "Running head")) for x in t)
    i = t.index("<H2>Financial support and sponsorship")
    assert t[i + 1] == "Eye Research Foundation, Hyderabad (grant 12)."
    assert t[i + 3] == "There are no conflicts of interest."


# --- abstract and key words --------------------------------------------------------------------------

def test_keywords_above_an_inline_abstract(tmp_path):
    path = make(tmp_path, ["A study of snake bite", "KEY WORDS: Snake envenomation, Optic neuritis",
                           "ABSTRACT: " + BODY + " " + BODY, "Introduction", BODY] + REFS)
    t, out, _ = run(path, ("sections_abstract", "sections_keywords"))
    i = t.index("Abstract")
    assert t[i + 1] == BODY + " " + BODY
    assert t[i + 2] == "Keywords: Optic neuritis, snake envenomation"
    assert "Abstract" in bold_texts(out) and not any(x.startswith("KEY WORDS") for x in t)


def test_abstract_heading_case_and_bold_text(tmp_path):
    doc = Document()
    doc.add_paragraph("A study of eyes")
    doc.add_paragraph("ABSTRACT")
    p = doc.add_paragraph("Purpose: ")
    p.add_run("Bold words").bold = True
    p.add_run(" " + BODY)
    doc.add_paragraph("Methods: " + BODY)
    doc.add_paragraph("Keywords: Glaucoma, cataract, cornea")
    for x in ["Introduction", BODY] + REFS:
        doc.add_paragraph(x)
    path = str(tmp_path / "a.docx")
    doc.save(path)
    t, out, _ = run(path, ("sections_abstract",))
    assert "Abstract" in t and "ABSTRACT" not in t
    root = etree.fromstring(zipfile.ZipFile(out).read("word/document.xml"))
    bold_words = ["".join(x.text or "" for x in r.iter(W + "t")) for r in root.iter(W + "r")
                  if r.find(f"{W}rPr/{W}b") is not None and r.find(f"{W}rPr/{W}b").get(W + "val") not in ("0", "false")]
    assert "Bold words" not in "".join(bold_words)


def test_missing_keywords_from_the_xml_without_query(tmp_path):
    path = make(tmp_path, ["A study", "Abstract", BODY, "Introduction", BODY] + REFS)
    meta = frontmatter.Meta(keywords="Orbital cellulitis, pansinusitis,intracranial extension")
    t, out, _ = run(path, ("sections_keywords",), meta=meta)
    assert "Keywords: Intracranial extension, orbital cellulitis, pansinusitis" in t
    qs = queries.collect(load(path, meta=meta), style.load("ijo"))
    assert not any(q.check == "keywords.missing" for q in qs.queries)


def test_keywords_label_on_its_own_line_runs_on(tmp_path):
    path = make(tmp_path, ["A study", "Abstract", BODY, ("Keywords:",), "glaucoma, cataract, cornea",
                           "Introduction", BODY] + REFS)
    t, out, _ = run(path, ("sections_keywords",))
    assert "Keywords: Cataract, cornea, glaucoma" in t and "glaucoma, cataract, cornea" not in t


def test_hyphenated_eponym_keeps_its_capitals():
    body = "This is a case of Miller-Fisher syndrome in a young man with ophthalmoplegia and ataxia."
    assert rules_sections._keyword_case("Miller Fisher syndrome", body) == "Miller Fisher syndrome"
    assert rules_sections._keyword_case("Turner-Kieser syndrome", "known as Turner-Kieser syndrome in the past.") \
        == "Turner-Kieser syndrome"


# --- queries and headings ------------------------------------------------------------------------------

def test_conclusion_query_with_a_colon_on_discussion(tmp_path):
    path = make(tmp_path, ["A study", "Introduction", BODY, "Methods", BODY, "Results", BODY, "Discussion:", BODY] + REFS)
    qs = queries.collect(load(path), style.load("ijo"))
    assert any(q.check == "sections.conclusion" for q in qs.queries)


def test_case_report_sections_are_first_level(tmp_path):
    path = make(tmp_path, ["A study", "Case Description", BODY + BODY, "Treatment and Outcome", BODY + BODY,
                           "Discussion", BODY] + REFS)
    t, _, _ = run(path, ("heading_tags",))
    assert "<H1>Case Description" in t and "<H1>Treatment and Outcome" in t


# --- front matter ------------------------------------------------------------------------------------------

def test_indian_state_abbreviations():
    assert frontmatter._state({"state": "U.p", "city": "Barabanki"}) == "Uttar Pradesh"
    assert frontmatter._country({"state": "U.P.", "city": "Barabanki"}) == "India"
    assert frontmatter._state({"state": "WI", "city": "Madison"}) == "WI"  # undotted codes are not read in the text


def test_xml_state_code_completes_an_affiliation():
    f = frontmatter.Fields(authors=[frontmatter.Author("Swasti", "Pandey", [1], True)],
                           affiliations={1: {"department": "Cornea", "institution": "Sankara Eye Hospital",
                                             "city": "Shimoga", "state": "", "country": ""}})
    meta = frontmatter.Meta(authors=[{"surname": "Pandey", "aff": {"city": "Shimoga", "state": "KA", "country": ""}}])
    frontmatter._xml_places(f, meta)
    assert f.affiliations[1]["state"] == "Karnataka" and f.affiliations[1]["country"] == "India"


def test_unnumbered_affiliations_linked_by_marks():
    f = frontmatter.Fields(authors=[frontmatter.Author("Varnika", "Sahi"), frontmatter.Author("Rajesh", "Joshi")],
                           affiliations={1: {"department": "Glaucoma Services"}, 2: {"department": "Cataract Services"}})
    src = frontmatter.Sources(first_page=["1) Dr. Varnika Sahi,1 MBBS", "2) Dr. Rajesh Joshi,2 DOMS",
                                          "Institutional Affiliations:", "Glaucoma Services", "Cataract Services"])
    frontmatter._link_affiliations(f, src, [])
    assert [a.affs for a in f.authors] == [[1], [2]]


def test_equal_contribution_mark_and_note():
    f = frontmatter.Fields(authors=[frontmatter.Author("Yijie", "Lu"), frontmatter.Author("Yao", "Tan"),
                                    frontmatter.Author("Yu", "Liu")])
    src = frontmatter.Sources(first_page=["Yijie Lu1,2†, Yao Tan1,3†, Yu Liu1*",
                                          "†These authors contributed equally to this work."])
    frontmatter._equal_contribution(f, src, [])
    assert f.equal_mark == "†" and f.equal_authors == [0, 1]
    assert f.equal_note == "These authors contributed equally to this work."


def test_contribution_statement_names_are_not_authors():
    f = frontmatter.Fields(authors=[frontmatter.Author("Soumya", "Grover"), frontmatter.Author("Radhika", "Paranjpe")])
    src = frontmatter.Sources(first_page=["Sincerely, Dr. Soumya Grover", "Author Contributions",
                                          "Concept and design: Dr. Soumya Grover",
                                          "Definition of intellectual content: Dr. Radhika Paranjpe"],
                              meta=frontmatter.Meta(authors=[{"given": "Soumya", "surname": "Grover"},
                                                             {"given": "Naina", "surname": "Chaturvedi"}]))
    assert frontmatter._names_from_contributions(f, src, [])


def test_named_unit_gets_no_department_prefix():
    assert frontmatter.department_form("Postdoctoral Station of Clinical Medicine") == \
        "Postdoctoral Station of Clinical Medicine"
    assert frontmatter.department_form("Ophthalmology") == "Department of Ophthalmology"


def test_article_file_label_is_a_title_page_label():
    assert frontmatter._LABEL.match("ARTICLE FILE")
    assert submission._FRONT_STATEMENT.match("FUNDING/SUPPORT- Hyderabad Eye Research Foundation")
