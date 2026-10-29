"""tools/defect_scan.py: the output defect checks, on synthetic documents (no client material)."""

import collections
import os
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))

import defect_scan as ds  # noqa: E402
from edit_os import docmodel, pipeline, style  # noqa: E402
from edit_os.changes import Change, Layer, Source, Status  # noqa: E402
from edit_os.copyedit import langtool  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)


def _doc(tmp_path, paras, name="clean.docx", bold=()):
    """A .docx with one paragraph per text; runs of the texts in `bold` are bold."""
    d = Document()
    for t in paras:
        p = d.add_paragraph()
        r = p.add_run(t)
        r.bold = t in bold
        r.font.name = "Times New Roman"
        r.font.size = 12 * 12700
    path = str(tmp_path / name)
    d.save(path)
    return docmodel.load(path)


def _article(tmp_path, clean_texts, source_texts=None, journal="ijo", bold=(), changes=(), comments=None,
             tracked=None, query_checks=None):
    clean = _doc(tmp_path, clean_texts, bold=bold)
    source = _doc(tmp_path, source_texts if source_texts is not None else clean_texts, "source.docx")
    with open(clean.path, "rb") as f:
        import zipfile
        z = zipfile.ZipFile(f)
        xml, parts = z.read("word/document.xml"), z.namelist()
    return ds.Article("A1", journal, style.load(journal), clean, source, tracked or xml, comments, xml, parts,
                      list(changes), [], {"qc": {"issues": []}, "render_conflicts": []},
                      collections.Counter(query_checks or {}))


def _checks(found):
    return collections.Counter(f.check for f in found)


# --- masking -----------------------------------------------------------------------------------------------------

def test_snippet_masks_names_numbers_and_keeps_tags():
    s = ds.snippet("<RH>Neeraja, et al.: Department of Orbit 2019 jane@x.org")
    assert s.startswith("<RH> Aa, et al.:")
    assert "Neeraja" not in s and "Orbit" not in s and "2019" not in s and "@" not in s
    assert ds.snippet("<H2>Financial support") == "<H2> Financial support"
    assert len(ds.snippet("one two three four five six seven eight nine ten").split()) == 8


# --- headings and fonts ----------------------------------------------------------------------------------------------

def test_tagged_heading_not_bold_is_an_error(tmp_path):
    a = _article(tmp_path, ["<H1>Introduction", "Some text here.", "<H1>Methods"], bold=("<H1>Methods",))
    out = []
    ds.check_headings(a, out)
    assert [(f.check, f.para) for f in out if f.check == "format.heading_not_bold"] == [("format.heading_not_bold", 0)]


def test_h2_at_heading_size_and_body_at_wrong_size(tmp_path):
    a = _article(tmp_path, ["Abstract", "<H1>Introduction", "<H2>Funding", "Body text."],
                 bold=("<H1>Introduction", "<H2>Funding"))
    ps = a.clean.paragraphs
    for r in ps[1].element.iter(qn("w:r")):  # the H1 at 14 pt: right
        r.find(qn("w:rPr")).find(qn("w:sz")).set(qn("w:val"), "28")
    for r in ps[2].element.iter(qn("w:r")):  # the H2 at 14 pt: wrong
        r.find(qn("w:rPr")).find(qn("w:sz")).set(qn("w:val"), "28")
    for r in ps[3].element.iter(qn("w:r")):  # body text at 10.5 pt: wrong
        r.find(qn("w:rPr")).find(qn("w:sz")).set(qn("w:val"), "21")
    out = []
    ds.check_headings(a, out)
    ds.check_fonts(a, out)
    assert [(f.check, f.para) for f in out] == [("format.heading_size", 2), ("format.size", 3)]


# --- queries ---------------------------------------------------------------------------------------------------------

def _tracked(body: str) -> bytes:
    return f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'.encode()


def _comments(texts: dict) -> bytes:
    cs = "".join(f'<w:comment w:id="{k}"><w:p><w:r><w:t>{v}</w:t></w:r></w:p></w:comment>' for k, v in texts.items())
    return f'<w:comments xmlns:w="{W}">{cs}</w:comments>'.encode()


def test_queries_on_deleted_text_duplicates_and_broken_texts(tmp_path):
    body = (
        '<w:p><w:commentRangeStart w:id="1"/><w:del><w:r><w:delText>gone words</w:delText></w:r></w:del>'
        '<w:commentRangeEnd w:id="1"/></w:p>'
        '<w:p><w:commentRangeStart w:id="2"/><w:commentRangeStart w:id="3"/><w:r><w:t>Title</w:t></w:r>'
        '<w:commentRangeEnd w:id="2"/><w:commentRangeEnd w:id="3"/></w:p>'
        '<w:p><w:commentRangeStart w:id="4"/><w:r><w:t>Text</w:t></w:r><w:commentRangeEnd w:id="4"/></w:p>'
        '<w:p><w:commentRangeStart w:id="5"/><w:commentRangeEnd w:id="5"/></w:p>'
    )
    texts = {"1": "AQ1: Kindly check.", "2": "AQ2: Kindly provide full name.", "3": "AQ3: Kindly provide full name.",
             "4": "AQ4: Kindly provide the None value ???", "5": "AQ5: Kindly check this."}
    a = _article(tmp_path, ["x"], tracked=_tracked(body), comments=_comments(texts))
    out = []
    ds.check_queries(a, out)
    assert sorted((f.check, f.para) for f in out) == [
        ("query.bad_text", 2), ("query.duplicate", 1), ("query.on_deleted", 0), ("query.on_empty_para", 3)]


# --- citations -------------------------------------------------------------------------------------------------------

def _ch(pid, start, end, old, new, reason="citations", status=Status.ACCEPTED, fmt=None):
    return Change(f"c{start}", pid, start, end, old, new, Layer.HOUSE_STYLE, Source.RULE, reason, status, format=fmt)


def test_citation_after_a_scale_value_and_in_a_table_count(tmp_path):
    src = ["Responses: strongly agree (5), agree (4) were common in the group.",
           "The patients were treated as described (2)."]
    a = _article(tmp_path, ["x"], source_texts=src)
    t0 = src[0]
    i5 = t0.index("(5)")
    t1 = src[1]
    i2 = t1.index(" (2)")
    a.changes = [
        _ch("p0001", i5, i5 + 3, "(5)", "[5]"),  # a Likert value turned into a citation
        _ch("p0002", i2, i2 + 4, " (2)", ""),  # a real citation moved before the full stop: delete + insert
        _ch("p0002", len(t1), len(t1), "", "[2]"),
    ]
    pairs = ds.citation_changes(a.changes)
    assert [(old, at) for _, old, at in pairs] == [("(5)", i5), (" (2)", i2)]
    out = []
    ds.check_citations(a, out)
    assert [(f.check, f.para) for f in out] == [("cite.after_scale_or_unit", 0)]


def test_citation_beyond_the_list_and_footnote_marker(tmp_path):
    src = ["Text cites a thing (7).", "<H1>References", "1. Lee K. Back pain. Spine 1997;22:1380-4."]
    a = _article(tmp_path, ["x"], source_texts=src)
    i = src[0].index("(7)")
    a.changes = [_ch("p0001", i, i + 3, "(7)", "[7]")]
    out = []
    ds.check_citations(a, out)
    assert [f.check for f in out] == ["cite.beyond_refs"]


# --- content ---------------------------------------------------------------------------------------------------------

BODY = ["Original Article", "Abstract", "Background text of the abstract here.", "<H1>Introduction",
        "The first body paragraph has quite a few words in it to be counted.",
        "The second body paragraph is a different one about other matters entirely.",
        "<H1>References", "1. Lee K. Back pain. Spine 1997;22:1380-4."]


def test_body_paragraph_lost_or_deleted_and_duplicates(tmp_path):
    clean = BODY[:4] + [BODY[5], BODY[5], "<H1>Introduction"] + BODY[6:]
    a = _article(tmp_path, clean, source_texts=BODY)
    out = []
    ds.check_content(a, out)
    assert sorted(_checks(out)) == ["content.dup_heading", "content.dup_para", "content.para_lost"]
    # deleted by a rule that is not front matter / submission: an error naming the rule
    a.changes = [_ch("p0005", 0, len(BODY[4]), BODY[4], "", reason="claude.other")]
    out = []
    ds.check_content(a, out)
    assert "content.para_deleted" in _checks(out) and "content.para_lost" not in _checks(out)
    # a standard-section rewrite or the submission rules are expected
    a.changes = [_ch("p0005", 0, len(BODY[4]), BODY[4], "", reason="submission_reviewer_reply")]
    out = []
    ds.check_content(a, out)
    assert not {"content.para_deleted", "content.para_lost"} & set(_checks(out))
    # a table's second caption deleted: expected when the query quotes it, an error otherwise
    a.changes = [_ch("p0005", 0, len(BODY[4]), BODY[4], "", reason="table_caption_duplicate")]
    out = []
    ds.check_content(a, out)
    assert "content.para_deleted" in _checks(out)
    a.query_checks["figures.table_caption_duplicate"] = 1
    out = []
    ds.check_content(a, out)
    assert "content.para_deleted" not in _checks(out)


# --- clean copy, figures, tables, references -----------------------------------------------------------------------------

def test_clean_copy_leftovers(tmp_path):
    texts = ["Technical Editing done on 27-October-2026:- ???", "<RH>Lee, et al.: ???", "Text with  two spaces.",
             "A stray <H2> tag inside.", "Space before , comma.", "Figure 1: ???", "Ends [[1]]."]
    a = _article(tmp_path, texts)
    out = []
    ds.check_clean(a, out)
    got = {(f.check, f.severity) for f in out}
    assert ("signoff.placeholder", "info") in got
    assert {"clean.placeholder_rh", "clean.stray_tag", "clean.space_before_punct", "clean.placeholder_legend",
            "clean.double_bracket", "clean.double_space"} <= {c for c, _ in got}
    assert "clean.no_running_head" not in {c for c, _ in got}
    a2 = _article(tmp_path, ["No running head here."])
    a2.clean_xml += b"<w:ins w:id='1'>"
    out = []
    ds.check_clean(a2, out)
    assert {"clean.no_running_head", "clean.residue"} <= set(_checks(out))


def test_double_bracket_made_by_us_is_an_error_the_authors_is_info(tmp_path):
    clean = ["Adopted from the earlier tools [B. Lee,[4]].", "4. Lee B. A review. J Test. 2021. [Link][Index]] more."]
    source = ["Adopted from the earlier tools [B. Lee,(2021)].", "4. Lee B. A review. J Test. 2021. [Link][Index]] more."]
    out = []
    ds.check_clean(_article(tmp_path, clean, source), out)
    got = sorted((f.para, f.severity) for f in out if f.check == "clean.double_bracket")
    assert got == [(0, "error"), (1, "info")]


def test_legend_label_and_table_caption(tmp_path):
    texts = ["<RH>Lee, et al.: Title", "Figure Legends", "Figure 1: A fine legend.", "Fig. 2. A legend typed wrong.",
             "Table 1", "<H1>References", "1. Lee K. Back pain. Spine 1997;22:1380-4."]
    a = _article(tmp_path, texts, bold=("Figure 1: A fine legend.",))
    t = a.clean.document.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Group"
    tbl = t._tbl
    a.clean.paragraphs[4].element.addnext(tbl)  # a table after a bare "Table 1" label, before the references
    out = []
    ds.check_figures_tables(a, out)
    c = _checks(out)
    assert c["figtab.legend_label"] == 1 and c["figtab.table_no_caption"] == 1
    assert any("label without caption" in f.loc for f in out if f.check == "figtab.table_no_caption")
    empty = a.clean.document.add_table(rows=1, cols=1)._tbl
    a.clean.paragraphs[1].element.addprevious(empty)
    out = []
    ds.check_figures_tables(a, out)
    assert _checks(out)["figtab.empty_table"] == 1


def test_tables_in_supplementary_material_are_not_checked(tmp_path):
    # S372 / G151: the supplementary material is retained as it is ("Appendix Table 1: ..." over its table)
    texts = ["<RH>Lee, et al.: Title", "Body text.", "<H2>Supplementary Material", "Appendix Table 1: Cases",
             "<H1>References", "1. Lee K. Back pain. Spine 1997;22:1380-4."]
    a = _article(tmp_path, texts)
    t = a.clean.document.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Case"
    a.clean.paragraphs[3].element.addnext(t._tbl)
    out = []
    ds.check_figures_tables(a, out)
    assert not [f for f in out if f.check.startswith("figtab.table")]


def test_reference_numbers_shown_and_entries(tmp_path):
    texts = ["<H1>References", "1. Lee K. Back pain. Spine 1997;22:1380-4.",
             "3. van Dam K, Roe J, et al Title of it. J Med 2001;2:3-4.", "4. smith J. Another. J 2001;1:1-2."]
    a = _article(tmp_path, texts)
    out = []
    ds.check_references(a, out)
    c = _checks(out)
    assert c["refs.number_sequence"] == 1 and c["refs.et_al_period"] == 1 and c["refs.lowercase_surname"] == 1
    assert "missing [2]" in next(f.loc for f in out if f.check == "refs.number_sequence")


# --- QC report -------------------------------------------------------------------------------------------------------

def test_qc_issues_queried_by_the_author_are_info(tmp_path):
    a = _article(tmp_path, ["x"], query_checks={"references.uncited": 1})
    a.qc = {"qc": {"issues": [
        {"severity": "error", "check": "references.cited", "message": "Reference 3 is not cited", "para_id": None},
        {"severity": "error", "check": "citations.order", "message": "Reference 4 before 3", "para_id": None},
        {"severity": "error", "check": "keywords.order", "message": "Keywords not in order", "para_id": None},
        {"severity": "error", "check": "signoff.name", "message": "no name", "para_id": None}]},
        "render_conflicts": [{"output": "tracked", "para_id": "p0001", "error": "x"}]}
    out = []
    ds.check_qc(a, out)
    sev = {f.check: f.severity for f in out}
    assert sev == {"qc.references.cited": "info", "qc.citations.order": "info", "qc.keywords.order": "error",
                   "qc.signoff.name": "info", "qc.render_conflict": "error"}


# --- end to end ----------------------------------------------------------------------------------------------------------

def test_scan_of_a_pipeline_run_and_summary(tmp_path):
    texts = ["Original Article", "Introduction", "Some patients were seen (1).", "Discussion",
             "Results were as expected (2).", "References",
             "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit Amst Neth. 2002;21(4):295-299.",
             "2. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4."]
    d = Document()
    for t in texts:
        d.add_paragraph(t)
    src = str(tmp_path / "ms.docx")
    d.save(src)
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    od = str(tmp_path / "out")
    pipeline.edit(src, st, od, accept_all=True, lt_client=None)
    found = ds.scan(od, "ms", "ijo")
    assert "scan.crash" not in _checks(found)
    assert "signoff.placeholder" in _checks(found)  # no editor on the run: counted, not an error
    rows = [f.to_dict() for f in found]
    summ = ds.summarize(rows, {"ms": "ijo"})
    assert summ["ijo"]["articles"] == 1
    assert summ["ijo"]["checks"]["signoff.placeholder"]["articles"] == 1


def test_daily_learning_tracks_defects():
    import learn

    r1 = {"defects": {"errors": 2, "checks": {"format.heading_not_bold": [2, "error"], "clean.double_space": [3, "info"]},
                      "examples": {}}}
    r2 = {"defects": {"errors": 0, "checks": {"signoff.placeholder": [1, "info"]}, "examples": {}}}
    m = learn.journal_metrics([r1, r2])
    assert m["defect_articles"] == 1 and m["defects"] == {"format.heading_not_bold": 1}
    worse = learn.journal_metrics([r1, r1])
    regs = learn.regressions(worse, m, "ijo")
    assert any(x["kind"] == "defects" and "format.heading_not_bold" in x["what"] for x in regs)
    md = learn.render_markdown({"date": "2026-10-27", "commit": "abc", "journals": {"ijo": {"heldout": m}}})
    assert "## Output defects" in md and "`format.heading_not_bold` 1" in md
