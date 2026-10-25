"""IJO citation gaps: "Fig." labels, figure citations moved before the full stop, bold citations made roman,
consecutive reference numbers as a range, citation forms the rules missed, figure parts with no number, and the
author-name query. Synthetic text only."""

from docx import Document

from edit_os import docmodel, queries, rules, rules_format, style, tracked
from edit_os.changes import Status


def _style(code="ijo"):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


def _refs(n, first=None):
    first = first or {}
    return ["References"] + [f"{i}. {first.get(i, f'Author{chr(64 + i)} A')}, Other B. A title {i}. J Test 2001;{i}:1-2."
                             for i in range(1, n + 1)]


def _save(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        if isinstance(t, list):  # [(text, {"bold": True}), ...] -> one run each
            para = doc.add_paragraph()
            for text, fmt in t:
                run = para.add_run(text)
                for k, v in fmt.items():
                    setattr(run.font, k, v)
        elif isinstance(t, tuple):  # (text, paragraph style)
            doc.add_paragraph(t[0], style=t[1])
        else:
            doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def _run(path, st, tmp_path, only=None):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")
    return [p.text for p in docmodel.load(str(out)).paragraphs], str(out)


FIGS = ["Figure 1: First", "Figure 2: Second", "Figure 3: Third"]


def test_ijo_fig_label_and_citation_before_the_full_stop(tmp_path):
    body = ["Results",
            "Vision improved at the last visit. (Figure 1A) The eye was quiet.",
            "There was a streak of hypopyon. [Figure 1A, B]",
            "Ulcer at presentation. [Fig 1a, 1b, 1c]. On history it was old.",
            "Pupils were dilated (Figure 1, top) Fundus was normal. They constricted with drops. (Figure 1, bottom). "
            "MRI was normal.",
            "( Figure 1 b- d)",
            "Both were seen (Figures 2 and 3)."] + FIGS
    t, _ = _run(_save(tmp_path, body), _style(), tmp_path, ["fig_citations", "guide_replace"])
    assert t[1] == "Vision improved at the last visit [Fig. 1A]. The eye was quiet."  # S300 "Fig. 1"; G245 upper
    assert t[2] == "There was a streak of hypopyon [Fig. 1A, B]."
    assert t[3] == "Ulcer at presentation [Fig. 1A-C]. On history it was old."  # G246: number not repeated
    assert t[4] == "Pupils were dilated [Fig. 1, top] Fundus was normal. They constricted with drops [Fig. 1, bottom]. " \
                   "MRI was normal."
    assert t[5] == "[Fig. 1B-D]"
    assert t[6] == "Both were seen [Figs. 2 and 3]."
    # Other journals keep "Figure" (G240) and the stop where it is (the setting is IJO's).
    t, _ = _run(_save(tmp_path, body, "b.docx"), _style("base"), tmp_path, ["fig_citations"])
    assert t[1] == "Vision improved at the last visit. [Figure 1A] The eye was quiet."
    assert t[3] == "Ulcer at presentation. [Figure 1A-C]. On history it was old."


def _bold_runs(path):
    doc = docmodel.load(path)
    out = []
    for p in doc.paragraphs:
        f = rules_format.char_formats(p)
        out.append("".join(ch.upper() if f[i].get("bold") else ch.lower() for i, ch in enumerate(p.text)))
    return out


def test_ijo_bold_figure_citations_become_roman(tmp_path):
    body = ["Results",
            [("Scores are shown [", {}), ("Table 1 and Figure 1", {"bold": True}), ("].", {})],
            [("Seen in the eye. (", {}), ("Figure 2", {"bold": True}), (")", {})],
            [("All of this line is bold [Figure 3].", {"bold": True})],
            "Table 1: A"] + FIGS
    t, out = _run(_save(tmp_path, body), _style(), tmp_path, ["fig_citations"])
    assert t[1] == "Scores are shown [Table 1 and Fig. 1]."
    assert t[2] == "Seen in the eye [Fig. 2]."
    b = _bold_runs(out)
    assert b[1] == "scores are shown [table 1 and fig. 1]."
    assert b[2] == "seen in the eye [fig. 2]."
    assert b[3] == "ALL OF THIS LINE IS BOLD [FIG. 3]."  # a bold line stays as it is
    t, out = _run(_save(tmp_path, body, "b.docx"), _style("base"), tmp_path, ["fig_citations"])
    assert _bold_runs(out)[1] == "scores are shown [TABLE 1 AND FIGURE 1]."  # off by default


def test_ijo_consecutive_reference_numbers_become_a_range(tmp_path):
    body = ["Introduction",
            "It is rare.[3,4,5] It is common.[1,3,4,5] Two stay.[1,2] Out of order stays.[5,3,4] Range kept.[6-8]"]
    path = _save(tmp_path, body + _refs(8))
    t, _ = _run(path, _style(), tmp_path, ["citations"])
    assert t[1] == "It is rare.[3-5] It is common.[1,3-5] Two stay.[1,2] Out of order stays.[5,3,4] Range kept.[6-8]"
    t, _ = _run(path, _style("base"), tmp_path, ["citations"])
    assert t[1] == body[1]


def test_reference_citation_forms_recognised(tmp_path):
    body = ["Introduction",
            "It predicts heart failure.[14.15] Later it was new.",
            "The syndrome, previously reported in 2008 (7), came back.",
            "Dislocation ranges from 0.85% to 2% [4], mostly early, and 0.01% [5] late.",
            ("It depends on local phagocytic mechanisms and on many other things that the text names at length, "
             "so that this paragraph typed in a heading style is running text all the same. [3]", "Heading 2"),
            "Others are cited.[1,2] And.[6] And.[8-13] And.[16]"]
    path = _save(tmp_path, body + _refs(16))
    st = _style()
    t, _ = _run(path, st, tmp_path, ["citations"])
    assert t[1] == "It predicts heart failure.[14,15] Later it was new."
    assert t[2] == "The syndrome, previously reported in 2008,[7] came back."
    assert t[3] == "Dislocation ranges from 0.85% to 2%,[4] mostly early, and 0.01%[5] late."
    assert t[4].endswith("all the same.[3]")
    qs = queries.collect(docmodel.load(path), st)
    assert not [q for q in qs.queries if q.check == "references.uncited"]  # 14 and 15 are cited


def test_values_are_still_not_citations(tmp_path):
    body = ["Introduction", "The mean was 2.5 (3) and in total 2008 (7) eyes; the ratio [1.5] held; CI [2.10]."]
    path = _save(tmp_path, body + _refs(9))
    t, _ = _run(path, _style(), tmp_path, ["citations"])
    assert t[1] == body[1]


def test_figure_parts_without_a_number(tmp_path):
    body = ["Case Report", "It was seen at presentation. (Figure A). Then it grew (Figure B-C) and (Figure  E-F).",
            "Figure 1: Clinical photographs"]
    path = _save(tmp_path, body + _refs(1))
    st = _style()
    t, _ = _run(path, st, tmp_path, ["fig_citations"])
    assert t[1] == "It was seen at presentation [Fig. 1A]. Then it grew [Fig. 1B-C] and [Fig. 1E-F]."
    assert not [q for q in queries.collect(docmodel.load(path), st).queries if q.check == "figures.uncited"]
    # Two figures: the number cannot be told, the citation stays as written.
    path = _save(tmp_path, body + ["Figure 2: More"] + _refs(1), "two.docx")
    t, _ = _run(path, st, tmp_path, ["fig_citations"])
    assert "(Figure B-C)" in t[1]


def test_author_name_against_the_reference_list(tmp_path):
    body = ["Introduction",
            "Rowan et al [1] described it. Nolan et al.[2] agreed. Mariot et al.[3] did too. "
            "Hale et al. [4] found more."]
    first = {1: "Hale JS", 2: "Nolan-Yáñez P", 3: "Marriot A", 4: "Hale JS"}
    path = _save(tmp_path, body + _refs(4, first))
    st = _style()
    qs = [q for q in queries.collect(docmodel.load(path), st).queries if q.check == "references.author_mismatch"]
    assert [q.text.split(": ", 1)[1] for q in qs] == [
        "Rowan et al. is cited as reference 1, whose first author is Hale JS.",
        "Mariot et al. is cited as reference 3, whose first author is Marriot A."]
    st.settings["citations"]["author_match_query"] = False
    assert not [q for q in queries.collect(docmodel.load(path), st).queries if q.check == "references.author_mismatch"]
