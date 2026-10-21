from docx import Document

from edit_os import biblio, docmodel, queries, refengine, renumber, rules, style, tracked
from edit_os.changes import Status


def _style():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["back_matter"]["add_missing"] = False  # keeps paragraph positions fixed
    return st


def _ref(n, title=None):
    return f"Author{n} A, Writer B. {title or f'Study number {n} of the eye'}. J Test 2001;{n}:1-2."


def _doc(tmp_path, body, refs, labels=True, tail=()):
    doc = Document()
    doc.add_paragraph("Introduction")
    for t in body:
        doc.add_paragraph(t)
    doc.add_paragraph("References")
    for i, r in enumerate(refs, 1):
        doc.add_paragraph(f"{i}. {r}" if labels else r)
    for t in tail:
        doc.add_paragraph(t)
    path = tmp_path / "r.docx"
    doc.save(path)
    return str(path)


def _apply(path, st, tmp_path):
    cs = rules.run_all(docmodel.load(path), st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")
    return [p.text for p in docmodel.load(str(out)).paragraphs]


def test_compress():
    st = _style()
    assert renumber.compress([5, 3, 4, 1], st) == "1,3-5"
    assert renumber.compress([2, 3], st) == "2,3"
    assert renumber.compress([7, 7, 2], st) == "2,7"
    assert renumber.expand("2-4, 7") == [2, 3, 4, 7] and renumber.expand("4-2") is None


def test_renumber_by_first_citation(tmp_path):
    refs = [_ref(n) for n in range(1, 7)]
    body = ["Alpha [1]. Beta [5]. Gamma [2-4]. Delta [4-6] and [1,5,6]."]
    path = _doc(tmp_path, body, refs)
    st = _style()
    pl = renumber.plan(docmodel.load(path), st)
    assert pl.apply and pl.mapping == {1: 1, 5: 2, 2: 3, 3: 4, 4: 5, 6: 6}
    t = _apply(path, st, tmp_path)
    # [5] -> [2]; [2-4] -> [3-5]; [4-6] = {5,2,6} -> [2,5,6]; [1,5,6] -> [1,2,6]
    assert t[1] == "Alpha.[1] Beta.[2] Gamma.[3-5] Delta[2,5,6] and.[1,2,6]"
    listed = t[3:9]
    assert listed[0].startswith("1. Author1 A")
    assert listed[1].startswith("2. Author5 A")
    assert listed[2].startswith("3. Author2 A") and listed[3].startswith("4. Author3 A")
    assert listed[4].startswith("5. Author4 A") and listed[5].startswith("6. Author6 A")
    assert all(x.endswith("2001;{}:1-2.".format(x.split("Author")[1][0])) for x in listed)
    qs = queries.collect(docmodel.load(path), st)
    assert "References are renumbered in text and list for citation sequence." in [q.text for q in qs.queries]


def test_round_bracket_document(tmp_path):
    """A document citing only in round brackets: the citation rule converts, renumbering remaps."""
    path = _doc(tmp_path, ["Alpha (2). Beta (1, 3). Six (6) weeks later (3-5)."], [_ref(n) for n in range(1, 6)])
    t = _apply(path, _style(), tmp_path)
    assert t[1] == "Alpha.[1] Beta.[2,3] Six (6) weeks later.[3-5]"
    assert t[3].startswith("1. Author2") and t[4].startswith("2. Author1")


def test_round_brackets_ignored_when_document_cites_in_square_brackets(tmp_path):
    path = _doc(tmp_path, ["Alpha [1]. Beta [2]. Klebsiella (3) grew. Gamma [3]. Then (2) eyes."],
                [_ref(n) for n in range(1, 4)])
    pl = renumber.plan(docmodel.load(path), _style())
    assert [o.numbers for o in pl.occurrences] == [[1], [2], [3]] and not pl.apply


def test_moved_entry_keeps_reference_formatting(tmp_path):
    refs = ["Lindqvist P., Hale R. Timing of surgery. J Med Genet 1987;24(3):616-620.",
            "Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit 2002;21:295-9."]
    path = _doc(tmp_path, ["First [2], then [1]."], refs)
    t = _apply(path, _style(), tmp_path)
    assert t[1] == "First,[1] then.[2]"
    assert t[3] == "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: A case report. Orbit 2002;21:295-9."
    assert t[4] == "2. Lindqvist P, Hale R. Timing of surgery. J Med Genet 1987;24:616-20."


def test_superscript_citations_renumbered(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph("Common worldwide.")
    p.add_run("2").font.superscript = True
    p.add_run(" Rare too")
    p.add_run("1,3").font.superscript = True
    p.add_run(".")
    doc.add_paragraph("References")
    for n in range(1, 4):
        doc.add_paragraph(f"{n}. {_ref(n)}")
    path = tmp_path / "s.docx"
    doc.save(path)
    t = _apply(str(path), _style(), tmp_path)
    assert t[1] == "Common worldwide.[1] Rare too.[2,3]"
    assert t[3].startswith("1. Author2") and t[4].startswith("2. Author1") and t[5].startswith("3. Author3")


def test_uncited_reference_blocks_renumbering(tmp_path):
    refs = [_ref(n) for n in range(1, 4)]
    path = _doc(tmp_path, ["Alpha [2]. Beta [1]."], refs)
    st = _style()
    pl = renumber.plan(docmodel.load(path), st)
    assert not pl.apply and pl.uncited == [3]
    cs = rules.run_all(docmodel.load(path), st, only=["renumber"])
    assert cs.changes == []
    texts = [q.text for q in queries.collect(docmodel.load(path), st).queries]
    assert "Reference no. 3 has not been cited. Kindly cite." in texts
    assert not any("renumbered" in x for x in texts)


def test_duplicate_reference_removed_and_list_renumbered(tmp_path):
    refs = [_ref(1), _ref(2), _ref(1).replace(",", " ,"), _ref(4)]  # 3 duplicates 1 (spacing differs)
    path = _doc(tmp_path, ["Alpha [1]. Beta [2]. Gamma [3]. Delta [4]; all [1-4]."], refs)
    st = _style()
    pl = renumber.plan(docmodel.load(path), st)
    assert pl.duplicates == [(1, 3)] and pl.mapping == {1: 1, 2: 2, 3: 1, 4: 3}
    t = _apply(path, st, tmp_path)
    assert t[1] == "Alpha.[1] Beta.[2] Gamma.[1] Delta;[3] all.[1-3]"
    assert t[3].startswith("1. Author1") and t[4].startswith("2. Author2")
    assert t[5].startswith("3. Author4") and t[6] == ""
    qs = queries.collect(docmodel.load(path), st)
    q = next(q for q in qs.queries if q.check == "references.duplicate_removed")
    assert q.text.startswith("References (1 and 3) are identical, so duplicate reference has been removed")


def test_duplicate_by_pmid(tmp_path, monkeypatch):
    refs = [_ref(1), _ref(2), _ref(3, "A differently worded title of study one")]
    path = _doc(tmp_path, ["Alpha [1]. Beta [2]. Gamma [3]."], refs)
    rec = biblio.Record(source="pubmed", title="x", pmid="123")

    def fake(ms, st):
        from edit_os import references
        out = []
        for r in references.extract(ms):
            pmid = "123" if r.number in (1, 3) else "999"
            out.append((r, refengine.Resolution(refengine.MATCHED, refengine.Parsed(r.text),
                                                biblio.Record(source="pubmed", title="x", pmid=pmid))))
        return out

    monkeypatch.setattr(rules, "lookup_references", fake)
    assert rec.pmid == "123"
    pl = renumber.plan(docmodel.load(path), _style())
    assert pl.duplicates == [(1, 3)] and pl.apply


def test_table_citations_follow_the_renumbering(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Alpha [2]. Beta [1].")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "As reported [2]"
    doc.add_paragraph("References")
    for n in range(1, 3):
        doc.add_paragraph(f"{n}. {_ref(n)}")
    path = tmp_path / "t.docx"
    doc.save(path)
    t = _apply(str(path), _style(), tmp_path)
    assert "As reported[1]" in t  # the citation rule now covers tables too
    assert t[1] == "Alpha.[1] Beta.[2]"


def test_bracketed_labels_and_auto_numbering_untouched_when_labels_out_of_sequence(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Alpha [2]. Beta [1].")
    doc.add_paragraph("References")
    doc.add_paragraph(f"[1] {_ref(1)}")
    doc.add_paragraph(f"[3] {_ref(3)}")  # list numbered 1, 3: inconsistent, leave it to the editor
    path = tmp_path / "b.docx"
    doc.save(path)
    assert not renumber.plan(docmodel.load(str(path)), _style()).apply


def test_bracketed_labels_renumbered(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Alpha [2]. Beta [1].")
    doc.add_paragraph("References")
    for n in (1, 2):
        doc.add_paragraph(f"[{n}] {_ref(n)}")
    path = tmp_path / "b.docx"
    doc.save(path)
    t = _apply(str(path), _style(), tmp_path)
    assert t[3].startswith("[1] Author2") and t[4].startswith("[2] Author1")


def test_citation_at_end_survives_added_back_matter(tmp_path):
    """The missing back-matter block and a citation are both inserted at the end of the last
    body paragraph: both must apply, the citation first."""
    path = _doc(tmp_path, ["Alpha (2). Beta (1, 3). Six (6) weeks later (3-5)."], [_ref(n) for n in range(1, 6)])
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    t = _apply(path, st, tmp_path)
    assert t[1] == "Alpha.[1] Beta.[2,3] Six (6) weeks later.[3-5]"
    assert t[2:6] == ["<H2>Financial support and sponsorship", "Nil.", "<H2>Conflicts of interest",
                      "There are no conflicts of interest."]


def _table_doc(tmp_path):
    """References cited only in a table at the end (after the list), as in submissions."""
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Alpha [1]. Beta [2]. Tables 1 and 2 summarise the reported cases.")
    doc.add_paragraph("Gamma [3].")
    doc.add_paragraph("References")
    for n in range(1, 7):
        doc.add_paragraph(f"{n}. {_ref(n)}")
    doc.add_paragraph("Table 1: Reported cases")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Rao et al., 2018[5]"
    table.cell(1, 0).text = "Lee et al., 2019[4]"
    doc.add_paragraph("Note. Adapted from Kim et al.[6]")
    path = tmp_path / "t.docx"
    doc.save(path)
    return str(path)


def test_table_citations_count_at_the_tables_first_callout(tmp_path):
    path = _table_doc(tmp_path)
    ms = docmodel.load(path)
    assert renumber.tables(ms)[0][0] == 1
    pl = renumber.plan(ms, _style())
    assert pl.apply and not pl.uncited
    assert pl.mapping == {1: 1, 2: 2, 5: 3, 4: 4, 6: 5, 3: 6}
    t = _apply(path, _style(), tmp_path)
    assert "Gamma.[6]" in t
    assert "Rao et al., 2018[3]" in t and "Lee et al., 2019[4]" in t and "Note. Adapted from Kim et al.[5]" in t
    refs = [x for x in t if x.startswith(tuple(f"{n}. " for n in range(1, 7)))]
    assert refs[2] == "3. " + _ref(5) and refs[5] == "6. " + _ref(3)


def test_qc_and_queries_count_table_citations(tmp_path):
    from edit_os import qc

    path = _table_doc(tmp_path)
    t = _apply(path, _style(), tmp_path)
    out = docmodel.load(str(tmp_path / "out.docx"))
    issues = {(i.check, i.message) for i in qc.check(out, _style()).issues}
    assert not [m for c, m in issues if c in ("references.cited", "citations.order")]
    qs = queries.collect(out, _style())
    assert not [q for q in qs.queries if q.check == "references.uncited"]
    assert t  # rendered


def test_entry_left_out_of_the_numbered_list_continues_it(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from edit_os import references

    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Alpha [1]. Beta [2]. Gamma [3].")
    doc.add_paragraph("References")
    for n in range(1, 4):
        p = doc.add_paragraph(_ref(n))
        if n < 3:  # Word numbering on the first two only
            num = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
                el = OxmlElement(tag)
                el.set(qn("w:val"), val)
                num.append(el)
            p._p.get_or_add_pPr().append(num)
    path = tmp_path / "n.docx"
    doc.save(path)
    assert [r.number for r in references.extract(docmodel.load(str(path)))] == [1, 2, 3]
