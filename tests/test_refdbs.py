"""Europe PMC, OpenAlex, Semantic Scholar and DOAJ behind PubMed / CrossRef (biblio clients, refengine.find_extra).
Invented papers only; HTTP is mocked (biblio._Http.get routed by host)."""

import io
import json
import urllib.error

import pytest

from edit_os import biblio, journals, provenance, refengine, style

ENTRY = ("Rostami K, Farahani M, Okonkwo C. Peer mentoring and clinical confidence of nursing students: "
         "a quasi-experimental study. J Nurs Teach Res. 2021;9(2):44-52.")
TITLE = "Peer mentoring and clinical confidence of nursing students: a quasi-experimental study"


def epmc_item(**kw):
    item = {"title": TITLE + ".", "pmid": "", "doi": "10.9999/jntr.2021.044", "pageInfo": "44-52",
            "authorList": {"author": [{"lastName": "Rostami", "initials": "K"}, {"lastName": "Farahani", "initials": "M"},
                                      {"lastName": "Okonkwo", "initials": "C"}]},
            "journalInfo": {"volume": "9", "issue": "2", "yearOfPublication": 2021,
                            "journal": {"title": "Journal of nursing teaching research", "issn": "1111-2222"}}}
    item.update(kw)
    return item


def openalex_work(**kw):
    w = {"id": "https://openalex.org/W1", "title": TITLE, "publication_year": 2021,
         "doi": "https://doi.org/10.9999/jntr.2021.044", "ids": {},
         "authorships": [{"author": {"display_name": n}} for n in ("Kaveh Rostami", "Maryam Farahani", "Chidi Okonkwo")],
         "primary_location": {"source": {"display_name": "Journal of Nursing Teaching Research", "issn_l": "1111-2222",
                                         "issn": ["1111-2222"]}},
         "biblio": {"volume": "9", "issue": "2", "first_page": "44", "last_page": "52"}}
    w.update(kw)
    return w


def s2_paper(**kw):
    d = {"title": TITLE, "year": 2021, "authors": [{"name": "K. Rostami"}, {"name": "Maryam Farahani"}],
         "journal": {"name": "Journal of Nursing Teaching Research", "volume": "9", "pages": "44 - 52"},
         "externalIds": {"DOI": "10.9999/jntr.2021.044"}}
    d.update(kw)
    return d


def doaj_article(**kw):
    b = {"title": TITLE, "year": "2021", "author": [{"name": "Kaveh Rostami"}, {"name": "Maryam Farahani"},
                                                   {"name": "Chidi Okonkwo"}],
         "journal": {"title": "Journal of Nursing Teaching Research", "volume": "9", "number": "2", "issns": ["1111-2222"]},
         "start_page": "44", "end_page": "52", "identifier": [{"type": "doi", "id": "10.9999/jntr.2021.044"}]}
    b.update(kw)
    return {"bibjson": b}


@pytest.fixture
def web(monkeypatch):
    """Routes biblio HTTP by host to `web.data[host]` (a callable(url) or a body); records every call. PubMed and
    CrossRef find nothing unless a test patches them."""
    class Web:
        data: dict = {}
        calls: list = []
        headers: list = []

    def get(self, url, host_gap, headers=None, timeout=30, breaker=False):
        Web.calls.append(url)
        Web.headers.append(headers or {})
        for host, body in Web.data.items():
            if host in url:
                return body(url) if callable(body) else body
        return None

    Web.data, Web.calls, Web.headers = {}, [], []
    monkeypatch.setattr(biblio._Http, "get", get)
    monkeypatch.setattr(biblio, "_http", None)
    monkeypatch.setattr(biblio._Http, "__init__", lambda self, path=None: None)
    monkeypatch.setattr(journals, "directory", lambda: None)
    for name in ("pubmed_search", "pubmed_best_match"):
        monkeypatch.setattr(biblio, name, lambda *a, **k: [])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [])
    monkeypatch.setattr(biblio, "crossref_search", lambda *a, **k: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    return Web


def epmc(*items):
    return json.dumps({"resultList": {"result": list(items)}})


def oa(*works):
    return json.dumps({"results": list(works)})


def doaj(*arts):
    return json.dumps({"results": list(arts)})


ALL = list(biblio.EXTRA_SOURCES)
REAL_GET = biblio._Http.get  # (conftest stubs it out for every test)


# --- clients ---------------------------------------------------------------------------------

@pytest.mark.parametrize("name,expected", [
    ("Zohreh Nabizadeh-Gharghozar", ("Nabizadeh-Gharghozar", "Z")),
    ("M. Adib-Hajbaghery", ("Adib-Hajbaghery", "M")),
    ("Ludwig van Beethoven", ("van Beethoven", "L")),
    ("Smith, John A.", ("Smith", "JA")),
    ("Smith JA", ("Smith", "JA")),
    ("Seyed Saeed Naghavi Alhosseini", ("Alhosseini", "SSN")),
    ("Farhati Farhati", ("Farhati", "F")),
    ("Sulistyorini", ("Sulistyorini", "")),
])
def test_split_name(name, expected):
    assert biblio.split_name(name) == expected


def test_europepmc_record_and_query(web):
    web.data["ebi.ac.uk"] = epmc(epmc_item())
    [r] = biblio.europepmc_search(TITLE, "Rostami", "2021")
    assert (r.source, r.authors[0], r.journal, r.journal_full) == ("europepmc", ("Rostami", "K"), "Journal of nursing teaching research",
                                                                    "Journal of nursing teaching research")
    assert (r.year, r.volume, r.issue, r.pages, r.doi, r.issn) == ("2021", "9", "2", "44-52", "10.9999/jntr.2021.044", ["1111-2222"])
    q = web.calls[0]
    assert "TITLE%3A%28" in q and "AUTH%3A%22Rostami%22" in q and "PUB_YEAR%3A%5B2020+TO+2022%5D" in q


def test_openalex_record_and_key_header(web, monkeypatch):
    monkeypatch.setenv("OPENALEX_API_KEY", "k-test")
    web.data["openalex.org"] = oa(openalex_work(ids={"pmid": "https://pubmed.ncbi.nlm.nih.gov/123456"}))
    [r] = biblio.openalex_search(TITLE, "2021")
    assert r.source == "openalex" and r.authors[:2] == [("Rostami", "K"), ("Farahani", "M")]
    assert (r.pages, r.doi, r.pmid, r.journal) == ("44-52", "10.9999/jntr.2021.044", "123456", "Journal of Nursing Teaching Research")
    assert "publication_year%3A2020-2022" in web.calls[0] and "k-test" not in web.calls[0]  # key never in the URL / cache key
    assert web.headers[0] == {"Authorization": "Bearer k-test"}


def test_semantic_scholar_match_and_doi(web, monkeypatch):
    monkeypatch.setenv("SEMANTIC_SCHOLAR_API_KEY", "s2-test")
    web.data["semanticscholar.org"] = lambda url: json.dumps({"data": [s2_paper()]} if "/search/match" in url else s2_paper())
    [r] = biblio.s2_match(TITLE, "2021")
    assert (r.source, r.authors[0], r.pages, r.volume) == ("semantic_scholar", ("Rostami", "K"), "44-52", "9")
    [d] = biblio.s2_doi("10.9999/jntr.2021.044")
    assert d.doi == "10.9999/jntr.2021.044" and "/DOI:10.9999/jntr.2021.044" in web.calls[1]
    assert web.headers[0] == {"x-api-key": "s2-test"}


def test_doaj_record(web):
    web.data["doaj.org"] = doaj(doaj_article(title="Online scenario- based learning: participants\\' views"))
    [r] = biblio.doaj_search(TITLE, "2021")
    assert r.title == "Online scenario-based learning: participants' views"
    assert (r.source, r.pages, r.issue, r.doi) == ("doaj", "44-52", "2", "10.9999/jntr.2021.044")
    assert "bibjson.title" in web.calls[0]


def test_journal_abbreviation_from_issn(web, monkeypatch):
    class Dir:
        def lookup(self, key):
            return journals.Journal("1", "Journal of Nursing Teaching Research", "J Nurs Teach Res", "J. Nurs. Teach. Res.",
                                    "1111-2222", "") if key == "1111-2222" else None
    monkeypatch.setattr(journals, "directory", lambda: Dir())
    web.data["doaj.org"] = doaj(doaj_article())
    [r] = biblio.doaj_search(TITLE, "2021")
    assert (r.journal, r.journal_full) == ("J Nurs Teach Res", "Journal of Nursing Teaching Research")


def test_user_agent_mailto_from_env(monkeypatch):
    monkeypatch.delenv("EDITOS_MAILTO", raising=False)
    assert "mailto" not in biblio.user_agent()
    monkeypatch.setenv("EDITOS_MAILTO", "refs@example.org")
    assert biblio.user_agent().endswith("(reference checking; mailto:refs@example.org)")


def test_breaker_skips_a_host_out_of_quota(monkeypatch, tmp_path):
    http = biblio._Http(str(tmp_path / "c.sqlite"))
    calls = []

    def urlopen(req, timeout=30):
        calls.append(req.full_url)
        raise urllib.error.HTTPError(req.full_url, 429, "Too Many", {"Retry-After": "35000"}, io.BytesIO(b""))

    monkeypatch.setattr(biblio.urllib.request, "urlopen", urlopen)
    monkeypatch.setattr(biblio._Http, "get", REAL_GET)
    monkeypatch.setattr(biblio.time, "sleep", lambda s: None)
    assert http.get("https://api.openalex.org/works?filter=a", 0, breaker=True) is None
    assert http.get("https://api.openalex.org/works?filter=b", 0, breaker=True) is None
    assert len(calls) == 1  # a used-up daily quota: one try, then the host is left alone


def test_not_found_body_is_cached_empty(monkeypatch, tmp_path):
    http = biblio._Http(str(tmp_path / "c.sqlite"))
    calls = []

    def urlopen(req, timeout=30):
        calls.append(1)
        raise urllib.error.HTTPError(req.full_url, 404, "Not Found", {}, io.BytesIO(b""))

    monkeypatch.setattr(biblio.urllib.request, "urlopen", urlopen)
    monkeypatch.setattr(biblio._Http, "get", REAL_GET)
    monkeypatch.setattr(biblio.time, "sleep", lambda s: None)
    assert http.get("https://api.semanticscholar.org/x", 0, breaker=True) == ""
    assert http.get("https://api.semanticscholar.org/x", 0, breaker=True) == "" and len(calls) == 1


# --- find: order, acceptance, preference ---------------------------------------------------------

def test_extra_source_used_only_when_pubmed_and_crossref_fail(web, monkeypatch):
    web.data["ebi.ac.uk"] = epmc(epmc_item())
    r, why = refengine.find(refengine.parse(ENTRY))
    assert r.source == "europepmc" and why == ""
    pub = biblio.Record(source="pubmed", pmid="9", title=TITLE + ".", authors=[("Rostami", "K")], journal="J Nurs Teach Res",
                        year="2021", volume="9", pages="44-52")
    monkeypatch.setattr(biblio, "pubmed_search", lambda *a, **k: ["9"])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [pub] if ids == ["9"] else [])
    web.calls.clear()
    r, _ = refengine.find(refengine.parse(ENTRY))
    assert r.source == "pubmed" and not web.calls  # PubMed has it: no other database asked


def test_sources_in_priority_order_and_per_journal(web):
    web.data["ebi.ac.uk"] = epmc(epmc_item())
    web.data["openalex.org"] = oa(openalex_work())
    r, _ = refengine.find(refengine.parse(ENTRY), ALL)
    assert r.source == "europepmc" and not any("openalex" in u for u in web.calls)
    r, _ = refengine.find(refengine.parse(ENTRY), ["pubmed", "crossref", "openalex", "doaj"])
    assert r.source == "openalex" and not any("ebi.ac.uk" in u for u in web.calls[-1:])
    r, why = refengine.find(refengine.parse(ENTRY), ["pubmed", "crossref"])
    assert r is None and why == "Not found in PubMed or CrossRef"


def test_style_default_lists_every_source():
    assert style.load("jehp").get("references.sources") == list(biblio.SOURCES)


def test_hook_none_with_reason_when_every_database_fails(web):
    r, why = refengine.find(refengine.parse(ENTRY), ALL)
    assert r is None
    assert why == "Not found in PubMed, CrossRef, Europe PMC, OpenAlex, Semantic Scholar or DOAJ"


@pytest.mark.parametrize("change", [
    {"authorList": {"author": [{"lastName": "Hosseini", "initials": "K"}]}},            # another first author
    {"journalInfo": {"volume": "9", "yearOfPublication": 2022, "journal": {"title": "J"}}},  # a year apart: not here
    {"journalInfo": {"volume": "31", "yearOfPublication": 2021, "journal": {"title": "J"}}, "pageInfo": "300-9"},  # vol + page
    {"title": "Peer mentoring and clinical confidence of medical residents: a cohort study."},  # another title
    {"journalInfo": {"yearOfPublication": 2021, "journal": {"title": "J"}}},             # no volume: not usable
])
def test_strict_acceptance_holds_for_extra_sources(web, change):
    web.data["ebi.ac.uk"] = epmc(epmc_item(**change))
    r, _ = refengine.find(refengine.parse(ENTRY), ["europepmc"])
    assert r is None


def test_two_papers_passing_is_no_match(web):
    other = epmc_item(doi="10.9999/other", pageInfo="44-52",
                      journalInfo={"volume": "9", "yearOfPublication": 2021, "journal": {"title": "Another journal"}})
    web.data["ebi.ac.uk"] = epmc(epmc_item(), other)
    r, why = refengine.find(refengine.parse(ENTRY), ["europepmc"])
    assert r is None and "Several papers in Europe PMC" in why


def test_organisation_as_author_still_rejected(web):
    web.data["doaj.org"] = doaj(doaj_article(author=[{"name": "Department of Nursing, Tehran University, Tehran"},
                                                    {"name": "Kaveh Rostami"}]))
    r, _ = refengine.find(refengine.parse(ENTRY), ["doaj"])
    assert r is None  # the first author is not the entry's


def test_pubmed_record_preferred_when_it_has_the_paper(web, monkeypatch):
    pub = biblio.Record(source="pubmed", pmid="555", doi="10.9999/jntr.2021.044", title=TITLE + ".",
                        authors=[("Rostami", "K"), ("Farahani", "M"), ("Okonkwo", "C")], journal="J Nurs Teach Res",
                        year="2021", volume="9", pages="44-52")
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [pub] if ids == ["555"] else [])
    web.data["openalex.org"] = oa(openalex_work(ids={"pmid": "https://pubmed.ncbi.nlm.nih.gov/555"}))
    r, _ = refengine.find(refengine.parse(ENTRY), ["openalex"])
    assert (r.source, r.found_via) == ("pubmed", "openalex")


def test_disagreeing_trusted_record_wins(web, monkeypatch):
    cr = biblio.Record(source="crossref", doi="10.9999/jntr.2021.044", title=TITLE, authors=[("Rostami", "K")],
                       journal="J Nurs Teach Res", year="2019", volume="9", pages="44-52")
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: cr)
    web.data["openalex.org"] = oa(openalex_work())
    r, why = refengine.find(refengine.parse(ENTRY), ["openalex"])
    assert r is None and "CrossRef record differs" in why


def test_doi_in_entry_checked_by_extra_source(web):
    entry = ENTRY + " doi:10.9999/jntr.2021.044"
    web.data["doaj.org"] = lambda url: doaj(doaj_article()) if "doi%3A" in url else doaj()
    r, _ = refengine.find(refengine.parse(entry), ["doaj"])
    assert r.source == "doaj" and r.doi == "10.9999/jntr.2021.044"
    web.data["doaj.org"] = lambda url: doaj(doaj_article(title="Hand hygiene in surgical wards")) if "doi%3A" in url else doaj()
    r, why = refengine.find(refengine.parse(entry), ["doaj"])
    assert r is None and "points to a different article" in why


def test_compound_surname_cut_as_the_entry_cuts_it(web):
    entry = ("Naghavi Alhosseini SS, Alavi-Majd H, Yanti RD. Peer mentoring and clinical confidence of nursing students: "
             "a quasi-experimental study. J Nurs Teach Res. 2021;9(2):44-52.")
    web.data["doaj.org"] = doaj(doaj_article(author=[{"name": "Seyed Saeed Naghavi Alhosseini"}, {"name": "Hamid Alavi Majd"},
                                                    {"name": "Risna Yanti"}]))
    r, _ = refengine.find(refengine.parse(entry), ["doaj"])
    assert r.authors == [("Naghavi Alhosseini", "SS"), ("Alavi-Majd", "H"), ("Yanti", "RD")]


def test_unwritable_journal_name_takes_the_entrys(web):
    web.data["doaj.org"] = doaj(doaj_article(journal={"title": "حیات", "volume": "9", "issns": ["1111-2222"]}))
    st = style.load("jehp")
    res = refengine.resolve(ENTRY, st)
    assert res.record.source == "doaj" and " J Nurs Teach Res 2021;9:44-52." in res.formatted
    bare = "Rostami K, Farahani M, Okonkwo C. " + TITLE + ". 2021"  # no journal to write instead: not used
    r, _ = refengine.find(refengine.parse(bare), ["doaj"])
    assert r is None


def test_non_indexed_journal_written_in_full(web):
    web.data["doaj.org"] = doaj(doaj_article())
    res = refengine.resolve(ENTRY, style.load("mjdrdypu"))
    assert res.formatted.endswith("Journal of Nursing Teaching Research 2021;9:44-52.")
    res = refengine.resolve(ENTRY, style.load("jehp"))  # ISO 4, as the JEHP editors write it
    assert res.formatted.endswith(" J Nurs Teach Res 2021;9:44-52.")


def test_provenance_chip_names_the_source(web):
    web.data["openalex.org"] = oa(openalex_work())
    res = refengine.resolve(ENTRY, style.load("jehp"))
    prov = provenance.reference(res, ENTRY)
    assert prov["kind"] == "openalex" and prov["label"] == "Verified in OpenAlex"
    assert prov["url"] == "https://doi.org/10.9999/jntr.2021.044" and prov["source_name"] == "OpenAlex"


def test_provenance_found_via(web):
    rec = biblio.Record(source="pubmed", pmid="555", title="T", authors=[("A", "B")], found_via="europepmc")
    res = refengine.Resolution("format_fixed", refengine.parse(ENTRY), rec, "x")
    prov = provenance.reference(res, ENTRY)
    assert prov["label"] == "PubMed 555" and prov["found_via"] == "Europe PMC"


def test_record_without_first_part_of_compound_surname_keeps_the_entrys(web):
    entry = ("Buckingham Shum S, Farahani M. Peer mentoring and clinical confidence of nursing students: "
             "a quasi-experimental study. J Nurs Teach Res. 2021;9(2):44-52.")
    web.data["semanticscholar.org"] = json.dumps({"data": [s2_paper(authors=[{"name": "S. Shum"}, {"name": "M. Farahani"}])]})
    r, _ = refengine.find(refengine.parse(entry), ["semantic_scholar"])
    assert r.authors[0] == ("Buckingham Shum", "S")


def test_entry_pages_kept_when_only_the_last_page_differs(web):
    web.data["semanticscholar.org"] = json.dumps({"data": [s2_paper(journal={"name": "journal of nursing teaching research",
                                                                              "volume": "9", "pages": "44-53"})]})
    res = refengine.resolve(ENTRY, style.load("jehp"))
    assert res.record.source == "semantic_scholar"
    assert res.formatted.endswith(" J Nurs Teach Res 2021;9:44-52.")



def test_record_differs_message_without_field_names_gives_a_query(monkeypatch, tmp_path):
    # "Found in Europe PMC, but the PubMed record differs ..." names no field: it crashed queries.collect (JEHP 1764)
    from docx import Document
    from edit_os import docmodel, queries, references, rules
    res = refengine.Resolution(refengine.QUERY, refengine.parse(ENTRY),
                               message=f"Found in Europe PMC, but the PubMed record {refengine.DIFFERS}")
    monkeypatch.setattr(rules, "lookup_references", lambda ms, st: [(r, res) for r in references.extract(ms)])
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph("1. " + ENTRY)
    doc.save(tmp_path / "m.docx")
    qs = queries.collect(docmodel.load(str(tmp_path / "m.docx")), style.load("jehp"))
    assert any(q.text == "The details given here differ from the published article. Kindly check." for q in qs.queries)
