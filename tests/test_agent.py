"""Editing agent (prototype), offline: invented text, a fake Claude that answers with tool calls."""

import json

import pytest
from docx import Document

from edit_os import agent, copyedit, docmodel, refengine, style, tracked
from edit_os.agent import context, loop, tools
from edit_os.changes import Source, Status


@pytest.fixture(autouse=True)
def private_cache(tmp_path, monkeypatch):
    monkeypatch.setattr(loop, "_cache", loop._Cache(str(tmp_path / "agent_cache.sqlite")))
    for env in ("EDITOS_CE_EXAMPLES", "EDITOS_AGENT_MODEL", "EDITOS_AGENT_BUDGET", "EDITOS_STYLESHEETS",
                "EDITOS_AGENT_MIN_CONFIDENCE"):
        monkeypatch.delenv(env, raising=False)


BODY = [
    "Introduction",
    "The patients was followed up for six months after the surgery in all of the participating centres.",
    "However the difference between the two groups was small and it was seen in 12 patients [3].",
    "Brown et al. reported a similar outcome in a larger cohort of older adults with diabetes [2].",
    "References",
    "1. Green A, White B. A first invented paper on eye care. J Invented Ophthalmol 2019;12:1-5.",
    "2. Black C, Grey D. Another invented paper on outcomes in diabetes. Invented Med J 2020;3:10-20.",
]


def make(tmp_path, paras=BODY):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return docmodel.load(str(path))


def journal(**cfg):
    st = style.load("ijo")
    st.settings["copyedit"]["agent"] = {"enabled": True, **cfg}
    return st


def use(name, inp, i=0):
    return {"type": "tool_use", "id": f"tu{name}{i}", "name": name, "input": inp}


class FakeClaude:
    """Answers each call with the next scripted list of content blocks."""

    def __init__(self, turns, usage=None):
        self.turns, self.calls = list(turns), []
        self.usage = usage or {"input_tokens": 1000, "output_tokens": 500, "cache_creation_input_tokens": 0,
                               "cache_read_input_tokens": 0}

    def __call__(self, params):
        self.calls.append(json.loads(json.dumps(params)))
        content = self.turns.pop(0) if self.turns else [{"type": "text", "text": "Done."}]
        stop = "tool_use" if any(b["type"] == "tool_use" for b in content) else "end_turn"
        return {"content": content, "stop_reason": stop, "usage": dict(self.usage), "model": params["model"]}


def pid(ms, text):
    """The agent's id ("P<n>") of the paragraph starting with `text`."""
    for i, p in enumerate(ms.paragraphs):
        if p.text.startswith(text):
            return context.label(i)
    raise KeyError(text)


def ed(p, old, new, kind="agreement", confidence="certain"):
    return dict(paragraph_id=p, old=old, new=new, kind=kind, reason="why", confidence=confidence)


def test_kept_edits_are_suggested_ai_changes_with_exact_offsets(tmp_path):
    ms = make(tmp_path)
    p1, p2 = pid(ms, "The patients"), pid(ms, "However")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p1, "patients was", "patients were"),
                                                        ed(p2, "However the", "However, the", "punctuation")]})]])
    res = agent.run(ms, journal(), call=fake)
    assert len(res.changes.changes) == 2 and res.stop == "done"
    for c in res.changes.changes:
        text = ms.by_id(c.para_id).text
        assert text[c.start:c.end] == c.old
        assert c.source == Source.AI and c.status == Status.SUGGESTED and c.reason.startswith("agent.")
        assert c.provenance["checker"] == "agent"
    # the tool result went back to Claude
    result = json.loads(fake.calls[1]["messages"][-1]["content"][0]["content"])
    assert all(r["kept"] for r in result["results"])
    # renders as tracked changes once accepted (agent output is never accepted by the agent itself)
    for c in res.changes.changes:
        c.status = Status.ACCEPTED
    tracked.render(ms, res.changes, str(tmp_path / "out.docx"))


@pytest.mark.parametrize("old,new,why", [
    ("12 patients", "twelve patients", "numbers"),
    ("patients [3]", "patients", "numbers"),
    ("small and it", "small (and it", "symbols"),
    ("Brown et al.", "Browne et al.", "name"),
    ("not in the text", "x", "not found"),
    ("the", "a", "more than once"),
])
def test_unsafe_edits_are_refused_with_a_reason(tmp_path, old, new, why):
    ms = make(tmp_path)
    p = pid(ms, "However") if "Brown" not in old else pid(ms, "Brown")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p, old, new)]})]])
    res = agent.run(ms, journal(), call=fake)
    assert res.changes.changes == []
    answer = json.loads(fake.calls[1]["messages"][-1]["content"][0]["content"])["results"][0]
    assert not answer["kept"] and why in answer["why"]


def test_headings_references_and_multi_sentence_edits_are_refused(tmp_path):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    edits = [ed(pid(ms, "Introduction"), "Introduction", "Background"),
             ed(pid(ms, "2. Black"), "Another invented", "An invented"),
             ed(p1, "centres.", "centres. They improved.", "other")]
    ms2 = make(tmp_path, BODY[:1] + ["The patients was seen. They was happy with the care they got there."] + BODY[2:])
    p_two = pid(ms2, "The patients")
    fake = FakeClaude([[use("propose_edits", {"edits": edits})]])
    res = agent.run(ms, journal(), call=fake)
    assert res.changes.changes == []
    fake2 = FakeClaude([[use("propose_edits", {"edits": [ed(p_two, "was seen. They was", "were seen. They were")]})]])
    assert agent.run(ms2, journal(), call=fake2).changes.changes == []


def test_overlapping_and_overlong_edits_are_refused(tmp_path):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    long_old = "followed up for six months after the surgery in all of the participating centres"
    fake = FakeClaude([[use("propose_edits", {"edits": [
        ed(p1, "patients was", "patients were"), ed(p1, "The patients was", "Patients were"),
        ed(p1, long_old, "followed up for half a year after their operations at every one of the participating "
                         "hospitals, clinics and other centres in the whole of the study network")]})]])
    res = agent.run(ms, journal(max_words=12), call=fake)
    assert [c.new for c in res.changes.changes] == ["were"]
    assert {e["why"] for e in res.log if not e["kept"]} >= {"overlaps an edit already kept in this paragraph"}


def test_min_confidence_certain_drops_likely_edits(tmp_path):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p1, "patients was", "patients were", confidence="likely")]})]])
    assert agent.run(ms, journal(min_confidence="certain"), call=fake).changes.changes == []


def test_reference_queries_need_a_code_verified_lookup(tmp_path):
    ms = make(tmp_path)
    p3 = pid(ms, "Brown")
    q = {"paragraph_id": p3, "anchor": "Brown et al.", "text": "Brown is not the first author of ref. 2. Please check."}
    calls = []

    def resolver(text, st):
        calls.append(text)
        return refengine.Resolution(refengine.QUERY, refengine.parse(text), message="Not found in PubMed or CrossRef")

    fake = FakeClaude([[use("add_queries", {"queries": [q]})],
                       [use("check_reference", {"number": 2})],
                       [use("add_queries", {"queries": [q]}, 1)]])
    res = agent.run(ms, journal(), call=fake, resolver=resolver)
    first = json.loads(fake.calls[1]["messages"][-1]["content"][0]["content"])["results"][0]
    assert not first["kept"] and "check_reference" in first["why"]
    looked = json.loads(fake.calls[2]["messages"][-1]["content"][0]["content"])
    assert looked["first_author_in_entry"] == "Black" and looked["lookup"].startswith("not verified")
    assert looked["cited_in"] == [p3]
    assert len(res.queries) == 1 and res.queries[0].check == "agent.query"
    q0 = res.queries[0]
    assert ms.by_id(q0.para_id).text[q0.start:q0.end] == "Brown et al."
    assert calls and "Black C" in calls[0]


def test_budget_stops_the_loop_and_cost_is_recorded(tmp_path):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    usage = {"input_tokens": 100000, "output_tokens": 50000, "cache_creation_input_tokens": 0,
             "cache_read_input_tokens": 0}  # $1.40 a call at Opus 5.5 prices
    turns = [[use("search_text", {"text": "patients"}, i)] for i in range(5)]
    turns[0] = [use("propose_edits", {"edits": [ed(p1, "patients was", "patients were")]})]
    fake = FakeClaude(turns, usage)
    res = agent.run(ms, journal(budget_usd=2.0), call=fake)
    assert res.stop == "budget" and len(fake.calls) == 2
    assert res.cost == pytest.approx(2.8) and res.spent == pytest.approx(2.8)
    assert len(res.changes.changes) == 1  # the edits of the last answer are still checked and kept
    # the warning to finish was sent after the first call
    assert any(b.get("type") == "text" and "Budget" in b.get("text", "") for b in fake.calls[1]["messages"][-1]["content"])


def test_rerun_is_served_from_the_disk_cache(tmp_path):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    script = [[use("propose_edits", {"edits": [ed(p1, "patients was", "patients were")]})]]
    first = agent.run(ms, journal(), call=FakeClaude(script))
    again = FakeClaude([])
    second = agent.run(ms, journal(), call=again)
    assert again.calls == [] and second.spent == 0 and second.cost == pytest.approx(first.cost)
    assert [c.new for c in second.changes.changes] == ["were"]


def test_examples_come_from_training_articles_only(tmp_path):
    from edit_os.copyedit import claude_pass

    names = [f"IJO_{n}_26" for n in range(1, 60)]
    test_name = next(n for n in names if claude_pass.is_test_article(n))
    train_name = next(n for n in names if not claude_pass.is_test_article(n))
    rows = [dict(article=a, stage="copy_edit", para=1, old="was", new="were", before="In this invented study the "
                 "patients ", after=" seen in the clinic every week.", in_refs=False, authors=[])
            for a in (test_name, train_name)]
    rows[0]["new"] = "LEAK"
    path = tmp_path / "changes.jsonl"
    path.write_text("\n".join(json.dumps(r) for r in rows))
    got = context.load_examples(str(path), "ijo")
    assert len(got) == 1 and got[0]["new"] == "were"
    assert context.load_examples(str(path), "ijo", exclude={train_name}) == []


def test_manuscript_view_tags_paragraphs(tmp_path):
    ms = make(tmp_path)
    text, ids = context.manuscript_text(ms, loop.editable_paragraphs(ms))
    assert f"[{pid(ms, 'The patients')}] The patients" in text
    assert f"[{pid(ms, 'Introduction')}] (heading) Introduction" in text
    assert f"[{pid(ms, '1. Green')}] (reference) 1. Green" in text


def test_run_all_hook_is_off_by_default(tmp_path, monkeypatch):
    ms = make(tmp_path)
    called = []
    monkeypatch.setattr(loop, "_api_call", lambda params: called.append(1))
    res = copyedit.run_all(ms, style.load("ijo"), only=["agent"])
    assert res.changes.changes == [] and called == []


def test_run_all_hook_runs_the_agent_when_enabled(tmp_path, monkeypatch):
    ms = make(tmp_path)
    p1 = pid(ms, "The patients")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p1, "patients was", "patients were")]})]])
    monkeypatch.setattr(loop, "_api_call", fake)
    monkeypatch.setattr(loop.ai_refs, "available", lambda: True)
    res = copyedit.run_all(ms, journal(), only=["agent"])
    assert [c.new for c in res.changes.changes] == ["were"]


def test_tools_are_strict_with_closed_schemas():
    for t in tools.TOOLS:
        assert t["strict"] and t["input_schema"]["additionalProperties"] is False


def test_splitting_a_run_on_sentence_is_allowed(tmp_path):
    ms = make(tmp_path, BODY[:1] + ["The scores rose in both groups, this was expected from the earlier work."] + BODY[2:])
    p = pid(ms, "The scores")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p, "groups, this was", "groups. This was", "punctuation")]})]])
    res = agent.run(ms, journal(), call=fake)
    assert [(c.old, c.new) for c in res.changes.changes] == [(", this", ". This")]


def test_rewording_next_to_numbers_is_refused(tmp_path):
    ms = make(tmp_path, BODY[:1] + ["Of 36 people, 4 were excluded (2 were ill, 1 had moved, 1 was away, 2 declined)."]
              + BODY[2:])
    p = pid(ms, "Of 36")
    fake = FakeClaude([[use("propose_edits", {"edits": [ed(p, "2 were ill, 1 had moved,",
                                                          "2 were ill, of whom 1 had moved and", "other")]})]])
    res = agent.run(ms, journal(), call=fake)
    assert res.changes.changes == [] and "numbers" in res.log[0]["why"]


def test_query_saying_a_listed_work_is_missing_is_refused(tmp_path):
    ms = make(tmp_path)
    p3 = pid(ms, "Brown")
    q = {"paragraph_id": p3, "anchor": "Brown et al.", "text": "Brown et al. does not appear in the reference list."}
    q_black = {"paragraph_id": pid(ms, "However"), "anchor": "the two groups",
               "text": "Black and Grey is not in the reference list. Please add it."}
    fake = FakeClaude([[use("add_queries", {"queries": [q, {**q_black, "anchor": "difference between"}]})]])
    res = agent.run(ms, journal(), call=fake)
    assert len(res.queries) == 2  # Brown is truly missing; the second anchor names no one the list holds
    ms2 = make(tmp_path, BODY[:3] + ["Black et al. reported a similar outcome in a cohort [2]."] + BODY[4:])
    q2 = {"paragraph_id": pid(ms2, "Black et al."), "anchor": "Black et al.",
          "text": "Black et al. does not appear in the reference list."}
    res2 = agent.run(ms2, journal(), call=FakeClaude([[use("add_queries", {"queries": [q2]})]]))
    assert res2.queries == [] and "Black" in res2.log[0]["why"]


def test_ijo_pilot_settings_and_one_reject_group():
    # owner, 2026-10-28: on for IJO, at most 5 queries, its edits left for the editor and grouped as one rule
    from edit_os import autopilot, rule_sources, style
    from edit_os.changes import Change, Layer, Source
    st = style.load("ijo")
    assert st.get("copyedit.agent.enabled") is True and st.get("copyedit.agent.max_queries") == 5
    assert not style.load("jehp").get("copyedit.agent.enabled")
    assert rule_sources.key_for("agent.grammar (subject-verb)") == rule_sources.key_for("agent.word_choice") == "agent"
    c = Change(id="a1", para_id="p", start=0, end=1, old="a", new="b", layer=Layer.LANGUAGE, source=Source.AI,
               reason="agent.grammar (x)")
    p = autopilot.plan([c], [], st, judge=lambda items: {i.id: ("accept", "") for i in items})
    assert p.accept == [] and p.left[0]["id"] == "a1"
