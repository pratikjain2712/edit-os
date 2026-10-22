"""Stylesheet-first settings (CLAUDE.md): synthetic examples of rules the journals' stylesheets and the pre-editing
guide ask for, applied whatever the editors' habits are."""

from docx import Document

from edit_os import docmodel, queries, rules, style, tracked
from edit_os.changes import Status

BODY = ("The disease is common in older patients and the patients were treated with surgery in this "
        "series of cases from the clinic.")


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def run(path, st, only):
    cs = rules.run_all(docmodel.load(path), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(path).replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(out).paragraphs]


def test_caption_final_period_per_kind(tmp_path):
    path = make(tmp_path, ["Results", BODY, "Table 1. Baseline data", "Figure 1. Axial scan.",
                           "Figure 2: (A) Axial and (B) sagittal views."])
    # Base (the shared sheet): table captions end with a period; legends lose theirs (G261); parts "(a)" (G263).
    t = run(path, style.load("jehp"), ["fig_captions"])
    assert t[2:5] == ["Table 1: Baseline data.", "Figure 1: Axial scan", "Figure 2: (a) Axial and (b) sagittal views"]
    # SGMJ: 'Figure caption | “Figure 1:” followed by caption in sentence case; bold; should have an end period'
    t = run(path, style.load("sgmj"), ["fig_captions"])
    assert t[2:4] == ["Table 1: Baseline data.", "Figure 1: Axial scan."]
    # ABR has no "The caption should end with a period": G200 deletes it.
    path = make(tmp_path, ["Results", BODY, "Table 1. Baseline data."], name="m2.docx")
    assert run(path, style.load("abr"), ["fig_captions"])[2] == "Table 1: Baseline data"


def test_h1_case_and_methods_heading(tmp_path):
    path = make(tmp_path, ["INTRODUCTION", BODY, "Methods", BODY, "Results", BODY, "References"])
    # G133 / S "H1 | Title case"; G138 'The heading "METHODS" ... should be changed to "Materials and Methods"'
    t = run(path, style.load("jehp"), ["heading_tags", "sections_headings"])
    assert t[0] == "<H1>Introduction" and t[2] == "<H1>Materials and Methods"
    # IJA: "Methods (not Materials and Methods ...)" and "H1 | All Caps"
    t = run(path, style.load("ija"), ["heading_tags", "sections_headings"])
    assert t[0] == "<H1>INTRODUCTION" and t[2] == "<H1>METHODS" and t[4] == "<H1>RESULTS"


def test_abstract_subheadings_title_case_and_methods(tmp_path):
    path = make(tmp_path, ["Abstract", "BACKGROUND: " + BODY, "Materials and methods: " + BODY,
                           "Conclusion: " + BODY, "Keywords: Cataract, glaucoma, retina", "Introduction", BODY])
    t = run(path, style.load("ni"), ["sections_abstract"])
    assert t[1].startswith("Background: ") and t[2].startswith("Methods: ")
    t = run(path, style.load("ipj"), ["sections_abstract"])  # IPJ lists "Materials and Methods:"
    assert t[2].startswith("Materials and Methods: ")
    t = run(path, style.load("kjo"), ["sections_abstract"])  # KJO: the guide's caps (G112, G114)
    assert t[1].startswith("BACKGROUND: ") and t[2].startswith("MATERIALS AND METHODS: ")


def test_word_limit_and_keyword_queries(tmp_path):
    long = " ".join(["word"] * 170) + "."
    path = make(tmp_path, ["Title", "Abstract", long, "Keywords: Alpha, beta, gamma, delta, epsilon, zeta",
                           "Introduction", BODY, "References"])
    ms = docmodel.load(path)
    checks = [q.check for q in queries.collect(ms, style.load("jehp")).queries]
    assert "abstract.too_long" in checks and "keywords.too_many" in checks  # S: 150 words; at most five keywords
    checks = [q.check for q in queries.collect(docmodel.load(path), style.load("ijo")).queries]
    assert "abstract.too_long" not in checks  # the IJO sheet has no word-limit line


def test_journal_specific_headings():
    st = style.load("lung")
    assert st.get("sections.abstract.heading") == "ABSTRACT" and st.get("sections.keywords.heading") == "KEY WORDS:"
    assert style.load("ijpm").get("sections.keywords.heading") == "KEYWORDS:"
    assert style.load("ni").get("front_matter.title_case") == "title"  # under "ni" and "neurol" alike
    assert style.load("neurol").get("front_matter.title_case") == "title"
