"""Output defects found by the defect scanner after the G2 fix batches (synthetic text only): tables imported from
another Word file in the journal's font, structured-abstract labels and the key words at the text size, the legends /
tables headings at the text size, queries off deleted text, back-matter statements not given twice, a table caption
with a line between it and its table, a typed label in a Word-numbered reference, an APA entry of an unlisted source."""

import io

from docx import Document
from docx.shared import Pt

from edit_os import docmodel, queries, refengine, refs_other, rules, rules_format, rules_sections, style
from edit_os.changes import Change, Layer, Source, Status

from test_legends import W, make
from test_structure_g2 import MANUSCRIPT, _accepted_run, _zip, journal, texts

FAMILY = "Times New Roman"


def _calibri_tables() -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_paragraph("Table 1: Socio-demographic characteristics of participants")
    t = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["Variable", "Percent"], ["Age", "79.9"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    run = t.cell(1, 0).paragraphs[0].runs[0]
    run.font.name = "Aptos"  # a direct font as well
    t.cell(1, 1).paragraphs[0].alignment = 1  # centred
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _sizes_fonts(path, text):
    ms = docmodel.load(str(path))
    fs = rules_format._FontStyles(ms)
    p = next(p for p in ms.paragraphs if p.text.strip() == text)
    f = rules_format.font_chars(p, fs)
    return {(x.get("font"), x.get("size")) for x, ch in zip(f, p.text) if not ch.isspace()}, p


def test_imported_tables_take_the_journal_font_size_and_cell_alignment(tmp_path):
    z = _zip(tmp_path, {"Table 1.docx": _calibri_tables()})
    got, out = _accepted_run(z, journal("jehp"), tmp_path)
    assert "Variable" in got
    for cell in ("Variable", "Age", "79.9"):
        fonts, p = _sizes_fonts(out.path("clean"), cell)
        assert fonts == {(FAMILY, 12.0)}, (cell, fonts)
    _, p = _sizes_fonts(out.path("clean"), "79.9")
    want = style.load("jehp").get("formatting.fonts")["table_align"]
    assert p.element.find(f"{W}pPr/{W}jc").get(f"{W}val") == want


def test_structured_abstract_labels_and_keywords_label_are_not_first_level_headings(tmp_path):
    paras = ["Original Article", "Abstract", "Background:", "Some words of the background here for the reader.",
             "Methods:", "Some words on the methods used in this study for the reader.", "Keywords",
             "Alpha, beta, gamma", "Introduction", "Body text of the introduction, long enough to be a paragraph.",
             "Methods", "Body text of the methods, long enough to be a paragraph of the article."]
    path = make(tmp_path, paras)
    ms = docmodel.load(path)
    levels = rules_format.heading_levels(ms, journal("jehp"))
    by_text = {p.text: p.id for p in ms.paragraphs}
    assert by_text["Background:"] not in levels and by_text["Methods:"] not in levels
    assert by_text["Keywords"] not in levels


def test_untagged_legends_and_tables_headings_are_not_first_level_in_a_tagging_journal(tmp_path):
    paras = ["Original Article", "Introduction", "Body text of the introduction, long enough to be a paragraph.",
             "Discussion", "Body text of the discussion, long enough to be a paragraph of the article.",
             "Figures", "Figure 1: A picture of the eye.", "Tables", "Table 1: The numbers."]
    ms = docmodel.load(make(tmp_path, paras))
    levels = rules_format.heading_levels(ms, journal("jehp"))
    by_text = {p.text: p.id for p in ms.paragraphs}
    assert by_text["Figures"] not in levels and by_text["Tables"] not in levels


def test_text_inserted_inside_a_font_change_takes_its_font(tmp_path):
    ms = docmodel.load(make(tmp_path, ["Keywords"]))
    p = ms.paragraphs[0]
    font = Change("f1", p.id, 0, 8, "Keywords", "Keywords", Layer.HOUSE_STYLE, Source.RULE, "fmt_font",
                  format={"size": 12, "font": FAMILY})
    rep = Change("k1", p.id, 0, 8, "Keywords", "Keywords:", Layer.STRUCTURE, Source.RULE, "sections_keywords",
                 format={"bold": True})
    ins = Change("k2", p.id, 8, 8, "", " alpha, beta", Layer.STRUCTURE, Source.RULE, "sections_keywords",
                 format={"bold": False})
    kept = rules_format.finish_fonts(ms, journal("jehp"), [rep, ins, font], [rep])
    assert ins in kept and ins.format == {"size": 12, "font": FAMILY, "bold": False}


def test_query_on_rewritten_text_takes_in_the_word_before(tmp_path):
    ms = docmodel.load(make(tmp_path, ["The author Nur5 wrote this."]))
    p = ms.paragraphs[0]
    s = p.text.index("Nur5")
    qs = queries.QuerySet()
    qs.queries.append(queries.Query(p.id, s, s + 3, "Kindly provide the full name.", "front.full_name"))
    c = Change("c1", p.id, s, s + 4, "Nur5", "Nur", Layer.STRUCTURE, Source.RULE, "front_matter")
    queries.keep_anchors_live(ms, qs, [c])
    q = qs.queries[0]
    assert p.text[q.start:q.end] == "author Nur5"


def test_query_on_text_run_on_elsewhere_follows_it(tmp_path):
    ms = docmodel.load(make(tmp_path, ["Figure 1: The eye.", "(A) Visual acuity (VA) after surgery."]))
    a, b = ms.paragraphs
    s = b.text.index("Visual")
    qs = queries.QuerySet()
    qs.queries.append(queries.Query(b.id, s, s + 18, "VA: defined twice.", "abbrev.defined_twice"))
    gone = Change("c1", b.id, 0, len(b.text), b.text, "", Layer.TABLES_FIGURES, Source.RULE, "fig_captions.run_on")
    moved = Change("c2", a.id, len(a.text), len(a.text), "", " " + b.text, Layer.TABLES_FIGURES, Source.RULE,
                   "fig_captions.run_on")
    queries.keep_anchors_live(ms, qs, [gone, moved])
    assert qs.queries[0].para_id == a.id


def test_query_on_a_deleted_reply_to_the_reviewers_goes_with_it(tmp_path):
    ms = docmodel.load(make(tmp_path, ["Reply to the reviewer, point seven of the reply.", "Body text."]))
    p = ms.paragraphs[0]
    qs = queries.QuerySet()
    qs.queries.append(queries.Query(p.id, 0, 0, "Separate rows?", "tables.cell_rows"))
    c = Change("c1", p.id, 0, len(p.text), p.text, "", Layer.STRUCTURE, Source.RULE, "submission_reviewer_reply")
    queries.keep_anchors_live(ms, qs, [c])
    assert qs.queries == []


def test_typed_label_in_a_word_numbered_reference_goes(tmp_path):
    from docx.oxml import parse_xml

    path = make(tmp_path, ["10. Smith AB, Jones CD. A study of the eye. Eye J 2020;1:1-5."])
    doc = Document(path)
    ppr = doc.paragraphs[0]._p.get_or_add_pPr()
    ppr.append(parse_xml('<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         '<w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'))
    doc.save(path)
    ms = docmodel.load(path)
    p = ms.paragraphs[0]
    if not p.numbered:  # no numbering definition in the test document: say Word numbers it
        p.numbered = True

    class Ref:
        unlabelled = False
        label = "10. "

    out = rules._label_changes(p, Ref(), journal("ijo"))
    assert len(out) == 1 and out[0].start == 0 and out[0].end == 4 and out[0].new == ""


def test_apa_entry_of_an_unlisted_source_is_styled_from_its_own_fields():
    st = journal("jehp")
    st.settings["references"].setdefault("other", {})["check_links"] = False
    text = ("Smith, A., Jones, B., Brown, C., Green, D., White, E., Black, F., & Grey, G. (2015). Deep tracing of "
            "learning. Advances in Neural Learning Systems, 28, 505–513.")
    res = refs_other.resolve(text, st, refengine.parse(text))
    assert res.formatted and res.formatted.startswith("Smith A, Jones B, Brown C, Green D, White E, Black F, et al.")
    assert ("references.unverified", {}) in res.queries


def test_title_page_statement_the_body_has_already_is_not_added_again():
    class P:
        def __init__(self, t):
            self.text = t

    class M:
        paragraphs = [P("The study was approved by the ethics committee of the university (ID 12/34) and all "
                        "participants gave consent.")]

    assert rules_sections._said_already(M(), "The study was approved by the ethics committee of the university "
                                             "(ID 12/34) and all participants gave consent.")
    assert not rules_sections._said_already(M(), "Funding was received from the council of the university.")


def test_back_matter_label_typed_twice_keeps_the_one_with_its_statement(tmp_path):
    paras = MANUSCRIPT[:3] + ["Conclusion", "The schools did well in the end."] + MANUSCRIPT[3:5] + ["Ethics approval and consent to participate: ",
                              "Ethics approval and consent to participate: The committee approved the study (ID 7)."] \
        + MANUSCRIPT[5:]
    z = _zip(tmp_path, {}, paras)
    got, _ = _accepted_run(z, journal("jehp"), tmp_path)
    heads = [t for t in got if "thics approval and consent" in t]
    assert len(heads) == 1, got


def test_real_funding_statement_beside_a_nil_one_under_the_journal_heading_is_not_given_twice(tmp_path):
    paras = MANUSCRIPT[:3] + ["Conclusion", "The schools did well in the end.", "Acknowledgement",
                              "We thank the schools.", "Funding:",
                              "This research was funded by the city council of the town (grant 12).",
                              "Finance support and sponsorship:", "The authors received no financial support.",
                              "Conflicts of interest:", "There are no conflicts of interest."] + MANUSCRIPT[7:]
    z = _zip(tmp_path, {}, paras)
    got, _ = _accepted_run(z, journal("jehp"), tmp_path)
    assert sum("funded by the city council" in t for t in got) == 1, got
    assert "The authors received no financial support." not in got
