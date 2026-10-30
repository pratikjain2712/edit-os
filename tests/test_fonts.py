"""Font family, font size and alignment (rules_format.rule_fonts, formatting.fonts) and their rendering."""

import copy

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from edit_os import docmodel, frontmatter, header, rules, rules_format, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status
from edit_os.tracked import PARAGRAPH_BREAK


def _ijo():
    st = style.load("ijo")
    st.settings = copy.deepcopy(st.settings)
    return st


def _doc(tmp_path, build):
    doc = Document()
    build(doc)
    path = str(tmp_path / "m.docx")
    doc.save(path)
    return path


def _body(doc):
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph()
    r = p.add_run("The eye was examined in detail and the pressure was ")
    r.font.size, r.font.name = Pt(10), "Arial"
    s = p.add_run("±")
    s.font.name = "Symbol"
    r = p.add_run(" 2 mmHg in both eyes of the patient.")
    r.font.size, r.font.name = Pt(10), "Arial"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph("Discussion")
    q = doc.add_paragraph("A second paragraph of running text that is long enough to be body text in the article.")
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def test_size_font_and_align_settings_are_on_for_ijo_jehp_and_mjdrdypu():
    for code in ("ijo", "jehp", "mjdrdypu"):  # MJDRDYPU: owner, 2026-10-26
        cfg = rules_format.fonts_settings(style.load(code))
        assert cfg and cfg["family"] == "Times New Roman" and cfg["size"] == 12 and cfg["heading_size"] == 14
        # JEHP justified like its editors (owner, 2026-10-30); IJO and MJDRDYPU left (G16)
        assert cfg["align"] == ("both" if code == "jehp" else "left") and cfg["running_head_align"] == "center"
    assert rules_format.fonts_settings(style.load("ijdr")) is None  # base: off until the owner confirms
    assert rules_format.insert_format(style.load("ijo"), "title") == {"size": 14, "font": "Times New Roman"}
    assert rules_format.insert_format(style.load("ijo"), "authors") == {"size": 12, "font": "Times New Roman"}
    assert rules_format.insert_align(style.load("ijo"), "running_head") == "center"
    assert rules_format.insert_format(style.load("ijdr")) is None


def test_change_format_values_are_checked():
    ok = dict(id="c", para_id="p", start=0, end=1, old="a", new="a", layer=Layer.HOUSE_STYLE, source=Source.RULE)
    Change(**ok, format={"size": 12, "font": "Times New Roman", "align": "left"})
    Change(**ok, format={"size": False, "font": False})
    for bad in ({"size": "12"}, {"size": True}, {"font": ""}, {"align": "middle"}, {"color": "FF0000"}):
        with pytest.raises(ValueError):
            Change(**ok, format=bad)


def test_rule_fonts_sizes_headings_body_and_keeps_symbol_font(tmp_path):
    ms = docmodel.load(_doc(tmp_path, _body))
    cs = rules_format.rule_fonts(ms, _ijo())
    by_para = {}
    for c in cs:
        by_para.setdefault(ms.by_id(c.para_id).text[:10], []).append(c)
    # G133: first-level headings 14 pt; G142: paragraphs 12 pt, left aligned
    assert by_para["Introducti"][0].format["size"] == 14
    body = by_para["The eye wa"]
    spans = [(c.start, c.end, c.format) for c in body]
    text = ms.by_id(body[0].para_id).text
    pm = text.index("±")
    assert all(not (s <= pm < e) for s, e, f in spans if "size" in f)  # the Symbol-font ± keeps its font
    assert {f["size"] for s, e, f in spans if "size" in f} == {12}
    assert {f.get("font") for s, e, f in spans if "size" in f} == {"Times New Roman"}
    assert any(f.get("align") == "left" for s, e, f in spans)  # justified -> left
    assert all(c.format_only for c in cs)


def test_rule_fonts_off_by_default():
    st = style.load("ijdr")
    assert rules_format.rule_fonts(None, st) == []


def test_rendered_sizes_fonts_and_alignment(tmp_path):
    path = _doc(tmp_path, _body)
    ms = docmodel.load(path)
    cs = ChangeSet(rules_format.rule_fonts(ms, _ijo()))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tracked.render(docmodel.load(path), cs, tracked=True)
    body = out.document.element.body
    p = [p for p in body.iter(qn("w:p")) if "examined" in docmodel.paragraph_text(p)][0]
    assert p.find(f"{qn('w:pPr')}/{qn('w:jc')}").get(qn("w:val")) == "left"
    ch = p.find(f"{qn('w:pPr')}/{qn('w:pPrChange')}")
    assert ch is not None and ch.find(f"{qn('w:pPr')}/{qn('w:jc')}").get(qn("w:val")) == "both"
    runs = docmodel.text_runs(p)
    sym = [r for r in runs if docmodel.run_text(r) == "±"][0]
    assert sym.find(f"{qn('w:rPr')}/{qn('w:rFonts')}").get(qn("w:ascii")) == "Symbol"
    for r in runs:
        if docmodel.run_text(r) != "±" and docmodel.run_text(r).strip():
            rpr = r.find(qn("w:rPr"))
            assert rpr.find(qn("w:sz")).get(qn("w:val")) == "24"
            assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "24"
            assert rpr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Times New Roman"
            assert rpr.find(qn("w:rPrChange")) is not None  # tracked formatting
    h = [p for p in body.iter(qn("w:p")) if docmodel.paragraph_text(p) == "Introduction"][0]
    assert docmodel.text_runs(h)[0].find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")) == "28"


def test_apply_format_sets_and_clears_size_and_font():
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    f.set(qn("w:asciiTheme"), "minorHAnsi")
    f.set(qn("w:eastAsia"), "SimSun")
    rpr.append(f)
    rpr.append(OxmlElement("w:b"))
    tracked.apply_format(rpr, {"size": 14, "font": "Times New Roman", "align": "left"})
    tags = [c.tag for c in rpr]
    assert tags == [qn("w:rFonts"), qn("w:b"), qn("w:sz"), qn("w:szCs")]  # schema order
    assert f.get(qn("w:asciiTheme")) is None and f.get(qn("w:eastAsia")) == "SimSun"
    assert f.get(qn("w:hAnsi")) == "Times New Roman" and rpr.find(qn("w:sz")).get(qn("w:val")) == "28"
    tracked.apply_format(rpr, {"size": False, "font": False})
    assert [c.tag for c in rpr] == [qn("w:b")]


def test_run_all_sizes_inserted_headings_and_carries_font_onto_replacements(tmp_path):
    def build(doc):
        _body(doc)
        doc.add_paragraph("Funding: none declared by the authors of this study.")
        for p in doc.paragraphs:
            for r in p.runs:
                if r.font.name != "Symbol":
                    r.font.size, r.font.name = Pt(10), "Arial"

    path = _doc(tmp_path, build)
    ms = docmodel.load(path)
    st = _ijo()
    keep = rules.RULES
    fake_texts = {"p0005": "Funding: none declared by the authors of this study."}

    def fake(ms_, st_):  # a section rule replacing a whole paragraph and opening an H1 heading after it
        p = [p for p in ms_.paragraphs if p.text == fake_texts["p0005"]][0]
        return [Change(id="x1", para_id=p.id, start=0, end=len(p.text), old=p.text,
                       new="Nil." + PARAGRAPH_BREAK + "<H1>Acknowledgements" + PARAGRAPH_BREAK + "None.",
                       layer=Layer.STRUCTURE, source=Source.RULE, reason="sections_back_matter",
                       segment_format=[None, {"bold": True}, None])]

    try:
        rules.RULES = {"sections_back_matter": fake, "format_fonts": rules_format.rule_fonts}
        cs = rules.run_all(ms, st)
    finally:
        rules.RULES = keep
    x = next(c for c in cs.changes if c.id == "x1")
    assert x.format == {"size": 12, "font": "Times New Roman"}  # the moot font change, carried over
    assert x.segment_format[1] == {"size": 14, "font": "Times New Roman", "bold": True}  # G146: 14 pt
    # the statement under the inserted heading: body size, not bold (rules_format.finish_headings)
    assert x.segment_format[2] == {"size": 12, "font": "Times New Roman", "bold": False}
    assert any(c.reason == "fmt_align" or "align" in (c.format or {}) for c in cs.changes
               if c.para_id == x.para_id) is False  # Normal style: already left
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tracked.render(docmodel.load(path), cs, tracked=False)
    ack = [p for p in out.document.element.body.iter(qn("w:p")) if docmodel.paragraph_text(p) == "<H1>Acknowledgements"][0]
    assert docmodel.text_runs(ack)[0].find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")) == "28"


def test_inserted_paragraphs_take_the_journal_font_size_and_alignment(tmp_path):
    path = _doc(tmp_path, _body)
    ms = docmodel.load(path)
    st = _ijo()
    first = ms.paragraphs[0].id
    inserts = [frontmatter.NewParagraph("i1", first, "before", [("<RH>Rao: A study", {})], "running_head",
                                        status=Status.ACCEPTED),
               frontmatter.NewParagraph("i2", first, "before", [("A study of the eye", {"bold": True})], "title",
                                        status=Status.ACCEPTED)]
    frontmatter.render_inserts(ms, inserts, tracked=True, run_format=lambda r: rules_format.insert_format(st, r),
                               align=lambda r: rules_format.insert_align(st, r))
    header.render(ms, [("p", "Technical Editing done on 1-January-2026:- A")], tracked=True, author="A",
                  date="2026-01-01T00:00:00Z", run_format=rules_format.insert_format(st),
                  align=rules_format.insert_align(st))
    out = tmp_path / "o.docx"
    ms.save(str(out))
    body = docmodel.load(str(out)).document.element.body
    got = {}
    for p in body.iter(qn("w:p")):
        t = docmodel.paragraph_text(p)
        if t.startswith(("<RH>", "A study", "Technical")):
            r = docmodel.text_runs(p)[0].find(qn("w:rPr"))
            got[t[:4]] = (p.find(f"{qn('w:pPr')}/{qn('w:jc')}").get(qn("w:val")),
                          r.find(qn("w:sz")).get(qn("w:val")), r.find(qn("w:rFonts")).get(qn("w:ascii")))
    assert got == {"<RH>": ("center", "24", "Times New Roman"), "A st": ("left", "28", "Times New Roman"),
                   "Tech": ("left", "24", "Times New Roman")}


def test_heading_italic_and_roman_levels(tmp_path):
    st = _ijo()
    st.settings["formatting"]["heading_italic_levels"] = [2]  # (IJO: H3; H2 here for a short test document)
    st.settings["formatting"]["heading_roman_levels"] = [1]

    def build(doc):
        doc.add_paragraph("Introduction")
        doc.add_paragraph("Some text about the study that is long enough to be a paragraph of running text. " * 3)
        doc.add_paragraph().add_run("Methods").italic = True
        doc.add_paragraph("Some text about the methods that is long enough to be a paragraph of running text. " * 3)
        doc.add_paragraph("Ethics approval")
        doc.add_paragraph("The study was approved by the institutional ethics committee of the hospital. " * 3)

    ms = docmodel.load(_doc(tmp_path, build))
    cs = rules_format.rule_cleanup(ms, st)
    got = {(ms.by_id(c.para_id).text, c.reason): c.format for c in cs if c.reason == "fmt_heading_italic"}
    assert got.get(("Methods", "fmt_heading_italic")) == {"italic": False}  # an H1: roman (S: "H1 | ... Roman")
    assert got.get(("Ethics approval", "fmt_heading_italic")) == {"italic": True}  # H2 as configured here
