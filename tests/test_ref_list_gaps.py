"""Reference-list gaps found on IJO articles: the list ends at the back matter after it, entries broken
over several paragraphs are one entry, the guide's citation form (no month, no "Pt 1"), IJO's "1. " labels
and italic "et al." in the list, Google Scholar chapters, renumbered entries. Synthetic references only;
PubMed, CrossRef and the NLM directory are patched or switched off."""

import pytest
from docx import Document

from edit_os import biblio, docmodel, journals, queries, references, refengine, refs_other, renumber, rules, style, tracked
from edit_os.changes import Status


def _style(journal="ijo"):
    st = style.load(journal)
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["back_matter"]["add_missing"] = False  # keeps paragraph positions fixed
    st.settings["sections"]["headings"]["drop"] = []
    return st


def _numbered(paragraph):
    """Make a python-docx paragraph a Word list item (w:numPr), as an auto-numbered reference list is."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        num.append(el)
    ppr.append(num)
    return paragraph


def _load(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def _apply(path, st, tmp_path, only=None):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")
    return [p.text for p in docmodel.load(str(out)).paragraphs], cs


def _ref(n):
    return f"Author{n} A, Writer B. Study number {n} of the eye. J Test 2001;{n}:1-2."


# --- 1. the list ends at the back matter after it; renumbering goes ahead ---------------------------

def test_consent_block_after_the_list_is_not_an_entry_and_renumbering_goes_ahead(tmp_path):
    path = _load(tmp_path, ["Introduction", "Alpha [2]. Beta [1].", "References",
                            "1. " + _ref(1), "2. " + _ref(2),
                            "Declaration of patient consent",
                            "The authors certify that they have obtained all appropriate patient consent forms."])
    ms = docmodel.load(path)
    assert [r.number for r in references.extract(ms)] == [1, 2]
    pl = renumber.plan(ms, _style())
    assert pl.apply and pl.mapping == {1: 2, 2: 1}
    t, _ = _apply(path, _style(), tmp_path)
    refs = [x for x in t if x[:3] in ("1. ", "2. ")]
    assert refs[0].startswith("1. Author2 A") and refs[1].startswith("2. Author1 A")
    # the consent block is not an entry; the back-matter rule moves it before the references (S392)
    assert t.index("<H2>Declaration of patient consent") < t.index("<H1>References")


def test_label_fix_does_not_travel_with_a_moved_entry(tmp_path):
    path = _load(tmp_path, ["Introduction", "Alpha [2]. Beta [1].", "References", "1." + _ref(1), "2. " + _ref(2)])
    t, _ = _apply(path, _style(), tmp_path)
    assert t[3] == "1. " + _ref(2) and t[4] == "2. " + _ref(1)


def test_unlabelled_heading_after_an_auto_numbered_list_ends_it(tmp_path):
    doc = Document()
    doc.add_paragraph("References")
    for n in (1, 2):
        _numbered(doc.add_paragraph(_ref(n)))
    doc.add_paragraph("Declaration")
    doc.add_paragraph("Patient’s consent has been obtained as per ICMJE guidelines.")
    doc.save(tmp_path / "a.docx")
    refs = references.extract(docmodel.load(str(tmp_path / "a.docx")))
    assert [r.number for r in refs] == [1, 2]


def test_list_out_of_sequence_is_queried_not_skipped_silently(tmp_path):
    path = _load(tmp_path, ["Introduction", "Alpha [2]. Beta [1]. Gamma [3].", "References",
                            "1. " + _ref(1), "3. " + _ref(2), "2. " + _ref(3)])
    ms = docmodel.load(path)
    pl = renumber.plan(ms, _style())
    assert not pl.apply and pl.blocked
    qs = queries.collect(ms, _style())
    assert any(q.check == "references.renumber_blocked" for q in qs.queries)


# --- 2. entries broken over several paragraphs ------------------------------------------------------

def test_broken_entries_are_joined_and_stray_labels_are_not_entries(tmp_path):
    doc = Document()
    doc.add_paragraph("References")
    lines = [("Gass JD. Diseases of the optic nerve that may simulate macular disease.", True),
             ("Trans Sect Ophthalmol Am Acad Ophthalmol Otolaryngol 1977; 83:763–770.", False),
             ("Shi J, Danesh-Meyer HV, Bhatti MT. Neuroretinitis: a comprehensive review. Eye (Lond).", True),
             ("2025 Feb;39:238-250.", False),
             ("6. ", False),
             ("Kahloun R, Khairallah-Ksiaa I, Abroug N, Mahmoud A, Ben Yahia S, Zaouali", True),
             ("S, et al. Final diagnosis in patients referred with a diagnosis of neuroretinitis.", False),
             ("Neuroophthalmology. 2015;39:266–70.", False),
             ("Choi SK, Byon IS, Kwon HJ, Park SW. Case series of neuroretinitis in Korea.", True),
             ("BMC Ophthalmol. 2024 Jan 18; 24:24.", False)]
    for text, numbered in lines:
        p = doc.add_paragraph(text)
        if numbered:
            _numbered(p)
    doc.save(tmp_path / "b.docx")
    refs = references.extract(docmodel.load(str(tmp_path / "b.docx")))
    assert [r.number for r in refs] == [1, 2, 3, 4]
    assert refs[0].full_text.endswith("Otolaryngol 1977; 83:763–770.")
    assert refs[1].full_text.endswith("Eye (Lond). 2025 Feb;39:238-250.")
    assert refs[2].full_text.endswith("Zaouali S, et al. Final diagnosis in patients referred with a diagnosis of "
                                      "neuroretinitis. Neuroophthalmology. 2015;39:266–70.")
    assert refs[3].full_text.endswith("Korea. BMC Ophthalmol. 2024 Jan 18; 24:24.")


def test_year_opening_a_line_is_not_a_label():
    assert not references._LABEL.match("2025 Feb;39:238-250.")
    assert not references._LABEL.match("2024 Apr 17; 16.")
    assert references._LABEL.match("12 Smith A. Title. J X 2001;1:2.")


def test_joined_entry_is_written_as_one_paragraph_in_the_guide_form(tmp_path):
    path = _load(tmp_path, ["References",
                            "1. Choi SK, Byon IS, Kwon HJ, Park SW. Case series of neuroretinitis in Korea.",
                            "BMC Ophthalmol. 2024 Jan 18; 24:24.",
                            "2. " + _ref(2)])
    st = _style()
    st.settings["references"]["journal_abbreviation"] = ""
    t, _ = _apply(path, st, tmp_path, only=["references"])
    assert t[1] == "1. Choi SK, Byon IS, Kwon HJ, Park SW. Case series of neuroretinitis in Korea. BMC Ophthalmol 2024;24:24."
    assert t[2] == "2. " + _ref(2)  # the emptied paragraph is gone


# --- 3-5. "et al." italic (IJO), labels, "Pt 1" -------------------------------------------------------

def test_et_al_italic_in_the_ijo_list_only(tmp_path):
    path = _load(tmp_path, ["References", "1. Okoro A, Lindqvist B, et al. Virtual reality. J Test 2020;12:40-6."])
    italic = [c for c in rules.run_all(docmodel.load(path), _style(), only=["references"]).changes
              if c.format and c.format.get("italic")]
    assert [c.old for c in italic] == ["et al"]
    italic = [c for c in rules.run_all(docmodel.load(path), _style("jehp"), only=["references"]).changes
              if c.format and c.format.get("italic")]
    assert not italic


def test_et_al_cut_to_six_authors_is_italic(tmp_path):
    names = ", ".join(f"Name{c} {c}" for c in "ABCDEFG")
    path = _load(tmp_path, ["References", f"1. {names}. Virtual reality. J Test 2020;12:40-6."])
    t, cs = _apply(path, _style(), tmp_path, only=["references"])
    assert t[1].startswith("1. NameA A, NameB B, NameC C, NameD D, NameE E, NameF F, et al. Virtual")
    assert any(c.new == "et al" and c.format == {"italic": True} for c in cs.changes)


def test_word_diff_puts_et_al_in_italics(tmp_path):
    path = _load(tmp_path, ["References", "1. A A, B B, C C, D D, E E, F F, G G. Title. J 2020;1:2-3."])
    p = docmodel.load(path).paragraphs[1]
    new = "A A, B B, C C, D D, E E, F F, et al. Title. J 2020;1:2-3."
    cs = rules.word_diff(p, 3, p.text[3:], new, "ref_lookup", et_al=True)
    assert any(c.new == "et al" and c.format == {"italic": True} for c in cs)
    text = p.text
    for c in sorted(cs, key=lambda c: (c.start, c.end), reverse=True):
        text = text[:c.start] + c.new + text[c.end:]
    assert text[3:] == new


@pytest.mark.parametrize("label, out", [("1.", "1. "), ("1)", "1. "), ("[1]", "1. "), ("1 ", "1. "), ("1.  ", "1. "),
                                        ("1.\t", "1.\t"), ("1)\t", "1.\t")])
def test_labels_written_as_number_period_space(tmp_path, label, out):
    path = _load(tmp_path, ["References", f"{label}Okoro A. Virtual reality. J Test 2020;12:40-6."])
    t, _ = _apply(path, _style(), tmp_path, only=["references"])
    assert t[1] == f"{out}Okoro A. Virtual reality. J Test 2020;12:40-6."


def test_part_goes_from_the_volume(tmp_path):
    path = _load(tmp_path, ["References", "1. Bai FY. Taxonomy of Candida. Int J Syst Evol Microbiol 2000;50 Pt 1:417-22.",
                            "2. Bai FY. Taxonomy of Candida. Int J Syst Evol Microbiol 2000;50(Pt 1):417-22.",
                            "3. Bai FY. Taxonomy of Candida. Int J Syst Evol Microbiol 2000;50(Suppl 1):417-22."])
    t, _ = _apply(path, _style(), tmp_path, only=["references"])
    assert t[1].endswith("2000;50:417-22.") and t[2].endswith("2000;50:417-22.")
    assert t[3].endswith("2000;50(Suppl 1):417-22.")  # guide: other content in the issue place stays
    r = biblio.Record(source="pubmed", title="Taxonomy of Candida.", authors=[("Bai", "FY")], journal="Int J Syst Evol Microbiol",
                      year="2000", volume="50 Pt 1", pages="417-422")
    assert refengine.format_record(r, _style()).endswith("2000;50:417-22.")


# --- 6. Google Scholar chapters -------------------------------------------------------------------

def test_scholar_chapter_in_the_edited_book_order():
    t = refs_other.template(
        "Shah R. Complications after SMILE. In: Small Incision Lenticule Extraction (SMILE) Principles and Future "
        "Concepts 2015 Aug 3 (pp. 97-105). Cham: Springer International Publishing.", "book_chapter")
    assert t.text == ("Shah R. Complications after SMILE. In: Small Incision Lenticule Extraction (SMILE) Principles and "
                      "Future Concepts. Cham: Springer International Publishing; 2015. p. 97-105.")
    t = refs_other.template("Ravishankar V, Jacob S. Corneal ring segments CAIRS. InKeratoconus: Current and Future "
                            "State-of-the-Art 2021 Nov 21 (pp. 167-175). Cham: Springer International Publishing.", "book")
    assert t.text == ("Ravishankar V, Jacob S. Corneal ring segments CAIRS. In: Keratoconus: Current and Future "
                      "State-of-the-Art. Cham: Springer International Publishing; 2021. p. 167-75.")


# --- 7. rebuilt entries: no double period, names, translated titles ---------------------------------

def test_moved_entry_takes_no_edit_of_the_entry_it_replaces(tmp_path, monkeypatch):
    """The entry rebuilt at position 1 adds a final "." to the author's; renumbering puts another entry
    there: the "." must not land after it ("2023;42:1247-56..")."""
    def resolve_all(texts, st):
        out = []
        for t in texts:
            p = refengine.parse(t)
            out.append(refengine.Resolution(refengine.FORMAT_FIXED, p, formatted=t.rstrip(".") + "."))
        return out
    monkeypatch.setattr(refengine, "resolve_all", resolve_all)
    st = _style()
    st.settings["references"]["lookup"] = True
    path = _load(tmp_path, ["Introduction", "Alpha [2]. Beta [1].", "References",
                            "1. Author1 A, Writer B. Study number 1 of the eye. J Test 2001;1:1-2",
                            "2. " + _ref(2)])
    t, _ = _apply(path, st, tmp_path, only=["renumber", "ref_lookup", "references"])
    assert t[3] == "1. " + _ref(2)
    assert t[4] == "2. " + _ref(1)


def test_brand_names_keep_their_capital():
    assert "Ziemer" in refengine.sentence_case("Lenticule Extraction Using the Ziemer FEMTO LDV Z8 Platform")


def test_transliterated_title_keeps_the_english_one_in_brackets():
    entry = ("Doga AV, Kostenev SV. Rezul'taty ekstraktsii lentikuly rogovitsy [Results of corneal lenticule extraction "
             "for correction]. Vestn Oftalmol. 2020;136:214-218.")
    r = biblio.Record(source="pubmed", title="[Results of corneal lenticule extraction for correction].",
                      authors=[("Doga", "AV"), ("Kostenev", "SV")], journal="Vestn Oftalmol", year="2020",
                      volume="136", pages="214-218")
    out = refengine.format_record(r, _style(), refengine.parse(entry))
    assert ("Rezul'taty ekstraktsii lentikuly rogovitsy [Results of corneal lenticule extraction for correction]. "
            "Vestn Oftalmol 2020;136:214-8.") in out


# --- 8. an abbreviated title found by its citation -------------------------------------------------

HOLEKAMP = ("Holekamp NM, de Beus A, Clark WL, Heier JS. Prospective trial of home OCT guided management of "
            "treatment-experienced nAMD patients. Retina 2024;44(5).")
HOLEKAMP_REC = biblio.Record(
    source="pubmed", title="PROSPECTIVE TRIAL OF HOME OPTICAL COHERENCE TOMOGRAPHY-GUIDED MANAGEMENT OF TREATMENT "
                           "EXPERIENCED NEOVASCULAR AGE-RELATED MACULAR DEGENERATION PATIENTS.",
    authors=[("Holekamp", "NM"), ("de Beus", "AM"), ("Clark", "WL"), ("Heier", "JS")], journal="Retina",
    year="2024", volume="44", issue="10", pages="1714-1731", pmid="39287534")


def test_acronyms_written_out_from_the_record():
    out, n = refengine.expand_acronyms("home OCT guided management of nAMD patients", HOLEKAMP_REC.title)
    assert n == 2 and "OPTICAL COHERENCE TOMOGRAPHY" in out and "NEOVASCULAR AGE-RELATED MACULAR DEGENERATION" in out


def test_abbreviated_title_is_accepted_only_with_author_year_journal_and_volume(monkeypatch):
    monkeypatch.setattr(journals, "directory", lambda: None)
    p = refengine.parse(HOLEKAMP)
    assert refengine.abbreviated_title_matches(p, HOLEKAMP_REC)
    other = biblio.Record(**{**HOLEKAMP_REC.__dict__, "volume": "43"})
    assert not refengine.abbreviated_title_matches(p, other)
    other = biblio.Record(**{**HOLEKAMP_REC.__dict__, "authors": [("Wykoff", "CC")]})
    assert not refengine.abbreviated_title_matches(p, other)
    other = biblio.Record(**{**HOLEKAMP_REC.__dict__, "title": "Home OCT in diabetic macular edema."})
    assert not refengine.abbreviated_title_matches(p, other)


def test_abbreviated_title_found_by_author_year_volume(monkeypatch):
    terms = []
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: terms.append(term) or
                        (["39287534"] if "[vi]" in term else []))
    monkeypatch.setattr(biblio, "pubmed_best_match", lambda term, retmax=10: [])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [HOLEKAMP_REC] if "39287534" in ids else [])
    monkeypatch.setattr(biblio, "crossref_search", lambda text, rows=5: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    monkeypatch.setattr(journals, "directory", lambda: None)
    res = refengine.resolve(HOLEKAMP, _style())
    assert res.status == refengine.FORMAT_FIXED and res.formatted.endswith("Retina 2024;44:1714-31.")
    assert any(t == "Holekamp[au] AND 2024[dp] AND 44[vi]" for t in terms)
