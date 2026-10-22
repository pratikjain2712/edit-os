"""Submission files API: listing with kinds, raw files, viewable forms, path safety, auth. Synthetic files only."""

import io
import json
import zipfile

import pytest
from docx import Document
from PIL import Image

from server import submission_files as sf
from test_api import KEY, client, wait  # noqa: F401  (fixture)
from test_pipeline import make_docx

XML = b"""<?xml version="1.0" encoding="UTF-8"?>
<manuscript>
  <id>test_1_26</id><title>A synthetic case of something rare</title><articletype>Case Report</articletype>
  <wordCount>1200</wordCount><keywords>rare case, imaging; pathology</keywords><doi>10.0000/test.1</doi>
  <firstpagefile>FirstPage.docx</firstpagefile><qr_code>qr.png</qr_code>
  <authors>
    <author><title>Dr</title><fname>Asha</fname><mname></mname><lname>Rao</lname><email>a@example.org</email>
      <department>Radiology</department><institute>City Hospital</institute><city>Pune</city><state>MH</state>
      <country>India</country><is_corresponding>Y</is_corresponding><address>Ward 4</address></author>
    <author><title>Dr</title><fname>Ben</fname><mname>K</mname><lname>Lee</lname><email>b@example.org</email>
      <department>Pathology</department><institute>City Hospital</institute><city>Pune</city><state></state>
      <country>India</country><is_corresponding>N</is_corresponding></author>
  </authors>
  <images><image><image_file_name>Figure 1.png</image_file_name><image_type>Figures</image_type>
    <legend>Chest X-ray</legend></image></images>
  <submission_date>13-Jul-26</submission_date><Accepted>06-Oct-26</Accepted>
</manuscript>"""

META = {"article_id": "test_1_26", "journal": "ijo", "files": {"manuscript_file": "Manuscript.docx", "fp_file": "missing"},
        "authors": [{"name": "Dr BEN K LEE", "email": "b@example.org", "author_order": "2", "corresponding_author": False},
                    {"name": "Dr ASHA RAO", "email": "a@example.org", "author_order": "Primary", "corresponding_author": True}]}


def _png(fmt="PNG", size=(40, 30)):
    b = io.BytesIO()
    Image.new("RGB", size, (200, 30, 30)).save(b, fmt)
    return b.getvalue()


def _first_page(path):
    d = Document()
    d.add_paragraph("A synthetic case of something rare")
    d.add_paragraph("Asha Rao, Ben K Lee")
    d.save(path)
    return path


@pytest.fixture
def subzip(tmp_path):
    z = tmp_path / "original.zip"
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, "w") as nz:
        nz.writestr("inner.txt", "never opened")
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(make_docx(tmp_path / "Manuscript.docx"), "Manuscript.docx")
        zf.write(_first_page(tmp_path / "FirstPage.docx"), "FirstPage.docx")
        zf.writestr("Figure 1.png", _png())
        zf.writestr("figure 2.tif", _png("TIFF"))
        zf.writestr("Figure 3.pdf", b"%PDF-1.4\n%stub\n")
        zf.writestr("test_1_26.xml", XML)
        zf.writestr("qr.png", _png(size=(8, 8)))
        zf.writestr("article_metadata.json", json.dumps(META))
        zf.writestr("Cover Letter.docx", (tmp_path / "Manuscript.docx").read_bytes())
        zf.writestr("Author Contribution Form.pdf", b"%PDF-1.4\n")
        zf.writestr("Table 1.docx", (tmp_path / "Manuscript.docx").read_bytes())
        zf.writestr("extra.zip", nested.getvalue())
        zf.writestr("page.html", "<script>alert(1)</script>")
        zf.writestr("../evil.txt", "outside")
        zf.writestr("__MACOSX/._Figure 1.png", b"x")
        zf.writestr("figs/", b"")
    return z


def _job(client, path):
    with open(path, "rb") as f:
        r = client.post("/api/jobs", headers=KEY, data={"journal": "ijo"},
                        files={"file": (path.name, f, "application/octet-stream")})
    assert r.status_code == 202, r.text
    job = wait(client, r.json()["id"])
    assert job["status"] == "done", job["error"]
    return r.json()


def test_kinds_and_order(subzip):
    files = sf.list_files(str(subzip))
    kinds = {e["path"]: e["kind"] for e in files}
    assert kinds == {
        "Manuscript.docx": "manuscript", "FirstPage.docx": "title_page", "Figure 1.png": "figure",
        "figure 2.tif": "figure", "Figure 3.pdf": "figure", "test_1_26.xml": "xml", "qr.png": "other",
        "article_metadata.json": "metadata", "Cover Letter.docx": "cover_letter",
        "Author Contribution Form.pdf": "form", "Table 1.docx": "table", "extra.zip": "other", "page.html": "other"}
    assert [e["kind"] for e in files][:5] == ["manuscript", "title_page", "figure", "figure", "figure"]
    assert [e["label"] for e in files if e["kind"] == "figure"] == ["Figure 1", "Figure 2", "Figure 3"]
    assert next(e for e in files if e["path"] == "qr.png")["label"] == "QR code"
    assert sum(e["manuscript"] for e in files) == 1


def test_listing_raw_and_views(client, subzip):
    job = _job(client, subzip)
    base = f"/api/jobs/{job['id']}/submission"
    r = client.get(base, headers=KEY)
    assert r.status_code == 200
    files = {e["path"]: e for e in r.json()["files"]}
    assert files["Figure 1.png"]["view"] == "image" and files["Figure 1.png"]["size"] == len(_png())
    assert files["FirstPage.docx"]["view"] == "docx" and files["test_1_26.xml"]["view"] == "metadata"
    assert files["extra.zip"]["view"] == "download"

    r = client.get(f"{base}/raw/Figure 1.png", headers=KEY)
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content == _png()
    assert r.headers["content-disposition"].startswith("inline")
    r = client.get(f"{base}/raw/figure 2.tif", headers=KEY)  # TIFF converted to PNG
    assert r.headers["content-type"] == "image/png" and r.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert Image.open(io.BytesIO(r.content)).size == (40, 30)
    r = client.get(f"{base}/raw/figure 2.tif", headers=KEY, params={"original": "true"})
    assert r.headers["content-type"] == "image/tiff" and r.headers["content-disposition"].startswith("attachment")
    r = client.get(f"{base}/raw/Figure 3.pdf", headers=KEY)
    assert r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")
    r = client.get(f"{base}/raw/extra.zip", headers=KEY)  # nested archive: a plain download
    assert r.headers["content-disposition"].startswith("attachment") and zipfile.is_zipfile(io.BytesIO(r.content))
    r = client.get(f"{base}/raw/page.html", headers=KEY)  # never served as a page
    assert r.headers["content-type"] == "application/octet-stream" and r.headers["content-security-policy"] == "sandbox"
    assert r.headers["content-disposition"].startswith("attachment")

    v = client.get(f"{base}/view/FirstPage.docx", headers=KEY).json()
    assert v["view"] == "docx" and [p["text"] for p in v["paragraphs"]][:2] == [
        "A synthetic case of something rare", "Asha Rao, Ben K Lee"]
    assert {"id", "kind", "runs"} <= set(v["paragraphs"][0])

    card = client.get(f"{base}/view/test_1_26.xml", headers=KEY).json()["card"]
    assert card["article_type"] == "Case Report" and card["title"] == "A synthetic case of something rare"
    assert card["keywords"] == ["rare case", "imaging", "pathology"] and card["doi"] == "10.0000/test.1"
    assert [a["name"] for a in card["authors"]] == ["Asha Rao", "Ben K Lee"]
    assert card["authors"][0]["affiliation"] == "Radiology, City Hospital, Pune, MH, India"
    assert card["corresponding"]["name"] == "Asha Rao" and card["corresponding"]["email"] == "a@example.org"
    assert card["dates"] == [{"label": "Received", "date": "2026-07-13"}, {"label": "Accepted", "date": "2026-10-06"}]
    assert card["figures"] == [{"file": "Figure 1.png", "type": "Figures", "legend": "Chest X-ray"}]

    card = client.get(f"{base}/view/article_metadata.json", headers=KEY).json()["card"]
    assert [a["name"] for a in card["authors"]] == ["Asha Rao", "Ben K Lee"]  # by author_order ("Primary" first), case fixed
    assert card["corresponding"]["name"] == "Asha Rao" and card["authors"][0]["prefix"] == "Dr"
    assert {"label": "Title page file", "value": "missing"} not in card["extra"]

    assert client.get(f"{base}/view/Figure 3.pdf", headers=KEY).json()["view"] == "pdf"


@pytest.mark.parametrize("path", ["../evil.txt", "..%2Fevil.txt", "%2E%2E/%2E%2E/jobs.sqlite", "/etc/passwd",
                                  "figs%2F..%2FFigure 1.png", "__MACOSX/._Figure 1.png", "extra.zip/inner.txt",
                                  "Manuscript.docx/word/document.xml", "nope.png", "figure 2.TIF"])
def test_path_traversal_and_unknown_names_refused(client, subzip, path):
    job = _job(client, subzip)
    for kind in ("raw", "view"):
        r = client.get(f"/api/jobs/{job['id']}/submission/{kind}/{path}", headers=KEY)
        assert r.status_code == 404, (kind, path, r.status_code)


def test_auth(client, subzip):
    job = _job(client, subzip)
    base = f"/api/jobs/{job['id']}/submission"
    for url in (base, f"{base}/raw/Figure 1.png", f"{base}/view/FirstPage.docx"):
        assert client.get(url).status_code == 401
        assert client.get(url, headers={"X-API-Key": "wrong"}).status_code == 401
        assert client.get(url, headers={"X-Review-Token": "wrong"}).status_code == 401
        assert client.get(url, headers={"X-Review-Token": job["review_token"]}).status_code == 200
        assert client.get(url, params={"token": job["review_token"]}).status_code == 200
    other = _job(client, subzip)  # one job's token does not open another job's files
    assert client.get(f"/api/jobs/{other['id']}/submission", headers={"X-Review-Token": job["review_token"]}).status_code == 401


def test_plain_docx_job_and_downloads_still_work(client, tmp_path):
    path = tmp_path / "ms.docx"
    make_docx(path)
    job = _job(client, path)
    files = client.get(f"/api/jobs/{job['id']}/submission", headers=KEY).json()["files"]
    assert [(e["path"], e["kind"], e["manuscript"]) for e in files] == [("ms.docx", "manuscript", True)]
    assert client.get(f"/api/jobs/{job['id']}/submission/raw/ms.docx", headers=KEY).content == path.read_bytes()
    assert client.get(f"/api/jobs/{job['id']}/files/tracked", headers=KEY).status_code == 200  # old route intact
    assert client.get(f"/api/jobs/{job['id']}/files/raw", headers=KEY).status_code == 404


def test_metadata_card_tolerates_junk():
    sf.metadata_card(b"not xml at all <<<", ".xml")  # recover mode: never raises
    assert sf.metadata_card(b"[1, 2]", ".json") is None
    assert sf.metadata_card(b"\xff\xfe", ".json") is None
    xxe = b'<?xml version="1.0"?><!DOCTYPE m [<!ENTITY x SYSTEM "file:///etc/passwd">]><manuscript><title>&x;</title></manuscript>'
    assert "root" not in sf.metadata_card(xxe, ".xml")["title"]


def test_xml_image_role_does_not_make_documents_figures():
    from server import submission_files as sf
    roles = {"authors contribution file.docx": "figure", "Cover Letter.docx": "figure", "Reply_template.doc": "figure",
             "Figure 1.png": "figure"}
    kinds = {n: sf._kind(n, roles, "Manuscript.docx") for n in roles}
    assert kinds == {"authors contribution file.docx": "form", "Cover Letter.docx": "cover_letter",
                     "Reply_template.doc": "other", "Figure 1.png": "figure"}
