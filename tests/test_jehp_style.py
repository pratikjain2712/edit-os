"""JEHP style fixes (synthetic text): key word casing, back-matter statement heads, the reference list heading,
table caption sentence case, numeric range dashes next to citations, bold tagged headings, author names from
the submission XML."""

from docx import Document

from edit_os import docmodel, frontmatter as fm, rules, rules_format, rules_guide, rules_sections, style, tracked
from edit_os.changes import Status

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail.")


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for text in paras:
        if isinstance(text, tuple):
            doc.add_paragraph().add_run(text[0]).bold = True
        else:
            doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def jehp():
    st = style.load("jehp")
    st.settings["references"]["lookup"] = False
    return st


def run(path, st, only):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    ms = tracked.render(docmodel.load(path), cs, tracked=False)
    out = str(path).replace(".docx", "_out.docx")
    ms.save(out)
    return [p.text for p in docmodel.load(out).paragraphs]


ARTICLE = ["Introduction", BODY, "Results", BODY, "Discussion", BODY, "Conclusion", BODY]


def test_keyword_case_inside_hyphenated_words():
    body = "We studied game-based learning and self-efficacy in a cross-sectional design. COVID-19 was ongoing."
    assert rules_sections._keyword_case("Game-Based Learning", body) == "game-based learning"
    assert rules_sections._keyword_case("Cross-Sectional Studies", body) == "cross-sectional studies"
    assert rules_sections._keyword_case("K-Means Clustering", body) == "K-means clustering"
    assert rules_sections._keyword_case("Non-ST Elevation", body) == "non-ST elevation"
    assert rules_sections._keyword_case("COVID-19", body) == "COVID-19"


def test_back_matter_statement_heads_tagged(tmp_path):
    path = make(tmp_path, ARTICLE + [
        "Conflicts of interest", "There are no conflicts of interest.",
        ("Data availability",), "The data are available from the corresponding author on request.",
        ("Consent for publication",), "All participants agreed to the publication of the results.",
        ("Declarations",), "Not applicable here at all.",
        "References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    out = run(path, jehp(), ["heading_tags", "subheadings"])
    assert "<H2>Data availability" in out and "<H2>Consent for publication" in out
    assert "<H2>Conflicts of interest" in out and "Declarations" in out  # a container head is not tagged
    assert not any(t.startswith("<H2><H2>") for t in out)
    base = style.load("base")
    base.settings["references"]["lookup"] = False
    assert "Data availability" in run(path, base, ["heading_tags", "subheadings"])  # off in base


def test_reference_list_heading_renamed(tmp_path):
    path = make(tmp_path, ARTICLE + ["Reference", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    out = run(path, jehp(), ["heading_tags", "sections_headings"])
    assert "<H1>References" in out


def test_table_caption_sentence_case(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " The Wilcoxon test was used [Table 1].",
                           "Table 1. Within-Group Comparison of Pain Scores (Wilcoxon Signed-Rank Test)",
                           "Table 2: comparison of scores: path coefficients by Group", BODY,
                           "References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    out = run(path, jehp(), ["fig_captions"])
    assert "Table 1: Within-group comparison of pain scores (Wilcoxon Signed-Rank Test)." in out
    assert "Table 2: Comparison of scores: path coefficients by group." in out
    base = style.load("base")
    base.settings["references"]["lookup"] = False
    assert any("Within-Group Comparison" in t for t in run(path, base, ["fig_captions"]))


def test_range_dash_leaves_round_bracket_citations():
    cfg = {"range_dash": True}
    text = "aged 46-55 years as reported earlier (11-13) and (23, 29-31) with scores (1.01-1.06)."
    got = [(text[s:e], new) for s, e, new, _ in rules_guide.propose(text, cfg)]
    assert got == [("-", "–"), ("-", "–")]
    assert text[text.index("46"):].startswith("46-55")


def test_tagged_headings_bold(tmp_path):
    path = make(tmp_path, ["Introduction", BODY, "Materials and Methods", "Study design and setting", BODY + BODY,
                           "References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    st = jehp()
    got = {(c.old, c.reason) for c in rules_format.rule_cleanup(docmodel.load(path), st) if c.format == {"bold": True}}
    assert ("Study design and setting", "fmt_heading_bold") in got
    st.settings["formatting"]["heading_bold_tagged"] = False
    got = {c.old for c in rules_format.rule_cleanup(docmodel.load(path), st) if c.format == {"bold": True}}
    assert "Study design and setting" not in got


XML = b"""<?xml version="1.0" encoding="UTF-8"?><manuscript><title>A study</title>
<articletype>Original Article</articletype><authors>
<author><fname>Amina</fname><mname>Amina</mname><lname>Benali</lname><is_corresponding>Y</is_corresponding>
<department>Psychiatry</department><institute>City University</institute><city>Agadir</city></author>
<author><fname>Omar</fname><mname> </mname><lname>Haddad</lname><is_corresponding>N</is_corresponding>
<department>Psychiatry</department><institute>City University</institute><city>Agadir</city></author>
<author><fname>Sara</fname><mname> </mname><lname>Idrissi</lname><is_corresponding>N</is_corresponding>
<department>Psychiatry</department><institute>City University</institute><city>Agadir</city></author>
</authors></manuscript>"""


def test_xml_middle_name_repeating_the_first():
    assert fm.parse_xml(XML).authors[0]["given"] == "Amina"


def test_single_name_author_kept():
    raw = {"authors": [{"given_names": "Sonali", "surname": "", "affiliations": [1]},
                       {"given_names": "Indu", "surname": "Rao", "affiliations": [1]}]}
    f = fm.verify(raw, "Sonali, Indu Rao", {})
    assert [(a.given, a.surname) for a in f.authors] == [("", "Sonali"), ("Indu", "Rao")]
    assert fm.format_name(f.authors[0], {})[0] == "Sonali"


def test_xml_authors_when_the_text_names_only_some(tmp_path, monkeypatch):
    raw = {"title": "A study",
           "authors": [{"given_names": "Omar", "surname": "Haddad", "affiliations": []},
                       {"given_names": "Sara", "surname": "Idrissi", "affiliations": []}]}
    monkeypatch.setattr(fm, "split", lambda text: raw)
    path = make(tmp_path, ["A study", "Dear Editor, Omar Haddad and Sara Idrissi submit this study.", "Abstract",
                           BODY, "Introduction", BODY])
    ms = docmodel.load(path)
    src = fm.Sources(meta=fm.parse_xml(XML))
    names = [b.text for b in fm.plan(ms, jehp(), src).blocks if b.role == "authors"]
    assert names and names[0].startswith("Amina Benali") and "Sara Idrissi" in names[0]


def test_lowercase_table_citation_capitalised(tmp_path):
    path = make(tmp_path, ["Introduction", BODY + " The results are presented in table 2 and figure 1.",
                           "References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."])
    out = run(path, jehp(), ["fig_citations"])
    assert any("presented in Table 2 and Figure 1." in t for t in out)
    base = style.load("base")
    base.settings["references"]["lookup"] = False
    assert any("presented in table 2" in t for t in run(path, base, ["fig_citations"]))


def test_table_footnote_final_period_deleted(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph(BODY + " See Table 1.")
    doc.add_paragraph("Table 1: Characteristics of the participants.")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Age"
    doc.add_paragraph("*P < 0.05.")
    doc.add_paragraph("SD, standard deviation; SE, standard error.")
    doc.add_paragraph(BODY)
    path = tmp_path / "t.docx"
    doc.save(path)
    out = run(str(path), jehp(), ["fig_captions"])
    assert "SD, standard deviation; SE, standard error" in out and "*P < 0.05." in out and BODY in out
