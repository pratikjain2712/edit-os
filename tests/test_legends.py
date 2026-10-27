"""Figure legends (part labels, parts run on, legends at the end under a heading), table captions in their table's
first row, table heads and notes (synthetic only)."""

import io
import zipfile

from docx import Document
from docx.shared import Inches

from edit_os import docmodel, figures_insert, legends, pipeline, rules, style, tracked
from edit_os.changes import Status

from test_figures_insert import picture

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ijo(**figures):
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    st.settings["figures"].update(figures)
    return st


def make(tmp_path, paras, name="m.docx"):
    """paras: text, ("bold", text), ("table", rows), ("picture", None)."""
    doc = Document()
    for item in paras:
        if isinstance(item, str):
            doc.add_paragraph(item)
        elif item[0] == "bold":
            doc.add_paragraph().add_run(item[1]).bold = True
        elif item[0] == "table":
            rows = item[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
        elif item[0] == "picture":
            doc.add_paragraph().add_run().add_picture(io.BytesIO(picture()), width=Inches(1))
    path = tmp_path / name
    doc.save(path)
    return str(path)


def apply(path, st, tmp_path, only=("fig_captions",)):
    cs = rules.run_all(docmodel.load(path), st, only=list(only))
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(tmp_path / "tracked.docx")
    return [p.text for p in docmodel.load(str(out)).paragraphs if p.text.strip()], cs


REFS = ["References", "1. Lee K. Back pain in adolescents. Spine 1997;22:1380-4."]


def test_parts_on_their_own_lines_run_on_into_the_legend(tmp_path):
    # S / G264: part labels and content in separate paragraphs -> run on; G263: "(a)" in round brackets.
    path = make(tmp_path, ["Results", "Seen at [Fig. 1] and [Fig. 2] and [Fig. 3].", *REFS, "Figure Legends",
                           "Figure 1    a. Pre-operative map", "        b. Post-operative OCT ",
                           "        c. Comparative map showing preoperative and 1 week  ",
                           "              postoperative maps",
                           "Fig.2 A Face photograph in primary gaze", "",
                           "Fig.2 B Face photograph in lateral view.",
                           "Figure 3: Fundus photographs due to A. Toxoplasmosis B. Tuberculosis C. Viral &",
                           "D. Bartonella."])
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[5:] == [
        "Figure 1: (a) Pre-operative map. (b) Post-operative OCT. (c) Comparative map showing preoperative and 1 week "
        "postoperative maps.",  # IJO: captions end with a period (S; owner, 2026-10-26)
        "Figure 2: (a) Face photograph in primary gaze. (b) Face photograph in lateral view.",
        "Figure 3: Fundus photographs due to (a) Toxoplasmosis (b) Tuberculosis (c) Viral & (d) Bartonella."]


def test_without_run_on_parts_keep_their_lines(tmp_path):
    path = make(tmp_path, [*REFS, "Figure 1A. Slit-lamp photograph.", "Figure 1B. B-scan, et al.",
                           "Figure 2 a. Pre-operative map", "   b. Post-operative OCT"])
    texts, _ = apply(path, ijo(legend_run_on=False), tmp_path)
    assert texts[2:] == ["Figure 1: (a) Slit-lamp photograph.", "(b) B-scan, et al.",
                         "Figure 2: (a) Pre-operative map", "   (b) Post-operative OCT."]


def test_unbracketed_labels_only_in_order_and_in_pairs(tmp_path):
    path = make(tmp_path, [*REFS, "Figure 1. MR results. A: Causal scatter plot. B: Funnel plot.",
                           "Figure 2. Rich in vitamin A. Shown here.",
                           "Figure 3. Case 1. A–C: right eye; D–F: left eye.",
                           "Figure 4 A shows the fundus."])  # a sentence, not a legend
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[2] == "Figure 1: MR results. (a): Causal scatter plot. (b): Funnel plot."
    assert texts[3] == "Figure 2: Rich in vitamin A. Shown here."
    assert texts[4] == "Figure 3: Case 1. (a–c): right eye; (d–f): left eye."
    assert texts[5] == "Figure 4 A shows the fundus."


def test_bold_part_labels_set_plain(tmp_path):
    doc = Document()
    for t in REFS:
        doc.add_paragraph(t)
    p = doc.add_paragraph()
    p.add_run("Figure 1: ").bold = True
    p.add_run("Photographs ")
    p.add_run("(A)").bold = True
    p.add_run(" right eye ")
    p.add_run("(B)").bold = True
    p.add_run(" left eye")
    doc.save(tmp_path / "b.docx")
    texts, cs = apply(str(tmp_path / "b.docx"), ijo(), tmp_path)
    assert texts[-1] == "Figure 1: Photographs (a) right eye (b) left eye."
    out = docmodel.load(str(tmp_path / "out.docx"))
    runs = [r for r in out.paragraphs[-1].element.iter(W + "r")]
    bold = {"".join(t.text for t in r.iter(W + "t")): r.find(f"{W}rPr/{W}b") is not None
            and r.find(f"{W}rPr/{W}b").get(f"{W}val") not in ("0", "false") for r in runs}
    assert not any(v for k, v in bold.items() if "(a)" in k or "(b)" in k)


def test_unnumbered_legends_become_figure_1(tmp_path):
    path = make(tmp_path, ["See the [Figure].", *REFS, ("bold", "Legends to Figures"),
                           "Figure (A, B): Slit-lamp photographs. ", "(C) Anterior segment OCT.", "(D) B-scan."])
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[3:] == ["Figure Legend",
                         "Figure 1: (a, b): Slit-lamp photographs. (c) Anterior segment OCT. (d) B-scan."]
    path = make(tmp_path, ["Case report text.", *REFS, ("bold", "Image description"),
                           "A 7-year-old girl with lesions. (A) Lower lid. (B) Upper lid."], "u2.docx")
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[3:] == ["Figure Legend", "Figure 1: A 7-year-old girl with lesions. (a) Lower lid. (b) Upper lid."]


def test_legends_heading_renamed(tmp_path):
    path = make(tmp_path, [*REFS, ("bold", "Legends:"), "Figure 1: One", "Figure 2: Two"])
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[2] == "Figure Legends"
    assert apply(path, ijo(legends_heading=None), tmp_path)[0][2] == "Legends:"


def test_all_caps_column_heads_title_case(tmp_path):
    path = make(tmp_path, ["After LASIK the BCVA improved.", "Table 1: Outcomes",
                           ("table", [["PARAMETER", "1 WEEK POSTOPERATIVE", "P VALUE", "SD", "LASIK OUTCOME"],
                                      ["BCVA", "0.1", "0.2", "0.3", "GOOD"]])])
    texts, _ = apply(path, ijo(), tmp_path)
    assert texts[2:7] == ["Parameter", "1 Week Postoperative", "P Value", "SD", "LASIK Outcome"]
    assert texts[7:] == ["BCVA", "0.1", "0.2", "0.3", "GOOD"]  # body rows as typed


def test_table_notes_comma_between_abbreviations_and_no_end_period(tmp_path):
    path = make(tmp_path, ["Table 1: Outcomes", ("table", [["A", "B"], ["1", "2"]]),
                           "CDVA = corrected distance visual acuity; SD = standard deviation; D = diopters.",
                           "Table 2: More", ("table", [["A", "B"], ["1", "2"]]),
                           "Note: Model 0: unadjusted; Model 1: adjusted for age, and sex."])
    texts, _ = apply(path, ijo(), tmp_path)
    # IJO: "term = definition" (owner, 2026-10-27: follow the guide, G211; replaces 2026-10-26 "term: definition")
    assert "CDVA = corrected distance visual acuity, SD = standard deviation, D = diopters" in texts
    assert "Note: Model 0: unadjusted; Model 1: adjusted for age, and sex" in texts  # not abbreviations: as typed


def _ijo_zip(tmp_path, paras, files=None, xml="<article/>"):
    ms = make(tmp_path, paras, "Manuscript.docx")
    z = tmp_path / "sub.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(ms, "ART_1/Manuscript.docx")
        zf.writestr("ART_1/ART_1.xml", xml)
        for name, data in (files or {}).items():
            zf.writestr(f"ART_1/{name}", data)
    return str(z)


def _body(path):
    out = []
    for el in Document(path).element.body:
        if el.tag == W + "tbl":
            first = el.find(f"{W}tr")
            out.append("TABLE: " + "".join(t.text or "" for t in first.iter(W + "t")))
        elif el.tag == W + "p":
            t = "".join(x.text or "" for x in el.iter(W + "t"))
            pic = el.find(f".//{W}drawing") is not None
            if t.strip() or pic:
                out.append(("[picture] " if pic else "") + t)
    return out


BODY = ["Original Article", "Introduction", "We saw it [Figure 1]. Then [Table 1] and [Figure 2].",
        ("picture", None), "Figure 1. Clinical photograph.", "Results", "More text.",
        "Figure 2: Scan of the orbit", "Discussion", "It grew.", *REFS,
        "Table 1. Baseline data", ("table", [["Age", "Sex"], ["30", "M"]])]


def test_legends_move_to_the_end_with_their_pictures_and_a_heading(tmp_path):
    out = pipeline.edit(_ijo_zip(tmp_path, BODY), ijo(), str(tmp_path / "out"), stages=("style",), accept_all=True)
    body = _body(out.path("clean"))
    i = next(k for k, t in enumerate(body) if "Lee K" in t)
    assert body[i + 1:] == ["First author:", "Figure Legends", "[picture] ", "Figure 1: Clinical photograph.",
                            "Figure 2: Scan of the orbit.",
                            "TABLE: Table 1: Baseline data."]
    assert "Figure 2: Scan of the orbit" not in body[:i]
    man, _, inserts, _ = pipeline.load_records(str(tmp_path / "out"))
    moves = [n for n in inserts if legends.is_move(n)]
    assert {m.kind for m in moves} == {"legends", "table_caption"}
    # The tracked copy: deleted where they were, inserted at the end; rejected: nothing moves.
    tr = Document(out.path("tracked")).element.body
    dels = [p for p in tr.iter(W + "p") if p.find(f"{W}pPr/{W}rPr/{W}del") is not None]
    assert any("Figure 2" in "".join(t.text or "" for t in p.iter(W + "delText")) for p in dels)
    for n in inserts:
        n.status = Status.REJECTED if legends.is_move(n) else n.status
    pipeline.save_records(str(tmp_path / "out"), inserts=inserts)
    again = pipeline.render_outputs(str(tmp_path / "out"), style=ijo())
    body = _body(again.path("clean"))
    assert body.index("Figure 2: Scan of the orbit.") < next(k for k, t in enumerate(body) if "Discussion" in t)
    assert "Table 1: Baseline data." in body
    # records round-trip
    d = moves[0].to_dict()
    assert pipeline.insert_from_dict(d).to_dict() == d


def test_first_author_block_goes_before_the_tables(tmp_path):
    path = make(tmp_path, ["Text [Table 1].", *REFS, "Table 1: Data", ("table", [["A", "B"], ["1", "2"]])])
    ms = docmodel.load(path)
    anchor, where = legends.back_slot(ms)
    assert (ms.by_id(anchor).text, where) == ("Table 1: Data", "before")
    path = make(tmp_path, ["Text [Table 1].", *REFS, "TABLES", "Table 1: Data", ("table", [["A", "B"], ["1", "2"]])],
                "h.docx")
    ms = docmodel.load(path)
    anchor, where = legends.back_slot(ms)
    assert (ms.by_id(anchor).text, where) == ("TABLES", "before")  # a heading over the tables starts them


def test_table_caption_below_its_table_goes_into_the_first_row(tmp_path):
    path = make(tmp_path, ["Text [Table 1] and [Table 2].", ("table", [["A", "B"], ["1", "2"]]), "Table 1: Below.",
                           "", ("table", [["C", "D"], ["3", "4"]]), "Table 2: Also below"])
    ms = docmodel.load(path)
    moves, _ = legends.plan(ms, ijo())
    assert [m.side for m in moves] == ["below", "below"]
    for m in moves:
        m.status = Status.ACCEPTED
    for is_tracked in (False, True):
        ms = docmodel.load(path)
        legends.render(ms, moves, (Status.ACCEPTED,), tracked=is_tracked, author="Edit OS", date="2026-10-25T00:00:00Z")
        ms.save(tmp_path / "t.docx")
        tables = Document(tmp_path / "t.docx").element.body.findall(f"{W}tbl")
        first = ["".join(t.text or "" for t in tb.find(f"{W}tr").iter(W + "t")) for tb in tables]
        assert first == ["Table 1: Below.", "Table 2: Also below"]
        span = tables[0].find(f"{W}tr/{W}tc/{W}tcPr/{W}gridSpan")
        assert span is not None and span.get(f"{W}val") == "2"
        tr0 = tables[0].find(f"{W}tr")
        assert (tr0.find(f"{W}trPr/{W}ins") is not None) == is_tracked


XML = """<manuscript><images>
<image><image_file_name>figure 1.png</image_file_name><image_type>Figures</image_type>
<legend>Figure 1 (A) Slit-lamp photograph. (B) At 1 week. </legend></image>
<image><image_file_name>figure 2.png</image_file_name><image_type>Figures</image_type><legend> </legend></image>
</images></manuscript>"""


def test_legends_from_the_xml_and_a_legends_file(tmp_path):
    paras = ["Introduction", "See [Figure 1], [Figure 2] and [Figure 3].", *REFS]
    leg = make(tmp_path, [("bold", "LEGENDS"), "Figure 3. Anterior segment OCT", "  a. Preoperative.",
                          "  b. Postoperative."], "LEGENDS.docx")
    files = {"figure 1.png": picture(seed=2), "figure 2.png": picture(seed=3), "figure 3.png": picture(seed=4)}
    z = _ijo_zip(tmp_path, paras, files, XML)
    got = figures_insert.supplied_legends(z, ijo())
    assert got[1] == ("(a) Slit-lamp photograph. (b) At 1 week.", "ART_1.xml")
    assert 2 not in got
    with zipfile.ZipFile(z, "a") as zf:
        zf.write(leg, "ART_1/LEGENDS.docx")
    got = figures_insert.supplied_legends(z, ijo())
    assert got[3] == ("Anterior segment OCT. (a) Preoperative. (b) Postoperative.", "LEGENDS.docx")
    out = pipeline.edit(z, ijo(), str(tmp_path / "out"), stages=("style",), accept_all=True)
    body = _body(out.path("clean"))
    i = body.index("Figure Legends")
    assert body[i + 1:] == ["[picture] ", "Figure 1: (a) Slit-lamp photograph. (b) At 1 week.", "[picture] ",
                            "Figure 2: ???", "[picture] ",
                            "Figure 3: Anterior segment OCT. (a) Preoperative. (b) Postoperative."]
    assert out.summary["queries"]["by_check"].get("figures.legend_missing") == 1
