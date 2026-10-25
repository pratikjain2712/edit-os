"""A record replaces an entry only when it is the entry's paper: first author, title, year (and, when
the entry gives them, journal, volume and pages) agree; Claude's fields are checked against the entry's
own text; a "doi:..." line is part of the entry above it; broken author lists are never written out.
Synthetic references only; PubMed, CrossRef, Claude and the NLM directory are patched."""

import pytest
from docx import Document

from edit_os import ai_refs, biblio, docmodel, journals, queries, references, refengine, rules, style, tracked
from edit_os.changes import Status


def rec(title, authors=(("Okoro", "A"),), journal="J Midwif Educ", year="2020", volume="12", pages="40-6", **kw):
    return biblio.Record(source=kw.pop("source", "pubmed"), title=title, authors=list(authors), journal=journal,
                         year=year, volume=volume, pages=pages, pmid=kw.pop("pmid", "1"), **kw)


ENTRY = ("Okoro A, Lindqvist B, Haddad C. Virtual reality in midwifery training: a review of outcomes. "
         "J Midwif Educ. 2020;12(3):40-6.")


@pytest.fixture
def db(monkeypatch):
    """PubMed returns `records` for every search, in order; CrossRef finds nothing."""
    records = []
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: [r.pmid for r in records])
    monkeypatch.setattr(biblio, "pubmed_best_match", lambda term, retmax=10: [])
    monkeypatch.setattr(biblio, "pubmed_fetch", lambda ids: [r for r in records if r.pmid in ids])
    monkeypatch.setattr(biblio, "crossref_search", lambda text, rows=5: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    monkeypatch.setattr(journals, "directory", lambda: None)
    return records


# --- acceptance ----------------------------------------------------------------------------------

def test_shorter_title_by_another_first_author_is_a_query(db):
    # The record's title is contained in the entry and the year agrees, but it is another author's paper.
    db.append(rec("Virtual reality in midwifery training", authors=[("Tanaka", "M")], pmid="7"))
    res = refengine.resolve(ENTRY, style.load("ijo"))
    assert res.status == refengine.QUERY and res.record is None


def test_shorter_title_with_author_and_year_is_taken(db):
    db.append(rec("Virtual reality in midwifery training: a review of outcomes.", pmid="7"))
    res = refengine.resolve(ENTRY, style.load("ijo"))
    assert res.status == refengine.FORMAT_FIXED and res.record.pmid == "7"


def test_year_one_apart_is_an_epub_two_apart_is_another_paper(db):
    p = refengine.parse(ENTRY)
    assert refengine.accept(p, rec("Virtual reality in midwifery training: a review of outcomes.", year="2021"))
    assert not refengine.accept(p, rec("Virtual reality in midwifery training: a review of outcomes.", year="2022"))


def test_same_title_in_another_journal_is_rejected_and_the_search_goes_on(db):
    # A statement printed in two journals: the entry's own volume and pages say which one.
    reprint = rec("Virtual reality in midwifery training: a review of outcomes.", journal="Nurse Educ Today",
                  volume="88", pages="101-7", pmid="8")
    right = rec("Virtual reality in midwifery training: a review of outcomes.", pmid="9")
    db.extend([reprint, right])
    assert refengine.citation_conflicts(ENTRY, reprint) and not refengine.citation_conflicts(ENTRY, right)
    res = refengine.resolve(ENTRY, style.load("ijo"))
    assert res.record.pmid == "9"


def test_journal_named_by_the_entry_must_be_the_records(db, monkeypatch):
    class Dir:
        J = {"Am J Prev Med": journals.Journal("1", "American journal of preventive medicine", "Am J Prev Med",
                                               "Am J Prev Med", "", ""),
             "J Gen Intern Med": journals.Journal("2", "Journal of general internal medicine", "J Gen Intern Med",
                                                  "J Gen Intern Med", "", "")}

        def lookup(self, name):
            return self.J.get(name) or next((j for j in self.J.values() if j.title.lower() == name.lower()), None)

    monkeypatch.setattr(journals, "directory", lambda: Dir())
    text = "Okoro A, Haddad C. Counselling smokers in hospital: a trial. American Journal of Preventive Medicine. 2023."
    p = refengine.parse(text)
    other = rec("Counselling smokers in hospital: a trial.", journal="J Gen Intern Med", year="2023")
    same = rec("Counselling smokers in hospital: a trial.", journal="Am J Prev Med", year="2023")
    assert refengine.journal_conflicts(p, other) and not refengine.accept(p, other)
    assert not refengine.journal_conflicts(p, same) and refengine.accept(p, same)


@pytest.mark.parametrize("record, entry, same", [
    ("Jørgensen", "Jorgensen OD", True),       # letters NFKD does not split
    ("Yılmaz", "Yilmaz M", True),
    ("Rüb", "Rüb U", True),
    ("De Potter", "De Potter P", True),
    ("van der Berg", "Berg JH", True),         # particles
    ("Abdel-Latif", "Abdel Latif M", True),    # hyphen
    ("Lucas Espeza", "Lucas L", True),         # one part of a compound surname
    ("Du", "Du Z", True),                      # a surname that is also a particle
    ("RIFE", "Rife DC", True),                 # an all-capitals record
    ("Roddy Mitchell", "A. Roddy Mitchell", True),
    ("Maiwald", "Maiwaid M", True),            # one-letter slip in a long name
    ("Larsen", "Moreau P", False),
    ("Zhu", "Xin Z", False),
])
def test_same_surname(record, entry, same):
    assert refengine.same_surname(record, entry) is same


def test_group_authors_and_acronyms():
    p = refengine.Parsed(text="ADA. Standards of care. 2024.", authors=["ADA"], title="Standards of care", year="2024")
    r = rec("Standards of care", authors=[("American Diabetes Association Professional Practice Committee", "")])
    assert refengine.author_agrees(p, r)
    # PubMed may list the group before the people; the entry's first person still has to be the record's.
    r = rec("x", authors=[("Vision Loss Expert Group", ""), ("Okoro", "A")])
    assert refengine.author_agrees(refengine.parse(ENTRY), r)


def test_entry_without_authors_needs_whole_title_year_and_journal(db):
    text = "Virtual reality in midwifery training: a review of outcomes. J Midwif Educ. 2020;12(3):40-6."
    p = refengine.Parsed(text=text, title="Virtual reality in midwifery training: a review of outcomes", year="2020")
    assert refengine.accept(p, rec("Virtual reality in midwifery training: a review of outcomes."))
    assert not refengine.accept(p, rec("Virtual reality in midwifery training: a review of outcomes.",
                                       journal="Nurse Educ Today"))
    assert not refengine.accept(p, rec("Virtual reality in midwifery training"))  # not the whole title


def test_journal_abbreviation_matches_the_full_name():
    r = rec("x", journal="Int J Gynaecol Obstet")
    assert refengine.journal_agrees("T. International Journal of Gynecology & Obstetrics. 2025;171:920-3.", r)
    assert not refengine.journal_agrees("T. Lancet. 2025;171:920-3.", r)


def test_doi_of_a_paper_by_another_first_author_is_queried(db, monkeypatch):
    db.append(rec("Virtual reality in midwifery training: a review of outcomes.", authors=[("Segaert", "S")],
                  doi="10.1000/vr.1"))
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: ["1"] if "[aid]" in term else [])
    res = refengine.resolve(ENTRY + " doi:10.1000/vr.1", style.load("ijo"))
    assert res.status == refengine.QUERY and refengine.DIFFERS in res.message and "first author" in res.message


def test_record_differs_query_text(db, monkeypatch, tmp_path):
    db.append(rec("Virtual reality in midwifery training: a review of outcomes.", authors=[("Segaert", "S")],
                  doi="10.1000/vr.1"))
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: ["1"] if "[aid]" in term else [])
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph("1. " + ENTRY + " doi:10.1000/vr.1")
    doc.save(tmp_path / "m.docx")
    qs = queries.collect(docmodel.load(str(tmp_path / "m.docx")), style.load("ijo"))
    assert any(q.text == "The first author given here differs from the published article. Kindly check."
               for q in qs.queries)


# --- Claude's fields -----------------------------------------------------------------------------

def test_fields_of_another_entry_are_not_used():
    neighbour = {"kind": "journal_article", "authors": ["Tanaka M", "Ruiz P"], "year": "2017",
                 "title": "Sutures versus staples for skin closure", "doi": "", "pmid": ""}
    assert ai_refs.check_fields(ENTRY, neighbour) is None
    p = refengine.from_fields(ENTRY, neighbour)
    assert p.first_surname == "Okoro" and p.year == "2020" and p.title.startswith("Virtual reality")


def test_each_field_must_be_in_the_entry():
    f = {"kind": "journal_article", "authors": ["Okoro A", "Lindqvist B"], "year": "2019",
         "title": "Virtual reality in midwifery training: a review of outcomes", "journal": "Nurse Educ Today"}
    checked = ai_refs.check_fields(ENTRY, f)
    assert checked["title"] and checked["authors"] == ["Okoro A", "Lindqvist B"]
    assert checked["year"] == "" and checked["journal"] == ""  # not in the entry: the regex parse decides
    # An author list borrowed from a neighbour may share one short name, not the list.
    assert not ai_refs.check_fields("Li B, Okoro A. Title of it. J X. 2020;1:2.",
                                     {"authors": ["Li W", "Tanaka M", "Ruiz P"], "title": "Title of it"})["authors"]
    # Text whose spaces were lost still carries Claude's title.
    assert ai_refs.check_fields("OkoroA.Virtualrealityinmidwiferytraining.JMidwifEduc.2020;12:40.",
                                {"authors": [], "title": "Virtual reality in midwifery training"})["title"]


def test_entry_that_starts_with_its_title_has_no_authors():
    text = "Virtual reality in midwifery training: a review of outcomes. J Midwif Educ. 2020;12(3):40-6."
    assert refengine.parse(text).authors  # the regex takes title words for names
    p = refengine.from_fields(text, {"kind": "journal_article", "authors": [], "year": "2020",
                                     "title": "Virtual reality in midwifery training: a review of outcomes"})
    assert p.authors == []


def test_parse_batch_realigns_answers_numbered_off_by_one(monkeypatch, tmp_path):
    texts = [ENTRY, "Tanaka M, Ruiz P. Sutures versus staples for skin closure. Surg Today. 2017;4:1-9."]
    # Claude skipped entry 1 and numbered entry 2's fields as 1.
    answer = [{"index": 1, "kind": "journal_article", "authors": ["Tanaka M", "Ruiz P"], "year": "2017",
               "title": "Sutures versus staples for skin closure", "doi": "", "pmid": ""}]
    monkeypatch.setattr(ai_refs, "_call", lambda numbered: answer)
    monkeypatch.setattr(ai_refs, "_cache", ai_refs._Cache(str(tmp_path / "c.sqlite")))
    out = ai_refs.parse_batch(texts)
    assert out[0] is None and out[1]["title"] == "Sutures versus staples for skin closure"
    assert ai_refs._cache.get(texts[0]) is None and ai_refs._cache.get(texts[1])


# --- splitting -----------------------------------------------------------------------------------

def _doc(tmp_path, *paras):
    doc = Document()
    doc.add_paragraph("References")
    for t in paras:
        doc.add_paragraph(t)
    doc.save(tmp_path / "m.docx")
    return docmodel.load(str(tmp_path / "m.docx"))


def test_doi_line_and_broken_entry_join_the_entry_above(tmp_path):
    ms = _doc(tmp_path,
              "1. Okoro A, Lindqvist B. Virtual reality in midwifery training. J Midwif Educ. 2020;12:40-6.",
              "doi:10.1000/vr.1",
              "2. Tanaka M, Ruiz P. Sutures versus staples for skin closure. Expert Rev Surg",
              "Med. 2017;4(2):1-9.",
              "3. Haddad C, Okoro A. Simulation in obstetrics. Lancet. 2019;3:4-5.",
              "Conflicts of interest: None, 2024")
    refs = references.extract(ms)
    assert [r.number for r in refs] == [1, 2, 3]  # the note after the list ends it
    assert refs[0].full_text.endswith("2020;12:40-6. doi:10.1000/vr.1")
    assert refs[1].full_text.endswith("Expert Rev Surg Med. 2017;4(2):1-9.")
    assert not refs[2].continuation  # a note after a complete entry stays apart


def test_next_entry_is_not_taken_for_a_continuation():
    prev = "Okoro A, Lindqvist B. Virtual reality in midwifery training. J Midwif Educ"
    assert not references.continues(prev, "Tanaka M, Ruiz P. Sutures versus staples. Surg Today. 2017;4:1-9.")
    assert not references.continues(prev, "Tanaka, M., & Ruiz, P. (2017). Sutures versus staples. Surg Today, 4, 1-9.")
    assert not references.continues(prev, "van der Berg JH. Staples. Surg Today. 2017;4:1-9.")
    assert references.continues(prev, "2020;12:40-6.")
    assert references.continues("Okoro A, Lindqvist B, Haddad C.", "Virtual reality in midwifery training")


def test_rebuilt_entry_takes_the_doi_line_with_it(db, tmp_path):
    db.append(rec("Virtual reality in midwifery training: a review of outcomes.",
                  authors=[("Okoro", "A"), ("Lindqvist", "B"), ("Haddad", "C")]))
    path = tmp_path / "m.docx"
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph("1. " + ENTRY)
    doc.add_paragraph("https://doi.org/10.1000/vr.1")
    doc.add_paragraph("2. Tanaka M. Sutures versus staples. Surg Today. 2017;4:1-9.")
    doc.save(path)
    cs = rules.run_all(docmodel.load(str(path)), style.load("ijo"), only=["ref_lookup"])
    assert any(c.para_id == "p0003" and c.old == "https://doi.org/10.1000/vr.1" and c.new == "" for c in cs.changes)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "o.docx"
    tracked.render(docmodel.load(str(path)), cs, tracked=False).save(out)
    texts = [p.text for p in docmodel.load(str(out)).paragraphs]
    assert texts[1].endswith("Virtual reality in midwifery training: A review of outcomes. J Midwif Educ 2020;12:40-6.")
    assert texts[2].startswith("2. Tanaka M.")


# --- author names from records -------------------------------------------------------------------

@pytest.mark.parametrize("surname, initials", [("B", "S"), ("S.B.", ""), ("P S", "M"), ("AP", "AS"),
                                               ("Sulistyorini,", "Y"), ("", "R"), ("Saima Zeb", ""),
                                               ("Department of Surgery, City University, Lagos", "")])
def test_malformed_record_names(surname, initials):
    assert refengine._bad_name(surname, initials)


@pytest.mark.parametrize("surname, initials", [("Li", "B"), ("Ng", "K"), ("De Potter", "P"),
                                               ("GBD 2019 Collaborators", ""), ("WHO Expert Consultation", "")])
def test_real_names_are_not_malformed(surname, initials):
    assert not refengine._bad_name(surname, initials)


def test_malformed_record_keeps_the_authors_names_and_asks(db):
    db.append(rec("Virtual reality in midwifery training: a review of outcomes.",
                  authors=[("Okoro", "A"), ("B", "S"), ("R", "S")], source="crossref"))
    res = refengine.resolve(ENTRY, style.load("ijo"))
    assert res.status == refengine.FORMAT_FIXED
    assert res.formatted.startswith("Okoro A, Lindqvist B, Haddad C. ")
    assert ("references.author_names", {}) in res.queries


def test_record_without_authors_is_written_from_its_title():
    r = rec("Clinical guidelines on obesity in adults.", authors=[])
    assert refengine.format_record(r, style.load("ijo")).startswith("Clinical guidelines on obesity in adults.")
