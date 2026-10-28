"""Fewer, merged author queries (owner, 2026-10-28: no query the editors would not write), synthetic text: the same
ask once, one query per reference, the uncited references and the uncited tables and figures each in one query.
Nothing the stylesheet asks for is dropped."""

from docx import Document

from edit_os import docmodel, qc, queries, style
from edit_os.copyedit import abbrev
from edit_os.queries import Query, QuerySet

from test_jehp_citations import _refs, jehp

BODY = "The scores were compared between the two groups with the appropriate tests for each outcome measure."


def _doc(tmp_path, paras):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = str(tmp_path / "m.docx")
    doc.save(path)
    return docmodel.load(path)


def test_uncited_references_are_one_query(tmp_path):
    ms = _doc(tmp_path, ["Introduction", "Scores rose [1].", "Methods", BODY]
              + _refs(*[f"Alpha B, Gamma D. Study number {n}. J Test 2020;{n}:1-9." for n in range(1, 7)]))
    st = jehp()
    qs = queries.condense(ms, st, queries.collect(ms, st))
    unc = [q for q in qs.queries if q.check == "references.uncited"]
    assert [q.text for q in unc] == ["Reference nos. 2–6 have not been cited. Kindly cite."]
    assert ms.by_id(unc[0].para_id).text.startswith("2. Alpha B")  # at the first uncited entry
    assert qc._queried_refs(unc) == ({2, 3, 4, 5, 6}, set())


def test_number_list_and_uncited_numbers():
    assert queries.number_list([19, 24, 25, 27, 28, 29]) == "19, 24, 25 and 27–29"
    assert queries.uncited_numbers("Reference nos. 19, 24, 25 and 27–29 have not been cited. Kindly cite.") == \
        [19, 24, 25, 27, 28, 29]
    assert queries.uncited_numbers("Reference no. 7 has not been cited. Kindly cite.") == [7]


def test_one_query_per_reference(tmp_path):
    ms = _doc(tmp_path, ["Introduction", "Scores rose [1,2].", "Methods", BODY]
              + _refs("Alpha B, Gamma D. Study one. J Test 2020;1:1-9.",
                      "Delta E. A handbook of tests. 2nd ed. 2019. Available from: http://example.org/x."))
    r1, r2 = [p for p in ms.paragraphs if p.text[:2] in ("1.", "2.")]
    qs = QuerySet()
    span1 = dict(start=3, end=len(r1.text))
    qs.add(r1.id, "references.not_found", **span1)
    qs.add(r1.id, "references.author_style", **span1)  # covered by "check and provide complete details"
    qs.add(r1.id, "references.no_pages", **span1)
    span2 = dict(start=3, end=len(r2.text))
    qs.add(r2.id, "references.link_dead", **span2)
    qs.add(r2.id, "references.access_date", **span2)
    qs.add(r2.id, "references.publisher_location", **span2)
    queries.condense(ms, jehp(), qs)
    assert [(q.para_id, q.check, q.text) for q in qs.queries] == [
        (r1.id, "references.not_found", queries.TEXT["references.not_found"]),
        (r2.id, "references.link_dead", "Link is not working. Kindly check. Kindly provide last accessed date and "
                                        "publisher location."),
    ]


def test_uncited_entry_keeps_its_ask_first_in_the_merged_query(tmp_path):
    ms = _doc(tmp_path, ["Introduction", BODY] + _refs("Alpha B. Study one. J Test 2020;1:1-9."))
    r1 = next(p for p in ms.paragraphs if p.text.startswith("1."))
    qs = QuerySet()
    qs.add(r1.id, "references.not_found", 3, len(r1.text))
    qs.add(r1.id, "references.uncited", 3, len(r1.text), n=1)
    queries.condense(ms, jehp(), qs)
    assert len(qs.queries) == 1 and qs.queries[0].check == "references.uncited"
    assert qs.queries[0].text.startswith("Reference no. 1 has not been cited. Kindly cite. ")
    assert qc._queried_refs(qs.queries) == ({1}, set())


def test_uncited_tables_and_figures_are_one_query(tmp_path):
    ms = _doc(tmp_path, ["Introduction", BODY, "Table 2: A", "Table 4: B", "Table 5: C", "Figure 2: D", "Legend"])
    ids = {p.text: p.id for p in ms.paragraphs}
    qs = QuerySet()
    for cap in ("Table 2: A", "Table 4: B", "Table 5: C", "Figure 2: D"):
        kind, n = cap.split(":")[0].split()
        qs.add(ids[cap], "figures.uncited", 0, 0, kind=kind, n=int(n))
    qs.add(ids["Legend"], "figures.uncited", 0, 0, kind="Figure", n=2)  # the same figure, from its supplied file
    queries.condense(ms, jehp(), qs)
    assert [(q.para_id, q.text) for q in qs.queries] == [
        (ids["Table 2: A"], "Tables 2, 4 and 5 and Figure 2 have not been cited in the text part. Kindly cite.")]


def test_condense_can_be_switched_off(tmp_path):
    ms = _doc(tmp_path, ["Introduction", BODY, "Table 2: A", "Table 4: B"])
    st = jehp()
    st.settings.setdefault("queries", {})["condense"] = False
    qs = QuerySet()
    for p in ms.paragraphs[2:]:
        qs.add(p.id, "figures.uncited", 0, 0, kind="Table", n=int(p.text[6]))
    assert len(queries.condense(ms, st, qs).queries) == 2


def test_abbreviation_defined_again_is_asked_once(tmp_path):
    ms = _doc(tmp_path, ["Abstract", "Background: The mean arterial pressure (MAP) rose.",
                         "Introduction",
                         "The mean arterial pressure (MAP) was measured. MAP rose.",
                         "Later the mean arterial pressure (MAP) fell. MAP fell.",
                         "Again the mean arterial pressure (MAP) was low. MAP was low.",
                         "Discussion", "MAP matters. " + BODY])
    st = style.load("ijo")
    _, qs = abbrev.collect(ms, st)
    twice = [q for q in qs if q.check == "abbrev.defined_twice"]
    assert len(twice) == 1 and ms.by_id(twice[0].para_id).text.startswith("Later")
    st.settings["copyedit"]["abbreviations"]["ask_once"] = False
    _, qs = abbrev.collect(ms, st)
    assert len([q for q in qs if q.check == "abbrev.defined_twice"]) == 2


def test_capital_heading_on_its_own_line_is_no_abbreviation(tmp_path):
    ms = _doc(tmp_path, ["Introduction", BODY, "DISCUSSION \nTo the best of our knowledge the XQZ score is new."])
    _, qs = abbrev.collect(ms, style.load("ijo"))
    assert [q.text for q in qs if q.check == "abbrev.undefined"] == ["XQZ: Kindly expand the abbreviation."]


def test_table_of_studies_by_author_is_not_asked_for_expansions(tmp_path):
    from edit_os import rules_figures

    doc = Document()
    for t in ["Introduction", "The studies are listed in Table 1 [Table 1]."]:
        doc.add_paragraph(t)
    doc.add_paragraph("Table 1: Studies reviewed")
    rows = [["Study", "Design", "Group"], ["QUILLAN CM", "Trial", "XQR"], ["DORSETT EM, 2014, PEDIATRICS", "Cohort", "A"]]
    t = doc.add_table(rows=3, cols=3)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.rows[r].cells[c].paragraphs[0].add_run(val)
    path = str(tmp_path / "m.docx")
    doc.save(path)
    ms = docmodel.load(path)
    asked = [fmt["abbrs"] for *_, check, fmt in rules_figures.footnote_queries(ms, jehp())
             if check.startswith("figures.footnote_abbreviations")]
    assert asked == ["XQR"]
