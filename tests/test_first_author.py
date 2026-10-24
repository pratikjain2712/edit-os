"""The "First author" block after the references (edit_os/first_author.py, sections.back_matter.first_author)."""

import json
import os
import zipfile

import pytest
from docx import Document

from edit_os import docmodel, first_author, frontmatter as fm, pipeline, rule_sources, style
from edit_os.changes import Layer, Source, Status

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

XML = b"""<?xml version="1.0" encoding="UTF-8"?><manuscript>
<title>Effect of Valsalva maneuver on intraocular pressure</title>
<articletype>Original Article</articletype>
<authors>
<author><title>Dr</title><fname>Asha</fname><mname>Lakshmi</mname><lname>Kumar</lname><email>asha@example.org</email>
<department>Ophthalmology</department><institute>City Eye Hospital</institute><city>Chennai</city>
<state>Tamil Nadu</state><country> </country><is_corresponding>Y</is_corresponding></author>
<author><title>Mr</title><fname>Ravi</fname><mname></mname><lname>Rao</lname><email>ravi@example.org</email>
<department>Physiology</department><institute>Erasmus University</institute><city>Rotterdam</city>
<state></state><country>Netherlands</country><is_corresponding>N</is_corresponding></author>
</authors></manuscript>"""

BODY = [
    "Original Article",
    "Introduction",
    "The lesion is shown in Figure 1.",
    "Discussion",
    "It grew fast.",
    "References",
    "1. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4.",
    "2. Rao S. Eye pain. Eye. 2001;3:1-4.",
    "Figure Legends",
    "Figure 1: Clinical photograph of the lesion",
]


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


@pytest.fixture
def ijo():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    return st


def plan(tmp_path, ijo, paras=BODY, xml=XML):
    ms = docmodel.load(make(tmp_path, paras))
    ms.front_sources = fm.Sources(meta=fm.parse_xml(xml), xml_name="ART.xml")
    return ms, first_author.plan(ms, ijo, fm.plan(ms, ijo))


def test_setting_is_off_in_base_and_on_for_ijo():
    assert style.load("jehp").get("sections.back_matter.first_author")["enabled"] is False
    cfg = style.load("ijo").get("sections.back_matter.first_author")
    assert cfg["enabled"] is True and cfg["label"] == "First author:"


def test_block_before_the_figure_legends(tmp_path, ijo):
    ms, pl = plan(tmp_path, ijo)
    assert [n.text for n in pl.inserts] == [
        "First author:", "Asha Lakshmi Kumar", "Department of Ophthalmology, City Eye Hospital, Chennai, Tamil Nadu, India"]
    legends = next(p for p in ms.paragraphs if p.text == "Figure Legends")
    assert all(n.anchor == legends.id and n.where == "before" for n in pl.inserts)
    assert pl.inserts[0].runs == [("First author:", {"bold": True})]
    assert pl.inserts[1].runs[0][1] == {"bold": True}
    assert pl.inserts[2].runs[0][1] == {}
    for n in pl.inserts:
        assert (n.reason, n.source, n.layer, n.status) == ("first_author", Source.RULE, Layer.STRUCTURE, Status.SUGGESTED)
        assert n.provenance
    assert pl.queries == []


def test_first_legend_without_a_heading(tmp_path, ijo):
    ms, pl = plan(tmp_path, ijo, [p for p in BODY if p != "Figure Legends"])
    first = next(p for p in ms.paragraphs if p.text.startswith("Figure 1:"))
    assert pl.inserts[0].anchor == first.id and pl.inserts[0].where == "before"


def test_legend_in_a_heading_style_after_the_references(tmp_path, ijo):
    doc = Document()
    for t in BODY[:8]:
        doc.add_paragraph(t)
    doc.add_paragraph("Figure 1", style="Heading 3")
    doc.add_paragraph("Nomogram of the irradiation time")
    doc.save(tmp_path / "h.docx")
    ms = docmodel.load(str(tmp_path / "h.docx"))
    ms.front_sources = fm.Sources(meta=fm.parse_xml(XML))
    pl = first_author.plan(ms, ijo, fm.plan(ms, ijo))
    head = next(p for p in ms.paragraphs if p.text == "Figure 1")
    assert pl.inserts[0].anchor == head.id and pl.inserts[0].where == "before"


def test_no_legends_goes_at_the_end(tmp_path, ijo):
    ms, pl = plan(tmp_path, ijo, BODY[:8] + ["", ""])
    last_ref = next(p for p in ms.paragraphs if p.text.startswith("2. Rao"))
    assert all(n.anchor == last_ref.id and n.where == "after" for n in pl.inserts)


def test_existing_block_is_not_added_again(tmp_path, ijo):
    _, pl = plan(tmp_path, ijo, BODY[:8] + ["First Author:", "Asha Kumar", "City Eye Hospital"] + BODY[8:])
    assert pl.inserts == [] and pl.queries == []
    assert first_author._EXISTING.match("First author's details: Asha Kumar")
    assert not first_author._EXISTING.match("First authored papers were excluded.")


def test_unknown_first_author_gets_the_label_and_a_query(tmp_path, ijo):
    ms, pl = plan(tmp_path, ijo, xml=b"<manuscript><title>A title</title></manuscript>")
    assert [n.text for n in pl.inserts] == ["First author:"]
    assert [k for _, k in pl.queries] == ["first_author.missing"]


def test_unknown_affiliation_gets_a_query(tmp_path, ijo):
    xml = (b"<manuscript><title>A title</title><authors><author><fname>Asha</fname><lname>Kumar</lname></author>"
           b"</authors></manuscript>")
    _, pl = plan(tmp_path, ijo, xml=xml)
    assert [n.text for n in pl.inserts] == ["First author:", "Asha Kumar"]
    assert [k for _, k in pl.queries] == ["first_author.affiliation"]


def test_full_name_from_the_xml_but_never_a_phone_number():
    f = fm.Fields(authors=[fm.Author("Asha L", "Kumar")])
    src = fm.Sources(meta=fm.Meta(authors=[{"given": "Asha Lakshmi", "surname": "Kumar"}]))
    assert first_author.full_name(f.authors[0], f, src, {"initial_period": False}) == "Asha Lakshmi Kumar"
    src = fm.Sources(meta=fm.Meta(authors=[{"given": "Meena 9382113646", "surname": "Lakshmipathy"}]))
    assert first_author.full_name(fm.Author("Meena", "Lakshmipathy"), f, src, {}) == "Meena Lakshmipathy"


def test_affiliation_country_from_the_others():
    f = fm.Fields(affiliations={1: {"department": "Ophthalmology", "institution": "Hacettepe University",
                                    "city": "Ankara", "state": "", "country": ""},
                                2: {"department": "Pathology", "institution": "Hacettepe University",
                                    "city": "Ankara", "state": "", "country": "Turkey"}})
    assert fm.affiliation_text(f, 1) == "Department of Ophthalmology, Hacettepe University, Ankara, Turkey"


def test_off_for_other_journals(tmp_path):
    st = style.load("jehp")
    ms = docmodel.load(make(tmp_path, BODY))
    ms.front_sources = fm.Sources(meta=fm.parse_xml(XML))
    assert first_author.plan(ms, st, fm.plan(ms, st)).inserts == []


def test_why_box_has_a_rule_for_it():
    assert rule_sources.key_for("first_author") == "first_author"
    assert rule_sources.locator("first_author")[1].label == "First author"


def _texts(path):
    doc = Document(path)
    return [("".join(t.text or "" for t in p.iter(W + "t")), p) for p in doc.element.body.iter(W + "p")]


def test_pipeline_tracked_insertion_and_clean_copy_only_when_accepted(tmp_path, ijo):
    ms_path = make(tmp_path, BODY, "Manuscript.docx")
    z = tmp_path / "sub.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(ms_path, "Manuscript.docx")
        zf.writestr("ART_1.xml", XML)
    out = tmp_path / "out"
    res = pipeline.edit(str(z), ijo, str(out), stages=("style",))
    man, _, inserts, _ = pipeline.load_records(str(out))
    fa = [n for n in inserts if n.reason == "first_author"]
    assert [n.text for n in fa][:2] == ["First author:", "Asha Lakshmi Kumar"]
    assert all(n.layer == Layer.STRUCTURE for n in fa)
    assert man["stage_info"]["style"]["first_author"] == 3
    saved = json.load(open(res.files["front_matter"], encoding="utf-8"))
    assert {d["layer"] for d in saved if d["reason"] == "first_author"} == {"structure"}

    texts = [t for t, _ in _texts(res.path("tracked"))]
    i = texts.index("First author:")
    assert texts[i - 1].startswith("2. Rao") and texts[i + 1] == "Asha Lakshmi Kumar"
    assert texts[i + 3] == "Figure Legends"
    para = _texts(res.path("tracked"))[i][1]
    assert para.find(f"{W}pPr/{W}rPr/{W}ins") is not None and para.find(f"{W}ins") is not None  # tracked insertion
    assert "First author:" not in [t for t, _ in _texts(res.path("clean"))]  # pending: not in the clean copy

    for n in inserts:
        n.status = Status.ACCEPTED if n.reason == "first_author" else n.status
    pipeline.save_records(str(out), inserts=inserts)
    again = pipeline.render_outputs(str(out), style=ijo)
    clean = [t for t, _ in _texts(again.path("clean"))]
    assert "First author:" in clean and os.path.exists(again.path("clean"))


def test_no_block_for_a_letter_to_the_editor(tmp_path, ijo):
    letter = XML.replace(b"<articletype>Original Article</articletype>", b"<articletype>Letter to Editor</articletype>")
    assert plan(tmp_path, ijo, xml=letter)[1].inserts == []
    standalone = XML.replace(b"<articletype>Original Article</articletype>", b"<articletype>Letter to Editor "
                             b"(Stand-alone brief communications unrelated to IJO publications)</articletype>")
    assert len(plan(tmp_path, ijo, xml=standalone)[1].inserts) == 3  # the editor adds it to these
