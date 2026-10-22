import pytest
from docx import Document

from edit_os import biblio, docmodel, refengine, rules, style, tracked
from edit_os.changes import Status

SHIELDS = biblio.Record(
    source="pubmed", pmid="8694700", doi="10.1001/archopht.1996.01100140164007",
    title="Diffuse choroidal melanoma. Clinical features predictive of metastasis.",
    authors=[("Shields", "CL"), ("Shields", "JA"), ("De Potter", "P"), ("Cater", "J"), ("Tardio", "D"),
             ("Barrett", "J"), ("Smith", "A")],
    journal="Arch Ophthalmol", year="1996", volume="114", issue="8", pages="956-63",
)
OTHER = biblio.Record(source="pubmed", pmid="1", title="Something else entirely about cataract surgery outcomes",
                      authors=[("Shields", "CL")], journal="Ophthalmology", year="1996", volume="1", pages="1-2")


@pytest.fixture
def fake_pubmed(monkeypatch):
    calls = []

    def search(term, retmax=5):
        calls.append(term)
        return ["1", "8694700"] if "choroidal" in term.lower() else []

    def fetch(ids):
        return [r for r in (OTHER, SHIELDS) if r.pmid in ids]

    monkeypatch.setattr(biblio, "pubmed_search", search)
    monkeypatch.setattr(biblio, "pubmed_fetch", fetch)
    monkeypatch.setattr(biblio, "crossref_search", lambda text, rows=5: [])
    monkeypatch.setattr(biblio, "crossref_doi", lambda doi: None)
    return calls


def test_parse():
    p = refengine.parse("Shields CL, Shields JA, De Potter P. Diffuse choroidal melanoma. Clinical features predictive "
                        "of metastasis. Arch Ophthalmol. 1996;114(8):956-63. doi: 10.1001/archopht.1996.01100140164007.")
    assert p.authors == ["Shields CL", "Shields JA", "De Potter P"]
    assert p.first_surname == "Shields" and p.year == "1996"
    assert p.doi == "10.1001/archopht.1996.01100140164007"


def test_match_and_format(fake_pubmed):
    text = ("Shields CL, Shields JA, De Potter P, Cater J, Tardio D, Barrett J. Diffuse choroidal melanoma. "
            "Clinical features predictive of metastasis. Arch Ophthalmol. 1996;114(8):95-63.")
    res = refengine.resolve(text, style.load("ijo"))
    assert res.status == refengine.FORMAT_FIXED and res.record.pmid == "8694700"
    assert res.formatted == ("Shields CL, Shields JA, De Potter P, Cater J, Tardio D, Barrett J, et al. Diffuse choroidal "
                             "melanoma. Clinical features predictive of metastasis. Arch Ophthalmol 1996;114:956-63.")


def test_wrong_paper_becomes_query(fake_pubmed):
    res = refengine.resolve("Shields CL. Uveal melanoma in children. Ophthalmology. 1996;1:1-2.", style.load("ijo"))
    assert res.status == refengine.QUERY  # OTHER shares author and year but not the title: never swapped in


def test_doi_pointing_elsewhere_is_queried(fake_pubmed, monkeypatch):
    monkeypatch.setattr(biblio, "pubmed_search", lambda term, retmax=5: ["1"] if "[aid]" in term else [])
    res = refengine.resolve("Rao S. Timing of surgery in children with cataract. J Pediatr. 2019;5:1-2. "
                            "doi:10.1000/xyz", style.load("ijo"))
    assert res.status == refengine.QUERY and "points to a different article" in res.message


def test_books_are_left_alone(fake_pubmed):
    res = refengine.resolve("Ringsven MK, Bond D. Gerontology and Leadership Skills for Nurses. 2nd ed. Albany, NY: "
                            "Delmar Publishers; 1996.", style.load("ijo"))
    assert res.status == refengine.NOT_JOURNAL and not fake_pubmed


def test_lookup_rule_writes_word_level_changes(fake_pubmed, tmp_path):
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph("1. Shields CL, Shields JA, De Potter P, Cater J, Tardio D, Barrett J. Diffuse choroidal melanoma. "
                      "Clinical features predictive of metastasis. Arch Ophthalmol. 1996;114(8):95-63.")
    path = tmp_path / "m.docx"
    doc.save(path)
    cs = rules.run_all(docmodel.load(str(path)), style.load("ijo"))
    entry = [c for c in cs.changes if c.para_id == "p0002"]
    assert entry and all(c.reason == "ref_lookup" for c in entry)  # small reference rules step aside
    assert all(len(c.old) < 20 for c in entry)  # small edits, not one whole-entry replacement
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "o.docx"
    tracked.render(docmodel.load(str(path)), cs, tracked=False).save(out)
    assert docmodel.load(str(out)).paragraphs[1].text.endswith("et al. Diffuse choroidal melanoma. Clinical features "
                                                                "predictive of metastasis. Arch Ophthalmol 1996;114:956-63.")


def test_claude_fields_drive_the_search(fake_pubmed, monkeypatch):
    # The regex parser would cut this title at "melanoma."; Claude's split keeps it whole.
    text = ("Shields CL , Shields JA,De Potter P, Cater J, Tardio D, Barrett J. Diffuse choroidal melanoma. Clinical "
            "features predictive of metastasis Arch Ophthalmol 1996 114 956")
    fields = {"kind": "journal_article", "authors": ["Shields CL", "Shields JA", "De Potter P"],
              "title": "Diffuse choroidal melanoma. Clinical features predictive of metastasis", "journal": "Arch Ophthalmol",
              "year": "1996", "volume": "114", "issue": "", "pages": "956", "doi": "10.9999/made.up", "pmid": ""}
    monkeypatch.setattr(refengine.ai_refs, "available", lambda: True)
    monkeypatch.setattr(refengine.ai_refs, "parse_batch", lambda texts: [fields])
    [res] = refengine.resolve_all([text], style.load("ijo"))
    assert res.status == refengine.FORMAT_FIXED and res.record.pmid == "8694700"
    assert res.parsed.doi == ""  # a DOI that is not in the text is ignored, not looked up
    assert res.formatted.endswith("Arch Ophthalmol 1996;114:956-63.")


def test_claude_kind_marks_books(fake_pubmed, monkeypatch):
    monkeypatch.setattr(refengine.ai_refs, "available", lambda: True)
    monkeypatch.setattr(refengine.ai_refs, "parse_batch", lambda texts: [{"kind": "book", "title": "Gerontology"}])
    [res] = refengine.resolve_all(["Ringsven MK. Gerontology. Delmar; 1996;3:4-5."], style.load("ijo"))
    assert res.status == refengine.NOT_JOURNAL and not fake_pubmed


def _record(title, journal="J Test", **kw):
    return biblio.Record(source=kw.pop("source", "pubmed"), pmid="9", title=title, authors=[("Rao", "S")],
                         journal=journal, year="2020", volume="75", pages="e1874", **kw)


def test_sentence_case_keeps_proper_nouns_from_a_sentence_cased_record():
    st = style.load("ijo")
    rec = _record("Placido disc imaging of the cornea in healthy adults and miners: a survey for the U.S. "
                  "Preventive Services Task Force.")
    title = refengine.format_record(rec, st).split(". ", 1)[1]
    assert title.startswith("Placido disc imaging of the cornea in healthy adults and miners: A survey for the U.S. "
                            "Preventive Services Task Force.")
    # The author's own sentence-cased title vouches for its capitals too ("first case in Colombia").
    p = refengine.parse("Rao S. Ocular findings in coal miners: first series from Malawi. J Test 2020;75:e1874.")
    rec = _record("Ocular Findings in Coal Miners: First Series from Malawi", source="crossref")
    assert "First series from Malawi." in refengine.format_record(rec, st, p)


def test_sentence_case_of_title_cased_titles():
    sc = refengine.sentence_case
    assert sc("Cataract in Infants-Screening, Treatment, and Outcomes: A Survey") == \
        "Cataract in infants-Screening, treatment, and outcomes: A survey"
    assert sc("Comparison of Colour Doppler Imaging in Open-Angle Glaucoma") == \
        "Comparison of colour Doppler imaging in open-angle glaucoma"
    assert sc("Retinoblastoma in an Adult") == "Retinoblastoma in an adult"
    assert sc("Effect of First-Line Drugs in Japan") == "Effect of first-line drugs in Japan"


def test_journal_qualifier_kept_for_ijo():
    rec = _record("A study of glaucoma drugs.", journal="Clinics (Sao Paulo)")
    assert "Clinics (Sao Paulo) 2020;75:e1874." in refengine.format_record(rec, style.load("ijo"))
    assert "Clinics (Sao Paulo) 2020;75:e1874." in refengine.format_record(rec, style.load("ija"))  # Index Medicus
    st = style.load("ija")
    st.settings["references"]["journal_drop_qualifier"] = True
    assert "Clinics 2020;75:e1874." in refengine.format_record(rec, st)


def test_erratum_kept_and_styled(fake_pubmed):
    text = ("Shields CL, Shields JA, De Potter P, Cater J, Tardio D, Barrett J. Diffuse choroidal melanoma. "
            "Clinical features predictive of metastasis. Arch Ophthalmol. 1996 Aug;114(8):956-63. "
            "doi: 10.1001/archopht.1996.01100140164007. Epub 1996 Jul 1. "
            "Erratum in: Arch Ophthalmol. 1997 Apr;115(4):e408-e410.")
    res = refengine.resolve(text, style.load("ijo"))
    assert res.formatted.endswith("Arch Ophthalmol 1996;114:956-63. Erratum in: Arch Ophthalmol 1997;115:e408-10.")


def test_record_title_brackets_and_all_caps():
    from edit_os import biblio, refengine
    rec = biblio.Record(source="pubmed", title="[Solid limbal dermoid in an elderly patient].", authors=[])
    assert refengine._usable_title(rec, None) == "Solid limbal dermoid in an elderly patient"
    caps = biblio.Record(source="pubmed", title="WIDE-ANGLE VIEWING SYSTEMS FOR VITREORETINAL SURGERY", authors=[])
    assert refengine._usable_title(caps, None) == "Wide-angle viewing systems for vitreoretinal surgery"
    p = refengine.Parsed(text="", title="Wide-angle viewing systems for vitreoretinal surgery")
    if p is not None:
        assert refengine._usable_title(caps, p) == p.title


def test_record_title_with_names_counts_as_sentence_case():
    # NJCP / JFMPC now sentence-case reference titles (S example): names in a PubMed title keep their capitals.
    t = ("Health-related quality of life in patients with dyspepsia presenting at the University of Benin "
         "Teaching Hospital, Benin City, South-South Nigeria: a cross-sectional study")
    assert refengine.is_sentence_cased(t)
    assert refengine.sentence_case(t, refengine.proper_nouns(t)) == t
    assert not refengine.is_sentence_cased("Treatment of Metastatic Cancer in Adult Patients With Chronic Disease")
