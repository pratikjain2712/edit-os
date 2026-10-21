"""Clean-up of the author's run formatting (rules_format.rule_cleanup) and the properties it uses."""

import copy
import io
import os
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from edit_os import docmodel, rules, rules_format, style, tracked
from edit_os.changes import Change, ChangeSet, Layer, Source, Status

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))
import eval_format  # noqa: E402

BLUE = RGBColor(0x44, 0x72, 0xC4)
REFS = ["References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."]


def base(**over):
    st = style.load("base")
    st.settings["references"]["lookup"] = False
    st.settings["formatting"].update(over)
    return st


def cleanup(path, st=None):
    return [(c.reason, c.old, c.format) for c in rules_format.rule_cleanup(docmodel.load(path), st or base())]


def save(doc, tmp_path, name="m.docx"):
    for t in REFS:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def accept_all(cs):
    for c in cs.changes:
        c.status = Status.ACCEPTED
    return cs


def runs(path, para=0):
    p = list(Document(path).element.body.iter(qn("w:p")))[para]
    return [r for r in p.iter(qn("w:r"))]


def rpr_has(run, tag):
    rpr = run.find(qn("w:rPr"))
    return rpr is not None and rpr.find(qn(tag)) is not None


# --- colour --------------------------------------------------------------------------------

def test_font_colour_removed_as_one_change(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    for text, colour in (("The lesion ", BLUE), ("was", None), (" excised in toto.", BLUE),
                         (" Black text ", RGBColor(0, 0, 0)), ("https://x.org/a", BLUE)):
        r = p.add_run(text)
        if colour is not None:
            r.font.color.rgb = colour
    path = save(doc, tmp_path)
    got = [g for g in cleanup(path) if g[0] == "fmt_color"]
    # The uncoloured short word is bridged; black and the URL keep their colour.
    assert got == [("fmt_color", "The lesion was excised in toto.", {"color": False})]


def test_colour_tracked_and_clean(tmp_path):
    doc = Document()
    r = doc.add_paragraph().add_run("Coloured words.")
    r.font.color.rgb = BLUE
    r.font.size = Pt(12)
    r.font.name = "Arial"
    path = save(doc, tmp_path)
    cs = accept_all(rules.run_all(docmodel.load(path), base(), only=["format_cleanup"]))
    out = tmp_path / "t.docx"
    tracked.render(docmodel.load(path), cs, tracked=True).save(out)
    run = runs(str(out))[0]
    assert not rpr_has(run, "w:color")
    old = run.find(qn("w:rPr")).find(qn("w:rPrChange")).find(qn("w:rPr"))
    # The revision holds the run's full previous properties, not just the colour.
    assert {c.tag for c in old} >= {qn("w:rFonts"), qn("w:color"), qn("w:sz")}
    Document(str(out))  # opens
    clean = tmp_path / "c.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(clean)
    run = runs(str(clean))[0]
    assert not rpr_has(run, "w:color") and run.find(qn("w:rPr")).find(qn("w:rPrChange")) is None


def test_colour_from_heading_style_set_to_automatic(tmp_path):
    doc = Document()
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    doc.add_paragraph("Methods", style="Heading 1")
    path = save(doc, tmp_path)
    assert ("fmt_color", "Methods", {"color": "auto"}) in cleanup(path)


# --- tables --------------------------------------------------------------------------------

def table_doc(rows, caption_row=False, span_head=False):
    doc = Document()
    doc.add_paragraph("Body text before the table.")
    t = doc.add_table(rows=len(rows) + caption_row, cols=len(rows[0]))
    if caption_row:
        merged = t.rows[0].cells[0].merge(t.rows[0].cells[-1])
        merged.text = "Table 1: Results"
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            t.rows[i + caption_row].cells[j].text = text
    if span_head:
        t.rows[0].cells[1].merge(t.rows[0].cells[2])
        t.rows[0].cells[1].paragraphs[0].runs[0].text = "Group"
        for extra in t.rows[0].cells[1].paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
    return doc


def test_table_column_heads_bold(tmp_path):
    path = save(table_doc([["Variable", "Mean (SD)", "p"], ["Age", "45 (3)", "0.04"]]), tmp_path)
    got = [g for g in cleanup(path) if g[0] == "fmt_table_head"]
    assert [o for _, o, _ in got] == ["Variable", "Mean (SD)", "p"]
    assert all(f == {"bold": True} for *_, f in got)


def test_table_head_below_caption_row(tmp_path):
    path = save(table_doc([["Variable", "Value"], ["Age", "45"]], caption_row=True), tmp_path)
    got = cleanup(path)
    assert [o for r, o, _ in got if r == "fmt_table_head"] == ["Variable", "Value"]
    assert ("fmt_label", "Table 1:", {"bold": True}) in got


def test_spanning_head_has_sub_heads(tmp_path):
    rows = [["Variable", "x", "x"], ["", "Number", "%"], ["Age", "12", "40.5"]]
    path = save(table_doc(rows, span_head=True), tmp_path)
    assert [o for r, o, _ in cleanup(path) if r == "fmt_table_head"] == ["Variable", "Group", "Number", "%"]


def test_table_body_bold_only_when_switched_on(tmp_path):
    doc = table_doc([["Study", "Year"], ["Smith", "2001"]])
    t = doc.tables[0]
    t.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    t.rows[0].cells[1].paragraphs[0].runs[0].bold = True
    t.rows[1].cells[0].paragraphs[0].runs[0].bold = True
    path = save(doc, tmp_path)
    assert [g for g in cleanup(path) if g[0] != "fmt_heading_bold"] == []  # ("References")
    assert ("fmt_table_body", "Smith", {"bold": False}) in cleanup(path, base(table_body_unbold=True))


# --- labels, headings, running text --------------------------------------------------------

def test_labels_bold(tmp_path):
    doc = Document()
    doc.add_paragraph("Keywords: cornea; dermoid")
    doc.add_paragraph("Figure 2: Slit-lamp photograph")
    doc.add_paragraph("Supplementary Table 1: Search strategy")
    p = doc.add_paragraph()
    p.add_run("Table 3").bold = True
    p.add_run(": already bold, the colon does not matter")
    doc.add_paragraph("Table 4 shows the results.")
    path = save(doc, tmp_path)
    assert [(r, o) for r, o, _ in cleanup(path)] == [("fmt_label", "Keywords:"), ("fmt_label", "Figure 2:"),
                                                     ("fmt_heading_bold", "References")]


def test_headings_bold_and_not_underlined(tmp_path):
    doc = Document()
    doc.add_paragraph("Some text in the abstract.")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Intro text.")
    r = doc.add_paragraph().add_run("Discussion")
    r.bold, r.underline = True, True
    long = doc.add_paragraph("A sentence formatted as a heading by mistake, which is long and ends with a period.",
                             style="Heading 1")
    assert long.text
    path = save(doc, tmp_path)
    got = cleanup(path)
    assert ("fmt_heading_bold", "Introduction", {"bold": True}) in got
    assert ("fmt_heading_underline", "Discussion", {"underline": False}) in got
    assert not any(r == "fmt_heading_bold" and o.startswith("A sentence") for r, o, _ in got)


def test_stray_bold_in_running_text(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph("Results are shown in ")
    p.add_run("Table 1").bold = True
    p.add_run(" and differed (")
    p.add_run("P < 0.05").bold = True
    p.add_run(")")
    p.add_run(".").bold = True
    p = doc.add_paragraph()
    p.add_run("Statistical analysis:").bold = True
    p.add_run(" SPSS was used.")
    p = doc.add_paragraph()
    p.add_run("ASMs").bold = True
    p.add_run(" are a list term at the start and stay bold.")
    path = save(doc, tmp_path)
    got = [(r, o) for r, o, _ in cleanup(path) if r == "fmt_body_bold"]
    assert got == [("fmt_body_bold", "Table 1"), ("fmt_body_bold", "P < 0.05).")]  # ")" is bridged
    assert not any(r == "fmt_body_bold" for r, *_ in cleanup(path, base(body_unbold=False)))


def test_switches(tmp_path):
    doc = Document()
    r = doc.add_paragraph("Introduction").runs[0]
    r.font.color.rgb = BLUE
    path = save(doc, tmp_path)
    assert {(g[0], g[1]) for g in cleanup(path)} == {("fmt_color", "Introduction"), ("fmt_heading_bold", "Introduction"),
                                                     ("fmt_heading_bold", "References")}
    assert cleanup(path, base(enabled=False)) == []
    assert cleanup(path, base(remove_font_color=False, heading_bold=False)) == []


# --- with text changes ---------------------------------------------------------------------

def test_format_and_text_changes_together(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    r = doc.add_paragraph().add_run("Smith et al found it in vitro.")
    r.font.color.rgb = BLUE
    path = save(doc, tmp_path)
    cs = rules.run_all(docmodel.load(path), base(), only=["formatting", "format_cleanup"])
    reasons = sorted(c.reason for c in cs.changes if c.para_id == "p0002")
    assert reasons == ["fmt_color", "fmt_et_al", "fmt_latin"]  # the colour change overlaps both
    accept_all(cs)
    for is_tracked in (True, False):
        out = tmp_path / f"o{is_tracked}.docx"
        tracked.render(docmodel.load(path), cs, tracked=is_tracked).save(out)
        assert docmodel.load(str(out)).paragraphs[1].text == "Smith et al. found it in vitro."
        live = [r for r in runs(str(out), 1) if r.find(qn("w:t")) is not None]
        assert live and not any(rpr_has(r, "w:color") for r in live)  # inserted "et al." too


def test_format_change_inside_replaced_text_dropped():
    class P:
        id, text = "p0001", "abc def"
    text = Change("t", "p0001", 0, 7, "abc def", "xyz", Layer.LANGUAGE, Source.RULE)
    fmt = Change("f", "p0001", 0, 3, "abc", "abc", Layer.HOUSE_STYLE, Source.RULE, format={"color": False})
    saved = dict(rules.RULES)
    try:
        rules.RULES.clear()
        rules.RULES.update(a=lambda ms, st: [text], b=lambda ms, st: [fmt])
        assert [c.id for c in rules.run_all(None, base()).changes] == ["t"]
    finally:
        rules.RULES.clear()
        rules.RULES.update(saved)


# --- renderer: the new properties ------------------------------------------------------------

def test_render_clear_properties(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("marked text")
    r.font.underline = True
    r.font.size = Pt(16)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "FFFF00")
    r._r.get_or_add_rPr().append(shd)
    path = tmp_path / "r.docx"
    doc.save(path)
    c = Change("f", "p0001", 0, 6, "marked", "marked", Layer.HOUSE_STYLE, Source.RULE, status=Status.ACCEPTED,
               format={"underline": False, "size": False, "highlight": False})
    out = tmp_path / "o.docx"
    tracked.render(docmodel.load(str(path)), ChangeSet([c]), tracked=True).save(out)
    first, rest = [r for r in runs(str(out)) if r.find(qn("w:t")) is not None]
    for tag in ("w:u", "w:sz", "w:highlight", "w:shd"):
        assert not rpr_has(first, tag) and rpr_has(rest, tag)
    old = first.find(qn("w:rPr")).find(qn("w:rPrChange")).find(qn("w:rPr"))
    assert {qn("w:u"), qn("w:sz"), qn("w:highlight"), qn("w:shd")} <= {c.tag for c in old}


def test_revision_ids_unique_with_existing_revisions(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Existing tracked text.")
    ins = OxmlElement("w:ins")
    for k, v in (("w:id", "9001"), ("w:author", "Author"), ("w:date", "2026-01-01T00:00:00Z")):
        ins.set(qn(k), v)
    run = copy.deepcopy(p.runs[0]._r)
    run.find(qn("w:t")).text = " Added."
    ins.append(run)
    p._p.append(ins)
    path = tmp_path / "e.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    c = Change("f", "p0001", 0, 8, "Existing", "Existing", Layer.HOUSE_STYLE, Source.RULE, status=Status.ACCEPTED,
               format={"bold": True})
    out = tmp_path / "o.docx"
    tracked.render(ms, ChangeSet([c]), tracked=True).save(out)
    ids = [el.get(qn("w:id")) for el in Document(str(out)).element.body.iter(qn("w:ins"), qn("w:rPrChange"))]
    assert len(ids) == len(set(ids)) == 2


# --- tools/eval_format.py --------------------------------------------------------------------

def test_eval_reads_editor_formatting(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Table 1").bold = True
    p.add_run(": Results ")
    blue = p.add_run("in blue")
    rpr = blue._r.get_or_add_rPr()
    ch = OxmlElement("w:rPrChange")
    for k, v in (("w:id", "1"), ("w:author", "E"), ("w:date", "2026-01-01T00:00:00Z")):
        ch.set(qn(k), v)
    old = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "4472C4")
    old.append(col)
    ch.append(old)
    rpr.append(ch)
    first = p.runs[0]._r.get_or_add_rPr()
    ch2 = copy.deepcopy(ch)
    ch2.find(qn("w:rPr")).remove(ch2.find(qn("w:rPr"))[0])
    first.append(ch2)  # bold added by the editor
    buf = io.BytesIO()
    doc.save(buf)
    rec = eval_format.editor_changes(buf.getvalue())[0]
    assert rec["before"] == "Table 1: Results in blue"
    assert [(s["label"], s["text"]) for s in rec["spans"]] == [("bold+", "Table 1"), ("color-", "in blue")]
