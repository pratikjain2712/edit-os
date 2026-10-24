import pytest
from docx import Document

from edit_os import docmodel, qc, rules, style, tracked
from edit_os.changes import Status


@pytest.fixture
def manuscript(tmp_path):
    doc = Document()
    for text in [
        "Abstract",
        "Key words: retina, cataract, amblyopia",
        "Introduction",
        "Dermoids are rare (1). Earlier series reported similar results (2–3), and others (2, 3).",
        "Mean age was 45 (12) years and p=0.03.",
        "Methods:",
        "Results",
        "Conclusion",
        "References",
        "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit Amst Neth. 2002;21(4):295-299.",
        "2. Duncan JL, Golabi M, Fredrick DR, Hoyt CS, Hwang DG, Kramer SG, Smith A. Complex limbal choristomas. "
        "Ophthalmology. 1998;105(8):1459-65.",
        "3. Lindqvist P., Hale R. Timing of surgery. J Med Genet 1987;24:616-20.",
    ]:
        doc.add_paragraph(text)
    path = tmp_path / "m.docx"
    doc.save(path)
    return str(path)


def apply(path, st, tmp_path):
    cs = rules.run_all(docmodel.load(path), st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "out.docx"
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    return [p.text for p in docmodel.load(str(out)).paragraphs]


def test_ijo_rules(manuscript, tmp_path):
    t = apply(manuscript, style.load("ijo"), tmp_path)
    assert t[0] == "Abstract"  # IJO does not tag "Abstract" and deletes "Introduction" (paragraph and all)
    assert t[1] == "Keywords: Amblyopia, cataract, retina"
    assert t[6:11] == ["<H1>Conclusion", "<H2>Financial support and sponsorship", "Nil.", "<H2>Conflicts of interest",
                       "There are no conflicts of interest."]
    del t[7:11]
    assert t[2] == "Dermoids are rare.[1] Earlier series reported similar results,[2,3] and others.[2,3]"
    # statistics, not a citation: operators spaced (G297), P capital (G294); H1 in Title case (S: "H1 | Caps", owner reading)
    assert t[3] == "Mean age was 45 (12) years and P = 0.03."
    assert t[4] == "<H1>Methods:" and t[5] == "<H1>Results" and t[7] == "<H1>References"  # G137 keeps the colon
    assert t[8] == "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: A case report. Orbit Amst Neth 2002;21:295-9."
    assert t[9] == ("2. Duncan JL, Golabi M, Fredrick DR, Hoyt CS, Hwang DG, Kramer SG, et al. "
                    "Complex limbal choristomas. Ophthalmology 1998;105:1459-65.")
    assert t[10] == "3. Lindqvist P, Hale R. Timing of surgery. J Med Genet 1987;24:616-20."


def test_base_rules_space_operators_and_tag_all_h1(manuscript, tmp_path):
    t = apply(manuscript, style.load("base"), tmp_path)
    del t[8:12]  # back matter added before the references
    assert t[0] == "<H1>Abstract" and t[2] == "<H1>Introduction"
    assert t[4] == "Mean age was 45 (12) years and P = 0.03."
    assert "<H1>Materials and Methods:" in t  # G138: "Methods" -> "Materials and Methods"; G137 keeps the colon
    assert t[9].endswith("2002;21:295-9.")


def test_qc_blocks_unedited_file(manuscript):
    report = qc.check(docmodel.load(manuscript), style.load("ijo"))
    checks = {i.check for i in report.issues}
    assert not report.passed
    assert {"citations.brackets", "sections.required", "references.authors", "keywords.order"} <= checks


def test_qc_after_rules_passes(manuscript, tmp_path):
    t = apply(manuscript, style.load("ijo"), tmp_path)
    doc = Document()
    for text in t:
        doc.add_paragraph(text)
    path = tmp_path / "edited.docx"
    doc.save(path)
    errors = {i.check for i in qc.check(docmodel.load(str(path)), style.load("ijo")).issues if i.severity == qc.ERROR}
    assert not errors  # the section rules add the back matter and sort the key words


def test_superscript_citations(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Colorectal cancer is common worldwide.¹ Rash is frequent⁴⁻⁶, and dry skin too")
    r = p.add_run("7,8")
    r.font.superscript = True
    p.add_run(". Leukocytes were 5 × 10")
    p.add_run("9").font.superscript = True
    p.add_run("/L in a 2 cm")
    p.add_run("2").font.superscript = True
    p.add_run(" area, 3 mm³ deep.")
    for t in ["References"] + [f"{n}. Author A. Title {n}. J Test 2001;{n}:1-2." for n in range(1, 9)]:
        doc.add_paragraph(t)
    path = tmp_path / "s.docx"
    doc.save(path)
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    t = apply(str(path), st, tmp_path)
    assert t[0] == ("Colorectal cancer is common worldwide.[1] Rash is frequent,[4-6] and dry skin too.[7,8] "
                    "Leukocytes were 5 × 109/L in a 2 cm2 area, 3 mm³ deep.")  # powers untouched
    out = Document(tmp_path / "out.docx")
    sup = [r.text for r in out.paragraphs[0].runs if r.font.superscript]
    assert "".join(sup) == "[1][4-6][7,8]92"  # citations superscript (guide); exponents stay superscript too


def _refs(doc, n):
    for t in ["References"] + [f"{k}. Author A. Title {k}. J Test 2001;{k}:1-2." for k in range(1, n + 1)]:
        doc.add_paragraph(t)


def _no_lookup(code="ijo"):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    st.settings["sections"]["back_matter"]["add_missing"] = False
    st.settings["sections"]["headings"]["drop"] = []  # keeps paragraph positions fixed
    return st


def test_statistic_indices_are_not_citations(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph("Flow was similar (31.4 ± 1.7%, p")
    p.add_run("1").font.superscript = True
    p.add_run(" = 0.412), and changed over time (p")
    p.add_run("2").font.superscript = True
    p.add_run(" < 0.05). The p")
    p.add_run("1").font.superscript = True
    p.add_run(" values compare groups; n² ≥ 4 and P¹=0.2. Glaucoma is common.")
    p.add_run("3").font.superscript = True
    _refs(doc, 3)
    path = tmp_path / "p.docx"
    doc.save(path)
    t = apply(str(path), _no_lookup(), tmp_path)
    assert t[1] == ("Flow was similar (31.4 ± 1.7%, p1 = 0.412), and changed over time (p2 < 0.05). "
                    "The p1 values compare groups; n² ≥ 4 and P¹=0.2. Glaucoma is common.[3]")


def test_front_matter_superscripts_are_not_citations(tmp_path):
    doc = Document()
    doc.add_paragraph("Running title: A rare tumour")
    p = doc.add_paragraph("Authors: Asha Rao¹, Meena Iyer², Ravi Menon")
    p.add_run("3").font.superscript = True
    doc.add_paragraph("¹Department of Ophthalmology, City Hospital, Pune, India²Junior Resident, City Hospital")
    p = doc.add_paragraph("")
    p.add_run("3").font.superscript = True
    p.add_run(" Professor, Department of Pathology, City Hospital, Pune, India")
    doc.add_paragraph("Corresponding Author: Dr. Meena Iyer²")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("A rare tumour in an adult is reported here.")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Tumours are rare in adults.¹ Others agree.²")
    _refs(doc, 3)
    path = tmp_path / "f.docx"
    doc.save(path)
    ms = docmodel.load(str(path))
    cs = rules.rule_citations(ms, _no_lookup())
    front = {p.id for p in ms.paragraphs[:5]}
    assert not [c for c in cs if c.para_id in front]
    t = apply(str(path), _no_lookup(), tmp_path)
    assert "Tumours are rare in adults.[1] Others agree.[2]" in t


def test_bracketed_unicode_superscripts(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Rigidity rises.[¹⁰] It is safe.[¹3] Many agree.[⁴-⁵]")
    _refs(doc, 13)
    path = tmp_path / "b.docx"
    doc.save(path)
    st = _no_lookup()
    st.settings["renumber"] = {"enabled": False}
    t = apply(str(path), st, tmp_path)
    assert t[1] == "Rigidity rises.[10] It is safe.[13] Many agree.[4,5]"


def test_round_bracket_counts_in_a_square_citing_manuscript(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Corneas are scarce.[1] Contamination is common.[2,3] It varies.[4]")
    doc.add_paragraph("Swabs help.[5] Media help.[6] Cold helps.[7] Speed helps.[8] So do others.[1,2]")
    doc.add_paragraph("Causes were trauma (7) and stroke (2); scar (3), foreign body (1).")
    doc.add_paragraph("Two risks were noted: (1) a long delay and (2) recent surgery.[5]")
    doc.add_paragraph("Lenses were described (6), as were rings (7).")
    _refs(doc, 8)
    path = tmp_path / "c.docx"
    doc.save(path)
    st = _no_lookup()
    st.settings["renumber"] = {"enabled": False}
    t = apply(str(path), st, tmp_path)
    assert t[3] == "Causes were trauma (7) and stroke (2); scar (3), foreign body (1)."
    assert t[4] == "Two risks were noted: (1) a long delay and (2) recent surgery.[5]"
    assert t[5] == "Lenses were described,[6] as were rings.[7]"  # rising: citations


def test_greek_and_operator_superscripts_are_not_citations(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Effects were large (F = 74.8, η² = 0.42; χ² = 23.2) with partial η² of 0.5, "
                      "and R²=0.45. Falls are common in older adults.²")
    _refs(doc, 3)
    path = tmp_path / "g.docx"
    doc.save(path)
    t = apply(str(path), _no_lookup(), tmp_path)
    assert t[1] == ("Effects were large (F = 74.8, η² = 0.42; χ² = 23.2) with partial η² of 0.5, "
                    "and R² = 0.45. Falls are common in older adults.[2]")


def test_numbered_questions_after_colon_are_not_citations(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Swabs help (1). Media help (2). Cold helps (3).")
    doc.add_paragraph("The review asks three questions: (1) How do schools act? (2) What do teachers do? "
                      "and (3) Which factors matter?")
    _refs(doc, 3)
    path = tmp_path / "q.docx"
    doc.save(path)
    st = _no_lookup()
    st.settings["renumber"] = {"enabled": False}
    t = apply(str(path), st, tmp_path)
    assert t[1] == "Swabs help.[1] Media help.[2] Cold helps.[3]"
    assert t[2] == ("The review asks three questions: (1) How do schools act? (2) What do teachers do? "
                    "and (3) Which factors matter?")


def test_comparison_after_a_word_left_as_typed_when_switched_off(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Patients aged ≥18 years with a score >5 were included (P<0.05, n=28, BMI ≥30).")
    path = tmp_path / "s.docx"
    doc.save(path)
    st = _no_lookup("base")  # base: unary signs closed (G301)
    assert apply(str(path), st, tmp_path)[1] == \
        "Patients aged ≥18 years with a score >5 were included (P < 0.05, n = 28, BMI ≥ 30)."
    st.settings["statistics"]["unary_after_word"] = True
    assert apply(str(path), st, tmp_path)[1] == \
        "Patients aged ≥ 18 years with a score > 5 were included (P < 0.05, n = 28, BMI ≥ 30)."


def test_superscript_citation_after_space_takes_the_space(tmp_path):
    import docx
    from edit_os import docmodel, rules, style, tracked
    from edit_os.changes import Status
    d = docx.Document()
    d.add_paragraph("Introduction")
    p = d.add_paragraph("It is rare among congenital malformations ")
    p.add_run("2").font.superscript = True
    p.add_run(". It has two forms.")
    d.add_paragraph("References")
    d.add_paragraph("1. Smith A. One. J X 2001;1:1-2.")
    d.add_paragraph("2. Jones B. Two. J Y 2002;2:3-4.")
    f = tmp_path / "m.docx"
    d.save(f)
    cs = rules.run_all(docmodel.load(str(f)), style.load("mjdrdypu"), only=["citations"])
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = tmp_path / "o.docx"
    tracked.render(docmodel.load(str(f)), cs, tracked=False).save(out)
    text = [q.text for q in docmodel.load(str(out)).paragraphs][1]
    assert text == "It is rare among congenital malformations.[2] It has two forms."
