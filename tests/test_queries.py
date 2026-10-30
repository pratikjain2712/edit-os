import zipfile

from docx import Document
from openpyxl import load_workbook

from edit_os import docmodel, queries, rules, style, tracked
from edit_os.changes import Status


def make(tmp_path):
    doc = Document()
    # an IMRaD paper without a type: an Original Article, so its missing Conclusion is queried (owner, 2026-10-27)
    for t in ["Key words: retina",
              "Introduction", "Methods", "Results",
              "Discussion",
              "Dermoids are rare (1). Some grow fast [4].",
              "References",
              "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit Amst Neth. 2002;21(4):295-299.",
              "2. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4."]:
        doc.add_paragraph(t)
    path = tmp_path / "m.docx"
    doc.save(path)
    return str(path)


def ijo_offline():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False  # reference-lookup queries are covered in test_refengine
    return st


def test_collect_queries(tmp_path):
    ms = docmodel.load(make(tmp_path))
    qs = queries.collect(ms, ijo_offline())
    got = [(q.id, q.check) for q in qs.queries]
    assert got == [("AQ1", "keywords.too_few"), ("AQ2", "sections.conclusion"),
                   ("AQ3", "references.missing_entry"), ("AQ4", "references.uncited")]
    assert qs.queries[3].text == "Reference no. 2 has not been cited. Kindly cite."
    assert ms.by_id(qs.queries[2].para_id).text[qs.queries[2].start:qs.queries[2].end] == "[4]"


def test_comments_and_tracked_changes_together(tmp_path):
    path = make(tmp_path)
    st = ijo_offline()
    changes = rules.run_all(docmodel.load(path), st)
    for c in changes.changes:
        c.status = Status.ACCEPTED
    ms = docmodel.load(path)
    qs = queries.collect(ms, st)
    assert queries.write_comments(ms, qs) == 4
    out = tmp_path / "t.docx"
    tracked.render(ms, changes, tracked=True).save(out)
    comments = zipfile.ZipFile(out).read("word/comments.xml").decode()
    assert "AQ4: Reference no. 2 has not been cited. Kindly cite." in comments
    assert "AQ3: Citation [4] has no entry" in comments
    body = Document(out).element.body
    assert len(body.findall(".//{*}commentRangeStart")) == 4 and body.findall(".//{*}ins")


def test_query_sheet(tmp_path):
    ms = docmodel.load(make(tmp_path))
    qs = queries.collect(ms, ijo_offline())
    queries.write_sheet(ms, qs, str(tmp_path / "q.xlsx"), "IJO_1_26")
    rows = list(load_workbook(tmp_path / "q.xlsx").active.values)
    assert rows[1] == ("No.", "Location", "Query", "Author's answer")
    assert rows[4][0] == "AQ3" and rows[4][1] == "[4]"


def test_reference_lookup_queries(tmp_path, monkeypatch):
    from edit_os import refengine

    def fake(texts, st):
        p = refengine.parse(texts[0])
        p.authors = ["Mohammad AEA", "Kroosh SS"]
        return [refengine.Resolution(refengine.QUERY, p, message="Not found in PubMed or CrossRef"),
                refengine.Resolution(refengine.QUERY, refengine.parse(texts[1]), message="DOI 10.1/x points to a different article")]

    monkeypatch.setattr(refengine, "resolve_all", fake)
    path = tmp_path / "r.docx"
    doc = Document()
    for t in ["References", "1. Mohammad AEA, Kroosh SS, et al. Huge dermoid. Orbit. 2002;21:295-9.",
              "2. Lee K. Back pain. Spine. 1997;22:1380-4. doi:10.1/x"]:
        doc.add_paragraph(t)
    doc.save(path)
    qs = queries.collect(docmodel.load(str(path)), style.load("ijo"))
    texts = [q.text for q in qs.queries]
    assert "Reference not found in PubMed. Hence kindly provide minimum 6 author names before et al." in texts
    assert "The DOI given points to a different article. Kindly check the reference." in texts


def test_manufacturer_missing():
    # JEHP editors: "Kindly provide complete manufacturer details ..." over "SPSS version 26" (19 of 302 articles)
    from types import SimpleNamespace as P

    from edit_os.queries import manufacturer_missing

    paras = [P(text="Data were analysed with SPSS version 26 and later in SPSS again.", in_table=False),
             P(text="We used SPSS Statistics 25 (IBM Corp., Armonk, NY, USA) and JASP.", in_table=False)]
    got = [(p.text[s:e]) for p, (s, e) in manufacturer_missing(paras)]
    assert got == ["SPSS", "JASP"]
    assert manufacturer_missing(paras[1:])[0][0].text[slice(*manufacturer_missing(paras[1:])[0][1])] == "JASP"


def test_jehp_standing_queries_and_dropped_checks(tmp_path):
    # JEHP editors' standing query on the headings (about 85 of 302 articles), on the first heading after the abstract;
    # queries the editors never raise are dropped (owner, 2026-10-30)
    from docx import Document

    from edit_os import docmodel, queries, rules_sections, style

    d = Document()
    for t in ["Abstract", "Background: We studied care. Methods: A survey. Results: Good. Conclusion: Fine.",
              "Keywords: Care; Survey", "Introduction", "Care matters [1].", "Methods", "We surveyed staff.",
              "References", "1. Smith A. A study. J Test 2020;1:1-2."]:
        d.add_paragraph(t)
    d.save(tmp_path / "a.docx")
    ms = docmodel.load(str(tmp_path / "a.docx"))
    st = style.load("jehp")
    qs = queries.QuerySet()
    rules_sections.add_queries(ms, st, qs)
    got = [(ms.by_id(q.para_id).text, q.text) for q in qs.queries if q.check == "sections.heading_levels_check"]
    assert got == [("Introduction", "Please ensure that the section heading levels are treated correctly.")]
    assert {"abstract.too_long", "front.running_head_long", "text.brackets"} <= set(st.get("queries.drop"))
    assert st.get("front_matter.queries.authors_check") is True
    assert style.load("ijo").get("queries.drop") == [] and not style.load("ijo").get("sections.heading_levels_query")
