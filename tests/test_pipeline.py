"""pipeline.edit / render_outputs and the `edit` / `render` commands, offline on a synthetic manuscript."""

import json
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from edit_os import __main__ as cli
from edit_os import docmodel, pipeline, style
from edit_os.changes import Change, ChangeSet, Layer, Source, Status
from edit_os.copyedit import langtool, spelling

VARCON = """\
# color <verified> (level 10)
A Cv: color / B C Z: colour

# randomized <verified> (level 50)
A Z: randomized / B: randomised
"""

TEXT = [
    "Original Article",
    "Introduction",
    "The colour of the lesion was noted in 40 randomised patients (1). Dermoids are rare [3].",
    "Methods",
    "We measured p=0.05 in all patients.",
    "Key words: retina",
    "Discussion",
    "Some grow fast (2).",
    "References",
    "1. Mohammad AEA, Kroosh SS. Huge corneal dermoid: a case report. Orbit Amst Neth. 2002;21(4):295-299.",
    "2. Lee K. Back pain in adolescents. Spine. 1997;22:1380-4.",
]


@pytest.fixture(autouse=True)
def offline_copyedit(monkeypatch, tmp_path):
    path = tmp_path / "varcon.txt"
    path.write_text(VARCON)
    v = spelling.parse(VARCON.splitlines())
    monkeypatch.setattr(spelling, "load", lambda *a, **k: v)
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)


def offline_style():
    st = style.load("ijo")
    st.settings["references"]["lookup"] = False
    return st


def make_docx(path, texts=TEXT):
    doc = Document()
    for t in texts:
        doc.add_paragraph(t)
    doc.save(path)
    return str(path)


def run(tmp_path, **kw):
    src = make_docx(tmp_path / "ms.docx")
    return pipeline.edit(src, offline_style(), str(tmp_path / "out"), **kw)


def test_all_outputs_exist_and_proposals_wait_for_the_editor(tmp_path):
    out = run(tmp_path)
    for kind in pipeline.FILES:
        assert (tmp_path / "out" / pipeline.FILES[kind].format(stem="ms")).exists(), kind
    assert (tmp_path / "out" / "source" / "ms.docx").exists()
    changes = ChangeSet.load(out.path("changes"))
    assert changes.changes and all(c.status == Status.SUGGESTED for c in changes.changes)
    assert len({c.id for c in changes.changes}) == len(changes.changes)
    reasons = {c.reason for c in changes.changes}
    assert any(r.startswith("spelling") for r in reasons), reasons  # copy-edit stage ran
    assert {Layer.LANGUAGE.value} < set(out.summary["changes"]["by_layer"]) and len(out.summary["changes"]["by_layer"]) > 1

    # Clean copy: nothing accepted yet, so the text is unchanged; tracked copy shows the suggestions.
    clean = [p.text for p in docmodel.load(out.path("clean")).paragraphs]
    assert "colour" in clean[2]
    body = Document(out.path("tracked")).element.body
    assert body.findall(".//{*}ins") and body.findall(".//{*}commentRangeStart")
    ws = load_workbook(out.path("queries")).active
    assert ws["A3"].value == "AQ1"


def test_clean_copy_drops_author_highlighting(tmp_path):
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    for t in TEXT:
        doc.add_paragraph(t)
    run_ = doc.paragraphs[2].runs[0]
    run_.font.highlight_color = WD_COLOR_INDEX.YELLOW
    shd = OxmlElement("w:shd")  # coloured shading, the other way authors mark text
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "00FFFF")
    run_._r.get_or_add_rPr().append(shd)
    doc.save(tmp_path / "ms.docx")
    out = pipeline.edit(str(tmp_path / "ms.docx"), offline_style(), str(tmp_path / "out"))

    clean = Document(out.path("clean")).element.body
    assert not clean.findall(".//{*}highlight") and not clean.findall(".//{*}rPr/{*}shd")
    assert Document(out.path("tracked")).element.body.findall(".//{*}highlight")  # editors still see it


def test_qc_report_content(tmp_path):
    out = run(tmp_path)
    rep = json.load(open(out.path("qc_json")))
    assert rep["journal"] == "IJO" and rep["stages"] == ["style", "copyedit"]
    assert "citations" in rep["stage_info"]["style"]["ran"]
    assert "spelling" in rep["stage_info"]["copyedit"]["ran"]
    assert rep["changes"]["pending"] == rep["changes"]["total"] + rep["front_matter_inserts"]["total"]
    assert rep["qc"]["basis"] == "accepted + pending suggestions"
    checks = {c["check"]: c for c in rep["qc"]["checks"]}
    assert set(checks) >= {"references.present", "citations.order", "statistics.spacing"}
    assert checks["references.present"]["passed"]
    assert not checks["references.exists"]["passed"]  # [3] has no entry, even after the edits
    assert rep["queries"]["by_check"]["references.missing_entry"] == 1
    assert rep["delivery_ready"] is False
    page = open(out.path("qc_html")).read()
    assert "QC report: ms" in page and "references.exists" in page and "language" in page


def test_rerender_reflects_decisions(tmp_path):
    out = run(tmp_path)
    outdir = str(tmp_path / "out")
    man, changes, inserts, qdicts = pipeline.load_records(outdir)
    colour = next(c for c in changes.changes if c.old == "colour")
    rand = next(c for c in changes.changes if c.old == "randomised")
    for c in changes.changes:
        c.status = Status.REJECTED
    colour.status = Status.ACCEPTED
    for n in inserts:
        n.status = Status.REJECTED
    qdicts[0]["status"] = "rejected"
    pipeline.save_records(outdir, changes, inserts, qdicts)

    out2 = pipeline.render_outputs(outdir, style=offline_style())
    text = docmodel.load(out2.path("clean")).paragraphs[2].text
    assert "color of the lesion" in text and "randomised" in text
    assert out2.summary["changes"]["pending"] == 0 and out2.summary["qc"]["basis"].startswith("clean")
    assert out2.summary["queries"]["rejected"] == 1
    tracked_xml = zipfile.ZipFile(out2.path("tracked")).read("word/document.xml").decode()
    assert "randomized" not in tracked_xml  # rejected: not in the tracked copy either
    comments = zipfile.ZipFile(out2.path("tracked")).read("word/comments.xml").decode()
    assert "AQ1:" in comments and qdicts[0]["text"] not in comments
    assert rand.id not in {c.id for c in ChangeSet.load(out2.path("changes")).changes if c.status == Status.ACCEPTED}


def test_stale_change_is_skipped_and_reported(tmp_path):
    run(tmp_path, stages=("style",))
    outdir = str(tmp_path / "out")
    _, changes, inserts, qdicts = pipeline.load_records(outdir)
    changes.add(Change("bad1", "p0003", 0, 3, "XYZ", "abc", Layer.LANGUAGE, Source.EDITOR, status=Status.ACCEPTED))
    pipeline.save_records(outdir, changes)
    out = pipeline.render_outputs(outdir, style=offline_style())
    assert any(c["para_id"] == "p0003" for c in out.summary["render_conflicts"])
    assert not out.summary["delivery_ready"]


def test_zip_input(tmp_path):
    src = make_docx(tmp_path / "Manuscript.docx")
    z = tmp_path / "sub.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(src, "IJO_1/Manuscript.docx")
    out = pipeline.edit(str(z), offline_style(), str(tmp_path / "out"), stages=("style",))
    assert out.stem == "Manuscript" and (tmp_path / "out" / "source" / "Manuscript.docx").exists()
    assert list(out.summary["stage_info"]) == ["style"]


def test_cli_edit_and_render_dir(tmp_path, monkeypatch, capsys):
    st = offline_style()
    monkeypatch.setattr(pipeline.style_mod, "load", lambda code: st)
    src = make_docx(tmp_path / "ms.docx")
    outdir = str(tmp_path / "o")
    assert cli.main(["edit", src, "ijo", outdir, "--accept-all"]) == 0
    printed = capsys.readouterr().out
    assert "0 pending" in printed and "ms_tracked.docx" in printed
    assert "color" in docmodel.load(f"{outdir}/ms_clean.docx").paragraphs[1].text  # "Introduction" deleted
    assert cli.main(["render", outdir, "--final"]) == 0


def test_cli_render_with_queries_and_format_changes(tmp_path):
    src = make_docx(tmp_path / "ms.docx")
    changes = ChangeSet([
        Change("c1", "p0003", 4, 10, "colour", "color", Layer.LANGUAGE, Source.RULE, status=Status.ACCEPTED),
        Change("c2", "p0003", 0, 3, "The", "The", Layer.HOUSE_STYLE, Source.RULE, status=Status.ACCEPTED,
               format={"italic": True}),
        Change("c3", "p0005", 0, 2, "We", "Us", Layer.LANGUAGE, Source.RULE, status=Status.REJECTED),
    ])
    changes.dump(tmp_path / "c.json")
    (tmp_path / "q.json").write_text(json.dumps([
        {"para_id": "p0003", "start": 4, "end": 10, "text": "Kindly check.", "check": "x"},
        {"para_id": "p0005", "start": 0, "end": 0, "text": "Dropped.", "check": "x", "status": "rejected"},
    ]))
    out = tmp_path / "r"
    assert cli.main(["render", src, str(tmp_path / "c.json"), str(out), "--queries", str(tmp_path / "q.json")]) == 0
    t = zipfile.ZipFile(out / "ms_tracked.docx")
    assert "AQ1: Kindly check." in t.read("word/comments.xml").decode()
    xml = t.read("word/document.xml").decode()
    assert "rPrChange" in xml and "Us" not in xml
    clean = docmodel.load(str(out / "ms_clean.docx")).paragraphs
    assert clean[2].text.startswith("The color") and clean[4].text.startswith("We")
    assert (out / "ms_queries.xlsx").exists()
