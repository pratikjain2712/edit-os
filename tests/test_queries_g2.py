"""Author queries: what they ask, when, and where they sit (invented text throughout)."""

from docx import Document

from edit_os import docmodel, frontmatter as fm, queries, refengine, rules, rules_sections, style
from edit_os.copyedit import abbrev


def make(tmp_path, paras, name="m.docx"):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def offline(code):
    st = style.load(code)
    st.settings["references"]["lookup"] = False
    return st


# --- 1, 2: reference queries -------------------------------------------------------------------

def test_volume_and_pages_read_from_the_entry():
    vp = queries.volume_pages
    assert vp("Rao K, Iyer M. A lid lesion. Int Surg J.  2025;12:991-3.") == (True, True)  # double space
    assert vp("Rao K. A trial. Nurs Today. 2020;7(2):87-96. DOI:10.18502/abc.v7i2.273") == (True, True)
    assert vp("Rao, K., A survey. Journal of Care, 2023. 13(1): p. 1-10.") == (True, True)  # EndNote
    assert vp("Rao K, Iyer M. (2021). A title. Journal Name, 12(3), 991–993.") == (True, True)  # APA
    assert vp("Rao K. A title. J Educ. 2025;12.") == (True, False)
    assert vp("Rao K. (2025). A title: 2 groups, 3 arms. Journal Name.") == (False, False)
    assert vp("Rao K. A title. J Educ 2025. doi:10.1000/123:45-6") == (False, False)  # DOI digits are no volume


class _Ref:
    def __init__(self, pid, text):
        self.para_id, self.text, self.offset, self.number = pid, text, 0, None


def _lookups(monkeypatch, pairs):
    monkeypatch.setattr(rules, "lookup_references", lambda ms, st: pairs)


def _res(status, message, text, **kw):
    return refengine.Resolution(status, refengine.parse(text), message=message, **kw)


def test_no_volume_record_but_the_entry_has_them(tmp_path, monkeypatch):
    text = "Rao K, Iyer M. A lid lesion in the young. Int Surg J.  2025;12:991-3."
    lacking = "Rao K, Iyer M. Another lid lesion in the young. Int Surg J. 2025."
    ms = docmodel.load(make(tmp_path, ["References", text, lacking]))
    p1, p2 = ms.paragraphs[1].id, ms.paragraphs[2].id
    msg = "Record has no volume yet (ahead of print?)"
    _lookups(monkeypatch, [(_Ref(p1, text), _res(refengine.QUERY, msg, text)),
                           (_Ref(p2, lacking), _res(refengine.QUERY, msg, lacking))])
    got = {(q.para_id, q.check) for q in queries.collect(ms, offline("jehp")).queries if q.check.startswith("references")}
    assert (p1, "references.no_volume") not in got and (p1, "references.no_pages") not in got
    assert (p2, "references.no_volume") in got


def test_et_al_with_few_authors_asked_besides_another_reference_query(tmp_path, monkeypatch):
    text = "Rao K, Iyer M, Das P, et al. A survey of nurses. J Nurs Pract. 2017;3(1):22-25. DOI:10.1000/xyz.1"
    ms = docmodel.load(make(tmp_path, ["References", text]))
    pid = ms.paragraphs[1].id
    _lookups(monkeypatch, [(_Ref(pid, text), _res(refengine.QUERY, "DOI 10.1000/xyz.1 points to a different article",
                                                  text))])
    got = [q.text for q in queries.collect(ms, offline("jehp")).queries if q.para_id == pid]
    assert got == ["The DOI given points to a different article. Kindly check the reference.",
                   "Kindly provide minimum 6 author names before et al."]


def test_reference_with_nothing_to_find_it_by_is_queried(tmp_path, monkeypatch):
    text = "Rao K, Iyer M. A short history of the words we use."
    ms = docmodel.load(make(tmp_path, ["References", text]))
    pid = ms.paragraphs[1].id
    _lookups(monkeypatch, [(_Ref(pid, text), _res(refengine.NOT_JOURNAL, "Not a recognisable reference type; left "
                                                                         "as written", text))])
    got = [q for q in queries.collect(ms, offline("ijo")).queries if q.para_id == pid]
    assert [q.text for q in got] == ["This reference is incomplete (no journal name, year, volume and page numbers). "
                                     "Kindly provide the complete details."]


# --- 3: abstract word count --------------------------------------------------------------------

def test_abstract_words_counted_as_word_does():
    assert rules_sections._abstract_words("Background: COVID-19 rose 3.5% (p<0.05) in 2020 – a rise.") == 8
    assert rules_sections._abstract_words("Methods") == 0
    assert rules_sections._abstract_words("Materials and Methods: Nurses were surveyed.") == 3


def test_abstract_of_240_words_is_not_queried(tmp_path):
    body = " ".join(["well-being 4.5%"] * 59)  # 118 Word words; 236 by \w+ splitting
    paras = ["Abstract", "Background: " + body, "Results: " + body + " end.", "Keywords: nursing; care; trial",
             "Introduction", "Text."]
    ms = docmodel.load(make(tmp_path, paras))
    qs = queries.collect(ms, offline("jehp"))
    assert "abstract.too_long" not in {q.check for q in qs.queries}
    paras[2] = "Results: " + body + " " + " ".join(["more"] * 20)
    ms = docmodel.load(make(tmp_path, paras, "m2.docx"))
    assert "abstract.too_long" in {q.check for q in queries.collect(ms, offline("jehp")).queries}


# --- 4, 6: front matter ------------------------------------------------------------------------

def _fields(authors, running_title=""):
    return {
        "article_type": "Original Article", "title": "Lid lesions in the young", "running_title": running_title,
        "authors": [{"given_names": g, "surname": s, "prefix": "", "affiliations": [1], "corresponding": i == 0}
                    for i, (g, s) in enumerate(authors)],
        "affiliations": [{"id": 1, "department": "Department of Ophthalmology", "institution": "City Eye Hospital",
                          "city": "Chennai", "state": "Tamil Nadu", "country": "India"}],
        "correspondence": {}, "proper_nouns": [],
        "paragraphs": [{"id": "M1", "role": "title"}, {"id": "M2", "role": "authors"}, {"id": "M3", "role": "affiliation"}],
    }


def test_full_name_query_names_the_author_on_the_author_group(tmp_path, monkeypatch):
    raw = _fields([("Asha", "Kumar"), ("S", "Natarajan")], running_title="Lid lesions")
    monkeypatch.setattr(fm, "split", lambda text: raw)
    ms = docmodel.load(make(tmp_path, ["Lid lesions in the young", "Asha Kumar1, S Natarajan1",
                                       "1Department of Ophthalmology, City Eye Hospital, Chennai, Tamil Nadu, India",
                                       "Introduction", "Text."]))
    pl = fm.plan(ms, offline("ijo"), fm.Sources())
    q = next(q for q in pl.queries if q.check == "front.full_name")
    assert q.text == "Kindly provide the full name of author “S Natarajan”."
    if q.insert_id:
        assert q.quote == "S Natarajan"
    else:
        p = ms.by_id(q.para_id)
        assert "Natarajan" in p.text[q.start:q.end] and p.text.startswith("Asha Kumar")


def test_running_title_over_50_characters_is_queried(tmp_path, monkeypatch):
    long = "Simulation and game based learning for resuscitation training"  # 61 characters
    raw = _fields([("Asha", "Kumar")], running_title=long)
    raw["paragraphs"] = [{"id": "M1", "role": "title"}, {"id": "M2", "role": "running_title"},
                         {"id": "M3", "role": "authors"}]
    monkeypatch.setattr(fm, "split", lambda text: raw)

    def plan(title, name):
        paras = ["Lid lesions in the young", f"Running title: {title}", "Asha Kumar1", "Introduction", "Text."]
        raw["running_title"] = title
        return fm.plan(docmodel.load(make(tmp_path, paras, name)), offline("jehp"), fm.Sources())

    q = next(q for q in plan(long, "m.docx").queries if q.check == "front.running_head_long")
    assert q.text == "The running title has more than 50 characters. Kindly shorten it to 50 characters or fewer."
    assert q.insert_id and q.quote.startswith("Simulation")  # on the running head, not the deleted label line
    assert not any(q.check == "front.running_head_long" for q in plan("Game learning for resuscitation", "m2.docx").queries)


def test_insert_comment_written_on_the_inserted_paragraph(tmp_path):
    ms = docmodel.load(make(tmp_path, ["Title of the paper", "Body text."]))
    n = fm.NewParagraph("fi1", ms.paragraphs[0].id, "before", [("<RH>Kumar: Short", {})], "running_head")
    qs = queries.QuerySet([queries.Query(n.anchor, 11, 16, "Kindly check.", "x", "AQ1", insert_id="fi1",
                                         quote="Short")])
    assert queries.write_comments(ms, qs, deferred={"fi1"}) == 0
    els = {}
    fm.render_inserts(ms, [n], statuses=(n.status,), elements=els)
    assert queries.write_insert_comments(ms, qs, els) == 1
    assert els["fi1"].xpath(".//w:commentRangeStart")


# --- 5: abbreviation scopes --------------------------------------------------------------------

def test_key_words_under_a_bare_heading_are_not_main_text(tmp_path):
    paras = ["Abstract", "Toxic anterior segment syndrome (TASS) followed lens surgery in one eye.", "Keywords",
             "Toxic anterior segment syndrome, Implantable Phakic Contact Lens (IPCL), lens surgery",
             "Introduction", "We report a case after Implantable Phakic Contact Lens (IPCL) surgery. The IPCL was clear.",
             "Discussion", "The IPCL was removed."]
    ms = docmodel.load(make(tmp_path, paras))
    _, qs = abbrev.collect(ms, style.load("ijo"))
    assert not any(q.check == "abbrev.defined_twice" for q in qs)


def test_used_before_its_definition_asks_to_expand_at_first_use(tmp_path):
    paras = ["Introduction", "The PBL course ran for a term. Problem-based learning (PBL) is common.",
             "Methods", "The problem-based learning (PBL) sessions ran weekly. PBL was graded."]
    ms = docmodel.load(make(tmp_path, paras))
    _, qs = abbrev.collect(ms, style.load("ijo"))
    checks = [(q.check, q.text) for q in qs if "PBL" in q.text]
    assert ("abbrev.before_definition", "PBL: Kindly expand the abbreviation at first use.") in checks
    assert not any(c == "abbrev.defined_twice" for c, _ in checks)


# --- 6: abstract headings, drafting instructions --------------------------------------------------

def test_abstract_heading_off_the_sheet_list_is_queried(tmp_path):
    paras = ["Abstract", "Purpose: To report a series.", "Design: Retrospective series.", "Methods: Records reviewed.",
             "Results: Ten eyes.", "Conclusion: Rare.", "Keywords: eye, injury, orbit", "Introduction", "Text."]
    ms = docmodel.load(make(tmp_path, paras))
    got = [q for q in queries.collect(ms, offline("ijo")).queries if q.check == "abstract.heading"]
    assert [ms.by_id(q.para_id).text[q.start:q.end] for q in got] == ["Design"]
    assert got[0].text == ("The abstract heading “Design” is not one of the journal’s headings (Purpose, Methods, "
                           "Results, Conclusion). Kindly check and revise the abstract under these headings.")


def test_drafting_instruction_in_a_table_note():
    d = queries.drafting_instruction
    text = "Cells show rho (p-value). Replace “<.001” with exact software p-values if the original output gives them."
    s, e = d(text)
    assert text[s:e].startswith("Replace “<.001”") and text[s:e].endswith("gives them.")
    assert d("Data were analysed as planned [insert reference].") is not None
    assert d("We replaced the missing values with the mean.") is None
    assert d("Replace the dressing with a new one daily.") is None


# --- 7: conclusion -----------------------------------------------------------------------------

def test_conclusion_only_for_original_articles_and_not_when_run_in(tmp_path):
    body = ["Introduction", "Text.", "Methods", "Text.", "Results", "Text.", "Discussion", "Text."]
    ms = docmodel.load(make(tmp_path, body))
    assert "sections.conclusion" in {q.check for q in queries.collect(ms, offline("ijo")).queries}
    ms = docmodel.load(make(tmp_path, body + ["CONCLUSION: The plan works."], "m2.docx"))
    assert "sections.conclusion" not in {q.check for q in queries.collect(ms, offline("ijo")).queries}
    ms = docmodel.load(make(tmp_path, ["Introduction", "Text.", "Case Report", "Text.", "Discussion", "Text."], "m3.docx"))
    assert "sections.conclusion" not in {q.check for q in queries.collect(ms, offline("ijo")).queries}


# --- 8: one query per anchor, never on a deleted paragraph ---------------------------------------

def test_identical_queries_on_one_anchor_asked_once(tmp_path):
    ms = docmodel.load(make(tmp_path, ["Text one.", "Text two."]))
    qs = queries.QuerySet()
    for _ in range(2):
        qs.add(ms.paragraphs[0].id, "text.brackets", 0, 4)
    assert len(qs.number(ms).queries) == 1


def test_query_moves_off_a_deleted_paragraph(tmp_path):
    from edit_os.changes import Change, Layer, Source

    ms = docmodel.load(make(tmp_path, ["Running title: Lid lesions", "Body text."]))
    p0, p1 = ms.paragraphs
    qs = queries.QuerySet([queries.Query(p0.id, 0, 0, "Kindly check.", "x")])
    gone = Change(id="c1", para_id=p0.id, start=0, end=len(p0.text), old=p0.text, new="", layer=Layer.STRUCTURE,
                  source=Source.RULE, reason="front_matter")
    assert queries.keep_anchors_live(ms, qs, [gone]) == 1
    assert qs.queries[0].para_id == p1.id
