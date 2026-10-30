"""Copy-edit text rules (edit_os/rules_text.py) and the spelling fixes that came with them. Synthetic text only."""

import pytest
from docx import Document

from edit_os import copyedit, docmodel, rules_text, style, tracked
from edit_os.changes import Status
from edit_os.copyedit import langtool, spelling
from edit_os.rules_text import Context


def fix(text, rule, ctx=None, **cfg):
    """Apply one rule to a string."""
    out = text
    for s, e, new, _ in sorted(rules_text.propose(text, {rule: {"enabled": True, **cfg}}, ctx or Context()),
                               reverse=True):
        assert text[s:e] != new
        out = out[:s] + new + out[e:]
    return out


@pytest.mark.parametrize("before, after", [
    ("The scan was non-invasive and non-inferior.", "The scan was noninvasive and noninferior."),
    ("Vision improved post operatively and pre-operative IOP fell.",
     "Vision improved postoperatively and preoperative IOP fell."),
    ("A non-glaucomatous, non-durable response.", "A nonglaucomatous, nondurable response."),
    ("Intraarterial and antiinflammatory drugs.", "Intra-arterial and anti-inflammatory drugs."),
    # kept: capitals, numbers, same vowel, adverbs, open compounds, nouns after post-/pre-, other prefixes
    ("non-ICU, post-COVID, pre-2000 and non-Hodgkin cases", "non-ICU, post-COVID, pre-2000 and non-Hodgkin cases"),
    ("intra-arterial, anti-inflammatory, re-examined", "intra-arterial, anti-inflammatory, re-examined"),
    ("non-normally distributed data", "non-normally distributed data"),
    ("non-small cell lung cancer and non-insulin-dependent diabetes",
     "non-small cell lung cancer and non-insulin-dependent diabetes"),
    ("comfortable post-surgery and post-transfusion", "comfortable post-surgery and post-transfusion"),
    ("a semi-structured, multi-centre, anti-fibrotic tool", "a semi-structured, multi-centre, anti-fibrotic tool"),
    ("a post hoc analysis and the post office", "a post hoc analysis and the post office"),
])
def test_prefixes(before, after):
    assert fix(before, "prefixes") == after


def test_prefixes_spaced_and_same_vowel_switches():
    assert fix("seen post operatively", "prefixes", spaced=False) == "seen post operatively"
    assert fix("intraarterial", "prefixes", same_vowel_hyphen=False) == "intraarterial"


@pytest.mark.parametrize("before, after", [
    ("Vision was stable at one month and three months.", "Vision was stable at 1 month and 3 months."),
    ("She had a five-year history of pain.", "She had a 5-year history of pain."),
    ("Drops were stopped after two to three weeks.", "Drops were stopped after 2 to 3 weeks."),
    ("Three months later the cyst recurred.", "Three months later the cyst recurred."),  # sentence start
    ("Pain resolved.[4] Two weeks later it recurred.", "Pain resolved.[4] Two weeks later it recurred."),
    ("one of three eyes and two patients", "one of three eyes and two patients"),     # not a time unit
    ("over three decades", "over three decades"),
])
def test_number_words(before, after):
    assert fix(before, "number_words") == after


def test_stat_semicolons():
    text = ("IOP was 17.0 ± 2.5 versus 16.1 ± 2.3 mmHg (p = 0.065), at 3 months it was 17.1 ± 2.9 versus "
            "16.3 ± 2.3 mmHg (p = 0.165), and at final visit, 17.1 ± 2.7 versus 16.0 ± 2.8 mmHg (p = 0.124).")
    assert fix(text, "stat_semicolons") == text.replace("(p = 0.065),", "(p = 0.065);").replace(
        "(p = 0.165),", "(p = 0.165);")
    kept = "Age differed (p = 0.01), whereas sex did not (p = 0.40). Scores (r = −0.42, P = 0.04) and flux."
    assert fix(kept, "stat_semicolons") == kept


@pytest.mark.parametrize("before, after", [
    ("Most were between 61–70 years.", "Most were between 61 and 70 years."),
    ("Ages ranged from 18-81 years and from 1919- 1999.", "Ages ranged from 18 to 81 years and from 1919 to 1999."),
    ("prevalence between 65%–80% of cases", "prevalence between 65% and 80% of cases"),
    ("between the ages of 6-21 years", "between the ages of 6 and 21 years"),
    # kept: already in words, rates, ranges not after between / from
    ("between 5 and 10, from 2 to 3", "between 5 and 10, from 2 to 3"),
    ("increased from 2-3/day to 10-12/day", "increased from 2-3/day to 10-12/day"),
    ("aged 18–81 years", "aged 18–81 years"),
])
def test_ranges(before, after):
    assert fix(before, "ranges") == after


def test_lowercase():
    ctx = Context.from_texts(["the retinal detachment and vitreous haemorrhage of the right eye was seen on "
                              "examination; the anterior chamber and the vitreous were clear; retinal and eye"])
    assert fix("B-scan showed features suggestive of Retinal Detachment and Vitreous haemorrhage.",
               "lowercase", ctx) == "B-scan showed features suggestive of retinal detachment and vitreous haemorrhage."
    # a full form before its abbreviation: its words need not appear elsewhere when `suffix` is on
    assert fix("The lesion was confirmed on Magnetic Resonance Imaging (MRI) of the orbits and the brain.", "lowercase", suffix=True) == \
        "The lesion was confirmed on magnetic resonance imaging (MRI) of the orbits and the brain."
    assert fix("The lesion was confirmed on Magnetic Resonance Imaging (MRI) of the orbits and the brain.", "lowercase") == \
        "The lesion was confirmed on Magnetic Resonance Imaging (MRI) of the orbits and the brain."
    kept = ["Approved by the World Health Organization (WHO) and the Institutional Ethics Committee.",
            "Seen with Doppler Imaging (DI) and the Snellen chart.",  # eponyms
            "Grading used Group A and Visit 4 and Zone I.",           # labels
            "Retinal Detachment And Vitreous Haemorrhage",            # heading-like line
            "The drugs were Carboplatin, Vincristine, and Etoposide.",
            "the paper “Retinal Detachment in Adults” was cited."]
    for t in kept:
        assert fix(t, "lowercase", ctx) == t, t
    # a sentence start keeps its capital
    assert fix("Retinal Detachment was seen.", "lowercase", ctx) == "Retinal detachment was seen."
    # a single word needs to be written in lower case often enough elsewhere
    assert fix("noted in the Anterior chamber.", "lowercase", Context.from_texts(["anterior"]), min_lower=2) == \
        "noted in the Anterior chamber."


def test_protected_spans_and_references():
    assert fix("see https://x.org/non-invasive and doi:10.1000/non-invasive", "prefixes") == \
        "see https://x.org/non-invasive and doi:10.1000/non-invasive"
    ref = "Rao A. Non-invasive imaging. Eye 2001;15:1-9."
    assert fix(ref, "prefixes") == ref


def make(tmp_path, paras):
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    doc.save(tmp_path / "m.docx")
    return docmodel.load(str(tmp_path / "m.docx"))


def test_checker_follows_journal_switches_and_renders(tmp_path, monkeypatch):
    monkeypatch.setattr(langtool, "is_running", lambda url=None: False)
    monkeypatch.setattr(spelling, "load", lambda *a, **k: spelling.Variants())
    ms = make(tmp_path, ["Abstract", "Purpose: To study a non-invasive test.",
                         "Methods", "Eyes were seen at three months, between 61-70 years of age.",
                         "References", "1. Rao A. A non-invasive test at three months. Eye 2001;15:1-9."])
    ijo = copyedit.run_all(ms, style.load("ijo"), only=["text"]).changes.changes
    assert sorted((c.old, c.new, c.reason) for c in ijo) == [
        ("-", "", "text.prefixes"), ("-", " and ", "text.ranges"), ("three", "3", "text.number_words")]
    assert all(c.layer.value == "language" and c.source.value == "rule" for c in ijo)
    # ABR: copy editors keep prefixes and number words; ranges are on everywhere
    abr = copyedit.run_all(ms, style.load("abr"), only=["text"]).changes.changes
    assert [c.reason for c in abr] == ["text.ranges"]
    for c in ijo:
        c.status = Status.ACCEPTED
    from edit_os.changes import ChangeSet
    out = tmp_path / "clean.docx"
    tracked.render(docmodel.load(str(tmp_path / "m.docx")), ChangeSet(ijo), tracked=False).save(out)
    paras = [p.text for p in docmodel.load(str(out)).paragraphs]
    assert paras[1] == "Purpose: To study a noninvasive test."
    assert paras[3] == "Eyes were seen at 3 months, between 61 and 70 years of age."
    assert paras[5] == "1. Rao A. A non-invasive test at three months. Eye 2001;15:1-9."


VARCON = """\
# grayish <verified> (level 40)
A: grayish / B: greyish

# analog <verified> (level 10)
A B: analogue / A.: analog | analogous
A B: analogues / A.: analogs | analogous
"""


def test_spelling_in_compounds_and_extra_variants(tmp_path):
    v = spelling.parse(VARCON.splitlines())
    ms = make(tmp_path, ["Results", "Gross section showed greyish-white areas; prostaglandin analogues were used.",
                         "Pain was scored on a visual analogue scale."])
    st = style.load("ijo")
    out = {c.para_id: c for c in spelling.check(ms, st, v)}
    assert [(c.old, c.new) for c in spelling.check(ms, st, v)] == [("greyish", "grayish"), ("analogues", "analogs")]
    assert ms.paragraphs[2].id not in out  # the instrument's name is kept


def test_compound_dash_and_aged_years_old():
    # JEHP_73_26 QC: "Inconsistent en dash & hyphen use"; "aged between 65 and 70 years old": "old" redundant
    from edit_os.rules_text import Context, find_aged_years_old, find_compound_dash

    ctx = Context.from_texts(["We measured oral health-related quality of life."])
    t = "Poor oral health–related quality of life; a 2015–2020 study; drug–drug interactions."
    assert [(t[s:e], n) for s, e, n in find_compound_dash(t, {}, ctx)] == [("–", "-")]
    t = "The majority were aged between 65 and 70 years old. A 45 years old man."
    assert [(t[s:e], n) for s, e, n in find_aged_years_old(t, {}, ctx)] == [(" old", "")]
