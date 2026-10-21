"""tools/stylesheet_check.py: edit-level disagreements, rule retrieval, quote checks, grouping, the report
numbers and the (patched) Claude judge. All text here is synthetic; no client stylesheet or manuscript."""

import os
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))

import stylesheet_check as sc  # noqa: E402

from edit_os import docmodel, tracked  # noqa: E402
from edit_os.changes import Change, ChangeSet, Layer, Source, Status  # noqa: E402

SHEET = """Short Title | Sentence case and centred, without a final full stop
Keywords | Placed on their own line below the summary
Listed alphabetically; reorder them when they are not
Spelling
Language | AMERICAN
Citations | Citation numbers go in square brackets, superscript, placed after punctuation
Inside citation brackets an en dash range becomes a hyphen, for instance [3–5] becomes [3-5]
Reference List | Check volume and page numbers against PubMed
Title | Sentence case and bold / Capital after a colon; small words in lowercase
Journals | Doe J, Roe K. A made-up study of eye drops. J Abbr Ophthalmol 2020;5:10-7.
"""

GUIDE = """REFERENCES
Drop the repeated leading digits of the last page in a page range.
HEADINGS
Section heads are set in capitals and bold.
"""


@pytest.fixture
def index(tmp_path):
    (tmp_path / "s.txt").write_text(SHEET, encoding="utf-8")
    (tmp_path / "g.txt").write_text(GUIDE, encoding="utf-8")
    return (sc.RuleIndex(sc.load_rules(str(tmp_path / "s.txt"), "S")),
            sc.RuleIndex(sc.load_rules(str(tmp_path / "g.txt"), "G")))


def item(kind="missed", old="colour", editor="color", tool="colour", section="text", left="", article="A_1",
         context="", record=""):
    return sc.Item(article, section, kind, "p0001", old, editor, tool, left, context or f"the [[{old}]] of", [], record)


# --- disagreements ------------------------------------------------------------------------------

def test_regions_classify_missed_extra_conflict_and_agreement():
    o = "The colour was red (1). Results: the eyes improved."
    t = "The color was red.[1] Results: The eyes improved."
    u = "The colour was red [1]. Results: the eyes improved."
    rs = sc.regions(o, t, u)
    got = {(r.kind, r.old.strip(), r.editor.strip(), r.tool.strip()) for r in rs}
    assert ("missed", "colour", "color", "colour") in got
    assert ("missed", "the", "The", "the") in got
    assert any(k == "conflict" for k, *_ in got)  # "(1)." -> ".[1]" vs "[1]."
    caps = next(r for r in rs if r.old == "the")
    assert caps.left == ":"
    agreed = sc.regions("a (1) b", "a [1] b", "a [1] b", agreed=True)
    assert [(r.kind, r.old, r.editor) for r in agreed] == [(None, "(", "["), (None, ")", "]")]


def test_whitespace_runs_are_not_disagreements():
    assert sc.regions("a  b", "a b", "a  b") == []


def test_patterns_group_numbers_and_punctuation_context():
    a = item(old="24µm", editor="24 µm", tool="24µm")
    b = item(old="30µm", editor="30 µm", tool="30µm", article="A_2")
    assert a.pattern() == b.pattern()
    after_colon = item(old="the", editor="The", tool="the", left=":")
    after_word = item(old="the", editor="The", tool="the", left="and")
    assert after_colon.pattern() != after_word.pattern()


def test_reference_pairing_by_number_then_by_words():
    orig = [(1, "Smith J. Glaucoma in the young. Eye 2010;1:1-2."),
            (2, "Jones K. Retinal detachment surgery outcomes. Retina 2011;2:3-4.")]
    ed = {1: "Jones K. Retinal detachment surgery outcomes. Retina 2011;2:3-4.",  # the editor renumbered
          2: "Smith J. Glaucoma in the young. Eye 2010;1:1-2."}
    assert sc.pair_refs(orig, ed) == {1: 2, 2: 1}
    assert sc.pair_refs(orig, {1: orig[0][1]}) == {1: 1}


# --- corpus counts --------------------------------------------------------------------------------

def _article(items, editor_regions, texts):
    return {"article": items[0].article, "items": items, "editor_regions": editor_regions, "texts": texts,
            "stats": {}}


def test_grouping_and_editor_inconsistency_counts():
    # The editor changed "colour" twice and left it once elsewhere: that is both ways.
    a1 = _article([item(), item()], [("text", "", "colour", "color")] * 2,
                  {"text": ["the colour of", "the colour of", "a colour here"], "ref": []})
    a2 = _article([item(kind="extra", old="–", editor="–", tool="-", section="ref", article="A_2")], [],
                  {"text": [], "ref": ["2010;1:1–2"]})
    groups = sc.group_items([a1, a2])
    by = {g["pattern"][1]: g for g in groups}
    assert by["missed"]["n"] == 2 and by["missed"]["articles"] == ["A_1"]
    c = by["missed"]["consistency"]
    assert (c["editor_made_elsewhere"], c["editor_left_unchanged"], c["inconsistent"]) == (2, 1, True)
    c = by["extra"]["consistency"]
    assert (c["editor_made_elsewhere"], c["inconsistent"]) == (0, False)


def test_a_rare_exception_is_not_inconsistency():
    assert sc._both_ways(1, 84) is False
    assert sc._both_ways(3, 10) is True
    assert sc._both_ways(0, 5) is False


# --- retrieval and quotes ---------------------------------------------------------------------------

def test_load_rules_keeps_sections_and_splits_long_lines(tmp_path):
    p = tmp_path / "s.txt"
    p.write_text("Keywords | Heading bold\nShould be in alphabetical order\nTable\nA | " + " / ".join(
        f"rule number {i} about tables and captions in detail" for i in range(8)), encoding="utf-8")
    rs = sc.load_rules(str(p), "S")
    assert rs[1].text == "Keywords | Should be in alphabetical order"  # labelled by the line above
    assert sum(r.text.startswith("A | rule number") for r in rs) == 8


def test_retrieval_finds_the_rule_for_the_edit(index):
    sheet, guide = index
    dash = item(kind="extra", old="–", editor="–", tool="-", section="ref",
                context="Eye 2010;1:[[1– => tool: -]]2.")
    top = sc.retrieve(dash, sheet, guide, k_sheet=3, k_guide=1)
    assert any("en dash range becomes a hyphen" in r.text for r in top)
    assert any(r.id.startswith("G") and "page range" in r.text for r in top)
    assert any(r.text.startswith("Journals |") for r in top)  # the example entry, for any reference
    caps = item(old="the", editor="The", tool="the", left=":")
    assert "Capital after a colon" in sc.retrieve(caps, sheet, None, k_sheet=2)[0].text
    spelling = item(old="enrollment", editor="enrolment", tool="enrollment")
    assert any("AMERICAN" in r.text for r in sc.retrieve(spelling, sheet, None, k_sheet=3))


def test_quote_must_occur_in_the_retrieved_text():
    rules = ["Title | Sentence case and bold / Capital after a colon; small words in lowercase"]
    assert sc.quote_ok("Capital  after a colon;\nsmall words in lowercase", rules)
    assert sc.quote_ok("“capital after a colon”", ["Capital after a colon"])  # quotes and case normalised
    assert sc.quote_ok("Sentence case ... small words in lowercase", rules)  # an elided quote, in order
    assert not sc.quote_ok("small words in lowercase ... Sentence case", rules)
    assert not sc.quote_ok("Lower case after a colon", rules)
    assert not sc.quote_ok("", rules) and not sc.quote_ok("case", rules)


def _group(kind="missed", section="text", inconsistent=False, record=""):
    it = item(kind=kind, section=section, record=record)
    return {"pattern": it.pattern(), "n": 1, "articles": [it.article], "examples": [it],
            "consistency": {"editor_made_elsewhere": 1, "editor_left_unchanged": 1, "inconsistent": inconsistent,
                            "change": ""}}


def test_check_verdict_needs_a_real_quote_and_a_possible_class():
    rules = [sc.Rule("S1", "Language | AMERICAN spelling throughout")]
    v = {"verdict": "editor_missed", "rule_id": "S1", "rule_quote": "AMERICAN spelling throughout", "why": ""}
    assert sc.check_verdict(v, _group(), rules)["verdict"] == "editor_missed"
    bad = {**v, "rule_quote": "British spelling throughout"}
    j = sc.check_verdict(bad, _group(), rules)
    assert j["verdict"] == "unclear" and "quote not found" in j["note"] and j["rule_quote"] == ""
    # The tool cannot be in error where it made no edit; a tool gap needs an editor edit.
    assert sc.check_verdict({"verdict": "tool_error"}, _group("missed"), rules)["verdict"] == "unclear"
    gap = {**v, "verdict": "tool_gap"}
    assert sc.check_verdict(gap, _group("extra"), rules)["verdict"] == "unclear"
    assert sc.check_verdict(gap, _group("conflict"), rules)["verdict"] == "tool_gap"
    # editor_inconsistent only when the counts show it.
    inc = {"verdict": "editor_inconsistent"}
    assert sc.check_verdict(inc, _group(), rules)["verdict"] == "unclear"
    assert sc.check_verdict(inc, _group(inconsistent=True), rules)["verdict"] == "editor_inconsistent"
    assert sc.check_verdict({"verdict": "house_habit"}, _group(), rules)["verdict"] == "house_habit"


def test_reference_editor_missed_can_rest_on_the_record():
    rec = "[tool lookup: format_fixed] pubmed PMID 1: Moreau SS. Ocular vascular occlusion. Eye 2010;1:1-2"
    g = _group("extra", "ref", record=rec)
    v = {"verdict": "editor_missed", "rule_id": "", "rule_quote": "", "record_quote": "Moreau SS", "why": ""}
    assert sc.check_verdict(v, g, [])["verdict"] == "editor_missed"
    assert sc.check_verdict({**v, "record_quote": "Moreau S, Jonas"}, g, [])["verdict"] == "unclear"
    assert sc.check_verdict(v, _group("extra", "text", record=""), [])["verdict"] == "unclear"


# --- the judge and the report numbers -------------------------------------------------------------

class MemCache:
    def __init__(self):
        self.d = {}

    def get(self, k):
        return self.d.get(k)

    def put(self, k, v):
        self.d[k] = v


def test_judge_batches_caches_and_scores(index):
    sheet, guide = index
    gap = item(old="the", editor="The", tool="the", left=":")
    ref_items = [item(kind="extra", section="ref", old=o, editor=o, tool=n, article="A_1",
                      record="[tool lookup: format_fixed] pubmed PMID 9: Moreau SS. Title. Eye 2010;1:1-2")
                 for o, n in (("Moreau S", "Moreau SS"), ("2009", "2010"))]
    for r in ref_items:
        r.where, r.entry = "ref 4", ("Moreau S. Title. Eye 2009;1:1-2", "same", "Moreau SS. Title. Eye 2010;1:1-2")
    err = item(kind="extra", old="Retina", editor="Retina", tool="retina")
    art = _article([gap, gap, err] + ref_items, [("text", ":", "the", "The")] * 2,
                   {"text": [": the a", ": the b"], "ref": []})
    art["stats"] = {"text:editor_edits": 10, "text:matched": 6, "text:missed": 2, "text:extra": 1,
                    "ref:editor_edits": 0, "ref:matched": 0, "ref:extra": 2}
    groups = sc.group_items([art])
    calls = []

    def fake(model, payload, effort):
        calls.append(payload)
        out = []
        for block in payload.split("### Item ")[1:]:
            iid = block.split("\n", 1)[0]
            if "REFERENCE ENTRY" in block:
                out += [{"id": f"{iid}.d1", "verdict": "editor_missed", "rule_id": "", "rule_quote": "",
                         "record_quote": "Moreau SS", "why": "record"},
                        {"id": f"{iid}.d2", "verdict": "editor_missed", "rule_id": "", "rule_quote": "",
                         "record_quote": "Eye 2010;1:1-2", "why": "record"}]
            elif "'Retina'" in block:
                out.append({"id": f"{iid}.d1", "verdict": "tool_error", "rule_id": "", "rule_quote": "",
                            "record_quote": "", "why": "proper noun"})
            else:
                rid = next(line.split(":")[0].strip() for line in block.splitlines() if "Capital after a colon" in line)
                out.append({"id": f"{iid}.d1", "verdict": "tool_gap", "rule_id": rid,
                            "rule_quote": "Capital after a colon", "record_quote": "", "why": "rule"})
        return {"verdicts": out}

    cache = MemCache()
    info = sc.judge(groups, sheet, guide, call=fake, cache=cache, workers=1, log=lambda *a: None)
    assert info == {"units": 3, "requests": 1}  # two patterns + one reference entry, in one request
    assert len(calls) == 1 and "REFERENCE ENTRY (ref 4)" in calls[0] and "Rules:" in calls[0]
    verdicts = {g["pattern"][4] if g["pattern"][1] == "missed" else g["pattern"][5]: g["judgement"]["verdict"]
                for g in groups}
    assert verdicts == {"The": "tool_gap", "Moreau SS": "editor_missed", "#": "editor_missed",
                        "retina": "tool_error"}
    gap_group = next(g for g in groups if g["judgement"]["verdict"] == "tool_gap")
    assert gap_group["judgement"]["rule_text"].endswith("Capital after a colon; small words in lowercase")

    s = sc.summarise("ijo", [art], groups)
    assert (s["editor_edits"], s["matched"], s["caught"]) == (10, 6, 2)
    assert s["raw_score"] == 0.6 and s["adjusted_score"] == round(8 / 12, 4)
    assert s["classes_by_occurrence"]["tool_gap"]["missed"] == 2
    assert s["fix_list"][0]["n"] == 2 and s["fix_list"][0]["judgement"]["rule_quote"] == "Capital after a colon"
    md = sc.markdown(s, groups, {"calls": 1, "cost": 0.001})
    assert "Adjusted score" in md and "Capital after a colon" in md
    assert "<table>" in sc.html_report(s, md)

    # A second run is answered from the cache.
    groups2 = sc.group_items([art])
    sc.judge(groups2, sheet, guide, call=fake, cache=cache, workers=1, log=lambda *a: None)
    assert len(calls) == 1
    assert [g["judgement"]["verdict"] for g in groups2] == [g["judgement"]["verdict"] for g in groups]


def test_judge_failure_leaves_patterns_unclear(index):
    sheet, guide = index
    groups = sc.group_items([_article([item()], [], {"text": [], "ref": []})])

    def boom(*a):
        raise ConnectionError("offline")

    cache = MemCache()
    sc.judge(groups, sheet, guide, call=boom, cache=cache, workers=1, log=lambda *a: None)
    assert groups[0]["judgement"]["verdict"] == "unclear" and "ConnectionError" in groups[0]["judgement"]["note"]
    assert cache.d == {}  # failures are not cached


def test_stylesheet_file_names_map_to_journal_codes(tmp_path):
    assert sc.stylesheet_codes("StyleSheet_Sample_IJN & DSHMJ.txt") == ["ijn_dshmj", "ijn_dshmj"]
    assert sc.stylesheet_codes("SGMJ journal stylesheet [updated 2Jan26].txt") == ["sgmj"]
    for name in ("StyleSheet_Sample_IJO.txt", "StyleSheet_Lung_New.txt", "StyleSheet_Sample_NI_New.txt"):
        (tmp_path / name).write_text("x", encoding="utf-8")
    assert sc.find_stylesheet("ijo", str(tmp_path)).endswith("IJO.txt")
    assert sc.find_stylesheet("lungindia", str(tmp_path)).endswith("Lung_New.txt")
    assert sc.find_stylesheet("neurol", str(tmp_path)).endswith("NI_New.txt")
    assert sc.find_stylesheet("xyz", str(tmp_path)) is None
    assert sc.journal_of("/d/Neurol-India-D-24-1") == "ni"


# --- one article end to end (rules run offline) -----------------------------------------------------

def test_run_article_lists_the_editors_edits_the_tool_did_not_make(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph("The colour of the optic disc was pale in all eyes we examined.")
    doc.add_paragraph("We enrolled forty patients over two years at one centre.")
    orig = tmp_path / "o.docx"
    doc.save(orig)
    ms = docmodel.load(str(orig))
    p = next(p for p in ms.paragraphs if p.text.startswith("The colour"))
    start = p.text.index("colour")
    cs = ChangeSet([Change("e1", p.id, start, start + 6, "colour", "color", Layer.LANGUAGE, Source.EDITOR,
                           status=Status.ACCEPTED)])
    te = tmp_path / "te.docx"
    tracked.render(docmodel.load(str(orig)), cs, tracked=True).save(te)

    r = sc.run_article("IJO_1_26", orig.read_bytes(), te.read_bytes(), "ijo")
    missed = [(i.kind, i.old, i.editor) for i in r["items"] if i.kind == "missed"]
    assert ("missed", "colour", "color") in missed
    st = r["stats"]
    assert st["text:editor_edits"] == st.get("text:matched", 0) + st.get("text:missed", 0) + st.get("text:conflict", 0)
    assert ("text", "", "colour", "color") in r["editor_regions"]


def test_a_moved_citation_is_marked_as_one_move(tmp_path):
    doc = Document()
    doc.add_paragraph("Smith et al. found the same in older patients with glaucoma.[4] Our results agree.")
    orig = tmp_path / "o.docx"
    doc.save(orig)
    p = docmodel.load(str(orig)).paragraphs[0]
    ins = p.text.index("al.") + 3
    dele = p.text.index("[4]")
    cs = ChangeSet([Change("e1", p.id, ins, ins, "", "[4]", Layer.LANGUAGE, Source.EDITOR, status=Status.ACCEPTED),
                    Change("e2", p.id, dele, dele + 3, "[4]", "", Layer.LANGUAGE, Source.EDITOR, status=Status.ACCEPTED)])
    te = tmp_path / "te.docx"
    tracked.render(docmodel.load(str(orig)), cs, tracked=True).save(te)
    r = sc.run_article("IJO_2_26", orig.read_bytes(), te.read_bytes(), "ijo")
    moved = [i for i in r["items"] if i.moved]
    assert len(moved) == 2 and {i.pattern()[1] for i in moved} == {"missed (move)"}
    g = {"pattern": moved[0].pattern(), "n": 1, "articles": ["IJO_2_26"], "examples": [moved[0]],
         "consistency": {"editor_made_elsewhere": 0, "editor_left_unchanged": 0, "inconsistent": False, "change": ""}}
    assert "Part of a MOVE" in sc._diff_line(g, "d1")
