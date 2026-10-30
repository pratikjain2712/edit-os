"""Headings and formatting as they come out (invented text): every paragraph the tool tags, renames, splits off or
inserts as a heading is bold, at its level's size and in the journal's font (rules_format.finish_headings); formatting
changes merge; statements under an inserted heading, figure caption text and a correspondence name are plain; running
text leaves Word's "Heading N" styles; sub-heading tagging (rules_sections.rule_subheadings); structured-abstract labels
with a dash; new paragraphs aligned."""

import copy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from edit_os import docmodel, frontmatter, rules, style, tracked
from edit_os.changes import Layer, Status

BODY = ("The participants were recruited from three schools and completed the questionnaires during class "
        "time with the help of trained staff who explained every item to them in detail over two weeks.")
FRONT = [("Abstract",), "Background: " + BODY, "Results: " + BODY, "Keywords: Health, nursing, students"]
REFS = ["References", "1. Lee K, Rao S. Back pain in adolescents. Spine 1997;22:1380-4."]


def _st(code):
    st = style.load(code)
    st.settings = copy.deepcopy(st.settings)
    st.settings["references"]["lookup"] = False
    return st


def _add(doc, item):
    """A paragraph: plain text, ("bold",), or [(text, {"bold": .., "italic": ..}), ...] runs."""
    if isinstance(item, tuple):
        doc.add_paragraph().add_run(item[0]).bold = True
        return doc.paragraphs[-1]
    if isinstance(item, list):
        p = doc.add_paragraph()
        for text, fmt in item:
            r = p.add_run(text)
            r.bold, r.italic = fmt.get("bold"), fmt.get("italic")
        return p
    return doc.add_paragraph(item)


def _make(tmp_path, items, name="m.docx"):
    doc = Document()
    for it in items:
        _add(doc, it)
    path = str(tmp_path / name)
    doc.save(path)
    return path


def _render(path, st, tracked_copy=False):
    cs = rules.run_all(docmodel.load(path), st)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    ms = tracked.render(docmodel.load(path), cs, tracked=tracked_copy)
    out = path.replace(".docx", "_t.docx" if tracked_copy else "_c.docx")
    ms.save(out)
    return cs, out


def _on(el):
    return el is not None and (el.get(qn("w:val")) or "true").lower() not in ("0", "false", "off", "none")


def _paras(path):
    """[(text, style id, alignment, [(run text, bold, size pt, font)])] of a rendered file, bold and size through the
    run, its character style, the paragraph style chain and the document defaults."""
    doc = Document(path)
    styles = {s.get(qn("w:styleId")): s for s in doc.styles.element.iter(qn("w:style"))}
    default = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")

    def chain(sid):
        out = []
        while sid in styles and len(out) < 10:
            out.append(styles[sid].find(qn("w:rPr")))
            b = styles[sid].find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None
        return out

    def look(rprs, tag):
        for r in rprs:
            if r is not None and r.find(qn(tag)) is not None:
                return r.find(qn(tag))
        return None

    out = []
    for p in doc.element.body.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        sid = ps.get(qn("w:val")) if ps is not None else "Normal"
        jc = ppr.find(qn("w:jc")) if ppr is not None else None
        runs = []
        for r in docmodel.text_runs(p):
            t = docmodel.run_text(r)
            if not t:
                continue
            rpr = r.find(qn("w:rPr"))
            rs = rpr.find(qn("w:rStyle")) if rpr is not None else None
            rprs = [rpr] + (chain(rs.get(qn("w:val"))) if rs is not None else []) + chain(sid) + [default]
            sz = look(rprs, "w:sz")
            f = look(rprs, "w:rFonts")
            runs.append((t, _on(look(rprs, "w:b")), int(sz.get(qn("w:val"))) / 2 if sz is not None else None,
                         f.get(qn("w:ascii")) if f is not None else None))
        out.append(("".join(r[0] for r in runs), sid, jc.get(qn("w:val")) if jc is not None else None, runs))
    return out


def _para(paras, start):
    return next(p for p in paras if p[0].startswith(start))


def _heading_ok(para, size):
    letters = [r for r in para[3] if r[0].strip()]
    return all(r[1] for r in letters) and all(r[2] == size and r[3] == "Times New Roman" for r in para[3])


# --- 1. formatting merges: a heading renamed keeps its bold --------------------------------------------------------

def test_a_formatting_change_merges_into_the_replacement(tmp_path):
    path = _make(tmp_path, ["Introduction", BODY, "Case report", BODY, "Discussion", BODY] + REFS)
    cs, out = _render(path, _st("ijo"))
    rep = next(c for c in cs.changes if c.old == "Case report" and not c.format_only)
    assert rep.new.endswith("Case Report") and rep.format.get("bold") is True and rep.format.get("size") == 14
    assert _heading_ok(_para(_paras(out), "<H1>Case Report"), 14)


# --- 2. a heading split off after a soft line break ----------------------------------------------------------------

def test_headings_split_off_a_paragraph_are_bold_at_their_size(tmp_path):
    def build(doc):
        for t in ["Introduction", BODY + " STEMI.\nMaterials and Methods", BODY, "Results", BODY]:
            doc.add_paragraph(t)
        doc.add_paragraph("Conclusion\n" + BODY).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for t in REFS:
            doc.add_paragraph(t)

    doc = Document()
    build(doc)
    path = str(tmp_path / "s.docx")
    doc.save(path)
    _, out = _render(path, _st("jehp"))
    paras = _paras(out)
    assert _heading_ok(_para(paras, "<H1>Materials and Methods"), 14)
    concl = _para(paras, "<H1>Conclusion")
    assert _heading_ok(concl, 14)
    after = paras[paras.index(concl) + 1]
    assert after[0] == BODY and not any(r[1] for r in after[3]) and all(r[2] == 12 for r in after[3])


# --- 3. back-matter headings: H2 at 12 pt, both of a run-on pair bold, the statement plain -------------------------

def test_back_matter_headings_bold_twelve_point_statements_plain(tmp_path):
    path = _make(tmp_path, FRONT + [
        "Introduction", BODY, "Discussion", BODY, BODY, "Conclusion", BODY, ("Conflicts of interest -Nil",),
        [("Financial Support:", {"bold": True}), (" Nil. ", {}), ("Conflicts:", {"bold": True}), (" None.", {})],
        "COMPLIANCE WITH JEHP AI POLICY STATEMENT: In accordance with the JEHP AI policy, the authors used no AI tools.",
        "Use of Artificial Intelligence:", "No artificial intelligence tools were used in writing this article at all.",
    ] + REFS)
    _, out = _render(path, _st("jehp"))
    paras = _paras(out)
    h2 = [p for p in paras if p[0].startswith("<H2>")]
    assert len(h2) >= 4 and all(_heading_ok(p, 12) for p in h2), [p[0] for p in h2]
    nil = next(p for p in paras if p[0] == "Nil")
    assert not any(r[1] for r in nil[3]) and all(r[2] == 12 for r in nil[3])
    assert any(p[0] == "<H2>Compliance with JEHP AI policy statement" for p in paras)  # S: "H2 | Sentence Case"
    assert any(p[0].startswith("<H2>Use of") for p in paras)


# --- 4. the abstract heading bold over the author's b=0 -------------------------------------------------------------

def test_abstract_heading_bold_over_an_explicit_not_bold(tmp_path):
    doc = Document()
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Abstract").bold = False
    for t in ["Background: " + BODY, "Results: " + BODY, "Keywords: Health, nursing", "Introduction", BODY] + REFS:
        doc.add_paragraph(t)
    path = str(tmp_path / "a.docx")
    doc.save(path)
    _, out = _render(path, _st("jehp"))
    assert _heading_ok(_para(_paras(out), "Abstract"), 14)


# --- 5. figure caption text plain -----------------------------------------------------------------------------------

def test_figure_caption_text_plain_label_bold(tmp_path):
    path = _make(tmp_path, FRONT + ["Introduction", BODY, "Discussion", BODY,
                            [("Figure 1:", {"bold": True, "italic": True}),
                             (" Study flow of the participants", {"bold": True, "italic": True})],
                            [("Figure 2:", {"bold": True}), (" Growth of ", {}), ("Candida albicans", {"italic": True})],
                            [("Table 1:", {"bold": True}), (" Characteristics of the participants", {})]] + REFS)
    cs, out = _render(path, _st("jehp"))
    fig = [c for c in cs.changes if c.reason == "fmt_figure_caption"]
    assert {tuple(sorted(c.format.items())) for c in fig} == {(("bold", False),), (("italic", False),)}
    paras = _paras(out)
    f1 = _para(paras, "Figure 1:")
    assert f1[3][0][1] and not any(r[1] for r in f1[3] if "Study" in r[0])
    doc = Document(out)
    runs = [r for p in doc.paragraphs if p.text.startswith("Figure 2") for r in p.runs]
    assert any(r.italic for r in runs if "andida" in r.text)  # an italic name inside the caption stays
    t1 = _para(paras, "Table 1:")
    assert all(r[1] for r in t1[3] if r[0].strip())  # G199: the table caption stays wholly bold


# --- 6. only the correspondence label is bold -------------------------------------------------------------------------

def test_text_a_rewrite_inserts_after_a_bold_label_is_plain(tmp_path):
    path = _make(tmp_path, [[("Address for correspondence:", {"bold": True}), (" Maryam Rao, Department of Eyes", {})]])
    p = docmodel.load(path).paragraphs[0]
    runs = [("Address for correspondence:", {"bold": True}), (" Dr. Maryam Rao, Department of Eyes", {})]
    cs = frontmatter._rewrite(p, runs, Layer.STRUCTURE)
    ins = [c for c in cs if c.new.strip() == "Dr." or "Dr." in c.new]
    assert ins and all(c.format == {"bold": False} for c in ins)


# --- 7. running text out of the "Heading 1" style; tag runs in the journal's font -------------------------------------

def test_running_text_in_a_heading_style_goes_to_normal(tmp_path):
    doc = Document()
    doc.add_paragraph("Introduction")
    doc.add_paragraph(BODY, style="Heading 1")
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("Discussion")
    r.font.name, r.font.size = "Arial", Pt(11)
    doc.add_paragraph(BODY)
    for t in REFS:
        doc.add_paragraph(t)
    path = str(tmp_path / "h.docx")
    doc.save(path)
    cs, out = _render(path, _st("jehp"))
    assert any(c.format_only and c.format.get("para_style") == "Normal" for c in cs.changes)
    paras = _paras(out)
    body = _para(paras, BODY[:30])
    assert body[1] == "Normal" and all(r[2] == 12 and r[3] == "Times New Roman" for r in body[3])
    disc = _para(paras, "<H1>Discussion")
    assert disc[1] == "Heading1" and _heading_ok(disc, 14)  # the heading keeps its style; its tag 14 pt
    _, tr = _render(path, _st("jehp"), tracked_copy=True)
    assert Document(tr).element.body.find(f".//{qn('w:pPrChange')}") is not None  # tracked


# --- 8. sub-heading tagging -----------------------------------------------------------------------------------------

def _texts(path, st):
    _, out = _render(path, st)
    return [p[0] for p in _paras(out)]


def test_plain_and_identifier_run_in_labels_stay_run_in(tmp_path):
    path = _make(tmp_path, FRONT + ["Introduction", BODY, "Materials and Methods", ("Study design",), BODY,
                            "Intervention Group (n=28): " + BODY, "Control Group (n=33): " + BODY,
                            [("Study/Project ID:", {"bold": True}), (" " + BODY, {})], "Results", BODY] + REFS)
    t = _texts(path, _st("jehp"))
    assert any(x.startswith("Intervention Group (n") for x in t)
    assert any(x.startswith("Study/Project ID: ") for x in t)


def test_bold_italic_heads_under_a_bold_head_are_third_level(tmp_path):
    path = _make(tmp_path, ["Introduction", BODY, ("Proposed System Architecture",), BODY,
                            [("User Interface Layer", {"bold": True, "italic": True})], BODY,
                            [("Learner Modeling Engine", {"bold": True, "italic": True})], BODY,
                            "Materials and Methods", [("Study design and setting", {"bold": True, "italic": True})],
                            BODY, "Results", BODY] + REFS)
    t = _texts(path, _st("jehp"))
    assert "<H2>Proposed system architecture" in t and "<H3>User interface layer" in t
    assert "<H2>Study design and setting" in t  # no bold roman head above it in its section


def test_word_numbered_bold_heads_lowercase_and_period_heads(tmp_path):
    doc = Document()
    for t in ["Introduction", BODY, "Materials and Methods"]:
        doc.add_paragraph(t)
    for head in ("Eligibility Criteria", "Search Strategy"):
        p = doc.add_paragraph(style="List Paragraph")
        p.add_run(head).bold = True
        num = OxmlElement("w:numPr")
        for tag, val in (("w:ilvl", "0"), ("w:numId", "7")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            num.append(el)
        p._p.get_or_add_pPr().append(num)
        doc.add_paragraph(BODY)
    doc.add_paragraph().add_run("study participants and sampling").bold = True
    doc.add_paragraph(BODY)
    doc.add_paragraph("Discussion")
    for head in ("Self-Efficacy", "Outcome Expectancy and outcome value associated with poor skills.", "Knowledge"):
        doc.add_paragraph().add_run(head).bold = True
        doc.add_paragraph(BODY)
    for t in REFS:
        doc.add_paragraph(t)
    path = str(tmp_path / "n.docx")
    doc.save(path)
    _, out = _render(path, _st("jehp"))
    paras = _paras(out)
    elig = _para(paras, "<H2>Eligibility criteria")
    assert elig[1] == "Normal"  # out of the list style
    assert all(x.find(qn("w:numPr")) is None for x in Document(out).element.body.iter(qn("w:pPr"))
               if "Eligibility" in "".join(t.text or "" for t in x.getparent().iter(qn("w:t"))))
    texts = [p[0] for p in paras]
    assert "<H2>Study participants and sampling" in texts
    assert "<H2>Outcome expectancy and outcome value associated with poor skills" in texts
    assert "<H2>Self-efficacy" in texts and "<H2>Knowledge" in texts


def test_label_only_head_over_its_run_in_parts(tmp_path):
    path = _make(tmp_path, ["Introduction", BODY, "Discussion", BODY, ("Strengths and Limitations:",),
                            [("Strengths:", {"bold": True}), (" " + BODY, {})],
                            [("Future Research:", {"bold": True}), (" " + BODY, {})], "Conclusion", BODY] + REFS)
    assert "<H2>Strengths and limitations" in _texts(path, _st("jehp"))


def test_review_italic_heads_under_the_main_heading_and_added_main_sections(tmp_path):
    path = _make(tmp_path, [("INTRODUCTION",), BODY, ("METHODS",), [("Data sources and search", {"italic": True})],
                            BODY, ("RESULTS",), BODY, ("OCULAR FINDINGS IN DISEASES",),
                            [("Trachoma", {"bold": True, "italic": True})], BODY,
                            [("Dry Eye Disease", {"bold": True, "italic": True})], BODY,
                            [("Diagnostic Patterns", {"italic": True})], BODY, ("DISCUSSION",), BODY] + REFS)
    t = _texts(path, _st("ijo"))
    assert "<H2>Data sources and search" in t  # right under "Methods": second level
    assert "<H1>Ocular Findings in Diseases" in t and "<H2>Trachoma" in t
    assert "<H3>Diagnostic patterns" in t  # italic under bold heads in its section


# --- 9. structured-abstract labels with a dash ---------------------------------------------------------------------

def test_abstract_labels_with_a_dash(tmp_path):
    path = _make(tmp_path, [("Abstract",), [("Purpose", {"bold": True}), ("–To report a series of cases.", {})],
                            [("Design-", {"bold": True}), ("Retrospective case series", {})],
                            [("Results", {"bold": True}), ("–Fifteen patients were included in the series.", {})],
                            "Keywords: Eye, trauma", "Introduction", BODY] + REFS)
    _, out = _render(path, _st("ijo"))
    paras = _paras(out)
    design = _para(paras, "Design")
    assert design[0] == "Design: Retrospective case series" and design[3][0][1]
    assert _para(paras, "Purpose")[0].startswith("Purpose: To report")


# --- 10. new paragraphs aligned -------------------------------------------------------------------------------------

def test_paragraphs_a_change_opens_take_the_journal_alignment(tmp_path):
    doc = Document()
    for t in ["Introduction", BODY, "Conclusion"]:
        doc.add_paragraph(t)
    p = doc.add_paragraph(BODY)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(" [18].\nAcknowledgement")
    doc.add_paragraph("The authors thank the staff of the clinic for their help with the study.")
    for t in REFS:
        doc.add_paragraph(t)
    path = str(tmp_path / "j.docx")
    doc.save(path)
    _, out = _render(path, _st("jehp"))
    ack = _para(_paras(out), "<H1>Acknowledgement")
    assert ack[2] == "both" and _heading_ok(ack, 14)  # JEHP justified (owner, 2026-10-30)
