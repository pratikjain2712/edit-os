"""Citations and references, round 2 (invented text): author-year articles numbered (owner, 2026-10-27), renumbered
citations keep their superscript, false and missed citations, citation spacing, reference cleanup (formatting,
corporate authors, record damage), number formatting (owner, 2026-10-27) and the non-breaking hyphen."""

import copy

from docx import Document
from docx.oxml import OxmlElement

from edit_os import (biblio, docmodel, queries, refengine, references, refs_other, refstyle, renumber, rules,
                     rules_guide, style, tracked)
from edit_os.changes import Status


def _style(journal="jehp"):
    st = style.load(journal)
    st.settings["references"]["lookup"] = False
    return st


def make(tmp_path, paras, name="m.docx", table=None):
    doc = Document()
    for i, item in enumerate(paras):
        p = doc.add_paragraph()
        for text, fmt in ([(item, {})] if isinstance(item, str) else item):
            run = p.add_run(text)
            for k, v in fmt.items():
                setattr(run.font, k, v)
        if i == 0 and table:
            t = doc.add_table(rows=1, cols=len(table))
            for cell, text in zip(t.rows[0].cells, table):
                cell.paragraphs[0].add_run(text)
    path = tmp_path / name
    doc.save(path)
    return str(path)


def run(path, st, only=None):
    cs = rules.run_all(docmodel.load(path), st, only=only)
    for c in cs.changes:
        c.status = Status.ACCEPTED
    out = str(path).replace(".docx", "_out.docx")
    tracked.render(docmodel.load(path), cs, tracked=False).save(out)
    tracked.render(docmodel.load(path), cs, tracked=True).save(str(path).replace(".docx", "_tr.docx"))
    return cs, [p.text for p in docmodel.load(out).paragraphs]


def para(text, superscript=()):
    return docmodel.Paragraph(id="p1", element=None, style="", text=text, superscript=list(superscript))


# --- 1. author-year articles numbered ------------------------------------------------------------------------------

AY_REFS = ["References",
           "Adams, J. (2020). A first title on school climate. Journal of Tests, 1(1), 1–9.",
           "Bravo, K., & Cole, M. (2019). A second title on teachers. Journal of Tests, 2(1), 3–14.",
           "Cole, M., Diaz, R., & Evans, T. (2021). A third title on bullying. Journal of Tests, 3(2), 5–16.",
           "Zed, Q. (2018). A title nobody cites. Journal of Tests, 4, 7–18."]


def test_author_year_article_numbered_in_order_of_citation(tmp_path):
    path = make(tmp_path, ["Introduction",
                           "Bullying persists in schools (Cole et al., 2021). Teachers matter "
                           "(Bravo & Cole, 2019; Adams, 2020).",
                           "Adams (2020) showed this, as Unknown et al. (2022) did."] + AY_REFS)
    st = _style()
    cs, t = run(path, st)
    assert t[1] == "Bullying persists in schools.[1] Teachers matter.[2,3]"
    assert t[2] == "Adams[3] showed this, as Unknown et al. (2022) did."
    head = t.index("References") if "References" in t else next(i for i, x in enumerate(t) if "References" in x)
    entries = [x for x in t[head + 1:head + 5]]
    assert [e.split(" ")[0].rstrip(",") for e in entries] == ["Cole", "Bravo", "Adams", "Zed"], t
    sup = [c for c in cs.changes if c.reason == "author_year"]
    assert all(c.format and c.format.get("superscript") for c in sup)
    assert any((c.format or {}).get("list_number") for c in cs.changes)  # JEHP: Word numbers the list
    qs = queries.collect(docmodel.load(path), st)
    texts = [q.text for q in qs.queries]
    assert any("Unknown et al. (2022)" in x for x in texts)
    assert any("Reference no. 4 has not been cited" in x for x in texts)
    assert any("numbered in order of first citation" in x for x in texts)


def test_author_year_group_keeps_the_unmatched_item(tmp_path):
    path = make(tmp_path, ["Introduction", "Overlooked (Adams, 2020; Missing et al., 2024)."] + AY_REFS)
    _, t = run(path, _style())
    assert t[1] == "Overlooked (Missing et al., 2024).[1]"


def test_numbered_list_is_left_to_the_numbered_rules(tmp_path):
    path = make(tmp_path, ["Introduction", "Text (Adams, 2020).", "References",
                           "1. Adams J. A title. J Test 2020;1:1-9."])
    from edit_os import authoryear
    assert not authoryear.plan(docmodel.load(path), _style()).apply


# --- 2. renumbered citations keep their superscript --------------------------------------------------------------

def test_renumbered_citation_keeps_superscript(tmp_path):
    path = make(tmp_path, ["Introduction", "First claim [2]. Second claim [1].", "References",
                           "1. Lee A, Kim B. Title one here. J Test 2020;1:1-9.",
                           "2. Park C, Choi D. Title two here. J Test 2021;2:1-9."])
    cs, t = run(path, _style())
    assert t[1] == "First claim.[1] Second claim.[2]"
    ren = [c for c in cs.changes if c.reason == "renumber" and "[" in c.new and c.para_id == "p0002"]
    assert ren and all(c.format and c.format.get("superscript") for c in ren)


# --- 3. false citations -------------------------------------------------------------------------------------------

def test_power_after_a_bracketed_sum_is_no_citation():
    assert rules.superscript_citations(para("The formula was: n = [(Z₁₋α/₂ + Z₁₋β)² × σ²] / δ².")) == []
    assert rules.superscript_citations(para("as Li et al. (2024)⁴ found"))  # a year's superscript still cites


def test_points_needs_a_scale():
    assert not rules.not_citation_before("such that cost and safety are the focal points", True, "5")
    assert rules.not_citation_before("each answer scores 4 points", True, "4")
    assert rules.not_citation_before("rated on a Likert scale where the highest level", True, "2")


def test_enumeration_in_abstract_is_no_citation():
    text = "four domains: (1) relational work; (2) moral framing; (3) policy; and (4) boundary-spanning collaboration."
    marks = rules.enumeration_markers(text)
    starts = {m.group(2): m.start() for m in rules._ROUND.finditer(text)}
    assert starts["4"] in marks and starts["1"] in marks


# --- 4. missed citations ------------------------------------------------------------------------------------------

NUM_REFS = ["References",
            "1. Gupta A, Rao B, Shah C. Art therapy trial. J Test 2024;1:1-9.",
            "2. Singh S, Sharma N, Patel A. Attitudes survey. J Test 2023;2:1-9.",
            "3. Patil R, Desai K, Kulkarni M. Caregiver study. J Test 2022;3:1-9."]


def test_bare_and_in_year_author_mentions_get_numbers(tmp_path):
    path = make(tmp_path, ["Introduction",
                           "Trials by Gupta et al. in 2024. Studies by Singh et al. 2023 and Patil et al. 2022 who "
                           "found the same."] + NUM_REFS)
    _, t = run(path, _style())
    assert t[1] == "Trials by Gupta et al.[1] in 2024. Studies by Singh et al.[2] and Patil et al.[3] who found the same."


def test_citation_cut_off_by_a_paragraph_break_goes_back(tmp_path):
    path = make(tmp_path, ["Introduction", "Earlier work [1] agrees with the vaginal microbiota",
                           "(2,3). Residing in rural areas mattered."] + NUM_REFS)
    _, t = run(path, _style())
    assert t[1] == "Earlier work[1] agrees with the vaginal microbiota.[2,3]"
    assert t[2] == "Residing in rural areas mattered."


def test_two_author_apa_mention_in_a_table(tmp_path):
    path = make(tmp_path, ["Introduction", "Shown in the table [1].", "References",
                           "[1] Corbett, A. T., & Anderson, J. R. (1994). Knowledge tracing. User Model, 4(4), 253–278."],
                table=["Corbett & Anderson (1994)", "BKT"])
    ms = docmodel.load(path)
    found = rules.author_year_mentions(ms, _style())
    assert [(a.number, a.new) for a in found] == [(1, "{n}")]


def test_unlabelled_entries_after_typed_labels_are_numbered(tmp_path):
    path = make(tmp_path, ["Introduction", "Guidance exists.¹⁻³ More.⁴"] + ["References",
                           "1. Lee A, Kim B. Title one here. J Test 2020;1:1-9.",
                           "2. Park C, Choi D. Title two here. J Test 2021;2:1-9.",
                           "Collins GS, Moons KG. Title three here. BMJ 2024;385:e1.",
                           "Liu X, Rivera SC. Title four here. BMJ 2020;370:m3."])
    st = _style("ijo")
    refs = references.extract(docmodel.load(path))
    assert [r.number for r in refs] == [1, 2, 3, 4]
    _, t = run(path, st)
    assert any(x.startswith("Guidance exists.[1-3] More.[4]") for x in t), t
    assert any(x.startswith("3. Collins GS") for x in t) and any(x.startswith("4. Liu X") for x in t)


def test_named_study_without_number_is_queried(tmp_path):
    path = make(tmp_path, ["Introduction", "Earlier work [1] and [2] exists.",
                           "A more recent study by Patil et al., on doctors found the same."] + NUM_REFS)
    qs = queries.collect(docmodel.load(path), _style())
    assert any("Patil et al." in q.text and q.check == "references.named_uncited" for q in qs.queries)


# --- 5. citation spacing ------------------------------------------------------------------------------------------

def test_citation_spacing_and_malformed_groups(tmp_path):
    refs = ["References"] + [f"{i}. Name{i} A. Title {i} here. J Test 2020;{i}:1-9." for i in range(1, 7)]
    path = make(tmp_path, ["Introduction", "Visual loss if untreated [1] .Subthreshold laser works.",
                           "Described first in 1964.[2]These tumours are rare.",
                           "Differences may affect outcomes. [3,4 5, 6],"] + refs)
    _, t = run(path, _style("ijo"))
    assert "Visual loss if untreated.[1] Subthreshold laser works." in t, t
    assert "Described first in 1964.[2] These tumours are rare." in t
    assert "Differences may affect outcomes.[3-6]" in t


# --- 6. PubMed / CrossRef damage -------------------------------------------------------------------------------

def _rec(**kw):
    base = dict(source="crossref", title="A title", authors=[("Lee", "A")], journal="J Test", year="2020",
                volume="1", issue="", pages="1-9")
    base.update(kw)
    return biblio.Record(**base)


def test_institutions_in_a_record_author_list_are_rejected():
    r = _rec(authors=[("International Islamic University Malaysia", ""), ("MOHD SHUKRI", "NA"), ("NORMAN", "MFH")])
    assert refengine.malformed_authors(r)
    assert not refengine.malformed_authors(_rec(authors=[("Lee", "A"), ("World Health Organization", "")]))


def test_compound_surname_kept_as_the_author_writes_it():
    assert refengine.compound_split("de Olaguer", "NP", "Pérez de Olaguer N")
    assert refengine.compound_split("Wijngaarden", "PV", "van Wijngaarden P")
    assert not refengine.compound_split("Naghavi Alhosseini", "SS", "Alhosseini SSN")
    p = refengine.parse("Skevas C, Pérez de Olaguer N, Lleó A. A study title. BMC Ophthalmol 2024;24:51.")
    r = _rec(authors=[("Skevas", "C"), ("de Olaguer", "NP"), ("Lleó", "A")])
    assert refengine._usable_authors(r, p, keep_compound=True)[1] == "Pérez de Olaguer N"  # IJO
    assert refengine._usable_authors(r, p)[1] == "de Olaguer NP"  # JEHP: the record's (its editors 4 of 5)


def test_trailing_provider_group_dropped():
    p = refengine.parse("Mansouri A, Sampat KM, Glaser BM. Efficacy of laser. Eye 2014;28:1418-24.")
    r = _rec(source="pubmed", authors=[("Mansouri", "A"), ("Sampat", "KM"), ("Glaser", "BM"), ("Medscape", "")])
    assert refengine._usable_authors(r, p) == ["Mansouri A", "Sampat KM", "Glaser BM"]


def test_article_numbers_and_volumes_as_written():
    p = refengine.parse("Wei J, Chen L. A case. J Int Med Res. 2025;53(8):03000605251363514.")
    assert refengine.record_pages(_rec(pages="3000605251363514", volume="53"), p) == "03000605251363514"
    p = refengine.parse("Singh A, Gupta N. A rarity. BMJ Case Rep. 2018 Nov;11(1):bcr-2018-225806.")
    assert refengine.record_pages(_rec(pages="bcr2018225806", volume="11"), p) == "bcr-2018-225806"
    out = refengine.format_record(_rec(volume="07", pages="143-145"), _style("ijo"), None)
    assert "2020;7:143-5" in out


def test_entry_tidy_before_lookup():
    assert refengine.corporate_author("Organization WH. Guidelines. Geneva: WHO; 2018.").startswith(
        "World Health Organization. Guidelines.")
    assert refengine.corporate_author("Control CfD, Prevention. STI guidelines, 2021.").startswith(
        "Centers for Disease Control and Prevention. STI")
    assert refengine.corporate_author("Weiss JS, Bouheraoua N. et al., IC3D classification. Cornea 2024;43:1-9.") \
        == "Weiss JS, Bouheraoua N, et al. IC3D classification. Cornea 2024;43:1-9."
    assert refengine.NLM_FIXES["J Educ Techno Soc"] == "J Educ Technol Soc"


def test_available_from_without_a_link_goes():
    fixed = refs_other.template("World Health Organization. A report. Geneva: World Health Organization; 2022. "
                                "Disponible sur: WHO.", "report")
    assert fixed.text == "World Health Organization. A Report. Geneva: World Health Organization; 2022."


def test_corporate_author_in_the_reference_rule(tmp_path):
    path = make(tmp_path, ["Introduction", "Guidance [1].", "References",
                           "1. Organization WH. Guidelines on sanitation and health. Geneva: WHO; 2018."])
    _, t = run(path, _style("ijo"), only=["references"])
    assert t[3] == "1. World Health Organization. Guidelines on sanitation and health. Geneva: WHO; 2018."


# --- 7. reference formatting -------------------------------------------------------------------------------------

def test_reference_italics_and_bold_cleared_et_al_kept(tmp_path):
    path = make(tmp_path, ["Introduction", "Text [1].", "References",
                           [("1. Lee A, Kim B, et al. ", {}), ("A trial of things.", {"italic": True}),
                            (" J Test 2020;", {}), ("12", {"bold": True}), (":1-9.", {})]])
    cs, _ = run(path, _style(), only=["references"])
    fmt = [(c.old, c.format) for c in cs.changes if c.reason == "ref_format"]
    assert ("A trial of things.", {"italic": False}) in fmt and ("12", {"bold": False}) in fmt


def test_apa_chapters_restyled_from_the_author_text():
    st = _style()
    out = refstyle.restyle_other("Branch, S., & Murray, J. P. (2018). Theoretical frameworks. In P. D’Cruz, & "
                                 "C. Rayner (Eds.), Concepts and Methods (Vol. 1, pp. 1–44). Springer.", st)
    assert out == ("Branch S, Murray JP. Theoretical frameworks. In: D’Cruz P, Rayner C, editors. Concepts and "
                   "Methods. Vol. 1. Springer; 2018. p. 1-44.")
    out = refstyle.restyle_other("Gutz, E., & Perren, S. (2022). The moral dimensions. In M. Killen & J. G. "
                                 "Smetana, Handbook of Moral Development (3rd ed., pp. 437–453). Routledge.", st)
    assert out == ("Gutz E, Perren S. The moral dimensions. In: Killen M, Smetana JG, editors. Handbook of Moral "
                   "Development. 3rd ed. Routledge; 2022. p. 437-53.")


def test_apa_one_name_author_and_no_half_conversion(tmp_path):
    t = ("Karyanti, Atmoko, A., & Hitipieuw, I. (2015). The effectiveness of training. Jurnal Pendidikan "
         "Humaniora, 3(2), 116–121.")
    res = refstyle.style_unverified(refengine.Resolution(refengine.QUERY, refengine.parse(t)), t, _style())
    assert res.formatted.startswith("Karyanti, Atmoko A, Hitipieuw I. The effectiveness of training.")
    path = make(tmp_path, ["Introduction", "Text [1].", "References",
                           "1. Piech, C., & Ng, A. (2015). Deep tracing. Advances in Things, 28, 505–513."])
    cs, _ = run(path, _style(), only=["references"])
    assert not [c for c in cs.changes if c.reason == "ref_initials"]


def test_proper_nouns_restored_in_titles():
    assert rules._restore_proper("Health workers organizing in tigray. Health Hum Rights 2025",
                                 "Health Workers Organizing in Tigray. Health", {"Tigray"}) \
        .startswith("Health workers organizing in Tigray.")
    assert rules._restore_proper("Prevalence among Indian adults", "Prevalence Among Indian Adults", set()) \
        == "Prevalence among Indian adults"


# --- 9. numbers -------------------------------------------------------------------------------------------------

def test_leading_zero_thousands_and_statistics_italic():
    cfg = {"leading_zero": True, "thousands_comma": True}
    got = rules_guide.propose("alpha was .87 (P < .05), n = 1,260 of 12,345.", cfg)
    assert [(s, e, new) for s, e, new, _ in got] == [(10, 10, "0"), (30, 31, "")]
    text = "F = 74.82, t(48) = 2.1, η² = 0.42, R² = 0.3, Cronbach’s α of 0.8; IL-6 r and the Fr = 2 index"
    spans = [text[s:e] for s, e in rules_guide.stat_italic_spans(text)]
    assert spans == ["F", "t", "η", "R", "α"]
    assert rules_guide.stat_italic_spans("Z value", True) == [(0, 1)]


def test_operator_spacing_in_tables_and_scale_anchors(tmp_path):
    path = make(tmp_path, ["Rated on a 5-point scale (1=very poor to 5=very good) by all.", "Then 28 (100 %)."],
                table=["Intervention (n=28)", "Mean±SD"])
    st = _style()
    _, t = run(path, st, only=["statistics", "guide_replace"])
    assert t[0] == "Rated on a 5-point scale (1 = very poor to 5 = very good) by all."
    ms = docmodel.load(str(tmp_path / "m_out.docx"))
    cells = [p.text for p in ms.paragraphs if p.in_table]
    assert cells == ["Intervention (n = 28)", "Mean ± SD"]


def test_p_value_head_capital(tmp_path):
    assert rules_guide.find_p_value("p Value", {}) == [(0, 1, "P")]


def test_non_breaking_hyphen_in_the_text_view(tmp_path):
    path = make(tmp_path, ["The GBL", "x"])
    doc = Document(path)
    p = doc.paragraphs[0]
    r = p.runs[0]
    nb = copy.deepcopy(r._r)
    for child in list(nb):
        if child.tag.endswith("}t"):
            nb.remove(child)
    nb.append(OxmlElement("w:noBreakHyphen"))
    r._r.addnext(nb)
    tail = p.add_run("WSQ group")
    doc.save(path)
    ms = docmodel.load(path)
    assert ms.paragraphs[0].text == "The GBL-WSQ group"
