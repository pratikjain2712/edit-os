"""Output defect scanner: check the files Edit OS delivers, independent of any editor's file.

The evaluation tools (evaluate.py, eval_final.py) score how close our text comes to the editors'; they never look at
the output itself. This tool runs the pipeline on each article (every suggestion accepted, as `--accept all`), then
checks the rendered tracked and clean copies and the records for defects an editor would see on opening them:

    qc.*         the pipeline's own QC report (errors and warnings, content.* preservation checks), render conflicts
    format.*     tagged headings not bold / not the journal's size, fonts other than the journal's, sizes off
    query.*      author queries on empty paragraphs or on deleted text, duplicates, broken texts ("???", "None", "{x}")
    cite.*       round / superscript numbers turned into citations next to scales, units, years, counts in tables,
                 footnote markers, numbers beyond the reference list
    content.*    whole body paragraphs deleted or emptied, paragraphs or headings appearing twice in the clean copy
    clean.*      leftovers in the clean copy: "???", stray tags, "[[", tracked-change residue, double spaces, space
                 before punctuation, the running head missing
    figtab.*     legends without "Figure N:", tables without a caption, captions not bold, tables not at the end
    refs.*       reference entries with a lower-case surname, "et al" without its period, empty entries, gaps
    signoff.*    the sign-off without the editor's name (expected when the run has no editor: counted, not an error)

Each finding: article, check, severity (error / warning / info), paragraph index and at most 8 words of masked
context (learn.snippet: names, numbers, e-mails and links masked). Manuscripts are client-confidential: keep the
output out of git (write it to a scratch directory).

Usage:
    python tools/defect_scan.py ARTICLE_DIR... --out OUT [--journal J] [--workers 4] [--reuse]
    python tools/defect_scan.py --out OUT --summary   # the per-journal / per-check tables of a finished run

ARTICLE_DIR holds original.zip (te.docx / ce.docx are not needed); a bare .zip is taken as its own article. The
journal comes from the folder name (IJO_1016_26 -> ijo) unless --journal is given. Results: OUT/<article>/ (the
pipeline's outputs), OUT/findings.jsonl (one finding per line), OUT/summary.json and the tables on stdout.
"""

from __future__ import annotations

import argparse
import collections
import concurrent.futures as cf
import glob
import json
import os
import re
import sys
import time
import zipfile
from dataclasses import asdict, dataclass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ERROR, WARNING, INFO = "error", "warning", "info"


@dataclass
class Finding:
    article: str
    check: str
    severity: str
    para: int | None = None  # index of the paragraph in the document checked (0-based), None: the whole file
    loc: str = ""  # at most 8 masked words
    doc: str = "clean"  # clean | tracked | records

    def to_dict(self) -> dict:
        return asdict(self)


_KEEP = {"figure", "figures", "fig", "table", "tables", "abstract", "introduction", "background", "methods", "method",
         "materials", "results", "discussion", "conclusion", "conclusions", "references", "reference", "keywords",
         "key", "words", "acknowledgement", "acknowledgements", "acknowledgments", "financial", "support",
         "sponsorship", "conflicts", "conflict", "interest", "interests", "patient", "consent", "declaration",
         "ethical", "ethics", "approval", "authors", "author", "contributions", "first", "legend", "legends",
         "case", "report", "description", "diagnosis", "the", "a", "an", "and", "of", "in", "for", "to", "on", "with",
         "kindly", "please", "provide", "check", "nil", "none", "et", "al", "dr", "mean", "sd", "or", "ci", "p", "n",
         "technical", "editing", "done", "copyediting", "correspondence", "e-mail", "email", "cite", "this",
         "article", "as", "submitted", "revised", "accepted", "study", "data", "statistical", "analysis", "aim",
         "aims", "objective", "objectives", "purpose", "participants", "setting", "design", "limitations",
         "strongly", "agree", "disagree", "neutral", "total", "score", "scale", "yes", "no"}


def snippet(text: str, n: int = 8) -> str:
    """At most `n` words of `text`, masked for the report: e-mails, links, digits (9) and every capitalised word
    that is not a common structural word ("Aa" / "AA"), the first word included."""
    text = re.sub(r"\S+@\S+", "<email>", text or "")
    text = re.sub(r"https?://\S+|www\.\S+", "<url>", text)
    text = re.sub(r"(<[A-Z]{1,3}\d?>)", r"\1 ", text)
    out = []
    for w in text.split()[:n]:
        if not re.fullmatch(r"</?[A-Z]{1,3}\d?>", w):
            w = re.sub(r"\d", "9", w)
        for core in re.findall(r"[^\W\d_][^\W_]*", w):
            if re.fullmatch(r"RH|H\d", core) or core.lower() in _KEEP or core.islower():
                continue
            w = w.replace(core, "AA" if core.isupper() else "Aa", 1)
        out.append(w)
    return " ".join(out)


def journal_of(name: str) -> str:
    from edit_os import style as style_mod

    return style_mod.resolve(re.split(r"[_\-]", os.path.basename(name.rstrip("/")))[0].lower())


# --- running the pipeline -----------------------------------------------------------------------------------------

def run_article(adir: str, journal: str, out: str, reuse: bool = False, editor: str | None = None) -> str:
    """Run the pipeline on `adir`/original.zip (or the zip `adir`) into OUT/<article>, accept every suggestion
    (changes and inserts; author queries stay as queries) and render the tracked and clean copies. No editor name:
    the sign-off keeps its placeholder, as on a job nobody signed."""
    from edit_os import pipeline, style as style_mod
    from edit_os.changes import Status

    name = article_name(adir)
    od = os.path.join(out, name)
    if reuse and glob.glob(os.path.join(od, "*_clean.docx")):
        return od
    src = adir if adir.lower().endswith(".zip") else os.path.join(adir, "original.zip")
    st = style_mod.load(journal)
    pipeline.edit(src, st, od, editor=editor)
    _, cs, ins, qd = pipeline.load_records(od)
    for item in list(cs.changes) + list(ins):
        if item.status == Status.SUGGESTED:
            item.status = Status.ACCEPTED
    pipeline.save_records(od, cs, ins, qd)
    pipeline.render_outputs(od, include_pending=False, style=st, editor=editor)
    return od


def article_name(adir: str) -> str:
    base = os.path.basename(adir.rstrip("/"))
    return base[:-4] if base.lower().endswith(".zip") else base


# --- reading the outputs ------------------------------------------------------------------------------------------

def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "").replace(" ", " ")).strip()


def _key(t: str) -> str:
    """A paragraph's text for matching: tags and punctuation spacing ignored, lower case."""
    return _norm(re.sub(r"<[A-Z]{1,3}\d?>", "", t)).lower()


def body_elements(doc_el):
    """Top-level <w:p> and <w:tbl> of the document body, in order."""
    body = doc_el.find(W + "body")
    return [el for el in body if el.tag in (W + "p", W + "tbl")]


class Styles:
    """Effective bold of each character, and font family / size (rules_format's style readers)."""

    def __init__(self, ms):
        from edit_os import rules_format as rf

        self.rf = rf
        self.bold = rf._Styles(ms)
        self.font = rf._FontStyles(ms)

    def bold_chars(self, p) -> list[bool]:
        from docx.oxml.ns import qn

        from edit_os.docmodel import run_text, text_runs

        ps = p.element.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        sid = ps.get(qn("w:val")) if ps is not None else self.font.default_para
        base = self.bold.props(sid).get("bold", False)
        out = []
        for r in text_runs(p.element):
            b = base
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                rs = rpr.find(qn("w:rStyle"))
                if rs is not None and "bold" in self.bold.props(rs.get(qn("w:val"))):
                    b = self.bold.props(rs.get(qn("w:val")))["bold"]
                d = self.rf._direct(rpr)
                if "bold" in d:
                    b = d["bold"]
            out += [bool(b)] * len(run_text(r))
        return out

    def font_chars(self, p) -> list[dict]:
        return self.rf.font_chars(p, self.font)


def _all_bold(p, styles: Styles, start: int = 0, end: int | None = None) -> float:
    """Share of the non-space characters of p.text[start:end] in bold (1.0 when there are none)."""
    b = styles.bold_chars(p)
    if len(b) != len(p.text):
        return 1.0
    idx = [i for i in range(start, len(p.text) if end is None else end) if not p.text[i].isspace()]
    return sum(b[i] for i in idx) / len(idx) if idx else 1.0


@dataclass
class Article:
    """Everything the checks read for one article."""
    name: str
    journal: str
    style: object
    clean: object  # docmodel.Manuscript
    source: object
    tracked_xml: bytes
    comments_xml: bytes | None
    clean_xml: bytes
    clean_parts: list[str]
    changes: list
    inserts: list
    qc: dict
    query_checks: collections.Counter
    imported: list = None  # texts of the paragraphs taken in from the submission's other files (supplied.py)

    @classmethod
    def load(cls, od: str, name: str | None = None, journal: str | None = None) -> "Article":
        from edit_os import docmodel, pipeline, style as style_mod

        man, cs, ins, _ = pipeline.load_records(od)
        paths = pipeline.file_paths(od, man["stem"])
        with open(paths["qc_json"], encoding="utf-8") as f:
            qc = json.load(f)
        z = zipfile.ZipFile(paths["tracked"])
        comments = z.read("word/comments.xml") if "word/comments.xml" in z.namelist() else None
        zc = zipfile.ZipFile(paths["clean"])
        j = journal or man["journal"]
        with open(paths["queries_json"], encoding="utf-8") as f:
            qchecks = collections.Counter(d["check"] for d in json.load(f) if d.get("status") != "rejected")
        return cls(name or os.path.basename(od.rstrip("/")), j, style_mod.load(j), docmodel.load(paths["clean"]),
                   docmodel.load(pipeline.source_path(od)), z.read("word/document.xml"), comments,
                   zc.read("word/document.xml"), zc.namelist(), list(cs.changes), list(ins), qc, qchecks,
                   _imported_texts(od, ins))


def _imported_texts(od: str, inserts) -> list[str]:
    """Paragraph texts of the accepted imports of other Word files (tables with their notes, supplementary
    material): two tables sent with the same notes ("p1: Comparison between ...") show them twice, as sent."""
    from docx.oxml.ns import qn

    from edit_os import docmodel, supplied
    from edit_os.changes import Status

    out = []
    for n in inserts:
        if not supplied.is_import(n) or n.status != Status.ACCEPTED:
            continue
        try:
            kids = list(docmodel.load(os.path.join(od, n.file)).document.element.body)
        except Exception:
            continue
        for block in n.blocks:
            for i in block:
                if i < len(kids):
                    els = [kids[i]] if kids[i].tag == qn("w:p") else list(kids[i].iter(qn("w:p")))
                    out += [docmodel.paragraph_text(e) for e in els]
    return out


# --- the checks -----------------------------------------------------------------------------------------------------

CHECKS: dict[str, str] = {}


def check(fn):
    CHECKS[fn.__name__.removeprefix("check_")] = (fn.__doc__ or "").strip().splitlines()[0]
    return fn


def _idx(ms) -> dict[str, int]:
    return {p.id: i for i, p in enumerate(ms.paragraphs)}


def _front_end(ms) -> int:
    """Index of the first paragraph of the abstract / main text in the clean copy (front matter before it)."""
    for i, p in enumerate(ms.paragraphs[:200]):
        t = re.sub(r"^<H\d>", "", p.text.strip())
        if re.match(r"(abstract|introduction|background|summary|dear\s+(editor|sir)|sir,|respected)\b", t, re.I) \
                or p.text.lstrip().startswith("<H1>"):
            return i
    return 0


def _refs_start(ms) -> int | None:
    from edit_os import references

    return references.find_heading(ms)


@check
def check_qc(a: Article, out: list):
    """The pipeline's QC report: every issue (content.* preservation included) and render conflicts."""
    idx = _idx(a.clean)
    found = {i["check"] for i in a.qc["qc"]["issues"]}
    # The manuscript's own reference problems are the author's to fix once queried: an uncited entry with its
    # "uncited" query, a citation without an entry with its "missing entry" query, and the order errors that follow
    # from them. Without the query they are a defect of ours (the author is never asked).
    queried = {"references.cited": a.query_checks.get("references.uncited", 0),
               "references.exists": a.query_checks.get("references.missing_entry", 0)}
    for i in a.qc["qc"]["issues"]:
        sev = ERROR if i["severity"] == "error" else WARNING
        if i["check"] == "signoff.name":
            sev = INFO  # no editor on this run: the placeholder is expected
        elif i["check"] in queried and queried[i["check"]]:
            sev = INFO
        elif i["check"] == "citations.order" and (queried["references.cited"] or queried["references.exists"]):
            sev = INFO
        out.append(Finding(a.name, "qc." + i["check"], sev, idx.get(i.get("para_id")), snippet(i["message"]),
                           "records"))
    for c in a.qc.get("render_conflicts", []):
        out.append(Finding(a.name, "qc.render_conflict", ERROR, None, c.get("output", ""), "records"))


_TAG = re.compile(r"^\s*<H(\d)>")


@check
def check_headings(a: Article, out: list):
    """Tagged <H1> / <H2> headings in bold (unless the journal unbolds the level) and at the journal's sizes."""
    st = a.style
    unbold = set(st.get("formatting.heading_unbold_levels", []) or [])
    bold_levels = set(st.get("formatting.heading_bold_levels", [1, 2, 3]) or [])
    want_bold = st.get("formatting.heading_bold_tagged", False)
    fonts = a.style.get("formatting.fonts", {}) or {}
    styles = Styles(a.clean)
    for i, p in enumerate(a.clean.paragraphs):
        m = _TAG.match(p.text)
        if not m or int(m.group(1)) > 2 or p.in_table:
            continue
        level, s = int(m.group(1)), m.end()
        e = len(p.text)
        run_in = re.search(r"[:.]\s+\S", p.text[s:])  # "<H2>Label: text" -- the label is the heading
        if run_in:
            e = s + run_in.start() + 1
        if not p.text[s:e].strip():
            out.append(Finding(a.name, "clean.stray_tag", ERROR, i, snippet(p.text)))
            continue
        if (want_bold and level in bold_levels or level == 1 and st.get("formatting.heading_bold", True)) \
                and level not in unbold:
            share = _all_bold(p, styles, s, e)
            if share == 0:
                out.append(Finding(a.name, "format.heading_not_bold", ERROR, i, snippet(p.text)))
            elif share < 1:
                out.append(Finding(a.name, "format.heading_part_bold", WARNING, i, snippet(p.text)))
        if fonts.get("enabled"):
            want = fonts.get("heading_size", 14) if level == 1 else fonts.get("size", 12)
            f = styles.font_chars(p)
            if len(f) == len(p.text):
                sizes = {f[k].get("size") for k in range(s, e) if not p.text[k].isspace()}
                if sizes - {want}:
                    out.append(Finding(a.name, "format.heading_size", ERROR, i,
                                       f"H{level} {'/'.join(str(x) for x in sorted(x for x in sizes if x))} pt "
                                       f"(want {want}): " + snippet(p.text, 5)))


@check
def check_fonts(a: Article, out: list):
    """Font family and size of every run against formatting.fonts (symbol fonts excepted); one finding per paragraph."""
    from edit_os import rules_format as rf

    cfg = a.style.get("formatting.fonts", {}) or {}
    if not cfg.get("enabled"):
        return
    family, size, big = cfg.get("family"), cfg.get("size", 12), cfg.get("heading_size", 14)
    styles = Styles(a.clean)
    front_end = _front_end(a.clean)
    for i, p in enumerate(a.clean.paragraphs):
        if not p.text.strip() or _TAG.match(p.text):
            continue
        f = styles.font_chars(p)
        if len(f) != len(p.text):
            continue
        front = i < front_end or p.text.lstrip().startswith("<RH>")
        # the legends heading: 14 pt when the author's heading was renamed (a first-level heading), 12 pt when added
        allowed = {size, big} if front or rf.heading_size_text(p.text.strip()) or _LEGEND_HEAD.match(p.text) \
            else {size}
        bad_font, bad_size = set(), set()
        for k, ch in enumerate(p.text):
            if ch.isspace() or rf._keeps_font(f[k].get("font") or "", ch):
                continue
            if family and f[k].get("font") != family:
                bad_font.add(f[k].get("font") or "?")
            if f[k].get("size") not in allowed:
                bad_size.add(f[k].get("size"))
        if bad_font:
            out.append(Finding(a.name, "format.font", ERROR, i, f"{'/'.join(sorted(bad_font))[:40]}: " + snippet(p.text, 5)))
        if bad_size:
            out.append(Finding(a.name, "format.size", ERROR, i,
                               f"{'/'.join(str(x) for x in sorted(x or 0 for x in bad_size))} pt: " + snippet(p.text, 5)))


# --- queries (Word comments in the tracked copy) ----------------------------------------------------------------------

_PIC = ("{http://schemas.openxmlformats.org/drawingml/2006/main}blip", "{urn:schemas-microsoft-com:vml}imagedata",
        W + "object", "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart")


def comment_anchors(document_xml: bytes) -> dict[str, dict]:
    """{comment id: {para: index of the paragraph it starts in, live: anchored text that stays, gone: anchored text
    deleted, para_live: that paragraph's text once the changes are accepted, picture: the paragraph shows one}}."""
    from lxml import etree

    root = etree.fromstring(document_xml)
    paras = [p for p in root.iter(W + "p")
             if not any(x.tag.endswith("}Fallback") for x in p.iterancestors())]
    pos = {p: i for i, p in enumerate(paras)}
    out: dict[str, dict] = {}
    open_ids: set[str] = set()
    para_of = None
    for el in root.iter(W + "p", W + "commentRangeStart", W + "commentRangeEnd", W + "t", W + "delText", W + "tab"):
        if el.tag == W + "p":
            if el in pos:
                para_of = el
            continue
        if el.tag == W + "commentRangeStart":
            cid = el.get(W + "id")
            p = next((x for x in el.iterancestors(W + "p") if x in pos), para_of)
            out[cid] = {"para": pos.get(p), "live": "", "gone": "", "el": p}
            open_ids.add(cid)
            continue
        if el.tag == W + "commentRangeEnd":
            open_ids.discard(el.get(W + "id"))
            continue
        if not open_ids:
            continue
        anc = {x.tag for x in el.iterancestors()}
        txt = "\t" if el.tag == W + "tab" else (el.text or "")
        gone = el.tag == W + "delText" or W + "del" in anc or W + "moveFrom" in anc
        for cid in open_ids:
            out[cid]["gone" if gone else "live"] += txt
    for d in out.values():
        p = d.pop("el")
        live = []
        if p is not None:
            for el in p.iter(W + "t", W + "tab"):
                anc = {x.tag for x in el.iterancestors()}
                if W + "del" not in anc and W + "moveFrom" not in anc:
                    live.append("\t" if el.tag == W + "tab" else el.text or "")
            d["picture"] = any(True for _ in p.iter(*_PIC))
        else:
            d["picture"] = False
        d["para_live"] = "".join(live)
    return out


def comment_texts(comments_xml: bytes | None) -> dict[str, str]:
    from lxml import etree

    if not comments_xml:
        return {}
    root = etree.fromstring(comments_xml)
    out = {}
    for c in root.iter(W + "comment"):
        out[c.get(W + "id")] = "\n".join("".join(t.text or "" for t in p.iter(W + "t")) for p in c.iter(W + "p"))
    return out


def _foreign_comments(comments_xml: bytes | None) -> set[str]:
    """Ids of the comments not written by Edit OS (no "AQn:" query number, another author)."""
    from lxml import etree

    from edit_os import queries

    if not comments_xml:
        return set()
    out = set()
    for c in etree.fromstring(comments_xml).iter(W + "comment"):
        text = "".join(t.text or "" for t in c.iter(W + "t"))
        if c.get(W + "author") != queries.AUTHOR and not re.match(r"\s*AQ\d+\s*:", text):
            out.add(c.get(W + "id"))
    return out


_BAD_QUERY = re.compile(r"\?\?\?|\bNone\b|\bnan\b|\{[a-z_]*\}|\[\]|''|\"\"|\(\s*\)")


@check
def check_queries(a: Article, out: list):
    """Author queries: on an empty paragraph, on deleted text, duplicated, or with a broken text."""
    anchors, texts = comment_anchors(a.tracked_xml), comment_texts(a.comments_xml)
    theirs = _foreign_comments(a.comments_xml)
    seen: dict[tuple, int] = {}
    for cid, text in texts.items():
        if cid in theirs:
            continue  # a comment the submission came with (the journal office's "Provided"): not a query of ours
        body = re.sub(r"^AQ\d+:\s*", "", text.strip())
        d = anchors.get(cid)
        if not body.strip():
            out.append(Finding(a.name, "query.empty_text", ERROR, d and d["para"], "", "tracked"))
        elif _BAD_QUERY.search(body):
            m = _BAD_QUERY.search(body)
            out.append(Finding(a.name, "query.bad_text", ERROR, d and d["para"],
                               f"{m.group(0)!r} in: " + snippet(body, 6), "tracked"))
        if d is None:
            out.append(Finding(a.name, "query.no_anchor", ERROR, None, snippet(body, 6), "tracked"))
            continue
        if d["gone"].strip() and not d["live"].strip():
            out.append(Finding(a.name, "query.on_deleted", ERROR, d["para"], snippet(body, 6), "tracked"))
        elif not d["para_live"].strip() and not d["picture"]:
            out.append(Finding(a.name, "query.on_empty_para", ERROR, d["para"], snippet(body, 6), "tracked"))
        k = (d["para"], _norm(body).lower())
        seen[k] = seen.get(k, 0) + 1
        if seen[k] == 2:
            out.append(Finding(a.name, "query.duplicate", ERROR, d["para"], snippet(body, 6), "tracked"))


# --- citations -------------------------------------------------------------------------------------------------------

_SCALE_WORDS = (r"agree|disagree|neutral|undecided|uncertain|never|always|sometimes|rarely|often|usually|seldom|"
                r"satisfied|dissatisfied|important|unimportant|excellent|good|fair|poor|strongly|yes|no|scored?|"
                r"points?|rating|rated|grade|stage|level|n|number|count|total|true|false|correct|incorrect")
_UNITS = (r"%|mm|cm|m|km|mg|g|kg|µg|mcg|ml|mL|l|L|dl|dL|mmol|mol|iu|IU|u|U|h|hr|hrs|hours?|min|mins|minutes?|s|sec|"
          r"seconds?|days?|weeks?|months?|years?|yrs?|y|°|°C|D|DS|DC|μm|um|nm|mmHg|kPa|Hz|dB|logMAR|times|fold|x|×")
_BEFORE_SCALE = re.compile(rf"(?:\b(?:{_SCALE_WORDS})|\d\s*(?:{_UNITS}))\s*[=:]?\s*$", re.I)
_BEFORE_NUMBER = re.compile(r"\d[\d.,]*\s*$")
_YEAR = re.compile(r"^(?:1[89]|20)\d\d$")


_SUPERS = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")


def _digits(t: str) -> tuple:
    return tuple(_numbers(t.translate(_SUPERS)))


def citation_changes(changes) -> list[tuple]:
    """(change, original text, its offset) for each accepted change that writes a numbered citation "[n...]" the
    original did not have in square brackets. A citation moved before the full stop is written by an insertion and
    a deletion of the original (round brackets or superscript) elsewhere in the paragraph: the original is that
    deletion's text, found by its numbers."""
    from edit_os.changes import Status

    acc = [c for c in changes if c.status == Status.ACCEPTED and c.reason.startswith("citations")]
    out = []
    for c in acc:
        if not re.search(r"\[\d[\d,\s–\-]*\]", c.new) or re.search(r"\[\d", c.old):
            continue
        if re.search(r"\d", c.old.translate(_SUPERS)):
            out.append((c, c.old, c.start))
            continue
        want = _digits(" ".join(re.findall(r"\[([^\]]*)\]", c.new)))
        partner = next((d for d in acc if d is not c and d.para_id == c.para_id and not re.search(r"\d", d.new)
                        and re.search(r"\d", d.old.translate(_SUPERS)) and _digits(d.old) == want), None)
        if partner is not None and not re.search(r"\[\d", partner.old):
            out.append((c, partner.old, partner.start))
    return out


def _numbers(s: str) -> list[int]:
    """The numbers a citation's text names, ranges expanded ("[3-5, 8]" -> 3, 4, 5, 8)."""
    nums: list[int] = []
    for m in re.finditer(r"(\d+)\s*[–\-]\s*(\d+)|(\d+)", s):
        if m.group(1):
            lo, hi = int(m.group(1)), int(m.group(2))
            nums += list(range(lo, hi + 1)) if 0 <= hi - lo < 200 else [lo, hi]
        else:
            nums.append(int(m.group(3)))
    return nums


@check
def check_citations(a: Article, out: list):
    """Numbers turned into citations next to scale words, units or counts, in table cells after a number, years,
    footnote markers (non-digit superscripts) and numbers beyond the reference list."""
    from edit_os import references

    src = {p.id: p for p in a.source.paragraphs}
    order = _idx(a.source)
    n_refs = len([r for r in references.extract(a.source) if r.number is not None]) or None
    for c, old, at in citation_changes(a.changes):
        p = src.get(c.para_id)
        if p is None:
            continue
        before = p.text[max(0, at - 40):at]
        old = old.strip()
        sup = any(s <= at < e for s, e in p.superscript) or bool(re.fullmatch(r"[⁰¹²³⁴⁵⁶⁷⁸⁹˒,\s\-–]+", old))
        loc = snippet(before[-30:] + " " + old + " -> " + c.new, 8)
        i = order.get(c.para_id)
        nums = _numbers(" ".join(re.findall(r"\[([^\]]*)\]", c.new)))
        # a superscript after a unit or a word is how such authors cite ("20 years²"); a round-bracket number there is a
        # value ("Strongly agree (5)", "12 (34%)", "n (%)") more often than a citation
        rnd = old.startswith("(") and not old.rstrip().endswith((".", ";"))  # "... nursing students (3)." cites
        if rnd and _BEFORE_SCALE.search(before):
            out.append(Finding(a.name, "cite.after_scale_or_unit", ERROR, i, loc, "records"))
        elif rnd and p.in_table and _BEFORE_NUMBER.search(before):
            out.append(Finding(a.name, "cite.table_count", ERROR, i, loc, "records"))
        if any(_YEAR.match(str(n)) for n in nums):
            out.append(Finding(a.name, "cite.year", ERROR, i, loc, "records"))
        elif n_refs and any(n > n_refs for n in nums):
            out.append(Finding(a.name, "cite.beyond_refs", ERROR, i, loc + f" (list has {n_refs})", "records"))
        if sup and re.search(r"[*†‡§¶#a-zA-Z]", old):
            out.append(Finding(a.name, "cite.footnote_marker", ERROR, i, loc, "records"))
        elif sup and p.in_table:
            out.append(Finding(a.name, "cite.superscript_in_table", INFO, i, loc, "records"))


# --- content: deletions and duplicates ---------------------------------------------------------------------------------

_EXPECTED_DELETION = re.compile(r"^(front_matter|submission|sections_keywords|sections_abstract|header|renumber)")


_STANDARD_SECTION = re.compile(r"(?i)conflicts?|competing|funding|financial|sponsor|grant|declarations?\b")


def _words(t: str) -> list[str]:
    return re.findall(r"[^\W\d_]+|\d+", re.sub(r"<[A-Z]{1,3}\d?>", " ", t).lower())


class Coverage:
    """How much of a text's words one clean-copy paragraph (with the next one: a heading split off) holds."""

    def __init__(self, ms):
        self.sets = [set(_words(p.text)) for p in ms.paragraphs]
        self.index: dict[str, set[int]] = collections.defaultdict(set)
        for i, s in enumerate(self.sets):
            for w in s:
                self.index[w].add(i)

    def best(self, text: str) -> float:
        ws = set(_words(text))
        if not ws:
            return 1.0
        rare = sorted(ws, key=lambda w: len(self.index.get(w, ())))[:6]
        cand = set().union(*(self.index.get(w, set()) for w in rare))
        best = 0.0
        for j in cand:
            got = self.sets[j] | (self.sets[j + 1] if j + 1 < len(self.sets) else set()) | \
                (self.sets[j - 1] if j > 0 else set())
            best = max(best, len(ws & got) / len(ws))
        return best


def _nil_ethics(a: "Article", text: str) -> bool:
    """An inline nil ethics statement the journal's settings delete (rules_sections, back_matter.delete_nil_ethics)."""
    from edit_os import rules_sections as rs

    m = rs._ETHICS.match(text)
    if not (rs._on(a.style, "back_matter.delete_nil_ethics") and m):
        return False
    content = text[m.end():]
    return bool(rs._BARE_NIL.match(content)) if rs._cfg(a.style, "back_matter.nil_ethics_bare_only") \
        else bool(rs._ETHICS_NIL.match(content))


def _nil_statement(a: "Article", text: str) -> bool:
    """A funding / conflicts / data statement that says there is nothing to state, which the journal deletes."""
    from edit_os import rules_sections as rs

    if not rs._on(a.style, "back_matter.delete_nil_statements"):
        return False
    return bool(rs._NO_DATA.match(text)) or bool(rs._NIL_TOPIC.search(text) and rs._nil(text))


def _source_body(ms) -> tuple[int, int]:
    from edit_os import references

    start = _front_end(ms)
    end = references.find_heading(ms)
    return start, end if end is not None else len(ms.paragraphs)


@check
def check_content(a: Article, out: list):
    """Whole body paragraphs gone from the clean copy (deleted by a rule that is not front matter / submission, or lost
    with no change at all), paragraphs cut to less than half, and paragraphs or headings appearing twice."""
    import evaluate
    from edit_os.changes import Status

    acc: dict[str, list] = collections.defaultdict(list)
    for c in a.changes:
        if c.status == Status.ACCEPTED:
            acc[c.para_id].append(c)
    cov = Coverage(a.clean)
    s0, s1 = _source_body(a.source)
    for i, p in enumerate(a.source.paragraphs):
        words = _words(p.text)
        if not s0 <= i < s1 or len(words) < 6:
            continue
        cs = acc.get(p.id, [])
        reasons = sorted({c.reason.split(" ")[0].split(":")[0] for c in cs if not c.format_only})
        if cov.best(p.text) >= 0.6:
            continue
        after = evaluate.apply_text(p.text, cs) if cs else p.text
        if p.in_table:
            kind, sev = ("content.cell_lost" if not reasons else "content.cell_deleted"), WARNING
        elif not reasons:
            kind, sev = "content.para_lost", ERROR
        elif all(_EXPECTED_DELETION.match(r) for r in reasons):
            continue
        elif all(r.startswith("sections_back_matter") for r in reasons) and _STANDARD_SECTION.search(p.text[:80]):
            continue  # a funding / conflicts statement the journal's standard wording replaces
        elif all(r.startswith("sections_back_matter") for r in reasons) and _nil_ethics(a, p.text):
            continue  # "Ethics: Not required ...": the journal deletes it (back_matter.delete_nil_ethics; editors 7 of 7)
        elif all(r.startswith("sections_back_matter") for r in reasons) and _nil_statement(a, p.text):
            continue  # "There were no specific funds allocated ...", "No new datasets were generated ...": a nil
            # statement the journal deletes / words its own way (back_matter.delete_nil_statements; IJO_2904 editors)
        elif not after.strip():
            kind, sev = "content.para_deleted", ERROR
        elif len(after.strip()) < 0.5 * len(p.text.strip()):
            kind, sev = "content.para_shrunk", ERROR
        else:
            kind, sev = "content.para_rewritten", WARNING
        out.append(Finding(a.name, kind, sev, i, f"{'/'.join(reasons)[:40]}: " + snippet(p.text, 6), "records"))

    # duplicates in the clean copy
    fe = _front_end(a.clean)
    skip = set()
    for i, p in enumerate(a.clean.paragraphs):
        if re.match(r"\s*First author\s*:", p.text, re.I):
            skip.update(range(i, i + 4))
    def wkey(t: str) -> str:  # words only, lower case, a typed entry number dropped: "5.\tSmith" == "Smith"
        return " ".join(_words(re.sub(r"^\s*\[?\d+[.)\]]?\s*", "", t)))

    src_count = collections.Counter(wkey(p.text) for p in a.source.paragraphs)
    src_count.update(wkey(t) for t in a.imported or [])
    from edit_os import legends
    refs_at, refs_last = _refs_start(a.clean), legends.refs_end(a.clean) or -1
    body = collections.defaultdict(list)
    heads = collections.defaultdict(list)
    for i, p in enumerate(a.clean.paragraphs):
        k = wkey(p.text)
        if not k or i < fe or i in skip:
            continue
        if _TAG.match(p.text) or re.fullmatch(r"figure legends?|tables?|references?|legends?", k):
            heads[k].append(i)
        elif len(k) >= 40 and not p.in_table and not (refs_at is not None and i > refs_at and i <= refs_last):
            body[k].append(i)  # (a duplicated reference entry is the author's: the "duplicate" query asks)
    for k, idx in heads.items():
        if len(idx) > max(1, src_count.get(k, 0)):
            out.append(Finding(a.name, "content.dup_heading", ERROR, idx[1], snippet(a.clean.paragraphs[idx[1]].text)))
    for k, idx in body.items():
        if len(idx) > max(1, src_count.get(k, 0)):
            out.append(Finding(a.name, "content.dup_para", ERROR, idx[1], snippet(a.clean.paragraphs[idx[1]].text)))


# --- leftovers in the clean copy ---------------------------------------------------------------------------------------

_SIGNOFF = re.compile(r"^\s*(technical editing|copy\s*editing) done on", re.I)
_RESIDUE = ("ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange", "sectPrChange", "tblPrChange",
            "commentRangeStart", "commentReference")
_SPACE_PUNCT = re.compile(r"(?<=[^\s(\[])[  ]+([,;:)\]]|\.(?![\d.]))")


def _dbl_context(t: str, i: int) -> str:
    """The words round a "[[" / "]]" at t[i], to find it in the submission's own text."""
    return t[max(0, i - 12):i + 14]


@check
def check_clean(a: Article, out: list):
    """Leftovers in the clean copy: "???" placeholders, stray tags, "[[", tracked-change residue, double spaces, a space
    before punctuation, no running head; the sign-off without a name is counted (info)."""
    for tag in _RESIDUE:
        n = len(re.findall(rf"<w:{tag}[ >/]".encode(), a.clean_xml))
        if n:
            out.append(Finding(a.name, "clean.residue", ERROR, None, f"w:{tag} x{n}"))
    if "word/comments.xml" in a.clean_parts:
        out.append(Finding(a.name, "clean.residue", ERROR, None, "comments part"))
    has_rh = False
    dbl = a.style.get("guide_replace.double_space", False)
    src_dbl = [p.text for p in a.source.paragraphs if "[[" in p.text or "]]" in p.text]
    for i, p in enumerate(a.clean.paragraphs):
        t = p.text
        if not t.strip():
            continue
        if _SIGNOFF.match(t):
            if "???" in t:
                out.append(Finding(a.name, "signoff.placeholder", INFO, i, snippet(t, 5)))
            continue
        if t.lstrip().startswith("<RH>"):
            has_rh = True
        if "???" in t:
            kind = "clean.placeholder_rh" if t.lstrip().startswith("<RH>") else \
                "clean.placeholder_legend" if re.match(r"\s*(Figure|Table)\s*\d+\s*:\s*\?\?\?", t) else "clean.placeholder"
            out.append(Finding(a.name, kind, WARNING, i, snippet(t)))
        tags = list(re.finditer(r"</?(?:H\d|RH|[A-Z]{1,3}\d?)>", t))
        for m in tags:
            if m.start() != len(t) - len(t.lstrip()) or not t[m.end():].strip():
                out.append(Finding(a.name, "clean.stray_tag", ERROR, i, f"{m.group(0)} " + snippet(t, 6)))
                break
        if "[[" in t or "]]" in t:
            # A "]]" the author typed ("[CrossRef][PubMed]] arrest ...": a reference pasted over another) is theirs,
            # queried with the entry: counted (info). One our changes made is an error.
            theirs = all(any(_dbl_context(t, m.start()) in s for s in src_dbl)
                         for m in re.finditer(r"\[\[|\]\]", t))
            out.append(Finding(a.name, "clean.double_bracket", INFO if theirs else ERROR, i, snippet(t)))
        if re.search(r"\S {2,}\S", t):  # an error where the journal's guide rule removes them, else a note
            out.append(Finding(a.name, "clean.double_space", WARNING if dbl else INFO, i, snippet(t)))
        m = _SPACE_PUNCT.search(t)
        if m and not re.search(r"\.\s*\.\s*\.", t[m.start():m.end() + 4]):
            out.append(Finding(a.name, "clean.space_before_punct", WARNING, i,
                               snippet(t[max(0, m.start() - 30):m.end() + 10], 6)))
    if not has_rh and a.style.get("front_matter.enabled", True):
        out.append(Finding(a.name, "clean.no_running_head", ERROR, None, ""))


# --- figures and tables ------------------------------------------------------------------------------------------------

_LEGEND = re.compile(r"^\s*Figure\s+\d+[A-Za-z]?\s*:")
_LEGEND_HEAD = re.compile(r"^\s*(?:<H\d>)?\s*(figure legends?|figure captions?|legends? (?:to|for) figures?)\s*:?\s*$", re.I)
_TABLE_CAP = re.compile(r"^\s*Table\s+S?\d+[A-Za-z]?\s*[:.]")


def _tbl_text(tbl) -> list[list[str]]:
    rows = []
    for tr in tbl.iter(W + "tr"):
        rows.append(["".join(t.text or "" for t in tc.iter(W + "t")) for tc in tr.iter(W + "tc")])
    return rows


@check
def check_figures_tables(a: Article, out: list):
    """Legends without the "Figure N:" label (or the label not bold), tables without a caption, captions not bold
    where the journal wants them bold, tables before the references when they go to the end."""
    ms = a.clean
    styles = Styles(ms)
    idx = {id(p.element): i for i, p in enumerate(ms.paragraphs)}
    in_legends = False
    label_bold = a.style.get("formatting.label_bold", True)
    for i, p in enumerate(ms.paragraphs):
        t = p.text.strip()
        if _LEGEND_HEAD.match(t):
            in_legends = True
            continue
        if not t or p.in_table:
            continue
        if in_legends and (_TAG.match(t) or _TABLE_CAP.match(t)):
            in_legends = False
        looks = re.match(r"(Fig\.?|Figure|FIGURE|FIG\.?)\s*\d+", t)
        if looks and not _LEGEND.match(t) and len(t) < 600 and (
                in_legends or re.match(r"(Fig\.?|Figure|FIGURE|FIG\.?)\s*\d+[A-Za-z]?\s*[.:\-–—)(]", t)):
            out.append(Finding(a.name, "figtab.legend_label", WARNING, i, snippet(t)))
        elif _LEGEND.match(t) and label_bold:
            m = _LEGEND.match(p.text)
            if _all_bold(p, styles, 0, m.end()) < 1:
                out.append(Finding(a.name, "figtab.legend_label_not_bold", WARNING, i, snippet(t, 4)))

    body = body_elements(ms.document.element)
    from edit_os import references
    ref_head = references.find_heading(ms)
    ref_el = ms.paragraphs[ref_head].element if ref_head is not None else None
    ref_pos = body.index(ref_el) if ref_el is not None and ref_el in body else None
    rh_pos = next((k for k, el in enumerate(body) if el.tag == W + "p" and
                   "".join(x.text or "" for x in el.iter(W + "t")).lstrip().startswith("<RH>")), 0)
    caption_bold = a.style.get("formatting.table_caption_bold", False)
    in_row = a.style.get("figures.table_caption_row", False)  # the caption goes into a merged first row
    to_end = a.style.get("figures.tables_to_end", False) or a.style.get("tables.to_end", False)
    from edit_os import qc as qc_mod
    supp_ids = qc_mod.supplementary_ids(ms)
    supp = {id(p.element) for p in ms.paragraphs if not p.in_table and
            (p.id in supp_ids or qc_mod._SUPP_HEAD.match(p.text))}
    in_supp = False
    for k, el in enumerate(body):
        if el.tag == W + "p":
            in_supp = id(el) in supp
        if el.tag != W + "tbl" or k < rh_pos:
            continue  # the header block's tables (renumbering) sit above the running head
        if in_supp:
            continue  # supplementary material, retained as it is (S372, G151): its tables are the author's
        rows = _tbl_text(el)
        first = " ".join(rows[0]).strip() if rows else ""
        if not any(c.strip() for r in rows for c in r):
            out.append(Finding(a.name, "figtab.empty_table", WARNING, None, f"table {k}"))
            continue
        cap_el = None
        if _TABLE_CAP.match(first):
            cap_el = next(iter(el.iter(W + "p")), None)
        else:
            for step in (-1, 1):  # up to two text paragraphs away (a "Panel A." line or a note between)
                j, seen = k + step, 0
                while 0 <= j < len(body) and body[j].tag == W + "p" and seen < 3:
                    t = "".join(x.text or "" for x in body[j].iter(W + "t"))
                    if _TABLE_CAP.match(t):
                        cap_el = body[j]
                        break
                    seen += bool(t.strip())
                    j += step
                if cap_el is not None:
                    break
        if re.search(r"(?i)reviewer|author'?s?\W*response|response to", first[:80]):
            out.append(Finding(a.name, "content.reviewer_reply_left", ERROR, None, f"table {k}: " + snippet(first, 6)))
            continue
        if cap_el is None:
            near = [body[j] for j in (k - 1, k - 2, k + 1) if 0 <= j < len(body) and body[j].tag == W + "p"]
            bare = any(re.match(r"\s*Table\s*\d+\s*[:.]?\s*$|\s*Table\d", "".join(x.text or "" for x in e.iter(W + "t")))
                       for e in near)
            out.append(Finding(a.name, "figtab.table_no_caption", WARNING, None,
                               ("label without caption text / colon; " if bare else "") + f"table {k}: " + snippet(first, 5)))
        elif in_row and not _TABLE_CAP.match(first):
            out.append(Finding(a.name, "figtab.caption_outside_table", WARNING, None,
                               f"table {k}: " + snippet("".join(x.text or "" for x in cap_el.iter(W + "t")), 4)))
        if cap_el is not None and caption_bold and id(cap_el) in idx:
            p = ms.paragraphs[idx[id(cap_el)]]
            if _all_bold(p, styles) < 1:
                out.append(Finding(a.name, "figtab.caption_not_bold", WARNING, idx[id(cap_el)], snippet(p.text, 5)))
        if to_end and ref_pos is not None and k < ref_pos and cap_el is not None:
            # Renumbered tables are left in place on purpose (legends._plan_tables: the renumbering and the move
            # are not done together), with a note telling the editor to move them: expected, counted (info).
            sev = INFO if any(str(n).startswith("Tables renumbered: left where they are")
                              for n in a.qc.get("notes") or []) else ERROR
            out.append(Finding(a.name, "figtab.table_not_at_end", sev, None, f"table {k}: " + snippet(first, 5)))


# --- references --------------------------------------------------------------------------------------------------------

_PARTICLES = {"van", "von", "de", "der", "den", "da", "di", "du", "dos", "das", "del", "della", "le", "la", "el", "al",
              "ten", "ter", "bin", "ibn", "y", "e", "et"}


@check
def check_references(a: Article, out: list):
    """Reference entries: a lower-case surname, "et al" without its period, empty entries, numbering gaps or a
    numbered list that restarts."""
    from docx.oxml.ns import qn

    from edit_os import references

    ms = a.clean
    refs = references.extract(ms)
    idx = _idx(ms)
    head = references.find_heading(ms)
    if not refs or head is None:
        return
    for r in refs:
        i = idx.get(r.para_id)
        first = r.text.strip().split(" ")[0] if r.text.strip() else ""
        if re.match(r"[a-z]", first) and first.lower().rstrip(",.") not in _PARTICLES \
                and not re.match(r"(https?:|www\.|doi)", first):
            out.append(Finding(a.name, "refs.lowercase_surname", WARNING, i, snippet(r.text, 3)))
        if re.search(r"\bet al\b(?!\.)", r.full_text):
            out.append(Finding(a.name, "refs.et_al_period", ERROR, i, snippet(r.text, 6)))
        if not r.text.strip() or len(r.text.strip()) < 5:
            out.append(Finding(a.name, "refs.empty_entry", ERROR, i, repr(r.text[:10])))
    last = idx.get(refs[-1].para_id, head)
    for i in range(head + 1, last):
        p = ms.paragraphs[i]
        if p.numbered and not p.text.strip() and not p.in_table:
            out.append(Finding(a.name, "refs.empty_entry", ERROR, i, "empty numbered paragraph"))
    shown = shown_numbers(ms, refs)
    gaps = [n for n in range(1, max(shown) + 1) if n not in set(shown)] if shown else []
    dup = sorted(n for n, c in collections.Counter(shown).items() if c > 1)
    if gaps or dup or shown != sorted(shown):
        out.append(Finding(a.name, "refs.number_sequence", ERROR, None,
                           f"missing {gaps[:5]} repeated {dup[:5]}" + (" out of order" if shown != sorted(shown) else "")))
    for r in refs:
        if r.label and ms.paragraphs[idx[r.para_id]].numbered:
            out.append(Finding(a.name, "refs.double_number", ERROR, idx[r.para_id], f"typed {r.label.strip()!r} in a "
                                                                                  "numbered list"))


def _list_starts(ms) -> dict[str, int]:
    """{numId: the number its first level starts at} from numbering.xml (startOverride, else the abstract's start)."""
    from docx.oxml.ns import qn

    try:
        root = ms.document.part.numbering_part.element
    except Exception:  # noqa: BLE001 - no numbering part
        return {}
    abstract = {}
    for an in root.iter(qn("w:abstractNum")):
        lvl = next((x for x in an.iter(qn("w:lvl")) if x.get(qn("w:ilvl")) == "0"), None)
        st = lvl.find(qn("w:start")) if lvl is not None else None
        abstract[an.get(qn("w:abstractNumId"))] = int(st.get(qn("w:val"))) if st is not None else 1
    out = {}
    for num in root.iter(qn("w:num")):
        an = num.find(qn("w:abstractNumId"))
        start = abstract.get(an.get(qn("w:val")) if an is not None else None, 1)
        for ov in num.iter(qn("w:lvlOverride")):
            so = ov.find(qn("w:startOverride"))
            if ov.get(qn("w:ilvl")) == "0" and so is not None:
                start = int(so.get(qn("w:val")))
        out[num.get(qn("w:numId"))] = start
    return out


def shown_numbers(ms, refs) -> list[int]:
    """The number a reader sees on each entry: the typed label, else Word's list number (its list's start plus the
    entries of that list before it). Entries neither typed nor in a list count as 0 (no number)."""
    from docx.oxml.ns import qn

    starts, seen = _list_starts(ms), collections.Counter()
    idx = _idx(ms)
    out = []
    for r in refs:
        p = ms.paragraphs[idx[r.para_id]]
        n = p.element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        listed = n is not None and n.get(qn("w:val")) != "0"
        if r.label:
            out.append(r.number)
            if listed:  # Word numbers it too ("10. 10. Smith"): the list's count goes on
                seen[n.get(qn("w:val"))] += 1
        elif listed:
            nid = n.get(qn("w:val"))
            out.append(starts.get(nid, 1) + seen[nid])
            seen[nid] += 1
        else:
            out.append(r.number or 0)
    return out


# --- driver ------------------------------------------------------------------------------------------------------------

def scan(od: str, name: str | None = None, journal: str | None = None) -> list[Finding]:
    """Every check on the outputs in `od`. A check that crashes is reported as `scan.crash` (its error type only)."""
    a = Article.load(od, name, journal)
    out: list[Finding] = []
    for cid, fn in [(k, globals()["check_" + k]) for k in CHECKS]:
        try:
            fn(a, out)
        except Exception as e:  # noqa: BLE001 - one broken check does not stop the others
            out.append(Finding(a.name, "scan.crash", ERROR, None, f"{cid}: {type(e).__name__}"))
    return out


def _work(args) -> tuple[str, str, list[dict], str]:
    adir, journal, outdir, reuse = args
    name = article_name(adir)
    j = journal or journal_of(name)
    try:
        od = run_article(adir, j, outdir, reuse=reuse)
    except Exception as e:  # noqa: BLE001
        return name, j, [Finding(name, "pipeline.crash", ERROR, None, type(e).__name__, "records").to_dict()], "crash"
    try:
        return name, j, [f.to_dict() for f in scan(od, name, j)], "ok"
    except Exception as e:  # noqa: BLE001
        return name, j, [Finding(name, "scan.crash", ERROR, None, type(e).__name__, "records").to_dict()], "crash"


def summarize(rows: list[dict], journals: dict[str, str]) -> dict:
    """{journal: {check: {"findings": n, "articles": n, "severity": s}}}, plus "articles" per journal."""
    out: dict = {}
    for j in sorted(set(journals.values())):
        out[j] = {"articles": sum(v == j for v in journals.values()), "checks": {}}
    for r in rows:
        j = journals.get(r["article"], "?")
        key = r["check"] if r["severity"] != INFO or r["check"].startswith(("signoff", "cite.superscript")) \
            else r["check"] + " (info)"  # a check that is an error in one article and expected in another
        c = out.setdefault(j, {"articles": 0, "checks": {}})["checks"].setdefault(
            key, {"findings": 0, "articles": set(), "severity": r["severity"]})
        c["findings"] += 1
        c["articles"].add(r["article"])
    for j in out.values():
        for c in j["checks"].values():
            c["articles"] = len(c["articles"])
    return out


def print_summary(summ: dict) -> None:
    for j, d in summ.items():
        print(f"\n== {j}: {d['articles']} articles")
        print(f"{'check':34s} {'sev':8s} {'findings':>8s} {'articles':>8s}")
        for k, c in sorted(d["checks"].items(), key=lambda kv: (kv[1]['severity'] != ERROR, -kv[1]["articles"])):
            print(f"{k:34s} {c['severity']:8s} {c['findings']:8d} {c['articles']:8d}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dirs", nargs="*")
    ap.add_argument("--out", required=True)
    ap.add_argument("--journal")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--reuse", action="store_true", help="scan outputs already in OUT instead of running again")
    ap.add_argument("--summary", action="store_true", help="only print the tables of OUT/findings.jsonl")
    a = ap.parse_args(argv)
    os.makedirs(a.out, exist_ok=True)
    fpath = os.path.join(a.out, "findings.jsonl")
    if a.summary:
        rows = [json.loads(x) for x in open(fpath, encoding="utf-8")]
        meta = json.load(open(os.path.join(a.out, "summary.json"), encoding="utf-8"))
        print_summary(summarize(rows, meta["journals"]))
        return 0
    jobs = [(d, a.journal, a.out, a.reuse) for d in a.dirs]
    rows, journals, t0 = [], {}, time.time()
    with cf.ProcessPoolExecutor(max(1, a.workers)) as ex:
        for name, j, found, status in ex.map(_work, jobs):
            journals[name] = j
            rows += found
            errs = sum(f["severity"] == ERROR for f in found)
            print(f"{name:20s} {j:8s} {status:5s} findings {len(found):4d} errors {errs:4d}", file=sys.stderr, flush=True)
    with open(fpath, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    summ = summarize(rows, journals)
    with open(os.path.join(a.out, "summary.json"), "w", encoding="utf-8") as f:
        json.dump({"journals": journals, "summary": summ, "seconds": round(time.time() - t0)}, f, indent=1)
    print_summary(summ)
    return 0


if __name__ == "__main__":
    sys.exit(main())
