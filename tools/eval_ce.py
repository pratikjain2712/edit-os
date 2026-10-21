"""Score the copy-edit checkers against a copy editor's real output.

Each article folder holds `te.docx` (the style editor's tracked file) and `ce.docx` (the copy
editor's tracked file, built on top of te.docx). Copy-edit changes are those whose author does
not appear in te.docx. For each paragraph of ce.docx:

  O = text with only the style editor's changes applied (the copy editor's starting point),
  T = text with every change applied (the copy editor's result),
  U = O with our checkers' changes applied.

The checkers run on a temporary .docx built from the O paragraphs. closure, helped and hurt are
as in tools/evaluate.py. Per checker, precision is the share of proposed edits that, applied on
their own, bring the paragraph closer to the copy editor's text. Author queries count as hits
when the copy editor's comments mention the abbreviation or their insertions expand it.

Usage: python tools/eval_ce.py DATA_DIR JOURNAL[,JOURNAL...] [--rules] [--show-hurt]
       [--lt-explore] [--only spelling,abbreviations,languagetool,claude]
       [--split train|test] [--articles A,B,...] [--model M] [--effort E]

--rules        precision per LanguageTool rule ID / reason
--lt-explore   accept every LanguageTool rule with one suggestion (for tuning the allowlist)
--split        only articles of that split (claude_pass.is_test_article); report Claude on `test`
--articles     only these article folders
--model/--effort  Claude model and effort for the claude checker (EDITOS_CE_MODEL / EDITOS_CE_EFFORT)

The Claude pass (paid) runs only when named in --only. Its queries count as hits when the copy
editor left a comment anchored in the same paragraph.
"""

from __future__ import annotations

import glob
import io
import os
import re
import sys
import tempfile
import zipfile
from collections import Counter, defaultdict

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402
from rapidfuzz.distance import Levenshtein  # noqa: E402

from edit_os import copyedit, docmodel, style as style_mod  # noqa: E402
from edit_os.copyedit import claude_pass, langtool  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def tokens(s: str) -> list[str]:
    return [" " if t.isspace() else t for t in re.findall(r"\w+|\s+|[^\w\s]", s)]


def distance(a: str, b: str) -> int:
    return Levenshtein.distance(tokens(a), tokens(b))


def authors(docx_bytes: bytes) -> set[str]:
    xml = zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml")
    root = etree.fromstring(xml)
    return {el.get(W + "author") for el in root.iter(W + "ins", W + "del") if el.get(W + "author")}


# Symbol-font characters (w:sym) written as text, so "40.75 ± 13.63" is not read as a double space.
_SYMBOL = {"B1": "±", "A3": "≤", "B3": "≥", "B4": "×", "B0": "°", "6D": "μ", "61": "α", "62": "β",
           "63": "χ", "67": "γ", "64": "δ", "44": "Δ", "6B": "κ", "6C": "λ", "70": "π", "73": "σ",
           "AE": "→", "AD": "↑", "AF": "↓", "BB": "≈", "B9": "≠", "2D": "−", "D6": "√", "A5": "∞"}


def closer(u: str, o: str, t: str) -> bool:
    """u is closer to t than o: fewer token edits, or (for spacing fixes, invisible to the token
    distance) the same token distance and fewer character edits."""
    du, do = distance(u, t), distance(o, t)
    return du < do or du == do and Levenshtein.distance(u, t) < Levenshtein.distance(o, t)


def views(docx_bytes: bytes, te_authors: set[str]):
    """(in_table, O, T, ce_inserted_text) per paragraph of the copy editor's tracked file."""
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml"))
    out = []
    for p in root.iter(W + "p"):
        o, t, ins = [], [], []
        for el in p.iter(W + "t", W + "delText", W + "tab", W + "sym"):
            if el.tag == W + "sym":
                txt = _SYMBOL.get(el.get(W + "char", "").upper()[-2:], "\u25a1")
            else:
                txt = "\t" if el.tag == W + "tab" else (el.text or "")
            ins_by = [a.get(W + "author") for a in el.iterancestors(W + "ins")]
            del_by = [a.get(W + "author") for a in el.iterancestors(W + "del")]
            deleted = el.tag == W + "delText" or bool(del_by)
            # O: style editor's insertions in, their deletions out; copy editor's changes undone.
            in_o = all(a in te_authors for a in ins_by) and not any(a in te_authors for a in del_by)
            if in_o:
                o.append(txt)
            if not deleted:
                t.append(txt)
            if ins_by and not all(a in te_authors for a in ins_by) and not deleted:
                ins.append(txt)
        in_table = any(True for _ in p.iterancestors(W + "tc"))
        out.append((in_table, "".join(o), "".join(t), "".join(ins)))
    return out


def comments(docx_bytes: bytes) -> str:
    """Comment texts, each with the text it is anchored to ("AQ: Kindly expand ..." sits on the ABBR)."""
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    if "word/comments.xml" not in z.namelist():
        return ""
    root = etree.fromstring(z.read("word/comments.xml"))
    texts = {c.get(W + "id"): "".join(t.text or "" for t in c.iter(W + "t")) for c in root.iter(W + "comment")}
    doc = z.read("word/document.xml").decode("utf-8")
    out = []
    for cid, text in texts.items():
        m = re.search(rf'<w:commentRangeStart w:id="{cid}"/>(.*?)<w:commentRangeEnd w:id="{cid}"/>', doc, re.S)
        anchor = "".join(re.findall(r"<w:(?:t|delText)(?: [^>]*)?>([^<]*)<", m.group(1))) if m else ""
        out.append(f"{text} [{anchor[:200]}]")
    return "\n".join(out)


def comment_anchors(docx_bytes: bytes) -> list[str]:
    return [m.group(1) for m in re.finditer(r"\[(.*)\]$", comments(docx_bytes), re.M) if len(m.group(1).strip()) >= 3]


def build_docx(paras, path: str) -> None:
    doc = Document()
    for in_table, o, _, _ in paras:
        if in_table:
            doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(o)
        else:
            doc.add_paragraph().add_run(o)
    doc.save(path)


def apply(text: str, changes) -> str:
    for c in sorted(changes, key=lambda c: (c.start, c.end), reverse=True):
        assert text[c.start:c.end] == c.old, (c, text[c.start:c.end])
        text = text[:c.start] + c.new + text[c.end:]
    return text


def journal_of(article_dir: str) -> str:
    return re.split(r"[_-]", os.path.basename(article_dir))[0].lower()


def evaluate(article_dir: str, only=None, lt_client=None, explore=False, show_hurt=False):
    with open(os.path.join(article_dir, "te.docx"), "rb") as f:
        te = f.read()
    with open(os.path.join(article_dir, "ce.docx"), "rb") as f:
        ce = f.read()
    paras = views(ce, authors(te))
    st = style_mod.load(journal_of(article_dir))
    if explore:
        st.settings["copyedit"]["languagetool"].update(categories=["*"], rules=[], disabled_rules=[])
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "o.docx")
        build_docx(paras, path)
        ms = docmodel.load(path)
        assert len(ms.paragraphs) == len(paras), (len(ms.paragraphs), len(paras))
        res = copyedit.run_all(ms, st, only=only, lt_client=lt_client)

    by_para = defaultdict(list)
    for c in res.changes.changes:
        by_para[c.para_id].append(c)
    r = dict(changes=len(res.changes.changes), queries=len(res.queries), dOT=0, dUT=0,
             helped=0, hurt=0, ce_paragraphs=0)
    hits: Counter = Counter()
    proposed: Counter = Counter()
    for p, (_, o, t, _) in zip(ms.paragraphs, paras):
        cs = by_para.get(p.id, [])
        d_o = distance(o, t)
        if o != t:
            r["ce_paragraphs"] += 1
        u = apply(o, cs)
        d_u = distance(u, t)
        r["dOT"] += d_o
        r["dUT"] += d_u
        if u != o:
            if d_u < d_o:
                r["helped"] += 1
            elif d_u > d_o:
                r["hurt"] += 1
                if show_hurt:
                    print(f"  HURT {[(c.old, c.new, c.reason) for c in cs]}\n    O: {o[:200]!r}\n    T: {t[:200]!r}")
        for c in cs:
            key = c.reason if c.reason.startswith("languagetool") else c.reason.split(" (")[0]
            proposed[key] += 1
            proposed["checker:" + checker_of(c)] += 1
            if closer(apply(o, [c]), o, t):
                hits[key] += 1
                hits["checker:" + checker_of(c)] += 1

    # Author queries: did the copy editor ask about, or expand, the same abbreviation?
    anchors = comment_anchors(ce)
    o_text = {p.id: o for p, (_, o, _, _) in zip(ms.paragraphs, paras)}
    ce_comments = "\n".join(c for c in comments(ce).split("\n")
                            if re.search(r"xpand|abbrevia|defin|full form", c, re.I))
    ce_ins = " ".join(x[3] for x in paras)
    for q in res.queries:
        proposed["query:" + q.check] += 1
        if q.check == "claude.query":
            if any(a.strip() in o_text.get(q.para_id, "") for a in anchors):
                hits["query:" + q.check] += 1
            continue
        abbr = q.text.split(":")[0]
        if re.search(rf"(?<!\w){re.escape(abbr)}(?!\w)", ce_comments) or f"({abbr})" in ce_ins \
                or re.search(rf"\w+ \w+[^()]*\({re.escape(abbr)}\)", ce_ins):
            hits["query:" + q.check] += 1
    return r, hits, proposed


def checker_of(c) -> str:
    if c.source.value == "ai":
        return "claude"
    if c.reason.startswith("languagetool"):
        return "languagetool"
    if c.reason.startswith("spelling"):
        return "spelling"
    return "abbreviations"


def main(argv):
    data, journals = argv[0], argv[1].lower().split(",")
    only = ["spelling", "abbreviations", "languagetool"]  # the Claude pass is paid: only on request
    if "--only" in argv:
        only = argv[argv.index("--only") + 1].split(",")
    for flag, env in (("--model", "EDITOS_CE_MODEL"), ("--effort", "EDITOS_CE_EFFORT")):
        if flag in argv:
            os.environ[env] = argv[argv.index(flag) + 1]
    split = argv[argv.index("--split") + 1] if "--split" in argv else None
    names = argv[argv.index("--articles") + 1].split(",") if "--articles" in argv else None
    lt_client = langtool.default_client()
    if lt_client is None:
        print("(LanguageTool server not running: start it with python -m edit_os.copyedit languagetool start)")
    dirs = sorted(d for d in glob.glob(os.path.join(data, "*"))
                  if journal_of(d) in journals or "all" in journals)
    tot = Counter()
    hits, proposed = Counter(), Counter()
    for d in dirs:
        if not os.path.exists(os.path.join(d, "ce.docx")):
            continue
        name = os.path.basename(d)
        if names and name not in names or split and claude_pass.is_test_article(name) != (split == "test"):
            continue
        r, h, pr = evaluate(d, only, lt_client, "--lt-explore" in argv, "--show-hurt" in argv)
        tot.update(r)
        tot["articles"] += 1
        hits.update(h)
        proposed.update(pr)
        print(f"{os.path.basename(d):26s} edits {r['changes']:4d} queries {r['queries']:3d}"
              f"  distance {r['dOT']:5d} -> {r['dUT']:5d}  helped {r['helped']:3d} hurt {r['hurt']:3d}")
    closure = (tot["dOT"] - tot["dUT"]) / tot["dOT"] if tot["dOT"] else 0
    print(f"\narticles {tot['articles']}  edits {tot['changes']}  queries {tot['queries']}  closure {closure:6.2%}"
          f"  paragraphs helped {tot['helped']} hurt {tot['hurt']}  (copy-edited paragraphs {tot['ce_paragraphs']})")
    keys = [k for k in proposed if k.startswith(("checker:", "query:"))]
    if "--rules" in argv:
        keys += [k for k in proposed if not k.startswith(("checker:", "query:"))]
    for k in sorted(keys, key=lambda k: (not k.startswith(("checker:", "query:")), -proposed[k])):
        print(f"  {k:55s} proposed {proposed[k]:5d}  precision {hits[k] / proposed[k]:6.1%}")
    for model, u in claude_pass.USAGE.items():
        n = max(1, tot["articles"])
        print(f"claude {model}: {dict(u)}  cost ${claude_pass.cost(model):.4f}  per article ${claude_pass.cost(model) / n:.4f}")
    if claude_pass.DROPS:
        print(f"claude edits dropped by verification: {dict(claude_pass.DROPS)}")


if __name__ == "__main__":
    main(sys.argv[1:])
