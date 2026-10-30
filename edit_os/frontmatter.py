"""Front matter: running head, manuscript type, title, author group, affiliations, correspondence,
history dates and the cite line, styled the way the pre-editing guide and the journal ask.

Where the inputs keep it: the title page is usually a separate file (FirstPage.docx, "Title Page.docx")
next to the manuscript; the submission system's XML holds the article type, the history dates and the
authors; some manuscripts also start with a title page of their own. The style editor rebuilds the
block at the top of the manuscript from all three, so most of it is new paragraphs.

How it works:
1. Claude Haiku splits the title-page text into fields (structured output) and names the proper
   nouns and abbreviations in the title. It never writes text that goes into the file: every name,
   affiliation part, address part and e-mail it returns is checked against the source text and
   dropped if it is not there (`verify`).
2. Code builds each paragraph from those fields and the XML: surnames for the running head,
   sentence or title case, initials, affiliation numbering and grouping, country, PIN format, dates.
3. The result is proposed, never applied: manuscript paragraphs that already hold front matter get
   `Change` records (text and tracked formatting); missing paragraphs are `NewParagraph` records,
   inserted as tracked paragraphs by `render_inserts`; what only the author can supply becomes an
   author query in the guide's wording.

Without an API key only the XML is used.
"""

from __future__ import annotations

import copy
import difflib
import hashlib
import io
import json
import os
import re
import sqlite3
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from itertools import count

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from . import ai_refs, provenance, queries, spend, submission
from .changes import Change, Layer, Source, Status
from .docmodel import Manuscript, Paragraph
from .style import JournalStyle, _merge

# Author queries, in the guide's wording.
queries.TEXT.update({
    "front.running_head": "Kindly provide running head.",
    "front.article_type": "Kindly provide article type.",
    "front.confirm_type": "Kindly confirm article type.",
    "front.full_name": "Kindly provide the full name of author “{name}”.",
    "front.authors_check": "Please ensure that the author names and affiliations are treated correctly.",
    "front.running_head_long": "The running title has more than {n} characters. Kindly shorten it to {n} characters "
                               "or fewer.",
    "front.country": "Kindly provide the country name for the affiliation.",
    "front.xml_only_author": "Author name(s) {names} given in the submission system, but missing in the author list "
                             "provided in the title page. Kindly check.",
    # in the editors' wording (front_matter.queries in the journal's style)
    "front.department": "Kindly provide department name.",
    "front.departments": "Kindly provide department names.",
    "front.missing_date": "Kindly provide missing {what} date.",
    "front.postal_code": "Kindly check and provide postal code.",
    "front.designator": "Kindly note that the corresponding author’s designator is missing. Please check and provide.",
    "front.xml_name": "Kindly note the Author name provided in XML as “{name}”. Please check and confirm.",
    "front.type_source": "Kindly note that the “Article type” provided in {source} file has been followed.",
})

MODEL = os.environ.get("EDITOS_MODEL", "claude-haiku-5-5")
USAGE = {"calls": 0, "input_tokens": 0, "output_tokens": 0}  # this process, for cost reporting
MISSING = "???"
_ids = count(1)

ROLES = ["running_title", "article_type", "title", "authors", "affiliation", "correspondence", "label",
         "other_front_matter", "body"]
_S = {"type": "string"}
SCHEMA = {
    "type": "object",
    "properties": {
        "article_type": _S, "title": _S, "running_title": _S,
        "authors": {"type": "array", "items": {
            "type": "object",
            "properties": {"given_names": _S, "surname": _S, "prefix": _S,
                           "affiliations": {"type": "array", "items": {"type": "integer"}},
                           "corresponding": {"type": "boolean"}},
            "required": ["given_names", "surname", "prefix", "affiliations", "corresponding"],
            "additionalProperties": False}},
        "affiliations": {"type": "array", "items": {
            "type": "object",
            "properties": {"id": {"type": "integer"}, "department": _S, "institution": _S, "city": _S,
                           "state": _S, "country": _S},
            "required": ["id", "department", "institution", "city", "state", "country"],
            "additionalProperties": False}},
        "correspondence": {
            "type": "object",
            "properties": {"prefix": _S, "given_names": _S, "surname": _S,
                           "address_parts": {"type": "array", "items": _S},
                           "city": _S, "state": _S, "postal_code": _S, "country": _S, "email": _S},
            "required": ["prefix", "given_names", "surname", "address_parts", "city", "state", "postal_code",
                         "country", "email"],
            "additionalProperties": False},
        "proper_nouns": {"type": "array", "items": _S},
        "paragraphs": {"type": "array", "items": {
            "type": "object",
            "properties": {"id": _S, "role": {"type": "string", "enum": ROLES}},
            "required": ["id", "role"], "additionalProperties": False}},
    },
    "required": ["article_type", "title", "running_title", "authors", "affiliations", "correspondence",
                 "proper_nouns", "paragraphs"],
    "additionalProperties": False,
}

INSTRUCTIONS = """You split the title page of a medical journal manuscript into fields.
Paragraphs are given with ids: F = the separate title-page file, M = the start of the manuscript file.
Copy every value exactly as it is written in the text (same spelling and capitals). Never add, guess,
complete, translate or correct anything; use "" or [] when something is not there.
- article_type: the manuscript type if stated (e.g. "Original Article", "Case Report"), without its label.
- title: the article title, without a label such as "Title:". If an M paragraph holds the title, copy
  it from there (the manuscript's own title wins over the title-page file).
- running_title: the running title / short title / running head, without its label.
- authors: in author-list order. given_names: first and middle names or initials as written;
  surname: the family name (keep particles such as van, von, de, Al, El and compound surnames whole).
  Leave out titles (Dr, Prof, Mr), degrees (MBBS, MD, MS, PhD, DNB, ...), designations, markers and
  asterisks. prefix: "Dr" or "Prof" if written before the name, else "". affiliations: ids of this
  author's affiliations, from superscript markers or from the affiliation written with the author's
  name. corresponding: true if the author is marked as corresponding author.
- affiliations: each institutional affiliation once (id 1, 2, ...), as parts copied from the text:
  department (department, unit, division or service, without designations such as "Professor",
  "Senior Resident", "Consultant", "Head"), institution (hospital, college or university; may contain
  commas), city, state, country. Leave out postal codes, street addresses, e-mail and phone.
- correspondence: the corresponding author's details from the correspondence section (or from that
  author's own entry if there is no such section): prefix, given_names, surname, address_parts
  (department, institution, street, area ... in order, each as written, without city, state, postal code,
  country, phone, fax, e-mail or web address), city, state, postal_code, country, email.
- proper_nouns: words in the title and running title that keep capital letters in sentence case:
  proper nouns (people, places, eponyms such as "Valsalva"), abbreviations and acronyms ("OCT",
  "COVID-19"), brand names. Disease, organism, anatomy and method names are not proper nouns unless
  they contain a person's or place's name ("granular dystrophy", "cytomegalovirus" are not; "Valsalva
  maneuver" is). Words capitalised only because the title is written in Title Case are not proper
  nouns. Give each word in its correct capitalisation.
- paragraphs: the role of every M paragraph: running_title, article_type, title, authors, affiliation,
  correspondence, label (a heading or label with no content of its own, e.g. "Title page", "Authors:"),
  other_front_matter (word counts, funding, conflicts, key words, any other title-page content) or
  body (abstract or main text)."""


# --- sources -------------------------------------------------------------------------------

@dataclass
class Meta:
    """What the submission system's XML says (portal format or JATS)."""
    article_type: str = ""
    title: str = ""
    running_title: str = ""
    authors: list[dict] = field(default_factory=list)  # given, surname, prefix, email, corresponding, aff
    dates: dict = field(default_factory=dict)  # received / revised / accepted -> datetime
    keywords: str = ""  # the key words as submitted ("Orbital cellulitis, pansinusitis,intracranial extension")

    def text(self) -> str:
        parts = [self.article_type, self.title, self.running_title]
        for a in self.authors:
            parts += [a.get("given", ""), a.get("surname", ""), a.get("email", "")]
            parts += list((a.get("aff") or {}).values())
        return "\n".join(p for p in parts if p)


@dataclass
class Sources:
    first_page: list[str] = field(default_factory=list)  # paragraphs of the title-page file
    meta: Meta = field(default_factory=Meta)
    # paragraphs of other submission files named for author contributions ("Author Contribution Form.docx"),
    # read by the back-matter rule (rules_sections) for the authors' contributions section
    contribution_forms: list[list[str]] = field(default_factory=list)
    # file names, for the review screen's "Source" chip
    first_page_name: str = ""
    xml_name: str = ""
    contribution_form_names: list[str] = field(default_factory=list)
    # the bold lead of title-page / contribution-form paragraphs ("Concept:"), by paragraph text
    bold_leads: dict[str, str] = field(default_factory=dict)
    # paragraphs of the submission's declaration files ("JEHP_EthicsDeclarations.docx", "Conflict of Interest.docx"):
    # funding, conflicts, ethics, data and AI statements the back matter uses when the manuscript has none
    declaration_files: list[list[str]] = field(default_factory=list)
    declaration_file_names: list[str] = field(default_factory=list)
    file_names: list[str] = field(default_factory=list)  # every file of the submission ("First author photograph.jpg")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_paragraphs(data: bytes) -> list[str]:
    return [t for t, _ in docx_paragraphs_bold(data)]


def _on_prop(rpr, tag: str) -> bool:
    el = rpr.find(_W + tag) if rpr is not None else None
    return el is not None and el.get(_W + "val") not in ("0", "false", "none")


def docx_paragraphs_bold(data: bytes) -> list[tuple[str, str]]:
    """(text, its bold lead) of each non-empty paragraph: "Concept:" for a paragraph that starts "**Concept:**
    Mehmet ..." (the back matter keeps a contribution list's bold labels: IJO_1347), "" when it does not start bold."""
    from .docmodel import paragraph_text
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    out = []
    for p in root.iter(_W + "p"):
        t = paragraph_text(p).strip()
        if not t:
            continue
        lead = ""
        for r in p.iter(_W + "r"):
            if any(a.tag == _W + "del" for a in r.iterancestors()):
                continue
            rt = "".join(x.text or "" for x in r.iter(_W + "t"))
            if not rt:
                continue
            if not _on_prop(r.find(_W + "rPr"), "b"):
                if rt.strip() or not lead.strip():
                    break
            lead += rt
        lead = lead.strip()
        out.append((t, lead if lead and t.startswith(lead) and lead != t else ""))
    return out


def doc_paragraphs(data: bytes) -> list[str]:
    """The paragraphs of a legacy Word 97-2003 .doc (the portal's "FirstPage.doc"): the main text read from the
    file's piece table (pure Python, `olefile`; the server image has no LibreOffice). Field instructions are left
    out (their results kept), a line break is "\\n", table cells are paragraphs. Where the file cannot be read
    that way (Word 6, encrypted, damaged), LibreOffice converts it when it is installed; else nothing."""
    try:
        out = _doc_text_ole(data)
    except Exception:  # noqa: BLE001 - any malformed file: try the converter, else nothing
        out = None
    if out is None:
        out = _doc_text_soffice(data)
    return [t for t in (re.sub(r"[ \t]+\n|\n[ \t]+", "\n", x).strip() for x in out or []) if t]


def _doc_text_ole(data: bytes) -> list[str] | None:
    import struct

    import olefile

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        wd = ole.openstream("WordDocument").read()
        ident, _, _, _, _, flags = struct.unpack_from("<HHHHHH", wd, 0)
        if ident != 0xA5EC or flags & 0x0100:  # not Word 97+, or encrypted
            return None
        table = "1Table" if flags & 0x0200 else "0Table"
        if not ole.exists(table):
            return None
        tbl = ole.openstream(table).read()
    finally:
        ole.close()
    ccp_text = struct.unpack_from("<i", wd, 0x4C)[0]       # FibRgLw97.ccpText: the main document's length
    fc_clx, lcb_clx = struct.unpack_from("<II", wd, 0x1A2)  # FibRgFcLcb97.fcClx / lcbClx
    clx = tbl[fc_clx:fc_clx + lcb_clx]
    i = 0
    while i < len(clx) and clx[i] == 0x01:  # Prc entries (property modifiers) before the piece table
        i += 3 + struct.unpack_from("<H", clx, i + 1)[0]
    if i >= len(clx) or clx[i] != 0x02:
        return None
    lcb = struct.unpack_from("<I", clx, i + 1)[0]
    plc = clx[i + 5:i + 5 + lcb]
    n = (lcb - 4) // 12
    cps = struct.unpack_from(f"<{n + 1}i", plc, 0)
    chars = []
    for k in range(n):
        if cps[k] >= ccp_text:
            break
        fc = struct.unpack_from("<I", plc, 4 * (n + 1) + 8 * k + 2)[0]
        cnt = min(cps[k + 1], ccp_text) - cps[k]
        if fc & 0x40000000:  # 8-bit text (cp1252) at fc / 2
            at = (fc & 0x3FFFFFFF) // 2
            chars.append(wd[at:at + cnt].decode("cp1252", errors="replace"))
        else:
            chars.append(wd[fc:fc + 2 * cnt].decode("utf-16-le", errors="replace"))
    text = "".join(chars)
    # fields: \x13 instruction \x14 result \x15 -> the result
    out, depth_skip = [], []
    for ch in text:
        if ch == "\x13":
            depth_skip.append(True)
        elif ch == "\x14":
            if depth_skip:
                depth_skip[-1] = False
        elif ch == "\x15":
            if depth_skip:
                depth_skip.pop()
        elif not (depth_skip and depth_skip[-1]):
            out.append(ch)
    text = "".join(out)
    text = text.replace("\x0b", "\n").replace("\x1e", "-").replace("\xa0", " ")
    text = re.sub(r"[\x00-\x06\x08\x0e-\x1f]", "", text)
    return re.split(r"[\r\x07\x0c]", text)


def _doc_text_soffice(data: bytes) -> list[str] | None:
    import shutil
    import subprocess
    import tempfile

    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "first.doc")
        with open(src, "wb") as f:
            f.write(data)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "docx", "--outdir", tmp, src], capture_output=True,
                           timeout=120, check=False)
            with open(os.path.join(tmp, "first.docx"), "rb") as f:
                return docx_paragraphs(f.read())
        except (OSError, subprocess.SubprocessError, zipfile.BadZipFile, KeyError, etree.XMLSyntaxError):
            return None


def file_paragraphs(name: str, data: bytes) -> list[str]:
    """The paragraphs of a submission's .docx or legacy .doc file."""
    if name.lower().endswith(".doc"):
        return doc_paragraphs(data)
    return docx_paragraphs(data)


_FIRST_PAGE = re.compile(r"first\s*page|title[\s._-]*page|^fp\.docx?$|^cover\s*page", re.I)
_WORD_FILE = re.compile(r"\.docx?$", re.I)
_DECLARATION_FILE = re.compile(r"(?i)declar|disclos|conflicts?\b|\bcoi\b|funding|ethics?\b|statements?\b")


def sources_from_zip(path: str) -> Sources:
    """Title-page file and XML from a submission zip (original.zip). The title page may be a legacy .doc
    ("FirstPage.doc", jehp_1039): read too; a .docx of the same kind wins."""
    z = zipfile.ZipFile(path)
    src = Sources()
    names = z.namelist()
    src.file_names = [os.path.basename(n) for n in names if not n.startswith("__MACOSX") and not n.endswith("/")]
    fp = [n for n in names if _WORD_FILE.search(n) and _FIRST_PAGE.search(os.path.basename(n))
          and not n.startswith("__MACOSX")]
    if fp:
        pick = sorted(fp, key=lambda n: (not n.lower().endswith(".docx"), len(n)))[0]
        try:
            src.first_page = file_paragraphs(pick, z.read(pick))
            src.first_page_name = os.path.basename(pick)
            if pick.lower().endswith(".docx"):
                src.bold_leads.update({t: b for t, b in docx_paragraphs_bold(z.read(pick)) if b})
        except (KeyError, zipfile.BadZipFile, etree.XMLSyntaxError):
            pass
    xml = [n for n in names if n.lower().endswith(".xml") and "/" not in n.strip("/")]
    if xml:
        src.meta = parse_xml(z.read(xml[0]))
        src.xml_name = os.path.basename(xml[0])
    meta_json = [n for n in names if os.path.basename(n).lower() == "article_metadata.json"]
    if meta_json and not src.meta.authors:  # no XML (or none with authors): the portal export's metadata file
        got = parse_metadata_json(z.read(meta_json[0]))
        if got.authors:
            got.article_type = src.meta.article_type or got.article_type
            got.title = src.meta.title or got.title
            got.dates = src.meta.dates or got.dates
            src.meta = got
            src.xml_name = os.path.basename(meta_json[0])
    for n in sorted(names):
        base = os.path.basename(n)
        if _WORD_FILE.search(n) and _DECLARATION_FILE.search(base) and not _FIRST_PAGE.search(base) \
                and not re.search(r"(?i)contribut|patient|consent\s*form", base) and not n.startswith("__MACOSX"):
            try:
                src.declaration_files.append(file_paragraphs(n, z.read(n)))
                src.declaration_file_names.append(base)
            except (KeyError, zipfile.BadZipFile, etree.XMLSyntaxError):
                pass
        if _WORD_FILE.search(n) and re.search(r"contribut", os.path.basename(n), re.I) \
                and not _FIRST_PAGE.search(os.path.basename(n)):
            try:
                src.contribution_forms.append(file_paragraphs(n, z.read(n)))
                if n.lower().endswith(".docx"):
                    src.bold_leads.update({t: b for t, b in docx_paragraphs_bold(z.read(n)) if b})
                src.contribution_form_names.append(os.path.basename(n))
            except (KeyError, zipfile.BadZipFile, etree.XMLSyntaxError):
                pass
    return src


def sources_from_files(first_page: str | None = None, xml: str | None = None) -> Sources:
    src = Sources()
    if first_page:
        with open(first_page, "rb") as f:
            src.first_page = file_paragraphs(first_page, f.read())
        src.first_page_name = os.path.basename(first_page)
    if xml:
        with open(xml, "rb") as f:
            src.meta = parse_xml(f.read())
        src.xml_name = os.path.basename(xml)
    return src


def _date(s: str) -> datetime | None:
    s = (s or "").strip()
    for fmt in ("%d-%b-%y", "%d-%b-%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def _txt(el, path: str) -> str:
    x = el.find(path)
    return re.sub(r"\s+", " ", "".join(x.itertext())).strip() if x is not None else ""


def _given(first: str, middle: str) -> str:
    """The portal's first and middle names; a middle name that repeats the first ("Khadija" / "Khadija") once."""
    if middle and middle.casefold() != first.casefold():
        return f"{first} {middle}".strip()
    return first


def parse_xml(data: bytes) -> Meta:
    try:
        root = etree.fromstring(data, etree.XMLParser(recover=True, resolve_entities=False, no_network=True))
    except etree.XMLSyntaxError:
        return Meta()
    if root is None:
        return Meta()
    m = Meta()
    if root.tag == "article":  # JATS
        m.article_type = _txt(root, ".//subj-group[@subj-group-type='Article Type']/subject") or root.get("article-type", "")
        m.title = _txt(root, ".//title-group/article-title")
        m.running_title = _txt(root, ".//alt-title[@alt-title-type='running-head']")
        m.keywords = ", ".join(_txt(k, ".") for k in root.iter("kwd") if _txt(k, "."))
        affs = {a.get("id"): {"institution": _txt(a, "institution") or _txt(a, "."), "country": _txt(a, "country")}
                for a in root.iter("aff")}
        for c in root.iter("contrib"):
            if c.get("contrib-type") != "author":
                continue
            x = c.find("xref[@ref-type='aff']")
            m.authors.append(dict(given=_txt(c, "name/given-names"), surname=_txt(c, "name/surname").rstrip("."),
                                  prefix=_txt(c, "name/prefix"), email=_txt(c, "email"),
                                  corresponding=c.get("corresp") == "yes",
                                  aff=affs.get(x.get("rid")) if x is not None else None))
        for d in root.iter("date"):
            try:
                when = datetime(int(_txt(d, "year")), int(_txt(d, "month")), int(_txt(d, "day")))
            except ValueError:
                continue
            kind = {"received": "received", "rev-recd": "revised", "accepted": "accepted"}.get(d.get("date-type"))
            if kind and (kind != "revised" or when > m.dates.get(kind, datetime.min)):
                m.dates[kind] = when
        return m
    m.article_type = _txt(root, "articletype")
    m.title = _txt(root, "title")
    m.keywords = _txt(root, "keywords")
    orders = []
    for a in root.iter("author"):
        orders.append(_txt(a, "author_order"))
        m.authors.append(dict(
            # The portal often repeats the first name as the middle name ("Khadija" / "Khadija"): once.
            given=_given(_txt(a, "fname"), _txt(a, "mname")),
            surname=_txt(a, "lname"),
            prefix=_txt(a, "title"), email=_txt(a, "email"), corresponding=_txt(a, "is_corresponding") == "Y",
            zip=_txt(a, "corresponding_zip") or _txt(a, "zip"),
            aff={k: _txt(a, t) for k, t in (("department", "department"), ("institution", "institute"),
                                            ("city", "city"), ("state", "state"), ("country", "country"))}))
    m.authors = _by_author_order(m.authors, orders)
    for key, tag in (("received", "submission_date"), ("revised", "revision_date"), ("accepted", "Accepted")):
        d = _date(_txt(root, tag))
        if d:
            m.dates[key] = d
    return m


def _by_author_order(authors: list[dict], orders: list[str]) -> list[dict]:
    """The portal's authors in author-list order. <author_order> numbers the co-authors ("0", "1", ... or "1", "2",
    ...) and marks the submitting author "Primary"; the submitting author takes the place the numbers leave free
    (orders 0, 1 and Primary: third; orders 1, 2 and Primary: first; orders 0, 2, 3 and Primary: second), which is
    where the editors put them. Without numbers the XML's own order stands."""
    nums = [int(o) if o.strip().isdigit() else None for o in orders]
    if len(orders) != len(authors) or not any(n is not None for n in nums):
        return authors
    used, free = {n for n in nums if n is not None}, count(0)
    keys = []
    for k, n in enumerate(nums):
        if n is None and orders[k].strip().lower() in ("primary", "first", "main"):
            n = next(x for x in free if x not in used)
            used.add(n)
        keys.append((n if n is not None else 10 ** 6, k))
    return [authors[k] for _, k in sorted(keys)]


_PREFIX = re.compile(r"(?i)^(dr|prof|mr|mrs|ms|miss)\.?\s+")


def parse_metadata_json(data: bytes) -> Meta:
    """The portal export's article_metadata.json, for a submission without its XML: the authors (name with its
    prefix, e-mail, corresponding flag) in author_order."""
    m = Meta()
    try:
        d = json.loads(data.decode("utf-8-sig"))
    except (ValueError, UnicodeDecodeError):
        return m
    if not isinstance(d, dict):
        return m
    for k in ("article_type", "title", "running_title"):
        if isinstance(d.get(k), str):
            setattr(m, k, d[k].strip())
    authors, orders = [], []
    for a in d.get("authors") or []:
        if not isinstance(a, dict):
            continue
        name = re.sub(r"\s+", " ", str(a.get("name") or "")).strip()
        pm = _PREFIX.match(name)
        prefix = pm.group(1).capitalize() if pm else ""
        name = name[pm.end():] if pm else name
        if not name:
            continue
        given, _, surname = name.rpartition(" ")
        corr = a.get("corresponding_author", a.get("corresponding"))
        authors.append(dict(given=given, surname=surname, prefix=prefix, email=str(a.get("email") or "").strip(),
                            corresponding=corr is True or str(corr).lower() in ("y", "yes", "true", "1"), aff={}))
        orders.append(str(a.get("author_order") or ""))
    m.authors = _by_author_order(authors, orders)
    return m


# --- Claude --------------------------------------------------------------------------------

class _Cache:
    def __init__(self, path: str = ai_refs.CACHE_DB):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.con = sqlite3.connect(path, check_same_thread=False)  # shared by worker threads (serialized SQLite)
        self.con.execute("CREATE TABLE IF NOT EXISTS frontmatter (key TEXT PRIMARY KEY, body TEXT)")

    @staticmethod
    def key(text: str) -> str:
        return hashlib.sha256(f"{MODEL}\n{INSTRUCTIONS}\n{json.dumps(SCHEMA)}\n{text}".encode()).hexdigest()

    def get(self, text: str) -> dict | None:
        row = self.con.execute("SELECT body FROM frontmatter WHERE key = ?", (self.key(text),)).fetchone()
        return json.loads(row[0]) if row else None

    def put(self, text: str, fields: dict) -> None:
        self.con.execute("INSERT OR REPLACE INTO frontmatter VALUES (?, ?)", (self.key(text), json.dumps(fields)))
        self.con.commit()


_cache: _Cache | None = None


def _call(text: str) -> dict | None:
    import anthropic

    client = anthropic.Anthropic(
        api_key=ai_refs._api_key(),
        base_url=os.environ.get("EDITOS_ANTHROPIC_BASE_URL", "https://api.anthropic.com"),
    )
    response = client.messages.create(
        model=MODEL, max_tokens=8000, system=INSTRUCTIONS,
        output_config={"effort": "low", "format": {"type": "json_schema", "schema": SCHEMA}},
        messages=[{"role": "user", "content": text}],
    )
    USAGE["calls"] += 1
    USAGE["input_tokens"] += response.usage.input_tokens
    USAGE["output_tokens"] += response.usage.output_tokens
    spend.add_usage("front_matter", MODEL, response.usage)
    if response.stop_reason in ("refusal", "max_tokens"):
        return None
    try:
        return json.loads(next(b.text for b in response.content if b.type == "text"))
    except (StopIteration, ValueError):
        return None


def split(text: str) -> dict | None:
    """Claude's field split of the numbered title-page text; cached by text."""
    global _cache
    if not ai_refs.available():
        return None
    if _cache is None:
        _cache = _Cache()
    got = _cache.get(text)
    if got is None:
        got = _call(text)
        if got is not None:
            _cache.put(text, got)
    return got


# --- reading the inputs --------------------------------------------------------------------

_BODY_START = re.compile(r"^\s*(?:<H\d>)?\s*(abstract|introduction|background|summary|key\s*-?\s*words?|"
                         r"purpose|aims?|objectives?|context|dear (?:editor|sir)|sir,)\b", re.I)


_BACK_MATTER = re.compile(r"^\s*(?:<H\d>)?\s*(references?|bibliography|declaration|financial support|conflicts? of "
                          r"interest|acknowledge?ments?|figure|table|legends?)\b", re.I)
# Title-page labels the editors delete ("Title Page", "Running title:", "Authors:").
_LABEL = re.compile(r"^\s*(title\s*page|running\s*(?:title|head)|short\s*title|(?:list of\s+)?authors?(?:\s*names?)?|"
                    r"author(?:'s|s')?\s*(?:names?|details|information)|contributors?|affiliations?|"
                    r"department\(?s?\)?\s*(?:and|&)\s*institution\(?s?\)?|corresponding\s*author|"
                    r"(?:article\s*)?title|manuscript\s*type|type\s*of\s*(?:manuscript|article)|"
                    r"(?:article|manuscript)\s+file|main\s+(?:document|manuscript)|blinded\s+manuscript)\b", re.I)


# Title-page blocks end here too: the key points / plain language summary IJO asks for come after it.
_LEAD_STOP = re.compile(r"^\s*(?:<H\d>)?\s*(?:key\s+(?:learning\s+)?(?:points?|messages?)|plain\s+language\s+summary)\b", re.I)


def leading_paragraphs(ms: Manuscript, limit: int = 30) -> list[Paragraph]:
    """Manuscript paragraphs before the abstract / main text: where a title page inside the
    manuscript would be. The submission's own parts (cover letter, reply to the
    reviewers, counts and phone lines `submission.py` deletes) are skipped, so a title page after a cover
    letter is found."""
    out = []
    skip = submission.deleted_ids(ms, getattr(ms, "_submission_style", None))
    for p in ms.paragraphs:
        if p.in_table or not p.text.strip() or p.id in skip:
            continue
        if _BODY_START.match(p.text) or _BACK_MATTER.match(p.text) or _LEAD_STOP.match(p.text) \
                or len(p.text) > 400 or len(out) >= limit:
            break
        out.append(p)
    return out


def prompt_text(src: Sources, lead: list[Paragraph]) -> tuple[str, dict[str, str]]:
    """The numbered text Claude sees, and M-id -> paragraph id."""
    lines, ids, size = [], {}, 0
    first_page = submission.texts_without_parts(src.first_page) if src.first_page else []
    if first_page:
        lines.append("Title-page file:")
        for i, t in enumerate(first_page[:80], 1):
            size += len(t)
            if size > 12000:
                break
            lines.append(f"[F{i}] {t}")
    if lead:
        lines.append("\nStart of the manuscript file:")
        for i, p in enumerate(lead, 1):
            ids[f"M{i}"] = p.id
            lines.append(f"[M{i}] {p.text[:1500]}")
    return "\n".join(lines), ids


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "").replace(" ", " ")
    s = re.sub(r"[‐-―−]", "-", s)
    return re.sub(r"[\s.,;:]+", "", s).casefold()


@dataclass
class Author:
    given: str
    surname: str
    affs: list[int] = field(default_factory=list)
    corresponding: bool = False
    prefix: str = ""
    written: str = ""  # the name as the title page writes it, before any expansion from the XML ("Dewi S. Rochmayani")


@dataclass
class Fields:
    article_type: str = ""
    title: str = ""
    running_title: str = ""
    authors: list[Author] = field(default_factory=list)
    affiliations: dict[int, dict] = field(default_factory=dict)
    correspondence: dict = field(default_factory=dict)
    proper_nouns: list[str] = field(default_factory=list)
    roles: dict[str, str] = field(default_factory=dict)  # manuscript para id -> role
    notes: list[str] = field(default_factory=list)  # for the editor: what was dropped and why
    # Equal contribution: the mark after the authors' names ("†"), its footnote as written, the authors (indexes)
    equal_mark: str = ""
    equal_note: str = ""
    equal_authors: list[int] = field(default_factory=list)
    # authors the submission XML names that the title page's author list leaves out (queried, not added)
    xml_only: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)  # for the editor (the sheet's "raise query internally")


def verify(raw: dict, source: str, para_ids: dict[str, str]) -> Fields:
    """Keep only what is in the source text. Model output never reaches the file unchecked."""
    src = _norm(source)
    f = Fields()

    def ok(value: str, what: str) -> str:
        value = re.sub(r"\s+", " ", (value or "").strip())
        if not value:
            return ""
        if _norm(value) in src:
            return value
        f.notes.append(f"{what} {value!r} is not in the source text; left out")
        return ""

    f.article_type = ok(raw.get("article_type", ""), "article type")
    f.title = ok(raw.get("title", ""), "title")
    f.running_title = ok(raw.get("running_title", ""), "running title")
    for a in raw.get("authors", []):
        given, surname = ok(a.get("given_names", ""), "given name"), ok(a.get("surname", ""), "surname")
        if not surname and given and " " not in given:
            given, surname = "", given  # one name only ("Sonakshi"): the sheet's "Punithavathy, ..., Stalin"
        elif not surname and given:  # "Retno Setyowati" with the surname left empty: the last word is the surname
            given, surname = given.rsplit(" ", 1)
        if not surname:
            continue
        prefix = a.get("prefix", "") if re.fullmatch(r"(?i)dr|prof", a.get("prefix", "").strip(" .")) else ""
        f.authors.append(Author(given, surname, [i for i in a.get("affiliations", []) if isinstance(i, int)],
                                bool(a.get("corresponding")), ok(prefix.strip(" ."), "prefix")))
    for aff in raw.get("affiliations", []):
        parts = {k: ok(aff.get(k, ""), k) for k in ("department", "institution", "city", "state", "country")}
        parts["department"] = strip_designation(parts["department"])
        parts = {k: ("" if re.search(r"\d{5,}|@", v) else v) for k, v in parts.items()}
        pin = affiliation_postal_code(source, parts)
        if pin:
            parts["postal_code"] = pin
        if parts["department"] or parts["institution"]:
            f.affiliations[aff.get("id", len(f.affiliations) + 1)] = parts
    c = raw.get("correspondence") or {}
    corr = {k: ok(c.get(k, ""), f"correspondence {k}") for k in
            ("given_names", "surname", "city", "state", "postal_code", "country", "email")}
    corr["prefix"] = ok(c.get("prefix", "").strip(" ."), "prefix") \
        if re.fullmatch(r"(?i)dr|prof|mrs|mr|ms", c.get("prefix", "").strip(" .")) else ""
    corr["address_parts"] = [x for x in (ok(p, "address part") for p in c.get("address_parts", [])) if x]
    corr["address_parts"] = [p for p in corr["address_parts"] if not _CONTACT.search(p)]
    corr["address_parts"] += _skipped_address_pieces(source, corr)
    if any(v for v in corr.values()):
        f.correspondence = corr
    title_words = _norm(" ".join((raw.get("title", ""), raw.get("running_title", ""))))
    f.proper_nouns = [w for w in raw.get("proper_nouns", []) if w.strip() and _norm(w) in title_words]
    for item in raw.get("paragraphs", []):
        if item.get("id") in para_ids and item.get("role") in ROLES:
            f.roles[para_ids[item["id"]]] = item["role"]
    return f


def _skipped_address_pieces(source: str, corr: dict) -> list[str]:
    """Pieces of the correspondence address the split left out between its last part and the city ("..., 1604. Cadde
    No: 9, Çankaya, 06800 Ankara, Türkiye": the district Çankaya, IJO_1347): read back from the same line, without
    the postal code, city, state and country."""
    parts = corr.get("address_parts") or []
    if not parts:
        return []
    last = parts[-1]
    places = [_norm(corr.get(k, "")) for k in ("city", "state", "country", "postal_code") if corr.get(k)]
    for line in source.splitlines():
        at = line.find(last)
        if at < 0:
            continue
        rest = line[at + len(last):]
        out = []
        for piece in re.split(r"\s*[,;]\s*", rest.strip(" ,.;")):
            piece = piece.strip(" .")
            n = _norm(piece)
            if not piece or not re.search(r"[^\W\d_]", piece) or _CONTACT.search(piece) or re.search(r"\d{4,}", piece) \
                    or any(x and (x in n or n in x) for x in places) or n in _norm(" ".join(parts)) \
                    or piece.casefold() in _COUNTRY_NAMES or len(piece.split()) > 4:
                break  # the city / postal code / country reached: nothing more is an address part
            out.append(piece)
        return out
    return []


_CORRESPONDENCE_LINE = re.compile(r"(?i)correspond|address\s+for|\b(?:phone|tel|mobile|mob|fax)\b")


def affiliation_postal_code(source: str, aff: dict) -> str:
    """The postal code the title page writes next to an affiliation's city ("Chennai - 600006", "Tokyo 102-8554",
    "75006 Paris", Iran's ten digits "Hamadan 6517838736"), on a line that also names the affiliation's institution
    or department; "" when none. The model is told to leave postal codes out; this reads them back from the text
    (S: "All the affiliations should be complete": owner, 2026-10-26)."""
    city = (aff.get("city") or "").strip()
    unit = _norm(aff.get("institution") or aff.get("department") or "")
    if not city or not unit:
        return ""
    c = re.escape(city)
    after = re.compile(rf"(?i)\b{c}\b\s*[,\-–]?\s*(\d{{5}}-?\d{{5}}|\d{{3}}\s?\d{{3}}|\d{{3}}-\d{{4}}|\d{{4,5}})(?![\d/])")
    before = re.compile(rf"(?i)(?<![\d\-/])(\d{{4,6}})\s+{c}\b")
    for line in source.splitlines():
        if unit not in _norm(line) or _CORRESPONDENCE_LINE.search(line):
            continue  # the correspondence address may be another building's (or a home's) postal code
        m = after.search(line) or before.search(line)
        if m:
            return m.group(1)
    return ""


def _complete_affiliations(affs: dict[int, dict], order: list[int], notes: list[str]) -> None:
    """Stylesheets: "All the affiliations should be complete". A department written without its institution and
    place, above or below departments that have them ("Dept. of Dermatology," / "Dept of Psychiatry, Dr. D.Y. Patil
    Medical College, ..., Pune, India"), shares that institution and place: the nearest following one, else the
    nearest before."""
    for k, i in enumerate(order):
        a = affs[i]
        if not a.get("department") or a.get("institution") or a.get("city") or a.get("state"):
            continue
        near = [affs[j] for j in order[k + 1:]] + [affs[j] for j in reversed(order[:k])]
        donor = next((d for d in near if d.get("institution") and (d.get("city") or d.get("country"))), None)
        if donor is None:
            continue
        for part in ("institution", "city", "state", "country"):
            if not a.get(part):
                a[part] = donor.get(part, "")
        notes.append(f"affiliation {a['department']!r}: institution and place taken from {donor['institution']!r}")


_CONTACT = re.compile(r"@|https?://|www\.|\b(?:phone|ph|tel|telephone|mobile|mob|fax|cell)\b|\+\d|\d{5,}\s*\d{3,}", re.I)


_COUNTRY_NAMES = {c.casefold() for c in (
    "India", "Iran", "Iraq", "Indonesia", "Malaysia", "Thailand", "Vietnam", "Philippines", "Singapore", "China",
    "Taiwan", "Japan", "Korea", "South Korea", "Nepal", "Pakistan", "Bangladesh", "Sri Lanka", "Bhutan", "Afghanistan",
    "Saudi Arabia", "UAE", "United Arab Emirates", "Oman", "Qatar", "Kuwait", "Bahrain", "Jordan", "Lebanon", "Syria",
    "Turkey", "Türkiye", "Egypt", "Morocco", "Algeria", "Tunisia", "Libya", "Nigeria", "Ghana", "Kenya", "Ethiopia",
    "Uganda", "Tanzania", "South Africa", "Cameroon", "Sudan", "USA", "United States", "Canada", "Mexico", "Brazil",
    "Chile", "Peru", "Colombia", "Argentina", "UK", "United Kingdom", "England", "Scotland", "Ireland", "Germany",
    "France", "Italy", "Spain", "Portugal", "Netherlands", "Belgium", "Sweden", "Norway", "Denmark", "Finland",
    "Switzerland", "Austria", "Poland", "Greece", "Israel", "Russia", "Ukraine", "Australia", "New Zealand",
    "Timor-Leste", "Myanmar", "Cambodia", "Laos", "Mongolia", "Kazakhstan", "Uzbekistan", "Yemen")}


def _xml_affiliation(aff: dict) -> dict:
    """A portal affiliation whose institution field carries the place ("Ilam University of Medical sciences, Ilam,
    Iran" with <city>ilam</city> <state>ilam</state>): the country and the city come out of the institution (as
    written there), and a state that only repeats the city goes."""
    if not aff:
        return aff
    aff = {k: v for k, v in aff.items() if re.search(r"[A-Za-z]", v or "")}  # "أبو ظبي", "گرگان": not in the file's script
    dept, inst = _norm(aff.get("department", "")), _norm(aff.get("institution", ""))
    if dept and inst and inst in dept:
        # the whole affiliation in both fields: the department field holds it
        aff.pop("institution")
        whole = [x.strip() for x in aff["department"].strip(" .,").split(",") if x.strip()]
        if len(whole) > 1 and whole[-1].casefold() in _COUNTRY_NAMES and not aff.get("country"):
            aff["country"] = whole.pop()
            aff["department"] = ", ".join(whole)
        for k in ("city", "state"):
            if aff.get(k) and _norm(aff[k]) in _norm(aff["department"]):
                aff.pop(k)
        return aff
    parts = [x.strip() for x in (aff.get("institution") or "").strip(" .,").split(",") if x.strip()]
    if len(parts) > 1 and parts[-1].casefold().strip(" .") in _COUNTRY_NAMES:
        country = parts.pop().strip(" .")
        if not aff.get("country") or aff["country"].casefold() == country.casefold():
            aff["country"] = country
        else:
            parts.append(country)
    if len(parts) > 1 and aff.get("city") and parts[-1].casefold() == aff["city"].strip().casefold():
        aff["city"] = parts.pop()
    if parts:
        aff["institution"] = ", ".join(parts)
    if aff.get("state") and aff["state"].strip().casefold() == (aff.get("city") or "").strip().casefold():
        del aff["state"]
    if aff.get("city") and aff["city"].islower():
        aff["city"] = aff["city"].title()
    return {k: v for k, v in aff.items() if v}


def _ends_with_country(text: str) -> bool:
    parts = [x.strip(" .") for x in text.strip(" .").split(",") if x.strip(" .")]
    return len(parts) > 1 and parts[-1].casefold() in _COUNTRY_NAMES


def _xml_chunks(text: str) -> list[str]:
    """A portal field holding two affiliations run together ("Psychiatric University Center Agadir, Morocco.Health
    Sciences Research Laboratory, ..., Agadir, Morocco."): split after each country."""
    out, rest = [], (text or "").strip()
    while True:
        m = next((m for m in re.finditer(r"(?<=[A-Za-z])\.\s*(?=[A-Z])", rest)
                  if _ends_with_country(rest[:m.start()])), None)
        if m is None:
            break
        out.append(rest[:m.start()].strip(" ,."))
        rest = rest[m.end():]
    if rest.strip(" ,."):
        out.append(rest.strip(" ,."))
    return out


def _covered(part: str, whole: str) -> bool:
    """Most words of `part` are in `whole`, typos allowed ("Faculty of medecine Agadir" in "..., Faculty of Medicine,
    Ibn Zohr University, Agadir")."""
    pw = [w for w in re.findall(r"\w+", part.casefold()) if w not in _MINOR]
    ww = re.findall(r"\w+", whole.casefold())
    if not pw:
        return True
    hit = sum(any(w == x or len(w) > 4 and difflib.SequenceMatcher(None, w, x).ratio() >= 0.8 for x in ww) for w in pw)
    return hit * 2 >= len(pw)


def _glued_city(parts: dict, city: str) -> dict:
    """"Psychiatric University Center Agadir" with the city Agadir: the city is split off."""
    inst, city = parts.get("institution", ""), (city or parts.get("city", "")).strip()
    if city and len(inst.split()) > 2 and inst.casefold().endswith(" " + city.casefold()):
        parts["institution"] = inst[:-len(city)].strip(" ,")
        parts["city"] = parts.get("city") or inst[-len(city):]
    return parts


def _xml_affiliations(aff: dict) -> list[dict]:
    """The portal author's affiliation(s): one (`_xml_affiliation`), or more when a field runs two together, or the
    department field holds a whole affiliation ("Department of Epidemiologiy, Faculty of medicine, Ibn Zohr
    University,Agadir,Morocco." with the institution "Faculty of medecine Agadir"): the whole one stands, and the
    institution field is added only when it says more (jehp_1402)."""
    aff = {k: v.strip() for k, v in aff.items() if v and v.strip()}
    dept, inst = aff.get("department", ""), aff.get("institution", "")
    dchunks, ichunks = _xml_chunks(dept), _xml_chunks(inst)
    xcity = aff.get("city", "") if re.search(r"[A-Za-z]", aff.get("city", "")) else ""
    if len(ichunks) > 1 and len(dchunks) <= 1:
        out = []
        for n, chunk in enumerate(ichunks):
            parts = _glued_city(_affiliation_parts(chunk), xcity)
            if n == 0 and dept and not parts["department"] and (_bare_specialty(dept) or _norm(dept) not in _norm(chunk)):
                # "Psychiatric" over "Psychiatric University Center ...": "Department of Psychiatric" (the editors)
                parts["department"] = "Department of " + dept if _bare_specialty(dept) else dept
            out.append({k: v for k, v in parts.items() if v})
        return out
    if len(dchunks) > 1 or _ends_with_country(dept) and _norm(inst) not in _norm(dept):
        out = []
        icity = _affiliation_parts(inst).get("city", "") if _ends_with_country(inst) else ""
        for n, chunk in enumerate(dchunks):
            parts = _glued_city(_affiliation_parts(chunk), xcity or icity)
            if n == len(dchunks) - 1 and inst and not _covered(inst, chunk):
                ip = _glued_city(_affiliation_parts(inst), xcity or icity) if _ends_with_country(inst) else \
                    dict(institution=re.sub(r"\s+" + re.escape(xcity) + r"$", "", inst, flags=re.I) if xcity else inst)
                parts["institution"] = ", ".join(x for x in (parts.get("institution"), ip.get("department"),
                                                              ip.get("institution")) if x)
                for k in ("city", "country"):
                    parts[k] = parts.get(k) or ip.get(k, "")
            if xcity and not parts.get("city"):
                parts["city"] = xcity.title() if xcity.isupper() or xcity.islower() else xcity
            out.append({k: v for k, v in parts.items() if v})
        return out
    one = _xml_affiliation(aff)
    return [one] if one else []


def fields_from_meta(m: Meta) -> Fields:
    """Fallback without Claude: the XML's own fields."""
    f = Fields(article_type=m.article_type, title=m.title, running_title=m.running_title)
    keys: dict[tuple, int] = {}
    for a in m.authors:
        if not a.get("surname"):
            continue
        ids = []
        for aff in _xml_affiliations({k: v for k, v in (a.get("aff") or {}).items() if v}):
            code = aff.get("state", "")
            if re.fullmatch(r"[A-Z]{2}", code) and code in _STATE_CODES_IN and aff.get("country", "India").casefold() \
                    in ("india", "in"):
                # the portal's <state>MH</state>; a code that does not fit a known city ("Lucknow" / "UK") gives way
                # to the city's state
                aff["state"] = _CITY_STATE_IN.get(_city_key(aff.get("city", "")), _STATE_CODES_IN[code])
                aff["country"] = "India"
            if aff:
                # one affiliation however the portal's fields split it ("School of Nursing and Midwifery Tabriz" /
                # "University of Medical Sciences, Tabriz, Iran." and "..., School of Nursing and Midwifery" / "Tabriz
                # University of Medical Sciences, Tabriz, Iran."): the first spelling stands
                k = _norm(" ".join(aff.get(x, "") for x in ("department", "institution", "city", "state", "country")))
                if k not in keys:
                    keys[k] = len(keys) + 1
                    f.affiliations[keys[k]] = {x: aff.get(x, "") for x in ("department", "institution", "city", "state",
                                                                           "country")}
                if keys[k] not in ids:
                    ids.append(keys[k])
        f.authors.append(Author(a.get("given", ""), a["surname"], ids, a.get("corresponding", False),
                                a.get("prefix", "") if re.fullmatch(r"(?i)dr|prof", a.get("prefix", "").strip(" .")) else ""))
    corr = next((a for a in m.authors if a.get("corresponding")), None)
    if corr:
        f.correspondence = dict(prefix="", given_names=corr.get("given", ""), surname=corr.get("surname", ""),
                                address_parts=[], city="", state="", postal_code="", country="",
                                email=corr.get("email", ""))
    return f


# --- text rules ----------------------------------------------------------------------------

_MINOR = {"a", "an", "the", "and", "but", "or", "nor", "for", "so", "yet", "as", "at", "by", "in", "of", "on",
          "to", "up", "via", "vs", "with", "without", "from", "into", "than", "per", "over", "versus", "&",
          "et", "al"}
_BREAK = re.compile(r"[:?!]$|^[–—-]$")


def _mostly_caps(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and sum(c.isupper() for c in letters) / len(letters) > 0.6


def _keep_map(keep: list[str]) -> dict[str, str]:
    out = {}
    for phrase in keep:
        for w in re.findall(r"[\w'’\-]+", phrase):
            if any(c.isupper() for c in w):
                out[w.casefold()] = w
                for part in re.split(r"[-\u2010\u2011]", w):  # "Timor-Leste": each part keeps its capital
                    if part[:1].isupper():
                        out.setdefault(part.casefold(), part)
    return out


_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
           "November", "December")


def case_body(ms: Manuscript) -> str:
    """The manuscript's running text before the references, for capitalisation evidence: sentences of 12 words
    or more, not written in Title Case (so the title and headings, wherever they stand, say nothing)."""
    from .references import find_heading

    if ms is None:
        return ""
    ref = find_heading(ms)
    out = []
    for p in ms.paragraphs[:ref] if ref is not None else ms.paragraphs:
        words = re.findall(r"[A-Za-z][\w'’-]*", p.text)
        if len(words) < 12:
            continue
        major = [w for w in words if w.lower() not in _MINOR]
        if major and sum(w[:1].isupper() for w in major) / len(major) >= 0.5:
            continue
        out.append(p.text)
    return "\n".join(out)


def _case_counts(word: str, body: str) -> tuple[int, int]:
    """(lower-case uses, capitalised uses inside a sentence) of `word` in `body`; a hyphenated name counts for
    each of its parts ("... in Timor-Leste")."""
    low, cap = word[0].lower() + word[1:], word[0].upper() + word[1:]
    lower = len(re.findall(r"(?<!\w)" + re.escape(low) + r"\b", body))
    caps = 0
    for m in re.finditer(r"(\S*)\s(?:[A-Z][\w'’]*[-‐‑])*" + re.escape(cap) + r"\b(?![-‐‑][a-z]|[\s\-‐‑]+[A-Z])", body):
        prev = m.group(1)
        if not prev or prev[0].isupper() or prev[-1] in ".!?:":
            continue
        before = re.search(r"(\S+)\s+$", body[:m.start()])
        if prev.lower() in _MINOR and before and before.group(1)[:1].isupper():
            continue  # "Materials and Methods": inside a capitalised phrase
        caps += 1
    return lower, caps


def proper_keep(texts: list[str], keep: list[str], body: str, notes: list[str] | None = None) -> list[str]:
    """The words that keep their capital in sentence case: the split's proper nouns, checked, and a deterministic
    backstop for the ones it misses (S: running head and title in sentence case "except proper nouns/abbreviations").
    A capitalised word of the title stays capitalised when the manuscript's running text writes it so inside its
    sentences more often than in lower case ("in Timor-Leste", "in Indonesia"), when the text says nothing and the
    dictionary lists it only capitalised (countries, nationalities, eponyms), or when it is a month next to a number
    ("2004-January 2026"). A proper noun of the split that the text writes in lower case and the dictionary does
    not list as a name ("Oral" in "Oral Disorders") is dropped, unless a model name or abbreviation follows it
    ("Femto LDV Z8")."""
    from .refengine import PROPER_NOUNS, _dictionary_words

    _, dict_caps = _dictionary_words()

    def named(w: str) -> bool:
        return w in PROPER_NOUNS or w in dict_caps

    out = []
    for k in keep:
        w = k.strip()
        brand = any(re.search(r"\b" + re.escape(w) + r"\s+(?:[A-Z]{2,}\s+)?[A-Z]*\d", t or "") for t in texts)
        if re.fullmatch(r"[A-Z][a-z'’]+", w) and not named(w) and not brand:  # "Femto LDV Z8": a product name
            lower, caps = _case_counts(w, body)
            if lower > caps:
                if notes is not None:
                    notes.append(f"{w!r} is not kept capitalised: the text writes it in lower case")
                continue
        out.append(k)
    for t in texts:
        for m in re.finditer(r"(?<![\w'’])[A-Z][a-z'’]+\b", t or ""):
            w = m.group(0)
            if w in _MONTHS:
                if re.search(r"\d[\s-]*$", t[:m.start()]) or re.match(r"[\s,-]*\d", t[m.end():]):
                    out.append(w)
                continue
            if re.search(r"(?:idae|aceae|inae|ales|oidea)$", w):
                continue  # a taxon above the genus ("Viperidae"): the editors write it in lower case in a title
            lower, caps = _case_counts(w, body)
            if caps > lower or (not lower and not caps and named(w)):
                out.append(w)
    return out


def recase(text: str, mode: str, keep: list[str], colon_capital: bool = True) -> str:
    """Sentence or title case, keeping proper nouns / abbreviations (`keep`) and words written with
    inner capitals or digits (OCT, KPro, COVID-19). The first word and the word after a colon,
    question mark or dash start with a capital."""
    if mode == "as_is":
        return text
    caps = _mostly_caps(text)
    known = _keep_map(keep)
    out, start, first = [], True, True
    for tok in re.split(r"(\s+)", text):
        if not tok or tok.isspace():
            out.append(tok)
            continue
        m = re.match(r"^([^\w]*)(.*?)([^\w]*)$", tok)
        lead, core, trail = m.groups()
        if core:
            pieces = re.split(r"([-\u2010\u2011\u2013\u2014])", core)
            parts, seps = pieces[::2], pieces[1::2]
            # An em dash inside a word starts a new phrase ("Interventions—A Systematic Review"); an en dash does
            # only before an article ("Interventions–A Review"), otherwise it joins a compound as a hyphen does
            # ("School Health–Related Outcomes" -> "school health–related outcomes": the author's dash kept).
            brk = [sep == "\u2014" or sep == "\u2013" and parts[k + 1].casefold() in ("a", "an", "the")
                   for k, sep in enumerate(seps)]
            new = []
            for i, part in enumerate(parts):
                low = part.casefold()
                if i and brk[i - 1]:
                    start = True
                if low in known:
                    w = known[low]
                elif not caps and (any(c.isupper() for c in part[1:]) or any(c.isdigit() for c in part)):
                    w = part
                elif mode == "title" and (start and i == 0 or low not in _MINOR):
                    w = part[:1].upper() + part[1:].lower() if caps or part[:1].islower() else part
                else:
                    w = part.lower()
                if start and (i == 0 or brk[i - 1]) and (first or colon_capital):
                    w = w[:1].upper() + w[1:]
                if i == 0 or brk[i - 1]:
                    start = False
                new.append(w)
            core = "".join(w + (seps[i] if i < len(seps) else "") for i, w in enumerate(new))
            start = first = False
        out.append(lead + core + trail)
        if _BREAK.search(tok.rstrip()):
            start = True
    return "".join(out)


def clean_title(text: str) -> str:
    t = re.sub(r"^\s*(?:article\s+)?title\s*[:\-–]\s*", "", text, flags=re.I)
    t = re.sub(r"\s+", " ", t).strip().strip("“”\"'‘’").strip()
    t = re.sub(r"\s+([.,:;?])", r"\1", t)
    return _drop_end_period(t)


def _drop_end_period(t: str) -> str:
    """Period at the end goes, unless the last word is an abbreviation."""
    t = t.rstrip()
    if t.endswith(".") and not re.search(r"\b(?:[A-Z]\.){2,}$|\bet al\.$", t):
        t = t[:-1].rstrip()
    return t


_PARTICLES = {"van", "von", "de", "der", "den", "da", "di", "del", "della", "dos", "du", "la", "le", "ter", "ten"}


def _name_case(word: str, first: bool) -> str:
    """ALL CAPS or all lower names get capitals: "MEERA" -> "Meera", "rao" -> "Rao"."""
    if not (word.isupper() and len(word) > 1 or word.islower()) or word in ("II", "III", "IV"):
        return word
    if not first and word.casefold() in _PARTICLES:
        return word.lower()
    return "-".join("'".join(x[:1].upper() + x[1:].lower() for x in part.split("'")) for part in word.split("-"))


def _is_initials(tok: str, all_caps_name: bool) -> bool:
    t = tok.replace(".", "")
    return bool(re.fullmatch(r"[A-Z]{1,3}", t)) and (not all_caps_name or len(t) == 1 or "." in tok)


def given_parts(given: str, all_caps: bool) -> list[tuple[str, bool]]:
    """Given names as (text, is_initial): "A.N." -> [("A", True), ("N", True)]."""
    out = []
    for tok in re.split(r"[\s]+", given.strip()):
        if not tok:
            continue
        if _is_initials(tok, all_caps):
            out += [(c, True) for c in tok.replace(".", "")]
        else:
            out.append((tok, False))
    return out


def format_name(a: Author, cfg: dict, abbreviate: bool | None = None) -> tuple[str, bool]:
    """'Given M. Surname' per the guide; returns (name, has_full_first_name)."""
    all_caps = (a.given + a.surname).isupper()
    parts = given_parts(a.given, all_caps)
    abbreviate = cfg.get("abbreviate_middle_names", False) if abbreviate is None else abbreviate
    dot = "." if cfg.get("initial_period", True) else ""
    words = []
    # a middle name stays in full when the first name is only an initial ("M. Banyameen Iqbal"): it is the one
    # name written out
    abbreviate = abbreviate and not (parts and parts[0][1])
    for i, (t, initial) in enumerate(parts):
        if initial or (i > 0 and abbreviate):
            words.append(t[0].upper() + dot)
        else:
            words.append(_name_case(t, i == 0))
    surname = " ".join(_name_case(w, False) for w in a.surname.split())
    full_first = bool(parts) and not parts[0][1]
    return " ".join(words + [surname]), full_first or not parts


def _initial_after_name(a: "Author") -> bool:
    """One name written out with its initial(s) after it ("Natarajan S", "Mathangi D.C", "Meera K R")."""
    toks = _name_tokens(a)
    if len(toks) < 2 or not re.fullmatch(r"[A-Z][a-z'’\-]{2,}|[A-Z]{3,}", toks[0].strip(",")):
        return False
    return all(re.fullmatch(r"(?:[A-Z]\.?){1,3}", t.strip(",")) for t in toks[1:])


def _initial_like(s: str) -> bool:
    return bool(re.fullmatch(r"[A-Z]\.?", s.strip()))


def _name_tokens(a: Author) -> list[str]:
    return [t for t in re.split(r"[\s]+", f"{a.given} {a.surname}".strip()) if t]


def _token_fits(short: str, full: str) -> bool:
    """A name word and its fuller form: the same word, or an initial of it ("A." / "Andi")."""
    s, f = short.strip(".").casefold(), full.strip(".").casefold()
    return s == f or len(s) == 1 and f[:1] == s


def unglued_name(a: Author, xml_names: list[Author]) -> tuple[str, str] | None:
    """(given, surname) for a name the title page runs together ("KavyaRaman") that the submission XML / notice
    gives as the same letters in more words (given "Kavya", surname "Raman"): the title page's letters, spaced and
    split as the XML has them ("Kavya", "Raman"). A glued name is a typo; only the same letters in the same order
    (case and spaces aside) count, and only when the XML has more words."""
    written = f"{a.given} {a.surname}".strip()
    letters = re.sub(r"\s+", "", written)
    if not letters or not re.fullmatch(r"[^\W\d_]+(?:[\s'’.-]+[^\W\d_]+)*", written):
        return None
    for x in xml_names:
        if not x.given.strip() or not x.surname.strip():
            continue
        xml = f"{x.given.strip()} {x.surname.strip()}"
        if re.sub(r"\s+", "", xml).casefold() != letters.casefold() or len(xml.split()) <= len(written.split()):
            continue
        out, k = [], 0
        for c in xml:
            if c.isspace():
                out.append(" ")
            else:
                out.append(letters[k])
                k += 1
        spaced = re.sub(r" +", " ", "".join(out))
        n = len(" ".join(x.given.split()))
        return spaced[:n].strip(), spaced[n:].strip()
    return None


def _expand_from(short: Author, full_names: list[Author]) -> Author | None:
    """The fuller form of a name with initials ("A. Fachruddin B" -> "Andi Fachruddin Benyamin", "Meera K R" ->
    "Meera K Rao") among names from the submission XML / notice: the same words in the same order, each
    the same or an initial of the fuller one (the XML's given name / surname split is used)."""
    toks = _name_tokens(short)
    if not any(len(t.strip(".")) == 1 for t in toks):
        return None
    for a in full_names:
        if not a.surname or not a.given:
            continue
        ft = _name_tokens(a)
        if len(ft) != len(toks):
            continue
        if all(_token_fits(s, f) for s, f in zip(toks, ft)) and len("".join(ft)) > len("".join(toks)) \
                and _token_fits(toks[0], ft[0]) and not all(len(f.strip(".")) == 1 for f in ft):
            return a
    return None


def initials_first(a: Author) -> bool:
    """Guide: initials written after the name go before it ("Swamy AN" -> "A. N. Swamy", "Suganthi. K" ->
    "K. Suganthi", "Helen Shaji J C" -> "J. C. Helen Shaji"); the name words stay together as the surname.
    Returns whether the name was changed."""
    toks = [t.rstrip(",") for t in _name_tokens(a)]
    k = len(toks)
    while k > 0 and re.fullmatch(r"(?:[A-Z]\.?){1,3}", toks[k - 1]) and not (
            len(toks[k - 1]) > 1 and "." not in toks[k - 1] and toks[k - 1] in ("III", "II", "IV", "JR", "SR")):
        k -= 1
    initials, words = toks[k:], [t.strip(".") for t in toks[:k]]
    if not initials or not words or len(initials) == len(toks):
        return False
    if (a.given + a.surname).isupper() and any(len(t) > 1 and "." not in t for t in initials) \
            or any(re.fullmatch(r"(?:[A-Z]\.?){1,3}", w + ".") for w in words):
        return False  # "MEERA K R": an all-caps name; "A. Fachruddin B": initials on both sides, left as written
    if len(a.surname.split()) != len(initials) and not set(a.surname.split()) <= set(initials):
        return False  # the split put a full word in the surname: not this pattern
    a.given = " ".join(c for t in initials for c in t.replace(".", ""))
    a.surname = " ".join(words)
    return True


def _expand(short: Author, full_names: list[Author]) -> Author | None:
    """The full form of a name written with an initial surname or fewer given names ("Meera R",
    "Meera K R"), from names written in full elsewhere in the sources."""
    g = short.given.split()
    for a in full_names:
        if not a.surname or not a.given:
            continue
        ag = a.given.split()
        if not g or g[0].casefold() != ag[0].casefold():
            continue
        if _norm(a.surname) == _norm(short.surname) and len(a.given) > len(short.given):
            return a
        if _initial_like(short.surname) and a.surname[:1].upper() == short.surname[:1].upper() and len(a.surname) > 1:
            return a
    return None


def cite_name(a: Author) -> str:
    all_caps = (a.given + a.surname).isupper()
    parts = given_parts(a.given, all_caps)
    initials = "".join(t[0].upper() for t, _ in parts)
    if parts and not parts[0][1] and re.fullmatch(r"\w+-[A-Z]\w*", parts[0][0]):  # "Wei-Ming" -> WM
        initials = initials[0] + parts[0][0].split("-")[1][0] + initials[1:]
    surname = " ".join(_name_case(w, False) for w in a.surname.split())
    return f"{surname} {initials}".strip()


_STATES_IN = {s.casefold() for s in (
    "Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh", "Goa", "Gujarat", "Haryana",
    "Himachal Pradesh", "Jharkhand", "Karnataka", "Kerala", "Madhya Pradesh", "Maharashtra", "Manipur",
    "Meghalaya", "Mizoram", "Nagaland", "Odisha", "Orissa", "Punjab", "Rajasthan", "Sikkim", "Tamil Nadu",
    "Tamilnadu", "Telangana", "Tripura", "Uttar Pradesh", "Uttarakhand", "West Bengal", "Delhi", "New Delhi",
    "Puducherry", "Pondicherry", "Chandigarh", "Jammu and Kashmir", "Ladakh")}


# Indian cities whose state the editors add when it is missing ("Pune" -> "Pune, Maharashtra").
_CITY_STATE_IN = {c.casefold(): st for st, cities in {
    "Maharashtra": "Mumbai Pune Nagpur Nashik Aurangabad Thane Navi_Mumbai Kolhapur Solapur Wardha Sangli Latur Pimpri "
                   "Karad Loni Sawangi Ahmednagar Amravati Akola Dhule Jalgaon Nanded Miraj Satara Yavatmal",
    "Karnataka": "Bengaluru Bangalore Mysuru Mysore Mangaluru Mangalore Manipal Hubballi Hubli Belagavi Belgaum Davangere",
    "Tamil Nadu": "Chennai Coimbatore Madurai Vellore Tiruchirappalli Trichy Salem Tirunelveli Puducherry_not Thanjavur",
    "Kerala": "Thiruvananthapuram Trivandrum Kochi Cochin Kozhikode Calicut Thrissur Kottayam Kollam Kannur Ernakulam",
    "Telangana": "Hyderabad Secunderabad Warangal",
    "Andhra Pradesh": "Visakhapatnam Vijayawada Guntur Tirupati Nellore Kurnool Kakinada",
    "West Bengal": "Kolkata Calcutta Siliguri Durgapur Howrah Barrackpore",
    "Uttar Pradesh": "Lucknow Kanpur Varanasi Agra Prayagraj Allahabad Noida Greater_Noida Ghaziabad Meerut Gorakhpur "
                     "Aligarh Bareilly Jhansi Saifai Moradabad",
    "Gujarat": "Ahmedabad Surat Vadodara Rajkot Gandhinagar Bhavnagar Jamnagar",
    "Rajasthan": "Jaipur Jodhpur Udaipur Ajmer Bikaner Kota",
    "Madhya Pradesh": "Bhopal Indore Gwalior Jabalpur Ujjain",
    "Punjab": "Ludhiana Amritsar Jalandhar Patiala Mohali Bathinda",
    "Haryana": "Gurugram Gurgaon Faridabad Rohtak Ambala Panipat Hisar",
    "Bihar": "Patna Gaya Bhagalpur Muzaffarpur Darbhanga",
    "Odisha": "Bhubaneswar Cuttack Rourkela Berhampur Sambalpur Burla",
    "Uttarakhand": "Dehradun Rishikesh Haridwar Haldwani",
    "Assam": "Guwahati Dibrugarh Silchar Jorhat",
    "Jharkhand": "Ranchi Jamshedpur Dhanbad",
    "Chhattisgarh": "Raipur Bilaspur Bhilai",
    "Himachal Pradesh": "Shimla",
    "Jammu and Kashmir": "Srinagar Jammu",
    "Goa": "Panaji",
}.items() for c in cities.replace("_not", "").split() for c in [c.replace("_", " ")]}
_CITY_STATE_IN.pop("puducherry", None)
_COUNTRY_CODES = {"IN": "India", "CN": "China", "US": "USA", "GB": "UK", "SA": "Saudi Arabia", "AE": "UAE",
                  "IR": "Iran", "IQ": "Iraq", "TR": "Turkey", "NP": "Nepal", "BD": "Bangladesh", "PK": "Pakistan",
                  "LK": "Sri Lanka", "EG": "Egypt", "NG": "Nigeria", "MY": "Malaysia", "SG": "Singapore",
                  "JP": "Japan", "KR": "South Korea", "AU": "Australia", "CA": "Canada", "DE": "Germany",
                  "FR": "France", "IT": "Italy", "NL": "Netherlands", "ES": "Spain", "BR": "Brazil", "TW": "Taiwan"}


# Indian states written short: "U.P." / "U.p" in the text (dotted: never a US state), and the submission portal's
# two-letter codes (<state>KA</state>), read only from its XML.
_STATE_DOTTED_IN = {"up": "Uttar Pradesh", "mp": "Madhya Pradesh", "ap": "Andhra Pradesh", "hp": "Himachal Pradesh",
                    "tn": "Tamil Nadu", "wb": "West Bengal", "j&k": "Jammu and Kashmir", "uk": "Uttarakhand"}
_STATE_CODES_IN = {"AP": "Andhra Pradesh", "AR": "Arunachal Pradesh", "AS": "Assam", "BR": "Bihar", "CG": "Chhattisgarh",
                   "CT": "Chhattisgarh", "GA": "Goa", "GJ": "Gujarat", "HR": "Haryana", "HP": "Himachal Pradesh",
                   "JH": "Jharkhand", "KA": "Karnataka", "KL": "Kerala", "MP": "Madhya Pradesh", "MH": "Maharashtra",
                   "MN": "Manipur", "ML": "Meghalaya", "MZ": "Mizoram", "NL": "Nagaland", "OR": "Odisha", "OD": "Odisha",
                   "PB": "Punjab", "RJ": "Rajasthan", "SK": "Sikkim", "TN": "Tamil Nadu", "TG": "Telangana",
                   "TS": "Telangana", "TR": "Tripura", "UP": "Uttar Pradesh", "UK": "Uttarakhand", "UT": "Uttarakhand",
                   "WB": "West Bengal", "DL": "Delhi", "PY": "Puducherry", "CH": "Chandigarh", "JK": "Jammu and Kashmir",
                   "LA": "Ladakh"}


def _state_name(state: str) -> str:
    """"U.P" / "U.p." -> "Uttar Pradesh" (a dotted Indian abbreviation); anything else as written."""
    st = (state or "").strip()
    if "." in st or "&" in st:
        return _STATE_DOTTED_IN.get(re.sub(r"[.\s]", "", st).casefold(), st)
    return st


def _city_key(city: str) -> str:
    """The town the state table knows: "Pimpri, Pune" -> "pune"."""
    return re.split(r"\s*,\s*", (city or "").strip(" ,"))[-1].casefold()


_COUNTRY_SAME = {"türkiye": "turkey", "turkiye": "turkey", "united states": "usa", "united states of america": "usa",
                 "u.s.a.": "usa", "united kingdom": "uk", "u.k.": "uk", "great britain": "uk", "kingdom of saudi arabia":
                 "saudi arabia", "ksa": "saudi arabia", "united arab emirates": "uae", "republic of korea": "south korea",
                 "korea": "south korea", "the netherlands": "netherlands", "holland": "netherlands"}


def _same_country(a: str, b: str) -> bool:
    """Two spellings of one country ("Turkey" / "Türkiye", "USA" / "United States")."""
    ka, kb = (a or "").strip(" .").casefold(), (b or "").strip(" .").casefold()
    return bool(ka) and bool(kb) and _COUNTRY_SAME.get(ka, ka) == _COUNTRY_SAME.get(kb, kb)


def _country(parts: dict) -> str:
    c = (parts.get("country") or "").strip()
    if c:
        return _COUNTRY_CODES.get(c, c)
    if _state_name(parts.get("state", "")).casefold() in _STATES_IN or _city_key(parts.get("city", "")) in _STATES_IN \
            or (not parts.get("state") and _city_key(parts.get("city", "")) in _CITY_STATE_IN):
        return "India"
    return ""


def _state(parts: dict) -> str:
    """The state as written, or the one an Indian city is in."""
    if parts.get("state"):
        return _state_name(parts["state"])
    if _country(parts).casefold() == "india":
        return _CITY_STATE_IN.get(_city_key(parts.get("city", "")), "")
    return ""


# A home address in the correspondence ("F 1106, Asawari", "F6 Staff Quarters", "A601, Watergreen Society"): no
# institution is put in front of it.
_RESIDENCE = re.compile(r"(?i)\b(?:flat|apartments?|apts?|society|residency|residence|quarters|house|tower|floor|villa|"
                        r"bungalow|plot|enclave|heights|[A-Z]\s?-?\d{2,4})\b")
_STREET = re.compile(r"(?i)^\s*(?:no\.?\s*)?\d[\w/-]*\b|\b(?:office|room|building|bldg|floor|campus|street|road|"
                     r"avenue|lane|p\.?\s*o\.?\s*box)\b")
_UNIT_WORDS = re.compile(r"(?i)\b(department|departments|division|unit|units|service|services|centre|center|institute|"
                         r"school|college|laborator(?:y|ies)|lab|clinic|section|faculty|program(?:me)?|hospital|"
                         r"university|foundation|cell|wing|ward|station)\b")


_DEPT_LEAD = re.compile(r"(?:deptt|dept|department)\b\.?(?:\s+of\b|\s*(?=[A-Z]))\s*", re.I)


def department_form(d: str) -> str:
    """"Cornea Department" and a bare specialty ("ophthalmology") -> "Department of ..."."""
    d = d.strip()
    lead = _DEPT_LEAD.match(d)
    if lead and d[lead.end():].strip():  # "Dept. of Psychiatry", "Department Of Pathology" -> "Department of ..."
        rest = d[lead.end():].strip()
        while _DEPT_LEAD.match(rest) and rest[_DEPT_LEAD.match(rest).end():].strip():
            rest = rest[_DEPT_LEAD.match(rest).end():].strip()  # "Department of Department of Pathology"
        return "Department of " + rest
    m = re.fullmatch(r"(?i)(.+?)\s+department", d)
    if m and not _UNIT_WORDS.search(m.group(1)):
        return "Department of " + m.group(1)
    if d and not _UNIT_WORDS.search(d) and not re.search(r"(?i)\bof\b", d) and len(d.split()) <= 5 and re.search(r"(?i)(?:logy|ics|ery|ine|ry|thesia|"
                                                                             r"thesiology|tics|ology)\b", d):
        return "Department of " + d
    return d


def _place_case(text: str) -> str:
    """"WUHAN" -> "Wuhan", "CHINA" -> "China"; short codes (USA, UK, WI) stay."""
    return " ".join(w.capitalize() if w.isupper() and len(w) > 3 else w for w in text.split())


# Function words of other languages inside a name ("Başakşehir Çam ve Sakura Şehir Hastanesi", "Universidade de
# São Paulo", "Hospital Universitario y Politécnico La Fe", "Klinikum rechts der Isar", "Karolinska Institutet för
# ..."): they stay lowercase when a name is recased.
_FOREIGN_MINOR = {"ve", "de", "del", "della", "delle", "dei", "degli", "da", "do", "dos", "das", "di", "du", "des",
                  "la", "le", "les", "y", "e", "und", "van", "von", "der", "den", "et", "för", "och", "og", "es", "el"}
# The conjunctions among them, written with a capital inside a name ("Çam Ve Sakura", "Universidade De São Paulo"):
# lowered between two capitalised words. Single letters (initials: "Dr Y S R"), articles and words that begin
# place and family names ("La Paz", "Van", "Der", "Des Moines", "Di Venere", "Da Nang") are left as written.
_FOREIGN_LOWER = {"ve", "und", "och", "og", "et", "för", "de", "del", "do", "dos", "das", "du"}


def _cap_words(text: str) -> str:
    """Capitalise lowercase major words: "city eye hospital" -> "City Eye Hospital"; a foreign connector inside
    a name stays (or becomes) lowercase: "Çam Ve Sakura" -> "Çam ve Sakura"."""
    words = re.split(r"(\s+)", re.sub(r"\s+", " ", text.strip()))
    out = []
    for i, w in enumerate(words):
        if i and w[:1].isupper() and w[1:].islower() and w.casefold() in _FOREIGN_LOWER \
                and i + 2 < len(words) and words[i + 2][:1].isupper() and out[-2][:1].isupper() \
                and out[-2].casefold() not in _MINOR:
            out.append(w.lower())
        elif w.isspace() or not w[:1].islower() or (i and (w in _MINOR or w in _PARTICLES or w in _FOREIGN_MINOR)):
            out.append(w)
        else:
            out.append(w[:1].upper() + w[1:])
    return "".join(out)


_DESIGNATIONS = ("professor and head", "professor & head", "head of the department", "head of department",
                 "senior consultant", "consultant", "senior resident", "junior resident", "chief resident", "resident",
                 "associate professor", "assistant professor", "additional professor", "professor", "head", "hod",
                 "senior lecturer", "lecturer", "dean", "director", "medical superintendent", "superintendent",
                 "medical officer", "senior registrar", "registrar", "research fellow", "clinical fellow", "fellow",
                 "postgraduate student", "postgraduate resident", "postgraduate", "pg student", "pg resident",
                 "intern", "research scholar", "scientist-[a-g]", "scientist")
_DESIGNATION = re.compile(r"^(?:(?:(?:retired|former|ex-?|emeritus)\s+)?(?:%s)\b\.?(?:\s*[,;/]\s*|\s+(?=(?:of\s+(?:the\s+)?)?department\b)|\s+of\s+(?=(?:the\s+)?department\b)))+"
                          % "|".join(_DESIGNATIONS), re.I)
# A title directly before a speciality without a comma: "Consultant cornea and general ophthalmology".
_TITLE_WORD = re.compile(r"^(?:senior\s+|junior\s+)?(?:consultant|resident|lecturer)\s+(?!of\b)(?=\w)", re.I)


def strip_designation(department: str) -> str:
    """"Senior Resident, Department of X" -> "Department of X"; "Consultant cornea" -> "cornea"."""
    d = department.strip()
    for pattern in (_DESIGNATION, _TITLE_WORD):
        m = pattern.match(d)
        if m and m.end() < len(d):
            d = d[m.end():].lstrip(" ,;")
    d = re.sub(r"(?i)^of\s+(?=(?:the\s+)?department\b)", "", d)
    return d


_DESIGNATION_ONLY = re.compile(r"(?i)^(?:(?:retired|former|ex-?|emeritus)\s+)?(?:%s)\.?(?:\s*(?:and|&)\s*(?:%s))?\.?$"
                               % ("|".join(_DESIGNATIONS), "|".join(_DESIGNATIONS)))


def drop_designations(text: str) -> str:
    """G69 "Designation and qualification of authors should be deleted", also inside a run-on affiliation or
    address ("..., Zahedan, Iran, Associate Professor, Community Nursing Research Center, ..."): a comma-separated
    piece that is only a designation goes, one that starts with it loses it."""
    pieces = [x for x in re.split(r"\s*,\s*", text.strip())]
    if len(pieces) < 2 and not _DESIGNATION.match(text.strip()):
        return text
    kept = [strip_designation(x) for x in pieces if not _DESIGNATION_ONLY.match(x.strip())]
    return ", ".join(x for x in kept if x)


def _clean_part(text: str) -> str:
    t = drop_designations(re.sub(r"\s+", " ", text))
    t = t.strip(" ,.;")
    t = re.sub(r"\s*&\s*", " and ", t)
    t = re.sub(r"(?i)^dept\.?(?:\s+of)?\s+", "Department of ", t)
    t = re.sub(r"\s+,", ",", t)
    t = re.sub(r",(?=\S)", ", ", t)
    return _cap_words(t)


# --- building the block --------------------------------------------------------------------

Run = tuple[str, dict]


# A postal code as a piece of an address ("Fez, 30000, Morocco", "Hamadan, 6517838695", "B-1200") or after / before
# the city in one piece ("Chennai - 600 006", "Rabat 10100", "75006 Paris"); never a number after a street word
# ("PO Box 1234", "Sector 12", "Building 2").
_POSTAL_PIECE = re.compile(r"^(?:(?:pin|zip|postal\s+code|post\s+code)\s*(?:code)?\s*[:.\-–]?\s*)?(?:[A-Z]{1,2}\s?-?\s?)?"
                           r"(?:\d{3}\s?\d{3}|\d{4,10}|\d{3}-\d{4}|\d{5}-\d{4}|\d{2}-\d{3})$", re.I)
_POSTAL_AFTER = re.compile(r"(?<=[A-Za-z.)])\s*[-–,]?\s*(?:pin\s*[:\-–]?\s*)?(\d{3}\s?\d{3}|\d{5}-\d{4}|\d{3}-\d{4}|\d{4,10})$", re.I)
_POSTAL_BEFORE = re.compile(r"^\d{4,6}\s+(?=[A-Z][a-z])")
_STREET_NUMBER = re.compile(r"(?i)\b(?:box|no\.?|number|street|st\.?|road|rd\.?|lane|plot|sector|block|building|bldg|floor|"
                            r"room|office|flat|house|suite|km|phase|ward|gate|avenue|ave\.?)\s*[-–:]?\s*\d[\d\s\-–]*$")


def drop_postal_codes(text: str) -> str:
    """S "delete pin/zip code" (IJO, JEHP; the editors delete it in the correspondence too): the postal codes in an
    affiliation or address, written as a piece of their own or next to the city."""
    out = []
    for piece in re.split(r"\s*,\s*", text.strip()):
        if _POSTAL_PIECE.match(piece) and not re.fullmatch(r"\d{1,3}", piece.strip()):
            continue
        if not _STREET_NUMBER.search(piece):
            piece = _POSTAL_AFTER.sub("", piece)
            piece = _POSTAL_BEFORE.sub("", piece)
        if piece.strip(" -–"):
            out.append(piece.strip())
    return ", ".join(out)


def _affiliation_body(a: dict, parts: list[str], postal: bool = False, drop_postal: bool = False) -> str:
    """An affiliation as printed, without the country: `parts` (the department, or nothing when it is printed
    with its group) + institution, city and state, cleaned, a place the institution already ends with dropped.
    `drop_postal`: postal codes written into the parts go (S "delete pin/zip code")."""
    parts = parts + [a.get("institution", ""), _place_case(a.get("city", "")), _place_case(_state(a))]
    parts = [_clean_part(p) for p in parts if p and p.strip(" ,.")]
    if drop_postal:
        parts = [q for q in (drop_postal_codes(p) for p in parts) if q]
    # drop a city or state the institution already names at its end ("AIIMS, New Delhi" + "New Delhi"); one
    # that is part of its name stays ("University of Phayao, Phayao", "Mohammed V University in Rabat, Rabat")
    dedup = []
    for p in parts:
        if not any(_norm(q).endswith(_norm(p))
                   and not re.search(r"(?i)\b(?:of|in|at|de|du|di)\s+" + re.escape(p) + "$", q) for q in dedup):
            dedup.append(p)
    city = _clean_part(_place_case(a.get("city", "")))
    pin = re.sub(r"\s+", "", a.get("postal_code", "")) if postal else ""
    at = next((k for k in range(len(dedup) - 1, -1, -1) if city and _norm(dedup[k]).endswith(_norm(city))), None)
    if pin and at is not None:  # after the city, as the correspondence prints it: "Chennai - 600 006" (India),
        # else "Tokyo, 102-8554"
        india = _country(a).casefold() == "india" and re.fullmatch(r"\d{6}", pin)
        dedup[at] += f" - {pin[:3]} {pin[3:]}" if india else f", {pin}"
    return ", ".join(dedup)


def affiliation_text(f: "Fields", i: int, postal: bool = False, drop_postal: bool = False) -> str:
    """Affiliation `i` of `f` printed complete, as the author group's affiliation line prints it on its own
    (department, institution, city, state, country). A missing country is the one the other affiliations all
    give, when the affiliation names no state that could say otherwise."""
    a = f.affiliations.get(i)
    if not a:
        return ""
    text = _affiliation_body(a, [department_form(a.get("department", ""))], postal, drop_postal)
    country = _country(a)
    if not country and not a.get("state"):
        others = {_country(b) for j, b in f.affiliations.items() if j != i} - {""}
        country = others.pop() if len(others) == 1 else ""
    if country:
        text += (", " if text else "") + _place_case(_clean_part(country))
    return text


@dataclass
class Block:
    role: str
    runs: list[Run]
    end: bool = False  # letters: authors to cite line go at the end, before the references
    provenance: dict | None = None  # which input the block was built from (set by `plan`)

    @property
    def text(self) -> str:
        return "".join(t for t, _ in self.runs)


def rules_text_on(style: JournalStyle, name: str) -> bool:
    from . import rules_text

    return rules_text.enabled(rules_text.settings(style), name)


def _settings(style: JournalStyle) -> dict:
    cfg = dict(style.get("front_matter", {}) or {})
    over = (cfg.pop("journals", {}) or {}).get(style.code.lower(), {})
    return _merge(cfg, over)


def _query_on(cfg: dict, name: str) -> bool:
    return bool((cfg.get("queries") or {}).get(name))


# A unit of the institution in an affiliation's department place: "Department of ...", "Division of Allergy",
# "Endovascular Intervention Unit", "Postdoctoral Station of ...", a research centre. A service ("Oculoplasty
# Services"), an institute ("Anant Bajaj Retina Institute"), a programme or a hospital there is no department.
_SUBUNIT = re.compile(r"(?i)\b(?:departments?|dept|division|units?|section|station|laborator(?:y|ies)|lab|clinic|cell|"
                      r"wing|ward|cent(?:er|re))\b")


def _no_department(a: dict) -> bool:
    """The affiliation names no department or unit of its institution ("Gokhale Eye Hospital, Mumbai"; the IJO
    editors ask "Kindly provide department" there: 13 of 13 articles)."""
    return not _SUBUNIT.search(department_form(a.get("department", "")))


def _name_tokens_of(text: str) -> list[str]:
    """Lower-case letters of a name, split at anything else ("Oktia W.K. Handayani" -> oktia, w, k, handayani)."""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c)).casefold()
    return [t for t in re.split(r"[^a-z]+", text) if t]


def xml_fuller_name(written: str, xml_authors: list[dict]) -> bool:
    """The submission XML gives in full a name the title page shortens: an initial the XML spells out ("Dewi S.
    Rochmayani" / "Dewi Sari Rochmayani", "Natarajan S" / "natarajan shanmugasundaram") or a single name to which it
    adds another ("Sonakshi" / "Sonakshi Rajput"). Not a name the XML only repeats ("Hartono Hartono"), shortens
    or spells otherwise (another person, or a spelling for the author to settle elsewhere)."""
    w = _name_tokens_of(written)
    full, inits = {t for t in w if len(t) > 1}, {t for t in w if len(t) == 1}
    if not full:
        return False
    best, score = None, 0
    for x in xml_authors:
        sur, giv = set(_name_tokens_of(x.get("surname", ""))), set(_name_tokens_of(x.get("given", "")))
        n = 2 * len(full & sur) + len(full & giv)
        if n > score:
            best, score = x, n
    if best is None:
        return False
    xt = set(_name_tokens_of(f"{best.get('given', '')} {best.get('surname', '')}"))
    if not full <= xt:
        return False
    extra = [t for t in xt - full if len(t) > 2]
    return any(t[0] in inits for t in extra) or (len(full) == 1 and not inits and bool(extra))


def _imrad(ms: Manuscript) -> bool:
    heads = {re.sub(r"^<H\d>|[:.]$", "", p.text.strip()).strip().lower() for p in ms.paragraphs if len(p.text) < 60}
    need = [{"introduction"}, {"materials and methods", "methods", "patients and methods", "subjects and methods",
                               "methodology"}, {"results"}, {"discussion"}, {"conclusion", "conclusions"}]
    return all(heads & n for n in need)


def _bare_specialty(d: str) -> bool:
    d = d.strip()
    return bool(d) and not _UNIT_WORDS.search(d) and not re.search(r"(?i)\bof\b", d) and len(d.split()) <= 5 \
        and not re.search(r"\d|,", d)


def _department_of_siblings(affs: dict[int, dict], order: list[int]) -> None:
    """A department named by its specialty alone gets "Department of" like the others of the block ("Professor
    (Anatomy)" -> "Anatomy" beside "Department of Ophthalmology" / "Department of Pathology": IJO_1745), so the
    affiliations read alike (G68 "Department name ... before institution")."""
    bare = [i for i in order if _bare_specialty(affs[i].get("department", ""))]
    if not bare or not any(department_form(affs[i].get("department", "")).startswith("Department of ") for i in order):
        return
    for i in bare:
        if not department_form(affs[i]["department"]).startswith("Department of "):
            affs[i]["department"] = "Department of " + affs[i]["department"].strip()


def _aff_words(a: dict) -> list[str]:
    text = " ".join([department_form(a.get("department", ""))] + [a.get(k, "") for k in ("institution", "city", "state",
                                                                                         "country")])
    return re.findall(r"\w+", unicodedata.normalize("NFKC", text).casefold())


def _near_duplicates(f: "Fields", order: list[int], key: dict[int, tuple]) -> None:
    """S "In case of slight changes in repeating affiliations, raise query internally": two affiliations that differ
    only in spelling ("Obstetrics and Gynaecology" / "Obstetrics and Gynecology", jehp_924) are one affiliation (the
    spelling most authors are linked to; the editors group them) and an editor warning says so; two that differ in a
    word or two ("Psychiatric University Center, Agadir" / "Psychiatric University Center, Ibn Zohr University,
    Agadir") are left as written with the warning."""
    uses = {i: sum(i in a.affs for a in f.authors) for i in order}
    words = {i: _aff_words(f.affiliations[i]) for i in order}
    for n, i in enumerate(order):
        for j in order[:n]:
            if key[i] == key[j]:
                break
            wi, wj = words[i], words[j]
            if len(wi) == len(wj) and wi != wj and all(
                    x == y or len(x) >= 5 and difflib.SequenceMatcher(None, x, y).ratio() >= 0.85 for x, y in zip(wi, wj)):
                keep, drop = (i, j) if uses[i] > uses[j] else (j, i)
                f.affiliations[drop] = dict(f.affiliations[keep])
                for k in order:
                    if key[k] == key[drop]:
                        key[k] = key[keep]
                words[drop] = words[keep]
                f.warnings.append(f"affiliations differing only in spelling grouped as one (check): "
                                  f"{affiliation_text(f, keep)!r}")
                break
            small, big = (set(wi), set(wj)) if len(set(wi)) <= len(set(wj)) else (set(wj), set(wi))
            if len(small) >= 4 and small < big and len(big - small) <= 3 and f.affiliations[i].get("city", "").casefold() \
                    == f.affiliations[j].get("city", "").casefold():
                f.warnings.append("slight changes in a repeated affiliation (check): "
                                  f"{affiliation_text(f, j)!r} / {affiliation_text(f, i)!r}")


def build(ms: Manuscript, f: Fields, meta: Meta, style: JournalStyle) -> tuple[list[Block], list[tuple[str, dict]]]:
    """The front-matter paragraphs in order, and the author queries (check, format args)."""
    cfg = _settings(style)
    blocks: list[Block] = []
    qs: list[tuple[str, dict]] = []
    B, I, SUP = {"bold": True}, {"italic": True}, {"superscript": True}
    plain: dict = {}

    # Manuscript type: the XML first (editors follow it), then the text, then IMRaD.
    first, second = (f.article_type, meta.article_type) if cfg.get("type_source") == "text" else (meta.article_type, f.article_type)
    atype = first or second
    if atype and cfg.get("type_qualifier") == "keep" and atype == meta.article_type:
        atype = re.sub(r"\s+", " ", atype).strip(" .:")  # "Case Report (for IJO Case Reports)" as the XML has it
    elif atype:
        atype = clean_article_type(atype)
    elif _imrad(ms):
        atype = "Original Article"
        qs.append(("front.confirm_type", {"_block": "article_type"}))
    else:
        qs.append(("front.article_type", {"_block": "article_type"}))
    if atype and _query_on(cfg, "type_source") and (f.article_type or meta.article_type):
        # "Kindly note that the “Article type” provided in FP file has been followed.": the title page's when it
        # names one, else the XML's
        fp = _norm(" ".join(getattr(getattr(ms, "front_sources", None), "first_page", None) or []))
        given = clean_article_type(f.article_type) if f.article_type else ""
        source = "FP" if given and _norm(given) in fp else "XML" if meta.article_type else ""
        if source:
            qs.append(("front.type_source", {"source": source, "_block": "article_type"}))
    letter = bool(re.search(r"letter", atype or "", re.I))

    # Title.
    raw_title = clean_title(f.title or meta.title)
    keep = proper_keep([raw_title, f.running_title or meta.running_title], f.proper_nouns, case_body(ms), f.notes)
    colon = cfg.get("capital_after_colon", True)
    title = recase(raw_title, cfg.get("title_case", "sentence"), keep, colon) if raw_title else ""
    if title and ms is not None and rules_text_on(style, "compound_dash"):
        # "oral health–related" where the text writes "health-related": the hyphen (JEHP_73_26, QC)
        from . import rules_text

        ctx = rules_text.Context.from_texts(p.text for p in ms.paragraphs if p.text.strip())
        for s_, e_, new in reversed(rules_text.find_compound_dash(title, {}, ctx)):
            title = title[:s_] + new + title[e_:]

    # Authors.
    names = []
    for a in f.authors:
        name, full = format_name(a, cfg)
        names.append(name)
        if a.written and _query_on(cfg, "xml_name") and xml_fuller_name(a.written, meta.authors) \
                and set(_name_tokens_of(name)) == set(_name_tokens_of(a.written)):
            # printed as the title page writes it, the XML's fuller form not taken: the editors note the XML's
            # (JEHP: "Kindly note the Author name provided in XML as “Sonakshi”. Please check and confirm.")
            qs.append(("front.xml_name", {"name": a.written, "_block": "authors", "_quote": name}))
        elif not full and (cfg.get("queries") or {}).get("full_name", True):  # on the author's name, naming them
            qs.append(("front.full_name", {"name": name, "_block": "authors", "_quote": name}))

    if names and _query_on(cfg, "authors_check"):
        # the JEHP editors' standing query on the author group (about 55 of 302 articles; owner, 2026-10-30)
        qs.append(("front.authors_check", {"_block": "authors"}))

    if f.xml_only:
        qs.append(("front.xml_only_author", {"names": ", ".join(f.xml_only)}))

    # Running head.
    surnames = [" ".join(_name_case(w, False) for w in a.surname.split()) for a in f.authors]
    short = f.running_title or meta.running_title
    short = re.sub(r"^\s*(?:running|short)\s*(?:title|head)\s*[:\-–]?\s*", "", short, flags=re.I)
    lead = re.match(r"^\s*([^:]{1,60}?)\s*:\s*(?=\S)", short)
    names_part = re.sub(r"(?:,?\s*et\.?\s*al\.?|\s+(?:and|&)\s+\S+)$", "", lead.group(1), flags=re.I) if lead else ""
    if lead and surnames and _norm(names_part) and _norm(names_part) == _norm(surnames[0]):
        short = short[lead.end():]  # "Das, et al: OSA risk ..." given as the running title: the names are added below
    short = _drop_end_period(recase(re.sub(r"\s+", " ", short).strip(), "sentence", keep)) if short.strip() else ""
    rh: list[Run] = [("<RH>", plain)]
    if not surnames:
        rh.append((MISSING, plain))
    elif len(surnames) == 1:
        rh.append((surnames[0], plain))
    elif len(surnames) == 2:
        rh.append((f"{surnames[0]} and {surnames[1]}", plain))
    else:
        rh += [(f"{surnames[0]}, ", plain), ("et al.", I)]
    rh.append((": " + (short or MISSING), plain))
    if not short and not (letter and cfg.get("ltr_no_running_head_query")):
        qs.append(("front.running_head", {"_block": "running_head"}))
    elif short and cfg.get("running_head_max") and len(short) > cfg["running_head_max"]:
        # S (IJO, JEHP): the running title is "Not more than 50 characters"
        qs.append(("front.running_head_long", {"n": cfg["running_head_max"], "_block": "running_head", "_quote": short}))
    blocks.append(Block("running_head", rh))
    if atype:
        blocks.append(Block("article_type", [(atype, B)]))
    if title:
        blocks.append(Block("title", [(title, B)]))

    # Affiliations: unique, in order of use; numbered unless there is only one; grouped by institution.
    affs = f.affiliations
    order: list[int] = []
    for a in f.authors:
        for i in a.affs:
            if i in affs and i not in order:
                order.append(i)
    if not order and len(affs) == 1:
        order = list(affs)
        for a in f.authors:
            a.affs = order[:]
    _complete_affiliations(affs, order, f.notes)
    _department_of_siblings(affs, order)
    # duplicates are judged on the printed form ("Cardiology Department" = "Department of Cardiology")
    key = {i: (_norm(department_form(affs[i].get("department", ""))),)
           + tuple(_norm(affs[i].get(k, "")) for k in ("institution", "city", "state", "country")) for i in order}
    _near_duplicates(f, order, key)
    uniq: list[int] = []
    same: dict[int, int] = {}
    for i in order:
        hit = next((u for u in uniq if key[u] == key[i]), None)
        same[i] = hit if hit is not None else i
        if hit is None:
            uniq.append(i)
    one_each = all(len({same[i] for i in a.affs if i in same}) <= 1 for a in f.authors)
    number: dict[int, str] = {}
    if len(uniq) > 1:
        base = 0 if one_each and f.authors and f.authors[0].affs and cfg.get("affiliation_first_unnumbered", True) else 1
        number = {u: (str(n) if n else "") for n, u in enumerate(uniq, base)}

    if names:
        runs: list[Run] = []
        for k, (a, name) in enumerate(zip(f.authors, names)):
            if k:
                runs.append((", ", B))
            runs.append((name, B))
            marks = [number.get(same[i], "") for i in a.affs if i in same]
            marks = ",".join(sorted(dict.fromkeys(m for m in marks if m), key=int))
            if k in f.equal_authors:
                marks += f.equal_mark
            if marks:
                runs.append((marks, {"bold": True, "superscript": True}))
        blocks.append(Block("authors", runs))

    if uniq:
        runs = []
        groups: list[list[int]] = []
        for u in uniq:  # departments of one institution and place go together
            d = department_form(affs[u].get("department", ""))
            loc = key[u][1:]
            g = next((g for g in groups if cfg.get("group_departments", True) and key[g[0]][1:] == loc and loc[0]
                      and d.lower().startswith("department of ")
                      and department_form(affs[g[0]].get("department", "")).lower().startswith("department of ")),
                     None)
            if g is not None:
                g.append(u)
            else:
                groups.append([u])
        countries = [_country(affs[g[0]]) for g in groups]
        no_dept: list[str] = []  # affiliations naming no department (front_matter.queries.department)
        for k, c in enumerate(countries):  # "Turkey" and "Türkiye" in one block: the first spelling throughout
            first = next((x for x in countries if _same_country(x, c)), c)
            countries[k] = first
        def may_share(k: int, c: str) -> bool:  # an affiliation without a country, printed with its neighbour's
            st = affs[groups[k][0]].get("state", "")
            return bool(c) and (not st or c.casefold() == "india" and st.casefold() in _STATES_IN)
        for k in list(range(1, len(countries))) + list(range(len(countries) - 2, -1, -1)):
            if not countries[k]:
                if k and may_share(k, countries[k - 1]):
                    countries[k] = countries[k - 1]
                elif k + 1 < len(countries) and may_share(k, countries[k + 1]):
                    countries[k] = countries[k + 1]
        for gi, g in enumerate(groups):
            if gi:
                runs.append((", ", plain))
            a0 = affs[g[0]]
            if len(g) > 1:
                runs.append(("Departments of ", plain))
                names = [_clean_part(re.sub(r"(?i)^department of ", "", department_form(affs[u]["department"])))
                         for u in g]
                last_join = ", " if any(" and " in n for n in names) else " and "
                for k, (u, name) in enumerate(zip(g, names)):
                    if k:
                        runs.append((last_join if k == len(g) - 1 else ", ", plain))
                    if number.get(u):
                        runs.append((number[u], SUP))
                    runs.append((name, plain))
                parts = []
            else:
                if number.get(g[0]):
                    runs.append((number[g[0]], SUP))
                parts = [department_form(a0.get("department", ""))]
            text = _affiliation_body(a0, parts, bool(cfg.get("affiliation_postal_code")),
                                     bool(cfg.get("delete_postal_codes")))
            if len(g) == 1 and text and _query_on(cfg, "department") and _no_department(a0):
                no_dept.append(text)
            if len(g) > 1 and text:
                text = ", " + text
            last_of_country = gi == len(groups) - 1 or countries[gi + 1] != countries[gi] \
                or cfg.get("affiliation_country_each", False)  # JEHP: every affiliation complete
            if countries[gi] and last_of_country:
                text += (", " if text else "") + _place_case(_clean_part(countries[gi]))
            elif not countries[gi] and last_of_country:
                qs.append(("front.country", {"_block": "affiliations"}))
            runs.append((text, plain))
        blocks.append(Block("affiliations", runs))
        if no_dept:  # one query on the line, on the first such affiliation ("Kindly provide the departments.": 1130)
            qs.append(("front.department" if len(no_dept) == 1 else "front.departments",
                       {"_block": "affiliations", "_quote": no_dept[0]}))
    if f.equal_note:
        blocks.append(Block("equal_contribution", [(f.equal_mark, {"bold": True, "superscript": True}),
                                                   (f.equal_note, B)]))

    # Correspondence.
    c = f.correspondence
    corr_author = next((a for a in f.authors if a.corresponding), None)
    if c or corr_author:
        given, surname = c.get("given_names", ""), c.get("surname", "")
        if not surname and corr_author:
            given, surname = corr_author.given, corr_author.surname
        full = _expand(Author(given, surname), f.authors) if surname else None
        if full is not None:  # "Author name(s) should be expanded"
            given, surname = full.given, full.surname
        elif surname:  # "Suganthi K" -> "K. Suganthi", as in the author group
            a = Author(given, surname)
            if initials_first(a):
                given, surname = a.given, a.surname
        name = format_name(Author(given, surname), cfg, abbreviate=False)[0] if surname else MISSING
        allowed = [x.strip(" .").lower() for x in cfg.get("correspondence_prefixes", ["Dr", "Prof"])]
        prefix = c.get("prefix") or (corr_author.prefix if corr_author else "")
        if prefix.strip(" .").lower() not in allowed:
            prefix = ""
        if not prefix and surname:
            # G91-92: "prefix like "Dr." ... can be present. Please retain/add period": the portal's <title>, for the
            # submission's author with this name (in either order: "Nguyen Duc Khanh" / given "Duc Khanh") or e-mail
            names = _norm(f"{given}{surname}")
            email = (c.get("email") or "").strip().casefold()
            def same(a: dict) -> bool:
                whole = (_norm(a.get("given", "") + a.get("surname", "")), _norm(a.get("surname", "") + a.get("given", "")))
                return bool(names) and names in whole or bool(email) and a.get("email", "").strip().casefold() == email

            xm = next((a for a in meta.authors if same(a)), None) or \
                next((a for a in meta.authors if _norm(a.get("surname", "")) == _norm(surname)), None)
            if xm and xm.get("prefix", "").strip(" .").lower() in allowed:
                prefix = xm["prefix"]
        prefix = {"dr": "Dr. ", "prof": "Prof. ", "mrs": "Mrs. ", "mr": "Mr. ", "ms": "Ms. "}.get(
            prefix.strip(" .").lower(), "") if prefix.strip(" .").lower() in allowed else ""
        blocks.append(Block("correspondence", [(cfg.get("correspondence_label", "Address for correspondence:"), B),
                                               (f" {prefix}{name},", plain)]))
        if surname and not prefix and _query_on(cfg, "designator"):
            # no Dr. / Prof. (nor another prefix the journal keeps) for the corresponding author (JEHP)
            qs.append(("front.designator", {"_block": "correspondence", "_quote": name}))
        parts = [_clean_part(p) for p in c.get("address_parts", [])]
        if cfg.get("correspondence_drop_street") and any(_UNIT_WORDS.search(p) for p in parts):
            # the institution's street, building and office ("Office 2110", "Building 2 (Bu Hasa)", "SAN Campus",
            # "215 Hong Bang Street", "4-16-11, Shimoochiai Shinjuku") go when the address names the institution
            parts = [p for p in parts if _UNIT_WORDS.search(p) or not _STREET.search(p)]
        city, pin = _place_case(_clean_part(c.get("city", ""))), re.sub(r"\s+", "", c.get("postal_code", ""))
        # the correspondence keeps its postal code where the journal's editors keep it (correspondence_postal_code)
        drop_postal = bool(cfg.get("delete_postal_codes")) and not cfg.get("correspondence_postal_code")
        pin_at_country = cfg.get("correspondence_postal_position") == "country"
        if drop_postal:  # S "delete pin/zip code": the editors delete it in the correspondence too
            parts = [q for q in (drop_postal_codes(p) for p in parts) if q]
            city, pin = drop_postal_codes(city), ""
        xm = next((a for a in meta.authors if surname and _norm(a.get("surname", "")) == _norm(surname)), None)
        if not pin and not drop_postal and xm and re.fullmatch(r"\d{6}", re.sub(r"\s+", "", xm.get("zip", ""))):
            pin = re.sub(r"\s+", "", xm["zip"])
        country = _country({"country": c.get("country", ""), "state": c.get("state", ""), "city": c.get("city", "")})
        if corr_author is None and surname:
            corr_author = next((a for a in f.authors if _norm(a.surname) == _norm(surname)), None)
        a = affs.get(corr_author.affs[0]) if corr_author is not None and corr_author.affs else None
        if not parts and a:  # S: "In case correspondence content is not present separately but the same is present
            # in affiliation then it should be used"
            parts = [_clean_part(department_form(a[k]) if k == "department" else a[k])
                     for k in ("department", "institution") if a.get(k)]
            city = city or _clean_part(a.get("city", ""))
            c = dict(c, state=c.get("state") or a.get("state", ""))
            country = country or _country(a)
        elif parts and a:
            same_city = bool(city) and _city_key(city) == _city_key(a.get("city", ""))
            if cfg.get("correspondence_add_institution") and a.get("institution") and same_city \
                    and not any(_UNIT_WORDS.search(p) or _RESIDENCE.search(p) for p in parts):
                parts = [_clean_part(a["institution"])] + parts  # a street address only: the institution first
            elif a.get("department") and _norm(parts[0]) == _norm(a.get("institution", "")):
                parts = [_clean_part(department_form(a["department"]))] + parts  # the institution alone: its department
            if not country and not c.get("state") and (same_city or not city):
                # no state or country ("Tevfik Saglam St. Number:1", "Sankara Eye Hospital, Shimoga"): the
                # affiliation's (S: "Country name must be added at the end if it is not present")
                city = city or _place_case(_clean_part(a.get("city", "")))
                c = dict(c, state=_state(a), city=c.get("city") or a.get("city", ""))
                country = _country(a)
        if pin and any(pin in re.sub(r"\s+", "", p) for p in parts):
            pin = ""  # written into the address already ("Jl. Ir. Sutami No.36A, Surakarta 57126")
        has_pin = bool(pin) or any(drop_postal_codes(p) != p for p in parts)
        if city:
            if pin and pin_at_country:
                pass
            elif pin and country.casefold() == "india" and re.fullmatch(r"\d{6}", pin):
                city += " - " + (f"{pin[:3]} {pin[3:]}" if cfg.get("pin_spaced", True) else pin)
            elif pin:
                city += ", " + pin
            parts.append(city)
        state = _state({"state": c.get("state", ""), "city": c.get("city", ""), "country": country})
        if state:
            parts.append(_place_case(_clean_part(state)))
        if pin and pin_at_country:
            parts.append(pin)  # "Semarang, Central Java, 50229, Indonesia" (the JEHP editors)
        if country:
            country = _COUNTRY_CODES.get(country, country)
            # the affiliations' spelling ("Turkey" there, "Türkiye" in the address: IJO_1347)
            country = next((_country(x) for x in affs.values() if _same_country(_country(x), country)), country)
            parts.append(_place_case(_clean_part(country)))
        if drop_postal:
            parts = [drop_postal_codes(p) for p in parts]
        dedup = []
        for p in parts:
            if p and not any(_norm(p) == _norm(q) for q in dedup):
                dedup.append(p)
        if dedup and not drop_postal and not has_pin and _query_on(cfg, "postal_code"):
            # on the address's last part, the country, where the JEHP editors anchor it
            qs.append(("front.postal_code", {"_block": "correspondence", "_quote": dedup[-1]}))
        if dedup and cfg.get("correspondence_one_line", False):  # "Address for correspondence: Dr. A, Dept, ..."
            blocks[-1].runs.append((" " + ", ".join(dedup) + ".", plain))
        elif dedup:
            blocks.append(Block("correspondence", [(", ".join(dedup) + ".", plain)]))
        email = c.get("email") or next((a.get("email", "") for a in meta.authors if a.get("corresponding")), "")
        if email:
            blocks.append(Block("correspondence", [(cfg.get("email_label", "E-mail:") + " ", plain), (email, plain)]))

    # History dates.
    dcfg = cfg.get("dates", {}) or {}
    labels = dcfg.get("labels", ["Received", "Revised", "Accepted"])
    no_xml = not (meta.article_type or meta.title or meta.authors or meta.dates)
    if meta.dates or no_xml and dcfg.get("without_xml", True) and f.authors:  # a plain .docx: "???" to fill in
        vals = [meta.dates.get(k) for k in dcfg.get("keys", ["received", "revised", "accepted"])]
        if meta.dates and _query_on(cfg, "missing_dates"):
            for lab, v in zip(labels, vals):  # the XML gives the history but not this date ("Revised: ???", IJO)
                if not v:
                    qs.append(("front.missing_date", {"what": lab.lower(), "_block": "dates",
                                                      "_quote": f"{lab}: {dcfg.get('missing', MISSING)}"}))
        items = [[(f"{lab}:", B), (" " + (v.strftime(dcfg.get("format", "%d-%m-%Y")) if v else dcfg.get("missing", MISSING)), plain)]
                 for lab, v in zip(labels, vals)]
        if dcfg.get("one_line"):
            runs = []
            for k, it in enumerate(items):
                if k:
                    runs.append(("\t", plain))
                runs += it
            blocks.append(Block("dates", runs))
        else:
            blocks += [Block("dates", it) for it in items]

    # Cite line.
    if f.authors and title:
        journal = cfg.get("cite_journal", "")
        if cfg.get("cite_journal_case_reports") and re.search(r"case reports\)", atype or "", re.I):
            journal = cfg["cite_journal_case_reports"]
        cites = [cite_name(a) for a in f.authors]
        mx = cfg.get("cite_max_authors", 6)
        runs = [(cfg.get("cite_label", "How to cite this article:"), B), (" " + ", ".join(cites[:mx]), plain)]
        if len(cites) > mx:
            runs += [(", ", plain), ("et al", I)]
        cite_title = recase(raw_title, "sentence", keep, colon)
        runs.append((f". {cite_title.rstrip('.')}. {journal + ' ' if journal else ''}XX;XX:XX.", plain))
        blocks.append(Block("cite", runs))
    if letter and cfg.get("letters_at_end", True):
        for b in blocks:
            b.end = b.role not in ("running_head", "article_type", "title")
    return blocks, qs


# --- proposing the edits -------------------------------------------------------------------

@dataclass
class NewParagraph:
    """A paragraph to insert next to an existing one (tracked as an inserted paragraph)."""
    id: str
    anchor: str  # para_id
    where: str  # "before" | "after"
    runs: list[Run]
    role: str
    status: Status = Status.SUGGESTED
    source: Source = Source.RULE
    reason: str = "front_matter"
    provenance: dict | None = None  # which input the paragraph was built from (edit_os/provenance.py)
    layer: Layer = Layer.STRUCTURE  # the review screen's layer

    @property
    def text(self) -> str:
        return "".join(t for t, _ in self.runs)

    def to_dict(self) -> dict:
        return dict(id=self.id, anchor=self.anchor, where=self.where, runs=[[t, f] for t, f in self.runs],
                    role=self.role, status=self.status.value, source=self.source.value, reason=self.reason,
                    provenance=self.provenance, layer=self.layer.value)


@dataclass
class Plan:
    changes: list[Change] = field(default_factory=list)
    inserts: list[NewParagraph] = field(default_factory=list)
    queries: list[queries.Query] = field(default_factory=list)
    blocks: list[Block] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)  # editor warnings (pipeline manifest)
    fields: Fields | None = None  # the authors and affiliations the blocks were built from (first_author.py)
    sources: Sources | None = None

    def add_queries(self, qs: queries.QuerySet, ms: Manuscript) -> queries.QuerySet:
        qs.queries += self.queries
        return qs.number(ms)


# Manuscript paragraph role -> the block role it holds.
_ROLE_TO_BLOCK = {"article_type": "article_type", "title": "title", "authors": "authors",
                  "affiliation": "affiliations", "correspondence": "correspondence"}
_DELETE_ROLES = {"running_title", "label"}


def _tokens_with_format(runs: list[Run]) -> list[tuple[str, dict]]:
    out = []
    for text, fmt in runs:
        for t in re.findall(r"\w+|\s+|[^\w\s]", text):
            out.append((t, fmt))
    return out


def _explicit_bold(text: str, fmt: dict | None) -> dict | None:
    """The format of text the rewrite inserts, its weight stated: inserted text otherwise takes the run it follows, and
    " Dr." after a bold "Address for correspondence:" came out bold (G87: only the label in bold; jehp_1755)."""
    if not text.strip():
        return fmt or None
    return {"bold": bool((fmt or {}).get("bold")), **(fmt or {})}


def _rewrite(p: Paragraph, runs: list[Run], layer: Layer) -> list[Change] | None:
    """Changes that turn paragraph `p` into `runs`: word-level text changes (inserted text carries
    its format) plus tracked formatting on kept text. None when that cannot be expressed without two
    insertions at one point (then the caller replaces the paragraph)."""
    from .rules_format import char_formats

    old = p.text
    a = re.findall(r"\w+|\s+|[^\w\s]", old)
    b = _tokens_with_format(runs)
    starts, pos = [], 0
    for t in a:
        starts.append(pos)
        pos += len(t)
    starts.append(pos)
    fmts = char_formats(p)
    out: list[Change] = []

    def mk(s, e, new, fmt=None):
        out.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=s, end=e, old=old[s:e], new=new, layer=layer,
                          source=Source.RULE, reason="front_matter", format=dict(fmt) if fmt else None))

    sm = difflib.SequenceMatcher(None, a, [t for t, _ in b], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            for k in range(i2 - i1):  # kept text: tracked formatting only
                want = b[j1 + k][1]
                s, e = starts[i1 + k], starts[i1 + k + 1]
                have = fmts[s] if s < len(fmts) else {}
                desired = {"bold": bool(want.get("bold")), "superscript": bool(want.get("superscript"))}
                if "italic" in want:
                    desired["italic"] = bool(want["italic"])
                delta = {prop: on for prop, on in desired.items() if bool(have.get(prop)) != on}
                if not delta:
                    continue
                if out and out[-1].format_only and out[-1].end == s and out[-1].format == delta:
                    s = out.pop().start
                mk(s, e, old[s:e], delta)
            continue
        segs: list[list] = []
        for t, f in b[j1:j2]:
            if segs and segs[-1][1] == f:
                segs[-1][0] += t
            else:
                segs.append([t, f])
        if not segs:
            mk(starts[i1], starts[i2], "")
            continue
        n_old = i2 - i1
        extra = None
        if len(segs) > 1 and n_old < len(segs):
            if len(segs) - n_old > 1:
                return None
            extra = segs.pop()  # one more piece than old words ("Pillai²" -> "Pillai" + superscript "1"): inserted after
        for k, (t, f) in enumerate(segs):
            lo = i1 + k if len(segs) > 1 else i1
            hi = lo + 1 if k < len(segs) - 1 else i2
            mk(starts[lo], starts[hi], t, _explicit_bold(t, f))
        if extra is not None:
            mk(starts[i2], starts[i2], extra[0], _explicit_bold(extra[0], extra[1]))
    return out


ARTICLE_TYPES = [
    "Original Article", "Case Report", "Case Series", "Review Article", "Systematic Review", "Meta-Analysis",
    "Letter to the Editor", "Short Communication", "Brief Communication", "Brief Report", "Research Letter",
    "Editorial", "Guest Editorial", "Commentary", "Perspective", "Viewpoint", "Photo Essay", "Surgical Technique",
    "Ophthalmic Images", "Images in Clinical Practice", "Book Review", "Correspondence", "Special Article",
]


def clean_article_type(raw: str) -> str:
    """Submission systems add qualifiers ("Case Report (for IJO Case Reports)"); the printed type
    is the plain name. Known types get their standard spelling."""
    t = re.sub(r"\([^)]*\)", " ", raw)
    t = re.split(r"\s+(?:for|-|–|:)\s+", t, maxsplit=1)[0]
    t = re.sub(r"\s+", " ", t).strip(" .:-–")
    for known in ARTICLE_TYPES:
        if t.lower() == known.lower() or t.lower().rstrip("s") == known.lower():
            return known
    return recase(t, "title", []) if t else raw.strip()

# --- where each block comes from (the review screen's "Source" chip) ----------------------------

_XML_FIELDS = {"article_type": "article type", "title": "title", "running_title": "running title"}


def _xml_hit(value: str, meta: Meta) -> str:
    """The submission XML field holding `value` ("article type", "author names", ...), or ""."""
    v = _norm(value)
    if not v:
        return ""
    for attr, name in _XML_FIELDS.items():
        if v and v in _norm(getattr(meta, attr)):
            return name
    for a in meta.authors:
        if v in _norm(" ".join([a.get("given", ""), a.get("surname", "")])):
            return "author names"
        if v in _norm(a.get("email", "")):
            return "author e-mail"
        if v in _norm(" ".join((a.get("aff") or {}).values())):
            return "affiliation"
    return ""


def locate(value: str, src: Sources, lead: list[Paragraph], prefer: tuple[str, ...] = ("title_page", "manuscript", "xml"),
           what: str = "") -> dict | None:
    """Which input `value` was read from: the title-page file, a manuscript paragraph or the submission XML."""
    v = _norm(value)
    if not v:
        return None
    for where in prefer:
        if where == "title_page":
            line = min((t for t in src.first_page if v in _norm(t)), key=len, default=None)  # the line it is most of
            if line is not None:
                return provenance.make("title_page", "Title page", line, file=src.first_page_name, field=what)
        elif where == "manuscript":
            p = min((p for p in lead if v in _norm(p.text)), key=lambda p: len(p.text), default=None)
            if p is not None:
                return provenance.make("manuscript", "Manuscript", p.text, para_id=p.id, field=what)
        elif where == "xml":
            hit = _xml_hit(value, src.meta)
            if hit:
                return provenance.make("xml", f"XML: {hit}", value, file=src.xml_name, field=hit)
    return None


def block_provenance(role: str, f: Fields, src: Sources, lead: list[Paragraph], cfg: dict, ms: Manuscript,
                     text: str = "") -> dict | None:
    """Where a front-matter block's content was read from (mirrors the choices `build` makes); `text` is the block's."""
    meta = src.meta
    if role == "article_type":
        first, second = (f.article_type, meta.article_type) if cfg.get("type_source") == "text" \
            else (meta.article_type, f.article_type)
        chosen = first or second
        if chosen and chosen == meta.article_type:
            return provenance.make("xml", "XML: article type", chosen, file=src.xml_name, field="article type")
        if chosen:
            return locate(chosen, src, lead, what="article type")
        if _imrad(ms):
            return provenance.make("inferred", "IMRaD headings", "Introduction, Methods, Results, Discussion and "
                                   "Conclusion headings: Original Article, with a query to confirm")
        return None
    if role == "title":
        if f.title:
            return locate(f.title, src, lead, what="title")
        return provenance.make("xml", "XML: title", meta.title, file=src.xml_name, field="title") if meta.title else None
    if role == "running_head":
        short = f.running_title or meta.running_title
        if not short:
            return provenance.make("inferred", "Not supplied", "author surnames from the author list; short title queried")
        if not f.running_title:
            return provenance.make("xml", "XML: running title", short, file=src.xml_name, field="running title")
        return locate(f.running_title, src, lead, what="running title")
    if role == "authors" and f.authors:
        a = f.authors[0]
        return locate(f"{a.given} {a.surname}".strip(), src, lead, what="authors") or \
            locate(a.surname, src, lead, what="authors")
    if role == "affiliations" and f.affiliations:
        aff = next(iter(f.affiliations.values()))
        for key in ("institution", "department", "city"):
            hit = locate(aff.get(key, ""), src, lead, what="affiliations")
            if hit:
                return hit
        return None
    if role == "correspondence":
        c = f.correspondence or {}
        if "@" in text:  # the e-mail line
            values = (c.get("email", ""),)
        elif text.startswith(cfg.get("correspondence_label", "Address for correspondence:")):  # the name line
            values = (c.get("surname", ""), c.get("email", ""))
        else:  # the address
            values = tuple(c.get("address_parts", [])[:2]) + (c.get("city", ""), c.get("email", ""))
        for value in values:
            hit = locate(value, src, lead, what="correspondence")
            if hit:
                return hit
        corr = next((a for a in meta.authors if a.get("corresponding")), None)
        if corr:
            return provenance.make("xml", "XML: corresponding author", corr.get("email", "") or corr.get("surname", ""),
                                   file=src.xml_name, field="corresponding author")
        return None
    if role == "dates":
        if meta.dates:
            return provenance.make("xml", "XML: history dates", ", ".join(f"{k} {v:%d-%m-%Y}" for k, v in meta.dates.items()),
                                   file=src.xml_name, field="history dates")
        return provenance.make("inferred", "Not supplied", "no submission XML: dates left to fill in")
    if role == "cite":
        return provenance.make("journal_style", "Title + authors", "built from the title and author names; journal name "
                               "from the journal's style settings")
    return None


def _consistent_name_order(f: Fields) -> None:
    """"Nguyen Thi Diep, Tran Hoa, Nguyen Duc Khanh" split as given "Nguyen Thi" / surname "Diep" but "Tran Hoa" as one
    surname with no given name: the same order for every author (the last word is the surname), so the cite line
    reads "Diep NT, Hoa T, Khanh ND", not "Tran Hoa"."""
    others = [a for a in f.authors if a.given.strip()]
    if not others:
        return
    for a in f.authors:
        if not a.given.strip() and len(a.surname.split()) == 2 and not any(
                w.casefold() in _PARTICLES | {"al", "el", "bin", "ben", "abu", "abd", "ibn", "mac", "mc", "st"}
                for w in a.surname.split()) and not a.surname[:1].islower():
            given, surname = a.surname.split()
            f.notes.append(f"{a.surname!r}: given name {given!r}, surname {surname!r}, as the co-authors' names are split")
            a.given, a.surname = given, surname


_RUNNING_LINE = re.compile(r"^\s*(?:running\s*(?:title|head)|short\s*(?:title|head))\s*[:\-–]\s*(\S.*?)\s*$", re.I)
_INTRO_HEAD = re.compile(r"^\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?introduction\b", re.I)


_RUNNING_LABEL = re.compile(r"^\s*(?:running\s*(?:title|head)|short\s*(?:title|head))\s*[:\-–]?\s*$", re.I)


def _short_title_line(p: Paragraph) -> bool:
    """A line that can be a short title on its own: not in a table, not a heading or body text, a few words."""
    t = p.text.strip()
    return bool(t) and not p.in_table and len(t) < 150 and len(t.split()) <= 15 and not _BODY_START.match(t) \
        and not _BACK_MATTER.match(t) and not re.search(r"[.;]\s*$|\.\s+[A-Z]", t)


def late_running_title(ms: Manuscript, lead: list[Paragraph]) -> tuple[Paragraph, str, Paragraph | None] | None:
    """A "Running Title: ..." line below the title page (after the abstract and key words, or below an "Abstract
    Page" heading), before the Introduction heading: front matter left in the body. Its text is the running title,
    and the line goes (the running head carries it). Also the label on a line of its own with the short title on
    the next ("Running title" / "Uric Acid-to-HDL Ratio in ACS"): both go. Returns (the paragraph with the short
    title, the short title, the label paragraph when separate)."""
    seen = {p.id for p in lead}
    paras = [p for p in ms.paragraphs[:80] if p.text.strip()]
    for k, p in enumerate(paras):
        if _INTRO_HEAD.match(p.text) and len(p.text) < 40:
            break
        if p.id in seen or p.in_table:
            continue
        if _RUNNING_LINE.match(p.text) and len(p.text) < 200:
            return p, _RUNNING_LINE.match(p.text).group(1), None
        if _RUNNING_LABEL.match(p.text) and k + 1 < len(paras) and paras[k + 1].id not in seen \
                and _short_title_line(paras[k + 1]):
            return paras[k + 1], paras[k + 1].text.strip(), p
    return None


def _xml_has_more(f: Fields, meta: Meta) -> bool:
    """The split names a strict subset of the XML's authors: the text is a cover letter signed by the corresponding
    author, or a title page that leaves authors out. The submission's author list is the complete one (S: the
    author group, running head and cite line name every author), even when the names the text gives are linked to
    an affiliation."""
    xml = [_norm(a.get("surname", "")) for a in meta.authors if a.get("surname")]
    got = [_norm(a.surname) for a in f.authors]
    return bool(got) and len(xml) > len(got) and set(got) < set(xml)


def plan(ms: Manuscript, style: JournalStyle, src: Sources | None = None) -> Plan:
    """Everything the front-matter step proposes for this manuscript."""
    cached = getattr(ms, "_front_plan", None)
    if cached is not None:
        return cached
    src = src or getattr(ms, "front_sources", None) or Sources()
    ms._submission_style = style
    cfg = _settings(style)
    pl = Plan()
    if not cfg.get("enabled", True):
        return pl
    lead = leading_paragraphs(ms)
    text, para_ids = prompt_text(src, lead)
    raw = split(text) if cfg.get("ai_split", True) and text.strip() else None
    if raw is not None:
        source_text = "\n".join(src.first_page + [p.text for p in lead]) + "\n" + src.meta.text()
        f = verify(raw, source_text, para_ids)
        for a in f.authors:
            a.written = " ".join(x for x in (a.given, a.surname) if x)
        xml_names = [Author(a.get("given", ""), a.get("surname", "")) for a in src.meta.authors]
        xml_names += [Author(g, sn) for g, sn in submission.notice_names(ms)]  # the portal's copyright notice
        written = _norm(" ".join(t for t in src.first_page + [p.text for p in lead] if not submission._COPYRIGHT.match(t)))
        for a in f.authors:  # "KavyaRaman": the submission XML gives "Kavya" "Raman"
            if (split_name := unglued_name(a, xml_names)) is not None:
                f.notes.append(f"name {a.given} {a.surname!r} spaced as the submission XML / notice has it: "
                               f"{split_name[0]} {split_name[1]!r}")
                a.given, a.surname = split_name
        for a in f.authors:  # "Meera K R": the submission XML has the surname in full
            if _initial_after_name(a) and not any(
                    _norm(f"{x.given}{x.surname}") in written for x in xml_names if x.surname and len(x.surname) > 1
                    and _norm(x.surname)[:1] == _norm(a.surname)[:1]):
                # "Natarajan S": a name with its initial after it (G43: the initial goes before it, "S. Natarajan"),
                # not expanded from the submission XML / copyright notice ("natarajan shanmugasundaram") when the
                # title page never writes the name in full: the name written stays the one in the running head and
                # cite line ("Natarajan", the editors in jehp_1683; owner, 2026-10-27)
                f.notes.append(f"{a.given} {a.surname}: kept as written (the initial is not expanded from the "
                               "submission XML / notice)")
                continue
            full = _expand(a, xml_names) if _initial_like(a.surname) else None
            if full is not None:
                f.notes.append(f"surname {a.surname!r} expanded to {full.surname!r} from the submission XML / notice")
                a.surname = full.surname
            elif any(len(t.strip(".")) == 1 for t in _name_tokens(a)):
                full = _expand_from(a, xml_names)  # "A. Fachruddin B": "Andi Fachruddin Benyamin" in the XML
                if full is not None:
                    f.notes.append(f"name {a.given} {a.surname!r} expanded to {full.given} {full.surname!r} "
                                   "from the submission XML / notice")
                    a.given, a.surname = full.given, full.surname
        if not f.authors and src.meta.authors:  # nothing usable from the text
            m = fields_from_meta(src.meta)
            f.authors, f.affiliations = m.authors, m.affiliations
            f.correspondence = f.correspondence or m.correspondence
        elif _xml_has_more(f, src.meta) and len(f.authors) >= 2 and all(a.affs for a in f.authors) and f.affiliations:
            # A title page whose author list leaves out an author the submission system names (IJO_505: the XML's 7th
            # author): the title page's list stands and the author is queried, never added (owner, 2026-10-27).
            have = {_norm(a.surname) for a in f.authors}
            f.xml_only = [" ".join(x for x in (a.get("prefix", ""), a.get("given", ""), a.get("surname", "")) if x)
                          for a in src.meta.authors if a.get("surname") and _norm(a["surname"]) not in have]
            f.notes.append(f"the submission XML names {', '.join(f.xml_only)}, not in the title page's author list: "
                           "queried, not added")
        elif _xml_has_more(f, src.meta):
            # The text gave only some of the submission's authors (no title page: a cover letter): the XML's list.
            m = fields_from_meta(src.meta)
            f.notes.append("authors taken from the submission XML: the text names only "
                           f"{len(f.authors)} of its {len(m.authors)} authors")
            f.authors, f.affiliations = m.authors, m.affiliations
        elif _names_from_contributions(f, src, lead):
            # No author list in the text (a cover letter as the title page): names read off the contribution
            # statements ("Definition of intellectual content: Dr. Radhika Paranjpe") are not the authors; the XML's are.
            m = fields_from_meta(src.meta)
            f.notes.append("authors taken from the submission XML: the text names authors only in contribution "
                           "statements, and not the submission's")
            f.authors, f.affiliations = m.authors, m.affiliations
    else:
        f = fields_from_meta(src.meta)
    _consistent_name_order(f)
    late = late_running_title(ms, lead)
    late_rt = late[0] if late is not None else None
    if late is not None and not f.running_title:
        f.running_title = late[1]
        f.notes.append(f"running title from the line after the key words: {f.running_title!r}")
    _fix_roles(f, lead)
    _link_affiliations(f, src, lead)
    _listed_affiliation_order(f, src, lead)
    _per_author_list(f)
    _inline_numbered_affiliations(f, src, lead)
    _equal_contribution(f, src, lead)
    _xml_places(f, src.meta)
    _xml_corresponding(f, src.meta)
    for a in f.authors:
        before = f"{a.given} {a.surname}"
        if initials_first(a):
            f.notes.append(f"initials moved before the name: {before!r} -> {a.given} {a.surname!r} (guide)")
    pl.notes = f.notes
    if not (f.authors or f.title or src.meta.title):
        ms._front_plan = pl
        return pl
    blocks, qspecs = build(ms, f, src.meta, style)
    pl.blocks = blocks
    pl.warnings = [{"para_id": "", "rule": "front_matter.affiliations", "message": w} for w in f.warnings]
    pl.fields, pl.sources = f, src
    for b in blocks:
        b.provenance = block_provenance(b.role, f, src, lead + ([late_rt] if late_rt is not None else []), cfg, ms,
                                        b.text)

    # Map existing manuscript paragraphs onto blocks; insert the rest in reading order.
    have: dict[str, list[Paragraph]] = {}
    front_roles = [p for p in lead if f.roles.get(p.id) in set(_ROLE_TO_BLOCK) | _DELETE_ROLES]
    for k, p in enumerate(lead):
        role = f.roles.get(p.id)
        if role == "running_title" and front_roles and front_roles[0] is p:
            have.setdefault("running_head", []).append(p)  # a running head on top: rewrite it in place
        elif role in _ROLE_TO_BLOCK:
            have.setdefault(_ROLE_TO_BLOCK[role], []).append(p)
        elif (role in _DELETE_ROLES and _LABEL.match(p.text) and (role == "running_title" or len(p.text) < 60)) \
                or (role not in _ROLE_TO_BLOCK and _RUNNING_LABEL.match(p.text)):
            pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=0, end=len(p.text), old=p.text,
                                     new="", layer=Layer.STRUCTURE, source=Source.RULE, reason="front_matter",
                                     provenance=provenance.make("manuscript", "Manuscript label", p.text, para_id=p.id,
                                                                role=role or "label")))
        elif role not in _ROLE_TO_BLOCK and _short_title_line(p) and (
                (f.running_title and _norm(p.text) == _norm(f.running_title) and _norm(p.text) != _norm(f.title))
                or (role == "running_title" and _norm(p.text) != _norm(f.title))
                or (k and _RUNNING_LABEL.match(lead[k - 1].text))):
            # The author's short title on a line of its own ("Marital Enrichment and Childbearing" below the title,
            # or below a bare "Running Title" label): the running head carries it, the line goes like the labels.
            pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=0, end=len(p.text), old=p.text,
                                     new="", layer=Layer.STRUCTURE, source=Source.RULE, reason="front_matter",
                                     provenance=provenance.make("manuscript", "Manuscript", "the short title on a line "
                                                                "of its own: it goes into the running head",
                                                                para_id=p.id, role="running_title")))
    for q in ([late[0], late[2]] if late is not None else []):
        if q is None:
            continue
        pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=q.id, start=0, end=len(q.text),
                                 old=q.text, new="", layer=Layer.STRUCTURE, source=Source.RULE,
                                 reason="front_matter",
                                 provenance=provenance.make("manuscript", "Manuscript", "the running title below the "
                                                            "title page: it goes into the running head", para_id=q.id,
                                                            role="running_title")))
    by_role: dict[str, list[Block]] = {}
    for b in blocks:
        by_role.setdefault(b.role, []).append(b)
    placed: dict[int, tuple[str, str]] = {}  # block index -> (para id, "in place")
    block_insert: dict[int, NewParagraph] = {}  # block index -> the paragraph inserted for it
    # Paragraphs rewritten in place keep their position, so they must already stand in the blocks' order: an
    # "Article type:" line below the title cannot become the article type above it. Keep the heaviest run of
    # paragraphs in block order (`_in_order`); the others are deleted and their blocks inserted where they belong.
    pos = {p.id: i for i, p in enumerate(ms.paragraphs)}
    assign = []  # (block index, paragraph, role)
    for role, paras in have.items():
        outs = [i for i, b in enumerate(blocks) if b.role == role]
        assign += [(outs[k], p, role) for k, p in enumerate(paras) if k < len(outs)]
    assign.sort(key=lambda a: a[0])
    keep = _in_order([pos[p.id] for _, p, _ in assign], [len(blocks[i].text) + 1 for i, _, _ in assign])
    moved = {id(assign[k][1]) for k in range(len(assign)) if k not in keep}
    for role, paras in have.items():
        outs = [i for i, b in enumerate(blocks) if b.role == role]
        for k, p in enumerate(paras):
            if k < len(outs) and id(p) in moved:
                pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=0, end=len(p.text), old=p.text,
                                         new="", layer=Layer.HOUSE_STYLE, source=Source.RULE, reason="front_matter",
                                         provenance=provenance.make("manuscript", "Manuscript", f"the {role.replace('_', ' ')} "
                                                                    "is moved to its place in the front matter",
                                                                    para_id=p.id, role=role)))
            elif k < len(outs):
                b = blocks[outs[k]]
                prov = dict(b.provenance, role=role) if b.provenance else {"kind": "manuscript", "label": "Manuscript",
                                                                           "para_id": p.id, "role": role}
                ch = _rewrite(p, b.runs, Layer.HOUSE_STYLE)
                if ch is None:  # replace the paragraph instead
                    pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=0, end=len(p.text), old=p.text,
                                             new="", layer=Layer.HOUSE_STYLE, source=Source.RULE, reason="front_matter",
                                             provenance=prov))
                    pl.inserts.append(NewParagraph(f"fi{next(_ids)}", p.id, "after", b.runs, role, provenance=b.provenance))
                    block_insert[outs[k]] = pl.inserts[-1]
                else:
                    for c in ch:
                        c.provenance = prov
                    pl.changes += ch
                placed[outs[k]] = (p.id, "in place")
            else:  # more paragraphs than the block needs
                pl.changes.append(Change(id=f"fm{next(_ids)}", para_id=p.id, start=0, end=len(p.text), old=p.text,
                                         new="", layer=Layer.HOUSE_STYLE, source=Source.RULE, reason="front_matter",
                                         provenance=provenance.make("manuscript", "Manuscript", "repeats the "
                                                                    f"{role.replace('_', ' ')} given above",
                                                                    para_id=p.id, role=role)))
    # The portal's copyright notice on top stays above the running head (it becomes the first author query there).
    notice = {p.id for p in ms.paragraphs[:5] if not p.in_table and submission._COPYRIGHT.match(p.text)}
    first_any = next((p for p in ms.paragraphs if p.text.strip() and p.id not in notice),
                     ms.paragraphs[0] if ms.paragraphs else None)
    first_placed = next((placed[i][0] for i in sorted(placed)), None)
    last: tuple[str, str] | None = None
    from .references import find_heading

    ref = find_heading(ms)
    end_anchor = (ms.paragraphs[ref].id, "before") if ref is not None else (
        (ms.paragraphs[-1].id, "after") if ms.paragraphs else None)
    for i, b in enumerate(blocks):
        if i in placed:
            last = (placed[i][0], "after")
            continue
        if b.end and end_anchor:
            pl.inserts.append(NewParagraph(f"fi{next(_ids)}", end_anchor[0], end_anchor[1], b.runs, b.role,
                                           provenance=b.provenance))
            block_insert[i] = pl.inserts[-1]
            continue
        if last is None:
            anchor = (first_placed or (first_any.id if first_any else None), "before")
        else:
            anchor = last
        if anchor[0] is None:
            continue
        pl.inserts.append(NewParagraph(f"fi{next(_ids)}", anchor[0], anchor[1], b.runs, b.role, provenance=b.provenance))
        block_insert[i] = pl.inserts[-1]
    anchor_q = first_placed or (first_any.id if first_any else None)
    for w in pl.warnings:
        w["para_id"] = anchor_q or ""
    pl.queries += _front_queries(ms, pl, blocks, qspecs, placed, block_insert, anchor_q)
    ms._front_plan = pl
    return pl


def _front_queries(ms: Manuscript, pl: Plan, blocks: list[Block], qspecs: list[tuple[str, dict]],
                   placed: dict[int, tuple[str, str]], block_insert: dict[int, "NewParagraph"],
                   fallback: str | None) -> list[queries.Query]:
    """The front-matter queries on the block they ask about: "Kindly provide the full name of author ..." on that
    name in the author group, the running-head queries on the running head. A block rewritten in place carries the
    query on its paragraph; a block inserted (or a paragraph replaced by one) carries it on the inserted paragraph
    (`Query.insert_id`), never on a paragraph the front matter deletes. Without the block, the first paragraph of
    the front matter that stays."""
    deleted = {c.para_id for c in pl.changes if c.start == 0 and c.end == len(c.old) and c.old.strip() and not c.new}
    stays = next((p.id for p in ms.paragraphs if p.text.strip() and p.id not in deleted), None)
    first_insert = pl.inserts[0] if pl.inserts else None

    def general(text: str, check: str) -> queries.Query | None:
        if fallback is not None and fallback not in deleted:
            return queries.Query(fallback, 0, 0, text, check)
        if first_insert is not None:  # the top of the front matter, as delivered
            return queries.Query(first_insert.anchor, 0, 0, text, check, insert_id=first_insert.id,
                                 quote=first_insert.text)
        return queries.Query(stays, 0, 0, text, check) if stays else None

    out = []
    for check, fmt in qspecs:
        role, quote = fmt.get("_block"), fmt.get("_quote", "")
        text = queries.TEXT[check].format(**{k: v for k, v in fmt.items() if not k.startswith("_")})
        roles = [k for k, b in enumerate(blocks) if b.role == role] if role else []
        # the block of that role that holds the quote ("Revised: ???" among the dates, the address's country)
        i = next((k for k in roles if quote and quote in blocks[k].text), roles[0] if roles else None)
        q = None
        if i is not None and i in block_insert:
            n = block_insert[i]
            s = n.text.find(quote) if quote else -1
            q = queries.Query(n.anchor, max(s, 0), s + len(quote) if s >= 0 else 0, text, check, insert_id=n.id,
                              quote=quote if s >= 0 else n.text)
        elif i is not None and i in placed and placed[i][0] not in deleted:
            p = ms.by_id(placed[i][0])
            s = p.text.find(quote) if quote else -1
            if s < 0 and quote:  # the name as written ("Natarajan S" for "S. Natarajan")
                sur = quote.split()[-1]
                s = p.text.find(sur)
                quote = sur if s >= 0 else quote
            q = queries.Query(p.id, max(s, 0), s + len(quote) if s >= 0 else 0, text, check)
        if q is None:
            q = general(text, check)
        if q is None:
            continue
        out.append(q)
    return out


def _in_order(positions: list[int], weights: list[int]) -> set[int]:
    """Indexes of the heaviest strictly increasing subsequence of `positions` (ties: the earliest items)."""
    n = len(positions)
    best, prev = [0] * n, [-1] * n
    for i in range(n):
        best[i] = weights[i]
        for j in range(i):
            if positions[j] < positions[i] and best[j] + weights[i] > best[i]:
                best[i], prev[i] = best[j] + weights[i], j
    if not n:
        return set()
    i = max(range(n), key=lambda k: (best[k], -k))
    out = set()
    while i >= 0:
        out.add(i)
        i = prev[i]
    return out


_LABEL_ROLE = [(re.compile(r"^\s*corresponding\s*author", re.I), "correspondence"),
               (re.compile(r"^\s*(?:list of\s+)?authors?(?:\s*names?)?\s*[:\-–]", re.I), "authors"),
               (re.compile(r"^\s*affiliations?\s*[:\-–]", re.I), "affiliation")]


# Authors' contribution statements ("Fatma Savur: Data acquisition, data analysis, manuscript review."): what an
# author did, not a title-page author line. The words that name a contribution (CRediT roles and the usual ones).
_ROLE_WORDS = re.compile(
    r"\b(?:concept\w*|design\w*|data|acquisition|analy[sz]\w*|interpretation|draft\w*|writing|wrote|review\w*|"
    r"edit\w*|supervis\w*|literature|search|manuscript|statistic\w*|investigation|methodology|curation|"
    r"visuali[sz]\w*|validation|resources|software|funding|administration|approv\w*|revis\w*|intellectual|"
    r"guarantor|collect\w*|prepar\w*|definition)\b", re.I)
# What stands before the colon of a label line, not an author's name ("Article Type: Review Article").
_NOT_NAME = re.compile(
    r"\b(?:title|type|article|manuscript|running|short|key|keywords?|words?|abstract|funding|conflicts?|source|"
    r"support|count|address|e-?mail|phone|tel|fax|correspond\w*|authors?|affiliations?|department|contributions?|"
    r"acknowledge?ments?|declarations?|ethic\w*|financial|competing|data|availability|orcid|received|accepted|"
    r"submitted|revised|figures?|tables?|category|section|specialty|subject|introduction|background|aims?|methods?|"
    r"results?|conclusions?|summary|approval|permissions?|consent|sub|re|note|cover|letter|generative|"
    r"statement|disclosure|revision|submission|declared)\b", re.I)


def is_contribution_line(text: str) -> bool:
    """"Name: role, role, role": an author's contribution statement ("Fatma Savur: Data acquisition, data analysis,
    manuscript review.", "A.B. – Conception and design"). Most of the items after the separator are short and name a
    contribution, with at least two contribution words; a label ("Title:", "Article Type:") is not a name. A paragraph
    of several lines ("YJ: Data collection\nAG: Editing") is one when every line is."""
    lines = [t.strip() for t in text.split("\n") if t.strip()]
    return bool(lines) and len(text) <= 600 and all(_contribution_line(t) for t in lines)


_CONTRIB_NAME = re.compile(
    r"^(?:[•➢▪*]\s*)?(?:(?:prof|dr|mr|mrs|ms)\.?\s*)*(?:(?:[A-ZÀ-ÖØ-Þİ][\w.’'\-]*\.?|(?:%s)\b)\s*){1,7}"
    r"(?:\([A-ZÀ-ÖØ-Þİ. ]{1,10}\)\s*)?$" % "|".join(sorted(_PARTICLES | {"al", "el", "bin", "binti"})))


def _contribution_line(text: str) -> bool:
    sep = re.search(r"\s*:\s*|\s+[–—-]\s*|[–—-]\s+", text)
    if not sep or sep.start() == 0:
        return False
    head, rest = text[:sep.start()], text[sep.end():].strip(" .")
    if not _CONTRIB_NAME.match(head) or _NOT_NAME.search(head) or _LABEL.match(head):
        return False
    if len(_ROLE_WORDS.findall(rest)) < 2 or re.match(r"(?:an?|the)\s+[A-Z]", rest, re.I):
        return False  # "...: A rare case report with review of the literature" is a title
    items = [t for t in re.split(r"\s*(?:[,;]|\band\b|&)\s*", rest) if t.strip()]
    roles = [t for t in items if _ROLE_WORDS.search(t) and len(t.split()) <= 7]
    return bool(items) and len(roles) * 2 >= len(items)


def contribution_paragraphs(lead: list[Paragraph]) -> set[str]:
    """Ids of the authors' contribution statements among the leading paragraphs: "Name: role" lines, an "Author(s)
    contribution(s)" heading or label line, and the statements after such a heading (up to the next label or a
    paragraph that names no contribution). They are not front matter: the front-matter rebuild leaves them alone."""
    from .rules_sections import _CONTRIB, _CONTRIB_WORDS, _OTHER_LABEL, _TEMPLATE

    out: set[str] = set()
    after_head = False
    for p in lead:
        text = p.text.strip()
        m = _CONTRIB.match(text)
        if m and not (m.group("sep").isspace() and not m.group("author")):
            out.add(p.id)
            after_head = not text[m.end():].strip()
            continue
        if is_contribution_line(text):
            out.add(p.id)
            continue
        if after_head and _CONTRIB_WORDS.search(text) and not (_TEMPLATE.search(text) or _OTHER_LABEL.match(text)
                                                               or _LABEL.match(text) or _BODY_START.match(text)):
            out.add(p.id)  # "All authors approved the final version ..." under the heading
            continue
        after_head = False
    return out


def _fix_roles(f: Fields, lead: list[Paragraph]) -> None:
    """Deterministic corrections to the roles Claude gives manuscript paragraphs: an author-written running
    head on top ("Rao, et al.: ...") is the running head, not the title; a label with its content after it
    ("Corresponding Author:" + the name on the next line) holds that content, so it is not deleted as a label;
    authors' contribution statements ("Fatma Savur: Data acquisition, ...") are not author lines or labels, so
    they are never rewritten or deleted here."""
    contrib = contribution_paragraphs(lead)
    title = _norm(f.title)
    for p in lead:
        if p.id in contrib and f.roles.get(p.id) not in (None, "body", "other_front_matter") \
                and not (title and title in _norm(p.text)):  # never the title ("Disease X: A review of ...")
            f.notes.append(f"kept as an authors' contribution statement, not {f.roles[p.id].replace('_', ' ')}: "
                           f"{p.text[:80]!r}")
            f.roles[p.id] = "other_front_matter"
    first = next((p for p in lead if f.roles.get(p.id) not in (None, "body", "other_front_matter")), None)
    for p in lead:
        role = f.roles.get(p.id)
        if p is first and role == "title" and re.match(r"^\s*(?:<RH>)?[^:\n]{1,60}?,?\s+et\s+al\.?\s*:", p.text):
            f.roles[p.id] = "running_title"
        elif role == "label":
            m = _LABEL.match(p.text)
            rest = p.text[m.end():].strip(" :.-–\n\t") if m else ""
            if rest:
                f.roles[p.id] = next((r for rx, r in _LABEL_ROLE if rx.match(p.text)), role)


def _contribution_mention(text: str) -> bool:
    """A contribution statement: "Name: role, role" or "Role, role: Name, Name"."""
    if is_contribution_line(text):
        return True
    m = re.match(r"^\W*([^:\n]{2,80}?)\s*:\s*\S", text)
    return bool(m and _ROLE_WORDS.search(m.group(1)))


def _names_from_contributions(f: Fields, src: Sources, lead: list[Paragraph]) -> bool:
    """Some of the text's authors are not in the submission XML, and the text names each of them only in authors'
    contribution statements."""
    xml = {_norm(a.get("surname", "")) for a in src.meta.authors if a.get("surname")}
    if not xml or len(src.meta.authors) < len(f.authors):
        return False
    odd = [a for a in f.authors if _norm(a.surname) not in xml]
    if not odd:
        return False
    lines = [t for text in list(src.first_page) + [p.text for p in lead] for t in text.split("\n") if t.strip()]
    for a in odd:
        hits = [t for t in lines if re.search(r"(?i)\b" + re.escape(a.surname) + r"\b", t)]
        if not hits or not all(_contribution_mention(t) for t in hits):
            return False
    return True


_EQUAL_NOTE = re.compile(r"^\s*([†‡§#¶])\s*(\S.*\b(?:contributed\s+equally|equal(?:ly)?\s+contribut\w*|joint\s+first|"
                         r"shared?\s+(?:the\s+)?first|co-?first)\b.*)$", re.I)


def _equal_contribution(f: Fields, src: Sources, lead: list[Paragraph]) -> None:
    """"Yijie Lu1,2†, Yao Tan1,3†" with "†These authors contributed equally to this work ...": the mark stays after the
    names (after the affiliation numbers) and the footnote under the affiliations, as written."""
    lines = [t for text in list(src.first_page) + [p.text for p in lead] for t in text.split("\n") if t.strip()]
    note = next((m for m in map(_EQUAL_NOTE.match, lines) if m), None)
    if note is None:
        return
    mark = note.group(1)
    marked = [k for k, a in enumerate(f.authors) if any(
        re.search(re.escape(a.surname) + r"[\s\d,*]*" + re.escape(mark), t) for t in lines if not _EQUAL_NOTE.match(t))]
    if marked:
        f.equal_mark, f.equal_note, f.equal_authors = mark, note.group(2).strip(), marked


def _link_affiliations(f: Fields, src: Sources, lead: list[Paragraph]) -> None:
    """Affiliations listed without numbers ("Institutional Affiliations:" then one per line) under authors marked
    "Dr. Varnika Sahi,1 MBBS": the marks number the list in order. Only when no author is linked yet, every mark
    names a listed affiliation, and each author's name is followed by marks somewhere in the text."""
    if len(f.affiliations) < 2 or any(a.affs for a in f.authors):
        return
    ids = sorted(f.affiliations)
    lines = list(src.first_page) + [p.text for p in lead]
    found: dict[int, list[int]] = {}
    for k, a in enumerate(f.authors):
        rx = re.compile(re.escape(a.surname) + r"\s*[,*]?\s*(\d{1,2}(?:\s*,\s*\d{1,2})*)(?!\d)")
        for t in lines:
            m = rx.search(t)
            if m:
                nums = [int(n) for n in re.findall(r"\d+", m.group(1))]
                if all(1 <= n <= len(ids) for n in nums):
                    found[k] = [ids[n - 1] for n in nums]
                break
    if len(found) != len(f.authors):
        return
    for k, a in enumerate(f.authors):
        a.affs = found[k]
    f.notes.append("authors linked to the unnumbered affiliation list by the marks after their names")


_AFF_LIST_LABEL = re.compile(r"(?i)^\W*(?:departments?|affiliations?|institutions?|institutional)\b[^:,]{0,40}:?\W*$")


def _author_marks(a: "Author", lines: list[str]) -> tuple[int, list[int]] | None:
    """(line index, the numbers marked after the author's name: "Dr. Ruhi Girish Jange DNB1", "Bipasha Mukherjee MS,
    DNB1", "A. Rao2,3")."""
    rx = re.compile(r"(?<![\w-])" + re.escape(a.surname) + r"(?:[\s,]+(?:[A-Z][A-Za-z.()]{0,9}))*?\s*,?\s*\*?\s*"
                    r"(\d{1,2}(?:\s*,\s*\d{1,2})*)\s*\*?(?=\s*(?:[,;]|\band\b|$))")
    for k, t in enumerate(lines):
        m = rx.search(t)
        if m:
            return k, [int(n) for n in re.findall(r"\d+", m.group(1))]
    return None


def _listed_affiliation_order(f: "Fields", src: "Sources", lead: list[Paragraph]) -> None:
    """Affiliations listed one per line without numbers under authors marked with numbers ("Dr. Ruhi Jange DNB1", ...,
    "Dr. Shahid Alam MS2" over "Orbit, Oculoplasty ... Services" / "Medical Research Foundation, ..., Chennai" /
    "Orbit, ..., Kolkata"): mark n is the n-th affiliation in the list's order, whatever ids the split gave (it can
    number them from the cover letter's signature, swapping them: IJO_2733). A department on its own line right above
    an affiliation the split gave none ("Orbit, Oculoplasty, Reconstructive and Aesthetic Services" over "Medical
    Research Foundation, Sankara Nethralaya, Chennai, India") is that affiliation's department."""
    if len(f.affiliations) < 2 or not f.authors:
        return
    lines = [t.strip() for text in submission.texts_without_parts(list(src.first_page)) + [p.text for p in lead]
             for t in text.split("\n") if t.strip()]
    found = [_author_marks(a, lines) for a in f.authors]
    if any(x is None for x in found):
        return
    start = max(k for k, _ in found) + 1
    marks = sorted({n for _, ns in found for n in ns})
    if marks != list(range(1, len(f.affiliations) + 1)):
        return
    pos: dict[int, int] = {}
    for i, aff in f.affiliations.items():
        key = _norm(aff.get("institution") or aff.get("department") or "")
        hit = next((k for k in range(start, len(lines)) if key and key in _norm(lines[k])), None)
        if hit is None:
            return
        pos[i] = hit
    if len(set(pos.values())) != len(pos) or any(re.match(r"^\W{0,2}\d", lines[k]) for k in pos.values()):
        return  # two affiliations on one line, or a numbered list: the numbers are written
    order = sorted(pos, key=pos.get)
    new = [[order[n - 1] for n in ns] for _, ns in found]
    if any(sorted(a.affs) != sorted(n) for a, n in zip(f.authors, new)):
        for a, n in zip(f.authors, new):
            a.affs = n
        f.notes.append("authors linked to the unnumbered affiliation list in its order by the marks after their names")
    insts = [_norm(a.get("institution", "")) for a in f.affiliations.values() if a.get("institution")]
    depts = {_norm(a.get("department", "")) for a in f.affiliations.values() if a.get("department")}
    for i in order:
        aff = f.affiliations[i]
        k = pos[i] - 1
        if aff.get("department") or k < start:
            continue
        above = lines[k]
        if k in pos.values() or len(above) > 150 or _AFF_LIST_LABEL.match(above) or _LABEL.match(above) \
                or any(x in _norm(above) for x in insts) or re.match(r"^\W{0,2}\d", above):
            continue
        if _norm(above) in depts or _UNIT_WORDS.search(above):
            aff["department"] = above.strip(" ,.;")
            f.notes.append(f"affiliation {aff.get('institution')!r}: department {aff['department']!r} from the line above it")


def _aff_key(a: dict) -> tuple:
    return (_norm(department_form(a.get("department", ""))),) + tuple(
        _norm(a.get(k, "")) for k in ("institution", "city", "state", "country"))


def _per_author_list(f: "Fields") -> None:
    """A numbered list with one affiliation per author, in author order ("1. MS, Assistant Professor, Guru Nanak Eye
    Center, ..." to "6. ..."), under names marked with the numbers of the distinct institutions instead ("saran,2
    kratee kratee2" for the pathology department, 1 for the eye centre: IJO_505): the split takes the marks as list
    numbers and links the pathologists to the eye centre. When the marks group the authors exactly as the per-author
    list does, author k has affiliation k."""
    ids = sorted(f.affiliations)
    if len(ids) != len(f.authors) or len(ids) < 2 or any(len(a.affs) != 1 for a in f.authors):
        return
    keys = [_aff_key(f.affiliations[i]) for i in ids]
    by_mark: dict[int, set] = {}
    for k, a in enumerate(f.authors):
        by_mark.setdefault(a.affs[0], set()).add(k)
    by_key: dict[tuple, set] = {}
    for k, key in enumerate(keys):
        by_key.setdefault(key, set()).add(k)
    same_partition = sorted(map(sorted, by_mark.values())) == sorted(map(sorted, by_key.values()))
    if not same_partition or len(by_key) == len(ids):
        return
    if all(keys[ids.index(a.affs[0])] == keys[k] for k, a in enumerate(f.authors)):
        return  # already linked so
    for k, a in enumerate(f.authors):
        a.affs = [ids[k]]
    f.notes.append("the numbered list gives one affiliation per author: author k linked to affiliation k (the marks "
                   "after the names group the authors the same way)")


_INLINE_AFFS = re.compile(r"^\s*(?:institutional\s+)?affiliations?\s*[:\-–]?\s*(?=1\s*[.)]?\s*[A-Za-z])", re.I)


def _affiliation_parts(text: str) -> dict:
    """"Radiography Department, University of Suffolk, Ipswich, United Kingdom" -> its parts."""
    parts = [x.strip() for x in text.strip(" .;,").split(",") if x.strip()]
    out = dict(department="", institution="", city="", state="", country="")
    if len(parts) > 1 and parts[-1].casefold() in _COUNTRY_NAMES:
        out["country"] = parts.pop()
    if len(parts) > 2 or (len(parts) == 2 and not _UNIT_WORDS.search(parts[-1])):
        out["city"] = parts.pop()
    if len(parts) > 1 and re.search(r"(?i)\b(?:department|dept|unit|division|section|service|program(?:me)?)\b",
                                    parts[0]):
        out["department"] = parts.pop(0)
    out["institution"] = ", ".join(parts)
    return out


def _inline_numbered_affiliations(f: Fields, src: Sources, lead: list[Paragraph]) -> None:
    """"Authors: A; B; C" over "Affiliations: 1 General Education Unit, ...; 2 General Education Department, ...;
    3 ..." (one paragraph, numbered inline, as many affiliations as authors, no marks after the names): the numbers
    follow the author order, so author k has affiliation k. The paragraph is split at its numbers when the split took
    it as one affiliation."""
    lines = [t for text in list(src.first_page) + [p.text for p in lead] for t in text.split("\n") if t.strip()]
    line = next((t for t in lines if _INLINE_AFFS.match(t)), None)
    if line is None or len(f.authors) < 2:
        return
    body = line[_INLINE_AFFS.match(line).end():]
    items = re.split(r"(?:^|\s*;\s*)(\d{1,2})\s*[.)]?\s+(?=[A-Za-z])", body)
    numbered = [(int(items[k]), items[k + 1].strip(" ;.")) for k in range(1, len(items) - 1, 2)]
    if len(numbered) != len(f.authors) or [n for n, _ in numbered] != list(range(1, len(numbered) + 1)):
        return
    for a in f.authors:  # marks after a name: the text links them itself
        if any(re.search(re.escape(a.surname) + r"\s*[,*]?\s*\d", t) for t in lines if t is not line):
            return
    if len(f.affiliations) != len(numbered):
        f.affiliations = {k: _affiliation_parts(t) for k, (k, t) in enumerate(numbered, 1)}
        f.notes.append(f"the inline numbered affiliations split into {len(numbered)}")
    ids = sorted(f.affiliations)
    if any(a.affs and a.affs != [ids[k]] for k, a in enumerate(f.authors)):
        return
    for k, a in enumerate(f.authors):
        a.affs = [ids[k]]
    f.notes.append("authors linked to the inline numbered affiliations in author order")


def _xml_places(f: Fields, meta: Meta) -> None:
    """An affiliation with no state or country ("Sankara Eye Hospital, Shimoga") completed from the submission XML when
    its author's entry there gives the same city and an Indian state code (<state>KA</state> -> Karnataka, India):
    S "Country name must be added at the end if it is not present"."""
    for a in f.authors:
        xa = next((x for x in meta.authors if _norm(x.get("surname", "")) == _norm(a.surname)), None)
        xaff = (xa or {}).get("aff") or {}
        code, city = (xaff.get("state") or "").strip(), (xaff.get("city") or "").strip()
        state = _STATE_CODES_IN.get(code.upper(), "") if re.fullmatch(r"[A-Za-z]{2}", code) else \
            (code if code.casefold() in _STATES_IN else "")
        if not state or (xaff.get("country") or "India").strip().casefold() not in ("india", "in"):
            continue
        for i in a.affs:
            aff = f.affiliations.get(i)
            if not aff or aff.get("state") or aff.get("country") or not aff.get("city"):
                continue
            if _city_key(aff["city"]) != _city_key(city):
                continue
            known = _CITY_STATE_IN.get(_city_key(city))
            if known and known != state:
                continue
            # the same place wherever it is listed (so repeated affiliations stay one)
            for other in f.affiliations.values():
                if not other.get("state") and not other.get("country") and _city_key(other.get("city", "")) == \
                        _city_key(aff["city"]) and _norm(other.get("institution", "")) == _norm(aff.get("institution", "")):
                    other["state"], other["country"] = state, "India"
            f.notes.append(f"affiliation {aff.get('institution') or aff.get('department')!r}: {state}, India from the "
                           f"submission XML (<state>{code}</state>)")


def _xml_corresponding(f: Fields, meta: Meta) -> None:
    """No corresponding author in the text: the one the submission XML marks (is_corresponding)."""
    if any(a.corresponding for a in f.authors) or (f.correspondence or {}).get("surname"):
        return
    xc = next((x for x in meta.authors if x.get("corresponding")), None)
    if xc is None:
        return
    a = next((a for a in f.authors if _norm(a.surname) == _norm(xc.get("surname", ""))), None)
    if a is not None:
        a.corresponding = True
        f.notes.append(f"corresponding author {a.given} {a.surname} from the submission XML")


def rule_front_matter(ms: Manuscript, style: JournalStyle) -> list[Change]:
    """In-place edits to front matter already in the manuscript (rules.RULES entry). The paragraphs
    to insert are in `plan(ms, style).inserts`."""
    if not style.get("front_matter.enabled", True):
        return []
    return list(plan(ms, style).changes)


# --- rendering inserted paragraphs ---------------------------------------------------------

def _run_el(text: str, fmt: dict, template_rpr):
    from .tracked import apply_format  # keeps the schema order of run properties
    r = OxmlElement("w:r")
    rpr = copy.deepcopy(template_rpr) if template_rpr is not None else OxmlElement("w:rPr")
    for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:u", "w:vertAlign", "w:rPrChange", "w:highlight"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    apply_format(rpr, {k: v for k, v in fmt.items() if v})
    if len(rpr):
        r.append(rpr)
    if text == "\t":
        r.append(OxmlElement("w:tab"))
    else:
        t = OxmlElement("w:t")
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
    return r


def render_inserts(ms: Manuscript, inserts: list[NewParagraph], tracked: bool = True, statuses=None,
                   author: str = "Edit OS", run_format=None, align=None, elements: dict | None = None) -> int:
    """Insert accepted NewParagraphs into the document; tracked as inserted paragraphs. Call after
    `tracked.render` (anchors are paragraph ids, which do not move). `run_format(role)` / `align(role)`: the
    journal's font and size for the runs and the paragraph's alignment (None: as the anchor's runs / the default).
    Returns how many were placed; `elements` (when given) collects each inserted w:p by its insert id."""
    statuses = statuses or (Status.ACCEPTED,)
    date = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    ids = count(8000)
    last_after: dict[str, object] = {}
    placed = 0
    for np_ in inserts:
        if np_.status not in statuses:
            continue
        try:
            anchor = ms.by_id(np_.anchor).element
        except KeyError:
            continue
        template = next((r.find(qn("w:rPr")) for r in anchor.iter(qn("w:r")) if r.find(qn("w:rPr")) is not None), None)
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        jc = align(np_.role) if align else None
        if jc:
            el = OxmlElement("w:jc")
            el.set(qn("w:val"), jc)
            ppr.append(el)
        base = (run_format(np_.role) if run_format else None) or {}
        if tracked:
            prpr = OxmlElement("w:rPr")
            mark = OxmlElement("w:ins")
            for k, v in ((qn("w:id"), str(next(ids))), (qn("w:author"), author), (qn("w:date"), date)):
                mark.set(k, v)
            prpr.append(mark)
            ppr.append(prpr)
        p.append(ppr)
        for text, fmt in np_.runs:
            pieces = [x for x in re.split(r"(\t)", text) if x]
            for piece in pieces:
                r = _run_el(piece, {**base, **(fmt or {})}, template)
                if tracked:
                    w = OxmlElement("w:ins")
                    for k, v in ((qn("w:id"), str(next(ids))), (qn("w:author"), author), (qn("w:date"), date)):
                        w.set(k, v)
                    w.append(r)
                    p.append(w)
                else:
                    p.append(r)
        if np_.where == "before":
            anchor.addprevious(p)
        else:
            prev = last_after.get(np_.anchor, anchor)
            prev.addnext(p)
            last_after[np_.anchor] = p
        if elements is not None:
            elements[np_.id] = p
        placed += 1
    return placed
