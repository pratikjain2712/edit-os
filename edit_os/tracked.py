"""Render change records into the .docx as Word tracked changes, or as a clean copy.

Formatting is kept by splitting the existing runs at change boundaries: deleted
text stays in its original run (wrapped in <w:del>), and inserted text copies the
formatting of the run it replaces or follows.

A paragraph break in inserted text (`PARAGRAPH_BREAK`, U+2029) starts a new paragraph there: the
text after it, and the rest of the original paragraph, move into new <w:p> elements. In tracked
mode the new paragraph marks are tracked insertions, so rejecting them joins the text back.

Formatting changes (`Change.format`) are applied first, while offsets still match the
original text: each run in the span gets the new run properties and, in tracked mode, a
<w:rPrChange> holding its previous properties. Text changes are applied afterwards.
"""

from __future__ import annotations

import copy
import re
from datetime import datetime, timezone

from docx.oxml import OxmlElement
from lxml import etree
from docx.oxml.ns import qn

from .changes import Change, ChangeSet
from .docmodel import SYMBOL_CODES, Manuscript, explode_runs, paragraph_text, run_text, text_runs


PARAGRAPH_BREAK = "\u2029"  # in Change.new: a new paragraph starts here


class ChangeConflict(ValueError):
    """A change no longer matches the text it was made against, or overlaps another."""


class _Ids:
    def __init__(self, start: int = 9000):
        self.n = start

    def next(self) -> str:
        self.n += 1
        return str(self.n)


def _split_run(run, offset: int):
    """Split a single-w:t run at `offset`; returns the right-hand run."""
    t = run.find(qn("w:t"))
    text = t.text or ""
    right = copy.deepcopy(run)
    t.text = text[:offset]
    right.find(qn("w:t")).text = text[offset:]
    for el in (t, right.find(qn("w:t"))):
        el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.addnext(right)
    return right


def _spans(p_el):
    """(run, start, end) for each text run of an exploded paragraph."""
    out, pos = [], 0
    for r in text_runs(p_el):
        n = len(run_text(r))
        out.append((r, pos, pos + n))
        pos += n
    return out


def _cut(p_el, offset: int) -> None:
    """Make sure a run boundary falls exactly at `offset`."""
    for r, s, e in _spans(p_el):
        if s < offset < e:
            if r.find(qn("w:t")) is None:
                raise ChangeConflict(f"cannot split non-text content at offset {offset}")
            _split_run(r, offset - s)
            return


# Child order inside <w:rPr> (CT_RPr); Word rejects out-of-order properties.
_RPR_ORDER = [qn(f"w:{t}") for t in (
    "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
    "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd "
    "fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange").split()]
_TOGGLES = {"italic": "w:i", "bold": "w:b", "small_caps": "w:smallCaps", "caps": "w:caps"}
_VERT = {"superscript": "superscript", "subscript": "subscript"}
# Properties a change clears (False): the elements removed for each.
_CLEARED = {"size": ("w:sz", "w:szCs"), "highlight": ("w:highlight",)}


def _rpr_insert(rpr, el) -> None:
    rank = _RPR_ORDER.index(el.tag)
    for child in rpr:
        if child.tag in _RPR_ORDER and _RPR_ORDER.index(child.tag) > rank:
            child.addprevious(el)
            return
    rpr.append(el)


def apply_format(rpr, fmt: dict) -> None:
    """Set (True) or clear (False) run properties on a <w:rPr>, keeping schema order."""
    for prop, on in fmt.items():
        if prop in _TOGGLES:
            tag = qn(_TOGGLES[prop])
            for el in rpr.findall(tag):
                rpr.remove(el)
            if on:
                _rpr_insert(rpr, OxmlElement(_TOGGLES[prop]))
        elif prop in _VERT:
            va = rpr.find(qn("w:vertAlign"))
            if on:
                if va is None:
                    va = OxmlElement("w:vertAlign")
                    _rpr_insert(rpr, va)
                va.set(qn("w:val"), _VERT[prop])
            elif va is not None and va.get(qn("w:val")) == _VERT[prop]:
                rpr.remove(va)
        elif prop == "underline":
            for el in rpr.findall(qn("w:u")):
                rpr.remove(el)
            if on:
                u = OxmlElement("w:u")
                u.set(qn("w:val"), "single")
                _rpr_insert(rpr, u)
        elif prop == "color":  # False: drop the direct colour; "auto": automatic over a coloured style
            for el in rpr.findall(qn("w:color")):
                rpr.remove(el)
            if on == "auto":
                c = OxmlElement("w:color")
                c.set(qn("w:val"), "auto")
                _rpr_insert(rpr, c)
        elif prop == "size" and on is not False:  # points -> half-points, the complex-script size too
            for tag in ("w:sz", "w:szCs"):
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)
                el = OxmlElement(tag)
                el.set(qn("w:val"), str(int(round(float(on) * 2))))
                _rpr_insert(rpr, el)
        elif prop == "font":  # the Latin and complex-script font; an East Asian font stays
            f = rpr.find(qn("w:rFonts"))
            if on is False:
                if f is not None:
                    rpr.remove(f)
                continue
            if f is None:
                f = OxmlElement("w:rFonts")
                _rpr_insert(rpr, f)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
                f.attrib.pop(qn(attr), None)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                f.set(qn(attr), on)
        elif prop in ("align", "symbol", "indent", "para_style", "list_number"):
            continue  # paragraph properties (`_align`, `_clear_indent`); the Symbol font character (`_new_run`, `_to_symbol`)
        elif prop in _CLEARED:
            for tag in _CLEARED[prop]:
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)
            if prop == "highlight":  # coloured run shading looks the same as a highlight
                for el in rpr.findall(qn("w:shd")):
                    if (el.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF"):
                        rpr.remove(el)
        else:
            raise ValueError(f"unknown format property {prop!r}")


def _rpr_of(run):
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run.insert(0, rpr)
    return rpr


def _format_run(run, fmt: dict, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    rpr = _rpr_of(run)
    before = copy.deepcopy(rpr)
    for el in before.findall(qn("w:rPrChange")):
        before.remove(el)
    apply_format(rpr, fmt)
    after = copy.deepcopy(rpr)
    for el in after.findall(qn("w:rPrChange")):
        after.remove(el)
    if etree.tostring(after) == etree.tostring(before):
        if not len(rpr):
            run.remove(rpr)
        return
    if tracked and rpr.find(qn("w:rPrChange")) is None:  # an earlier revision keeps its baseline
        ch = OxmlElement("w:rPrChange")
        ch.set(qn("w:id"), ids.next())
        ch.set(qn("w:author"), author)
        ch.set(qn("w:date"), date)
        ch.append(before)
        rpr.append(ch)


def _sym(ch: str):
    """<w:sym w:font="Symbol" w:char="F0B1"/> for a character of docmodel.SYMBOL_CODES (Word's inserted symbol)."""
    el = OxmlElement("w:sym")
    el.set(qn("w:font"), "Symbol")
    el.set(qn("w:char"), f"F0{SYMBOL_CODES[ch]:02X}")
    return el


def _new_run(template, text: str, fmt: dict | None = None):
    r = OxmlElement("w:r")
    if template is not None and template.find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(template.find(qn("w:rPr")))
        for el in rpr.findall(qn("w:rPrChange")):  # inserted text carries no format revision
            rpr.remove(el)
        r.append(rpr)
    if fmt:
        apply_format(_rpr_of(r), fmt)
    if fmt and fmt.get("symbol") and text in SYMBOL_CODES:
        r.append(_sym(text))  # formatting.symbol_font: the character in the Symbol font, as Word inserts it
        return r
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def _to_symbol(run, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """A one-character text run ("±") -> the same character as a Symbol-font <w:sym>: tracked, the typed character
    is deleted and the symbol inserted (Word has no format revision for it)."""
    t = run.find(qn("w:t"))
    if t is None or (t.text or "") not in SYMBOL_CODES:
        return
    new = _new_run(run, t.text, {"symbol": True})
    if tracked:
        run.addnext(new)
        _to_deleted(run)
        _wrap("w:del", run, ids, author, date)
        _wrap("w:ins", new, ids, author, date)
    else:
        run.addnext(new)
        run.getparent().remove(run)


def _wrap(tag: str, run, ids: _Ids, author: str, date: str):
    w = OxmlElement(tag)
    w.set(qn("w:id"), ids.next())
    w.set(qn("w:author"), author)
    w.set(qn("w:date"), date)
    run.addprevious(w)
    w.append(run)
    return w


def _to_deleted(run) -> None:
    for t in run.findall(qn("w:t")):
        t.tag = qn("w:delText")


def _apply_paragraph(p_el, changes: list[Change], tracked: bool, ids: _Ids, date: str) -> None:
    text = paragraph_text(p_el)
    formats = [c for c in changes if c.format_only]
    changes = [c for c in changes if not c.format_only]
    # Text changes must not overlap each other, nor formatting changes of the same property each
    # other; a formatting change may cover text that a text change replaces.
    for c in changes + formats:
        if text[c.start:c.end] != c.old:
            raise ChangeConflict(
                f"{c.id}: expected {c.old!r} at {c.start}-{c.end} in {c.para_id}, found {text[c.start:c.end]!r}")
    last_end = -1
    for c in changes:
        if c.start < last_end:
            raise ChangeConflict(f"{c.id}: overlaps an earlier change in {c.para_id}")
        last_end = c.end
    for i, c in enumerate(formats):
        if any(d.start < c.end and c.start < d.end and set(d.format) & set(c.format) for d in formats[:i]):
            raise ChangeConflict(f"{c.id}: overlaps an earlier formatting change in {c.para_id}")

    explode_runs(p_el)
    for c in formats + changes:
        _cut(p_el, c.start)
        _cut(p_el, c.end)

    # Formatting first: it does not move any offsets (the Symbol-font characters last: they replace their run).
    formats.sort(key=lambda c: bool(c.format.get("symbol")))
    for c in formats:
        if "align" in c.format:
            _align(p_el, c.format["align"], tracked, ids, c.author, date)
        if "indent" in c.format:
            _clear_indent(p_el, tracked, ids, c.author, date)
        if "para_style" in c.format:
            _para_style(p_el, c.format["para_style"], tracked, ids, c.author, date)
        if c.format.get("symbol"):
            for r, s, e in _spans(p_el):
                if s >= c.start and e <= c.end and e > s:
                    _to_symbol(r, tracked, ids, c.author, date)
            continue
        if c.format.get("list_number") is False:
            _unnumber(p_el, tracked, ids, c.author, date)
        if not set(c.format) - {"align", "indent", "para_style", "list_number"}:
            continue
        for r, s, e in _spans(p_el):
            if s >= c.start and e <= c.end and e > s:
                _format_run(r, c.format, tracked, ids, c.author, date)

    # Work right to left so earlier offsets stay valid; at a shared start, the
    # replacement goes before a pure insertion so the insertion lands in front of it.
    # An insertion that opens new paragraphs goes first at its point, so a plain insertion at
    # the same point (a citation at the end of the text) still ends the original paragraph.
    for c in sorted(changes, key=lambda c: (c.start, c.end, c.new.startswith(PARAGRAPH_BREAK)), reverse=True):
        spans = _spans(p_el)
        doomed = [r for r, s, e in spans if s >= c.start and e <= c.end and e > s]
        before = [r for r, s, e in spans if e <= c.start]
        after = [r for r, s, e in spans if s >= c.end]
        template = doomed[0] if doomed else (before[-1] if before else (after[0] if after else None))

        anchor = None  # element the inserted run goes after
        for r in doomed:
            if tracked:
                _to_deleted(r)
                anchor = _wrap("w:del", r, ids, c.author, date)
            else:
                anchor = r.getprevious()
                r.getparent().remove(r)

        if not c.new:
            continue
        first, *rest = c.new.split(PARAGRAPH_BREAK)
        runs = [_new_run(template, text, fmt) for text, fmt in c.pieces(0, first)] or [_new_run(template, "", c.segment(0))]
        ins = runs[0]
        if doomed and tracked:
            anchor.addnext(ins)
        elif after:
            after[0].addprevious(ins)
        elif before:
            before[-1].addnext(ins)
        else:
            p_el.append(ins)
        if tracked:
            _wrap("w:ins", ins, ids, c.author, date)
        last = ins
        for run in runs[1:]:  # [length, format] pieces of the first segment
            (last.getparent() if tracked else last).addnext(run)
            if tracked:
                _wrap("w:ins", run, ids, c.author, date)
            last = run
        ins = last
        if rest:
            _break_paragraph(p_el, ins, after, rest, template, c, tracked, ids, date)
        if not first:  # only there to mark the position of the break
            top = ins.getparent() if tracked else ins
            top.getparent().remove(top)


_PPR = qn("w:pPr")
# Child order inside <w:pPr> (CT_PPr), up to the alignment and what follows it.
_PPR_ORDER = [qn(f"w:{t}") for t in (
    "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
    "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
    "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange").split()]
_JC_SAME = {"start": "left", "end": "right", "distribute": "both"}


def _ppr_of(p_el):
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    return ppr


def _ppr_insert(ppr, el) -> None:
    rank = _PPR_ORDER.index(el.tag)
    for child in ppr:
        if child.tag in _PPR_ORDER and _PPR_ORDER.index(child.tag) > rank:
            child.addprevious(el)
            return
    ppr.append(el)


def _align(p_el, value: str, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """Set the paragraph's alignment (<w:jc>); tracked as a paragraph property change (<w:pPrChange>)."""
    ppr = _ppr_of(p_el)
    jc = ppr.find(qn("w:jc"))
    have = jc.get(qn("w:val")) if jc is not None else None
    if have is not None and _JC_SAME.get(have, have) == value:
        return
    before = _ppr_snapshot(ppr)
    if jc is None:
        jc = OxmlElement("w:jc")
        _ppr_insert(ppr, jc)
    jc.set(qn("w:val"), value)
    _track_ppr(ppr, before, tracked, ids, author, date)


def _para_style(p_el, value: str, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """Set the paragraph's style (<w:pStyle>); tracked as a paragraph property change (<w:pPrChange>)."""
    ppr = _ppr_of(p_el)
    ps = ppr.find(qn("w:pStyle"))
    if ps is not None and ps.get(qn("w:val")) == value:
        return
    before = _ppr_snapshot(ppr)
    if ps is None:
        ps = OxmlElement("w:pStyle")
        _ppr_insert(ppr, ps)
    ps.set(qn("w:val"), value)
    _track_ppr(ppr, before, tracked, ids, author, date)


def _new_paragraph_props(p_el, props: dict, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """Paragraph properties a change gives a paragraph it opens (`Change.paragraph_props`): alignment and style."""
    if props.get("align"):
        _align(p_el, props["align"], tracked, ids, author, date)
    if props.get("para_style"):
        _para_style(p_el, props["para_style"], tracked, ids, author, date)


def _unnumber(p_el, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """Take the paragraph out of Word's automatic numbering (a heading the author numbered with a list); tracked as a
    paragraph property change."""
    ppr = _ppr_of(p_el)
    num = ppr.find(qn("w:numPr"))
    if num is None:
        return
    before = _ppr_snapshot(ppr)
    ppr.remove(num)
    _track_ppr(ppr, before, tracked, ids, author, date)


def _ppr_snapshot(ppr):
    """CT_PPrBase copy of a <w:pPr>: its properties without rPr / sectPr / pPrChange (a pPrChange's content)."""
    before = OxmlElement("w:pPr")
    for child in ppr:
        if child.tag not in _PPR_TAIL:
            before.append(copy.deepcopy(child))
    return before


def _track_ppr(ppr, before, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    if tracked and ppr.find(qn("w:pPrChange")) is None:  # an earlier change keeps the original properties
        ch = OxmlElement("w:pPrChange")
        ch.set(qn("w:id"), ids.next())
        ch.set(qn("w:author"), author)
        ch.set(qn("w:date"), date)
        ch.append(before)
        ppr.append(ch)


_IND_LEFT = ("w:left", "w:start", "w:leftChars", "w:startChars", "w:firstLine", "w:firstLineChars", "w:hanging",
             "w:hangingChars")


def _clear_indent(p_el, tracked: bool, ids: "_Ids", author: str, date: str) -> None:
    """Set the paragraph's left and first-line indentation to zero (<w:ind w:left="0" w:firstLine="0"/>, over any the
    style gives; the right indentation stays); tracked as a paragraph property change (<w:pPrChange>)."""
    ppr = _ppr_of(p_el)
    ind = ppr.find(qn("w:ind"))
    before = _ppr_snapshot(ppr)
    if ind is None:
        ind = OxmlElement("w:ind")
        _ppr_insert(ppr, ind)
    for a in _IND_LEFT:
        ind.attrib.pop(qn(a), None)
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:firstLine"), "0")
    _track_ppr(ppr, before, tracked, ids, author, date)


# Child order inside <w:pPr> (CT_PPr); the paragraph mark's <w:rPr> goes before sectPr / pPrChange.
_PPR_TAIL = (qn("w:rPr"), qn("w:sectPr"), qn("w:pPrChange"))


def _top(p_el, el):
    while el.getparent() is not p_el:
        el = el.getparent()
    return el


def _mark_inserted(p_el, ids: _Ids, author: str, date: str) -> None:
    """Record the paragraph mark of `p_el` as a tracked insertion."""
    ppr = p_el.find(_PPR)
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        tail = [c for c in ppr if c.tag in _PPR_TAIL[1:]]
        if tail:
            tail[0].addprevious(rpr)
        else:
            ppr.append(rpr)
    if rpr.find(qn("w:ins")) is not None:  # already a new paragraph
        return
    mark = OxmlElement("w:ins")
    mark.set(qn("w:id"), ids.next())
    mark.set(qn("w:author"), author)
    mark.set(qn("w:date"), date)
    rpr.insert(0, mark)


def _break_paragraph(p_el, ins, after, segments, template, c: Change, tracked, ids, date) -> None:
    """Split `p_el` after the inserted run `ins`: each segment becomes a new paragraph, and
    whatever followed the insertion point moves into the last one."""
    tail = []
    if after:
        start = _top(p_el, after[0])
        if start is _top(p_el, ins):
            raise ChangeConflict(f"{c.id}: cannot start a paragraph inside a hyperlink or field")
        tail = [start] + [s for s in start.itersiblings()]
    ppr = p_el.find(_PPR)
    new_paras = []
    for i, seg in enumerate(segments):
        p = OxmlElement("w:p")
        last = i == len(segments) - 1
        if ppr is not None:
            pp = copy.deepcopy(ppr)
            for tag in ("w:pPrChange",) + (() if last else ("w:sectPr",)):
                for el in pp.findall(qn(tag)):
                    pp.remove(el)
            if not (last and tail):  # new text is not a list item, whatever the paragraph was
                for el in pp.findall(qn("w:numPr")):
                    pp.remove(el)
            p.append(pp)
        for text, fmt in c.pieces(i + 1, seg):
            run = _new_run(template, text, fmt)
            p.append(run)
            if tracked:
                _wrap("w:ins", run, ids, c.author, date)
        if last:
            for el in tail:
                p.append(el)
        props = c.paragraph_props(i + 1)
        if props:  # a new paragraph's own properties are part of its insertion; the last one holds the old mark
            _new_paragraph_props(p, props, tracked and last and bool(tail), ids, c.author, date)
        new_paras.append(p)
    # The original paragraph mark now ends the last part: the first part's mark, and every
    # mark but the last, are new.
    if ppr is not None:
        for el in ppr.findall(qn("w:sectPr")):
            ppr.remove(el)
    if tracked:
        _mark_inserted(p_el, ids, c.author, date)
        for p in new_paras[:-1]:
            _mark_inserted(p, ids, c.author, date)
    anchor = p_el
    for p in new_paras:
        anchor.addnext(p)
        anchor = p


# --- whole paragraphs and tables deleted ----------------------------------------------------
#
# A change deleting all of a paragraph's text (start 0, end len, new "") deletes the paragraph: in the
# tracked copy its paragraph mark is a tracked deletion (<w:pPr><w:rPr><w:del/></w:rPr></w:pPr>), so
# accepting it leaves no empty line; in the clean copy the paragraph goes. Not a paragraph holding an image,
# object, field or note reference, a section break, the last paragraph of a table cell or of the document,
# nor one whose removal would join two tables. Empty paragraphs between two deleted ones go with them, and a
# table whose every paragraph is deleted (or empty) is deleted as rows (<w:trPr><w:del/></w:trPr>).
#
# The test is on what the changes leave in the paragraph, not on one change: a paragraph whose text is
# deleted by several changes, or replaced by text that starts a new paragraph ("\u2029<H2>Authors'
# contributions..."), is deleted the same way. A heading tag inserted at the start of such a paragraph
# (another rule tagging the heading that this one deletes) has nothing left to tag: it is dropped, so no
# paragraph holding only "<H2>" is left behind.

_KEEP = tuple(qn(t) for t in ("w:drawing", "w:pict", "w:object", "w:sym", "w:footnoteReference",
                              "w:endnoteReference", "w:fldSimple", "w:fldChar", "w:tbl", "w:sectPr")) + (
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent",)
_BLOCKS = (qn("w:p"), qn("w:tbl"), qn("w:sectPr"))


_TAG_ONLY = re.compile(r"(?:<H\d>)+")


def _tag_only(c: Change) -> bool:
    """A change that only puts a heading tag in front of the text (it replaces nothing but spaces)."""
    return not c.format_only and not c.old.strip() and bool(_TAG_ONLY.fullmatch(c.new))


def _left_in_paragraph(text: str, edits: list[Change]) -> str:
    """The text the (non-overlapping) `edits` leave in the paragraph itself: up to the first new paragraph.
    Order at one point as `_apply_paragraph` builds it: a plain insertion, then one opening new paragraphs,
    then a replacement starting there."""
    out, pos = [], 0
    for c in sorted(edits, key=lambda c: (c.start, c.end > c.start, c.new.startswith(PARAGRAPH_BREAK))):
        out.append(text[pos:c.start] if c.start > pos else "")
        out.append(c.new)
        pos = max(pos, c.end)
    out.append(text[pos:])
    return "".join(out).split(PARAGRAPH_BREAK)[0]


def _emptying(text: str, cs: list[Change]) -> tuple[list[Change], bool]:
    """(the changes to apply, whether they leave the paragraph empty). A paragraph with text that the
    changes empty loses its heading tag insertions too: a tag on no text is never wanted."""
    edits = [c for c in cs if not c.format_only]
    if not text or not any(c.end > c.start for c in edits):
        return cs, False
    rest = [c for c in edits if not _tag_only(c)]
    left = _left_in_paragraph(text, rest)
    if left.strip() or (left and not text.strip()):  # a run of spaces tidied is not a paragraph deleted
        return cs, False
    return [c for c in cs if c.format_only or not _tag_only(c)], True


def _removable(el) -> bool:
    return not any(True for x in el.iter(*_KEEP) if x is not el)


def _neighbour(el, gone: set, step: str):
    sib = getattr(el, step)()
    while sib is not None and (sib in gone or sib.tag not in _BLOCKS):
        sib = getattr(sib, step)()
    return sib


def _mark(parent, tag: str, ids: _Ids, author: str, date: str, first: bool) -> None:
    mark = OxmlElement(tag)
    mark.set(qn("w:id"), ids.next())
    mark.set(qn("w:author"), author)
    mark.set(qn("w:date"), date)
    if first:  # CT_ParaRPr: w:ins, w:del come first
        ins = parent.find(qn("w:ins"))
        if ins is not None:
            ins.addnext(mark)
        else:
            parent.insert(0, mark)
    else:  # CT_TrPr: before w:trPrChange
        change = parent.find(qn("w:trPrChange"))
        if change is not None:
            change.addprevious(mark)
        else:
            parent.append(mark)


def _mark_paragraph_deleted(p_el, ids: _Ids, author: str, date: str) -> None:
    ppr = p_el.find(_PPR)
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        tail = [c for c in ppr if c.tag in _PPR_TAIL[1:]]
        if tail:
            tail[0].addprevious(rpr)
        else:
            ppr.append(rpr)
    if rpr.find(qn("w:del")) is None:
        _mark(rpr, "w:del", ids, author, date, first=True)


def _mark_table_deleted(tbl, ids: _Ids, author: str, date: str) -> None:
    for tr in tbl.findall(qn("w:tr")):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            prex = tr.find(qn("w:tblPrEx"))
            if prex is not None:
                prex.addnext(trpr)
            else:
                tr.insert(0, trpr)
        if trpr.find(qn("w:del")) is None:
            _mark(trpr, "w:del", ids, author, date, first=False)


def _delete_blocks(ms: Manuscript, deleted: dict, tracked: bool, ids: _Ids, date: str) -> None:
    """`deleted`: paragraph element -> author of the change deleting all its text."""
    if not deleted:
        return
    gone: set = set()
    by_el = {p.element: p for p in ms.paragraphs}
    # Tables: every paragraph deleted or empty, at least one deleted, no nested table or object.
    for tbl in ms.document.element.body.iter(qn("w:tbl")):
        if any(a.tag == qn("w:tbl") for a in tbl.iterancestors()) or tbl.find(f".//{qn('w:tbl')}") is not None:
            continue
        paras = list(tbl.iter(qn("w:p")))
        hits = [p for p in paras if p in deleted]
        if hits and all(p in deleted or (not paragraph_text(p).strip() and _removable(p)) for p in paras) \
                and all(_removable(p) for p in hits):
            gone.add(tbl)
            if tracked:
                _mark_table_deleted(tbl, ids, deleted[hits[0]], date)
            else:
                ms.removed.append(tbl)
    order = [p.element for p in ms.paragraphs]
    for i, p_el in enumerate(order):
        author = deleted.get(p_el)
        if author is None and not paragraph_text(p_el).strip():
            # An empty paragraph between two deleted ones (blank lines inside a deleted cover letter).
            prev = _neighbour(p_el, set(), "getprevious")
            nxt = _neighbour(p_el, set(), "getnext")
            author = deleted.get(prev) if prev in deleted and nxt in deleted else None
        if author is None or not _removable(p_el) or any(a in gone for a in p_el.iterancestors()):
            continue
        parent = p_el.getparent()
        if parent is None or by_el.get(p_el) is None:
            continue
        prev, nxt = _neighbour(p_el, gone, "getprevious"), _neighbour(p_el, gone, "getnext")
        if nxt is None or nxt.tag == qn("w:sectPr"):
            continue  # the last paragraph of a cell or of the document
        if prev is not None and prev.tag == qn("w:tbl") and nxt.tag == qn("w:tbl"):
            continue  # the two tables would join
        gone.add(p_el)
        if tracked:
            _mark_paragraph_deleted(p_el, ids, author, date)
        else:
            ms.removed.append(p_el)


# Parts holding text runs: the body, headers and footers, footnotes and endnotes.
_STORY_TYPES = ("document.main+xml", "header+xml", "footer+xml", "footnotes+xml", "endnotes+xml")


def strip_highlighting(ms: Manuscript) -> int:
    """Remove the author's text highlighting (Word highlight colours and coloured run shading) from
    every story of the document; returns how many marks were removed.

    For the clean copy only: authors mark revised text in yellow or cyan for the reviewers, and the
    delivered file must not carry it. The tracked copy keeps it, so editors still see what was marked.
    """
    n = 0
    for part in ms.document.part.package.iter_parts():
        if not part.content_type.endswith(_STORY_TYPES):
            continue
        root = getattr(part, "_element", None)
        parsed = root is None
        if parsed:
            root = etree.fromstring(part.blob)
        for rpr in list(root.iter(qn("w:rPr"))):
            for el in rpr.findall(qn("w:highlight")):
                rpr.remove(el)
                n += 1
            for el in rpr.findall(qn("w:shd")):
                if (el.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF"):
                    rpr.remove(el)
                    n += 1
        if parsed:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return n


# Child order inside <w:pPr> before <w:numPr> (CT_PPr).
_BEFORE_NUMPR = tuple(qn(t) for t in ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
                                      "w:widowControl"))


def _numbering_part(doc_part):
    """The document's numbering part, added (empty) when the file has none (python-docx cannot create one)."""
    from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.oxml import parse_xml
    from docx.parts.numbering import NumberingPart

    try:
        return doc_part.part_related_by(RT.NUMBERING)
    except KeyError:
        el = parse_xml('<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING, el, doc_part.package)
        doc_part.relate_to(part, RT.NUMBERING)
        return part


def _list_num_id(ms: Manuscript) -> str:
    """A decimal "1." list of its own in the document's numbering (one per render): the reference list whose
    typed labels became Word's automatic numbering starts at 1."""
    num_id = getattr(ms, "_label_num_id", None)
    if num_id:
        return num_id
    numbering = _numbering_part(ms.document.part).element
    abstract = [int(x) for x in numbering.xpath("./w:abstractNum/@w:abstractNumId") if x.isdigit()]
    nums = [int(x) for x in numbering.xpath("./w:num/@w:numId") if x.isdigit()]
    a_id, num_id = str(max(abstract, default=-1) + 1), str(max(nums, default=0) + 1)
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    a = etree.fromstring(
        f'<w:abstractNum xmlns:w="{w}" w:abstractNumId="{a_id}"><w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(a)
    else:
        numbering.append(a)
    numbering.append(etree.fromstring(
        f'<w:num xmlns:w="{w}" w:numId="{num_id}"><w:abstractNumId w:val="{a_id}"/></w:num>'))
    ms._label_num_id = num_id
    return num_id


def _number_paragraph(ms: Manuscript, p_el, tracked: bool, ids: _Ids, author: str, date: str) -> None:
    """Make `p_el` an item of the document's own "1." list (a <w:pPrChange> keeps the old paragraph properties)."""
    ppr = p_el.find(_PPR)
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    if ppr.find(qn("w:numPr")) is not None:
        return
    before = copy.deepcopy(ppr)
    for el in list(before):
        if el.tag in _PPR_TAIL:
            before.remove(el)
    num_pr = OxmlElement("w:numPr")
    ilvl, nid = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    ilvl.set(qn("w:val"), "0")
    nid.set(qn("w:val"), _list_num_id(ms))
    num_pr.append(ilvl)
    num_pr.append(nid)
    lead = [c for c in ppr if c.tag in _BEFORE_NUMPR]
    if lead:
        lead[-1].addnext(num_pr)
    else:
        ppr.insert(0, num_pr)
    if tracked and ppr.find(qn("w:pPrChange")) is None:
        ch = OxmlElement("w:pPrChange")
        ch.set(qn("w:id"), ids.next())
        ch.set(qn("w:author"), author)
        ch.set(qn("w:date"), date)
        ch.append(before)
        ppr.append(ch)


def render(ms: Manuscript, changes: ChangeSet, tracked: bool = True, statuses=None,
           on_conflict=None) -> Manuscript:
    """Apply changes to the manuscript in place. Use a freshly loaded Manuscript per output.

    A `ChangeConflict` stops the render, unless `on_conflict(para_id, error)` is given: then that
    paragraph's changes are skipped (its text is left as it was) and rendering goes on.
    """
    from .changes import Status
    statuses = statuses or (Status.ACCEPTED,)
    # Revision ids must be unique in the document: start above any the file already holds.
    taken = [int(v) for el in ms.document.element.body.iter() if (v := el.get(qn("w:id"))) and v.isdigit()]
    ids = _Ids(max([9000] + taken))
    date = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    deleted = {}  # paragraphs whose whole text the changes delete -> the author of that change
    for p in ms.paragraphs:
        cs = changes.for_paragraph(p.id, statuses)
        if not cs:
            continue
        text = paragraph_text(p.element)
        cs, whole = _emptying(text, cs)
        if on_conflict is None:
            _apply_paragraph(p.element, cs, tracked, ids, date)
        else:
            backup = copy.deepcopy(p.element)
            try:
                _apply_paragraph(p.element, cs, tracked, ids, date)
            except ChangeConflict as e:
                p.element.getparent().replace(p.element, backup)
                p.element = backup
                on_conflict(p.id, e)
                continue
        if whole:
            deleted[p.element] = next(c.author for c in cs if not c.format_only and c.end > c.start)
        else:
            numbered = next((c for c in cs if (c.format or {}).get("list_number")), None)
            if numbered is not None:  # a typed label deleted: Word's automatic numbering instead
                _number_paragraph(ms, p.element, tracked, ids, numbered.author, date)
    _delete_blocks(ms, deleted, tracked, ids, date)
    return ms
