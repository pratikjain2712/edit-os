"""Template leftovers: what a submission template or a converter left in the manuscript that is not the article.

The editors delete these by hand (tracked); here each is a suggestion like any other.

Text paragraphs (`Change` records deleting the paragraph, reason `template_leftovers`; rules.RULES runs this right
after the front matter, and no other rule's edit stays on a paragraph it deletes):
  - template labels in the front of the article: "Abstract Page", "Main article" (jehp_1738, 1153);
  - the article title typed again above the abstract or the introduction (jehp_1738, 1771, IJO_2733) when the front
    matter sets the title from the title page / XML;
  - a paragraph holding only "." (jehp_1588); Word's "Top of Form" / "Bottom of Form" markers (jehp_997);
  - statement labels with nothing under them right above the reference list: "Ethical approve", "Using AI:" (1280);
  - a list of table titles after the reference list whose tables carry the same captions (IJO_2733);
  - a block of heading lines typed again a little further on, a figure drawn with shapes between (jehp_1793).

Objects (`BlockDelete` records, rendered after the text changes):
  - floating text boxes that hold only a template word: "Text", "Abstract Page" (jehp_1738);
  - pictures after the reference list that repeat a picture of the article body (jehp_997).

    rule_leftovers(ms, style) -> [Change]
    plan(ms, style)           -> [BlockDelete]
    render(ms, inserts, statuses, tracked=..., author=..., date=...)
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field

from docx.oxml.ns import qn

from . import provenance, references
from .changes import Change, Layer, Source, Status
from .docmodel import Manuscript, Paragraph, has_picture
from .frontmatter import NewParagraph
from .style import JournalStyle

REASON = "template_leftovers"
DELETE = "template_leftovers_objects"

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"

# The first main heading of the body: what comes before it is the title page and the abstract.
_BODY_START = re.compile(r"^\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?(?:introduction|background|case\s+(?:report|presentation|"
                         r"description))\s*[:.]?\s*$", re.I)
_BODY_START_LATE = re.compile(r"^\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?(?:materials?\s+and\s+methods?|methods?|methodology)"
                              r"\s*[:.]?\s*$", re.I)
_FRONT_LABEL = re.compile(r"^\s*(?:abstract\s+page|main\s+(?:article|text|manuscript|document|file)|blinded\s+manuscript|"
                          r"manuscript\s+(?:file|text|body))\s*[:.\-]?\s*$", re.I)
_LONE_MARK = re.compile(r"^\s*[.,;:]\s*$")
_FORM_MARK = re.compile(r"^\s*(?:top|bottom)\s+of\s+form\s*$", re.I)
# A statement label of the template with no statement under it.
_EMPTY_LABEL = re.compile(r"^\s*(?:<H\d>)?\s*(?:ethic\w*(?:\s+(?:approv\w*|statement|clearance|considerations?|declarations?))?|"
                          r"(?:using|use\s+of)\s+(?:ai|artificial\s+intelligence|generative\s+ai)(?:\s+tools?)?|"
                          r"(?:generative\s+)?ai\s+(?:use|usage|statement|declaration|disclosure)|"
                          r"declarations?(?:\s+of\s+(?:ai|generative\s+ai)(?:\s+use)?)?|trial\s+registration|"
                          r"clinical\s+trial\s+(?:number|registration))\s*[:.\-–]?\s*$", re.I)
# Text boxes holding only a template word.
_BOX_WORDS = {"text", "abstract page", "main article", "title", "insert text here", "type text", "click here to enter text",
              "text box", "textbox"}


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", re.sub(r"<H\d>", "", text, flags=re.I).lower())


def in_textbox(p: Paragraph) -> bool:
    return any(a.tag == f"{_W}txbxContent" for a in p.element.iterancestors())


def body_start(ms: Manuscript) -> int:
    """Index of the first main heading of the body (len(paragraphs) when there is none)."""
    end = references.find_heading(ms)
    for rx in (_BODY_START, _BODY_START_LATE):
        for i, p in enumerate(ms.paragraphs[:end]):
            if not p.in_table and not in_textbox(p) and rx.match(p.text):
                return i
    return len(ms.paragraphs)


def _change(p: Paragraph, why: str) -> Change:
    from .rules import _change as make

    c = make(p, 0, len(p.text), "", Layer.STRUCTURE, REASON)
    c.provenance = provenance.make("rule", "Template leftover", why, rule=REASON)
    return c


# --- text paragraphs --------------------------------------------------------------------------

def _title_texts(ms: Manuscript, style: JournalStyle) -> tuple[set[str], set[str]]:
    """(normalised titles the front matter sets, ids of the paragraphs it edits or deletes itself)."""
    from . import frontmatter

    if not style.get("front_matter.enabled", True):
        return set(), set()
    try:
        plan = frontmatter.plan(ms, style)
    except Exception:  # the front matter is not this module's business: no title, no deletion
        return set(), set()
    titles = {_norm(n.text) for n in plan.inserts if getattr(n, "role", "") == "title" and n.text.strip()}
    src = getattr(ms, "front_sources", None)
    if titles and src is not None and getattr(src.meta, "title", ""):
        titles.add(_norm(src.meta.title))
    return {t for t in titles if len(t) >= 20}, {c.para_id for c in plan.changes}


def _repeated_headings(ms: Manuscript, start: int, end: int) -> list[Paragraph]:
    """The first copy of a block of heading lines that the text gives again within 25 paragraphs, a figure drawn with
    shapes (or a text box) between the two copies (jehp_1793: "Materials and Methods" / "Study Design and Setting"
    typed among the shapes of Figure 1, then again under it)."""
    paras = ms.paragraphs
    top = [i for i in range(start, end) if not paras[i].in_table and not in_textbox(paras[i])]

    def heading_like(p: Paragraph) -> bool:
        t = re.sub(r"<H\d>", "", p.text).strip()
        return bool(t) and len(t.split()) <= 8 and not t.endswith((".", ":", ",")) and t[:1].isupper()

    out: list[Paragraph] = []
    k = 0
    while k < len(top):
        i = top[k]
        if not heading_like(paras[i]):
            k += 1
            continue
        block = [i]
        m = k + 1
        while m < len(top) and heading_like(paras[top[m]]) and len(block) < 4:
            block.append(top[m])
            m += 1
        keys = [_norm(paras[j].text) for j in block]
        found = False
        for n in range(m, min(len(top), m + 25)):
            cand = top[n:n + len(block)]
            if len(cand) == len(block) and [_norm(paras[j].text) for j in cand] == keys:
                between = paras[block[-1] + 1:cand[0]]
                if any(_shapes(q.element) for q in between) or any(in_textbox(q) for q in between):
                    found = True
                break
        if found:
            out += [paras[j] for j in block]
            k = m
        else:
            k += 1
    return out


def _shapes(el) -> bool:
    """A paragraph anchoring drawn shapes or text boxes (no picture of its own)."""
    return (el.find(f".//{_W}drawing") is not None or el.find(f".//{_W}pict") is not None
            or el.find(f".//{_MC}") is not None) and not has_picture(el)


def _table_captions_near(ms: Manuscript) -> dict[int, list[Paragraph]]:
    """Table number -> its captions that sit right above or below a table (blank lines aside)."""
    from . import rules_figures

    out: dict[int, list[Paragraph]] = {}
    for cap in rules_figures.captions(ms):
        p = ms.by_id(cap.para_id)
        if cap.kind == rules_figures.TABLE and not p.in_table and rules_figures._beside_table(p.element):
            out.setdefault(cap.number, []).append(p)
    return out


def rule_leftovers(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not style.get("leftovers.enabled", True):
        return []
    from . import rules_figures

    paras = ms.paragraphs
    ref = references.find_heading(ms)
    end = ref if ref is not None else len(paras)
    start = body_start(ms)
    out: list[Change] = []
    taken: set[str] = set()

    def drop(p: Paragraph, why: str) -> None:
        if p.id not in taken and p.text.strip():
            taken.add(p.id)
            out.append(_change(p, why))

    titles, front_ids = _title_texts(ms, style)
    earlier: set[str] = set()
    for p in paras[:min(start, end)]:
        if p.in_table or in_textbox(p):
            continue
        key = _norm(p.text)
        if p.id in front_ids:
            earlier.add(key)
            continue
        if _FRONT_LABEL.match(p.text):
            drop(p, f"template label \"{p.text.strip()}\"")
        elif len(key) >= 30 and len(p.text.split()) >= 5 and (key in titles or key in earlier) \
                and not re.search(r"[.!?]\s*$", p.text.strip()):
            drop(p, "the article title typed again (the front matter sets the title)")
        earlier.add(key)
    for p in paras:
        if p.in_table or in_textbox(p):
            continue
        if _FORM_MARK.match(p.text):
            drop(p, "Word's form marker")
        elif _LONE_MARK.match(p.text) and (ref is None or paras.index(p) < ref):
            drop(p, "a paragraph holding only a punctuation mark")
    if ref is not None:  # statement labels with nothing under them, right above the reference list
        j = ref - 1
        while j >= 0 and (not paras[j].text.strip() and not has_picture(paras[j].element)):
            j -= 1
        while j >= 0 and not paras[j].in_table and _EMPTY_LABEL.match(paras[j].text):
            drop(paras[j], f"template label \"{paras[j].text.strip()}\" with no statement")
            j -= 1
            while j >= 0 and not paras[j].text.strip() and not has_picture(paras[j].element):
                j -= 1
        # table titles listed again after the reference list, their tables captioned already
        near = _table_captions_near(ms)
        for cap in rules_figures.captions(ms):
            p = ms.by_id(cap.para_id)
            i = paras.index(p)
            if cap.kind != rules_figures.TABLE or i <= ref or p.in_table or rules_figures._beside_table(p.element):
                continue
            twins = [q for q in near.get(cap.number, []) if q is not p]
            if twins and any(_norm(cap.match.group("rest"))[:40] == _norm(rules_figures.caption_match(q).group("rest"))[:40]
                             for q in twins):
                drop(p, f"Table {cap.number}'s title listed again (the table has its caption)")
    for p in _repeated_headings(ms, start, end):
        drop(p, "heading typed again among the shapes of a drawn figure")
    return out


# --- objects ----------------------------------------------------------------------------------

@dataclass
class BlockDelete(NewParagraph):
    """Objects to delete: text boxes (`boxes`: the paragraph ids of their text) and picture paragraphs (`pictures`)."""
    boxes: list[str] = field(default_factory=list)
    pictures: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {**super().to_dict(), "delete_objects": True, "boxes": self.boxes, "pictures": self.pictures}

    @classmethod
    def from_dict(cls, d: dict) -> "BlockDelete":
        return cls(id=d["id"], anchor=d["anchor"], where=d.get("where", "before"),
                   runs=[(t, f) for t, f in d.get("runs") or []], role=d.get("role", "delete"),
                   status=Status(d.get("status", "suggested")), source=Source(d.get("source", "rule")),
                   reason=d.get("reason", DELETE), provenance=d.get("provenance") or None,
                   layer=Layer(d.get("layer", Layer.STRUCTURE.value)), boxes=list(d.get("boxes") or []),
                   pictures=list(d.get("pictures") or []))


def is_delete(n) -> bool:
    return isinstance(n, BlockDelete)


def _box_text(box) -> str:
    from .docmodel import paragraph_text
    return " ".join(paragraph_text(p) for p in box.iter(f"{_W}p")).strip()


def _blobs(ms: Manuscript, p: Paragraph) -> list[str]:
    out = []
    part = ms.document.part
    for b in p.element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        rid = b.get(qn("r:embed"))
        try:
            out.append(hashlib.sha1(part.related_parts[rid].blob).hexdigest())
        except (KeyError, AttributeError):
            return []
    return out


def plan(ms: Manuscript, style: JournalStyle) -> list[BlockDelete]:
    if not style.get("leftovers.enabled", True):
        return []
    paras = ms.paragraphs
    boxes = []
    seen_boxes = set()
    for p in paras:
        if not in_textbox(p):
            continue
        box = next(a for a in p.element.iterancestors() if a.tag == f"{_W}txbxContent")
        if any(b is box for b in seen_boxes):
            continue
        seen_boxes.add(box)
        if re.sub(r"\s+", " ", _box_text(box)).strip(" :.").lower() in _BOX_WORDS:
            boxes.append(p.id)
    pictures = []
    ref = references.find_heading(ms)
    if ref is not None:
        body = {h for q in paras[:ref] if has_picture(q.element) for h in _blobs(ms, q)}
        from . import rules_figures

        def captioned(i: int) -> bool:  # a legend right above or below the picture (blank lines aside)
            for step in (-1, 1):
                j = i + step
                while 0 <= j < len(paras) and not paras[j].text.strip() and abs(j - i) < 4:
                    j += step
                if 0 <= j < len(paras) and paras[j].text.strip() and (
                        rules_figures.caption_match(paras[j]) is not None
                        or re.match(r"\s*(?:<H\d>)?\s*fig(?:ure)?s?\.?\s*\d", paras[j].text, re.I)):
                    return True
            return False

        for i in range(ref + 1, len(paras)):
            q = paras[i]
            if q.in_table or q.text.strip() or not has_picture(q.element) or captioned(i):
                continue
            hs = _blobs(ms, q)
            if hs and all(h in body for h in hs):
                pictures.append(q.id)  # the same picture again after the references, no legend with it (997)
    out = []
    if boxes:
        what = ", ".join(sorted({f"\"{ms.by_id(i).text.strip()}\"" for i in boxes}))
        out.append(BlockDelete(id="leftover_boxes", anchor=boxes[0], where="before",
                               runs=[(f"[Template text boxes deleted: {what}]", {})], role="delete", reason=DELETE,
                               provenance=provenance.make("rule", "Template leftover", "text boxes holding only a "
                                                          "template word", rule=DELETE),
                               layer=Layer.STRUCTURE, boxes=boxes))
    if pictures:
        out.append(BlockDelete(id="leftover_pictures", anchor=pictures[0], where="before",
                               runs=[(f"[Repeated pictures after the references deleted: {len(pictures)}]", {})],
                               role="delete", reason=DELETE,
                               provenance=provenance.make("rule", "Template leftover", "pictures after the reference "
                                                          "list that repeat a picture of the article", rule=DELETE),
                               layer=Layer.STRUCTURE, pictures=pictures))
    return out


def render(ms: Manuscript, inserts, statuses, *, tracked: bool, author: str, date: str) -> list[str]:
    """Delete the objects of the records whose status is in `statuses` (call after tracked.render)."""
    from .header import _Marks, _mark_block

    notes: list[str] = []
    mk = _Marks(ms, author, date, tracked)
    body = ms.document.element.body
    for n in inserts:
        if not is_delete(n) or n.status not in statuses:
            continue
        for pid in n.boxes:
            try:
                p = ms.by_id(pid)
            except KeyError:
                continue
            run = next((a for a in p.element.iterancestors() if a.tag == f"{_W}r"
                        and a.getparent() is not None and a.getparent().tag in (f"{_W}p", f"{_W}hyperlink")
                        and not any(x.tag == f"{_W}txbxContent" for x in a.iterancestors())), None)
            if run is None or run.getparent() is None:
                continue
            host = run.getparent()
            if tracked:
                if run.getparent().tag == f"{_W}del":
                    continue
                w = mk.mark("w:del")
                run.addprevious(w)
                w.append(run)
            else:
                host.remove(run)
            _drop_if_empty(host, body, mk, tracked, ms)
        for pid in n.pictures:
            try:
                el = ms.by_id(pid).element
            except KeyError:
                continue
            if el.getparent() is not body or el.getnext() is None:
                continue
            if tracked:
                _mark_block(el, mk, "w:del")
            else:
                ms.removed.append(el)  # (it leaves the XML on save: its id stays a valid anchor until then)
    return notes


def _drop_if_empty(host, body, mk, tracked: bool, ms: Manuscript) -> None:
    """A paragraph left with nothing but the deleted box goes too (not the last of the document)."""
    from .docmodel import paragraph_text

    if host.tag != f"{_W}p" or host.getparent() is not body or host.getnext() is None:
        return
    live = [r for r in host.iter(f"{_W}r") if not any(a.tag == f"{_W}del" for a in r.iterancestors())
            and not any(a.tag == f"{_W}txbxContent" for a in r.iterancestors())]
    if paragraph_text(host).strip() or any(has_picture(r) or r.find(f".//{_W}drawing") is not None or
                                           r.find(f".//{_MC}") is not None for r in live):
        return
    if tracked:
        from .header import _mark_paragraph
        ppr = host.find(f"{_W}pPr")
        if ppr is not None and ppr.find(f"{_W}rPr/{_W}del") is not None:
            return
        _mark_paragraph(host, mk, "w:del")
    else:
        ms.removed.append(host)
