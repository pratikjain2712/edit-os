"""Figure legends at the end of the article under a "Figure Legends" heading; table captions in the table's first row.

The guide: G265 "The caption(s) should be placed at end of document (before table if present)", G195 the tables "at
end of document after figure captions"; the stylesheets: "Image should be present first and below it the respective
caption" (a picture moves with its legend). The back of the article then reads: references, the First author block
(first_author.py), the "Figure Legends" heading and the legends, the tables.

  - Legends in the body (or after the tables) move there, each with the paragraphs that continue it (a line broken
    mid-phrase, a "Note. ..." under it, its prose after the references: rules_figures.legends): one `BlockMove` record
    the editor accepts or rejects. In the
    tracked copy each legend is deleted where it was and inserted, as edited, at the end (header.move_blocks does the
    same for renumbered tables); in the clean copy it is moved. A heading over the legends moves with them; with none,
    the move adds one.
  - Legends already in place with no heading get the heading as an inserted paragraph (`NewParagraph`); an author's
    heading ("Legends:", "Legends to Figures", "Image description") is renamed by `rules_figures` (fig_captions).
  - Figures supplied only as files (figures_insert.py) go under the heading too, among the legends in figure order.
  - Table captions go into a merged first row of their table (`figures.table_caption_row`): the IJO technical editors'
    form in every article with tables (8 of 8); the sheet is silent on where the caption sits.

  - Captioned tables in the body move to the end, after the legends, with their captions and notes, under a
    "Tables" heading (`figures.tables_to_end`, `figures.tables_heading`; the JEHP editors in every such article). A caption
    typed as "Table 1." over its title line is one caption (rules_figures.split_table_captions); a note right under a
    table that no footnote pattern knows ("Cells show ...") goes with it too.
    Tables already at the end stay; a moved one goes before the first of them with a higher number. Uncaptioned
    tables stay where they are; renumbered tables are left (a note) since header.move_blocks reorders them.

Settings: `figures.legends_to_end`, `figures.legends_heading`, `figures.legends_heading_single`,
`figures.table_caption_row`, `figures.tables_to_end`, `figures.tables_heading`.

    back_slot(ms)                       where the back blocks start (First author, legends, tables)
    plan(ms, style, figure_inserts)     -> list of records (NewParagraph / BlockMove) and notes
    render(ms, inserts, statuses, ...)  carry out the accepted (tracked copy: also pending) moves
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import provenance, references, rules_figures
from .changes import Layer, Source, Status
from .docmodel import Manuscript, has_picture, paragraph_text
from .frontmatter import NewParagraph
from .style import JournalStyle

MOVE_LEGENDS = "fig_legends_move"
MOVE_TABLE_CAPTION = "table_caption_row"
MOVE_TABLES = "tables_move"
HEADING = "fig_legends_heading"


@dataclass
class BlockMove(NewParagraph):
    """Paragraphs to move (legends to the end, a table caption into its table). `anchor` / `where` say where they go
    for the review screen; `render` works out the place in the rendered document."""
    kind: str = "legends"  # legends | table_caption | tables
    blocks: list[list[str]] = field(default_factory=list)  # paragraph ids, one list per block, in order
    heading: str = ""  # legends: a heading to put over them ("" = none, or the author's heading moves with them)
    refs_end: str = ""  # legends: the last paragraph of the reference list
    side: str = ""  # table_caption: "above" | "below" -- where the manuscript's table captions sit
    notes: list[str] = field(default_factory=list)  # tables: notes split out of the captions (placed under the tables)

    def to_dict(self) -> dict:
        return {**super().to_dict(), "move": self.kind, "blocks": self.blocks, "heading": self.heading,
                "refs_end": self.refs_end, "side": self.side, "notes": self.notes}

    @classmethod
    def from_dict(cls, d: dict) -> "BlockMove":
        return cls(id=d["id"], anchor=d["anchor"], where=d.get("where", "before"),
                   runs=[(t, f) for t, f in d.get("runs") or []], role=d.get("role", "move"),
                   status=Status(d.get("status", "suggested")), source=Source(d.get("source", "rule")),
                   reason=d.get("reason", MOVE_LEGENDS), provenance=d.get("provenance") or None,
                   layer=Layer(d.get("layer", Layer.TABLES_FIGURES.value)), kind=d.get("move", "legends"),
                   blocks=[list(b) for b in d.get("blocks") or []], heading=d.get("heading", ""),
                   refs_end=d.get("refs_end", ""), side=d.get("side", ""), notes=list(d.get("notes") or []))


def is_move(n) -> bool:
    return isinstance(n, BlockMove)


# --- where things are ------------------------------------------------------------------------

def refs_end(ms: Manuscript) -> int | None:
    """Index of the last paragraph of the reference list (its heading when it has no entries)."""
    head = references.find_heading(ms)
    if head is None:
        return None
    index = {p.id: i for i, p in enumerate(ms.paragraphs)}
    refs = references.extract(ms)
    ids = [r.para_id for r in refs] + [pid for r in refs for pid, _ in r.continuation]
    return max([head] + [index[i] for i in ids if i in index])


def _is_table_caption(p) -> bool:
    m = rules_figures.caption_match(p)
    return m is not None and m.group("word").lower().startswith("tab")


def legend_heading(ms: Manuscript, first: int) -> int | None:
    """Index of the heading right above the legend (or its picture) at index `first`, if it has one."""
    return rules_figures.heading_above(ms, first)


def _block_start(ms: Manuscript, i: int) -> int:
    """A legend's first paragraph: the picture right above it, if any (blank lines between)."""
    start = i
    for j in range(i - 1, max(-1, i - 4), -1):
        p = ms.paragraphs[j]
        if p.in_table:
            break
        if has_picture(p.element) and not p.text.strip():
            start = j
            continue
        if p.text.strip():
            break
    return start


def back_slot(ms: Manuscript) -> tuple[str, str] | None:
    """(anchor paragraph id, "before" | "after"): where the back of the article starts after the reference list -- the
    legends' heading (or the picture over the first legend, or the legend) when legends follow the references before
    any table, else the first table caption (or the paragraph above the first table) -- with none of these, after the
    last paragraph with text or a picture. None for an empty document."""
    paras = ms.paragraphs
    if not paras:
        return None
    end = refs_end(ms)
    index = {p.id: i for i, p in enumerate(paras)}
    legs = {index[leg.para.id] for leg in rules_figures.legends(ms)}
    if end is not None:  # after the references, a legend in a heading style or written as a sentence counts too
        legs |= {i for i in range(end + 1, len(paras)) if not paras[i].in_table and _LEGEND_START.match(paras[i].text)}
    if end is not None and paras[end].in_table:
        # a reference list laid out as a table (references._reference_table): the back starts after that table,
        # never inside its last cell
        while end + 1 < len(paras) and paras[end + 1].in_table:
            end += 1
    if end is not None:
        for i in range(end + 1, len(paras)):
            p = paras[i]
            if p.in_table:
                above = [j for j in range(end, i) if not paras[j].in_table and (paras[j].text.strip() or
                                                                                   has_picture(paras[j].element))]
                if not above and paras[end].in_table:  # between the reference table and the next one
                    gap = [j for j in range(end + 1, i) if not paras[j].in_table]
                    if gap:
                        return paras[gap[0]].id, "before"
                return paras[above[-1] if above else end].id, "after"
            if i in legs:
                h = legend_heading(ms, _block_start(ms, i))
                return paras[h if h is not None else _block_start(ms, i)].id, "before"
            if rules_figures.LEGENDS_HEADING.match(p.text) and any(j > i for j in legs):
                return p.id, "before"
            if _is_table_caption(p) or _TABLES_HEADING.match(p.text):
                return p.id, "before"
    body = [i for i, p in enumerate(paras) if not p.in_table]
    if not body:
        return None
    last = body[-1]
    filled = [i for i in body if paras[i].text.strip() or has_picture(paras[i].element)]
    if filled and not any(paras[j].in_table for j in range(filled[-1] + 1, last + 1)):
        last = filled[-1]  # trailing empty paragraphs stay below the block (a table after it does not)
    return paras[last].id, "after"


def _first_table_after(ms: Manuscript, end: int) -> int:
    """Index of the first table (or table caption) after the reference list; len(paragraphs) when there is none."""
    for i in range(end + 1, len(ms.paragraphs)):
        p = ms.paragraphs[i]
        if p.in_table or _is_table_caption(p) or _TABLES_HEADING.match(p.text):
            return i
    return len(ms.paragraphs)


# A heading over the tables at the end: "Tables", "TABLES:", "Table Legends".
_TABLES_HEADING = re.compile(r"^\s*(?:<H\d>)?\s*tables?(?:\s+(?:legends?|captions?))?\s*[:.\-–—]*\s*$", re.I)
# "Figure 1", "Fig. 2:", "Figure 1 presents ...": a figure legend after the reference list, whatever its style.
_LEGEND_START = re.compile(r"^\s*(?:<H\d>)?\s*fig(?:ure)?s?\.?\s*\d{1,2}(?!\d)", re.I)


def heading_text(style: JournalStyle, n_figures: int) -> str:
    many = style.get("figures.legends_heading") or ""
    return (style.get("figures.legends_heading_single") or many) if n_figures == 1 else many


# --- planning --------------------------------------------------------------------------------

def plan(ms: Manuscript, style: JournalStyle, figure_inserts=()) -> tuple[list[NewParagraph], list[str]]:
    """Records for the legends (move / heading) and the table captions; notes for the editor."""
    out: list[NewParagraph] = []
    notes: list[str] = []
    if style.get("figures.legends_to_end"):
        out += _plan_legends(ms, style, figure_inserts, notes)
    from . import rules_tables
    out += rules_tables.plan_cell_blanks(ms, style)  # before the tables move: the blanks go in the table moved
    split = _split_notes(ms, style)
    if style.get("figures.table_caption_row"):
        out += _plan_table_captions(ms, split)
    if style.get("figures.tables_to_end"):
        out += _plan_tables(ms, style, notes, split)  # after the caption rows: a caption goes into its row, then moves along
    return out, notes


_BARE_TABLE_LABEL = re.compile(r"^\s*tables?\s*(\d{1,2}|[IVX]{1,5})\s*[.:]?\s*$", re.I)


def _tbl_number(ms: Manuscript, tbl, cap) -> int | None:
    if cap is not None:
        m = rules_figures.caption_match(cap)
        if m is None:
            b = _BARE_TABLE_LABEL.match(cap.text)
            return _num(b.group(1).upper()) if b else None
        return _num(m.group("num")) if m else None
    tr = tbl.find(qn("w:tr"))
    m = _CAPTION_START.match(_row_text(tr)) if tr is not None else None
    return _num(re.search(r"[\dIVX]+", m.group(0)).group(0)) if m else None


def _row_text(tr) -> str:
    return " ".join(paragraph_text(p) for p in tr.iter(qn("w:p")))


def _num(s: str) -> int | None:
    return int(s) if s.isdigit() else rules_figures.roman_to_int(s)


def _split_notes(ms: Manuscript, style: JournalStyle) -> list[str]:
    """The notes rules_figures splits out of table captions (figures.table_caption_notes), as they will read."""
    if style.get("figures.table_caption_notes") != "split":
        return []
    from .tracked import PARAGRAPH_BREAK

    out = []
    mode = style.get("figures.table_caption_end_period") or "delete"
    for cap in rules_figures.captions(ms):
        if cap.kind != rules_figures.TABLE:
            continue
        got = rules_figures._caption_note(ms, ms.by_id(cap.para_id), cap.match, mode, style)
        if got:
            out += [t.strip() for c in got[1] for t in c.new.split(PARAGRAPH_BREAK) if t.strip()]
    return out


def _plan_tables(ms: Manuscript, style: JournalStyle, notes: list[str], split: list[str]) -> list[NewParagraph]:
    """Captioned tables in the body (before the reference list's end) move to the end of the article, after the
    figure legends (G195), each with its caption and the notes under it, under a "Tables" heading. A table with no
    caption stays (the editors leave those: 1580, 1751)."""
    paras = ms.paragraphs
    end = refs_end(ms)
    if end is None:
        return []
    body = ms.document.element.body
    stop = _top(paras[end].element, body)
    by_el = {}
    for p in paras:
        by_el.setdefault(_top(p.element, body), []).append(p)
    note_ids = {q.id for block in rules_figures.table_notes(ms) for q in block}
    split_title = {t.element: lab for lab, t in rules_figures.split_table_captions(ms)}  # "Table 1." / its title
    kids = list(body)
    blocks, numbers = [], []
    for k, el in enumerate(kids):
        if el is stop:
            break
        if el.tag != qn("w:tbl") or el not in by_el:
            continue
        cap, side, title, between = None, "", None, None
        for s_, step in (("above", -1), ("below", 1)):
            j = k + step
            while 0 <= j < len(kids) and kids[j].tag == qn("w:p") and _empty(kids[j]):
                j += step
            if 0 <= j < len(kids) and kids[j].tag == qn("w:p") and kids[j] in by_el:
                q = by_el[kids[j]][0]
                m = rules_figures.caption_match(q)
                if s_ == "above" and kids[j] in split_title:  # "Table 1." / "Baseline ..." over it: one caption
                    cap, side, title = split_title[kids[j]], s_, q
                    break
                if m is not None and m.group("word").lower().startswith("tab") or _BARE_TABLE_LABEL.match(q.text):
                    cap, side = q, s_  # (a bare "Table 1" over its table, the title missing, is its caption too)
                    break
                if s_ == "above" and m is None and len(q.text) < 400:
                    # one line between the caption and its table ("Table 1. Summary of the protocol" / "Session
                    # content was adapted from ..." / the table: jehp_1764): the caption above it, the line with it
                    i = j - 1
                    while 0 <= i < len(kids) and kids[i].tag == qn("w:p") and _empty(kids[i]):
                        i -= 1
                    if 0 <= i < len(kids) and kids[i].tag == qn("w:p") and kids[i] in by_el:
                        c2 = by_el[kids[i]][0]
                        m2 = rules_figures.caption_match(c2)
                        if m2 is not None and m2.group("word").lower().startswith("tab"):
                            cap, side, between = c2, s_, q
                            break
        tr = el.find(qn("w:tr"))
        in_row = tr is not None and _CAPTION_START.match(_row_text(tr))
        if cap is None and not in_row:
            continue
        ids = ([cap.id] if cap is not None and side == "above" else []) + ([title.id] if title is not None else []) \
            + ([between.id] if between is not None else []) + [by_el[el][0].id]
        if cap is not None and side == "below":
            ids.append(cap.id)
        last = kids.index(cap.element) if cap is not None and side == "below" else k
        for j in range(last + 1, len(kids)):  # the notes under it
            x = kids[j]
            if x.tag != qn("w:p") or x not in by_el:
                break
            q = by_el[x][0]
            if _empty(x):
                continue
            if q.id in note_ids or j == last + 1 and rules_figures.loose_table_note(q.text):
                ids.append(q.id)  # (a note right under the table that no footnote pattern knows: jehp_1402)
                continue
            break
        blocks.append(ids)
        numbers.append(_tbl_number(ms, el, cap))
    if not blocks:
        return []
    if "Table" in rules_figures.renumber_plan(ms, style):
        notes.append("Tables renumbered: left where they are (move them to the end after the figure legends by hand)")
        return []
    head = style.get("figures.tables_heading") or ""
    what = ", ".join(f"Table {n}" for n in numbers if n is not None) or f"{len(blocks)} tables"
    prov = provenance.make("guide", "Pre-editing guide", "G195: the tables at the end of the document, after the "
                           "figure captions", rule=MOVE_TABLES)
    return [BlockMove(id="tables_move", anchor=paras[end].id, where="after",
                      runs=[(f"[Tables moved to the end, after the figure legends: {what}]", {})],
                      role="move", reason=MOVE_TABLES, provenance=prov, layer=Layer.TABLES_FIGURES,
                      kind="tables", blocks=blocks, heading=head, refs_end=paras[end].id, notes=split)]


def _plan_legends(ms: Manuscript, style: JournalStyle, figure_inserts, notes: list[str]) -> list[NewParagraph]:
    paras = ms.paragraphs
    index = {p.id: i for i, p in enumerate(paras)}
    legs = rules_figures.legends(ms)
    end = refs_end(ms)
    no_legend = [f for f in figure_inserts if getattr(f, "legend", "")]
    numbers = {leg.number for leg in legs} | {f.figure for f in no_legend}
    head = heading_text(style, len(numbers))
    if end is None or not (legs or no_legend):
        return []
    if any(q.in_table for leg in legs for q in leg.paragraphs):
        notes.append("Figure legends in a table: left where they are")
        return []
    stop = _first_table_after(ms, end)
    in_place = all(end < index[leg.para.id] < stop for leg in legs)
    starts = [_block_start(ms, index[leg.para.id]) for leg in legs]
    h = legend_heading(ms, min(starts)) if legs else None
    prov = provenance.make("guide", "Pre-editing guide", "G265: the captions at the end of the document, before the "
                           "tables; G195: the tables after the figure captions", rule=MOVE_LEGENDS)
    if legs and not in_place:
        heads = sorted({x for x in (legend_heading(ms, s) for s in starts) if x is not None})
        blocks = [[paras[j].id for j in heads]] if heads else []
        blocks += [[q.id for q in leg.paragraphs] for leg in legs]  # with its notes and the lines that continue it
        slot = back_slot(ms)
        anchor, where = slot if slot else (paras[-1].id, "after")
        what = ", ".join(f"Figure {leg.number}" for leg in legs)
        return [BlockMove(id="legends_move", anchor=anchor, where=where,
                          runs=[(f"[Figure legends moved to the end, after the references: {what}]", {})],
                          role="move", reason=MOVE_LEGENDS, provenance=prov, layer=Layer.TABLES_FIGURES,
                          kind="legends", blocks=blocks, heading="" if heads else head,
                          refs_end=paras[end].id)]
    if h is not None or not head:
        return []
    # In place (or only supplied figures) with no heading: add one.
    if legs:
        anchor, where = paras[min(starts)].id, "before"
    else:
        slot = back_slot(ms)
        if slot is None:
            return []
        anchor, where = slot
    hp = provenance.make("journal_style", "Journal practice", f"\"{head}\" over the figure legends (the editors' "
                         "heading; the sheet is silent)", rule=HEADING)
    return [NewParagraph("legends_heading", anchor, where, [(head, {"bold": True})], "legends_heading",
                         source=Source.RULE, reason=HEADING, provenance=hp, layer=Layer.TABLES_FIGURES)]


def _plan_table_captions(ms: Manuscript, split: list[str] = ()) -> list[NewParagraph]:
    """A table caption right above its table (or right below it, notes aside) goes into a merged first row."""
    out = []
    body = ms.document.element.body
    side = _side(ms) or ("above" if rules_figures.split_table_captions(ms) else None)
    if side is None:
        return out
    claimed: list = []  # (elements, not id(): lxml proxies are not stable)
    for lab, title in rules_figures.split_table_captions(ms):  # "Table 1." / "Baseline ...": the label goes in, joined
        tbl = title.element.getnext()
        while tbl is not None and _empty(tbl):
            tbl = tbl.getnext()
        if tbl is None or tbl.tag != qn("w:tbl") or lab.element.getparent() is not body:
            continue
        claimed.append(tbl)
        num = re.search(r"[\dIVX]+\s*[.:\-–—]?\s*$", lab.text)
        prov = provenance.make("journal_style", "Journal practice", "Table caption in a merged first row of the table "
                               "(the editors' form; the sheet is silent)", rule=MOVE_TABLE_CAPTION)
        out.append(BlockMove(id=f"tcap{len(out) + 1}", anchor=lab.id, where="before",
                             runs=[(f"[{lab.text.strip().rstrip('.:')} caption moved into the table's first row]", {})],
                             role="move", reason=MOVE_TABLE_CAPTION, provenance=prov, layer=Layer.TABLES_FIGURES,
                             kind="table_caption", blocks=[[lab.id]], side="above", notes=list(split)))
    for cap in rules_figures.captions(ms):
        if cap.kind != rules_figures.TABLE:
            continue
        p = ms.by_id(cap.para_id)
        if p.in_table or p.element.getparent() is not body or has_picture(p.element):
            continue
        here = side
        tbl = _caption_table(p.element, side)
        if tbl is None or any(t is tbl for t in claimed):  # this caption sits on the other side of its table
            here = "below" if side == "above" else "above"
            tbl = _caption_table(p.element, here)
            if tbl is not None and _captioned(tbl, side):
                tbl = None
        if tbl is None or any(t is tbl for t in claimed):
            continue
        claimed.append(tbl)
        prov = provenance.make("journal_style", "Journal practice", "Table caption in a merged first row of the table "
                               "(the editors' form; the sheet is silent)", rule=MOVE_TABLE_CAPTION)
        out.append(BlockMove(id=f"tcap{len(out) + 1}", anchor=p.id, where="before",
                             runs=[(f"[Table {cap.number} caption moved into the table's first row]", {})],
                             role="move", reason=MOVE_TABLE_CAPTION, provenance=prov, layer=Layer.TABLES_FIGURES,
                             kind="table_caption", blocks=[[p.id]], side=here, notes=list(split)))
    return out


def _side(ms: Manuscript) -> str | None:
    """Whether the manuscript's table captions sit above or below their tables, read from its first table that has a
    caption next to it (None: no table caption next to a table)."""
    body = ms.document.element.body
    for tbl in body.iterchildren(qn("w:tbl")):
        for side, step in (("above", "getprevious"), ("below", "getnext")):
            el = getattr(tbl, step)()
            while el is not None and _empty(el):
                el = getattr(el, step)()
            if el is not None and el.tag == qn("w:p") and _CAPTION_START.match(paragraph_text(el)):
                return side
    return None


_CAPTION_START = re.compile(r"\s*tables?\s*[\dIVX]{1,4}(?![\w])", re.I)


def _empty(el) -> bool:
    """A blank paragraph (no text, no picture, text box or object)."""
    return el.tag == qn("w:p") and not paragraph_text(el).strip() and not has_picture(el) and \
        el.find(f".//{qn('w:drawing')}") is None and el.find(f".//{qn('w:pict')}") is None and \
        el.find(f".//{qn('w:object')}") is None


def _captioned(tbl, side: str) -> bool:
    """Whether a table has a caption paragraph on that side ("above" | "below")."""
    step = "getprevious" if side == "above" else "getnext"
    el = getattr(tbl, step)()
    while el is not None and _empty(el):
        el = getattr(el, step)()
    return el is not None and el.tag == qn("w:p") and bool(_CAPTION_START.match(paragraph_text(el)))


def _caption_table(p_el, side: str, notes=()):
    """The table a caption paragraph belongs to: the table right below it (captions above their tables), or right
    above it (captions below), blank lines (and a note split out of the caption, `notes`) aside."""
    step = "getnext" if side == "above" else "getprevious"
    el = getattr(p_el, step)()
    while el is not None and (_empty(el) or el.tag == qn("w:p") and paragraph_text(el).strip() in notes
                              and paragraph_text(el).strip()):
        el = getattr(el, step)()
    return el if el is not None and el.tag == qn("w:tbl") else None


# --- rendering -------------------------------------------------------------------------------

def render(ms: Manuscript, inserts, statuses, *, tracked: bool, author: str, date: str,
           run_format: dict | None = None, align: str | None = None) -> list[str]:
    """Carry out the moves whose status is in `statuses` (call after tracked.render and the inserts); a heading it
    adds has its runs in `run_format` (the journal's font and size) and is aligned `align`. Returns notes."""
    from .header import _Marks

    notes: list[str] = []
    mk = _Marks(ms, author, date, tracked, run_format, align)
    for n in inserts:
        if not is_move(n) or n.status not in statuses:
            continue
        try:
            if n.kind == "legends":
                notes += _move_legends(ms, n, mk)
            elif n.kind == "table_caption":
                notes += _caption_row(ms, n, mk)
            elif n.kind == "tables":
                notes += _move_tables(ms, n, mk)
            elif n.kind == "cell_blanks":
                from . import rules_tables
                notes += rules_tables.render_cell_blanks(ms, n, mk)
        except KeyError:
            notes.append(f"{n.text}: not done (a paragraph is gone)")
    _order_back_figures(ms)
    return notes


_LEGEND_NUMBER = re.compile(r"^\s*fig(?:ure)?s?\.?\s*(\d{1,2})(?!\d)", re.I)


def _order_back_figures(ms: Manuscript) -> None:
    """Figures supplied only as files, with their legends from other material (figures_insert.render records them),
    were placed where the back of the article starts, which is above the legends' heading. Each goes under the
    heading, among the legends in figure order: before the first legend with a higher number (its picture above it),
    else after the last legend (and the lines that go with it)."""
    blocks = sorted(getattr(ms, "_back_figures", None) or [], key=lambda b: b[0])
    body = ms.document.element.body
    for num, els in blocks:
        if not els or els[-1].getparent() is not body:
            continue
        mine = {id(e) for e in els}
        heading = None
        legs: list[tuple[int, object]] = []
        gap = False  # a blank line since the last legend (or a line of it)
        x = els[-1].getnext()
        while x is not None:
            if x.tag != qn("w:p"):
                break  # a table, the section properties
            t = paragraph_text(x)
            if id(x) in mine or _is_deleted(x):
                x = x.getnext()
                continue
            if _CAPTION_START.match(t) or _TABLES_HEADING.match(t) and t.strip():
                break
            m = _LEGEND_NUMBER.match(t)
            if m:
                legs.append((int(m.group(1)), x))
            elif heading is None and not legs and t.strip() and rules_figures.LEGENDS_HEADING.match(t):
                heading = x
            elif t.strip() and (not legs and heading is None or legs and gap):
                break  # other text: not over the legends, or past them
            if not t.strip() and not has_picture(x):
                gap = True
            elif t.strip():
                gap = False
            x = x.getnext()
        if heading is None and not any(k < num for k, _ in legs):
            continue  # already under the heading, before the legends with higher numbers
        higher = next((el for k, el in legs if k > num), None)
        if higher is not None:
            above = _picture_above(higher)
            target = above[0] if above else higher
            for e in els:
                target.addprevious(e)
            continue
        if not legs:  # only the heading under it
            at = heading
        else:  # after the last legend and the lines that go with it (up to a blank line)
            at = legs[-1][1]
            y = at.getnext()
            while y is not None and y.tag == qn("w:p") and paragraph_text(y).strip() and \
                    not _LEGEND_NUMBER.match(paragraph_text(y)) and not _CAPTION_START.match(paragraph_text(y)) \
                    and not _TABLES_HEADING.match(paragraph_text(y)):
                at = y
                y = y.getnext()
        for e in els:
            at.addnext(e)
            at = e


def _top(el, body):
    while el is not None and el.getparent() is not body:
        el = el.getparent()
    return el


def _is_deleted(el) -> bool:
    return el.tag == qn("w:p") and el.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None


def _is_inserted(el) -> bool:
    return el.tag == qn("w:p") and el.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}") is not None


def _detach(els: list, mk) -> list:
    """The elements to place at the new position; in the tracked copy the originals stay, deleted, and copies with
    the edits accepted go in as insertions (a paragraph that is itself an insertion just moves)."""
    from .header import _mark_block, _reject, moved_copy

    if not mk.tracked:
        return els
    out = []
    for e in els:
        if _is_inserted(e):
            out.append(e)
            continue
        c = moved_copy(e)
        _mark_block(c, mk, "w:ins")
        _reject(e)
        _mark_block(e, mk, "w:del")
        out.append(c)
    return out


def _dest(ms: Manuscript, refs_end_id: str):
    """The element the legends go before: the first table or table caption after the reference list; None: the end."""
    body = ms.document.element.body
    el = _top(ms.by_id(refs_end_id).element, body)
    nxt = el.getnext()
    while nxt is not None:
        if nxt.tag == qn("w:tbl"):
            return nxt
        if nxt.tag == qn("w:p") and (_CAPTION_START.match(paragraph_text(nxt)) or
                                     _TABLES_HEADING.match(paragraph_text(nxt))):
            return nxt
        if nxt.tag == qn("w:sectPr"):
            return nxt
        nxt = nxt.getnext()
    return None


def _picture_above(el) -> list:
    """Picture-only paragraphs right above a legend (blank lines between), top-down."""
    out, prev, steps = [], el.getprevious(), 0
    while prev is not None and prev.tag == qn("w:p") and steps < 3:
        if has_picture(prev) and not paragraph_text(prev).strip():
            out.insert(0, prev)
        elif paragraph_text(prev).strip():
            break
        prev = prev.getprevious()
        steps += 1
    return out


def _move_legends(ms: Manuscript, n: BlockMove, mk) -> list[str]:
    body = ms.document.element.body
    groups = []
    for block in n.blocks:
        els = []
        for k, pid in enumerate(block):
            el = ms.by_id(pid).element
            if el.getparent() is not body:
                return [f"{n.text}: not done (a legend is not a paragraph of its own)"]
            if k and (_is_deleted(el) or any(el is r for r in ms.removed)):
                continue  # a line joined to the legend (run on): its deleted paragraph stays where it was
            els += (_picture_above(el) if not k else []) + [el]
        if els:
            groups.append(els)
    dest = _dest(ms, n.refs_end)
    place = []
    if n.heading:
        place.append(_heading(n.heading, mk))
    seen = set()
    for els in sorted(groups, key=lambda g: list(body).index(g[-1])):
        for e in els:
            if id(e) not in seen:
                seen.add(id(e))
                place.append(e)
    if any(_is_deleted(e) for e in place):
        return [f"{n.text}: not done (the legends were moved with a renumbering)"]
    moved = _detach(place, mk)
    for e in moved:
        if dest is not None:
            dest.addprevious(e)
        else:
            body.append(e)
    return []


def _heading(text: str, mk):
    from .header import aligned

    p = OxmlElement("w:p")
    ppr = aligned(OxmlElement("w:pPr"), mk)
    if mk.tracked:
        rpr = OxmlElement("w:rPr")
        rpr.append(mk.mark("w:ins"))
        ppr.append(rpr)
    p.append(ppr)
    from .header import styled

    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    r.append(rpr)
    styled(r, mk)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    if mk.tracked:
        w = mk.mark("w:ins")
        w.append(r)
        p.append(w)
    else:
        p.append(r)
    return p


def _grid_columns(tbl) -> int:
    grid = tbl.find(qn("w:tblGrid"))
    n = len(grid.findall(qn("w:gridCol"))) if grid is not None else 0
    if n:
        return n
    tr = tbl.find(qn("w:tr"))
    total = 0
    for tc in (tr.findall(qn("w:tc")) if tr is not None else []):
        span = tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
        total += int(span.get(qn("w:val"))) if span is not None else 1
    return max(total, 1)


def _caption_row(ms: Manuscript, n: BlockMove, mk) -> list[str]:
    body = ms.document.element.body
    p_el = ms.by_id(n.blocks[0][0]).element
    if p_el.getparent() is not body:
        return [f"{n.text}: not done (the caption is not a paragraph of its own)"]
    if _is_deleted(p_el):
        return [f"{n.text}: not done (the caption was moved with its table)"]
    tbl = _caption_table(p_el, n.side or "above", n.notes)
    if tbl is None:
        return [f"{n.text}: not done (no table next to the caption)"]
    cols = _grid_columns(tbl)
    tr = OxmlElement("w:tr")
    trpr = OxmlElement("w:trPr")
    if mk.tracked:
        trpr.append(mk.mark("w:ins"))
    tr.append(trpr)
    tc = OxmlElement("w:tc")
    tcpr = OxmlElement("w:tcPr")
    if cols > 1:
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), str(cols))
        tcpr.append(span)
    tc.append(tcpr)
    tr.append(tc)
    split = []  # a note split out of the caption (between it and the table) goes under the table
    if n.side != "below":
        x = p_el.getnext()
        while x is not None and x is not tbl:
            if x.tag == qn("w:p") and paragraph_text(x).strip() in n.notes and paragraph_text(x).strip():
                split.append(x)
            x = x.getnext()
    (moved,) = _detach([p_el], mk)
    tc.append(moved)
    first = tbl.find(qn("w:tr"))
    if first is not None:
        first.addprevious(tr)
    else:
        tbl.append(tr)
    at = tbl
    for e in _detach(split, mk):
        at.addnext(e)
        at = e
    return []


def _move_tables(ms: Manuscript, n: BlockMove, mk) -> list[str]:
    """Each table block (caption, table, notes -- whatever of them is still a body paragraph: a caption already put
    into its table's first row travels inside it) goes to the end of the document, under the heading."""
    body = ms.document.element.body
    kids = list(body)
    known = {p.element for p in ms.paragraphs}
    groups = []
    for block in n.blocks:
        els = []
        for pid in block:
            el = _top(ms.by_id(pid).element, body)
            if el is None:
                continue
            if el.tag == qn("w:p") and _is_deleted(el):
                continue  # a caption moved into its table (tracked copy): the deleted original stays
            if not any(e is el for e in els):
                els.append(el)
            if el.tag == qn("w:p"):  # paragraphs a change split off it (a note's abbreviations moved first) go too
                x = el.getnext()
                while x is not None and x.tag == qn("w:p") and x not in known:
                    if not any(e is x for e in els):
                        els.append(x)
                    x = x.getnext()
        tbls = [e for e in els if e.tag == qn("w:tbl")]
        if not tbls:
            return [f"{n.text}: not done (a table is gone)"]
        if any(_is_deleted_tbl(t) for t in tbls):
            return [f"{n.text}: not done (the tables were moved with a renumbering)"]
        tbl = tbls[0]
        split = []  # notes split out of the caption: between the caption and the table, or under the table
        for step in ("getprevious", "getnext"):
            x = getattr(tbl, step)()
            while x is not None and x.tag == qn("w:p"):
                t = paragraph_text(x).strip()
                if t and t in n.notes:
                    split.append(x)
                elif t:
                    break
                x = getattr(x, step)()
        els = [e for e in els if not any(e is x for x in split)]
        order = sorted(els, key=lambda e: list(body).index(e))
        k = next(i for i, e in enumerate(order) if e is tbl)
        groups.append(order[:k + 1] + sorted(split, key=lambda e: list(body).index(e)) + order[k + 1:])
    groups.sort(key=lambda g: min(list(body).index(e) for e in g))
    heading = _heading(n.heading, mk) if n.heading and not _has_tables_heading(ms, n.refs_end) else None
    # Tables already at the end keep their place; a moved table goes before the first of them with a higher number.
    later = _end_tables(ms, n.refs_end, groups)
    sect = body.find(qn("w:sectPr"))
    first = True
    for els in groups:
        num = _block_number(els)
        dest = next((start for k, start in later if num is not None and k is not None and k > num), None)
        moved = _detach(els, mk)
        if first and heading is not None:
            moved.insert(0, heading)
            first = False
        for e in moved:
            if dest is not None:
                dest.addprevious(e)
            elif sect is not None and sect.getparent() is body:
                sect.addprevious(e)
            else:
                body.append(e)
    return []


def _block_number(els) -> int | None:
    """The number of a table block from its caption (a paragraph, or the table's first row)."""
    for e in els:
        tr = e.find(qn("w:tr")) if e.tag == qn("w:tbl") else None
        t = paragraph_text(e) if e.tag == qn("w:p") else \
            " ".join(paragraph_text(p) for p in tr.iter(qn("w:p"))) if tr is not None else ""
        m = _CAPTION_START.match(t)
        if m:
            return _num(re.search(r"[\dIVX]+$", m.group(0).strip()).group(0))
    return None


def _end_tables(ms: Manuscript, refs_end_id: str, moving) -> list[tuple[int | None, object]]:
    """(number, first element) of the table blocks already after the reference list (caption above or in its row)."""
    body = ms.document.element.body
    skip = {id(e) for g in moving for e in g}
    el = _top(ms.by_id(refs_end_id).element, body)
    out = []
    x = el.getnext() if el is not None else None
    while x is not None:
        if x.tag == qn("w:tbl") and id(x) not in skip and not _is_deleted_tbl(x):
            start = x
            prev = x.getprevious()
            while prev is not None and prev.tag == qn("w:p") and _empty(prev):
                prev = prev.getprevious()
            if prev is not None and prev.tag == qn("w:p") and _CAPTION_START.match(paragraph_text(prev)) \
                    and not _is_deleted(prev):
                start = prev
            out.append((_block_number([start, x]), start))
        x = x.getnext()
    return out


def _is_deleted_tbl(tbl) -> bool:
    rows = tbl.findall(qn("w:tr"))
    return bool(rows) and all(r.find(f"{qn('w:trPr')}/{qn('w:del')}") is not None for r in rows)


def _has_tables_heading(ms: Manuscript, refs_end_id: str) -> bool:
    """Whether a "Tables" heading already stands after the reference list."""
    body = ms.document.element.body
    el = _top(ms.by_id(refs_end_id).element, body)
    nxt = el.getnext() if el is not None else None
    while nxt is not None:
        if nxt.tag == qn("w:p") and not _is_deleted(nxt) and _TABLES_HEADING.match(paragraph_text(nxt)) \
                and paragraph_text(nxt).strip():
            return True
        nxt = nxt.getnext()
    return False
