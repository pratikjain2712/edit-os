"""Duplicate references and renumbering by order of first citation (pre-editing guide, E.g. 3-5).

  E.g. 3  duplicate references are deleted and the list renumbered;
  E.g. 4  references first cited out of order are renumbered in the text and the list;
  E.g. 5  if any reference is never cited, nothing is renumbered: the uncited query goes instead.

The change model only edits text inside paragraphs, so the list is reordered in place: each
reference paragraph gets the entry that belongs at its position (with its label number), and the
paragraphs left over by deleted duplicates are emptied. A moved entry carries the formatting the
reference rules propose for it, since its new position's own reference edits no longer apply.

Citations in every form `rules.rule_citations` handles ([n], [2-5], (n), superscript) are
renumbered on top of that rule's proposal, then sorted and re-compressed ([4,5,6] -> [4-6]).
The rule is registered before "citations", so its proposals win where they overlap.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from . import provenance, refengine, references
from .changes import Change, Layer
from .docmodel import Manuscript
from .style import JournalStyle

QUERY_TEXT = {
    "references.renumbered": "References are renumbered in text and list for citation sequence.",
    "references.duplicate_removed": "References ({a} and {b}) are identical, so duplicate reference has been "
                                    "removed and references are renumbered in text and list.",
    "references.duplicate": "References ({a} and {b}) are identical. Kindly check.",
    "references.renumber_blocked": "References are not cited in sequential order, but the reference list could not "
                                   "be renumbered automatically ({why}). Kindly renumber the references in text and "
                                   "list.",
    "references.not_renumbered": "This paragraph in the reference list is not a numbered reference and was left out "
                                 "of the renumbering. Kindly check.",
}


def _key(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", refengine.fold(text))


def expand(inner: str) -> list[int] | None:
    out: list[int] = []
    for part in re.split(r"\s*,\s*", inner.strip()):
        r = re.fullmatch(r"(\d+)\s*[-–—]\s*(\d+)", part)
        if r:
            a, b = int(r.group(1)), int(r.group(2))
            if b <= a or b - a > 200:
                return None
            out += range(a, b + 1)
        elif part.isdigit():
            out.append(int(part))
        else:
            return None
    return out


def compress(nums: list[int], style: JournalStyle) -> str:
    """Sorted, unique numbers as a citation: runs of `compress_from`+ become ranges."""
    nums = sorted(set(nums))
    least = style.get("renumber.compress_from", 3)
    dash = style.get("citations.range_dash", "-")
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i + 1 >= least:
            parts.append(f"{nums[i]}{dash}{nums[j]}")
        else:
            parts += [str(n) for n in nums[i:j + 1]]
        i = j + 1
    return style.get("citations.multiple_delimiter", ",").join(parts)


# --- tables: where their citations stand in the citation order ---------------------------

_TABLE_NOTE = re.compile(r"^\s*(?:(?:notes?|sources?|adapted|abbreviations?|data from|reproduced|modified|reprinted)\b|"
                         r"[*†‡§¶#])", re.I)


def _outer_table(p) -> object | None:
    from docx.oxml.ns import qn

    outer = None
    for a in p.element.iterancestors(qn("w:tbl")):
        outer = a
    return outer


def tables(ms: Manuscript) -> list[tuple[int | None, list[str]]]:
    """(table number from its caption, ids of its paragraphs: cells and the notes under it) for
    each table, in document order."""
    from . import rules_figures

    paras = ms.paragraphs
    groups: list[tuple[object, list[int]]] = []
    for i, p in enumerate(paras):
        if not p.in_table:
            continue
        t = _outer_table(p)
        if groups and groups[-1][0] is t:
            groups[-1][1].append(i)
        else:
            groups.append((t, [i]))

    def caption_number(i: int) -> int | None:
        m = rules_figures.caption_match(paras[i])
        if m and m.group("word").lower().startswith("tab"):
            return int(m.group("num")) if m.group("num").isdigit() else rules_figures.roman_to_int(m.group("num"))
        return None

    out = []
    for _, idxs in groups:
        first, last = idxs[0], idxs[-1]
        number = None
        k, seen = first - 1, 0  # caption above (the usual place), within a few lines
        while k >= 0 and seen < 3 and not paras[k].in_table:
            if paras[k].text.strip():
                number = caption_number(k)
                seen += 1
                if number is not None:
                    break
            k -= 1
        if number is None:  # caption in the first cell
            number = next((n for i in idxs[:3] if (n := caption_number(i)) is not None), None)
        members = [paras[i].id for i in idxs]
        k = last + 1
        while k < len(paras) and not paras[k].in_table:  # notes under the table
            text = paras[k].text
            if text.strip():
                if not _TABLE_NOTE.match(text):
                    if number is None:
                        number = caption_number(k)  # caption below the table
                    break
                members.append(paras[k].id)
            k += 1
        out.append((number, members))
    return out


def reading_keys(ms: Manuscript) -> dict[str, tuple]:
    """Sort key (paragraph position, offset, ...) for each paragraph whose citations count: a
    table's cells and notes stand at the end of the table's first callout in the text
    ("... role.[7] Tables 1-3 summarise", as the editors renumber them), or where the table is
    when it is never called out. Combine with a citation's offset: `citation_key`."""
    from . import rules_figures

    cached = getattr(ms, "_reading_keys", None)
    if cached is not None:
        return cached
    index = {p.id: i for i, p in enumerate(ms.paragraphs)}
    ms._para_index = index
    callout: dict[int, tuple[int, int]] = {}
    for c in rules_figures.citations(ms):
        if c.kind == rules_figures.TABLE:
            for n in c.numbers:
                callout.setdefault(n, (index[c.para_id], c.end))
    keys: dict[str, tuple] = {}
    for t_order, (number, members) in enumerate(tables(ms)):
        anchor = callout.get(number) if number is not None else None
        for j, pid in enumerate(members):
            keys[pid] = (anchor[0], anchor[1], 1, t_order, j) if anchor else (index[pid], 0, 1, t_order, j)
    ms._reading_keys = keys
    return keys


def citation_key(ms: Manuscript, para_id: str, start: int) -> tuple:
    keys = reading_keys(ms)
    if para_id in keys:
        return keys[para_id] + (start,)
    return (ms._para_index[para_id], start, 0, 0, 0, 0)


def citation_paragraphs(ms: Manuscript):
    """Paragraphs whose citations count, in citation order (by paragraph; use `citation_key` for
    the order of single citations): the body before the reference list, and tables with their
    notes wherever they are."""
    from . import rules

    keys = reading_keys(ms)
    ref_paras = {r.para_id for r in references.extract(ms)}
    body = {p.id for p in rules.body_paragraphs(ms)}
    paras = [p for p in ms.paragraphs if p.text.strip() and p.id not in ref_paras and (p.id in body or p.id in keys)]
    return sorted(paras, key=lambda p: citation_key(ms, p.id, 0))


@dataclass
class Occurrence:
    para_id: str
    start: int
    end: int
    old: str  # text in the paragraph
    new: str  # after the citation rule (bracketed, punctuation placed), before renumbering
    numbers: list[int]


@dataclass
class Plan:
    refs: list[references.Reference] = field(default_factory=list)
    occurrences: list[Occurrence] = field(default_factory=list)
    duplicates: list[tuple[int, int]] = field(default_factory=list)  # (kept, duplicate), old numbers
    uncited: list[int] = field(default_factory=list)
    mapping: dict[int, int] = field(default_factory=dict)  # old number -> new number
    apply: bool = False  # renumber (reorder and/or drop duplicates)
    skipped: list[references.Reference] = field(default_factory=list)  # unnumbered paragraphs left in place
    blocked: str = ""  # why a list cited out of order could not be renumbered (queried)

    @property
    def reordered(self) -> bool:
        return any(k != v for k, v in self.mapping.items())


def _formula_superscript(text: str, start: int, end: int) -> bool:
    """Z²P, d², η²H: a power in a formula, not a citation."""
    nxt = text[end:end + 1]
    if text[start - 1:start] in (".", ",", ";", ":") and start >= 2 and text[start - 2].isalpha():
        return False  # after a word's punctuation: "manner.⁴Materials" (a run-on heading) is a citation
    if nxt.isalnum():
        return True
    if start >= 1 and text[start - 1].isalpha() and (start < 2 or not text[start - 2].isalpha()):
        return True
    return "=" in text[max(0, start - 15):start]


def occurrences(ms: Manuscript, style: JournalStyle, n_refs: int) -> list[Occurrence]:
    """Every numbered citation, in citation order: the forms the citation rule converts (its decision
    on round brackets included: `rules.round_citations`), plus [n] citations anywhere outside the
    reference list (tables, legends). Square-bracket numbers beyond the list count too: a citation
    with no entry stops renumbering and is queried."""
    from . import rules

    proposals: dict[str, list[Change]] = {}
    for c in rules.rule_citations(ms, style):
        proposals.setdefault(c.para_id, []).append(c)
    ref_paras = {r.para_id for r in references.extract(ms)}
    heading = references.find_heading(ms)
    front = rules.front_matter_ids(ms)  # author and affiliation numbers
    out = []
    for idx, p in enumerate(ms.paragraphs):
        if p.id in ref_paras or idx == heading or not p.text.strip():
            continue
        if p.id in front and len(p.text.split()) < 12:
            continue  # title-page lines: author and affiliation numbers (a running sentence there still cites)
        found = []
        mine = proposals.get(p.id, [])
        for c in mine:
            # A converted round or superscript citation; its punctuation move ("(1)." -> ".") carries no number.
            if not c.old.lstrip().startswith(("(", "[")) and not any(
                    d.end == c.start and d.old.lstrip().startswith("(") for d in mine) \
                    and _formula_superscript(p.text, c.start, c.end):
                continue
            m = re.search(r"\[([^\]]*)\]", c.new)
            if m and (nums := expand(m.group(1))):
                found.append(Occurrence(p.id, c.start, c.end, c.old, c.new, nums))
        for m in rules._SQUARE.finditer(p.text):
            if any(c.start < m.end() and m.start() < c.end for c in mine):
                continue
            if rules.not_citation_before(p.text[:m.start()].rstrip(), False):
                continue
            nums = expand(m.group(2))
            if nums and min(nums) > 0:
                found.append(Occurrence(p.id, m.start(), m.end(), m.group(0), m.group(0), nums))
        out += found
    # Tables and their notes count where they are first called out (`reading_keys`).
    return sorted(out, key=lambda o: citation_key(ms, o.para_id, o.start))


def duplicates(ms: Manuscript, style: JournalStyle, refs: list[references.Reference]) -> list[tuple[int, int]]:
    """(kept, duplicate) pairs: same text once normalised, or the same verified PMID / DOI."""
    from . import rules

    ids: dict[str, set[str]] = {}
    for ref, res in rules.lookup_references(ms, style):
        if res.record is not None and res.status in (refengine.MATCHED, refengine.FORMAT_FIXED):
            rec = res.record
            ids[ref.para_id] = ({f"pmid:{rec.pmid}"} if rec.pmid else set()) | \
                               ({f"doi:{rec.doi.lower()}"} if rec.doi else set())
    seen: dict[str, int] = {}
    out = []
    for r in refs:
        keys = {k for k in [_key(r.text)] if len(k) >= 25} | ids.get(r.para_id, set())
        first = next((seen[k] for k in keys if k in seen), None)
        if first is not None:
            out.append((first, r.number))
        else:
            for k in keys:
                seen[k] = r.number
    return out


def plan(ms: Manuscript, style: JournalStyle) -> Plan:
    cached = getattr(ms, "_renumber_plan", None)
    if cached is not None:
        return cached
    pl = Plan()
    ms._renumber_plan = pl
    if not style.get("renumber.enabled", True):
        return pl
    every = references.extract(ms)
    refs = [r for r in every if r.number is not None]
    if not refs or len(refs) * 2 < len(every):
        return pl  # no list, or an author-year one
    if [r.number for r in refs] != list(range(1, len(refs) + 1)):
        # labels out of sequence: leave it to the editor, but say so when the citations need renumbering
        pl.occurrences = occurrences(ms, style, len(refs))
        first = []
        for o in pl.occurrences:
            first += [n for n in sorted(o.numbers) if n not in first]
        if first != sorted(first):
            pl.blocked = "the list's numbers are not 1 to {}".format(len(refs))
        pl.occurrences = []
        return pl
    pl.refs = refs
    pl.skipped = [r for r in every if r.number is None]  # an unnumbered paragraph in a numbered list: left in place
    n = len(refs)
    pl.occurrences = occurrences(ms, style, n)
    if not pl.occurrences:
        return pl
    pl.duplicates = duplicates(ms, style, refs) if style.get("renumber.remove_duplicates", True) else []
    canon = {r.number: r.number for r in refs}
    for kept, dup in pl.duplicates:
        canon[dup] = canon[kept]
    order: list[int] = []
    for o in pl.occurrences:
        for num in sorted(o.numbers):
            c = canon.get(num)
            if c is None:
                return pl  # cites a reference that is not in the list
            if c not in order:
                order.append(c)
    removed = {d for _, d in pl.duplicates}
    pl.uncited = [r.number for r in refs if r.number not in removed and r.number not in order]
    if pl.uncited:
        return pl  # E.g. 5: query the missing citations, do not renumber
    if not style.get("renumber.reorder", True):
        order = sorted(order)
    new = {c: i + 1 for i, c in enumerate(order)}
    pl.mapping = {r.number: new[canon[r.number]] for r in refs}
    pl.apply = bool(pl.duplicates) or pl.reordered
    return pl


def _formatted_entries(ms: Manuscript, style: JournalStyle) -> dict[str, str]:
    """Each reference paragraph's text with the reference rules' own proposals applied."""
    from . import rules

    cs = rules.run_all(ms, style, only=["ref_lookup", "references"])
    out = {}
    for r in references.extract(ms):
        text = ms.by_id(r.para_id).text
        # (the label's own edits stay with the position: only the entry moves)
        mine = (c for c in cs.changes if c.para_id == r.para_id and c.start >= r.offset and c.reason != "ref_label")
        for c in sorted(mine, key=lambda c: c.start, reverse=True):
            text = text[:c.start] + c.new + text[c.end:]
        out[r.para_id] = text
    return out


def rule_renumber(ms: Manuscript, style: JournalStyle) -> list[Change]:
    from . import rules

    pl = plan(ms, style)
    if not pl.apply:
        return []
    out: list[Change] = []
    # Citations.
    for o in pl.occurrences:
        mapped = sorted({pl.mapping[n] for n in o.numbers})
        if mapped == sorted(set(o.numbers)) and not (len(set(o.numbers)) < len(o.numbers)):
            continue
        new = re.sub(r"\[[^\]]*\]", lambda m: f"[{compress(mapped, style)}]", o.new, count=1)
        if new != o.old:
            p = ms.by_id(o.para_id)
            c = rules._change(p, o.start, o.end, new, Layer.REFERENCES, "renumber")
            olds = ",".join(str(n) for n in o.numbers)
            news = compress(mapped, style)
            c.provenance = provenance.make("renumber", f"Citation {olds} → {news}", _why(pl),
                                           old=list(o.numbers), new=mapped)
            out.append(c)
    # Reference list: position k gets the entry numbered k.
    by_new: dict[int, references.Reference] = {}
    removed = {d for _, d in pl.duplicates}
    for r in pl.refs:
        if r.number not in removed:
            by_new[pl.mapping[r.number]] = r
    formatted = None
    for k, pos in enumerate(pl.refs, 1):
        p = ms.by_id(pos.para_id)
        src = by_new.get(k)
        if src is None:  # left over after deleting duplicates
            c = rules._change(p, 0, len(p.text), "", Layer.REFERENCES, "renumber")
            dups = ", ".join(f"{d} = {kept}" for kept, d in pl.duplicates)
            c.provenance = provenance.make("renumber", "Duplicate removed", f"duplicate references: {dups}" if dups else "")
            out.append(c)
            continue
        prov = provenance.make("renumber", f"Reference {src.number} → {k}", _why(pl), old=[src.number], new=[k])
        if pos.label:
            m = re.search(r"\d+", pos.label)
            if m and int(m.group(0)) != k:
                c = rules._change(p, m.start(), m.end(), str(k), Layer.REFERENCES, "renumber")
                c.provenance = prov
                out.append(c)
        if src is not pos:
            if formatted is None:
                formatted = _formatted_entries(ms, style)
            src_p = ms.by_id(src.para_id)
            full = formatted[src.para_id]
            tail = len(src_p.text) - (src.offset + len(src.text))  # trailing whitespace kept as is
            entry = full[src.offset:len(full) - tail]
            c = rules._change(p, pos.offset, pos.offset + len(pos.text), entry, Layer.REFERENCES, "renumber")
            c.provenance = prov
            out.append(c)
    return out


def _why(pl) -> str:
    if pl.duplicates:
        return "duplicate references merged; the list renumbered in order of first citation"
    return "references renumbered in order of first citation in the text"


def add_queries(ms: Manuscript, style: JournalStyle, qs) -> None:
    from . import queries

    queries.TEXT.update({k: v for k, v in QUERY_TEXT.items() if k not in queries.TEXT})
    pl = plan(ms, style)
    if pl.blocked:
        heading = references.find_heading(ms)
        qs.add(ms.paragraphs[heading].id, "references.renumber_blocked", why=pl.blocked)
    if not pl.refs:
        return
    if pl.apply:
        for r in pl.skipped:  # never skipped silently
            qs.add(r.para_id, "references.not_renumbered", r.offset, r.offset + len(r.text))
    if pl.apply and style.get("header.renumber.enabled", True):
        return  # the query and the old -> new table go at the top of the file (header.py), as the editors put them
    by_num = {r.number: r for r in pl.refs}
    for kept, dup in pl.duplicates:
        r = by_num[dup]
        check = "references.duplicate_removed" if pl.apply else "references.duplicate"
        qs.add(r.para_id, check, r.offset, r.offset + len(r.text), a=kept, b=dup)
    if pl.apply and pl.reordered and not pl.duplicates:
        heading = references.find_heading(ms)
        qs.add(ms.paragraphs[heading].id, "references.renumbered")
