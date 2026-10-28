"""Tables and material the author sent as other Word files of the submission, merged into the manuscript.

Authors send tables (and figures) as separate Word files -- "Revised Figure,Table .docx" (jehp_1683), "table apocrine
hidrocystoma IJO.docx" (IJO_505), "Table 1 and 2.docx" (jehp_997) -- and the editors merge them into the article.
Figures in such files are placed by figures_insert.py (their pictures and captions, `docx_figures`); this module
takes the tables:

  - each table of the file with a caption ("Table 2: ...", above or below it, or "Table 2." and its title on two lines)
    whose number the manuscript has no table caption for goes to the end of the article (G195: the tables at the end,
    after the figure captions), with its notes, as a tracked insertion (`DocImport`, one suggestion per file), the
    caption in a merged first row when the journal does that (figures.table_caption_row) and the label as the journal
    writes it ("Table 2: " / end period: figures.caption_separator, figures.table_caption_end_period);
  - a table the manuscript has already (the same number captioned there) is left out;
  - a table that cannot be copied safely (pictures, text boxes, embedded objects, footnote references) is not
    inserted: an author query names the file.

A questionnaire / instrument / appendix sent as its own file is supplementary material (S372: 'please add the heading
"Supplementary Material" after Acknowledgement and retain the supplementary material as it is'; owner, 2026-10-27):
the heading goes after the Acknowledgement section, the file's content under it when it copies safely (plain
paragraphs and tables), and a query names the file either way.

    docx_members(zip_path, manuscript_name)        the other Word files worth reading
    plan(ms, style, zip_path, outdir, manuscript_name) -> Plan(inserts, queries, warnings)
    render(ms, inserts, base_dir, statuses, ...)    place the accepted (tracked copy: also pending) imports
"""

from __future__ import annotations

import copy
import os
import re
import zipfile
from dataclasses import dataclass, field

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import provenance, queries as queries_mod, rules_figures
from .changes import Layer, Source, Status
from .docmodel import Manuscript, has_picture, paragraph_text
from .frontmatter import NewParagraph, _FIRST_PAGE
from .style import JournalStyle

REASON_TABLES = "supplied_tables"
REASON_SUPPLEMENT = "supplied_supplementary"
DIR = os.path.join("source", "supplied")

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"

# Files that are not part of the article whatever they hold.
_NOT_ARTICLE = re.compile(r"reply|respon|review|referee|rebuttal|comment|contribut|credit|declaration|\bforms?\b|cevap|"
                          r"consent|copyright|cover|letter|checklist|guideline|reporting|strobe|consort|prisma\s*checklist|"
                          r"ethic|approval|certificate|conflict|\bcoi\b|undertaking|author|manuscript|revision\s*notes?",
                          re.I)
# Supplementary material sent as its own file.
# A questionnaire / study instrument sent as its own file: supplementary material to merge (owner, 2026-10-27).
# Owner, 2026-10-27 (later the same day): files named Supplementary / Appendix / Annexure are merged the same way.
_SUPPLEMENT = re.compile(r"questionnaire|instrument|\bsurvey\b|interview\s*guide|data\s*collection\s*tool|proforma|"
                         r"pro-forma|supplement|appendix|annex", re.I)
# A file the author names as supplementary material ("Supplementary File 1", "Appendix 1"). The editors left these
# out of the manuscript (jehp 1153, 1193, 1486, 1580, 1772, 1787), but the owner chose S372 (2026-10-27): Word files
# are merged under "Supplementary Material" like a questionnaire (_SUPPLEMENT); other formats get an editor note.
_SUPPLEMENTARY_FILE = re.compile(r"supplement|appendix|annex", re.I)

QUERY_TEXT = {
    "supplied.table_not_inserted": "Table {n} is provided in the separate file “{file}” but could not be placed in "
                                   "the article. Kindly check and provide it in editable form.",
    "supplied.table_no_caption": "The separate file “{file}” has a table without a caption. Kindly provide the "
                                 "caption and cite the table in the text.",
    "supplied.supplementary_included": "Please note that the “{what}” provided in the separate file “{file}” has been "
                                       "included as Supplementary Material. Kindly check.",
    "supplied.supplementary_file": "The “{what}” is provided in the separate file “{file}”. Kindly check and provide "
                                   "it for inclusion as Supplementary Material.",
}


@dataclass
class DocImport(NewParagraph):
    """Body elements of another Word file of the submission to put into the manuscript.

    `blocks`: per block the indexes of the file's body elements (paragraphs / tables) in order; `captions`: element
    index -> the caption as it should read; `kind`: "tables" (to the end) or "supplementary" (after the anchor, under
    `heading`)."""
    kind: str = "tables"
    file: str = ""  # path of the saved copy, relative to the output folder
    name: str = ""  # the file's name in the submission
    blocks: list[list[int]] = field(default_factory=list)
    captions: dict = field(default_factory=dict)
    heading: str = ""
    caption_row: bool = False

    def to_dict(self) -> dict:
        return {**super().to_dict(), "import_file": self.file, "kind": self.kind, "name": self.name,
                "blocks": self.blocks, "captions": {str(k): v for k, v in self.captions.items()},
                "heading": self.heading, "caption_row": self.caption_row}

    @classmethod
    def from_dict(cls, d: dict) -> "DocImport":
        return cls(id=d["id"], anchor=d["anchor"], where=d.get("where", "after"),
                   runs=[(t, f) for t, f in d.get("runs") or []], role=d.get("role", "import"),
                   status=Status(d.get("status", "suggested")), source=Source(d.get("source", "rule")),
                   reason=d.get("reason", REASON_TABLES), provenance=d.get("provenance") or None,
                   layer=Layer(d.get("layer", Layer.TABLES_FIGURES.value)), kind=d.get("kind", "tables"),
                   file=d.get("import_file", ""), name=d.get("name", ""),
                   blocks=[list(b) for b in d.get("blocks") or []],
                   captions={int(k): v for k, v in (d.get("captions") or {}).items()},
                   heading=d.get("heading", ""), caption_row=bool(d.get("caption_row")))


def is_import(n) -> bool:
    return isinstance(n, DocImport)


@dataclass
class Plan:
    inserts: list[DocImport] = field(default_factory=list)
    queries: list[queries_mod.Query] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)


# --- the submission's other Word files ----------------------------------------------------------

def docx_members(zip_path: str, manuscript_name: str = "") -> list[str]:
    """Names of the zip's Word files other than the manuscript and the title page, not named as correspondence,
    forms or declarations."""
    try:
        z = zipfile.ZipFile(zip_path)
    except (OSError, zipfile.BadZipFile):
        return []
    out = []
    for n in z.namelist():
        base = os.path.basename(n)
        if not base.lower().endswith(".docx") or n.startswith("__MACOSX") or base.startswith((".", "~$")):
            continue
        if base == manuscript_name or _FIRST_PAGE.search(base):
            continue
        stem = os.path.splitext(base)[0]
        if _NOT_ARTICLE.search(stem) and not _SUPPLEMENT.search(stem) and not re.search(r"\b(?:tables?|fig\w*)\b", stem, re.I) \
                or _SUPPLEMENTARY_FILE.search(stem) and not _SUPPLEMENT.search(stem):
            continue
        out.append(n)
    return out


def load_member(zip_path: str, name: str, outdir: str | None) -> tuple[Manuscript, str] | None:
    """The member read as a Manuscript, saved under OUTDIR/source/supplied (path relative to OUTDIR)."""
    import tempfile

    from . import docmodel

    data = zipfile.ZipFile(zip_path).read(name)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.basename(name))[:80] or "file.docx"
    rel = os.path.join(DIR, safe)
    if outdir:
        os.makedirs(os.path.join(outdir, DIR), exist_ok=True)
        path = os.path.join(outdir, rel)
    else:
        path = os.path.join(tempfile.mkdtemp(), safe)
    with open(path, "wb") as f:
        f.write(data)
    try:
        return docmodel.load(path), rel
    except Exception:  # not a Word file python-docx can read
        return None


def _unsafe(el) -> str:
    """Why an element cannot be copied into another document ("" when it can)."""
    if el.find(f".//{_W}txbxContent") is not None or el.find(f".//{_MC}") is not None:
        return "text boxes"
    if has_picture(el) or el.find(f".//{_W}drawing") is not None or el.find(f".//{_W}pict") is not None:
        return "pictures"
    if el.find(f".//{_W}object") is not None:
        return "embedded objects"
    if el.find(f".//{_W}footnoteReference") is not None or el.find(f".//{_W}endnoteReference") is not None:
        return "footnotes"
    return ""


def _empty(el) -> bool:
    return el.tag == f"{_W}p" and not paragraph_text(el).strip() and not _unsafe(el)


# Next to a table in a file of tables, "Table 1 detailed description of ..." is its caption whatever its case (IJO_505).
_LOOSE_CAPTION = re.compile(r"^\s*(?P<word>Tables?|TABLES?)\s*(?P<num>\d{1,2}|[IVX]{1,5})(?![\w])\s*[.:\-–—]?\s*(?P<rest>\S.*)$",
                            re.S)


def _caption(sms: Manuscript, el):
    """The caption match of a body paragraph element next to a table (rules_figures', else a looser one), or None."""
    p = next((q for q in sms.paragraphs if q.element is el), None)
    if p is None:
        return None, None
    return p, rules_figures.caption_match(p) or _LOOSE_CAPTION.match(p.text)


def _table_blocks(sms: Manuscript) -> list[dict]:
    """The file's tables: {"number", "caption" (element index or None), "split" (title index), "table", "notes",
    "side"}."""
    body = sms.document.element.body
    kids = list(body)
    split = {lab.element: title.element for lab, title in rules_figures.split_table_captions(sms)}
    note_ids = {q.element for block in rules_figures.table_notes(sms) for q in block}
    out = []
    used = set()
    for k, el in enumerate(kids):
        if el.tag != f"{_W}tbl":
            continue
        blk = {"table": k, "caption": None, "split": None, "notes": [], "number": None, "side": ""}
        j = k - 1
        while j >= 0 and _empty(kids[j]):
            j -= 1
        if j >= 0 and kids[j].tag == f"{_W}p" and j not in used:
            p, m = _caption(sms, kids[j])
            prev = j - 1
            while prev >= 0 and _empty(kids[prev]):
                prev -= 1
            if m is not None and m.group("word").lower().startswith("tab"):
                blk.update(caption=j, side="above", number=_num(m.group("num")))
            elif prev >= 0 and kids[prev] in split and split[kids[prev]] is kids[j]:
                lab = rules_figures._SPLIT_LABEL.match(paragraph_text(kids[prev]))
                blk.update(caption=prev, split=j, side="above", number=_num(lab.group("num")))
        last = k
        if blk["caption"] is None:  # the caption under the table (IJO_505)
            j = k + 1
            while j < len(kids) and _empty(kids[j]):
                j += 1
            if j < len(kids) and kids[j].tag == f"{_W}p":
                p, m = _caption(sms, kids[j])
                if m is not None and m.group("word").lower().startswith("tab"):
                    blk.update(caption=j, side="below", number=_num(m.group("num")))
                    last = j
        for j in range(last + 1, min(len(kids), last + 5)):  # its notes
            x = kids[j]
            if x.tag != f"{_W}p" or _unsafe(x):
                break
            t = paragraph_text(x)
            if not t.strip():
                continue
            p, m = _caption(sms, x)
            if m is not None or rules_figures._SPLIT_LABEL.match(t):
                break
            if x in note_ids or rules_figures.loose_table_note(t) or rules_figures._FOOTNOTE.match(t.strip()):
                blk["notes"].append(j)
                continue
            break
        used.update(i for i in (blk["caption"], blk["split"]) if i is not None)
        out.append(blk)
    return out


def _num(s: str) -> int | None:
    return int(s) if s.isdigit() else rules_figures.roman_to_int(s.upper())


def _caption_text(sms: Manuscript, style: JournalStyle, blk: dict) -> str:
    """The caption as the journal writes it: "Table 2: Title" (separator, case and end period of the journal), from
    the caption rules run on the file itself."""
    kids = list(sms.document.element.body)
    cap_el = kids[blk["caption"]]
    if blk["split"] is not None:
        changes = rules_figures._split_caption_changes(sms, style)
        lab = next((p for p in sms.paragraphs if p.element is cap_el), None)
        for c in changes:
            if lab is not None and c.para_id == lab.id and c.new:
                return lab.text[:c.start] + c.new
        return paragraph_text(cap_el).strip() + " " + paragraph_text(kids[blk["split"]]).strip()
    p = next(q for q in sms.paragraphs if q.element is cap_el)
    if rules_figures.caption_match(p) is None:  # "Table 1 detailed description ...": the label and case set here
        m = _LOOSE_CAPTION.match(p.text)
        rest = m.group("rest").strip()
        rest = rest[:1].upper() + rest[1:]
        mode = style.get("figures.table_caption_end_period") or "keep"
        if mode == "delete" and rules_figures._end_period(rest) is not None:
            rest = rest[:rules_figures._end_period(rest)].rstrip()
        elif mode == "add" and not rest.endswith((".", "?", "!")):
            rest += "."
        return f"Table {_num(m.group('num'))}{style.get('figures.caption_separator', ': ')}{rest}"
    try:
        edits = [(c.start, c.end, c.new) for c in rules_figures.rule_captions(sms, style)
                 if c.para_id == p.id and not c.format_only and " " not in c.new]
    except Exception:  # the caption as typed
        edits = []
    text = p.text
    for s, e, new in sorted(edits, reverse=True):
        text = text[:s] + new + text[e:]
    return text.strip()


# --- planning ---------------------------------------------------------------------------------

def _query(para_id: str, key: str, **fmt) -> queries_mod.Query:
    return queries_mod.Query(para_id, 0, 0, QUERY_TEXT[key].format(**fmt), key)


def _manuscript_tables(ms: Manuscript) -> set[int]:
    """Table numbers the manuscript has a caption for (a line, a first row, or "Table 1." over its title)."""
    nums = {c.number for c in rules_figures.captions(ms) if c.kind == rules_figures.TABLE}
    for lab, _ in rules_figures.split_table_captions(ms):
        m = rules_figures._SPLIT_LABEL.match(lab.text)
        if m and _num(m.group("num")):
            nums.add(_num(m.group("num")))
    return nums


def _what(name: str) -> str:
    """What the file is, for the query: "Questionnaire", "Appendix", ..."""
    stem = os.path.splitext(os.path.basename(name))[0]
    for word in ("questionnaire", "interview guide", "proforma", "survey", "instrument", "appendix", "annexure", "annex",
                 "supplementary material", "supplement"):
        if re.search(word, stem, re.I):
            return word[:1].upper() + word[1:]
    return "Supplementary material"


def _supplement_anchor(ms: Manuscript) -> tuple[str, str] | None:
    """Where "Supplementary Material" goes (S372: after Acknowledgement): after the Acknowledgement section; with none,
    before the reference list."""
    from . import references

    paras = ms.paragraphs
    ref = references.find_heading(ms)
    end = ref if ref is not None else len(paras)
    ack = re.compile(r"^\s*(?:<H\d>)?\s*acknowledge?ments?\b", re.I)
    for i in range(end - 1, -1, -1):
        p = paras[i]
        if p.in_table or not ack.match(p.text):
            continue
        last = i
        if len(p.text.split()) <= 3:  # a heading: its text follows, up to the next heading-like line
            for j in range(i + 1, end):
                q = paras[j]
                t = q.text.strip()
                if q.in_table or not t:
                    continue
                if len(t.split()) <= 6 and not t.endswith((".", ";", ",")) or t.startswith("<H"):
                    break
                last = j
        return paras[last].id, "after"
    if ref is not None:
        return paras[ref].id, "before"
    return None


def plan(ms: Manuscript, style: JournalStyle, zip_path: str | None, outdir: str | None,
         manuscript_name: str = "") -> Plan:
    out = Plan()
    if not style.get("supplied.enabled", True) or not zip_path or not zipfile.is_zipfile(zip_path):
        return out
    from . import legends, references

    have = _manuscript_tables(ms)
    ref = references.find_heading(ms)
    end_i = legends.refs_end(ms)
    end_id = ms.paragraphs[end_i].id if end_i is not None else (ms.paragraphs[-1].id if ms.paragraphs else "")
    sep_mode = style.get("figures.table_caption_row", False)
    n_imp = 0
    for name in zipfile.ZipFile(zip_path).namelist():
        stem = os.path.splitext(os.path.basename(name))[0]
        if stem and _SUPPLEMENTARY_FILE.search(stem) and not name.endswith("/") and \
                not name.lower().endswith(".docx") and not name.startswith("__MACOSX"):
            out.warnings.append({"para_id": "", "rule": "supplied", "message": f"{os.path.basename(name)}: a "
                                 "supplementary file that is not a Word file, kept separate (not merged)."})
    for name in docx_members(zip_path, manuscript_name):
        base = os.path.basename(name)
        got = load_member(zip_path, name, outdir)
        if got is None:
            out.warnings.append({"para_id": "", "rule": "supplied", "message": f"{base}: not readable as a Word file."})
            continue
        sms, rel = got
        if _SUPPLEMENT.search(os.path.splitext(base)[0]) and not re.search(r"\btables?\b", base, re.I):
            _plan_supplement(ms, style, sms, rel, base, out)
            continue
        blocks = _table_blocks(sms)
        if not blocks:
            continue
        kids = list(sms.document.element.body)
        take, captions = [], {}
        for blk in blocks:
            n = blk["number"]
            if n is None:  # (asked only of a file named for its tables: other files' tables are not the article's)
                if re.search(r"tab", base, re.I) and any(b["number"] is not None for b in blocks) or \
                        re.search(r"\btables?\b", base, re.I):
                    out.queries.append(_query(end_id, "supplied.table_no_caption", file=base))
                continue
            if n in have:
                continue  # the manuscript has this table (the file repeats it)
            idx = ([blk["caption"]] if blk["side"] == "above" else []) + \
                ([blk["split"]] if blk["split"] is not None else []) + [blk["table"]] + \
                ([blk["caption"]] if blk["side"] == "below" else []) + blk["notes"]
            why = next((_unsafe(kids[i]) for i in idx if _unsafe(kids[i])), "")
            if why:
                out.queries.append(_query(end_id, "supplied.table_not_inserted", n=n, file=base))
                out.warnings.append({"para_id": end_id, "rule": "supplied",
                                     "message": f"Table {n} in {base} not inserted: it holds {why}."})
                continue
            captions[blk["caption"]] = _caption_text(sms, style, blk)
            take.append((n, idx, blk))
            have.add(n)
        if not take:
            continue
        take.sort(key=lambda t: t[0])
        n_imp += 1
        what = ", ".join(f"Table {n}" for n, _, _ in take)
        out.inserts.append(DocImport(
            id=f"import{n_imp}", anchor=end_id, where="after",
            runs=[(f"[{what} from {base} put at the end of the article]", {})], role="import", reason=REASON_TABLES,
            provenance=provenance.make("submission_file", "Submission file", f"{base}: {what}", files=[base]),
            layer=Layer.TABLES_FIGURES, kind="tables", file=rel, name=base,
            blocks=[[i for i in idx if i != blk["split"]] for _, idx, blk in take],
            captions={k: v for k, v in captions.items() if any(k in idx for _, idx, _ in take)},
            heading=style.get("figures.tables_heading") or "", caption_row=bool(sep_mode)))
    return out


def _plan_supplement(ms: Manuscript, style: JournalStyle, sms: Manuscript, rel: str, base: str, out: Plan) -> None:
    heading = style.get("sections.back_matter.supplementary_heading") or ""
    if heading and style.get("headings.tag"):
        heading = f"<H2>{heading}"  # a back-matter heading (S: H2)
    what = _what(base)
    slot = _supplement_anchor(ms)
    if slot is None:
        return
    anchor, where = slot
    kids = [el for el in sms.document.element.body if el.tag in (f"{_W}p", f"{_W}tbl")]
    body = list(sms.document.element.body)
    idx = [body.index(el) for el in kids]
    while idx and _empty(body[idx[-1]]):
        idx.pop()
    while idx and _empty(body[idx[0]]):
        idx.pop(0)
    lines = [re.sub(r"\W+", " ", paragraph_text(body[i])).strip().lower() for i in idx if body[i].tag == f"{_W}p"]
    lines = [t for t in lines if len(t) > 15][:5]
    have = {re.sub(r"\W+", " ", p.text).strip().lower() for p in ms.paragraphs}
    if lines and all(t in have for t in lines):
        out.warnings.append({"para_id": anchor, "rule": "supplied", "message": f"{base}: the manuscript has this "
                             "material already; not merged again."})
        return
    unsafe = next((_unsafe(body[i]) for i in idx if _unsafe(body[i])), "")
    numbered = any(body[i].find(f".//{_W}numPr") is not None for i in idx)
    if not heading:
        out.warnings.append({"para_id": anchor, "rule": "supplied", "message": f"{base}: supplementary material sent "
                             "as a separate file; not merged (the journal sets no Supplementary Material heading)."})
        return
    blocks = [] if (unsafe or numbered or not idx) else [idx]
    out.inserts.append(DocImport(
        id=f"supplement{len(out.inserts) + 1}", anchor=anchor, where=where,
        runs=[(f"[{heading}: {what} from {base}" + ("" if blocks else " (heading only)") + "]", {})],
        role="import", reason=REASON_SUPPLEMENT,
        provenance=provenance.make("stylesheet", "Stylesheet", "S372: the heading “Supplementary Material” after "
                                   "Acknowledgement, the material retained as it is", files=[base]),
        layer=Layer.STRUCTURE, kind="supplementary", file=rel, name=base, blocks=blocks, heading=heading))
    key = "supplied.supplementary_included" if blocks else "supplied.supplementary_file"
    out.queries.append(_query(anchor, key, what=what, file=base))
    if not blocks:
        why = unsafe or ("numbered lists" if numbered else "no content")
        out.warnings.append({"para_id": anchor, "rule": "supplied",
                             "message": f"{base}: \"{heading}\" heading added; the content not merged ({why})."})


# --- rendering --------------------------------------------------------------------------------

def _copy_styles(ms: Manuscript, sms: Manuscript, els) -> None:
    """Styles the copied elements use that the manuscript lacks (a table's grid lines live in its style)."""
    try:
        target = ms.document.styles.element
        source = sms.document.styles.element
    except (AttributeError, KeyError, NotImplementedError):
        return
    have = {s.get(qn("w:styleId")) for s in target.iter(qn("w:style"))}
    by_id = {s.get(qn("w:styleId")): s for s in source.iter(qn("w:style"))}
    want = []
    for el in els:
        for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
            for x in el.iter(qn(tag)):
                want.append(x.get(qn("w:val")))
    while want:
        sid = want.pop()
        if not sid or sid in have or sid not in by_id:
            continue
        st = copy.deepcopy(by_id[sid])
        target.append(st)
        have.add(sid)
        for rel in ("w:basedOn", "w:link", "w:next"):
            b = st.find(qn(rel))
            if b is not None:
                want.append(b.get(qn("w:val")))


def _clean_copy(el):
    """A copy of a body element that can live in the manuscript: list numbering, bookmarks, comments and
    hyperlinks (their relationships belong to the other file) taken out, the text kept."""
    c = copy.deepcopy(el)
    for tag in ("w:numPr", "w:bookmarkStart", "w:bookmarkEnd", "w:commentRangeStart", "w:commentRangeEnd",
                "w:proofErr", "w:permStart", "w:permEnd"):
        for x in list(c.iter(qn(tag))):
            x.getparent().remove(x)
    for x in list(c.iter(qn("w:commentReference"))):
        r = x.getparent()
        r.getparent().remove(r)
    for h in list(c.iter(qn("w:hyperlink"))):
        parent = h.getparent()
        i = parent.index(h)
        for k in list(h):
            parent.insert(i, k)
            i += 1
        parent.remove(h)
    for x in c.iter():  # relationship ids of the other file
        for a in [a for a in x.attrib if a.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")]:
            del x.attrib[a]
    return c


def _set_caption(p_el, text: str, mk, bold_label: bool = True) -> None:
    """Replace the paragraph's runs by the caption as it should read, the label bold (G260 / the sheets)."""
    from .header import styled

    first = p_el.find(f".//{_W}r")
    rpr = copy.deepcopy(first.find(qn("w:rPr"))) if first is not None and first.find(qn("w:rPr")) is not None else None
    for r in [x for x in p_el if x.tag in (qn("w:r"), qn("w:ins"), qn("w:del"), qn("w:smartTag"), qn("w:fldSimple"))]:
        p_el.remove(r)
    m = re.match(r"\s*(?:Table|Figure)\s+[\dIVX]+[A-Za-z]?\s*[:.]\s*", text)
    pieces = [(text[:m.end()], True), (text[m.end():], None)] if m and bold_label else [(text, None)]
    for t, bold in pieces:
        if not t:
            continue
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        styled(r, mk, {"bold": True} if bold else None)
        x = OxmlElement("w:t")
        x.text = t
        x.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(x)
        p_el.append(r)


def _normalise(el, mk, table_align: str | None, align: str | None) -> None:
    """Give imported content the journal's font and size (`mk.run_format`, formatting.fonts) like every other run of
    the article: the other file's fonts (Calibri 10 pt, Aptos) never reach the deliverable. Runs in a symbol font keep
    it (their letters are Greek / dingbats) and take the size. Table cells take `fonts.table_align` (G193-194), other
    paragraphs the journal's alignment."""
    from .rules_format import _SYMBOL_FONT, _PLAIN_CHAR
    from .tracked import _ppr_insert, apply_format

    fmt = dict(getattr(mk, "run_format", None) or {})
    if not fmt and not table_align and not align:
        return
    for p in el.iter(qn("w:p")):
        in_cell = any(a.tag == qn("w:tc") for a in p.iterancestors())
        want = table_align if in_cell else align
        if want:
            ppr = p.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                p.insert(0, ppr)
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc")
                _ppr_insert(ppr, jc)
            jc.set(qn("w:val"), want)
        if not fmt:
            continue
        ppr = p.find(qn("w:pPr"))
        mark = ppr.find(qn("w:rPr")) if ppr is not None else None
        if mark is not None:
            apply_format(mark, fmt)
        for r in p.iter(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                r.insert(0, rpr)
            f = rpr.find(qn("w:rFonts"))
            font = (f.get(qn("w:ascii")) or f.get(qn("w:hAnsi")) or "") if f is not None else ""
            text = "".join(t.text or "" for t in r.iter(qn("w:t")))
            symbol = bool(_SYMBOL_FONT.search(font)) and not (re.search(r"math", font, re.I) and text
                                                               and all(_PLAIN_CHAR.match(c) for c in text))
            apply_format(rpr, {k: v for k, v in fmt.items() if not (symbol and k == "font")})


def _caption_into_row(tbl, cap_el, mk) -> None:
    from .legends import _grid_columns

    cols = _grid_columns(tbl)
    tr = OxmlElement("w:tr")
    tr.append(OxmlElement("w:trPr"))
    tc = OxmlElement("w:tc")
    tcpr = OxmlElement("w:tcPr")
    if cols > 1:
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), str(cols))
        tcpr.append(span)
    tc.append(tcpr)
    tc.append(cap_el)
    tr.append(tc)
    first = tbl.find(qn("w:tr"))
    (first.addprevious(tr) if first is not None else tbl.append(tr))


def render(ms: Manuscript, inserts, base_dir: str, statuses, *, tracked: bool, author: str, date: str,
           run_format: dict | None = None, align: str | None = None, style: JournalStyle | None = None) -> list[str]:
    """Place the imports whose status is in `statuses` (after tracked.render and the inserts, before the legend and
    table moves). Returns notes on what could not be placed."""
    from . import docmodel
    from .header import _Marks, _mark_block
    from .legends import _heading, _has_tables_heading

    notes: list[str] = []
    todo = [n for n in inserts if is_import(n) and n.status in statuses]
    if not todo:
        return notes
    mk = _Marks(ms, author, date, tracked, run_format, align)
    body = ms.document.element.body
    sect = body.find(qn("w:sectPr"))
    headed = False
    supp_headed = False  # one "Supplementary Material" heading over all the supplementary files merged
    tail: dict[str, object] = {}  # anchor id -> the last element placed after it (the files keep their order)
    for n in todo:
        path = os.path.join(base_dir, n.file)
        try:
            sms = docmodel.load(path)
            anchor = ms.by_id(n.anchor).element
        except Exception:
            notes.append(f"{n.text}: not done (the file or its place is gone)")
            continue
        skids = list(sms.document.element.body)
        place: list = []
        if n.kind == "supplementary" and n.heading and not supp_headed:
            place.append(_heading(n.heading, mk))
        for block in n.blocks:
            els = []
            for i in block:
                if i >= len(skids):
                    continue
                c = _clean_copy(skids[i])
                if i in n.captions:
                    _set_caption(c, n.captions[i], mk)
                els.append(c)
            _copy_styles(ms, sms, els)
            if n.kind == "tables" and n.caption_row:
                tbl = next((e for e in els if e.tag == qn("w:tbl")), None)
                caps = [e for k, e in zip([i for i in block if i < len(skids)], els) if k in n.captions]
                if tbl is not None and caps:
                    els = [e for e in els if e is not caps[0]]
                    _caption_into_row(tbl, caps[0], mk)
            place += els
        if not place:
            continue
        if n.kind == "supplementary":
            supp_headed = supp_headed or bool(n.heading)
        cfg = (style.get("formatting.fonts") or {}) if style is not None and run_format else {}
        for e in place:
            _normalise(e, mk, cfg.get("table_align") or None, align)
        if tracked:
            for e in place:
                if e.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}") is None:
                    _mark_block(e, mk, "w:ins")
        if n.kind == "tables":
            if n.heading and not headed and not _has_tables_heading(ms, n.anchor):
                place.insert(0, _heading(n.heading, mk))
            headed = headed or bool(n.heading)  # one heading over all the imported tables
            for e in place:
                (sect.addprevious(e) if sect is not None and sect.getparent() is body else body.append(e))
        else:
            at = anchor
            if n.where == "before":
                for e in place:
                    anchor.addprevious(e)
            else:
                at = tail.get(n.anchor, at)
                while at.getparent() is not body:
                    at = at.getparent()
                for e in place:
                    at.addnext(e)
                    at = e
                tail[n.anchor] = at
    return notes
