"""Table body layout from the pre-editing guide (TABLE, G193-G209) and the stylesheets ("Table Body"):

  - indentation in the stub column (G208, S "Table Body"): "In case clear/more indentation is present for cells then
    the indentation should be removed and 5 spaces (and increment of 5 spaces) should be added in its place". The
    first column's leading spaces / tabs and its direct paragraph indentation are read as levels (the distinct amounts,
    smallest first); each level becomes 5 typed spaces more, the paragraph indentation is cleared (`rule_indent`);
  - rows under a sub-heading row (G209 "In case of headings present only for few rows then please add indentation
    where necessary"): a body row with text only in its first cell, followed by data rows, gets its rows indented by 5
    spaces -- only in a stub column the author did not indent at all (else the author's levels are kept);
  - empty paragraphs inside a cell go (G196 "... or deleted (without content)"): one `legends.BlockMove` record
    (kind "cell_blanks") the editor accepts or rejects (`plan_cell_blanks`, `render_cell_blanks`);
  - queries where the change cannot be made safely as a tracked change (`add_queries`): empty column-head cells to be
    merged horizontally (G206: Word has no tracked change for a cell merge), and rows whose cells hold their entries
    as separate paragraphs, which may be separate rows (G196: "paragraphs (with content) should be changed to separate
    rows" -- a value with its range on a second line is one row, so this is asked, not done).

Settings: `tables:` (base.yaml).
"""

from __future__ import annotations

import re
from collections import defaultdict

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .changes import Change, Layer
from .docmodel import Manuscript, Paragraph, has_picture, paragraph_text
from .style import JournalStyle

QUERY_TEXT = {
    # G206: "In case of empty cells, they should be merged horizontally and not vertically"
    "tables.empty_head_cells": "The empty cells in the column heads of this table should be merged horizontally. "
                               "Kindly check.",
    # G196: "Where necessary paragraphs (with content) should be changed to separate rows"
    "tables.cell_rows": "The entries in this row are typed as separate lines within the cells. Kindly check whether "
                        "they should be separate rows.",
}

_SPACE_TWIPS = 60  # a space at 12 pt Times New Roman is about 3 pt
_MIN_IND = 100  # direct indentation below this (twips) is padding, not an indent
_TAB = 12  # a tab to Word's default stop (0.5 inch, 720 twips), in spaces
_TOTAL = re.compile(r"^\s*(?:total|overall|grand total)\b", re.I)


def _on(style: JournalStyle, key: str):
    return style.get(f"tables.{key}", False)


# --- the tables' cells -------------------------------------------------------------------------

class _Grid:
    """One table: {row: {col: [paragraphs]}} of its own cells (not nested tables), the header rows and the tc's."""

    def __init__(self, tbl, paras: list[Paragraph]):
        self.tbl = tbl
        self.rows: dict[int, dict[int, list[Paragraph]]] = defaultdict(lambda: defaultdict(list))
        self.tcs: dict[tuple[int, int], object] = {}
        trs = tbl.findall(qn("w:tr"))
        for r, tr in enumerate(trs):
            for c, tc in enumerate(tr.findall(qn("w:tc"))):
                self.tcs[(r, c)] = tc
        by_el = {p.element: p for p in paras}
        for (r, c), tc in self.tcs.items():
            for p_el in tc.findall(qn("w:p")):
                if p_el in by_el:
                    self.rows[r][c].append(by_el[p_el])
        self.n_rows = len(trs)
        self.heads: set[int] = set()

    def cell_text(self, r: int, c: int) -> str:
        return " ".join(p.text for p in self.rows[r].get(c, [])).strip()

    def cols(self, r: int) -> list[int]:
        return sorted(self.rows[r])


def _grids(ms: Manuscript) -> list[_Grid]:
    """The data tables of the document (top-level tables, two rows and two columns at least), with their header
    rows as the formatting rules read them (rules_format._header_rows)."""
    from . import docview
    from .rules_format import _header_rows

    body = ms.document.element.body
    paras = [p for p in ms.paragraphs if p.in_table]
    out = []
    names = {}
    for i, tbl in enumerate(body.iter(qn("w:tbl")), 1):
        names[tbl] = f"t{i}"
        if any(a.tag == qn("w:tbl") for a in tbl.iterancestors()) or tbl.find(f".//{qn('w:tbl')}") is not None:
            continue
        out.append(_Grid(tbl, paras))
    if not out:
        return []
    classified = docview.classify(ms)
    by_id = {p.id: p for p in ms.paragraphs}
    for d in classified:
        d["el"] = by_id[d["id"]].element
    heads = _header_rows(classified)
    for g in out:
        g.heads = set(heads.get(names.get(g.tbl), set()))
    return [g for g in out if g.heads]


def _caption_row(g: _Grid, r: int) -> bool:
    from .rules_figures import caption_match

    cols = g.cols(r)
    return len(cols) == 1 and any(caption_match(p) is not None for p in g.rows[r][cols[0]])


def _body_rows(g: _Grid) -> list[int]:
    start = max(g.heads) + 1
    return [r for r in range(start, g.n_rows) if not _caption_row(g, r)]


# --- indentation -------------------------------------------------------------------------------

def _direct_indent(p: Paragraph) -> int:
    """The paragraph's own left + first-line indentation in twips (0 when none is set directly)."""
    ind = p.element.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    if ind is None:
        return 0

    def val(name):
        v = ind.get(qn(name))
        try:
            return int(v) if v is not None else 0
        except ValueError:
            return 0
    left = val("w:left") or val("w:start")
    return left + val("w:firstLine") - val("w:hanging")


def _numbered(p: Paragraph) -> bool:
    return p.element.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None


def _lead(text: str) -> str:
    return text[:len(text) - len(text.lstrip(" \t\u00a0"))]


def _width(p: Paragraph) -> float:
    """How far the paragraph's text starts in, in spaces: typed spaces, tabs and direct indentation."""
    lead = _lead(p.text)
    w = lead.count(" ") + lead.count("\u00a0") + _TAB * lead.count("\t")
    ind = _direct_indent(p)
    if ind >= _MIN_IND:
        w += ind / _SPACE_TWIPS
    return w


def _levels(widths: list[float]) -> dict[float, int]:
    """Distinct indentation amounts -> levels 1, 2, ... (amounts within a space or so of each other are one level)."""
    out, level, prev = {}, 0, None
    for w in sorted({w for w in widths if w > 0}):
        if prev is None or w - prev >= 2:
            level += 1
        out[w] = level
        prev = w
    return out


def _subheading_row(g: _Grid, r: int) -> bool:
    """A body row with text in its first cell only (or one cell across the table): a heading over the rows below."""
    cols = g.cols(r)
    if not cols or not g.cell_text(r, cols[0]):
        return False
    if any(has_picture(p.element) for p in g.rows[r][cols[0]]):
        return False
    return all(not g.cell_text(r, c) for c in cols[1:])


def rule_indent(ms: Manuscript, style: JournalStyle) -> list[Change]:
    """G208 / G209 in the stub (first) column of each table's body; settings `tables.indent`."""
    from .rules import _change

    cfg = style.get("tables.indent") or {}
    if not cfg.get("enabled"):
        return []
    step = int(cfg.get("spaces", 5))
    out: list[Change] = []
    for g in _grids(ms):
        rows = _body_rows(g)
        stub = {r: [p for p in g.rows[r].get(0, []) if p.text.strip() and not _numbered(p)] for r in rows}
        if not any(stub.values()):
            continue
        widths = {p.id: _width(p) for ps in stub.values() for p in ps}
        base = min(widths.values())  # the whole column set in alike (cell padding typed as spaces): no level
        # (a single stray space is no level either)
        widths = {pid: (w - base if w - base >= 2 else 0) for pid, w in widths.items()}
        levels = _levels(list(widths.values()))
        want: dict[str, int] = {pid: levels.get(w, 0) for pid, w in widths.items()}
        if not levels and cfg.get("subheading_rows", True):
            # No indentation in the stub at all: rows under a sub-heading row (a head given for a few rows only)
            # get one level.
            heading = False
            top = max(g.heads)
            heads_again = [g.cell_text(top, c) for c in g.cols(top)]  # the column heads repeated inside the table
            data_rows = [r for r in rows if not _subheading_row(g, r)]
            if data_rows and len(data_rows) < len(rows):
                for r in rows:
                    if _subheading_row(g, r):
                        nxt = [x for x in rows if x > r]
                        heading = bool(nxt) and not _subheading_row(g, nxt[0])
                        continue
                    if [g.cell_text(r, c) for c in g.cols(r)] == heads_again:
                        continue  # a repeated head row ("Coefficient | Estimate | CI") is not a row of the group
                    if _TOTAL.match(g.cell_text(r, 0)):
                        heading = False  # "Total" closes the group
                        continue
                    if heading:
                        for p in stub[r]:
                            want[p.id] = 1
        for ps in stub.values():
            for p in ps:
                level = want.get(p.id, 0)
                lead = _lead(p.text)
                new = " " * (step * level)
                if lead != new:
                    out.append(_change(p, 0, len(lead), new, Layer.TABLES_FIGURES, "tables.indent"))
                if _direct_indent(p) >= _MIN_IND:
                    c = _change(p, 0, len(p.text), p.text, Layer.TABLES_FIGURES, "tables.indent", {"indent": False})
                    out.append(c)
    return out


# --- queries -----------------------------------------------------------------------------------

def _empty_cell(g: _Grid, r: int, c: int) -> bool:
    tc = g.tcs.get((r, c))
    if tc is None:
        return False
    vm = tc.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}")
    if vm is not None and vm.get(qn("w:val"), "continue") != "restart":
        return False  # the lower part of a vertically merged cell
    return not g.cell_text(r, c) and not any(has_picture(p.element) for p in g.rows[r].get(c, []))


def add_queries(ms: Manuscript, style: JournalStyle, qs) -> None:
    from . import queries

    if not (_on(style, "empty_head_query") or _on(style, "cell_rows_query")):
        return
    queries.TEXT.update({k: v for k, v in QUERY_TEXT.items() if k not in queries.TEXT})
    for g in _grids(ms):
        anchor = next((p for r in sorted(g.heads) for c in g.cols(r) for p in g.rows[r][c] if p.text.strip()), None)
        if anchor is None:
            continue
        if _on(style, "empty_head_query"):
            # two or more empty cells side by side in a head row: one cell across them (G206)
            if any(_empty_cell(g, r, c) and _empty_cell(g, r, c + 1) for r in g.heads for c in g.cols(r)):
                qs.add(anchor.id, "tables.empty_head_cells")
        if _on(style, "cell_rows_query"):
            for r in _body_rows(g):
                counts = [len([p for p in g.rows[r][c] if p.text.strip()]) for c in g.cols(r)]
                # two cells or more with lines of their own, the filled cells holding as many lines (blank ones
                # keeping the values level with their labels): rows typed inside one row
                lines = {len(g.rows[r][c]) for c, n in zip(g.cols(r), counts) if n}
                if sum(1 for n in counts if n >= 2) >= 2 and len(lines) == 1:
                    first = next(p for c in g.cols(r) for p in g.rows[r][c] if p.text.strip())
                    qs.add(first.id, "tables.cell_rows")
                    break  # one query per table


# --- empty paragraphs inside cells (G196) --------------------------------------------------------

def _blank(p_el) -> bool:
    return not paragraph_text(p_el).strip() and not has_picture(p_el) and \
        not any(True for _ in p_el.iter(qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:fldChar"),
                                        qn("w:fldSimple"), qn("w:footnoteReference"), qn("w:sectPr"),
                                        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))


def cell_blanks(ms: Manuscript) -> list[str]:
    """Ids of the empty paragraphs inside table cells that also hold text (a cell keeps one paragraph)."""
    by_el = {p.element: p for p in ms.paragraphs}
    out = []
    for tc in ms.document.element.body.iter(qn("w:tc")):
        ps = tc.findall(qn("w:p"))
        if len(ps) < 2 or not any(paragraph_text(p).strip() for p in ps) or tc.find(qn("w:tbl")) is not None:
            continue
        # A row whose cells hold several lines each (rows typed as lines inside the cells): the blank lines keep the
        # values level with their labels, so they stay (the row is queried: tables.cell_rows)
        tr = tc.getparent()
        if tr is not None and any(sum(1 for p in other.findall(qn("w:p")) if paragraph_text(p).strip()) >= 2
                                  for other in tr.findall(qn("w:tc"))):
            continue
        out += [by_el[p].id for p in ps if _blank(p) and p in by_el]
    return out


def plan_cell_blanks(ms: Manuscript, style: JournalStyle):
    from . import provenance
    from .legends import BlockMove

    if not _on(style, "delete_empty_cell_paragraphs"):
        return []
    ids = cell_blanks(ms)
    if not ids:
        return []
    prov = provenance.make("guide", "Pre-editing guide", "G196: paragraphs without content inside table cells are "
                           "deleted", rule="table_cell_blanks")
    return [BlockMove(id="cell_blanks", anchor=ids[0], where="before",
                      runs=[(f"[{len(ids)} empty paragraph{'s' * (len(ids) > 1)} inside table cells deleted]", {})],
                      role="move", reason="table_cell_blanks", provenance=prov, layer=Layer.TABLES_FIGURES,
                      kind="cell_blanks", blocks=[ids])]


def render_cell_blanks(ms: Manuscript, n, mk) -> list[str]:
    """Delete the empty paragraphs: in the tracked copy their paragraph marks are tracked deletions (an empty last
    paragraph of a cell takes the paragraph before it instead: that one's mark is deleted, and the last one takes its
    properties, so the text keeps its formatting when accepted); in the clean copy they go."""
    from copy import deepcopy

    from .tracked import _PPR_TAIL

    done = 0
    for pid in n.blocks[0] if n.blocks else []:
        el = ms.by_id(pid).element
        tc = el.getparent()
        if tc is None or tc.tag != qn("w:tc") or not _blank(el):
            continue
        ps = [p for p in tc.findall(qn("w:p")) if not _gone(p, ms)]
        if len(ps) < 2 or el not in ps:
            continue
        if not mk.tracked:
            ms.removed.append(el)
            done += 1
            continue
        if el is not ps[-1]:
            _mark_deleted(el, mk)
            done += 1
            continue
        prev = next((p for p in reversed(ps[:ps.index(el)]) if not _deleted(p)), None)
        if prev is None:
            continue
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            el.insert(0, ppr)
        for child in list(ppr):
            if child.tag not in _PPR_TAIL:
                ppr.remove(child)
        src = prev.find(qn("w:pPr"))
        if src is not None:
            for k, child in enumerate(x for x in src if x.tag not in _PPR_TAIL):
                ppr.insert(k, deepcopy(child))
        _mark_deleted(prev, mk)
        done += 1
    return [] if done else [f"{n.text}: not done (the paragraphs are gone)"]


def _gone(p_el, ms: Manuscript) -> bool:
    return any(p_el is e for e in ms.removed)


def _deleted(p_el) -> bool:
    return p_el.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None


def _mark_deleted(p_el, mk) -> None:
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        tail = [c for c in ppr if c.tag in (qn("w:sectPr"), qn("w:pPrChange"))]
        if tail:
            tail[0].addprevious(rpr)
        else:
            ppr.append(rpr)
    if rpr.find(qn("w:del")) is None:
        m = mk.mark("w:del")
        ins = rpr.find(qn("w:ins"))
        if ins is not None:
            ins.addnext(m)
        else:
            rpr.insert(0, m)
