"""Figures supplied as separate files in the submission zip, inserted above their legends.

Authors often send figures as separate files ("Figure 1.jpg", "Fig2.tif", "1a.jpg") and only the
legends in the manuscript. The stylesheets say the image comes first with its caption below it
("Image should be present first and below it the respective caption"), so for each figure legend
("Figure 1: ...") that has no picture next to it, the file that clearly carries that number is
proposed as a new paragraph holding the picture, directly above the legend (`position: below` puts
it under the legend, as the IDOJ editors did). It is a tracked insertion (inserted paragraph mark +
inserted run) the editor accepts or rejects like any other suggestion; nothing reaches the clean
copy until it is accepted. Off by default: most journals' editors leave separate figure files out.

Only clear cases are inserted:
  - the manuscript shows no picture next to any legend of that number, and no picture of its own
    that no legend accounts for (then which figure it is cannot be told, and nothing is inserted);
  - exactly one file carries the number, or several that are parts of it (1a.jpg, 1b.jpg: one
    paragraph holding the parts in order);
  - JPEG / PNG as they are; TIFF, BMP, GIF and WebP converted to PNG (Pillow).
Anything else -- no number in the file name, several candidates, a number with no legend, a
supplementary figure, PowerPoint / PDF / Word / vector files -- becomes an editor warning naming
the file. Pictures already in the manuscript are never touched.

Figures sent in a Word file with their captions ("Revised Figures.docx") are read as such files too (`docx_figures`):
each picture numbered from its caption, the caption standing in for the legend when the manuscript has none.

Settings live under `figures.insert:` in the style files. The picture files go into the output
folder (`source/figures/`) so later renders (after the editor's decisions) find them.

    plan(ms, style, zip_path, outdir)   -> Plan(inserts, warnings)
    render(ms, inserts, base_dir, ...)  place the accepted (or pending, in the tracked copy) inserts
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from itertools import count

from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import provenance, rules_figures
from .changes import Source, Status
from .docmodel import Manuscript, has_picture
from .frontmatter import NewParagraph
from .style import JournalStyle

REASON = "figure_insert"
FIG_DIR = os.path.join("source", "figures")

DIRECT = {".jpg": "jpg", ".jpeg": "jpg", ".png": "png"}
CONVERT = {".tif", ".tiff", ".bmp", ".gif", ".webp"}
# Figure files that cannot be placed as a picture: an editor warning names them.
UNSUPPORTED = {".pptx", ".ppt", ".pdf", ".docx", ".doc", ".eps", ".svg", ".emf", ".wmf", ".ai", ".psd", ".heic"}
_IMAGE_EXT = set(DIRECT) | CONVERT | {".eps", ".svg", ".emf", ".wmf", ".ai", ".psd", ".heic"}

# Files in a submission that are not figures, whatever their format.
_NOT_FIGURE = re.compile(
    r"author|consent|signature|\bsign\b|logo|copyright|cover\s*letter|letter|\bform\b|declaration|"
    r"checklist|reply|response|reviewer|rebuttal|conflict|\bcoi\b|ethic|approval|certificate|"
    r"manuscript|title\s*page|first\s*page|graphical\s*abstract|supplement|\bsuppl?\b|appendix|annex|"
    r"table", re.I)
# Figure number in a file name: "Figure 1", "Fig_2a", "FIG2B", "F3", "Picture 1", "figure-4 edited".
_NAME_NUM = re.compile(
    r"(?<![a-z])(?P<label>fig(?:ure)?s?|pictures?|pic|images?|img|f(?=\s*[-_.]?\s*\d))\s*[-_.]?\s*(?:no\.?\s*)?"
    r"(?P<num>\d{1,2})(?!\d)(?:\s*[-_]?\s*\(?(?P<part>[a-z])\)?(?![a-z]))?", re.I)
# A bare number, with or without a part letter: "1a.jpg", "2.png".
_BARE = re.compile(r"^\s*(?P<num>\d{1,2})\s*[-_]?\s*(?P<part>[a-z])?\s*$", re.I)
_SUPP_NAME = re.compile(r"(?<![a-z])(?:supp\w*|s\d{1,2}\b|fig(?:ure)?\s*[-_]?\s*s\d)", re.I)

_EMU_PER_PX = 9525  # at 96 dpi


@dataclass
class FigureFile:
    name: str  # path inside the zip
    ext: str
    number: int | None
    part: str = ""
    problem: str = ""  # why it cannot be placed by number ("" when it can)
    data: bytes | None = None  # a picture taken out of a Word file of the submission (docx_figures)

    @property
    def base(self) -> str:
        return os.path.basename(self.name)


@dataclass
class FigureInsert(NewParagraph):
    """A new paragraph holding the picture(s) of one figure, placed before its legend."""
    images: list[dict] = field(default_factory=list)  # {"file": rel path, "name": zip name, "width": emu, "height": emu}
    figure: int = 0
    legend: str = ""  # the manuscript has no legend for this figure: "Figure N: ???" goes under the picture

    @property
    def text(self) -> str:  # what the review screen lists for this suggestion
        pics = f"[Figure {self.figure} image: " + ", ".join(os.path.basename(i["name"]) for i in self.images) + "]"
        return f"{pics} {self.legend}" if self.legend else pics

    def to_dict(self) -> dict:
        return {**super().to_dict(), "images": self.images, "figure": self.figure, "legend": self.legend}

    @classmethod
    def from_dict(cls, d: dict) -> "FigureInsert":
        return cls(id=d["id"], anchor=d["anchor"], where=d.get("where", "before"), runs=[], role=d.get("role", "figure"),
                   status=Status(d.get("status", "suggested")), source=Source(d.get("source", "rule")),
                   reason=d.get("reason", REASON), images=list(d.get("images") or []), figure=int(d.get("figure", 0)),
                   legend=d.get("legend", ""),
                   provenance=d.get("provenance") or None)


def is_figure_insert(n) -> bool:
    return isinstance(n, FigureInsert)


@dataclass
class Plan:
    inserts: list[FigureInsert] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    files: list[FigureFile] = field(default_factory=list)
    queries: list[tuple[str, str, dict]] = field(default_factory=list)  # (para_id, queries.TEXT key, format values)
    legends: dict[int, str] = field(default_factory=dict)  # figure -> where its legend came from (no legend in the text)


def settings(style: JournalStyle) -> dict:
    return style.get("figures.insert") or {}


# --- the zip's figure files -------------------------------------------------------------------

def parse_name(name: str) -> tuple[int | None, str, str]:
    """(figure number, part letter, problem) read from a file name."""
    stem = os.path.splitext(os.path.basename(name))[0]
    if _SUPP_NAME.search(stem):
        return None, "", "supplementary figure"
    hits = list(_NAME_NUM.finditer(stem))
    nums = {int(m.group("num")) for m in hits}
    if len(nums) > 1 or any(m.group("label").lower() in ("figures", "pictures", "images") for m in hits):
        return None, "", "several figures in one file name"
    if hits:
        m = hits[0]
        return int(m.group("num")), (m.group("part") or "").lower(), ""
    m = _BARE.match(stem)
    if m:
        return int(m.group("num")), (m.group("part") or "").lower(), ""
    return None, "", "no figure number in the file name"


def supplied_files(zip_path: str) -> list[FigureFile]:
    """The submission's figure files: images, and figure-named PowerPoint / PDF / Word files."""
    z = zipfile.ZipFile(zip_path)
    names = [n for n in z.namelist() if not n.endswith("/") and not n.startswith("__MACOSX")
             and not os.path.basename(n).startswith(".")]
    xml_stems = {os.path.splitext(os.path.basename(n))[0].lower() for n in names if n.lower().endswith(".xml")}
    out, seen = [], set()
    for n in names:
        base = os.path.basename(n)
        stem, ext = os.path.splitext(base)
        ext = ext.lower()
        if ext not in _IMAGE_EXT and ext not in UNSUPPORTED:
            continue
        if stem.lower() in xml_stems or z.getinfo(n).file_size < 2048:
            continue  # the platform's article-id image beside the submission XML, icons
        if _NOT_FIGURE.search(stem) and not re.search(r"^fig", stem, re.I):
            continue
        num, part, problem = parse_name(n)
        if ext not in _IMAGE_EXT and num is None:
            continue  # a PDF / Word / PowerPoint file not named as a figure: some other material
        digest = hashlib.sha1(z.read(n)).hexdigest()
        if digest in seen:
            continue  # the same file twice
        seen.add(digest)
        if ext in UNSUPPORTED:
            problem = problem or f"{ext[1:].upper()} file: cannot be placed as a picture"
        out.append(FigureFile(n, ext, num, part, problem))
    return out


# --- figures in a Word file of the submission ("Revised Figures.docx") ----------------------------

_CONTENT_EXT = {"image/png": ".png", "image/jpeg": ".jpg", "image/jpg": ".jpg", "image/tiff": ".tiff", "image/gif": ".gif",
                "image/bmp": ".bmp", "image/x-emf": ".emf", "image/x-wmf": ".wmf", "image/emf": ".emf", "image/wmf": ".wmf"}


def docx_figures(zip_path: str, style: JournalStyle, manuscript_name: str = "") -> tuple[list[FigureFile],
                                                                                         dict[int, tuple[str, str]], set[str]]:
    """Figures sent in a Word file of the submission with their captions ("Revised Figures.docx": IJO_1347; "Revised
    Figure,Table .docx": jehp_1683): (one FigureFile per picture, numbered from its caption -- several pictures under
    one caption are its parts a, b, ...; figure number -> (legend text, file name); the files read). A caption goes
    with the pictures on its side: the file either puts each picture above its caption or each caption above its
    picture, as its first figure shows. A file whose pictures and captions do not pair up gives nothing."""
    from . import supplied

    files: list[FigureFile] = []
    legends: dict[int, tuple[str, str]] = {}
    used: set[str] = set()
    for name in supplied.docx_members(zip_path, manuscript_name):
        got = supplied.load_member(zip_path, name, None)
        if got is None:
            continue
        sms = got[0]
        base = os.path.basename(name)
        caps = {c.para_id: c for c in rules_figures.captions(sms) if c.kind in (rules_figures.FIG, rules_figures.GRAPH)}
        seq = []  # ("P", paragraph) | ("C", number)
        for p in sms.paragraphs:
            if p.in_table or rules_figures.in_textbox(p):
                continue
            if p.id in caps:
                seq.append(("C", caps[p.id].number))
            elif has_picture(p.element) and not p.text.strip():
                seq.append(("P", p))
        if not any(k == "C" for k, _ in seq) or not any(k == "P" for k, _ in seq):
            continue
        groups: dict[int, list] = {}
        if seq[0][0] == "P":  # pictures, then their caption
            pending = []
            for k, x in seq:
                if k == "P":
                    pending.append(x)
                elif pending and x not in groups:
                    groups[x], pending = pending, []
        else:  # a caption, then its pictures
            cur = None
            for k, x in seq:
                if k == "C":
                    cur = x if x not in groups else None
                    if cur is not None:
                        groups[cur] = []
                elif cur is not None:
                    groups[cur].append(x)
        groups = {n: ps for n, ps in groups.items() if ps}
        if not groups:
            continue
        used.add(name)
        part_rel = sms.document.part.related_parts
        for num, ps in sorted(groups.items()):
            blips = [(p, b) for p in ps for b in p.element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")]
            for k, (p, b) in enumerate(blips):
                try:
                    img = part_rel[b.get(qn("r:embed"))]
                except KeyError:
                    continue
                ext = _CONTENT_EXT.get(getattr(img, "content_type", ""), os.path.splitext(img.partname)[1].lower())
                part = "abcdefghijkl"[k] if len(blips) > 1 and k < 12 else ""
                problem = "" if ext in DIRECT or ext in CONVERT else f"{ext[1:].upper()} picture: cannot be placed"
                files.append(FigureFile(f"{base}#Figure {num}{part}", ext, num, part, problem, data=img.blob))
        for num, text in rules_figures.legend_texts(sms, style).items():
            legends.setdefault(num, (text, base))
    return files, legends, used


# --- legends sent outside the manuscript -------------------------------------------------------

_LEGENDS_FILE = re.compile(r"legend|caption", re.I)


def supplied_legends(zip_path: str, style: JournalStyle,
                     extra: dict[int, tuple[str, str]] | None = None) -> dict[int, tuple[str, str]]:
    """Figure number -> (legend text after "Figure N: ", where it came from), for figures the manuscript has no legend
    for: a separate legends file in the zip ("LEGENDS.docx", "Figures Legends.docx"), else the submission XML's
    <image><legend> for the figure's file (the IJO editor of 2449: "Followed XML file for Figure captions")."""
    import tempfile

    from lxml import etree

    from . import docmodel

    out: dict[int, tuple[str, str]] = {}
    try:
        z = zipfile.ZipFile(zip_path)
    except (OSError, zipfile.BadZipFile):
        return out
    names = [n for n in z.namelist() if not n.startswith("__MACOSX") and not os.path.basename(n).startswith(".")]
    for n in names:
        base = os.path.basename(n)
        if not n.lower().endswith(".docx") or not _LEGENDS_FILE.search(base) or re.search(r"manuscript", base, re.I):
            continue
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "legends.docx")
            with open(path, "wb") as fh:
                fh.write(z.read(n))
            try:
                found = rules_figures.legend_texts(docmodel.load(path), style)
            except Exception:  # an unreadable file: nothing from it
                found = {}
        for num, text in found.items():
            out.setdefault(num, (text, base))
    for num, got in (extra or {}).items():  # the captions of a figures file ("Revised Figures.docx")
        out.setdefault(num, got)
    xmls = [n for n in names if n.lower().endswith(".xml")]
    for n in xmls:
        try:
            root = etree.fromstring(z.read(n))
        except etree.XMLSyntaxError:
            continue
        for img in root.iter("image"):
            fname = (img.findtext("image_file_name") or "").strip()
            legend = (img.findtext("legend") or "").strip()
            if not fname or not legend:
                continue
            num, _, problem = parse_name(fname)
            if num is None or problem or num in out:
                continue
            text = rules_figures.legend_from_text(legend, num, style)
            if text:
                out[num] = (text, os.path.basename(n))
    return out


# --- the manuscript's legends and pictures ----------------------------------------------------

def _legends(ms: Manuscript) -> dict[int, list[int]]:
    """Figure number -> indexes of its legend paragraphs (outside tables)."""
    index = {p.id: i for i, p in enumerate(ms.paragraphs)}
    out: dict[int, list[int]] = {}
    for c in rules_figures.captions(ms):
        i = index[c.para_id]
        if c.kind == rules_figures.FIG and c.number and not ms.paragraphs[i].in_table:
            out.setdefault(c.number, []).append(i)
    return out


def _end_slot(ms: Manuscript) -> tuple[str, str]:
    """Where figures with no legend go: where the back of the article starts after the reference list, before the
    tables (legends.back_slot; G195: tables at the end, after the figure captions), else after the last paragraph with
    text."""
    from . import legends
    return legends.back_slot(ms) or (ms.paragraphs[-1].id, "after")


def _neighbours(ms: Manuscript, i: int, pictures: set[int]) -> list[int]:
    """The nearest paragraphs before and after paragraph i that hold text or a picture."""
    out = []
    for step in (-1, 1):
        j = i + step
        while 0 <= j < len(ms.paragraphs) and not ms.paragraphs[j].text.strip() and j not in pictures:
            j += step
        if 0 <= j < len(ms.paragraphs):
            out.append(j)
    return out


def _covered(ms: Manuscript, legends: dict[int, list[int]], pictures: set[int]) -> tuple[set[int], set[int]]:
    """(figure numbers that have a picture beside a legend, picture paragraphs accounted for)."""
    covered, used = set(), set()
    for num, idxs in legends.items():
        for i in idxs:
            near = [j for j in [i] + _neighbours(ms, i, pictures) if j in pictures]
            if near:
                covered.add(num)
                used.update(near)
    return covered, used


# --- picture sizes ---------------------------------------------------------------------------

def _text_area(ms: Manuscript) -> tuple[int, int]:
    """(width, height) in EMU of the last section's text area (A4 with 1 inch margins if unknown)."""
    try:
        s = ms.document.sections[-1]
        w = int(s.page_width) - int(s.left_margin) - int(s.right_margin)
        h = int(s.page_height) - int(s.top_margin) - int(s.bottom_margin)
        if w > 0 and h > 0:
            return w, h
    except (IndexError, TypeError, ValueError):
        pass
    return 5731510, 8863330


def fit(px_w: int, px_h: int, max_w: int, max_h: int) -> tuple[int, int]:
    """EMU size that fills `max_w` keeping the aspect ratio, shrunk further if taller than `max_h`."""
    w = max_w
    h = round(w * px_h / max(px_w, 1))
    if h > max_h:
        h = max_h
        w = round(h * px_w / max(px_h, 1))
    return int(w), int(h)


def _prepare(data: bytes, ext: str, max_px: int) -> tuple[bytes, str] | None:
    """(image bytes Word can show, extension), converting / downscaling with Pillow when needed."""
    needs_convert = ext in CONVERT
    try:
        from PIL import Image
    except ImportError:
        return None if needs_convert else (data, DIRECT[ext])
    try:
        img = Image.open(io.BytesIO(data))
        too_big = max_px and max(img.size) > max_px
        if not needs_convert and not too_big:
            return data, DIRECT[ext]
        img.seek(0)
        if too_big:
            img.thumbnail((max_px, max_px))
        if not needs_convert and ext in (".jpg", ".jpeg") and img.mode in ("RGB", "L", "CMYK"):
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=90)
            return buf.getvalue(), "jpg"
        if img.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            img = img.convert("RGBA" if "A" in img.mode else "RGB")
        buf = io.BytesIO()
        img.save(buf, "PNG", optimize=True)
        return buf.getvalue(), "png"
    except Exception:  # unreadable or exotic file: leave it to the editor
        return None


# --- planning --------------------------------------------------------------------------------

def _warn(para_id: str, message: str) -> dict:
    return {"para_id": para_id, "rule": "figures.insert", "message": message}


def cited_figures(ms: Manuscript) -> set[int]:
    """The figure numbers the text cites (rules_figures' citations)."""
    from .rules_figures import FIG, citations

    return {n for c in citations(ms) if c.kind == FIG for n in c.numbers}


def plan(ms: Manuscript, style: JournalStyle, zip_path: str | None, outdir: str | None,
         manuscript_name: str = "") -> Plan:
    """Figure inserts for the legends that have no picture, from the zip's figure files.

    The picture files are written to OUTDIR/source/figures (OUTDIR None: not written; the plan is
    for counting only)."""
    cfg = settings(style)
    out = Plan()
    if not cfg.get("enabled") or not zip_path or not zipfile.is_zipfile(zip_path):
        return out
    files = supplied_files(zip_path)
    extra, docx_legends, used_docx = docx_figures(zip_path, style, manuscript_name)
    numbered = {f.number for f in files if f.number is not None and not f.problem}
    files = [f for f in files if f.name not in used_docx] + [f for f in extra if f.number not in numbered]
    out.files = files
    if not files:
        return out
    legends = _legends(ms)
    pictures = {i for i, p in enumerate(ms.paragraphs) if not p.in_table and has_picture(p.element)}
    covered, used = _covered(ms, legends, pictures)
    stray = sorted(pictures - used)
    names = lambda fs: ", ".join(f.base for f in fs)  # noqa: E731

    def block(num) -> tuple[int, int]:
        """First and last paragraph of the figure's legend ("Figure 1A: ..." + "Figure 1B: ..."): its
        last occurrence, with the legends of the same number right before it."""
        idxs = legends[num]
        first = last = idxs[-1]
        for i in reversed(idxs[:-1]):
            if all(not ms.paragraphs[j].text.strip() for j in range(i + 1, first)):
                first = i
            else:
                break
        return first, last

    def legend_id(num):
        return ms.paragraphs[block(num)[0]].id if num in legends else ""

    placeable = [f for f in files if not f.problem]
    by_num: dict[int, list[FigureFile]] = {}
    for f in placeable:
        by_num.setdefault(f.number, []).append(f)
    # A legend with no picture and no numbered file: an unnumbered file may be it (the editor decides).
    orphans = sorted(n for n in legends if n not in covered and n not in by_num)
    unnumbered = [f for f in files if f.number is None and f.problem == "no figure number in the file name"]
    for f in files:
        if f.problem:
            hint = f" (Figure {orphans[0]} has no picture)" if len(orphans) == 1 and len(unnumbered) == 1 \
                and f is unnumbered[0] else ""
            out.warnings.append(_warn(legend_id(f.number) if f.number in legends else
                                      (legend_id(orphans[0]) if hint else ""),
                                      f"Figure file {f.base} not inserted: {f.problem}{hint}."))
    if stray and any(n not in covered for n in by_num):
        pids = ", ".join(ms.paragraphs[i].id for i in stray[:5])
        out.warnings.append(_warn(ms.paragraphs[stray[0]].id,
                                  f"Figure files not inserted ({names(placeable)}): the manuscript has pictures "
                                  f"no figure legend accounts for ({pids}), so which figures are missing is unclear."))
        return out

    supplied = supplied_legends(zip_path, style, docx_legends) if any(n not in legends for n in by_num) else {}
    max_w, max_h = _text_area(ms)
    max_h = int(max_h * float(cfg.get("max_height", 0.85)))
    per_row = max(1, int(cfg.get("parts_per_row", 2)))
    max_px = int(cfg.get("max_pixels", 2400) or 0)
    z = zipfile.ZipFile(zip_path)
    n_ins = count(1)
    for num, fs in sorted(by_num.items()):
        if num in covered:
            continue  # the manuscript already shows it: never touched
        no_legend = num not in legends
        if no_legend and not cfg.get("without_legend", True):
            out.warnings.append(_warn("", f"Figure file(s) {names(fs)} not inserted: the manuscript has no "
                                          f"legend for Figure {num}."))
            continue
        parts = sorted(fs, key=lambda f: f.part)
        if len(fs) > 1 and (any(not f.part for f in fs) or len({f.part for f in fs}) != len(fs)):
            out.warnings.append(_warn(legend_id(num), f"Figure {num} not inserted: several files match "
                                                      f"({names(fs)}); place the right one by hand."))
            continue
        prepared = []
        for f in parts:
            got = _prepare(f.data if f.data is not None else z.read(f.name), f.ext, max_px)
            if got is None:
                break
            prepared.append((f, *got))
        if len(prepared) != len(parts):
            bad = parts[len(prepared)]
            out.warnings.append(_warn(legend_id(num), f"Figure {num} not inserted: {bad.base} could not be "
                                                      f"read or converted to PNG."))
            continue
        cols = min(len(prepared), per_row)
        slot_w = max_w if cols == 1 else int(max_w / cols * 0.97)
        images = []
        for f, data, ext in prepared:
            try:
                img = DocxImage.from_blob(data)
                pw, ph = img.px_width, img.px_height
            except Exception:
                pw = ph = 0
            if not pw or not ph:
                images = []
                break
            w, h = fit(pw, ph, slot_w, max_h)
            stem = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.splitext(f.base)[0])[:60] or "figure"
            rel = os.path.join(FIG_DIR, f"fig{num}{f.part}_{stem}.{ext}")
            if outdir:
                os.makedirs(os.path.join(outdir, FIG_DIR), exist_ok=True)
                with open(os.path.join(outdir, rel), "wb") as fh:
                    fh.write(data)
            images.append({"file": rel, "name": f.name, "width": w, "height": h})
        if not images:
            out.warnings.append(_warn(legend_id(num), f"Figure {num} not inserted: {names(parts)} has no "
                                                      f"readable picture size."))
            continue
        below = cfg.get("position", "above") == "below"
        legend = ""
        if no_legend:  # picture, then its legend (S/G266: image first, its caption below)
            anchor, where = _end_slot(ms)
            text, origin = supplied.get(num, ("", ""))
            if text:  # the legend from the submission's other material (a legends file, the XML)
                legend = f"Figure {num}: {text}"
                out.legends[num] = origin
            else:  # "Figure N: ???" and a query to the author
                legend = f"Figure {num}: ???"
                out.queries.append((anchor, "figures.legend_missing", {"n": num}))
            if num not in cited_figures(ms):  # a picture sent as a file that the text never cites
                out.queries.append((anchor, "figures.uncited", {"kind": "Figure", "n": num}))
        else:
            first, last = block(num)
            anchor, where = ms.paragraphs[last if below else first].id, "after" if below else "before"
        files_ = [os.path.basename(i["name"]) for i in images]
        detail = ", ".join(files_) + (f"; legend from {out.legends[num]}" if num in out.legends else "")
        out.inserts.append(FigureInsert(
            id=f"fig_ins{next(n_ins)}", anchor=anchor, legend=legend,
            where=where, runs=[], role="figure", reason=REASON, images=images, figure=num,
            provenance=provenance.make("figure_file", "Figure file", detail, files=files_)))
    return out


# --- rendering -------------------------------------------------------------------------------

def _mark(el, ids, author: str, date: str):
    el.set(qn("w:id"), str(next(ids)))
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), date)
    return el


def render(ms: Manuscript, inserts, base_dir: str, tracked: bool = True, statuses=None,
           author: str = "Edit OS", align: str = "center", run_format: dict | None = None,
           legend_align: str | None = None) -> int:
    """Place figure inserts whose status is in `statuses` (default accepted) next to their anchors,
    as tracked insertions when `tracked`. Call after `tracked.render`. Returns how many were placed;
    an insert whose anchor or picture file is gone is skipped."""
    statuses = statuses or (Status.ACCEPTED,)
    date = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    taken = [int(v) for el in ms.document.element.body.iter() if (v := el.get(qn("w:id"))) and v.isdigit()]
    ids = count(max([7000] + taken) + 1)
    placed = 0
    tail: dict[str, object] = {}
    for n in inserts:
        if not is_figure_insert(n) or n.status not in statuses:
            continue
        try:
            anchor = ms.by_id(n.anchor).element
        except KeyError:
            continue
        paths = [os.path.join(base_dir, i["file"]) for i in n.images]
        if not paths or not all(os.path.exists(p) for p in paths):
            continue
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        if n.where != "after":
            ppr.append(OxmlElement("w:keepNext"))  # the picture stays on the page of the legend below it
        if align:
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), align)
            ppr.append(jc)
        if tracked:
            prpr = OxmlElement("w:rPr")
            prpr.append(_mark(OxmlElement("w:ins"), ids, author, date))
            ppr.append(prpr)
        p.append(ppr)
        for k, (img, path) in enumerate(zip(n.images, paths)):
            r = OxmlElement("w:r")
            if k:  # a space between the parts of one figure
                t = OxmlElement("w:t")
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                t.text = " "
                sep = OxmlElement("w:r")
                sep.append(t)
                p.append(_wrap(sep, tracked, ids, author, date))
            with open(path, "rb") as fh:
                inline = ms.document.part.new_pic_inline(fh, img["width"], img["height"])
            drawing = OxmlElement("w:drawing")
            drawing.append(inline)
            r.append(drawing)
            p.append(_wrap(r, tracked, ids, author, date))
        block = [p]
        if n.legend:
            block.append(_legend_paragraph(n.legend, tracked, ids, author, date, run_format, legend_align))
        if n.legend:  # a figure with no legend in the text: legends.render puts it among the legends, in number order
            ms.__dict__.setdefault("_back_figures", []).append((n.figure, block))
        if n.where == "after" and n.legend:  # a figure with no legend in the text, at the end of the document:
            body = ms.document.element.body     # after the First author block and the legends heading
            last = body[-1] if len(body) and body[-1].tag == qn("w:sectPr") else None
            for el in block:
                (last.addprevious(el) if last is not None else body.append(el))
        elif n.where == "after":
            prev = tail.get(n.anchor, anchor)  # several figures after one paragraph keep their order
            for el in block:
                prev.addnext(el)
                prev = el
            tail[n.anchor] = prev
        else:
            for el in block:
                anchor.addprevious(el)
        placed += 1
    return placed


def _legend_paragraph(text: str, tracked: bool, ids, author: str, date: str, run_format: dict | None = None,
                      align: str | None = None):
    """"Figure N: ???" under a picture whose legend the manuscript lacks; the label bold (G260). `run_format`,
    `align`: the journal's font, size and alignment."""
    p = OxmlElement("w:p")
    if tracked or align:
        ppr = OxmlElement("w:pPr")
        if align:
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), align)
            ppr.append(jc)
        if tracked:
            prpr = OxmlElement("w:rPr")
            prpr.append(_mark(OxmlElement("w:ins"), ids, author, date))
            ppr.append(prpr)
        p.append(ppr)
    label, _, rest = text.partition(":")
    for part, bold in ((label + ":", True), (rest, False)):
        r = OxmlElement("w:r")
        fmt = {**(run_format or {}), **({"bold": True} if bold else {})}
        if fmt:
            from .tracked import apply_format
            rpr = OxmlElement("w:rPr")
            apply_format(rpr, fmt)
            r.append(rpr)
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = part
        r.append(t)
        p.append(_wrap(r, tracked, ids, author, date))
    return p


def _wrap(run, tracked: bool, ids, author: str, date: str):
    if not tracked:
        return run
    w = _mark(OxmlElement("w:ins"), ids, author, date)
    w.append(run)
    return w
