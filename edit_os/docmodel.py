"""Read a .docx into paragraphs with stable IDs, without losing formatting.

The original XML stays the source of truth. Each paragraph gets an ID (p0001, ...)
and a plain-text view; edits are expressed as character offsets into that text
and written back by `tracked.write`.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn

# Run content elements and the text each contributes to the paragraph's text view.
_TEXT = qn("w:t")
_TAB = qn("w:tab")
_BREAKS = (qn("w:br"), qn("w:cr"))
_RPR = qn("w:rPr")
_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_NOT_OWN = {qn("w:del"), qn("w:p"), _FALLBACK}
_SYM = qn("w:sym")
# Word's non-breaking hyphen (Ctrl+Shift+-): a character of its own, read as "-" so "GBL‑WSQ" stays one word in the
# text view (1039). One element, one character: offsets and run splitting treat it as they treat a <w:sym>.
_NB_HYPHEN = qn("w:noBreakHyphen")

# Characters Word inserts from the Symbol font (Insert > Symbol: <w:sym w:font="Symbol" w:char="F0B1"/>), by their
# Symbol code: the Greek letters, ±, ≤, ≥ and the minus sign (owner, 2026-10-26: these, never ×, ®, ©, µ). The text
# view shows them as the Unicode character, so rules and QC see "±" whichever way it was typed.
_GREEK_UPPER = "ΑΒΧΔΕΦΓΗΙϑΚΛΜΝΟΠΘΡΣΤΥςΩΞΨΖ"  # Symbol 0x41-0x5A
_GREEK_LOWER = "αβχδεφγηιϕκλμνοπθρστυϖωξψζ"  # Symbol 0x61-0x7A
SYMBOL_CODES: dict[str, int] = {ch: 0x41 + k for k, ch in enumerate(_GREEK_UPPER)}
SYMBOL_CODES.update({ch: 0x61 + k for k, ch in enumerate(_GREEK_LOWER)})
SYMBOL_CODES.update({"±": 0xB1, "≤": 0xA3, "≥": 0xB3, "−": 0x2D})
_FROM_SYMBOL = {code: ch for ch, code in SYMBOL_CODES.items()}


def sym_char(el) -> str:
    """The Unicode character of a <w:sym> in the Symbol font (one of SYMBOL_CODES), else ""."""
    if (el.get(qn("w:font")) or "").strip().lower() != "symbol":
        return ""
    try:
        code = int(el.get(qn("w:char")) or "", 16)
    except ValueError:
        return ""
    return _FROM_SYMBOL.get(code - 0xF000 if code >= 0xF000 else code, "")


@dataclass
class Paragraph:
    id: str
    element: object  # lxml <w:p>
    style: str
    text: str = ""
    in_table: bool = False
    numbered: bool = False  # auto-numbered list item (w:numPr)
    superscript: list[tuple[int, int]] = field(default_factory=list)  # spans formatted as superscript


@dataclass
class Manuscript:
    path: str
    document: object  # docx.Document
    paragraphs: list[Paragraph] = field(default_factory=list)
    # Elements a clean render removed (whole paragraphs and tables deleted by accepted changes); they leave
    # the XML only on save, so paragraph ids stay valid anchors until then (frontmatter.render_inserts).
    removed: list = field(default_factory=list)

    def by_id(self, para_id: str) -> Paragraph:
        for p in self.paragraphs:
            if p.id == para_id:
                return p
        raise KeyError(para_id)

    def save(self, path: str) -> None:
        for el in self.removed:
            if el.getparent() is not None:
                el.getparent().remove(el)
        self.removed = []
        self.document.save(path)


def run_text(run) -> str:
    """Text a single <w:r> contributes; non-text content (images, fields) counts as ''."""
    out = []
    for child in run:
        if child.tag == _TEXT:
            out.append(child.text or "")
        elif child.tag == _TAB:
            out.append("\t")
        elif child.tag in _BREAKS:
            out.append("\n")
        elif child.tag == _SYM:
            out.append(sym_char(child))
        elif child.tag == _NB_HYPHEN:
            out.append("-")
    return "".join(out)


def text_runs(p_el) -> list:
    """Runs that make up the visible text, in order, skipping existing tracked deletions."""
    runs = []
    for r in p_el.iter(qn("w:r")):
        # Only this paragraph's own runs: skip tracked deletions, runs of paragraphs nested in a
        # text box (they are paragraphs of their own), and Word's duplicate fallback copy.
        for a in r.iterancestors():
            if a is p_el:
                runs.append(r)
                break
            if a.tag in _NOT_OWN:
                break
    return runs


def superscript_spans(p_el) -> list[tuple[int, int]]:
    """Text spans whose runs are formatted superscript (direct formatting), merged when adjacent."""
    spans, pos = [], 0
    for r in text_runs(p_el):
        n = len(run_text(r))
        va = r.find(f"{_RPR}/{qn('w:vertAlign')}")
        if n and va is not None and va.get(qn("w:val")) == "superscript":
            if spans and spans[-1][1] == pos:
                spans[-1] = (spans[-1][0], pos + n)
            else:
                spans.append((pos, pos + n))
        pos += n
    return spans


def paragraph_text(p_el) -> str:
    return "".join(run_text(r) for r in text_runs(p_el))


_PICTURE_TAGS = (
    "{http://schemas.openxmlformats.org/drawingml/2006/main}blip",  # picture in a w:drawing
    "{urn:schemas-microsoft-com:vml}imagedata",  # older w:pict picture
    "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart",  # embedded chart
    qn("w:object"),  # OLE object shown as a picture (MathType equations included)
)


def has_picture(p_el) -> bool:
    """True when the paragraph shows a picture, chart or embedded object of its own (not one in a
    tracked deletion, Word's duplicate fallback copy or a paragraph nested in a text box)."""
    for el in p_el.iter(*_PICTURE_TAGS):
        for a in el.iterancestors():
            if a is p_el:
                return True
            if a.tag in _NOT_OWN:
                break
    return False


def explode_runs(p_el) -> None:
    """Split every run so it holds exactly one content element (keeping its rPr).

    After this, a w:t run can be split by character and tab/break runs are one character.
    """
    for r in text_runs(p_el):
        content = [c for c in r if c.tag != _RPR]
        if len(content) <= 1:
            continue
        parent = r.getparent()
        idx = parent.index(r)
        for i, c in enumerate(content):
            new = copy.deepcopy(r)
            for child in list(new):
                if child.tag != _RPR:
                    new.remove(child)
            new.append(copy.deepcopy(c))
            parent.insert(idx + i, new)
        parent.remove(r)


def load(path: str) -> Manuscript:
    doc = Document(path)
    ms = Manuscript(path=path, document=doc)
    body = doc.element.body
    n = 0
    for p_el in body.iter(qn("w:p")):
        if any(a.tag == _FALLBACK for a in p_el.iterancestors()):
            continue  # duplicate copy of a text box for old Word versions
        n += 1
        style_el = p_el.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        style = style_el.get(qn("w:val")) if style_el is not None else ""
        in_table = any(a.tag == qn("w:tc") for a in p_el.iterancestors())
        numbered = p_el.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
        ms.paragraphs.append(Paragraph(
            id=f"p{n:04d}", element=p_el, style=style,
            text=paragraph_text(p_el), in_table=in_table, numbered=numbered,
            superscript=superscript_spans(p_el),
        ))
    return ms
