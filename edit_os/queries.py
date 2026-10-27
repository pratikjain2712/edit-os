"""Author queries: what cannot be fixed without the author, in the pre-editing guide's wording.

Queries are records like changes: anchored to exact text, numbered AQ1, AQ2, ... in reading
order, written into the file as Word comments, and listed on a query sheet. Add comments before
rendering tracked changes, so anchors are taken from the original text.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run

from . import refengine, references, refs_other
from .docmodel import Manuscript
from .style import JournalStyle

AUTHOR = "Edit OS"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Standard wording, from the pre-editing style guide where it gives one.
TEXT = {
    "references.uncited": "Reference no. {n} has not been cited. Kindly cite.",
    "references.missing_entry": "Citation [{n}] has no entry in the reference list. Kindly provide the reference.",
    "references.not_found_et_al": "Reference not found in PubMed. Hence kindly provide minimum 6 author names before et al.",
    "references.not_found": "Reference not found in PubMed or CrossRef. Kindly check and provide complete details.",
    "references.wrong_id": "The {what} given points to a different article. Kindly check the reference.",
    "references.no_volume": "Kindly provide volume and page numbers.",
    "references.no_pages": "Kindly provide page numbers.",
    "references.record_differs": "The {what} given here {differ} from the published article. Kindly check.",
    "references.author_style": "The author names could not be set in the journal's style (surname followed by initials). Kindly check and provide them.",
    "references.author_names": "The author names differ from the published article ({pairs}). Kindly check.",
    "references.author_order": "The order of the authors differs from the published article ({names}). Kindly "
                               "check.",
    "references.author_count": "The published article lists {n} authors ({names}). Kindly check.",
    "references.author_names_incomplete": "The published record gives incomplete author names ({names}), so the "
                                          "author names could not be verified. Kindly check.",
    "references.author_names_unlisted": "The published record lists no authors, so the author names could not be "
                                        "verified. Kindly check.",
    "references.full_names": "The author names were given in full and have been changed to surname and initials. "
                             "Kindly check.",
    "references.book_et_al": "Kindly provide minimum 6 author names before et al.",
    "references.link_dead": "Link is not working. Kindly check.",
    "references.link_redirect": "The link now opens a different page ({url}). Kindly check and provide the correct URL.",
    "references.access_date": "Kindly provide last accessed date.",
    "references.link_broken": "The URL is broken by spaces. Kindly check and provide the complete URL.",
    "references.publisher_name": "Kindly provide publisher name and location.",
    "references.publisher_location": "Kindly provide publisher location.",
    "references.book_year": "Kindly check the year of publication; the book is listed with {years}.",
    "references.book_unverified": "This book could not be verified. Kindly check the authors, title, edition, publisher, city and year.",
    "references.unverified": "This reference could not be verified. Kindly check and confirm the details.",
    "keywords.missing": "Kindly provide key words.",
    "keywords.too_few": "Kindly provide minimum three key words.",
    "sections.conclusion": "Kindly provide the conclusion.",
    "figures.legend_missing": "Kindly provide the legend for Figure {n}.",
    "text.brackets": "The brackets here do not match. Kindly check.",
    "text.instruction": "This looks like a drafting instruction left in the text (“{quote}”). Kindly check and "
                        "delete it, or provide the final text.",
    "references.incomplete": "This reference is incomplete (no {missing}). Kindly provide the complete details.",
}


_LINK_TAIL = re.compile(r"\b(?:doi\b|https?://|www\.|available\s+(?:from|at)|retrieved\s+from|\[?cited\b|10\.\d{4,9}/)", re.I)
# volume, an optional issue, then the pages: "2025;12:991-3", "2020;7(2):87-96", "2023. 13(1): p. 1-10" (EndNote),
# "Journal, 12(3), 991–993" (APA), "2024;15:e1234"
_VOL_PAGES = re.compile(r"(?<![\d.\-/])(\d{1,4}[A-Za-z]?)\s*(?:\(\s*[^()]{1,20}\))?\s*[:,]\s*(?:pp?\.\s*)?"
                        r"([A-Za-z]{0,2}\d+(?:\s*[-–—]\s*[A-Za-z]{0,2}\d+)?)(?![\d/])")
_VOL_ONLY = re.compile(r"(?:19|20)\d{2}[a-z]?(?:\s+[A-Z][a-z]{2,8}(?:\s+\d{1,2})?)?\s*[;.]\s*(\d{1,4})(?:\s*\([^()]{1,20}\))?"
                       r"\s*(?:[.:]\s*$|\.|$)")


def volume_pages(entry: str) -> tuple[bool, bool]:
    """(has a volume, has page numbers) in a reference entry as written: "Int Surg J.  2025;12:991-3." -> both;
    "J Educ. 2025;12." -> volume only. The DOI / URL tail is not read (its digits are no volume)."""
    m = _LINK_TAIL.search(entry)
    body = entry[:m.start()] if m else entry
    year = re.search(r"\b(?:19|20)\d{2}\b", body)
    tail = body[year.start():] if year else body
    tail = tail.rstrip(" .;")
    for v in _VOL_PAGES.finditer(tail):
        # near the end of the citation, where volume and pages stand (not "2 groups: 3, 4" in an APA title)
        if not re.fullmatch(r"(?:19|20)\d{2}", v.group(1)) and len(tail) - v.end() <= 25:
            return True, True
    return bool(_VOL_ONLY.search(tail)), False


def original_article(ms: Manuscript) -> bool:
    """Is this an Original Article? The type the front matter sets (the XML's, the manuscript's "Article type:"
    line, else "Original Article" for an IMRaD paper), else the XML's. "Kindly provide the conclusion" is asked of
    Original Articles only (owner, 2026-10-27: the editors' practice)."""
    from .rules_sections import article_type

    plan = getattr(ms, "_front_plan", None)
    kind = next((b.text for b in getattr(plan, "blocks", []) or [] if b.role == "article_type"), "") or article_type(ms)
    if kind.strip():
        return bool(re.search(r"\boriginal\b", kind, re.I))
    heads = {re.sub(r"^<H\d>|[:.]$|^\d+\.?\s*", "", p.text.strip()).strip().lower() for p in ms.paragraphs
             if len(p.text) < 60}  # no type given: an IMRaD paper (its Conclusion is what may be missing)
    need = [{"introduction", "background"}, {"materials and methods", "methods", "patients and methods",
                                             "subjects and methods", "methodology"}, {"results"}, {"discussion"}]
    return all(heads & n for n in need)


def unmatched_brackets(text: str) -> tuple[int, int] | None:
    """The span of a stray bracket: "(−1.20,3.40))", "((" / "))" doubled or a closing one with nothing open, in a
    paragraph whose round or square brackets do not pair up. List markers ("a) ...", "1) ...", "i)") are not
    strays."""
    for o, c in ("()", "[]"):
        if text.count(o) == text.count(c):
            continue
        depth = 0
        for i, ch in enumerate(text):
            if ch == o:
                depth += 1
            elif ch == c:
                if depth == 0:
                    if re.match(r"(?:^|.*[\s;,:])(?:[a-zA-Z]|\d{1,2}|[ivxIVX]{1,4})$", text[:i]) and not \
                            text[i + 1:i + 2] == c:
                        continue  # "a) first item", "2) ..."
                    return i, i + 1
                depth -= 1
        m = re.search(re.escape(o) * 2 + "|" + re.escape(c) * 2, text)
        if m:
            return m.start(), m.end()
    return None


# A drafting note left in the text: an instruction sentence ("Replace “<.001” with exact software p-values if the
# original output provides them." in a table note, jehp_1402) or a bracketed placeholder ("[insert reference]").
_INSTRUCTION = re.compile(
    r"(?:^|(?<=[.!?]\s)|(?<=[.!?][”\"]\s))(?P<s>(?:Replace|Insert|Fill\s+in|Update)\b(?:[^.!?]|\.(?=\d)){3,160}?"
    r"\b(?:with\s+(?:the\s+)?(?:exact|actual|final|correct|real)\b|if\s+(?:the\s+)?original\b|once\s+(?:available|known)"
    r"|before\s+submission)(?:[^.!?]|\.(?=\d))*[.!?]?)")
_PLACEHOLDER = re.compile(r"\[(?:insert|add|citation\s+needed|ref(?:erence)?s?\s+needed|TBD|TODO|to\s+be\s+added)[^\]]{0,60}\]",
                          re.I)


def drafting_instruction(text: str) -> tuple[int, int] | None:
    """The span of a drafting instruction or placeholder left in a paragraph, or None."""
    m = _INSTRUCTION.search(text)
    if m:
        return m.start("s"), m.end("s")
    m = _PLACEHOLDER.search(text)
    return (m.start(), m.end()) if m else None


@dataclass
class Query:
    para_id: str
    start: int  # anchor in Paragraph.text; start == end anchors the whole paragraph
    end: int
    text: str
    check: str
    id: str = ""  # AQ1, AQ2, ... assigned in reading order
    # A query on a paragraph the front matter inserts (frontmatter.NewParagraph id): start / end are offsets in that
    # paragraph's text and `quote` is the text anchored; para_id is the manuscript paragraph it stands next to (for
    # the reading order). Written after the inserts are rendered (write_insert_comments).
    insert_id: str = ""
    quote: str = ""


@dataclass
class QuerySet:
    queries: list[Query] = field(default_factory=list)

    def add(self, para_id: str, check: str, start: int = 0, end: int = 0, **fmt) -> Query:
        from .rules_figures import QUERY_TEXT as FIGURE_TEXT
        from .rules_sections import TEXT as SECTION_TEXT

        if check == "references.record_differs" and "differ" not in fmt:  # "the journal and volume ... differ"
            what = fmt.get("what", "")
            fmt["differ"] = "differ" if " and " in what or what.endswith("s") else "differs"
        q = Query(para_id, start, end, {**FIGURE_TEXT, **SECTION_TEXT, **TEXT}[check].format(**fmt), check)
        self.queries.append(q)
        return q

    def number(self, ms: Manuscript) -> "QuerySet":
        order = {p.id: i for i, p in enumerate(ms.paragraphs)}
        seen, unique = set(), []
        for q in self.queries:  # the same query twice on one anchor is asked once
            key = (q.para_id, q.insert_id, q.start, q.end, q.text)
            if key not in seen:
                seen.add(key)
                unique.append(q)
        self.queries[:] = unique
        # a paragraph inserted before its anchor (the front matter) reads first, in the order the inserts were made
        self.queries.sort(key=lambda q: (order.get(q.para_id, 10 ** 9), 0 if q.insert_id else 1,
                                         int(re.sub(r"\D", "", q.insert_id) or 0), q.start))
        for i, q in enumerate(self.queries, 1):
            q.id = f"AQ{i}"
        return self


def keep_anchors_live(ms: Manuscript, qs: QuerySet, changes, inserts=()) -> int:
    """Move each query off a paragraph that is empty or that a change deletes whole (a running-title line, a "Title:"
    line the front matter rewrites as a new paragraph): onto the paragraph inserted in its place when there is one,
    else onto the next paragraph that stays. Returns how many moved."""
    gone = set()
    if changes:
        writes = {c.para_id for c in changes if c.new.strip()}
        gone = {c.para_id for c in changes if c.start == 0 and c.old.strip() and not c.new.strip()
                and c.para_id not in writes and c.end == len(ms.by_id(c.para_id).text)}
    # paragraphs a block move takes elsewhere (legends, a caption into its table): their comments go with them
    gone -= {pid for n in inserts for b in getattr(n, "blocks", None) or [] for pid in b}
    paras = ms.paragraphs
    index = {p.id: i for i, p in enumerate(paras)}
    moved = 0
    # A query on material the tool takes out as not the article's (a reply to the reviewers, a template leftover)
    # asks nothing the author has to answer: it goes with the material.
    discarded = {c.para_id for c in changes or () if c.para_id in gone
                 and (c.reason.startswith("submission_reviewer_reply") or c.reason == "template_leftovers")}
    drop = {k for k, q in enumerate(qs.queries) if not q.insert_id and q.para_id in discarded}
    moved += _off_replaced_spans(ms, qs, changes or (), skip=drop)
    if drop:
        qs.queries[:] = [q for k, q in enumerate(qs.queries) if k not in drop]
    for k, q in enumerate(qs.queries):
        if q.insert_id or q.para_id not in index:
            continue
        p = paras[index[q.para_id]]
        empty = not p.text.strip()
        if empty and _picture_only(p) or not empty and p.id not in gone:
            continue  # a picture (figure or equation image) carries the comment on its run
        repl = next((n for n in inserts if getattr(n, "anchor", None) == p.id and hasattr(n, "runs")
                     and getattr(n, "reason", "") == "front_matter"), None)
        if repl is not None and repl.text.strip():
            qs.queries[k] = Query(p.id, 0, 0, q.text, q.check, q.id, insert_id=repl.id, quote=repl.text)
            moved += 1
            continue
        nxt = next((x for x in paras[index[p.id] + 1:] if x.text.strip() and x.id not in gone), None)
        if nxt is not None:
            qs.queries[k] = Query(nxt.id, 0, 0, q.text, q.check, q.id)
            moved += 1
    return moved


def _off_replaced_spans(ms: Manuscript, qs: QuerySet, changes, skip=frozenset()) -> int:
    """A query on text a change deletes or rewrites (the citation "(Saini et al., 2024)" renumbered, "Nur5" -> "Nur", a
    caption part run on to the legend above) would sit on struck-through text only: it moves to where the text goes
    (the paragraph a change writes it into), else it takes in the word next to the change that stays."""
    edits = [c for c in changes if not c.format_only and c.old and c.reason != "renumber"]
    moved = 0
    for k, q in enumerate(qs.queries):
        if k in skip or q.insert_id or q.end <= q.start:
            continue
        c = next((c for c in edits if c.para_id == q.para_id and c.start <= q.start and q.end <= c.end), None)
        if c is None:
            continue
        try:
            p = ms.by_id(q.para_id)
        except KeyError:
            continue
        quote = p.text[q.start:q.end]
        if not c.new.strip() and quote.strip():
            # the text written into another paragraph (a run-on caption part): the query goes with it
            to = next((x for x in changes if x is not c and x.para_id != c.para_id and not x.format_only
                       and quote in x.new), None)
            if to is not None:
                qs.queries[k] = Query(to.para_id, 0, 0, q.text, q.check, q.id)
                moved += 1
                continue
        before = re.search(r"(\S+\s*)$", p.text[:c.start])
        after = re.match(r"\s*\S+", p.text[c.end:])
        if before:
            start, end = before.start(1), c.end
        elif after:
            start, end = c.start, c.end + after.end()
        else:
            start = end = 0  # the whole paragraph rewritten: on the paragraph
        # (the neighbour word must not be another change's deleted text either)
        if any(x is not c and x.para_id == c.para_id and x.old and x.start <= start < x.end for x in edits) \
                and after:
            start, end = c.start, c.end + after.end()
        qs.queries[k] = Query(q.para_id, start, end, q.text, q.check, q.id)
        moved += 1
    return moved


def _picture_only(p) -> bool:
    return any(r.find(_W + t) is not None for r in p.element.iter(_W + "r") for t in ("drawing", "pict", "object"))


def rules_sections_xml(ms: Manuscript) -> list[str]:
    from .rules_sections import xml_keywords  # the submission's own key words fill an empty list

    return xml_keywords(ms)


def collect(ms: Manuscript, style: JournalStyle) -> QuerySet:
    from .rules import lookup_references  # rules imports this module's neighbours; avoid a cycle

    qs = QuerySet()
    refs = references.extract(ms)
    numbered = {r.number: r for r in refs if r.number is not None}

    # Citations against the reference list: the citations renumbering reads (renumber.occurrences: the
    # citation rule's own decisions on round brackets and superscripts), so the queries, renumbering and QC
    # agree. Counts, list markers, statistics and the reference list's own "(4)" issues are not citations.
    from .renumber import occurrences

    cited: set[int] = set()
    for o in occurrences(ms, style, 0):
        cited.update(o.numbers)
        missing = [n for n in o.numbers if numbered and n not in numbered]
        if missing:
            # anchored on the brackets as the author wrote them (a converted citation's own text)
            text = ms.by_id(o.para_id).text
            lo = o.start if o.end > o.start else max(0, o.start - 40)
            spans = list(re.finditer(r"[\[(][^\[\]()]*[\])]", text[lo:max(o.end, o.start)]))
            s, e = (lo + spans[-1].start(), lo + spans[-1].end()) if spans else (o.start, o.end)
            for n in missing:
                qs.add(o.para_id, "references.missing_entry", s, e, n=n)
    if cited:
        for n, r in numbered.items():
            if n not in cited:
                qs.add(r.para_id, "references.uncited", r.offset, r.offset + len(r.text), n=n)

    # Reference lookups that need the author.
    for ref, res in lookup_references(ms, style):
        span = dict(start=ref.offset, end=ref.offset + len(ref.text))
        has_et_al = re.search(r"\bet al\b", ref.text) is not None
        few_authors = len(res.parsed.authors) < 6
        if res.status == refengine.QUERY:
            # One reference can need several things from the author (S: "If et al. is present and less than 6
            # author names are present ... (Query: Please provide minimum 6 author names before et al.)" holds
            # whatever else is asked about it): the record's problem first, then the author names.
            found = True  # a record was found (another article, other details, no volume yet)
            if "points to a different article" in res.message:
                what = "PMID" if res.message.startswith("PMID") else "DOI"
                qs.add(ref.para_id, "references.wrong_id", what=what, **span)
            elif refengine.DIFFERS in res.message:
                what = res.message.split(": the ", 1)[1].split(" " + refengine.DIFFERS)[0]
                qs.add(ref.para_id, "references.record_differs", what=what, **span)
            elif "no volume" in res.message:
                # the record has no volume yet: ask only for what the entry as delivered still lacks (the author
                # often gives them: "Int Surg J.  2025;12:991-3")
                vol, pages = volume_pages(ref.text)
                if res.styled and res.formatted:
                    v2, p2 = volume_pages(res.formatted)
                    vol, pages = vol or v2, pages or p2
                if not vol:
                    qs.add(ref.para_id, "references.no_volume", **span)
                elif not pages:
                    qs.add(ref.para_id, "references.no_pages", **span)
            else:
                found = False
            if has_et_al and few_authors:
                qs.add(ref.para_id, "references.book_et_al" if found else "references.not_found_et_al", **span)
            elif not found:
                qs.add(ref.para_id, "references.not_found", **span)
        elif res.status in refs_other.STATUSES and refs_other.classify(ref.text, res.parsed.kind) in ("book", "book_chapter") \
                and has_et_al and few_authors:
            qs.add(ref.para_id, "references.book_et_al", **span)
        added = {q.check for q in qs.queries if q.para_id == ref.para_id}
        if added & {"references.no_volume", "references.not_found", "references.not_found_et_al"}:
            added.add("references.no_pages")  # "volume and page numbers" / "complete details" ask for the pages too
        for check, fmt in res.queries:  # books, links, reports (refs_other); fields missing from a styled entry
            if check == "references.full_names" and res.status == refengine.QUERY:
                continue  # already queried as not found: "provide complete details" covers the names
            if check == "references.no_volume" and "references.no_volume" not in added:
                vol, pages = volume_pages(ref.text)  # the styled entry lost what the author gave: not missing
                if vol:
                    check = "references.no_pages" if not pages else ""
            if check == "references.no_pages" and check not in added and all(volume_pages(ref.text)):
                continue
            if check and check not in added:
                added.add(check)
                qs.add(ref.para_id, check, **fmt, **span)
        if not (added - {"references.uncited"}) and not re.search(r"\b(?:1[6-9]|20)\d{2}[a-z]?\b", ref.text) \
                and not re.search(r"https?://|www\.|\bdoi\b|10\.\d{4,9}/", ref.text, re.I):
            # nothing to find it by: no year, no link ("Lederberg J, McCray AT. ’Ome Sweet ’Omics-- A Genealogical
            # Treasury of Words.", IJO_2623 ref 1) -- left as written, so the author must complete it
            vol, pages = volume_pages(ref.text)
            missing = "year of publication" if vol and pages else "journal name, year, volume and page numbers"
            qs.add(ref.para_id, "references.incomplete", missing=missing, **span)

    # Key words.
    from .rules_sections import layout, listed_keywords

    kw = layout(ms).keywords  # also a "Keywords:" line ending an abstract paragraph
    if kw is None:
        kw = next((p for p in ms.paragraphs if re.match(r"^\s*(?:<H\d>)?\s*key\s*-?words?\b", p.text, re.I)), None)
    if kw is not None:
        # the list as the key-words rule reads it ("Diabetes Mellitus, Type 2; ..." is one key word with its comma);
        # "Key words" on its own line, the list below it
        words = listed_keywords(ms) if kw is layout(ms).keywords else [
            w for w in re.split(r"[;,]", re.sub(r"^\s*(?:<H\d>)?\s*key\s*-?words?\s*[:.\-–]?", "", kw.text, flags=re.I))
            if w.strip()]
        if not words and not (style.get("sections.keywords.from_xml") and rules_sections_xml(ms)):
            qs.add(kw.id, "keywords.missing")
        elif len(words) < style.get("keywords.min", 3):
            qs.add(kw.id, "keywords.too_few")

    # Abstract, key words and back matter.
    from . import rules_sections

    rules_sections.add_queries(ms, style, qs)

    # Expected sections.
    # each line: a heading after a manual line break in a longer paragraph ("... [18].\nConclusion") is one too
    heads = {re.sub(r"^<H\d>|:$", "", line.strip()).strip().lower() for p in ms.paragraphs
             for line in p.text.split("\n") if len(line) < 80}
    # a run-in "CONCLUSION: Based on ..." in the body (not the structured abstract's "Conclusion:") is one too
    abstract_ids = {p.id for p in layout(ms).abstract}
    run_in = any(re.match(r"\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?conclusions?\s*[:.\-–]\s*\S", line, re.I)
                 for p in ms.paragraphs if p.id not in abstract_ids and not p.in_table for line in p.text.split("\n"))
    if "conclusion" in style.get("expected_h1", []) and not heads & {"conclusion", "conclusions"} and not run_in \
            and original_article(ms):
        # the Discussion heading as written: "Discussion", "<H1>Discussion", "Discussion:", "4. Discussion"
        target = next((p for p in ms.paragraphs if re.fullmatch(r"\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?discussion\s*[:.]?\s*",
                                                                p.text, re.I)), None)
        if target is not None:
            qs.add(target.id, "sections.conclusion")
    if style.get("queries.unmatched_brackets", False):
        ref_ids = {r.para_id for r in refs}
        for p in ms.paragraphs:
            if p.id in ref_ids or not p.text.strip():
                continue
            span = unmatched_brackets(p.text)
            if span:
                qs.add(p.id, "text.brackets", *span)
    if style.get("queries.drafting_instructions", True):
        ref_ids = {r.para_id for r in refs}
        for p in ms.paragraphs:
            if p.id in ref_ids or not p.text.strip():
                continue
            span = drafting_instruction(p.text)
            if span:
                quote = p.text[span[0]:span[1]]
                qs.add(p.id, "text.instruction", *span, quote=quote if len(quote) <= 80 else quote[:77].rstrip() + "...")
    from . import renumber, rules_figures
    renumber.add_queries(ms, style, qs)
    rules_figures.add_queries(ms, style, qs)
    from . import rules_tables
    rules_tables.add_queries(ms, style, qs)  # G206 empty head cells to merge, G196 entries on lines of their own
    from .rules import add_author_queries  # "Radziwill et al. [6]" where reference 6 is Bae: S "Match author names"
    add_author_queries(ms, style, qs)
    return qs.number(ms)


def write_comments(ms: Manuscript, qs: QuerySet, author: str = AUTHOR, deferred: set[str] = frozenset()) -> int:
    """Add each query as a Word comment on its anchor. Returns how many were placed. Queries on an inserted paragraph
    in `deferred` (the inserts this copy renders) wait for `write_insert_comments`; those whose insert is not
    rendered sit on the whole paragraph the insert stands next to."""
    from .docmodel import explode_runs
    from .tracked import _cut, _spans

    placed = 0
    body = ms.document._body
    for q in qs.queries:
        if q.insert_id and q.insert_id in deferred:
            continue
        if q.insert_id:
            q = Query(q.para_id, 0, 0, q.text, q.check, q.id)
        try:
            p = ms.by_id(q.para_id)
        except KeyError:
            continue
        explode_runs(p.element)
        if q.end > q.start:
            _cut(p.element, q.start)
            _cut(p.element, q.end)
            runs = [r for r, s, e in _spans(p.element) if s >= q.start and e <= q.end and e > s]
        else:
            runs = [r for r, s, e in _spans(p.element) if e > s]
        if not runs:  # a paragraph holding only a picture (an equation or figure given as an image): its run
            runs = [r for r in p.element.iter(_W + "r") if r.find(_W + "drawing") is not None
                    or r.find(_W + "pict") is not None or r.find(_W + "object") is not None]
        if not runs:
            continue
        para = DocxParagraph(p.element, body)
        ms.document.add_comment([Run(runs[0], para), Run(runs[-1], para)], text=f"{q.id}: {q.text}",
                                author=author, initials="EO")
        placed += 1
    return placed


def write_insert_comments(ms: Manuscript, qs: QuerySet, elements: dict, author: str = AUTHOR) -> int:
    """The queries on inserted paragraphs, as Word comments on those paragraphs (`elements`: insert id -> its w:p,
    from frontmatter.render_inserts): on the quoted text when it is there, else on the whole paragraph."""
    from docx.oxml.ns import qn

    placed = 0
    body = ms.document._body
    for q in qs.queries:
        el = elements.get(q.insert_id) if q.insert_id else None
        if el is None:
            continue
        runs = [r for r in el.iter(qn("w:r")) if r.find(qn("w:t")) is not None or r.find(qn("w:tab")) is not None]
        if not runs:
            continue
        pos, chosen = 0, []
        for r in runs:
            t = "".join((x.text or "") if x.tag == qn("w:t") else "\t" for x in r if x.tag in (qn("w:t"), qn("w:tab")))
            if q.end > q.start and pos < q.end and pos + len(t) > q.start:
                chosen.append(r)
            pos += len(t)
        chosen = chosen or runs
        para = DocxParagraph(el, body)
        ms.document.add_comment([Run(chosen[0], para), Run(chosen[-1], para)], text=f"{q.id}: {q.text}",
                                author=author, initials="EO")
        placed += 1
    return placed


def write_sheet(ms: Manuscript, qs: QuerySet, path: str, article: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Queries"
    ws.append([f"Author queries {article}".strip()])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["No.", "Location", "Query", "Author's answer"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for q in qs.queries:
        try:
            text = ms.by_id(q.para_id).text
        except KeyError:
            text = ""
        where = q.quote if q.insert_id else (text[q.start:q.end] if q.end > q.start else text) or "[image]"
        ws.append([q.id, where[:120], q.text, ""])
    for col, width in zip("ABCD", (8, 60, 70, 40)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
