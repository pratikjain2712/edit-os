"""What the editors put at the top of the file, and tables / figures put in their new order.

Top of the tracked and clean copies (tracked insertions), rendered with the current statuses:

  Technical Editing done on 22-October-2026:- <editor>   pre-editing guide, SIGNATURE: "Technical Editing done on
                                                         <date>:- <name>" (every one of 236 editor files signs off)
  Copyediting done on 22-10-26                           when the copy-edit stage ran (the guide is silent; the copy
                                                         editors' most common form, 44 of 154)
  AQ: <editor queries>                                   optional per journal (`header.editor_queries`), off by default
  AQ: <renumbering query>  + a 2-column old -> new table one per renumbering that is rendered: references
                                                         (renumber.py), tables and figures (rules_figures.py)

The sign-off date is the render date (when the editor finalises); the name is the editor's (the review
service passes it in); without one the placeholder goes in and QC flags it (`qc.check`: signoff.name).
The renumbering lines appear only while the renumbering changes they describe are rendered: accepted (clean
copy) or accepted / pending (tracked copy). The queries go below the portal's copyright-notice query when the
file opens with one, as the editors place them; the sign-off lines go first.

`move_blocks` puts renumbered tables (caption, table, notes) and figure legends (with a picture right above)
in their new order: in the tracked copy the block is deleted where it was and inserted, as edited, where it
belongs; in the clean copy it is moved. Blocks that cannot be told apart safely are left where they are.

Settings: `header:` in the style files (base.yaml).
"""

from __future__ import annotations

import copy
import re
from datetime import datetime
from itertools import count

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .changes import ChangeSet
from .docmodel import Manuscript, paragraph_text
from .style import JournalStyle

PLACEHOLDER = "???"
AUTHOR = "Edit OS"
KINDS = {"references": "references", "Table": "tables", "Figure": "figures"}
_SIGNOFF = re.compile(r"^\s*Technical Editing done on\b", re.I)


def settings(style: JournalStyle) -> dict:
    return style.get("header", {}) or {}


def query_label(style: JournalStyle) -> str:
    """"AQ: " or "Query: ": the journal's label for queries written into the text."""
    from . import submission

    label = settings(style).get("query_label") or submission.settings(style).get("copyright_notice_prefix") or "AQ: "
    return label if label.endswith(" ") else label + " "


# --- at edit time: what was renumbered (saved in the manifest) ---------------------------

def plan(ms: Manuscript, style: JournalStyle) -> dict:
    """The renumberings the rules proposed and where the header queries go, for the manifest."""
    from . import renumber, rules_figures, submission

    out: dict = {"renumbering": [], "after": None}
    pl = renumber.plan(ms, style)
    if pl.apply:
        removed = {d for _, d in pl.duplicates}
        out["renumbering"].append({
            "kind": "references", "duplicates": [list(d) for d in pl.duplicates],
            "mapping": [[old, new] for old, new in sorted(pl.mapping.items()) if old not in removed]
            + [[d, pl.mapping[k]] for k, d in pl.duplicates]})
    tables = {n: ids for n, ids in renumber.tables(ms) if n is not None}
    for kind, rp in rules_figures.renumber_plan(ms, style).items():
        blocks = {}
        for old, cap in rp.captions.items():
            blocks[str(old)] = [cap] + (tables.get(old, []) if kind == rules_figures.TABLE else [])
        out["renumbering"].append({"kind": kind, "mapping": [[o, n] for o, n in sorted(rp.mapping.items())],
                                   "blocks": blocks})
    for p in ms.paragraphs[:5]:
        if not p.in_table and submission._COPYRIGHT.match(p.text):
            out["after"] = p.id
            break
    return out


def _rendered(item: dict, changes: ChangeSet, statuses) -> bool:
    """Whether the changes behind a renumbering are in this output."""
    if item["kind"] == "references":
        return any(c.reason == "renumber" and c.status in statuses for c in changes.changes)
    return any(c.reason == "fig_renumber" and c.status in statuses and (c.provenance or {}).get("figure_kind") == item["kind"]
               for c in changes.changes)


# --- the lines -----------------------------------------------------------------------------

def _date(when: datetime, fmt: str) -> str:
    return when.strftime(fmt)


def lines(style: JournalStyle, info: dict, changes: ChangeSet, statuses, *, stages, editor: str | None,
          when: datetime) -> list[tuple[str, object]]:
    """[("p", text) | ("table", rows)] in order. `info` is `plan`'s output from the manifest."""
    cfg = settings(style)
    if not cfg.get("enabled", True):
        return []
    out: list[tuple[str, object]] = []
    so = cfg.get("signoff", {}) or {}
    if so.get("enabled", True):
        name = (editor or "").strip() or so.get("placeholder", PLACEHOLDER)
        out.append(("p", so.get("technical_editing", "Technical Editing done on {date}:- {name}").format(
            date=_date(when, so.get("date_format", "%d-%B-%Y")), name=name)))
        if "copyedit" in stages and so.get("copyediting"):
            out.append(("p", so["copyediting"].format(date=_date(when, so.get("copyediting_date_format", "%d-%m-%y")),
                                                      name=name)))
    label = query_label(style)
    eq = cfg.get("editor_queries", {}) or {}
    texts = cfg.get("texts", {}) or {}
    for key in ("names_affiliations", "heading_levels"):
        if eq.get(key):
            out.append(("q", label + texts[key]))
    rcfg = cfg.get("renumber", {}) or {}
    if rcfg.get("enabled", True):
        items = sorted(info.get("renumbering", []), key=lambda d: ["Table", "Figure", "references"].index(d["kind"]))
        for item in items:
            if not _rendered(item, changes, statuses):
                continue
            key = KINDS[item["kind"]]
            text = rcfg.get(key, "")
            if item.get("duplicates"):
                pairs = ", ".join(f"({a} and {b})" for a, b in item["duplicates"])
                text = rcfg.get("references_duplicates", text).format(pairs=pairs)
            out.append(("q", label + text))
            head = (rcfg.get("columns", {}) or {}).get(key, ["Old", "New"])
            rows = sorted(item["mapping"], key=lambda r: (r[1], r[0]))
            out.append(("table", [list(head)] + [[str(o), str(n)] for o, n in rows]))
    return out


# --- rendering -------------------------------------------------------------------------------

class _Marks:
    def __init__(self, ms: Manuscript, author: str, date: str, tracked: bool, run_format: dict | None = None,
                 align: str | None = None):
        self.run_format = run_format  # the journal's font and size for the runs of new paragraphs (or None)
        self.align = align  # and their alignment (None: the paragraph style's)
        taken = [int(v) for el in ms.document.element.body.iter() if (v := el.get(qn("w:id"))) and v.isdigit()]
        self.ids = count(max([20000] + taken) + 1)
        self.author, self.date, self.tracked = author, date, tracked

    def mark(self, tag: str):
        m = OxmlElement(tag)
        m.set(qn("w:id"), str(next(self.ids)))
        m.set(qn("w:author"), self.author)
        m.set(qn("w:date"), self.date)
        return m


def styled(run, mk: _Marks, fmt: dict | None = None):
    """Give a new run the journal's font and size (`mk.run_format`) and `fmt` (bold, ...)."""
    from .tracked import apply_format

    want = {**(getattr(mk, "run_format", None) or {}), **(fmt or {})}
    if not want:
        return run
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run.insert(0, rpr)
    apply_format(rpr, want)
    if not len(rpr):
        run.remove(rpr)
    return run


def aligned(ppr, mk: _Marks):
    """Give a new paragraph's <w:pPr> the alignment `mk.align` (before its <w:rPr>)."""
    if getattr(mk, "align", None):
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), mk.align)
        ppr.insert(0, jc)
    return ppr


def _para(text: str, mk: _Marks):
    p = OxmlElement("w:p")
    ppr = aligned(OxmlElement("w:pPr"), mk)
    if mk.tracked:
        rpr = OxmlElement("w:rPr")
        rpr.append(mk.mark("w:ins"))
        ppr.append(rpr)
    p.append(ppr)
    if text:
        r = OxmlElement("w:r")
        styled(r, mk)
        t = OxmlElement("w:t")
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        if mk.tracked:
            w = mk.mark("w:ins")
            w.append(r)
            p.append(w)
        else:
            p.append(r)
    return p


def _table(rows: list[list[str]], mk: _Marks):
    tbl = OxmlElement("w:tbl")
    tblpr = OxmlElement("w:tblPr")
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "0")
    w.set(qn("w:type"), "auto")
    tblpr.append(w)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "auto")):
            b.set(qn(k), v)
        borders.append(b)
    tblpr.append(borders)
    tbl.append(tblpr)
    grid = OxmlElement("w:tblGrid")
    for _ in rows[0]:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), "1800")
        grid.append(gc)
    tbl.append(grid)
    for row in rows:
        tr = OxmlElement("w:tr")
        if mk.tracked:
            trpr = OxmlElement("w:trPr")
            trpr.append(mk.mark("w:ins"))
            tr.append(trpr)
        for cell in row:
            tc = OxmlElement("w:tc")
            tcpr = OxmlElement("w:tcPr")
            tcw = OxmlElement("w:tcW")
            tcw.set(qn("w:w"), "1800")
            tcw.set(qn("w:type"), "dxa")
            tcpr.append(tcw)
            tc.append(tcpr)
            tc.append(_para(cell, mk))
            tr.append(tc)
        tbl.append(tr)
    return tbl


def render(ms: Manuscript, items: list[tuple[str, object]], *, tracked: bool, author: str, date: str,
           after: str | None = None, run_format: dict | None = None, align: str | None = None) -> int:
    """Insert the header lines at the top of the body (queries below the paragraph `after` when it is still
    in the file), their runs in `run_format` (the journal's font and size), the paragraphs aligned `align`.
    Returns how many elements were inserted."""
    if not items:
        return 0
    mk = _Marks(ms, author, date, tracked, run_format, align)
    body = ms.document.element.body
    signoff, rest = [], []
    for kind, val in items:
        if kind == "table":
            rest += [_table(val, mk), _para("", mk)]  # a blank line after the table, as the editors leave it
        else:
            (signoff if kind == "p" else rest).append(_para(val, mk))
    els = signoff + rest
    first = next((c for c in body if c.tag in (qn("w:p"), qn("w:tbl"))), None)
    anchor_el = None
    if after:
        try:
            anchor_el = ms.by_id(after).element
        except KeyError:
            anchor_el = None
        if anchor_el is not None and (anchor_el.getparent() is not body or anchor_el in ms.removed):
            anchor_el = None
    for e in signoff:
        if first is not None:
            first.addprevious(e)
        else:
            body.insert(0, e)
    prev = anchor_el if anchor_el is not None else (signoff[-1] if signoff else None)
    for e in rest:
        if prev is not None:
            prev.addnext(e)
        elif first is not None:
            first.addprevious(e)
        else:
            body.insert(0, e)
        prev = e
    return len(els)


# --- moving tables and figure legends ---------------------------------------------------------

def _top(el, body):
    while el is not None and el.getparent() is not body:
        el = el.getparent()
    return el


def _has_media(el) -> bool:
    return any(x.tag in (qn("w:drawing"), qn("w:pict"), qn("w:object")) for x in el.iter())


# A table's notes: "Abbreviations: ...", "*P<0.05", or a list of abbreviations ("pp, percentage points; CI, ...").
_NOTE = re.compile(r"^\s*(?:(?:notes?|sources?|adapted|abbreviations?|data from|reproduced|modified|reprinted)\b|[*†‡§¶#]|"
                   r"(?-i:[\w%.\-–]{1,15}(?:,| =|:)\s+[a-z][^;]{0,60};))", re.I)


def _blocks(ms: Manuscript, item: dict) -> dict[int, list] | None:
    """old number -> the body-level elements of its block, or None when the blocks cannot be told apart."""
    body = ms.document.element.body
    kids = list(body)
    index = {id(e): i for i, e in enumerate(kids)}
    spans: dict[int, tuple[int, int]] = {}
    for old, pids in item.get("blocks", {}).items():
        tops = []
        for pid in pids:
            try:
                t = _top(ms.by_id(pid).element, body)
            except KeyError:
                return None
            if t is None:
                return None
            tops.append(index[id(t)])
        lo, hi = min(tops), max(tops)
        if item["kind"] == "Figure":  # a picture right above its legend moves with it
            k = lo - 1
            while k >= 0 and lo - k <= 2 and kids[k].tag == qn("w:p") and _has_media(kids[k]) \
                    and not paragraph_text(kids[k]).strip():
                lo = k
                k -= 1
        if item["kind"] == "Table":  # notes under the table that renumber.tables does not count ("pp, percentage points;")
            k = hi + 1
            while k < len(kids) and kids[k].tag == qn("w:p"):
                t = paragraph_text(kids[k])
                if t.strip() and not _NOTE.match(t):
                    break
                if t.strip():
                    hi = k
                k += 1
        members = {index[id(_top(ms.by_id(pid).element, body))] for pid in pids}
        for k in range(lo, hi + 1):  # only the block's own paragraphs, notes, blank lines and pictures in between
            e = kids[k]
            if k in members or e.tag == qn("w:tbl") or (e.tag == qn("w:p") and (not paragraph_text(e).strip()
                                                                                or _NOTE.match(paragraph_text(e)))):
                continue
            return None
        spans[int(old)] = (lo, hi)
    ordered = sorted(spans.values())
    if any(a[1] >= b[0] for a, b in zip(ordered, ordered[1:])):
        return None
    if any(kids[k].tag == qn("w:sectPr") for lo, hi in ordered for k in range(lo, hi + 1)):
        return None
    return {old: kids[lo:hi + 1] for old, (lo, hi) in spans.items()}


def _accept(el) -> None:
    """The element as it reads with every revision in it accepted (for the inserted copy)."""
    for tag in ("w:del", "w:moveFrom"):
        for x in list(el.iter(qn(tag))):
            par = x.getparent()
            if par is not None:
                par.remove(x)
    for tag in ("w:ins", "w:moveTo"):
        for x in list(el.iter(qn(tag))):
            par = x.getparent()
            if par is None:
                continue
            if par.tag in (qn("w:rPr"), qn("w:trPr")):
                par.remove(x)
                continue
            i = par.index(x)
            for c in list(x):
                par.insert(i, c)
                i += 1
            par.remove(x)
    for tag in ("w:rPrChange", "w:pPrChange", "w:trPrChange", "w:tcPrChange", "w:tblPrChange",
                "w:commentRangeStart", "w:commentRangeEnd", "w:commentReference", "w:bookmarkStart", "w:bookmarkEnd"):
        for x in list(el.iter(qn(tag))):
            x.getparent().remove(x)


def _reject(el) -> None:
    """The element as it was before the rules' revisions (for the deleted original)."""
    paras = [el] if el.tag == qn("w:p") else list(el.iter(qn("w:p")))
    for p in paras:
        mark = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}")
        if mark is not None:
            mark.getparent().remove(mark)
    for tag in ("w:ins", "w:moveTo"):
        for x in list(el.iter(qn(tag))):
            par = x.getparent()
            if par is not None and par.tag not in (qn("w:rPr"), qn("w:trPr")):
                par.remove(x)
    for tag in ("w:del", "w:moveFrom"):
        for x in list(el.iter(qn(tag))):
            par = x.getparent()
            if par is None:
                continue
            if par.tag in (qn("w:rPr"), qn("w:trPr")):
                par.remove(x)
                continue
            i = par.index(x)
            for c in list(x):
                par.insert(i, c)
                i += 1
            par.remove(x)
    for x in list(el.iter(qn("w:delText"))):
        x.tag = qn("w:t")
    for tag, holder in (("w:rPrChange", "w:rPr"), ("w:pPrChange", "w:pPr")):
        for x in list(el.iter(qn(tag))):
            parent = x.getparent()
            old = x.find(qn(holder))
            keep = {qn("w:rPr"), qn("w:sectPr")} if tag == "w:pPrChange" else set()
            for c in list(parent):
                if c is not x and c.tag not in keep:
                    parent.remove(c)
            parent.remove(x)
            if old is not None:
                for i, c in enumerate(list(old)):
                    parent.insert(i, c)


def _runs(p):
    return [r for r in p.iter(qn("w:r")) if r.getparent().tag in (qn("w:p"), qn("w:hyperlink"), qn("w:smartTag"))]


def _mark_paragraph(p, mk: _Marks, tag: str) -> None:
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        tail = [c for c in ppr if c.tag in (qn("w:sectPr"), qn("w:pPrChange"))]
        (tail[0].addprevious(rpr) if tail else ppr.append(rpr))
    if rpr.find(qn(tag)) is None:
        rpr.insert(0, mk.mark(tag))
    for r in _runs(p):
        if tag == "w:del":
            for t in r.findall(qn("w:t")):
                t.tag = qn("w:delText")
            for t in r.findall(qn("w:instrText")):
                t.tag = qn("w:delInstrText")
        w = mk.mark(tag)
        r.addprevious(w)
        w.append(r)


def _mark_block(el, mk: _Marks, tag: str) -> None:
    for p in ([el] if el.tag == qn("w:p") else list(el.iter(qn("w:p")))):
        _mark_paragraph(p, mk, tag)
    for tr in el.iter(qn("w:tr")):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            prex = tr.find(qn("w:tblPrEx"))
            (prex.addnext(trpr) if prex is not None else tr.insert(0, trpr))
        if trpr.find(qn(tag)) is None:
            trpr.append(mk.mark(tag))


def move_blocks(ms: Manuscript, info: dict, changes: ChangeSet, statuses, *, tracked: bool, author: str,
                date: str) -> list[str]:
    """Put renumbered tables / figure legends in their new order. Returns notes on what was not moved."""
    from .frontmatter import _in_order

    notes = []
    mk = _Marks(ms, author, date, tracked)
    for item in info.get("renumbering", []):
        if item["kind"] == "references" or not _rendered(item, changes, statuses):
            continue
        blocks = _blocks(ms, item)
        if not blocks:
            notes.append(f"{item['kind']}s renumbered but left where they are: their blocks could not be told apart")
            continue
        new = {int(o): n for o, n in item["mapping"]}
        body = ms.document.element.body
        kids = list(body)
        doc = sorted(blocks, key=lambda o: kids.index(blocks[o][0]))  # old numbers in document order
        keep = {doc[k] for k in _in_order([new[o] for o in doc], [1] * len(doc))}
        moving = sorted((o for o in doc if o not in keep), key=lambda o: new[o])
        if not moving:
            continue
        kept = sorted(keep, key=lambda o: new[o])
        tail: dict[int, object] = {}
        for o in moving:
            later = next((k for k in kept if new[k] > new[o]), None)
            els = blocks[o]
            if tracked:
                # A paragraph already deleted (a caption moved into its table, legends.py) stays where it is, deleted.
                els = [e for e in els if not (e.tag == qn("w:p") and e.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}")
                                              is not None and not paragraph_text(e).strip())]
                copies = [copy.deepcopy(e) for e in els]
                for c in copies:
                    _accept(c)
                    _mark_block(c, mk, "w:ins")
                for e in els:
                    _reject(e)
                    _mark_block(e, mk, "w:del")
                place = copies
            else:
                place = els
            if later is not None:
                anchor = blocks[later][0]
                for e in place:
                    anchor.addprevious(e)
            else:
                before = max((k for k in kept if new[k] < new[o]), key=lambda k: new[k])
                prev = tail.get(before, blocks[before][-1])
                for e in place:
                    prev.addnext(e)
                    prev = e
                tail[before] = prev
    return notes


def has_placeholder(ms: Manuscript) -> str | None:
    """Para id of a sign-off line still waiting for the editor's name."""
    for p in ms.paragraphs[:8]:
        if _SIGNOFF.match(p.text) and PLACEHOLDER in p.text:
            return p.id
    return None


def signed(ms: Manuscript) -> bool:
    return any(_SIGNOFF.match(p.text) for p in ms.paragraphs[:8])
