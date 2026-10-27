"""Section rules: body headings, the abstract, key words and the back matter.

What the style editors do (measured on their tracked files, IJO / ABR / NJCP / JCRT / AIAN):

- Headings: numbering is removed ("1. Introduction" -> "Introduction"), minor words in title
  case are lowered ("Materials And Methods"); IJO renames "Materials and Methods" to "Methods" and
  drops the "Introduction" heading. Nobody renames "Methods" to "Materials and Methods" (the
  generic guide's rule), so renames are per-journal settings.
- Abstract: "Abstract:" loses its colon, a missing "Abstract" heading is added, sub-headings get
  a colon and a space ("Purpose:To" -> "Purpose: To"; a sub-heading on its own line gets ": "),
  "Materials and Methods" / "Methodology" -> "Methods" (stylesheets and data).
- Key words: "Keywords:" label, comma + space between, alphabetical order, first one capitalised
  and the rest lower case unless the manuscript writes them capitalised (proper nouns), no final
  period, duplicates removed.
- Back matter: author statements that say nothing ("Funding: None", "Conflict of interest: Nil")
  are deleted and the standard "Financial support and sponsorship / Nil." and "Conflicts of
  interest / There are no conflicts of interest." sections are added before the references, as
  new paragraphs (tracked.PARAGRAPH_BREAK). A real statement keeps its text under the standard
  heading. A nil statement becomes the standard section in place when that keeps the order. A bare
  "Declarations" heading and inline nil ethics statements ("Ethics: Not required") go.
- Patient consent: an author's statement without a heading gets one; with none at all and
  identifiable patient images (decided from the article type and figure legends), the standard
  statement is added with the stylesheet's query. Authors' contributions: an unlabelled list gets
  its heading; per journal, the list from the title page / contribution form / cover letter is added.
  Per journal (IJO), statements on the manuscript's own title page move there: deleted on top, added under the
  heading (only the top copy is deleted when the back already holds the same statements).
  New sections go in the journal's order (consent, contributions, financial support, conflicts).

Everything is a `Change` on the original text; settings live under `sections:` in the style.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx.oxml.ns import qn

from . import provenance, references, submission
from .changes import Change, Layer
from .docmodel import Manuscript, Paragraph
from .style import JournalStyle
from .tracked import PARAGRAPH_BREAK as BR

_NUMBERING = re.compile(r"^\s*(?:\d+(?:\.\d+)*\.?|[IVX]+\.)\s*(?=[A-Za-z])")
_TAG = re.compile(r"^\s*<H\d>\s*")
_MINOR = {"a", "an", "and", "as", "at", "but", "by", "for", "in", "nor", "of", "on", "or", "the", "to",
          "vs", "via", "with", "from", "into"}


def _change(p: Paragraph, start: int, end: int, new: str, layer: Layer, rule: str, fmt=None) -> Change:
    from .rules import _change as make  # rules registers this module; import late to avoid a cycle

    return make(p, start, end, new, layer, rule, fmt)


def _cfg(style: JournalStyle, key: str, default=None):
    return style.get(f"sections.{key}", default)


def _on(style: JournalStyle, key: str) -> bool:
    return bool(_cfg(style, "enabled", True) and _cfg(style, key, False))


def _key(text: str) -> str:
    """'1. Materials And Methods:' -> 'materials and methods'."""
    return _NUMBERING.sub("", _TAG.sub("", text)).strip().rstrip(":.-–— \t").strip().lower()


def in_textbox(p: Paragraph) -> bool:
    """A paragraph inside a drawing's text box (the boxes of a PRISMA flow diagram: "Identification", "Screening",
    "Included"): never a heading."""
    return any(a.tag == qn("w:txbxContent") for a in p.element.iterancestors())


def heading_line(text: str, known) -> tuple[int, int, int, int] | None:
    """A main heading sharing its paragraph with running text through a manual line break: "DISCUSSION \\nTo the best
    ..." (the heading first) or "... reported [18].\\nConclusion" (the heading last line). Returns (heading start,
    heading end, break start, break end): the break (with the spaces around it) becomes a paragraph break. `known`:
    heading keys (`_key`) that count."""
    if "\n" not in text:
        return None
    nl = text.index("\n")
    line = text[:nl]
    if line.strip() and len(line.strip()) <= 60 and _key(line) in known and text[nl:].strip():
        s = len(line) - len(line.lstrip())
        e = len(line.rstrip())
        b = nl + 1
        while b < len(text) and text[b] in " \t\n":
            b += 1
        return s, e, e, b
    nl = text.rindex("\n")
    line = text[nl + 1:]
    if line.strip() and len(line.strip()) <= 60 and _key(line) in known and text[:nl].strip():
        b0 = len(text[:nl].rstrip(" \t\n"))
        s = nl + 1 + len(line) - len(line.lstrip())
        return s, len(text.rstrip()), b0, s
    return None


def main_heading_keys(style: JournalStyle) -> set[str]:
    """Keys of the first-level headings: the H1 list and the renamings that end in one ("Result" -> "Results")."""
    levels = _levels(style)
    h1 = {k for k, v in levels.items() if v == "H1"} | {k.lower() for k in style.get("headings.h1", [])}
    rename = {k.lower(): v for k, v in (_cfg(style, "headings.rename") or {}).items()}
    return h1 | {k for k, v in rename.items() if v.lower() in h1}


def _levels(style: JournalStyle) -> dict[str, str]:
    if not style.get("headings.tag"):
        return {}
    levels = {name: "H1" for name in style.get("headings.h1", [])}
    levels.update({name: "H2" for name in style.get("headings.h2", [])})
    return levels


def _match_case(src: str, new: str) -> str:
    """Write `new` in the case style of `src` (ALL CAPS stays caps)."""
    letters = [c for c in src if c.isalpha()]
    if letters and all(c.isupper() for c in letters) and len(letters) > 1:
        return new.upper()
    return new


def _title_case(text: str) -> str:
    words = text.split(" ")
    out = []
    for i, w in enumerate(words):
        if i and w.lower() in _MINOR:
            out.append(w.lower())
        elif w[:1].islower():
            out.append(w[:1].upper() + w[1:])
        else:
            out.append(w)
    return " ".join(out)


def _sentence_case(text: str) -> str:
    words = text.split(" ")
    return " ".join(w if i == 0 or not w[1:].islower() else w.lower() for i, w in enumerate(words))


def _cased(text: str, case: str) -> str:
    if case == "upper":
        return text.upper()
    if case == "title":
        return _title_case(text.lower() if text.isupper() else text)
    if case == "sentence":
        t = text.lower() if text.isupper() else text
        return t[:1].upper() + _sentence_case(t)[1:]
    return text


# --- structure: where the abstract, key words and back matter are ---------------------

_KW = re.compile(r"^(\s*(?:<H\d>)?\s*)(key\s*-?\s*words?)(\s*[:.\-–]?\s*)", re.I)
_ABSTRACT = re.compile(r"^(\s*)(abstract)\b(\s*[:.\-–]?\s*)", re.I)
_SUBHEADS = ("purpose|aims?|aims? and objectives?|background(?: and aims?| and objectives?)?|objectives?|"
             "introduction|context|importance|design|settings?(?: and design)?|study design|participants|"
             "patients|subjects|methods?|methodology|materials? and methods?|methods? and materials?|subjects? and methods?|"
             "patients? and methods?|statistical analysis(?: used)?|interventions?|main outcome measures?|"
             "observations?|case(?: presentation| description| report| summary)?|results?|findings|"
             "discussion|conclusions?(?: and relevance)?|clinical relevance|significance")
# "Purpose: To ...", or with a dash ("Purpose–To report ...", "Design-Retrospective ...": IJO_2733) before a capital
_SUB_INLINE = re.compile(r"^(\s*)(" + _SUBHEADS + r")(\s*:\s*|\s*[-–—]\s*(?=(?-i:[A-Z])))(?=\S)", re.I)
_SUB_ALONE = re.compile(r"^(\s*)(" + _SUBHEADS + r")(\s*:?\s*)$", re.I)


def _words(text: str) -> int:
    return len(re.findall(r"\w+", text))


def _word_count(text: str) -> int:
    """Words as Word counts them: whitespace-separated tokens ("COVID-19", "3.5", "45%" and "p<0.05" are one
    each); a lone dash or symbol is not a word."""
    return sum(1 for t in re.sub(r"</?H\d>", " ", text).split() if re.search(r"\w", t))


def _abstract_words(text: str) -> int:
    """`_word_count` of an abstract paragraph without its structured sub-heading ("Background:", "Methods" on its own
    line): the 250-word limit is on the text."""
    if _SUB_ALONE.match(text):
        return 0
    m = _SUB_INLINE.match(text)
    return _word_count(text[m.end():] if m else text)


@dataclass
class Layout:
    abstract_head: Paragraph | None = None   # the "Abstract" heading paragraph (or inline "Abstract: ...")
    abstract: list[Paragraph] = field(default_factory=list)  # abstract text paragraphs
    keywords: Paragraph | None = None
    keywords_first: bool = False  # the key words stand above the abstract (S: "Keywords ... after the Abstract")
    keywords_at: int = 0  # the key words are a later line of `keywords` ("... addiction.\nKeywords: ..."): its offset


def layout(ms: Manuscript) -> Layout:
    """Locate the abstract and key words in the front of the manuscript."""
    cached = getattr(ms, "_sections_layout", None)
    if cached is not None:
        return cached
    paras = [p for p in ms.paragraphs if not p.in_table]
    ref_head = references.find_heading(ms)
    limit = ms.paragraphs[ref_head].id if ref_head is not None else None
    front = []
    for p in paras:
        if p.id == limit:
            break
        front.append(p)
    front = front[: max(30, len(front) * 2 // 3)]
    lay = Layout()
    kw_i = next((i for i, p in enumerate(front) if _KW.match(p.text) and len(p.text) < 400), None)
    if kw_i is None:
        # the key words on the last line of an abstract paragraph, after a manual line break
        for i, p in enumerate(front):
            m = re.search(r"\n[ \t]*(?=key\s*-?\s*words?\s*[:.\-–])", p.text, re.I)
            if m and len(p.text) - m.end() < 400 and "\n" not in p.text[m.end():].strip():
                kw_i, lay.keywords_at = i, m.end()
                break
    if kw_i is not None:
        lay.keywords = front[kw_i]
    def is_head(text: str) -> bool:
        m = _ABSTRACT.match(text)
        if not m:
            return False
        rest = text[m.end():]
        return not rest.strip() or bool(_SUB_ALONE.match(rest) or _SUB_INLINE.match(rest)) or _words(rest) > 25

    head_i = next((i for i, p in enumerate(front[:kw_i] if kw_i is not None else front) if is_head(p.text)), None)
    if head_i is not None:
        lay.abstract_head = front[head_i]
        if _words(front[head_i].text) > 25:
            lay.abstract.append(front[head_i])  # inline "Abstract: text"
        stop = kw_i if kw_i is not None else len(front)
        if lay.keywords_at:
            stop += 1  # the paragraph ending in the key words is the abstract's last
        for p in front[head_i + 1:stop]:
            if not p.text.strip():
                continue
            if kw_i is None and _key(p.text) in ("introduction", "background") and not _SUB_INLINE.match(p.text):
                break
            if kw_i is None and len(lay.abstract) >= 12:
                break
            lay.abstract.append(p)
        lay.abstract = lay.abstract[:15]
    elif kw_i is not None and (after := next((i for i in range(kw_i + 1, min(kw_i + 4, len(front)))
                                              if is_head(front[i].text)), None)) is not None:
        # The key words above the abstract ("KEY WORDS: ..." then "ABSTRACT: We report ..."): the abstract is the
        # heading's own text (or the paragraph under it) and the structured parts that follow.
        lay.abstract_head = front[after]
        lay.keywords_first = True
        rest = [p for p in front[after + 1:] if p.text.strip()]
        if _words(front[after].text) > 25:
            lay.abstract.append(front[after])
        elif rest:
            lay.abstract.append(rest.pop(0))
        if lay.abstract and (lay.abstract[-1] is front[after] or _SUB_INLINE.match(lay.abstract[-1].text)
                             or _SUB_ALONE.match(lay.abstract[-1].text)):
            for p in rest[:12]:
                if not (_SUB_INLINE.match(p.text) or _SUB_ALONE.match(p.text)) \
                        or _key(p.text) in ("introduction", "background"):
                    break
                lay.abstract.append(p)
    elif kw_i is not None:
        # No heading: walk back from the key words over sub-headings and long paragraphs.
        found = []
        for p in reversed(front[:kw_i]):
            if not p.text.strip():
                continue
            if _SUB_ALONE.match(p.text) or _SUB_INLINE.match(p.text) or _words(p.text) >= 15:
                found.append(p)
                if len(found) >= 15:
                    break
            else:
                break
        lay.abstract = list(reversed(found))
    ms._sections_layout = lay
    return lay


def abstract_ids(ms: Manuscript) -> set[str]:
    """Paragraphs inside the abstract (sub-headings there are not body headings)."""
    return {p.id for p in layout(ms).abstract}


def _structured(lay: Layout) -> bool:
    return sum(1 for p in lay.abstract if _SUB_INLINE.match(p.text) or _SUB_ALONE.match(p.text)) >= 2


# --- body headings ---------------------------------------------------------------------

def _tag_start(text: str, s: int) -> int:
    """Where a heading tag inserted before text[s:] starts: at the very start of the paragraph when only
    spaces come before it (they go with it: " Methods" -> "<H1>Methods", not " <H1>Methods")."""
    return 0 if not text[:s].strip() else s


def rule_headings(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not _cfg(style, "enabled", True):
        return []
    levels = _levels(style)
    rename = {k.lower(): v for k, v in (_cfg(style, "headings.rename") or {}).items()}
    drop = {k.lower() for k in _cfg(style, "headings.drop") or []}
    case = _cfg(style, "headings.case", "keep")
    known = set(levels) | set(rename) | drop | {k.lower() for k in style.get("headings.h1", [])}
    skip = abstract_ids(ms) | {p.id for p in [layout(ms).abstract_head, layout(ms).keywords] if p}
    # The reference list's own heading in the journal's wording ("Reference" -> "References"); only that
    # paragraph: "Reference" elsewhere is a table's reference category.
    ref_name = _cfg(style, "headings.reference_heading")
    ref_i = references.find_heading(ms) if ref_name else None
    ref_id = ms.paragraphs[ref_i].id if ref_i is not None else None
    mains = main_heading_keys(style)
    out = []
    for p in ms.paragraphs:
        full = p.text
        if p.in_table or p.id in skip or not full.strip() or in_textbox(p):
            continue
        base, end = 0, len(full)
        if len(full) > 60 or "\n" in full.strip():
            # a main heading on a line of its own inside a longer paragraph ("DISCUSSION \nTo the best ..."): that
            # line (heading_tags breaks the paragraph there)
            hl = heading_line(full, mains)
            if hl is None:
                if len(full) > 60:
                    continue
            else:
                base, end = hl[0], hl[1]
        text = full[base:end]
        whole = base == 0 and end == len(full)

        def ch(s: int, e: int, new: str, fmt=None) -> Change:
            return _change(p, base + s, base + e, new, Layer.STRUCTURE, "sections_headings", fmt)

        key = _key(text)
        if whole and p.id == ref_id and key != ref_name.lower() and (
                re.fullmatch(r"references?|bibliography|reference list", key) or references.is_heading(text)):
            rename = {**rename, key: ref_name}
        elif key not in known:
            continue
        tag = _TAG.match(text)
        lead = tag.end() if tag else len(text) - len(text.lstrip())
        num = _NUMBERING.match(text[lead:])
        core_start = lead + (num.end() if num else 0)
        core = text[core_start:].rstrip()
        core = core[:-1].rstrip() if core.endswith(":") else core
        if p.id == ref_id and whole and core.rstrip(":.-–—_ ") != core:
            # "References-", "REFERENCES." (IJO 2733): the reference list's heading loses its trailing marks
            core = core.rstrip(":.-–—_ ")
            out.append(ch(core_start + len(core), len(text), ""))
        core_end = core_start + len(core)
        if key in drop and key not in levels:
            if whole:
                out.append(ch(0, len(text), ""))
            continue
        if num and _on(style, "headings.strip_numbering"):
            s = lead + num.start()
            out.append(ch(s, core_start, ""))
        new = core
        if key in rename:
            new = _match_case(core, rename[key])
        if case and case != "keep" and levels.get(key, "H1") == "H1":
            new = _cased(new, case)
        elif _on(style, "headings.title_minor_words") and not new.isupper():
            new = " ".join(w.lower() if i and w.lower() in _MINOR and w[:1].isupper() else w
                           for i, w in enumerate(new.split(" ")))
        if new != core:
            level = levels.get(new.lower())
            if level and key not in levels and not tag:  # heading_tags did not know the old name
                new = f"<{level}>{new}"
            start = _tag_start(full, base + core_start) - base if new.startswith("<H") else core_start
            out.append(ch(start, core_end, new))
        if _on(style, "headings.bold") and new == core:
            out.append(ch(core_start, core_end, core, {"bold": True}))
    return out


# --- letters: salutation ------------------------------------------------------------------

_SALUTATION = re.compile(
    r"^(?P<lead>\s*)(?P<sal>(?:(?:my\s+)?dear|respected)\s+(?:sir|madam|sir\s*(?:/|or|and)\s*madam|editor(?:[- ]in[- ]chief)?)"
    r"|sirs?|madam|sir\s*(?:/|or|and)\s*madam)(?P<punct>\s*[,:;!]?)(?=\s*$|\s+[A-Z])", re.I)


def rule_salutation(ms: Manuscript, style: JournalStyle) -> list[Change]:
    """Letters to the editor open with the journal's salutation (sections.letter_salutation, e.g. "Dear Editor,"):
    "Sir,", "Dear Sir/Madam,", "Respected Editor-in-Chief," -> "Dear Editor," (IJO, IDOJ, IJD, IJN / DSHMJ, JOACP
    sheets). Only an existing salutation is replaced; the first one in the body."""
    want = _cfg(style, "letter_salutation")
    if not want or not _cfg(style, "enabled", True):
        return []
    from .rules import body_paragraphs

    for p in body_paragraphs(ms)[:60]:
        m = _SALUTATION.match(p.text)
        if not m or not m.group("punct").strip() and p.text[m.end():].strip():
            continue  # "Sir Ronald Fisher showed": a salutation ends the paragraph or has its comma
        old = p.text[m.start("sal"):m.end("punct")]
        if old != want:
            return [_change(p, m.start("sal"), m.end("punct"), want, Layer.STRUCTURE, "sections_salutation")]
        return []
    return []


# --- abstract ---------------------------------------------------------------------------

# --- sub-headings --------------------------------------------------------------------
#
# "Statistical analysis" on its own line inside Methods -> "<H2>Statistical analysis". Measured on the
# style editors' files (160 articles): a short line (at most 8 words) inside a main section, without a
# final period, starting with a capital (or "2.1"), followed by running text, is tagged in 92% of cases
# when it is bold or italic, and in 93% when it is plain and the next paragraph is long (>= 200 chars).
# Statements ("Funding", "Not applicable"), legends and "ABBR: expansion" lines are left out.

_SUB_SKIP = re.compile(
    r"^(acknowledg|funding|financial|conflicts?\b|competing|authors?['’]?\s+contrib|declaration|ethic|"
    r"availability|data availability|source of support|figure legends?|legends?\b|key\s*-?\s*words?|"
    r"not applicable|not statistically|none\b|nil\b|consent|abbreviations?|list of abbreviations|tables?\b|"
    r"figures?\b|statements?\b|authors?\b|text\b)", re.I)
# Nested sub-headings: "Author Impact Analysis", "Keyword Co-occurrence Network" are headings (a bibliometric study's)
_SUB_SKIP_NESTED = re.compile(_SUB_SKIP.pattern.replace("key\\s*-?\\s*words?|", "key\\s*-?\\s*words?\\s*(?::|$)|")
                              .replace("authors?\\b|", "authors?[’']?s?\\s+(?:contrib|declar|statement|detail|info)|"), re.I)
_SUB_NUMBER = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+(?=[A-Z])")
_RUN_IN = re.compile(r"^[A-Z][\w’'\- ]{2,40}:\s+[A-Z]")
_ABBR_LINE = re.compile(r"^[A-Z][A-Za-z]*[A-Z]\w*\s*[:=–,-]\s")


def _all_fmt(p: Paragraph, prop: str) -> bool:
    from .rules_format import char_formats

    letters = [f for f, c in zip(char_formats(p), p.text) if c.strip()]
    return bool(letters) and all(f.get(prop) for f in letters)


_NOT_BACK_HEAD = re.compile(r"^(?:figures?|tables?|(?:figure\s+)?legends?|appendix|appendices|supplementary|"
                            r"abbreviations?|list of abbreviations|references?|bibliography|declarations?$)\b", re.I)


def _back_matter_head(p: Paragraph, after: list[Paragraph], max_words: int, min_next: int) -> bool:
    """A statement heading after the Conclusion ("Data availability", "Consent for publication"): a short line
    of its own, bold or in a Word heading style, over the statement's text. The labels the back-matter rules
    know (conflicts of interest, funding, acknowledgement) are theirs."""
    text = p.text.strip()
    if p.in_table or len(text.split()) > max_words or text.endswith((".", ",", ";")) or ":" in text.rstrip(":")[:-1] \
            or not text[:1].isupper() or _NOT_BACK_HEAD.match(text) or _LABEL.match(text) and _LABEL.match(text).group("label"):
        return False
    from . import rules_figures
    if rules_figures.caption_match(p) is not None:
        return False
    nxt = next((q for q in after if q.text.strip()), None)
    if nxt is None or nxt.in_table or len(nxt.text.strip()) < 15 or _all_fmt(nxt, "bold"):
        return False
    # A plain line needs running text after it, as a plain sub-heading in the body does; a statement's label on a line
    # of its own does not ("Use of Artificial Intelligence:" over its statement, 1683).
    return _all_fmt(p, "bold") or (p.style or "").lower().startswith("heading") or len(nxt.text) >= min_next \
        or text.endswith(":") and bool(_STATEMENT_WORDS.search(text))


def _emph(p: Paragraph, prop: str, s: int = 0, e: int | None = None) -> bool:
    """Every letter and digit of p.text[s:e] has `prop` (punctuation and spaces aside: "**Research Instruments**:")."""
    from .rules_format import char_formats

    fm = char_formats(p)
    if len(fm) != len(p.text):
        return False
    e = len(p.text) if e is None else e
    letters = [f for f, c in zip(fm[s:e], p.text[s:e]) if c.isalnum()]
    return bool(letters) and all(f.get(prop) for f in letters)


def _style_level(p: Paragraph) -> int | None:
    """The level of a Word heading style ("Heading 2" -> 2)."""
    m = re.match(r"heading\s*(\d)", (p.style or "").strip().lower())
    return int(m.group(1)) if m else None


# A methods section's usual sub-heading, given as a run-in label ("Study Design: Eligible ...", "Statistical analysis :")
_METHODS_LABEL = re.compile(r"\b(?:design|setting|participants?|sampl\w*|data\s+(?:collection|sources?|extraction|analysis)"
                            r"|procedures?|interventions?|ethic\w*|statistic\w*|instruments?|tools?|measures?|outcomes?"
                            r"|inclusion|exclusion|recruitment|population|search\s+strategy|eligibility|quality\s+assessment"
                            r"|study\s+selection|questionnaires?|validity|reliability)\b", re.I)
# The methods sub-sections every study has, as a plain "Statistical Analysis:" line over their text (not an instrument's
# name, "Attitudes toward Addiction Questionnaire:", which the editors leave as typed: 1764).
_METHODS_SECTION = re.compile(r"(?:statistical|data)\s+analys[ie]s|study\s+(?:design|setting|population|area)(?:\s+and\s+"
                              r"setting)?|sample\s+size(?:\s+calculation)?|sampling(?:\s+(?:method|technique))?|"
                              r"data\s+collection(?:\s+(?:tools?|procedure|methods?))?|ethical\s+considerations?|"
                              r"inclusion\s+(?:and\s+exclusion\s+)?criteria|exclusion\s+criteria|study\s+procedure")
_SUB_NUMBER1 = re.compile(r"^(\d+)\.\s+(?=[A-Z])")  # "1. Demographic Characteristics" (nested: bold / styled only)
_SUB_LETTER = re.compile(r"^([a-h])\.\s+(?=[A-Z])")  # "d.    Synthesis of Results" (nested: bold / styled only; 997)
# A run-in identifier label ("Study/Project ID:", "Registration No.:") is a statement, not a heading (1153: editors keep it)
_ID_LABEL = re.compile(r"/|\bID\b|\bNo\.?$|\bnumbers?\b|\bcodes?\b", re.I)
_CASE_HEAD = re.compile(r"[Cc]ase\s*\d+\b|[Cc]ase\s+(?:[IVX]+|[A-Z])\b")  # "Case 1", "Case II": one of a series' cases
_RUN_IN_HEAD = re.compile(r"^(\s*)([A-Z][^:.;?!]{1,60}?)(\s*:\s+)(?=[A-Z(\d])")


@dataclass
class _Sub:
    p: Paragraph
    bold: bool
    italic: bool
    num: object
    section: int
    end: int                  # end of the heading line in p.text
    brk: tuple | None = None  # (start, end) of the break after it, made a paragraph break
    parent: bool = False      # another heading follows right below it
    styled: int | None = None
    back: bool = False        # a back-matter statement head: always <H2>
    run_in: bool = False      # "Study Design: Eligible ..." split into heading + text
    num_id: str | None = None  # Word's automatic numbering list the head is an item of
    strip_period: bool = False  # "Outcome Expectancy and ... skills." among sibling heads: the period goes
    promote: bool = False     # a main section the author added (<H1>)


def _all_caps(text: str) -> bool:
    letters = [c for c in _TAG.sub("", text) if c.isalpha()]
    return len(letters) >= 4 and all(c.isupper() for c in letters)


def _num_id(p: Paragraph) -> str | None:
    from docx.oxml.ns import qn

    el = p.element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    return el.get(qn("w:val")) if el is not None else None


def rule_subheadings(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not style.get("headings.tag") or not style.get("headings.subheadings", False):
        return []
    from . import rules_figures

    max_words = style.get("headings.subheading_max_words", 8)
    back_matter = style.get("headings.back_matter_subheadings", False)
    min_next = style.get("headings.subheading_plain_min_next", 200)
    # Nested sub-headings (per journal): a head over another head is the parent (<H2>) of those below it (<H3>), heads
    # over a figure / table caption or a short paragraph, "Antimicrobial Activity: Synthesis and ..." heads, single-level
    # numbered heads, a bold run-in label alone in its section split off, Word heading styles giving the level.
    nested = style.get("headings.nested_subheadings", False)
    main = ({n.lower() for n in style.get("headings.h1", [])} | set(_levels(style))
            | {k.lower() for k in (_cfg(style, "headings.rename") or {})}
            | {k.lower() for k in (_cfg(style, "headings.drop") or [])}
            | {"introduction", "background", "materials and methods", "material and methods", "methods",
               "results", "discussion", "case report", "case presentation", "case description"})
    known = set(_levels(style)) | {"abstract", "references", "conclusion", "conclusions"}
    ref_i = references.find_heading(ms)
    paras = ms.paragraphs[:ref_i] if ref_i is not None else ms.paragraphs
    skip = abstract_ids(ms) | {p.id for p in [layout(ms).abstract_head, layout(ms).keywords] if p}

    def shape(p: Paragraph, text: str, emph: bool, words: int | None = None) -> bool:
        """A heading's shape: short, no final period, capital (or numbered) start, not a statement / legend / ABBR."""
        if text.endswith((".", ",", ";")) or len(text.split()) > (words or max_words) \
                or (_SUB_SKIP_NESTED if nested else _SUB_SKIP).match(text) \
                or _ABBR_LINE.match(text) or rules_figures.caption_match(p) is not None:
            return False
        if ":" in text[:-1]:
            # "Comparator (C): Non-smokers" is a list item; "Antimicrobial Activity: Synthesis and Contextual
            # Comparison" in bold is a heading
            before, after = text[:text.index(":")].split(), text[text.index(":") + 1:].split()
            # "Mechanism: Why can OVM cause enophthalmos?": a topic and its question, a heading in any journal
            question = text.endswith("?") and 1 <= len(before) <= 4 and len(after) >= 2 and after[0][:1].isupper()
            if not (question or nested and emph and len(before) >= 2 and len(after) >= 2 and after[0][:1].isupper()):
                return False
        # nested: a bold / italic head typed in lower case ("study participants and sampling": 1039) is one too
        return bool(text[:1].isupper() or _SUB_NUMBER.match(text) or nested and emph and (
            _SUB_NUMBER1.match(text) or _SUB_LETTER.match(text) or text[:1].islower()))

    def emphasis(p: Paragraph, e: int) -> tuple[bool, bool, int | None]:
        return _emph(p, "bold", 0, e), _emph(p, "italic", 0, e), _style_level(p) if nested else None

    def head_like(q: Paragraph) -> bool:
        """`q` could be a heading itself (for the one above it to be its parent)."""
        t = q.text.strip()
        if not t or q.in_table or len(t) > 100 or in_textbox(q) or _key(q.text) in known or _key(q.text) in main:
            return False
        b, it, sl = emphasis(q, len(q.text))
        return (b or it or sl is not None) and shape(q, t, b or it or sl is not None)

    # A section of its own that the author added to the main ones (a review's "OCULAR MICROBIOME IN DISEASES", "The Role
    # of Systemic Factors in Ocular Microbiome Research" after bold italic heads, IJO_2623: editors <H1>): a bold head
    # over another head, in capitals like the main headings, or (one level of sub-heads) bold roman like them after
    # bold italic heads.
    first_keys = {n.lower() for n in style.get("headings.h1", [])} | {"materials and methods", "methods"}
    mains = [q for q in paras if _key(q.text) in first_keys and q.text.strip() and not q.in_table]
    caps_main = bool(mains) and 2 * sum(_all_caps(q.text) for q in mains) > len(mains)
    roman_main = bool(mains) and all(_emph(q, "bold") and not _emph(q, "italic") for q in mains)
    found: list[_Sub] = []
    near_miss: list[_Sub] = []
    in_body = in_back = False
    back_tagged = None
    section = 0
    section_key: dict[int, str] = {}  # section -> its main heading's key ("methods")
    run_in: dict[int, int] = {}  # section -> paragraphs opening with a run-in label ("Histopathology: The ...")
    run_in_bold: dict[int, list] = {}  # nested: section -> bold run-in labels (split off when alone)
    last_first = ""
    for i, p in enumerate(paras):
        text = p.text.strip()
        if not text or p.in_table or p.id in skip or in_textbox(p):
            continue
        key = _key(p.text)
        if key in main and key in first_keys:
            last_first = key  # the first-level section a statement head ("Ethical considerations") sits in
        if key in main:
            section_key[section + 1] = key
            # The statement heads after the Conclusion ("Conflicts of interest") do not open a body section.
            in_body = key not in ("conclusion", "conclusions") and not (
                back_matter and in_back and (_levels(style).get(key) == "H2" or re.fullmatch(_ACK, key)))
            in_back = not in_body and (in_back or key in ("conclusion", "conclusions"))
            section += 1
            continue
        if _RUN_IN.match(text) and len(text) > 100:
            run_in[section] = run_in.get(section, 0) + 1
        if back_matter and in_back and key not in known and _back_matter_head(p, paras[i + 1:], max_words, min_next):
            if back_tagged is None:  # the statement heads the back-matter rule tags itself (consent, contributions)
                back_tagged = {c.para_id for c in rule_back_matter(ms, style) if "<H" in c.new}
            if p.id not in back_tagged:  # always <H2>; not counted as a bold body head
                found.append(_Sub(p, False, False, None, section, len(p.text.rstrip()), back=True))
            continue
        if not in_body or key in known or _TAG.match(text):
            continue
        end, brk = len(p.text.rstrip()), None
        if nested and "\n" in text:
            # "1. Demographic Characteristics and Baseline Homogeneity\t\nA total of 68 ..." in a heading style (or
            # a bold first line): the first line is the heading, the rest its text
            nl = p.text.index("\n")
            first = p.text[:nl].rstrip()
            rest_at = nl + 1 + len(p.text[nl + 1:]) - len(p.text[nl + 1:].lstrip())
            if first.strip() and len(p.text[rest_at:]) >= 100 and (_style_level(p) or _emph(p, "bold", 0, len(first))):
                end, brk, text = len(first), (len(first), rest_at), first.strip()
        if len(text) > 100:
            if nested and (m := _RUN_IN_HEAD.match(p.text)) and len(m.group(2).split()) <= 6 \
                    and not _emph(p, "bold", m.end(), len(p.text)) \
                    and not _SUB_SKIP.match(m.group(2)) and _key(m.group(2)) not in known | main:
                bold_label = _emph(p, "bold", m.start(2), m.end(2))
                # (one word, "Intervention:", "Outcomes:", is a PICOS list item). A plain label stays run in ("Intervention
                # Group (n=28):" beside its plain "Control Group (n=33):", 1039; "Inclusion and exclusion criteria:",
                # 1559: editors 4 of 4), as does an identifier ("Study/Project ID:", 1153).
                methods = _METHODS_LABEL.search(m.group(2)) and 2 <= len(m.group(2).split()) <= 5
                if bold_label and not _ID_LABEL.search(m.group(2)):
                    run_in_bold.setdefault(section, []).append(
                        _Sub(p, bold_label, False, None, section, m.end(2), (m.end(2), m.end()), run_in=bool(methods)))
            continue
        bold, italic, styled = emphasis(p, end)
        emph = bold or italic or styled is not None
        num_id = None
        if nested and p.numbered and styled is None:
            # An item of a numbered list ("1. Family and Dyadic Based Partnership Care" over its indented text, 1580);
            # a bold item over running text of its own is a head in Word's numbering (997: editors tag 22 of 22).
            nxt0 = next((q for q in paras[i + 1:] if q.text.strip()), None)
            if not bold or nxt0 is None or nxt0.numbered or (nxt0.style or "") == (p.style or "") or len(nxt0.text) < 100:
                continue
            num_id = _num_id(p)
        near = strip_period = False
        if not shape(p, text, emph):
            # A near miss, a head only if a sibling head in its section is set the same way: a longer one ("Distribution
            # of Ocular Microbiome Studies by Disease and Sample Size", IJO_2623) or one ending in a period ("Outcome
            # Expectancy and outcome value associated with poor communication skills.", 1153).
            alt = text[:-1].rstrip() if text.endswith(".") and text.count(".") == 1 else text
            if not (emph and brk is None and shape(p, alt, emph, 14)):
                continue
            near, strip_period = True, alt != text
        num = _SUB_NUMBER.match(text) or (nested and emph and (_SUB_NUMBER1.match(text) or _SUB_LETTER.match(text))) \
            or None
        if brk is None:
            later = [q for q in paras[i + 1:] if q.text.strip()]
            nxt = later[0] if later else None
            parent = False
            if bold and nxt is not None and (caps_main and _all_caps(text) and head_like(nxt) or (
                    not nested and roman_main and not italic and sum(
                        1 for g in found if g.section == section and g.bold and g.italic and not g.back) >= 2)):
                found.append(_Sub(p, bold, italic, num, section, end, None, True, styled, promote=True))
                continue
            if nested and emph and nxt is not None and head_like(nxt):
                parent = True  # "Publication Performance and Growth Trends" over "Research Performance Overview"
            elif nested and emph and nxt is not None and not nxt.in_table and rules_figures.caption_match(nxt) is not None:
                # a head over a figure or table: its text comes after the caption (and the table)
                body = next((q for q in later if not q.in_table and rules_figures.caption_match(q) is None
                             and not in_textbox(q)), None)
                if body is None or len(body.text) < 100 and not head_like(body):
                    continue
            elif nxt is None or nxt.in_table or len(nxt.text) < 100:
                # a list item or a label over a table, not a heading over running text; nested: a bold head over a
                # short paragraph that is a sentence
                if not (nested and emph and nxt is not None and not nxt.in_table and len(nxt.text.strip()) >= 40
                        and nxt.text.rstrip().endswith(".") and not _emph(nxt, "bold")):
                    continue
            # "Statistical Analysis:" on its own line over the methods' running text is a heading however it is set
            # (a plain "Results:" over a list is not): a methods sub-section's name of two to five words
            methods_head = text.endswith(":") and _METHODS_SECTION.fullmatch(_key(text)) is not None \
                and len(nxt.text) >= 100 and "method" in last_first
            if not (bold or italic or styled is not None or methods_head) \
                    and (text.endswith(":") or len(nxt.text) < min_next):
                continue
            sub = _Sub(p, bold, italic, num, section, end, None, parent, styled, num_id=num_id, strip_period=strip_period)
            (near_miss if near else found).append(sub)
        else:
            found.append(_Sub(p, bold, italic, num, section, end, brk, False, styled))
    split_sections = set()
    for sec, items in run_in_bold.items():
        # A bold run-in label alone in its section, or one a methods section uses as a heading ("Study Design:",
        # "Ethical consideration:", "Statistical analysis :"), is split off; a series of other run-in labels
        # ("Triadic Ecosystem Integration: ...") is a list, left as typed.
        alone = len(items) == 1 and run_in.get(sec, 0) <= 1 and items[0].bold
        for f in items:
            if alone or f.run_in:
                f.run_in = True
                found.append(f)
                split_sections.add(sec)
    for f in near_miss:  # a sibling in the same section set the same way
        if any(g.section == f.section and not g.back and not g.run_in and (g.bold, g.italic) == (f.bold, f.italic)
               for g in found):
            found.append(f)
    order = {p.id: k for k, p in enumerate(ms.paragraphs)}
    found.sort(key=lambda f: order[f.p.id])
    # A case report's Methods section holds a statement or two (ethics, consent), not sub-sections: the first own-line
    # head in it that is not one of the cases opens a section of its own (<H1>), as the case report's other sections
    # are first level (IJO G139); the heads after it ("Case 1", "Summary of Previously Reported Cases") are its <H2>s.
    # IJO_702: "Methods" (the ethics statement) / "Literature search strategy" <H1> / "Case 1" ... <H2>.
    first_level = set()
    if style.get("headings.case_report_methods_sections", False) and _is_case_report(ms):
        seen = set()
        for f in found:
            if f.back or f.section in seen:
                continue
            seen.add(f.section)
            if "method" in section_key.get(f.section, "") and not _CASE_HEAD.match(f.p.text[:f.end].strip()):
                first_level.add(f.p.id)
    head_case = _cfg(style, "headings.case", "keep")
    pos = {p.id: k for k, p in enumerate(paras)}

    def heads_run_ins(f: _Sub) -> bool:
        """"Strengths and Limitations:" over "Strengths: First ..." (1683): the head of the run-in labels after it."""
        nxt = next((q for q in paras[pos[f.p.id] + 1:] if q.text.strip()), None)
        m = _RUN_IN_HEAD.match(nxt.text) if nxt is not None else None
        return bool(m) and m.group(2).split()[0].lower() in f.p.text[:f.end].lower().split()

    out = []
    parent: _Sub | None = None
    last_sec = None
    before: dict[int, list[_Sub]] = {}  # section -> the heads tagged before this one
    for f in found:
        p = f.p
        if f.section != last_sec:
            parent, last_sec = None, f.section
        if p.text[:f.end].rstrip().endswith(":") and run_in.get(f.section, 0) >= 2 and not f.run_in \
                and f.section not in split_sections and not heads_run_ins(f):
            continue  # "Staging and Treatment:" among run-in labels ("Ocular examination: The ...") is one of them
        prev = before.setdefault(f.section, [])
        # Italic heads under bold ones (earlier in their section: an italic head right under the main heading is
        # second level, IJO_2623), and "2.1.1"-numbered heads, are third level.
        under_bold = any(g.bold for g in prev if not g.back)
        deep = f.num is not None and f.num.re is _SUB_NUMBER and f.num.group(1).count(".") >= 2
        level = "H3" if (f.italic and not f.bold and under_bold) or deep else "H2"
        if nested and not f.back:
            if f.bold and f.italic and any(g.bold and not g.italic and not g.run_in and not g.back for g in prev):
                level = "H3"  # bold italic heads under a bold roman one (1209)
            elif f.num_id and any(g.num_id and g.num_id != f.num_id for g in prev):
                level = "H3"  # a second numbered list of heads under the first one's last (997)
            elif f.run_in and any(not g.run_in and not g.back for g in prev) \
                    and not _METHODS_LABEL.search(p.text[:f.end]):
                level = "H3"  # a run-in label split off under the section's own-line heads ("Community engagement", 1153)
        promote = f.promote
        if not f.back:
            prev.append(f)
        if nested and not f.back:
            numbered = f.num is not None and f.num.re is _SUB_NUMBER1
            if f.parent or numbered:
                # a numbered head ("4. Summary of Quantitative Findings") is the parent of the unnumbered ones after it
                parent = f
            elif parent is not None:
                if parent.styled is not None and f.styled is not None and f.styled <= parent.styled:
                    parent = None  # a sibling in the parent's own heading style
                else:
                    level = "H3"
        if p.id in first_level or promote:
            level = "H1"
        lead = len(p.text) - len(p.text.lstrip())  # the tag goes at the very start: leading spaces go
        if f.num and _on(style, "headings.strip_sub_numbering"):
            # "3.2 Association ..." -> "<H2>Association ...": the number goes with the tag
            out.append(_change(p, 0, lead + f.num.end(), f"<{level}>", Layer.STRUCTURE, "subheadings"))
        else:
            out.append(_change(p, 0, lead, f"<{level}>", Layer.STRUCTURE, "subheadings"))
        first_at = lead + (f.num.end() if f.num else 0)
        if p.text[first_at:first_at + 1].islower() and level != "H1":
            out.append(_change(p, first_at, first_at + 1, p.text[first_at].upper(), Layer.STRUCTURE, "subheadings"))
        if level == "H1":
            core = p.text[lead:f.end]
            cased = _cased(core, head_case) if head_case != "keep" else core
            if cased != core:  # the first level's case (IJO S: "H1 | Title Caps"), not a sub-heading's
                out.append(_change(p, lead, f.end, cased, Layer.STRUCTURE, "subheadings"))
        elif style.get("headings.subheading_case", "keep") == "sentence":
            out += sentence_case_changes(ms, p, "subheadings",
                                         hyphen_parts=style.get("headings.sentence_case_hyphen_parts", False),
                                         end=f.end, brackets=style.get("headings.sentence_case_brackets", False),
                                         start=first_at if f.num and f.num.re is _SUB_LETTER else 0)
        if f.brk is not None:  # the text after the heading line (or run-in label) becomes its own paragraph
            out.append(_change(p, f.brk[0], f.brk[1], BR, Layer.STRUCTURE, "subheadings"))
            continue
        if style.get("headings.strip_trailing_colon", True) and p.text.rstrip().endswith(":") or f.strip_period:
            k = len(p.text.rstrip()) - 1  # with the spaces after it
            out.append(_change(p, k, len(p.text), "", Layer.STRUCTURE, "subheadings"))
        elif p.text != p.text.rstrip() and lead < len(p.text.rstrip()):  # the colon stays (G137), the spaces go
            out.append(_change(p, len(p.text.rstrip()), len(p.text), "", Layer.STRUCTURE, "subheadings"))
    return out


def sentence_case_changes(ms: Manuscript, p: Paragraph, rule: str, start: int = 0,
                          hyphen_parts: bool = False, end: int | None = None, brackets: bool = False) -> list[Change]:
    """A Title Case heading in sentence case, word by word ("Statistical Analysis" -> "Statistical analysis").
    The first word and the word after a colon keep their capital; a word is lowered unless the manuscript's
    running text capitalises it inside a sentence more often than not (so "Iran", "Bandura" stay); acronyms and
    mixed-case words ("RTIs", "NVivo") are left alone."""
    body = getattr(ms, "_sentence_body", None)
    if body is None:
        body = ms._sentence_body = _body_text(ms, set())
    out = []
    first = True
    words = [m for m in re.finditer(r"[^\s]+", p.text) if end is None or m.start() < end]
    # An expansion before its abbreviation keeps its capitals: "Borderline Personality Inventory (BPI)".
    keep = set()
    for k, m in enumerate(words):
        a = re.match(r"\(([A-Z][A-Za-z]*[A-Z])s?\)", m.group(0))
        if a and k >= len(a.group(1)):
            before = words[k - len(a.group(1)):k]
            if all(w.group(0)[:1].upper() == c.upper() for w, c in zip(before, a.group(1))):
                keep.update(w.start() for w in before)
    lead_word = next((m.start() for m in words if m.start() >= start), None) if start else None
    # Brackets whose capitalised words the text all writes in lower case ("(In Vitro)", "(Dermal Model)") are lowered;
    # any other ("(Wilcoxon Signed-Rank Test)", "(Heine, Germany)") keeps its capitals.
    low_brackets = []
    for g in re.finditer(r"\(([^()]*)\)", p.text) if brackets else ():
        caps_words = [w for w in g.group(1).split() if w[:1].isupper()]
        if caps_words and all(re.fullmatch(r"[A-Z][a-z]+[,.;:]?", w) and _counts(w.rstrip(",.;:"), body)[0]
                              and _counts(w.rstrip(",.;:"), body)[1] <= _counts(w.rstrip(",.;:"), body)[0]
                              for w in caps_words):
            low_brackets.append((g.start(), g.end()))
    for m in words:
        w = m.group(0)
        if m.start() < start or w.startswith("<H") or re.fullmatch(r"[\d.]+", w):
            continue  # before `start`: a caption's "Table 1:" label
        hyphened = re.fullmatch(r"\(?[A-Za-z0-9]+(?:[-‐‑][A-Za-z0-9]+)+\W*", w)
        if hyphen_parts and hyphened and m.start() not in keep and p.text[:m.start()].count("(") <= p.text[:m.start()].count(")"):
            # "Within-Group Comparison", "Cross-Sectional Study": the parts after a hyphen too, unless the text
            # writes them capitalised after a hyphen ("Kaplan-Meier", "Mann-Whitney").
            after = p.text[m.end():].lstrip()
            parts = list(re.finditer(r"[^-‐‑]+", w))
            head0 = re.search(r"[A-Za-z]+", parts[0].group(0))
            # "Kaplan-Meier", "Mann-Whitney": a name (never lower case in the text) keeps its capitals;
            # "Activity Inventory-7 (CAHAI-7)": an expansion before its abbreviation keeps its capitals.
            name = head0 is not None and re.fullmatch(r"[A-Z][a-z]+", head0.group(0)) is not None and \
                _counts(head0.group(0), body)[0] == 0 and not _ordinary_word(head0.group(0))  # never in lower case: a name
            if name or re.match(r"\([A-Z]", after):
                first = w.endswith(":")
                continue
            prefixed = head0 is not None and _ordinary_word(head0.group(0))  # "Inter-Group", not "Anti-Candida"
            for k, part in enumerate(parts):
                text_part = part.group(0)
                if not re.fullmatch(r"\(?[A-Z][a-z]+\W*", text_part) or (k == 0 and first):
                    continue
                head = re.search(r"[A-Za-z]+", text_part).group(0)
                if prefixed and k and _dictionary_lower() and head.lower() not in _dictionary_lower():
                    continue  # the name after a prefix keeps its capital ("Anti-Candida", "Pre-Columbian")
                if k == 0:
                    lower, caps = _counts(head, body)
                else:
                    lower = len(re.findall(r"[-‐‑]" + head.lower() + r"\b", body))
                    caps = len(re.findall(r"[-‐‑]" + head + r"\b", body)) \
                        - len(re.findall(r"[-‐‑]" + head + r"\b", p.text))  # not the heading itself
                if caps <= lower:
                    i = m.start() + part.start() + text_part.index(head[0])
                    out.append(_change(p, i, i + 1, head[0].lower(), Layer.STRUCTURE, rule))
            first = w.endswith(":")
            continue
        if not first and m.start() not in keep and re.match(r"\([A-Z][a-z]+\W*$", w):
            # "(In Vitro)": the bracket's first word
            if any(a <= m.start() < b for a, b in low_brackets):
                out.append(_change(p, m.start() + 1, m.start() + 2, w[1].lower(), Layer.STRUCTURE, rule))
            first = w.endswith(":")
            continue
        if first or m.start() in keep or not re.match(r"[A-Z][a-z]+(?:[’'][a-z]+)?\W*$", w):
            if m.start() == lead_word and w[:1].islower() and re.fullmatch(r"[a-z]+\W*", w):
                out.append(_change(p, m.start(), m.start() + 1, w[0].upper(), Layer.STRUCTURE, rule))  # caption's first word
            first = w.endswith(":")
            continue
        head = re.match(r"[A-Za-z]+", w).group(0)
        lower, caps = _counts(head, body)
        inside = p.text[:m.start()].count("(") > p.text[:m.start()].count(")")  # "(Heine, Germany)": maker, place
        named = head in _proper_nouns() or (head == "United" and re.match(r"\s+(?:States|Kingdom|Arab|Nations)\b", p.text[m.end():]))
        if caps <= lower and not (proper_word(head) and caps) and not named and (not inside or any(a <= m.start() < b for a, b in low_brackets)):
            # never capitalised inside a sentence of the text, or more often lower case; in brackets only when all
            # its words are ("(In Vitro)" -> "(in vitro)", not the maker's "(Heine, Germany)")
            out.append(_change(p, m.start(), m.start() + 1, w[0].lower(), Layer.STRUCTURE, rule))
        first = w.endswith(":")
    return out


def _held_here(style: JournalStyle, reason: str) -> bool:
    """Autopilot leaves changes with this reason for the editor in this journal (autopilot.hold)."""
    return any(reason == h or reason.startswith((h + ":", h + ".", h + " ")) for h in style.get("autopilot.hold") or [])


def rule_abstract(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not _cfg(style, "enabled", True):
        return []
    lay = layout(ms)
    if not lay.abstract:
        return []
    out = []
    heading = _cfg(style, "abstract.heading", "Abstract")
    head = lay.abstract_head
    if head is not None:
        m = _ABSTRACT.match(head.text)
        if head in lay.abstract:  # inline "Abstract: text" -> heading paragraph + text
            if _on(style, "abstract.split_inline_heading"):
                out.append(_change(head, m.start(2), m.end(3), heading + BR, Layer.STRUCTURE, "sections_abstract"))
                if _on(style, "abstract.heading_bold"):  # S: 'The word "Abstract" ... in bold style'
                    out[-1].segment_format = [{"bold": True, "caps": False}, None]
        else:
            if _on(style, "abstract.heading_case") and m.group(2) != heading:
                # "ABSTRACT" -> "Abstract" (S: 'The word "Abstract" should be added as heading')
                out.append(_change(head, m.start(2), m.end(2), heading, Layer.STRUCTURE, "sections_abstract",
                                   {"caps": False, "italic": False}))
            if _on(style, "abstract.strip_heading_colon") and m.group(3).strip():
                out.append(_change(head, m.end(2), m.end(2) + len(m.group(3).rstrip()), "", Layer.STRUCTURE,
                                   "sections_abstract"))
            if _on(style, "abstract.heading_case") and (_has_caps(head) or _any_fmt_span(head, m.start(2), m.end(2),
                                                                                           "italic")):
                # the all-caps attribute shows "ABSTRACT" just the same; S: bold, not italic
                out.append(_change(head, m.start(2), m.end(2), head.text[m.start(2):m.end(2)], Layer.STRUCTURE,
                                   "sections_abstract", {"caps": False, "italic": False}))
    elif _on(style, "abstract.add_heading") and _structured(lay):  # an unstructured one is too unsure
        first = lay.abstract[0]
        i = ms.paragraphs.index(first)
        prev = ms.paragraphs[i - 1] if i else None
        if prev is not None and not prev.text and not prev.in_table:
            out.append(_change(prev, 0, 0, heading, Layer.STRUCTURE, "sections_abstract",
                               {"bold": True, "italic": False}))
        else:
            out.append(_change(first, 0, 0, heading + BR, Layer.STRUCTURE, "sections_abstract"))

    rename = {k.lower(): v for k, v in (_cfg(style, "abstract.rename") or {}).items()}
    sub_case = _cfg(style, "abstract.subheading_case", "keep")
    for p in lay.abstract:
        if p is head:
            continue
        m = _SUB_INLINE.match(p.text) or _SUB_ALONE.match(p.text)
        if not m:
            continue
        label = m.group(2)
        new = label
        reason = "sections_abstract"
        if label.lower() in rename:
            new = _match_case(label, rename[label.lower()])
            # its own reason, so a journal can hold one renaming for the editor ("sections_abstract.rename:methods")
            reason = f"sections_abstract.rename:{rename[label.lower()].lower()}"
            cased = _cased(label, sub_case) if sub_case != "keep" else label
            if cased != label and _held_here(style, reason):
                # The journal holds this renaming for the editor; a held change would also hold the case fix it
                # carries ("MATERIALS AND METHODS" left in capitals). The two cannot be separate changes on one
                # label, so the case fix is the change and the renaming goes with it as a note for the editor.
                new, reason = cased, "sections_abstract"
                note = provenance.make("journal_style", "Stylesheet", f'Sub-heading "{cased}": the stylesheet also '
                                       f'renames it "{rename[label.lower()]}" (held for the editor in this journal)')
                out.append(_change(p, m.start(2), m.end(2), new, Layer.STRUCTURE, reason))
                out[-1].provenance = note
                new = label = cased  # done: the colon check below goes on; no bold change over the case fix
        new = _cased(new, sub_case) if sub_case != "keep" else new
        if new != label:
            # the new sub-heading is bold as the others ("Materials and Methods" -> a bold "Methods")
            out.append(_change(p, m.start(2), m.end(2), new, Layer.STRUCTURE, reason,
                               {"bold": True} if _on(style, "abstract.subheading_bold") else None))
        sep = m.group(3)
        alone = m.re is _SUB_ALONE
        if _on(style, "abstract.subheading_colon"):
            want = ": "
            if alone and sep != want:
                out.append(_change(p, m.end(2), m.end(3), want, Layer.STRUCTURE, "sections_abstract"))
            elif not alone and sep != want and sep.strip() in (":", "-", "–", "—"):
                out.append(_change(p, m.end(2), m.end(3), want, Layer.STRUCTURE, "sections_abstract"))
        if _on(style, "abstract.subheading_bold") and new == label and label == m.group(2):
            out.append(_change(p, m.start(2), m.end(2), label, Layer.STRUCTURE, "sections_abstract", {"bold": True}))
    if _on(style, "abstract.unbold_text"):
        out += _unbold_abstract(lay, style)
    return out


def _has_caps(p: Paragraph) -> bool:
    """A run of the paragraph has the all-caps attribute (<w:caps/>)."""
    for el in p.element.iter(qn("w:caps")):
        if (el.get(qn("w:val")) or "true").lower() not in ("0", "false", "off", "none"):
            return True
    return False


def _unbold_abstract(lay: Layout, style: JournalStyle) -> list[Change]:
    """S: "Abstract text should not be in bold": bold words in the abstract's text go roman (a structured abstract's
    sub-heading labels and the heading stay bold)."""
    from .rules_format import char_formats

    out = []
    for p in lay.abstract:
        fm = char_formats(p)
        if len(fm) != len(p.text):
            continue
        start = 0
        m = _SUB_INLINE.match(p.text) or _SUB_ALONE.match(p.text)
        if m:
            start = m.end()
        if p is lay.abstract_head:
            start = max(start, _ABSTRACT.match(p.text).end())
        i = start
        while i < len(p.text):
            if fm[i].get("bold") and not p.text[i].isspace():
                j = i
                while j < len(p.text) and (fm[j].get("bold") or p.text[j].isspace()):
                    j += 1
                e = len(p.text[:j].rstrip())
                if any(ch.isalnum() for ch in p.text[i:e]):
                    out.append(_change(p, i, e, p.text[i:e], Layer.STRUCTURE, "sections_abstract", {"bold": False}))
                i = j
            else:
                i += 1
    return out


# --- key words --------------------------------------------------------------------------

def _body_text(ms: Manuscript, skip: set[str]) -> str:
    """Running text of the manuscript: no titles or headings (Title Case), no reference list."""
    skip = skip | {r.para_id for r in references.extract(ms)}
    return "\n".join(p.text for p in ms.paragraphs if p.id not in skip and _words(p.text) >= 12)


def _counts(word: str, body: str) -> tuple[int, int]:
    """(lower-case uses, capitalised uses inside a sentence) of a word in the manuscript.
    Capitals inside a capitalised phrase ("Complete Androgen Insensitivity Syndrome (CAIS)")
    are the author's emphasis, not a proper noun, and do not count."""
    low = word[0].lower() + word[1:]
    cap = word[0].upper() + word[1:]
    lower = len(re.findall(r"\b" + re.escape(low) + r"\b", body))
    caps = 0
    for m in re.finditer(r"(\S*)( )" + re.escape(cap) + r"\b(?![\s\-‑]+[A-Z])", body):
        prev = m.group(1)
        # "Shiraz, Iran": after a comma the capital is the word's own, not a capitalised phrase's
        if prev and (not prev[0].isupper() or prev[-1] == ",") and prev[-1] not in ".!?:":
            caps += 1
    # An eponym written with a hyphen inside a sentence ("... of Miller-Fisher syndrome", "Turner-Kieser syndrome"):
    # a proper noun for each of its names, however the key word joins them ("Miller Fisher syndrome").
    for m in re.finditer(r"(\S*) ((?:[A-Z][a-zà-ÿ]+[-‐‑])+[A-Z][a-zà-ÿ]+)\b(?![\s\-‑]+[A-Z])", body):
        prev = m.group(1)
        if cap in re.split(r"[-‐‑]", m.group(2)) and prev and not prev[0].isupper() and prev[-1] not in ".!?:":
            caps += 1
    return lower, caps


# Prefixes written with a hyphen ("Inter-Group", "Post-Intervention"): never a name.
_HYPHEN_PREFIXES = frozenset("""inter intra extra post pre non cross multi self co sub semi pseudo meta para peri
    trans ultra over under well follow long short high low mid mini micro macro neo re de un""".split())


def _proper_nouns() -> frozenset[str]:
    from .refengine import PROPER_NOUNS

    return PROPER_NOUNS


def _dictionary_lower() -> set[str]:
    from .refengine import _dictionary_words

    return _dictionary_words()[0]


def _ordinary_word(word: str) -> bool:
    """A word that is not a name: a hyphen prefix ("Inter-", "Intra-")."""
    return word.lower() in _HYPHEN_PREFIXES


def proper_word(word: str) -> bool:
    """A word that is a name wherever it stands: a place or person the reference rules know (refengine.PROPER_NOUNS:
    "Iran", "Yazd") or one the dictionary lists only capitalised ("Doppler", "Bowen")."""
    from .refengine import PROPER_NOUNS, _dictionary_words

    base = re.sub(r"[’']s?$", "", word)
    return base in PROPER_NOUNS or base in _dictionary_words()[1]


def _keyword_case(kw: str, body: str) -> str:
    """Lower-case a keyword's words unless the manuscript writes them capitalised (proper nouns:
    "Goldenhar syndrome", "Iran"); capitalise lower-case words the manuscript capitalises."""
    out = []
    for w in kw.split(" "):
        # Each part of a hyphenated word too: "Game-Based Learning" -> "game-based learning",
        # "K-Means" -> "K-means" (not "Non-ST", "COVID-19"; the part after an apostrophe stays as it is).
        parts = re.split(r"([-‐‑/])", w)
        for i in range(0, len(parts), 2):
            head = re.split(r"[’']", parts[i])[0].rstrip(",:;")  # "Mellitus," in "Diabetes Mellitus, Type 2"
            if len(head) > 1 and head.isalpha() and head[1:].islower():
                lower, caps = _counts(head, body)
                if head[0].isupper() and caps <= lower and not proper_word(head):
                    parts[i] = parts[i][0].lower() + parts[i][1:]
                elif head[0].islower() and i == 0 and (caps >= 2 and not lower or proper_word(head.capitalize())
                                                       and not lower):
                    parts[i] = parts[i][0].upper() + parts[i][1:]
        out.append("".join(parts))
    return " ".join(out)


def keyword_list(text: str) -> tuple[int, list[str]] | None:
    """(offset where the list starts, keywords) for a key words paragraph."""
    m = _KW.match(text)
    if not m:
        return None
    return m.end(), _split_keywords(text[m.end():])


def _split_keywords(rest: str) -> list[str]:
    if ";" in rest:
        # a list delimited by semicolons keeps the commas of inverted MeSH terms: "Diabetes Mellitus, Type 2",
        # "Behavior, Addictive" are one key word each (one comma, a capital after it); "Child, pediatrics" is two
        parts = [k.strip() for k in re.split(r";(?![^()]*\))", rest) if k.strip()]
        if any("," in k for k in parts) and all(k.count(",") <= 1 and re.search(r",\s*[A-Z0-9]", k) for k in parts
                                                 if "," in k):
            return parts
    parts = re.split(r"[;,](?![^()]*\))", rest)  # not inside "(CPFE; IPF)"
    return [k.strip() for k in parts if k.strip()]


def _list_paragraph(ms: Manuscript, kw: Paragraph) -> Paragraph | None:
    """The paragraph after a bare "Key words" line, when it holds the list."""
    i = ms.paragraphs.index(kw)
    nxt = next((q for q in ms.paragraphs[i + 1:i + 3] if q.text.strip()), None)
    if nxt is None or _words(nxt.text) > 40 or not re.search(r"[;,]", nxt.text) or _SUB_INLINE.match(nxt.text):
        return None
    return nxt


def listed_keywords(ms: Manuscript) -> list[str]:
    lay = layout(ms)
    kw = lay.keywords
    if kw is None:
        return []
    _, words = keyword_list(kw.text[lay.keywords_at:])
    if not words and (nxt := _list_paragraph(ms, kw)) is not None:
        words = _split_keywords(nxt.text)
    return words


def xml_keywords(ms: Manuscript) -> list[str]:
    """The key words given in the submission XML (<keywords>), when there is one."""
    meta = getattr(getattr(ms, "front_sources", None), "meta", None)
    return _split_keywords(getattr(meta, "keywords", "") or "")


def _styled_keywords(kws: list[str], ms: Manuscript, style: JournalStyle, skip: set[str]) -> list[str]:
    """The key words as the journal writes them: no final period, no repeats, case, order."""
    kws = list(kws)
    if _on(style, "keywords.drop_final_period") and kws and kws[-1].endswith(".") \
            and not re.search(r"\b[A-Za-z]\.$", kws[-1]):
        kws[-1] = kws[-1][:-1].rstrip()
    if _on(style, "keywords.dedupe"):
        seen, uniq = set(), []
        for k in kws:
            if k.lower() not in seen:
                seen.add(k.lower())
                uniq.append(k)
        kws = uniq
    if _on(style, "keywords.case"):
        body = _body_text(ms, skip)
        kws = [_keyword_case(k, body) for k in kws]
    if style.get("keywords.alphabetical") and _on(style, "keywords.sort"):
        kws = sorted(kws, key=lambda k: k.lower())
    if _on(style, "keywords.case") and kws and kws[0][:1].islower():
        first = kws[0].split(" ")[0]
        if first.lower() == first:  # not "mRNA", "pH"
            kws[0] = kws[0][0].upper() + kws[0][1:]
    return kws


def _keywords_line(last: Paragraph, label: str, words: str, style: JournalStyle) -> Change:
    """A new "Keywords: ..." paragraph after `last`: the label bold (S: 'Heading should be "Keywords:" followed by a
    colon, bold'), the list roman."""
    c = _change(last, len(last.text), len(last.text), BR + label + " " + words, Layer.STRUCTURE, "sections_keywords",
                {"italic": False, "bold": False})
    if _on(style, "keywords.label_bold"):
        c.segment_format = [None, [[len(label), {"bold": True}]]]
    return c


def _keywords_inline(ms: Manuscript, style: JournalStyle, p: Paragraph, at: int, label: str, delim: str) -> list[Change]:
    """The key words as the last line of an abstract paragraph ("... addiction.\nKeywords: Emotion Regulation, ..."):
    the line becomes its own paragraph (S: "Should appear in a separate line after the Abstract"), styled as usual."""
    text = p.text
    m = _KW.match(text[at:])
    end_label = at + m.end()
    rest = text[end_label:]
    b0 = len(text[:at].rstrip(" \t\n"))
    out = [_change(p, b0, at, BR, Layer.STRUCTURE, "sections_keywords")]
    have = text[at + m.start(2):end_label]
    want = label + " "
    label_bold = _on(style, "keywords.label_bold")
    if _on(style, "keywords.fix_heading") and have != want:
        out.append(_change(p, at + m.start(2), end_label, want, Layer.STRUCTURE, "sections_keywords",
                           {"bold": True} if label_bold else None))
    kws = _styled_keywords(_split_keywords(rest), ms, style, {p.id})
    new_rest = _keyword_delimiter(kws, delim).join(kws)
    lead = len(rest) - len(rest.lstrip())
    old_rest = rest.strip()
    if old_rest and new_rest != old_rest:
        out.append(_change(p, end_label + lead, end_label + lead + len(old_rest), new_rest, Layer.STRUCTURE,
                           "sections_keywords"))
    return out


def _keyword_delimiter(kws: list[str], delim: str) -> str:
    """"; " between key words that hold a comma themselves (MeSH "Diabetes Mellitus, Type 2"), as the editors keep
    them; the journal's delimiter otherwise."""
    return "; " if any("," in k for k in kws) else delim


def rule_keywords(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not _cfg(style, "enabled", True):
        return []
    lay = layout(ms)
    out = []
    label = _cfg(style, "keywords.heading", "Keywords:")
    delim = _cfg(style, "keywords.delimiter", ", ")
    from_xml = xml_keywords(ms) if _on(style, "keywords.from_xml") else []
    xml_prov = provenance.make("xml", "Submission XML", "<keywords>: " + ", ".join(from_xml)) if from_xml else None
    p = lay.keywords
    if p is None:
        if lay.abstract and _on(style, "keywords.add_missing"):
            # None in the manuscript: the submission's own (editors: "Followed XML file for Keywords"), else the
            # label alone and a query.
            words = delim.join(_styled_keywords(from_xml, ms, style, set())) if from_xml else ""
            out.append(_keywords_line(lay.abstract[-1], label, words, style))
            out[-1].provenance = xml_prov
        return out
    if lay.keywords_at:
        return _keywords_inline(ms, style, p, lay.keywords_at, label, delim)
    text = p.text
    m = _KW.match(text)
    end_label = m.end()
    have = text[m.start(2):end_label]
    rest = text[end_label:]
    label_bold = _on(style, "keywords.label_bold")
    # S: "Should appear in a separate line after the Abstract": key words above the abstract move under it.
    move = lay.keywords_first and bool(lay.abstract) and _on(style, "keywords.after_abstract")
    list_p, base = p, end_label
    separate = False
    if not rest.strip():
        # "Key words" on its own line, the list in the next paragraph
        nxt = _list_paragraph(ms, p)
        if nxt is not None:
            list_p, base, rest, separate = nxt, 0, nxt.text, True
        elif from_xml:
            rest, list_p, base = "", None, end_label
        else:
            if _on(style, "keywords.fix_heading") and have != label and not move:
                out.append(_change(p, m.start(2), end_label, label, Layer.STRUCTURE, "sections_keywords",
                                   {"bold": True} if label_bold else None))
            return out
    kws = _split_keywords(rest) if list_p is not None else list(from_xml)
    kws = _styled_keywords(kws, ms, style, {p.id} | ({list_p.id} if list_p is not None else set()))
    new_rest = _keyword_delimiter(kws, delim).join(kws)
    if move:
        out.append(_change(p, 0, len(p.text), "", Layer.STRUCTURE, "sections_keywords"))
        if separate:
            out.append(_change(list_p, 0, len(list_p.text), "", Layer.STRUCTURE, "sections_keywords"))
        out.append(_keywords_line(lay.abstract[-1], label, new_rest, style))
        out[-1].provenance = provenance.make("manuscript", "Manuscript", "the key words, moved under the abstract "
                                             "(stylesheet: after the Abstract)", para_id=p.id) if list_p is not None \
            else xml_prov
        return out
    runon = list_p is None or (separate and _on(style, "keywords.run_on"))
    if _on(style, "keywords.fix_heading"):
        want = label + (" " if rest.strip() and not separate else "")
        if have != want:
            out.append(_change(p, m.start(2), end_label, want, Layer.STRUCTURE, "sections_keywords",
                               {"bold": True} if label_bold else None))
        elif label_bold:
            ls, le = m.start(2), m.start(2) + len(have.rstrip())
            if not _all_fmt_span(p, ls, le, "bold") or _any_fmt_span(p, ls, le, "italic"):
                out.append(_change(p, ls, le, text[ls:le], Layer.STRUCTURE, "sections_keywords",
                                   {"bold": True, "italic": False}))
    if runon:
        # S182 "Keywords: Acceptability, Hall’s technique, ...": the list runs on after the label
        out.append(_change(p, len(text), len(text), " " + new_rest, Layer.STRUCTURE, "sections_keywords",
                           {"bold": False, "italic": False}))
        if list_p is None:
            out[-1].provenance = xml_prov
        else:
            out.append(_change(list_p, 0, len(list_p.text), "", Layer.STRUCTURE, "sections_keywords"))
        return out
    lead = len(rest) - len(rest.lstrip())
    trail = len(rest) - len(rest.rstrip())
    old_rest = rest[lead:len(rest) - trail]
    if new_rest != old_rest:
        out.append(_change(list_p, base + lead, base + lead + len(old_rest), new_rest, Layer.STRUCTURE,
                           "sections_keywords"))
    return out


def _any_fmt_span(p: Paragraph, s: int, e: int, prop: str) -> bool:
    from .rules_format import char_formats

    fm = char_formats(p)
    return len(fm) == len(p.text) and any(fm[i].get(prop) for i in range(s, e) if not p.text[i].isspace())


def _all_fmt_span(p: Paragraph, s: int, e: int, prop: str) -> bool:
    from .rules_format import char_formats

    fm = char_formats(p)
    if len(fm) != len(p.text):
        return True
    return all(fm[i].get(prop) for i in range(s, e) if not p.text[i].isspace())


# --- back matter -----------------------------------------------------------------------

_PL = r"(?:\(s\)|s)?"  # "Conflict(s)", "Source(s)", "Disclosure(s)"
_CONFLICT = (r"(?:declaration\s+of\s+)?conflict" + _PL + r"\s+of\s+interest" + _PL + r"(?:\s+statement)?"
             r"(?:\s+and\s+(?:any\s+)?disclosure" + _PL + r")?|"
             r"competing\s+interest" + _PL + r"(?:\s+(?:statement|declaration))?|"
             r"declaration\s+of\s+(?:competing\s+|conflicting\s+)?interest" + _PL + r"|"
             r"(?:financial\s+)?disclosure" + _PL + r"(?:\s+statement)?|"
             r"conflicting\s+interest" + _PL + r"(?:\s+statement)?|"
             r"financial\s+(?:relationship|interest)" + _PL + r"(?!\s+and\s+spons[eo]r)")
_FUNDING = (r"funding(?:\s+(?:information|sources?|statement|support|details))?|"
            r"financial\s+support(?:\s+and\s+sponsorship)?|finance\s+support(?:\s+and\s+sponsorship)?|"
            r"source" + _PL + r"\s+of\s+(?:funding|support)(?:\s*/\s*funding)?|"
            r"sponsorship|grant\s+support|support(?:\s*/\s*funding)?|"
            # "Financial relationship and sponsership: None" (the template's "Financial support and sponsorship")
            r"financial\s+relationships?\s+and\s+spons(?:or|er)ship")
_ACK = r"acknowledge?ments?|acknowledge(?=\s*[:\-–—])"
_LABEL = re.compile(
    r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?(?P<label>(?P<conflict>" + _CONFLICT + r")(?:\s*(?:&|and)\s*(?P<both>"
    + _FUNDING + r"))?|(?P<funding>" + _FUNDING + r")|(?P<ack>" + _ACK + r"))"
    r"(?:\s*\([^)]{0,60}\))?(?P<sep>\s*[:.\-–—]\s*|\s+|$)", re.I)
_NIL = re.compile(r"\b(?:none|nil|no|not|nothing|n/?a|non[- ]?existent)\b", re.I)
_REAL = re.compile(r"\b(?:grant|supported|funded|funding from|award|fellowship|scheme|project no|received (?:a|funding|support)"
                   r"|thank|grateful|gratitude)\b", re.I)
_NEGATION = re.compile(r"\b(?:no|not|nothing|none|neither|nor|without)\b", re.I)


def _sep_ok(sep: str | None) -> bool:
    """A label's separator: punctuation, a tab, or a run of spaces ("Ethical approval    All procedures ...": IJO_1745)."""
    sep = sep or ""
    return bool(sep.strip()) or "\t" in sep or len(sep) >= 2


def _nil(content: str) -> bool:
    """A statement that says nothing: "None", "Nil.", "The authors declare no conflicts of interest",
    "This review has no specific grant from any funding agency" (a negated "grant" is not real support)."""
    c = content.strip()
    if not c:
        return True
    if _words(c) > 30 or not _NIL.search(c):
        return False
    return not any(not _NEGATION.search(c[max(0, m.start() - 60):m.start()]) for m in _REAL.finditer(c))


def _kind(m) -> str:
    if m.group("both"):
        return "both"
    return "conflict" if m.group("conflict") else ("funding" if m.group("funding") else "ack")


_CAPTION = re.compile(r"^\s*(?:<[A-Z]+\d*>)?\s*(?:table|figure|fig\.?|graph|chart|supplementary)\s*[S]?\d", re.I)

# A sentence with no label that only says there is nothing to declare ("No conflict of interest.",
# "The authors have nothing to disclose.", "This work did not receive any funding ..."): editors delete it.
_NIL_SENTENCE = re.compile(r"^\s*(?:the\s+authors?(?:\(s\))?|all\s+(?:the\s+)?authors|authors|we|none\s+of\s+the\s+authors|no"
                           r"|there\s+(?:is|are|was|were)\s+no|this\s+(?:work|study|research|article|review|manuscript"
                           r"|systematic\s+review|case(?:\s+report)?))\b", re.I)
_NIL_TOPIC = re.compile(r"conflict|competing|\binterests?\b|funding|\bfunds?\b|financial|\bgrant|sponsor|disclos|relationship", re.I)

# "Declarations" / "Statements and declarations": a container heading over the author statements; editors
# delete it (13 of 14 in the training files) as the statements become the journal's own sections.
_DECLARATIONS = re.compile(r"^\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?(?:(?:ethical\s+|authors?[’'`]?s?[’'`]?\s+)?declarations|statements?\s+and\s+declarations"
                           r"|declarations?\s+and\s+statements?|compliance\s+with\s+ethical\s+standards)\s*:?\s*$", re.I)
# Ethics / consent statements that say nothing ("Ethics: Not required for single case report.", "Consent for
# publication: Not applicable"): editors delete them (7 of 7).
_ETHICS = re.compile(
    r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?(?P<label>(?:ethic(?:s|al)|ethics\s+committee)(?:\s+(?:approval|approve|statement|clearance"
    r"|declaration|committee\s+approval|considerations?))?(?:\s+and\s+(?:informed\s+)?consent(?:\s+to\s+participate)?(?:\s+and\s+publication)?)?"
    r"|consent\s+(?:for|to)\s+(?:publication|publish|participate)(?:\s+and\s+(?:publication|publish))?"
    r"|institutional\s+review\s+board(?:\s+statement)?"
    r"|clinical\s+trial(?:\s+(?:number|registration|registry))?|trial\s+registration|study\s+id)(?P<sep>\s*[:.\-–—]\s*|\s*$|"
    r"\t\s*|\s{2,})", re.I)
_NO_DATA = re.compile(r"^\s*(?:no\s+(?:new\s+)?data(?:sets?)?\s+(?:were|was|have\s+been)\s+(?:generated|created|collected)|data\s+"
                      r"sharing\s+(?:is\s+)?not\s+applicable)", re.I)
_BARE_NIL = re.compile(r"^\s*(?:not\s+applicable|n/?a|nil|none|not\s+required)\s*[.;]?\s*$", re.I)
_ETHICS_NIL = re.compile(r"^\s*(?:not\s+(?:applicable|required|needed|necessary|sought)|n/?a\b|none\b|nil\b|exempt|waived)", re.I)

# Consent: a heading ("Declaration of patient consent", "Patient Consent:", "Consent for publication") or a
# statement without one ("Written informed consent was obtained from the patient for the publication ...").
_CONSENT_HEAD = re.compile(
    r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?(?P<label>(?:authors?[’'`]?s?[’'`]?\s+)?declaration(?:\s+statement)?\s+of\s+"
    r"(?:the\s+)?(?:patients?[’'`]?s?[’'`]?\s+)?consent|patients?[’'`]?s?[’'`]?\s+consent(?:\s+(?:statement|form|for\s+publication))?"
    r"|participants?[’'`]?s?[’'`]?\s+consent|informed\s+consent(?:\s+statement)?)(?P<sep>\s*[:.\-–—]\s*|\s*$|\t\s*|"
    r"\s{2,})", re.I)
_BARE_CONSENT = re.compile(r"^\s*(?:<H\d>)?\s*(?:\d+\.?\s*)?(?:authors?[’'`]?s?[’'`]?\s+)?(?:consent|declarations?)"
                           r"(?:\s+statements?)?\s*[:.\-–—]?\s*$", re.I)  # "Author’s Declaration statement-" (IJO_1432)
_CONSENT_GIVEN = re.compile(r"\bconsent\b.{0,200}\b(?:obtained|given|taken|signed|provided|received)\b|"
                            r"\b(?:obtained|given|gave|gives|granted|taken|signed|provided|certify)\b.{0,120}\bconsent\b", re.I)
_CONSENT_PUBLICATION = re.compile(r"\b(?:publication|publish\w*|images?|photo\w*|clinical\s+(?:information|details)"
                                  r"|case\s+report|this\s+case|reported\s+in\s+the\s+journal)\b", re.I)

# Authors' contributions: a heading ("Authors' contributions", "Author Contributions (CRediT)") or an inline label.
_CONTRIB = re.compile(
    r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?(?P<label>contribution\s+list(?=\s*:)|contributorship(?:\s+statement)?|(?P<author>(?:authors?[’'`]?s?[’'`]?|authorship|individual)\s+)?"
    r"contributions?(?:\s+of\s+(?:the\s+)?authors?)?(?:\s+(?:statement|details))?(?:\s*\(credit(?:\s+statement)?\))?"
    r"|credit\s+(?:authorship\s+)?(?:contribution\s+)?statement)(?P<sep>\s*[:.\-–—]\s*|\s+|$)", re.I)
_CONTRIB_WORDS = re.compile(
    r"\b(?:concept\w*|design\w*|intellectual|literature|data|manuscript|writing|wrote|draft\w*|review\w*|edit\w*|analy[sz]\w*"
    r"|statistic\w*|supervis\w*|investigation|methodology|clinical|surg\w*|acquisition|interpretation|curation|visuali[sz]\w*"
    r"|validation|resources|software|approv\w*|responsible|performed|collect\w*|critical\w*|revis\w*|stud(?:y|ies)|patients?)\b",
    re.I)
_UP = "A-ZÀ-ÖØ-Þİ"
# "Prof. Dr. Asha K. R. – Concept, design, ..."; "A.B.: Conceptualization, ..."; "Ravi Kumar (RK): ..."
_CONTRIB_LINE = re.compile(rf"^\s*(?:[•➢▪*]\s*)?(?:(?:prof|dr|mr|mrs|ms)\.?\s+)*(?:[{_UP}][\w.’'\-]*\.?\s*){{1,6}}"
                           rf"(?:\([{_UP}. ]{{1,10}}\)\s*)?[:–—\-]\s+\S")
# "Ravider kumar saran – Histopthological diagnosis", "Kratee- Histopthological ...": a name (any case) and what they did
_NAME_ROLE = re.compile(r"^\s*(?:[•➢▪*]\s*)?[^\W\d_][\w.’']*(?:\s+[^\W\d_][\w.’']*){0,3}\s*[–—\-:]+\s*[^\W\d_]")
_TEMPLATE = re.compile(r"\[\s*[xX✓✔]?\s*\]|[□☐✓✔]|^author\s+name|^contributor\s*\d|write the first letter|^role\b|tick", re.I)
_OTHER_LABEL = re.compile(r"^\s*(?:corresponding|address|e-?mail|word\s+count|no\.?\s+of|number\s+of|thank|sincerely|yours|"
                          r"regards|declaration|statement\s+on|data\s+(?:availability|sharing|supporting)|availability|key\s*words?|"
                          r"running\s+title|orcid|ethic|consent|funding|conflict|acknowledg|financial|source|competing|"
                          r"meeting|presentation|copyright|guarantor|clinical\s+trial|trial\s+regist|generative|"
                          r"use\s+of\s+(?:ai|artificial)|ai\s+(?:statement|policy|declaration|usage|use)|"
                          r"we\s+(?:hope|look)|the\s+manuscript\s+(?:has|is)\s+not)",
                          re.I)
_PHOTO_LEGEND = re.compile(r"clinical\s+(?:photo(?:graph)?s?|pictures?|images?)|(?:facial|external|full[- ]face)\s+"
                           r"(?:photo(?:graph)?s?|appearance|view)|photo(?:graph)?s?\s+of\s+the\s+(?:patient|face|child|baby|infant)"
                           r"|patient[’']?s\s+face", re.I)
_LEGEND = re.compile(r"^\s*(?:<[A-Z]+\d*>)?\s*(?:figure|fig\.?)\s*\d", re.I)
_CASE_TYPE = re.compile(r"\bcase\s+(?:reports?|series|snippets?|history)\b|clinicopathologic", re.I)
_IMAGE_TYPE = re.compile(r"\bimages?\b|neuroimage|photo\s*(?:essay|quiz)", re.I)
_LETTER_TYPE = re.compile(r"\bletters?\b|correspondence", re.I)


_LEGENDS_HEAD = re.compile(r"^\s*(?:<H\d>)?\s*(?:figure\s+legends?|legends?(?:\s+(?:to|for|of)\s+(?:the\s+)?figures?)?"
                           r"|figure\s+captions?|captions?)\s*:?\s*$", re.I)


def _back_region(ms: Manuscript) -> tuple[list[Paragraph], Paragraph | None]:
    """Paragraphs that may hold back-matter statements (from the discussion on, including any
    after the reference list), and the paragraph the missing sections are added after: the last
    text paragraph before the references, ahead of any table or figure captions."""
    ref_i = references.find_heading(ms)
    end = ref_i if ref_i is not None else len(ms.paragraphs)
    paras = ms.paragraphs[:end]
    start = 0
    for i, p in enumerate(paras):
        if _key(p.text) in ("discussion", "conclusion", "conclusions") and len(p.text) < 40:
            start = i + 1
    if not start:
        start = len(paras) * 2 // 3
    # A label statement read as the last entries of a reference list ("Conflicts of interests: None" after
    # the references) is back matter, not a reference.
    refs = references.extract(ms)
    ref_ids = {r.para_id for r in refs
               if not (_LABEL.match(r.text) and _words(r.text) <= 12 and _LABEL.match(r.text).group("sep").strip())}
    # So is a consent section read as the list's last entries ("Declaration of Patient Consent" and its statement).
    for k, r in enumerate(refs):
        if (_CONSENT_HEAD.match(r.text) or _BARE_CONSENT.match(r.text)) and _words(r.text) <= 12:
            ref_ids.discard(r.para_id)
            if k + 1 < len(refs) and _CONSENT_GIVEN.search(refs[k + 1].text):
                ref_ids.discard(refs[k + 1].para_id)
    region = [p for p in ms.paragraphs[start:] if not p.in_table and p.id not in ref_ids
              and (ref_i is None or p is not ms.paragraphs[ref_i])]
    anchor = None
    if ref_i is not None:
        for p in paras[start:] if start < len(paras) else []:
            if p.in_table or _CAPTION.match(p.text) or _LEGENDS_HEAD.match(p.text):
                break
            if re.search(r"\w", p.text):
                anchor = p
        if anchor is None:
            anchor = next((p for p in reversed(paras) if p.text.strip() and not p.in_table
                           and not _CAPTION.match(p.text)), None)
    return region, anchor


def _said_already(ms: Manuscript, text: str) -> bool:
    """The statement (read from the title page) stands in the manuscript already, word for word."""
    want = " ".join(re.findall(r"\w+", text.lower()))
    if len(want) < 20:
        return False
    return any(want in " ".join(re.findall(r"\w+", p.text.lower())) for p in ms.paragraphs)


def _names_an_author(text: str, ms: Manuscript) -> bool:
    """The head of a "Name: role, role" line names one of the submission's authors (a name word, or their initials:
    "SVK", "A.B."), so "Future Research: 1] Multi-centric ..." in the Discussion is not a contributions list
    (jehp_1683). Without an author list (no submission XML, no title page) any head counts."""
    names = []
    src = getattr(ms, "front_sources", None)
    for a in getattr(getattr(src, "meta", None), "authors", None) or []:
        names.append((a.get("given", ""), a.get("surname", "")))
    plan = getattr(ms, "_front_plan", None)
    for a in getattr(getattr(plan, "fields", None), "authors", None) or []:
        names.append((a.given, a.surname))
    if not names:
        return True
    m = re.match(r"^\s*(?:[•➢▪*]\s*)?(?:(?:prof|dr|mr|mrs|ms)\.?\s+)*(?P<head>[^:–—\-]{1,80}?)\s*[:–—\-]", text, re.I)
    if not m:
        return False
    head = m.group("head")
    words = {w.casefold() for g, sn in names for w in re.findall(r"[^\W\d_]{2,}", f"{g} {sn}")}
    if any(w.casefold() in words for w in re.findall(r"[^\W\d_]{2,}", head)):
        return True
    inits = set()
    for g, sn in names:
        parts = re.findall(r"[^\W\d_]+", f"{g} {sn}")
        if parts:
            inits.add("".join(w[0] for w in parts).upper())
            inits.add("".join(w[0] for w in parts[-1:] + parts[:-1]).upper())
    key = re.sub(r"[\s.()]", "", head).upper()
    return bool(key) and key.isalpha() and key in inits


def article_type(ms: Manuscript) -> str:
    """The article type from the submission XML ("Case Report", "Letter to Editor"), when there is one."""
    src = getattr(ms, "front_sources", None)
    return (getattr(getattr(src, "meta", None), "article_type", "") or "").strip()


def _is_case_report(ms: Manuscript) -> bool:
    if _CASE_TYPE.search(article_type(ms)):
        return True
    return any(_key(p.text) in ("case report", "case reports", "case presentation", "case description",
                                "case history", "case summary")
               for p in ms.paragraphs if len(p.text) < 40 and not p.in_table)  # not a table's "Case report" study design


def has_consent_statement(ms: Manuscript) -> bool:
    return any(re.search(r"\bconsent\b", p.text, re.I) for p in ms.paragraphs)


def identifiable_patient_images(ms: Manuscript) -> bool:
    """Does the article print images of a patient (the stylesheets' "images with any identifiable patient
    information")? Decided from the article type and the figure legends only:
    an image article (Images in Neurology, Ophthalmic Images, ...) with figures, or a case report / case series /
    letter whose legends show clinical photographs ("Clinical photograph showing ...", "facial appearance")."""
    legends = [p.text for p in ms.paragraphs if _LEGEND.match(p.text)]
    if not legends:
        return False
    kind = article_type(ms)
    if _IMAGE_TYPE.search(kind):
        return True
    patient_article = bool(_CASE_TYPE.search(kind) or _LETTER_TYPE.search(kind)) or (not kind and _is_case_report(ms))
    return patient_article and any(_PHOTO_LEGEND.search(t) for t in legends)


def _contribution_items(paras: list[str]) -> list[str]:
    """The contribution statements after an "Author contributions" label in a title page, contribution
    form or cover letter: the label's own text and the lines after it that name what authors did. A list of
    names, a tick-box table or a form template gives nothing (nothing is invented)."""
    for i, text in enumerate(paras):
        m = _CONTRIB.match(text)
        if not m or (m.group("sep").isspace() and not m.group("author")) or len(text) > 2000:
            continue
        items = [t.strip() for t in re.split(r"\s*\n\s*", text[m.end():]) if t.strip()]
        if items and not all(_CONTRIB_WORDS.search(t) and not _TEMPLATE.search(t) for t in items):
            continue
        for nxt in paras[i + 1:i + 30]:
            lines = [t.strip() for t in re.split(r"\s*\n\s*", nxt) if t.strip()]
            if not lines or any(len(t) > (2000 if len(_CONTRIB_WORDS.findall(t)) >= 6 else 700) or _TEMPLATE.search(t)
                                or _OTHER_LABEL.match(t) or _CONTRIB.match(t) or _CONSENT_HEAD.match(t)
                                or _CONSENT_GIVEN.search(t) and not re.search(r"\b[Aa]uthors?\b|\b[A-Z]{2,3}\b", t)
                                or not (_CONTRIB_WORDS.search(t) or _CONTRIB_LINE.match(t) or _NAME_ROLE.match(t))
                                for t in lines):
                break
            items += lines
        if items:
            return items
    return []


def contributions_from_sources(ms: Manuscript, skip: set[str] = frozenset()) -> list[str]:
    """Authors' contributions given with the submission: the title-page file, a contribution form, or a cover
    letter inside the manuscript (paragraphs in `skip`, the back matter, are not read)."""
    return contributions_with_origin(ms, skip)[0]


def contributions_with_origin(ms: Manuscript, skip: set[str] = frozenset()) -> tuple[list[str], dict | None]:
    """`contributions_from_sources` and where the statements were read (a provenance dict)."""
    src = getattr(ms, "front_sources", None)
    forms = [list(f) for f in getattr(src, "contribution_forms", None) or []]
    form_names = list(getattr(src, "contribution_form_names", None) or [])
    paras = [p for p in ms.paragraphs if p.id not in skip and not p.in_table]
    pools = [(list(getattr(src, "first_page", None) or []), ("title_page", "Title page", getattr(src, "first_page_name", "")))]
    pools += [(f, ("contribution_form", "Contribution form", form_names[i] if i < len(form_names) else ""))
              for i, f in enumerate(forms)]
    pools.append(([p.text for p in paras], ("manuscript", "Manuscript", "")))
    for pool, (kind, label, name) in pools:
        items = _contribution_items(pool)
        if items:
            para = None
            if kind == "manuscript":
                para = next((p for p in paras if items[0] in p.text), None)
                letter = set(submission.cover_letter_paragraphs(ms, getattr(ms, "_submission_style", None)))
                at = paras.index(para) if para is not None else -1
                if para is not None and (para.text in letter or any(
                        re.match(r"^\s*cover(?:ing)?\s+letter\b", q.text, re.I) for q in paras[max(0, at - 10):at])):
                    kind, label = "cover_letter", "Cover letter"
            return items, provenance.make(kind, label, items[0], file=name, para_id=para.id if para else "")
    return [], None


# A closing line of a contribution list ("All authors read and approved the final manuscript.").
_APPROVAL = re.compile(r"^\s*(?:all\s+(?:the\s+)?|the\s+|each\s+)?authors?\b.{0,200}\b(?:read|approv\w*|agree\w*|"
                       r"accountab\w*|contribut\w*)\b", re.I)


def top_contributions(ms: Manuscript, skip: set[str] = frozenset()) -> tuple[list[Paragraph], list[str]]:
    """Authors' contribution statements on the manuscript's own title page (before the abstract / main text): the
    paragraphs (an "Authors' contributions" label or heading, the "Name: role, role" lines, a closing "All authors
    read and approved ..." line right after them) and the statements they hold, one per line. Only when the
    leading paragraphs end at the abstract / main text; the first one (the title) and those in `skip` (the back
    matter) are never statements. Nothing when no statement is found."""
    from . import frontmatter

    lead = frontmatter.leading_paragraphs(ms)
    if not lead:
        return [], []
    after = ms.paragraphs[ms.paragraphs.index(lead[-1]) + 1:]
    nxt = next((p for p in after if p.text.strip() and not p.in_table), None)
    if nxt is None or not (frontmatter._BODY_START.match(nxt.text) or frontmatter._LEAD_STOP.match(nxt.text)):
        return [], []  # no title page ahead of the abstract / main text (a short piece read from the top)
    ids = frontmatter.contribution_paragraphs([p for p in lead[1:] if p.id not in skip])
    paras = [p for p in lead if p.id in ids]
    if not paras:
        return [], []
    for q in lead[lead.index(paras[-1]) + 1:]:
        if not (_APPROVAL.match(q.text) and len(q.text) <= 400 and _CONTRIB_WORDS.search(q.text)):
            break
        paras.append(q)
    items = []
    for p in paras:
        text = p.text
        m = _CONTRIB.match(text)
        if m and not (m.group("sep").isspace() and not m.group("author")):
            text = text[m.end():]
        items += [t.strip() for t in re.split(r"\s*\n\s*", text) if t.strip()]
    if not any(frontmatter.is_contribution_line(t) or _CONTRIB_WORDS.search(t) and not _APPROVAL.match(t)
               for t in items):
        return [], []  # a heading or an approval line alone: nothing to move
    return paras, items


def _plain(text: str) -> str:
    return re.sub(r"\W+", " ", text).strip().lower()


# Data availability statement: "Data availability statement: ...","Availability of data and materials", "Data
# sharing statement" (a heading line, or a label with the statement after it).
_DATA_HEAD = re.compile(
    r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?(?P<label>data\s+(?:availability|sharing)(?:\s+statements?)?"
    r"|availability\s+of\s+(?:the\s+)?data(?:\s+and\s+materials?)?)(?P<sep>\s*[:.\-–—]\s*|\s*$)", re.I)
_DATA_TEMPLATE = re.compile(r"\b(?:if\s+applicable|please\s+(?:state|provide|mention|specify)|e\.g\.|tick)\b", re.I)


def _statement(text: str) -> str:
    """"not applicable" -> "Not applicable.": first letter capital, a final period."""
    text = re.sub(r"\s+", " ", text).strip(" -–—:")
    if not text:
        return ""
    text = text[0].upper() + text[1:]
    return text if text.endswith((".", "!", "?")) else text + "."


def data_availability_from_sources(ms: Manuscript) -> str:
    """The data availability statement given on the title page (the first-page file): the label's own text, or the
    paragraph after a bare label. A form's prompt or an empty label gives nothing (nothing is invented)."""
    src = getattr(ms, "front_sources", None)
    paras = [t for t in (getattr(src, "first_page", None) or []) if t.strip()]
    for i, text in enumerate(paras):
        m = _DATA_HEAD.match(text)
        if not m:
            continue
        content = text[m.end():].strip()
        if not content and i + 1 < len(paras):
            nxt = paras[i + 1].strip()
            if not (_OTHER_LABEL.match(nxt) or _CONTRIB.match(nxt) or _LABEL.match(nxt) or _CONSENT_HEAD.match(nxt)
                    or _ETHICS.match(nxt)) and len(nxt) <= 800:
                content = nxt
        if content and len(content) <= 800 and not _DATA_TEMPLATE.search(content):
            return _statement(content)
    return ""


def acknowledgement_from_sources(ms: Manuscript) -> tuple[str, str] | None:
    """(label, text) of the acknowledgement on the title page (first-page file): the label's own text or the paragraph
    after a bare "Acknowledgements:" label. Nothing for a nil one ("None") or a form's prompt."""
    src = getattr(ms, "front_sources", None)
    paras = [t for t in submission.texts_without_parts(list(getattr(src, "first_page", None) or [])) if t.strip()]
    for i, text in enumerate(paras):
        m = _LABEL.match(text)
        if not m or _kind(m) != "ack" or not m.group("sep").strip() and text[m.end():].strip():
            continue
        content = text[m.end():].strip()
        if not content and i + 1 < len(paras):
            nxt = paras[i + 1].strip()
            if not (_LABEL.match(nxt) or _OTHER_LABEL.match(nxt) or _CONTRIB.match(nxt) or _CONSENT_HEAD.match(nxt)
                    or _ETHICS.match(nxt) or _DATA_HEAD.match(nxt)) and len(nxt) <= 1500:
                content = nxt
        if content and len(content) <= 1500 and not _nil(content) and not _DATA_TEMPLATE.search(content):
            return m.group("ack"), _statement(content)
    return None


# A statement on the use of AI tools ("Use of AI tools", "AI policy statement", "Declaration of generative AI in
# scientific writing", "Disclosure: ChatGPT was used ..."): a short label naming AI.
_AI_WORDS = re.compile(r"\b(?:AI|artificial\s+intelligence|generative|chat\s?gpt|LLMs?|large\s+language\s+models?)\b", re.I)
_AI_TEXT = re.compile(r"\b(?:chat\s?gpt|claude|gemini|grammarly|copilot|quillbot|generative\s+AI|AI[- ](?:assisted|based|powered|"
                      r"tools?|writing)|artificial\s+intelligence|large\s+language\s+models?|LLMs?)\b", re.I)


def ai_label(text: str) -> tuple[str, str] | None:
    """(label, statement) of an AI-use statement's label line ("Use of AI tools" / "AI policy statement: ...")."""
    t = re.sub(r"^\s*(?:<H\d>)?\s*(?:\d+(?:\.\d+)*\.?\s*)?", "", text)
    m = re.match(r"(?P<label>[^:.\n]{2,90}?)(?:\s*[:\-–—]\s+|\s*:\s*|\s*$)", t)
    if not m or not _AI_WORDS.search(m.group("label")) or len(m.group("label").split()) > 10 \
            or not re.search(r"(?i)\b(?:use|usage|used|using|policy|statement|declaration|disclosure|tools?|assist\w*|generated|"
                             r"technolog\w*|writing)\b", m.group("label")) \
            or re.search(r"(?i)\b(?:is|are|was|were|has|have|had|can|may)\b", m.group("label")):
        return None
    rest = t[m.end():].strip()
    if not rest and t.rstrip().endswith("."):
        return None
    return m.group("label").strip(), rest


def _section_contrib_label(text: str) -> bool:
    """An "Authors' contributions" section label (not the "Contribution: This author helped ..." under each name)."""
    km = _CONTRIB.match(text)
    return bool(km) and bool(km.group("author") or km.group("label").lower().startswith(
        ("credit", "contribution list", "contributorship")))


def ai_disclosure(text: str) -> bool:
    """"Disclosure: OpenAI ChatGPT was used during revision ...": a statement on AI use, not a conflicts disclosure."""
    return bool(re.match(r"(?i)^\s*(?:<H\d>)?\s*disclosures?\s*(?:statement)?\s*[:\-–—]", text) and _AI_TEXT.search(text))


def statement_label(text: str) -> tuple[str, str, str] | None:
    """(kind, label, statement written after it) of a back-matter statement label: consent, contributions, data,
    ethics, funding, conflict, both, ack, ai. A label without a separator is one only when nothing follows it."""
    for kind, rx in (("consent", _CONSENT_HEAD), ("contributions", _CONTRIB), ("data", _DATA_HEAD), ("ethics", _ETHICS),
                     ("label", _LABEL)):
        m = rx.match(text)
        if not m:
            continue
        rest = text[m.end():].strip()
        if not _sep_ok(m.group("sep")) and rest:
            continue
        k = _kind(m) if kind == "label" else kind
        if k == "ethics" and re.match(r"(?i)consent\s+(?:for|to)\s+publi", m.group("label")):
            k = "consent"
        return k, m.group("label"), rest
    got = ai_label(text)
    return ("ai", got[0], got[1]) if got else None


def declared_statements(ms: Manuscript) -> dict[str, tuple[str, str, dict]]:
    """The author's statements given with the submission outside the manuscript (owner, 2026-10-27, G283: a declaration
    found only on the title page is used in the back matter): kind -> (label, statement, provenance), for funding,
    conflict, ethics, data, ai. Read from the title-page file (without its cover letter), then the submission's
    declaration files ("JEHP_EthicsDeclarations.docx"); the first one of each kind stands. A form's prompt gives
    nothing (nothing is invented)."""
    cached = getattr(ms, "_declared_statements", None)
    if cached is not None:
        return cached
    src = getattr(ms, "front_sources", None)
    pools = [(submission.texts_without_parts(list(getattr(src, "first_page", None) or [])),
              ("title_page", "Title page", getattr(src, "first_page_name", "")))]
    names = list(getattr(src, "declaration_file_names", None) or [])
    pools += [(list(f), ("declaration_file", "Declaration file", names[i] if i < len(names) else ""))
              for i, f in enumerate(getattr(src, "declaration_files", None) or [])]
    out: dict[str, tuple[str, str, dict]] = {}
    for pool, (pk, plabel, pname) in pools:
        paras = [t.strip() for t in pool if t.strip()]
        for i, text in enumerate(paras):
            hit = statement_label(text)
            if hit is None or hit[0] not in ("funding", "conflict", "ethics", "data", "ai") or hit[0] in out:
                continue
            kind, label, content = hit
            if not content:
                lines = []
                for nxt in paras[i + 1:i + 4]:
                    if statement_label(nxt) or label_only_text(nxt) \
                            or re.match(r"(?i)^\W*(?:to\b|dear\b|cover(?:ing)?\s+letter|sincerely|yours)", nxt) or len(nxt) > 1500:
                        break
                    lines.append(nxt)
                content = " ".join(lines)
            content = re.sub(r"\s+", " ", content).strip()
            if not content or len(content) > 2500 or _DATA_TEMPLATE.search(content) or _TEMPLATE.search(content[:40]):
                continue
            out[kind] = (label, _statement(content), provenance.make(pk, plabel, content[:200], file=pname))
    ms._declared_statements = out
    return out


def has_data_availability(ms: Manuscript) -> bool:
    ref_ids = {r.para_id for r in references.extract(ms)}
    return any(_DATA_HEAD.match(p.text) for p in ms.paragraphs if p.id not in ref_ids and len(p.text) <= 1000)


@dataclass
class _Section:
    kind: str           # consent | contributions | funding | conflict | data
    para: Paragraph     # its first paragraph (heading, or the statement itself)
    heading: bool       # para is a heading line (heading_tags may insert at its start)
    end: Paragraph | None = None  # its last paragraph, when known


_LABEL_TYPOS = {"ethical approve": "Ethical approval", "ethics approve": "Ethics approval"}  # jehp_1755, jehp_1280


def _sentence_label(label: str) -> str:
    label = re.sub(r"\s+", " ", label).strip()
    label = _LABEL_TYPOS.get(label.casefold(), label)
    return label if label.isupper() else label[:1] + _sentence_case(label)[1:]


def back_matter_plan(ms: Manuscript, style: JournalStyle, run: bool = False) -> dict:
    """What the back-matter rule does to the consent and contributions sections, for the queries:
    "exists" | "added" | "missing" | None (not checked). With `run`, works it out if the rule has not run."""
    cached = getattr(ms, "_back_matter_plan", None)
    if cached is not None and cached[0] is style:
        return cached[1]
    plan = {"consent": None, "contributions": None, "data": None}
    ms._back_matter_plan = (style, plan)
    if run and _on(style, "back_matter.enabled"):
        _statements(ms, style)
    return plan


def section_headings(sec: dict) -> set[str]:
    """The headings (lower case) that stand for a required section: the journal's own and the variants
    it keeps as written (`keep_labels`: JEHP keeps the author's "Conflict of interest")."""
    return {sec["heading"].lower()} | {k.lower() for k in (sec.get("keep_labels") or [])}


def rule_back_matter(ms: Manuscript, style: JournalStyle) -> list[Change]:
    if not _on(style, "back_matter.enabled"):
        return []
    return _statements(ms, style) + _supplementary(ms, style)


def _statements(ms: Manuscript, style: JournalStyle) -> list[Change]:
    region, anchor = _back_region(ms)
    tag = "<H2>" if style.get("headings.tag") else ""
    levels = _levels(style)
    plan = back_matter_plan(ms, style)
    plan.update(consent=None, contributions=None, data=None, twice=[])
    std = {}
    for sec in style.get("required_sections", []) or []:
        h = sec["heading"].lower()
        kind = "conflict" if "conflict" in h else ("funding" if "financ" in h or "fund" in h else h)
        std[kind] = sec
    out = []
    satisfied: set[str] = set()
    deleted: set[str] = set()   # paragraphs this rule deletes
    consumed: set[str] = set()  # statement paragraphs read with the label before them
    sections: list[_Section] = []
    stated: dict[str, tuple[Paragraph, int]] = {}  # kind -> (the paragraph holding its statement, where it starts)
    ack_seen = False            # the manuscript has an acknowledgement (or a nil one)
    ack_at = None               # (paragraph, label match, statement paragraph) of the acknowledgement
    candidates = []             # nil statements of a required section, which may become that section in place
    delete_nil = _on(style, "back_matter.delete_nil_statements")
    ack_heading = _cfg(style, "back_matter.acknowledgement_heading")
    consent_cfg = _cfg(style, "back_matter.consent") or {}
    contrib_cfg = _cfg(style, "back_matter.contributions") or {}
    consent_on = _cfg(style, "enabled", True) and consent_cfg.get("enabled", False)
    contrib_on = _cfg(style, "enabled", True) and contrib_cfg.get("enabled", False)
    data_cfg = _cfg(style, "back_matter.data_availability") or {}
    data_on = _cfg(style, "enabled", True) and data_cfg.get("enabled", False)
    # The journal's H2 is bold (IJO S: "H2 | Sentence Case, Bold"): the section headings this rule adds, renames or
    # splits off are bold, their statements roman.
    head_bold = _on(style, "back_matter.heading_bold")
    hfmt = {"bold": True} if head_bold else None
    nonempty = [p for p in region if p.text.strip()]

    def standard_label(kind: str, label: str) -> bool:
        """The journal's heading for a required section, a variant it keeps, or one it renames ("Finance support and
        sponsorship")."""
        sec = std.get(kind)
        label = re.sub(r"\s+", " ", label).strip().lower()
        return sec is not None and (label in section_headings(sec) or label in {v.lower() for v in sec.get("variants") or []})

    # Where the manuscript has the section under the journal's heading (or a variant of it), another statement of
    # that kind ("Funding: This study received no external funding." beside "Financial support and sponsorship: Nil")
    # keeps its own label: renaming it would give the section twice.
    standard_at = {}
    for q in nonempty:
        if (mq := _LABEL.match(q.text)) and _kind(mq) in ("funding", "conflict") and standard_label(_kind(mq), mq.group("label")):
            standard_at.setdefault(_kind(mq), set()).add(q.id)
    by_text = {p.id: p.text for p in ms.paragraphs}
    ref_i = references.find_heading(ms)
    # After the reference list (legends, leftovers) a consent or contributions section counts as present, but
    # editors do not restyle it there.
    after_refs = {p.id for p in ms.paragraphs[ref_i:]} if ref_i is not None else set()

    def delete(p: Paragraph):
        out.append(_change(p, 0, len(p.text), "", Layer.STRUCTURE, "sections_back_matter"))
        deleted.add(p.id)

    def label_only(p: Paragraph) -> bool:
        return bool(p.text.strip()) and _words(p.text) <= 8 and not p.text.rstrip().endswith(".")

    def bold_line(p: Paragraph, rule: str):
        """A heading line already there, in bold when the journal's H2 is bold (format only)."""
        if not head_bold or _all_fmt(p, "bold"):
            return
        s = len(p.text) - len(p.text.lstrip())
        e = len(p.text.rstrip())
        if e > s:
            out.append(_change(p, s, e, p.text[s:e], Layer.STRUCTURE, rule, {"bold": True}))

    def tag_heading(p: Paragraph, m, name: str, rename: str | None = None):
        """Tag a heading line the journal's heading list does not know, drop its colon, sentence case."""
        s, e = m.start("label"), m.end("label")
        label = p.text[s:e]
        known = levels.get(_key(p.text)) or p.text.lstrip().startswith("<H")
        new = rename or _sentence_label(label)
        if not known and tag:
            new = tag + new
            s = _tag_start(p.text, s)
            label = p.text[s:e]
        if new != label:
            out.append(_change(p, s, e, new, Layer.STRUCTURE, f"sections_{name}", hfmt))
        end = len(p.text.rstrip())
        if not known and p.text[e:end].strip() in (":", ".", "-", "–", "—"):
            out.append(_change(p, e, end, "", Layer.STRUCTURE, f"sections_{name}"))
        bold_line(p, f"sections_{name}")

    def split_ethics(p: Paragraph, em, content: str):
        """"Ethical approval    All procedures ..." (a label typed before its statement): heading + statement
        (IJO_1745); with a colon, split_run_in does it where the journal splits run-in labels."""
        sep = em.group("sep")
        if not content.strip() or not _sep_ok(sep) or _ETHICS_NIL.match(content):
            return
        if sep.strip() and not _on(style, "back_matter.split_ethics"):
            return
        if not sep.strip() or not _on(style, "back_matter.split_run_in"):
            split_inline(p, em, "back_matter")

    def split_inline(p: Paragraph, m, name: str, rename: str | None = None):
        """"Patient Consent: Written informed ..." -> heading paragraph + statement."""
        c = _change(p, 0, m.end(), f"{tag}{rename or _sentence_label(m.group('label'))}{BR}", Layer.STRUCTURE,
                    f"sections_{name}")
        if head_bold:
            c.segment_format = [{"bold": True}, None]
        out.append(c)

    # where the statements end: the tables, figure legends or captions after the text (a "Note: Data are presented
    # as ..." under a table is not an unlabelled contributions list)
    back_end = next((k for k, q in enumerate(nonempty) if _CAPTION.match(q.text) or _LEGENDS_HEAD.match(q.text)
                     or re.match(r"(?i)^\s*(?:<H\d>)?\s*(?:tables?|figures?)\s*:?\s*$", q.text)), len(nonempty))
    # A section stated twice, once with nothing in it ("None") and once for real: the nil one goes
    nil_twice: set[str] = set()
    hits: dict[str, list[tuple[Paragraph, bool]]] = {}
    for k, q in enumerate(nonempty[:back_end]):
        mq = _LABEL.match(q.text)
        if mq is None or ai_disclosure(q.text) or q.id in after_refs or _kind(mq) not in ("funding", "conflict") or not _sep_ok(mq.group("sep")) \
                and q.text[mq.end():].strip():
            continue
        body = q.text[mq.end():]
        if not body.strip() and k + 1 < len(nonempty):
            body = nonempty[k + 1].text
        hits.setdefault(_kind(mq), []).append((q, _nil(body)))
    for kind, hs in hits.items():
        if len(hs) > 1 and any(not nil for _, nil in hs):
            nil_twice |= {q.id for q, nil in hs if nil}
    # A nil statement under the journal's heading that goes (the real one stays) does not make the real one a second
    # statement: "Funding: This research received no external funding ..." beside "Finance support and sponsorship:
    # The authors received no financial support" (jehp_1580) is the section, under the journal's heading.
    for kind in standard_at:
        standard_at[kind] -= nil_twice
    standard_at = {k: v for k, v in standard_at.items() if v}
    # a paragraph holding only spaces among the statements (a stray bold space left above "Acknowledge- None": jehp_924)
    for q in region:
        if q.text and not q.text.strip() and not q.in_table and q.id not in after_refs and not submission.has_media(q):
            out.append(_change(q, 0, len(q.text), "", Layer.STRUCTURE, "sections_back_matter"))
    for i, p in enumerate(nonempty):
        if p.id in deleted or p.id in consumed:
            continue
        nxt = nonempty[i + 1] if i + 1 < len(nonempty) else None
        # "Declarations" container heading, nil ethics statements
        if _on(style, "back_matter.delete_declarations_heading") and _DECLARATIONS.match(p.text):
            delete(p)
            continue
        # a statement on the use of AI tools ("Use of AI tools", "Disclosure: ChatGPT was used ..."): its own section,
        # never the authors' contributions or a conflicts disclosure
        if ai_label(p.text) or ai_disclosure(p.text):
            said = (ai_label(p.text) or ("", "x"))[1]
            if _on(style, "back_matter.delete_nil_ethics") and (_BARE_NIL.match(said) or not said.strip() and nxt is not None
                                                                and _BARE_NIL.match(nxt.text)):
                # "Use of AI: none": says nothing, goes like a bare "Not applicable." ethics statement (jehp_924)
                delete(p)
                if not said.strip():
                    delete(nxt)
                continue
            bare = not (ai_label(p.text) or ("", "x"))[1]
            sections.append(_Section("ai", p, bare))
            if bare and p.id not in after_refs and not p.text.lstrip().startswith("<H"):
                # "Use of Artificial Intelligence:" over its statement: a heading like the others (editors: 1683)
                tag_heading(p, re.match(r"^\s*(?P<label>.*?[^\s:.\-–—])(?P<sep>\s*[:.\-–—]?\s*)$", p.text), "back_matter")
                if nxt is not None:
                    consumed.add(nxt.id)
            continue
        em = _ETHICS.match(p.text)
        nil_ethics = _on(style, "back_matter.delete_nil_ethics")
        if em and nil_ethics and _cfg(style, "back_matter.nil_ethics_bare_only"):
            # Only a bare "Not applicable." goes, inline or on the line under its heading (with the heading); one that
            # gives a reason stays (split into its own section with back_matter.split_run_in).
            content = p.text[em.end():]
            if content.strip() and _sep_ok(em.group("sep")):
                if _BARE_NIL.match(content):
                    delete(p)
            elif not content.strip() and nxt is not None and _BARE_NIL.match(nxt.text):
                delete(p)
                delete(nxt)
            if p.id not in deleted:
                sections.append(_Section("ethics", p, not content.strip()))
                split_ethics(p, em, content)
            continue
        if em and nil_ethics:
            # Inline only: a heading with "Not applicable" on its own line is kept (ABR editors tag it).
            content = p.text[em.end():]
            if content.strip() and _sep_ok(em.group("sep")) and _ETHICS_NIL.match(content) and _words(content) <= 25:
                delete(p)
            if p.id not in deleted:
                sections.append(_Section("ethics", p, not content.strip()))
                split_ethics(p, em, content)
            continue
        # an unlabelled statement that there are no data ("No new datasets were generated or analysed, so there are no
        # data available."): goes like a nil ethics statement (IJO_2904)
        if delete_nil and nil_ethics and _NO_DATA.match(p.text) and _words(p.text) <= 30 and not data_on:
            delete(p)
            continue
        # data availability: kept as written; it places the sections added around it
        dm = _DATA_HEAD.match(p.text)
        if dm and not data_on and not p.text[dm.end():].strip() and p.id not in after_refs:
            # "Data availability" over its statement: the journal's heading (data_heading), tagged like the others
            sections.append(_Section("data", p, True))
            want = _cfg(style, "back_matter.data_heading")
            if want and not p.text.lstrip().startswith("<H"):
                tag_heading(p, dm, "back_matter", want)
            if nxt is not None:
                consumed.add(nxt.id)
            continue
        if data_on and _DATA_HEAD.match(p.text):
            sections.append(_Section("data", p, not p.text[_DATA_HEAD.match(p.text).end():].strip()))
            continue
        # consent
        cm = _CONSENT_HEAD.match(p.text)
        if cm and consent_on:
            content = p.text[cm.end():]
            sections.append(_Section("consent", p, not content.strip()))
            if p.id in after_refs:
                continue
            stated_text = content if content.strip() else (nxt.text if nxt is not None else "")
            # "Informed consent" over "Patient consent for publication of the images was obtained." is the patient
            # consent section too (S388; IJO_1745)
            rename = consent_cfg.get("heading") if consent_cfg.get("rename_heading") and (re.search(
                r"patient|declaration", cm.group("label"), re.I) or _CONSENT_PUBLICATION.search(stated_text)
                and re.search(r"(?i)\bpatient|\bimages?\b|photo", stated_text)) else None
            if content.strip() and _sep_ok(cm.group("sep")) and _words(content) > 3:
                split_inline(p, cm, "consent", rename)
            elif not content.strip():
                tag_heading(p, cm, "consent", rename)
                if nxt is not None:
                    consumed.add(nxt.id)
            continue
        if consent_on and consent_cfg.get("head_unlabelled", True) and _CONSENT_GIVEN.search(p.text) \
                and _CONSENT_PUBLICATION.search(p.text) and _words(p.text) <= 90 \
                and not any(s.kind == "consent" for s in sections):
            prev = nonempty[i - 1] if i else None
            if prev is not None and prev.id not in deleted and _BARE_CONSENT.match(prev.text) \
                    and not any(x.para is prev for x in sections):
                # "Consent" / "Declaration" over the statement: the section's heading, in the journal's words
                # (IJO_702, IJO_2449: the editors wrote "Declaration of patient consent").
                sections.append(_Section("consent", prev, True))
                consumed.add(p.id)
                if prev.id not in after_refs:
                    s0, e0 = len(prev.text) - len(prev.text.lstrip()), len(prev.text.rstrip())
                    out.append(_change(prev, s0, e0, f"{tag}{consent_cfg.get('heading', 'Declaration of patient consent')}",
                                       Layer.STRUCTURE, "sections_consent", hfmt))
            elif p.id in after_refs:
                sections.append(_Section("consent", p, False))
            elif prev is None or prev.id in deleted or not label_only(prev):
                sections.append(_Section("consent", p, False))
                out.append(_change(p, 0, 0, f"{tag}{consent_cfg.get('heading', 'Declaration of patient consent')}{BR}",
                                   Layer.STRUCTURE, "sections_consent", {"bold": False, "italic": False}))
                out[-1].segment_format = [hfmt] if head_bold else None
            continue
        # authors' contributions
        km = _CONTRIB.match(p.text)
        if km and contrib_on and not km.group("author") and not km.group("label").lower().startswith(
                ("credit", "contribution list", "contributorship")):
            km = None  # "Contribution: This author helped ..." under each author's name is not the section's label
        if km and contrib_on:
            content = p.text[km.end():]
            sections.append(_Section("contributions", p, not content.strip()))
            if p.id in after_refs:
                continue
            if content.strip() and _sep_ok(km.group("sep")):
                if contrib_cfg.get("split_inline", True):  # "Contribution list: A.S.: ..." -> the journal's heading
                    split_inline(p, km, "contributions", contrib_cfg.get("heading", "Author contributions")
                                 if km.group("label").lower().startswith("contribution list") else None)
            elif not content.strip():
                tag_heading(p, km, "contributions")
            continue
        if contrib_on and contrib_cfg.get("head_unlabelled", True) and _CONTRIB_LINE.match(p.text) \
                and len(_CONTRIB_WORDS.findall(p.text)) >= 2 and not _OTHER_LABEL.match(p.text) \
                and not any(s.kind == "contributions" for s in sections) and i < back_end \
                and _names_an_author(p.text, ms):
            prev = nonempty[i - 1] if i else None
            if p.id in after_refs:
                sections.append(_Section("contributions", p, False))
            elif prev is None or prev.id in deleted or not label_only(prev) or _LABEL.match(prev.text):
                sections.append(_Section("contributions", p, False))
                out.append(_change(p, 0, 0, f"{tag}{contrib_cfg.get('heading', 'Author contributions')}{BR}",
                                   Layer.STRUCTURE, "sections_contributions", {"bold": False, "italic": False}))
                out[-1].segment_format = [hfmt] if head_bold else None
            continue
        # funding / conflicts / acknowledgement
        m = _LABEL.match(p.text)
        if not m or (m.group("sep") and m.group("sep").isspace() and _words(p.text[m.end():]) > 30):
            if delete_nil and _NIL_SENTENCE.match(p.text) and _NIL_TOPIC.search(p.text) and _words(p.text) <= 35 \
                    and _nil(p.text):
                prev = nonempty[i - 1] if i else None
                if prev is None or prev.id in deleted or not label_only(prev):
                    delete(p)
            continue
        kind = _kind(m)
        if kind == "funding" and m.group("label").lower() == "support" and not m.group("sep").strip():
            continue  # a sentence starting "Support ..."
        if p.id in nil_twice:
            # "Financial relationship and sponsership: None" beside "Source of funding- This research was supported
            # ...": the empty one goes, the real one is the section (editors: jehp_924)
            delete(p)
            if not p.text[m.end():].strip() and nxt is not None and not _LABEL.match(nxt.text):
                delete(nxt)
            continue
        twice = kind in ("funding", "conflict") and kind in stated and stated[kind][0].id != p.id \
            and p.id not in after_refs and kind in std
        if twice or kind in standard_at and standard_at[kind] - {p.id} and not standard_label(kind, m.group("label")) \
                and p.id not in after_refs:
            # A second statement of a section already there ("Financial Support: Nil." and later "Funding: Nil. No
            # financial support was received ...": jehp_1683) keeps its own label, never the same journal heading
            # twice; the author is asked which one stays.
            if twice:
                plan.setdefault("twice", []).append((p.id, std[kind]["heading"]))
            if p.text[m.end():].strip() and _sep_ok(m.group("sep")):
                split_inline(p, m, "back_matter")
            elif not p.text[m.end():].strip():
                tag_heading(p, m, "back_matter")
                if nxt is not None and not _LABEL.match(nxt.text):
                    consumed.add(nxt.id)
            continue
        content = p.text[m.end():]
        content_p = None
        if not content.strip():
            if nxt is not None and not _LABEL.match(nxt.text):
                content_p, content = nxt, nxt.text
                consumed.add(nxt.id)
        nil = _nil(content)
        kinds = ("conflict", "funding") if kind == "both" else (kind,)
        if kind == "ack":
            ack_seen = True
            if nil and (delete_nil or _BARE_NIL.match(content)) and (content.strip() or content_p is None
                                                                     and i + 1 < len(nonempty)):
                # "Acknowledge- None": nothing to acknowledge, the line goes (S "Retain if given"; editors: jehp_924)
                delete(p)
                if content_p is not None:
                    delete(content_p)
            elif ack_heading and not p.text[m.end():].strip():
                s, e = m.start("ack"), m.end("ack")
                label = p.text[s:e]
                # "singular": the stylesheets' 'If heading "Acknowledgements" is present, then please delete "s"'
                # (the author's spelling otherwise kept: "Acknowledgments" -> "Acknowledgment").
                want = (label[:-1] if label[-1:] in "sS" else label) if ack_heading == "singular" else ack_heading
                if label != want:
                    out.append(_change(p, s, e, want, Layer.STRUCTURE, "sections_back_matter", hfmt))
                bold_line(p, "sections_back_matter")
            if p.id not in deleted and p.id not in after_refs and ack_at is None:
                ack_at = (p, m, content_p)
            continue
        kept_labels = section_headings(std[kind]) - {std[kind]["heading"].lower()} if kind in std else set()
        standard = kind in std and m.group("label").lower() in section_headings(std[kind])
        if nil and content_p is not None and standard:
            # The standard heading with a "None" under it: keep both, word the statement as the journal does.
            satisfied.add(kind)
            sections.append(_Section(kind, p, True))
            stated.setdefault(kind, (content_p, 0))
            bold_line(p, "sections_back_matter")
            body = content_p.text.strip()
            default = std[kind].get("default", "")
            short = body.rstrip(".").lower() in ("none", "nil", "not applicable", "na", "n/a")
            reword = _on(style, "back_matter.reword_nil") or short
            if reword and body and body.rstrip(".").lower() != default.rstrip(".").lower() and (short or _nil(body)):
                # "None" -> "Nil" / the journal's wording; "The authors declare no conflict of interest" -> the same
                s = content_p.text.index(body)
                new = ("Nil" + body[len(body.rstrip(".")):] if short else "Nil.") if kind == "funding" else default
                out.append(_change(content_p, s, s + len(body), new, Layer.STRUCTURE, "sections_back_matter"))
            continue
        if nil and delete_nil:
            delete(p)
            if content_p is not None:
                delete(content_p)
            if kind in std and p.id not in after_refs and (_sep_ok(m.group("sep")) or content_p is not None):
                candidates.append((kind, p, m, content_p))
            continue
        if kind == "both" or not all(k in std for k in kinds):
            satisfied.update(kinds)  # a real joint statement: leave it to the editor
            continue
        sec = std[kind]
        satisfied.add(kind)
        sections.append(_Section(kind, p, content_p is not None or not content.strip()))
        stated.setdefault(kind, (content_p or p, 0 if content_p is not None else m.end()))
        if content_p is None and content.strip():
            # "Funding: Eye Research Foundation grant 12." -> heading paragraph + statement
            out.append(_change(p, 0, m.end(), f"{tag}{sec['heading']}{BR}", Layer.STRUCTURE, "sections_back_matter"))
            if head_bold:
                out[-1].segment_format = [{"bold": True}, None]
        else:
            s, e = m.start("label"), m.end("label")
            head_end = len(p.text.rstrip())
            label_text = p.text[s:e]
            heading = label_text if label_text.lower() in kept_labels else sec["heading"]
            tagged = levels.get(_key(p.text)) and not p.text.lstrip().startswith("<H")
            if label_text != heading or (not tagged and tag and not p.text.lstrip().startswith("<H")):
                new = heading if tagged or p.text.lstrip().startswith("<H") else tag + heading
                if tag and new.startswith(tag):
                    s = _tag_start(p.text, s)
                out.append(_change(p, s, e, new, Layer.STRUCTURE, "sections_back_matter", hfmt))
            if not tagged and p.text[e:head_end].strip() in (":", "."):
                out.append(_change(p, e, head_end, "", Layer.STRUCTURE, "sections_back_matter"))
            if content_p is not None or not content.strip():
                bold_line(p, "sections_back_matter")

    if _on(style, "back_matter.split_run_in"):
        _split_run_in(style, nonempty, std, levels, tag, deleted, consumed, after_refs, satisfied, out, standard_label)

    moved_consent = _move_consent(ms, style, sections, nonempty, deleted, after_refs, out) \
        if consent_on and consent_cfg.get("move") else None
    moved_ack = _move_acknowledgement(ms, style, ack_at, sections, deleted, after_refs, levels, tag, out) \
        if ack_at is not None and _on(style, "back_matter.acknowledgement_first") else None

    # A nil statement becomes the journal's section in place ("Funding: None" -> "Financial support and sponsorship
    # / Nil.") when that keeps the journal's order; the rest stay deleted and the sections are added below.
    order = list(_cfg(style, "back_matter.order") or ["consent", "contributions", "funding", "conflict"])
    if "ack" not in order:  # an acknowledgement added from the title page goes before the financial support, or
        # (G148, back_matter.acknowledgement_first) before all the statements
        order.insert(0 if _on(style, "back_matter.acknowledgement_first") else
                     order.index("funding") if "funding" in order else len(order), "ack")
    rank = {k: i for i, k in enumerate(order)}
    pos = {p.id: i for i, p in enumerate(ms.paragraphs)}
    done: set[str] = set(satisfied)
    for kind, p, m, content_p in candidates:
        if kind in done or kind not in std or kind not in rank:
            continue
        here = [x for x in sections if x.para.id not in deleted]
        if any(rank.get(k, -1) < rank[kind] and k not in done for k in std) \
                or any(rank.get(x.kind, -1) > rank[kind] and pos[x.para.id] < pos[p.id] for x in here) \
                or any(rank.get(x.kind, -1) < rank[kind] and pos[x.para.id] > pos[p.id] for x in here):
            continue
        sec = std[kind]
        out[:] = [c for c in out if not (c.para_id in (p.id, getattr(content_p, "id", None))
                                          and c.start == 0 and c.end == len(by_text[c.para_id]) and not c.new)]
        deleted.discard(p.id)
        s0, e0 = m.start("label"), m.end("label")
        lead = len(p.text) - len(p.text.lstrip())
        known = levels.get(_key(p.text)) or p.text.lstrip().startswith("<H")  # heading_tags tags it: no second tag
        out.append(_change(p, lead, e0, f"{'' if known else tag}{sec['heading']}", Layer.STRUCTURE,
                           "sections_back_matter", hfmt))
        if content_p is None:
            out.append(_change(p, e0, len(p.text), BR + sec.get("default", ""), Layer.STRUCTURE, "sections_back_matter",
                               {"bold": False} if head_bold else None))
        else:
            if e0 < len(p.text):
                out.append(_change(p, e0, len(p.text), "", Layer.STRUCTURE, "sections_back_matter"))
            deleted.discard(content_p.id)
            out.append(_change(content_p, 0, len(content_p.text), sec.get("default", ""), Layer.STRUCTURE,
                               "sections_back_matter"))
        done.add(kind)
        satisfied.add(kind)
        sections.append(_Section(kind, p, True, content_p or p))

    # --- sections to add: the standard statements, patient consent, authors' contributions ---
    new: dict[str, list[str]] = {}
    origins: dict[str, dict | None] = {}  # kind -> where its text comes from (provenance)
    # Statements given on the title page left above the main text ("FUNDING/SUPPORT- Hyderabad Eye Research
    # Foundation ...", deleted there by submission.py): the section's text when the back has none.
    given: dict[str, tuple[str, dict]] = {}
    for label, content, fp in submission.front_statements(ms, getattr(ms, "_submission_style", None) or style):
        m = _LABEL.match(f"{label.replace('/', ' ')}: {content}")
        if m and _kind(m) in ("funding", "conflict") and not _nil(content):
            given.setdefault(_kind(m), (_statement(content), provenance.make(
                "manuscript", "Manuscript title-page statement", _statement(content), para_id=fp.id)))
    # A funding sentence inside the conflicts statement ("The authors declare no conflict of interest. The study received
    # no commercial funding.": jehp_1209) is the funding statement: it moves under the funding heading (owner,
    # 2026-10-27: never "Nil." when the submission states it).
    if "funding" in std and "funding" not in satisfied and "conflict" in stated:
        cp, at = stated["conflict"]
        moved = _funding_sentence(cp.text, at)
        if moved is not None:
            s0, e0 = moved
            out.append(_change(cp, s0, e0, "", Layer.STRUCTURE, "sections_back_matter"))
            given["funding"] = (_statement(cp.text[s0:e0]), provenance.make(
                "manuscript", "Manuscript conflicts statement", cp.text[s0:e0].strip(), para_id=cp.id))
    # Statements the submission gives outside the manuscript (title page, declaration files): owner, 2026-10-27 (G283)
    from_sources = set(_cfg(style, "back_matter.statements_from_sources") or [])
    declared = declared_statements(ms) if from_sources else {}
    reword = _on(style, "back_matter.reword_nil")
    for kind in ("funding", "conflict"):
        # a bare "None." / "Nil." is the journal's default anyway (reworded where the journal rewords)
        if kind in declared and kind in from_sources and kind not in given \
                and not (reword and _BARE_NIL.match(declared[kind][1])):
            given[kind] = (declared[kind][1], declared[kind][2])
    for kind, sec in std.items():
        if kind not in satisfied:
            if kind in given:
                new[kind] = [f"{tag}{sec['heading']}", given[kind][0]]
                origins[kind] = given[kind][1]
                continue
            new[kind] = [f"{tag}{sec['heading']}", sec.get("default", "")]
            origins[kind] = provenance.make("journal_style", "Journal's standard text",
                                            f"{sec['heading']}: the journal's wording for a missing section")
    if ack_heading is not None and not ack_seen and _on(style, "back_matter.acknowledgement_from_sources"):
        # S382 "Acknowledgments | ... Retain if given": an acknowledgement given only on the title page
        got = acknowledgement_from_sources(ms)
        if got:
            label, text = got
            want = (label[:-1] if label[-1:] in "sS" else label) if ack_heading == "singular" else ack_heading
            ack_tag = f"<{levels[want.lower()]}>" if tag and want.lower() in levels else tag  # IJO, JEHP: <H1>
            new["ack"] = [f"{ack_tag}{_sentence_label(want)}", text]
            fsrc = getattr(ms, "front_sources", None)
            origins["ack"] = provenance.make("title_page", "Title page", text, file=getattr(fsrc, "first_page_name", ""))
    have = {s.kind for s in sections}
    if moved_consent:
        new["consent"], origins["consent"] = moved_consent
    if consent_on:
        if "consent" in have or has_consent_statement(ms):
            plan["consent"] = "exists"
        elif consent_cfg.get("add_missing", True) and identifiable_patient_images(ms):
            new["consent"] = [f"{tag}{consent_cfg.get('heading', 'Declaration of patient consent')}",
                              consent_cfg.get("text", "")]
            origins["consent"] = provenance.make("inferred", "Patient images in legends", "a figure legend shows the "
                                                 "patient and no consent statement was found: the journal's standard text")
            plan["consent"] = "added"
        else:
            plan["consent"] = "missing"
    if contrib_on:
        # Statements on the manuscript's own title page move to the back (per journal: IJO): a tracked delete there and
        # the statements inserted under the heading. When the back already has the section, the top copy is deleted
        # only when every statement is already there (a duplicate); otherwise it is left to the editor.
        top, top_items = top_contributions(ms, {p.id for p in region}) if contrib_cfg.get("move_from_top") \
            else ([], [])
        if "contributions" in have:
            plan["contributions"] = "exists"
            back = _plain(" ".join(p.text for p in region))
            if top and all(_plain(t) in back for t in top_items if not _APPROVAL.match(t)):
                for p in top:
                    out.append(_change(p, 0, len(p.text), "", Layer.STRUCTURE, "sections_contributions"))
        elif top:
            new["contributions"] = [f"{tag}{contrib_cfg.get('heading', 'Author contributions')}"] + top_items
            origins["contributions"] = provenance.make("manuscript", "Manuscript title page", top_items[0],
                                                       para_id=top[0].id)
            plan["contributions"] = "added"
            for p in top:
                out.append(_change(p, 0, len(p.text), "", Layer.STRUCTURE, "sections_contributions"))
        else:
            items, origin = contributions_with_origin(ms, {p.id for p in region}) \
                if contrib_cfg.get("from_sources") else ([], None)
            if items:
                new["contributions"] = [f"{tag}{contrib_cfg.get('heading', 'Author contributions')}"] + items
                origins["contributions"] = origin
                plan["contributions"] = "added"
            else:
                plan["contributions"] = "missing"
    if data_on:
        if "data" in have or has_data_availability(ms):
            plan["data"] = "exists"
        else:
            text = data_availability_from_sources(ms) if data_cfg.get("from_sources", True) else ""
            if text:
                new["data"] = [f"{tag}{data_cfg.get('heading', 'Data availability statement')}", text]
                fsrc = getattr(ms, "front_sources", None)
                origins["data"] = provenance.make("title_page", "Title page", text, file=getattr(fsrc, "first_page_name", ""))
                plan["data"] = "added"
            else:
                plan["data"] = "missing"
    region_text = [p.text for p in nonempty if p.id not in deleted]
    for kind in ("ethics", "data", "ai"):
        if kind not in declared or kind not in from_sources or kind in new or any(s.kind == kind for s in sections):
            continue
        label, text, origin = declared[kind]
        if _ETHICS_NIL.match(text) or _BARE_NIL.match(text):
            continue  # "Not applicable.": nothing to state
        if kind == "ethics" and any(_ETHICS.match(t) for t in region_text):
            continue
        if _said_already(ms, text):
            continue  # the body states it already (the methods' "Ethical consideration": jehp_1153): not twice
        if kind == "data" and has_data_availability(ms):
            continue
        if kind == "ai" and any(ai_label(t) or _AI_TEXT.search(t) and re.search(r"(?i)\b(?:used|use|assist\w*)\b", t)
                                for t in region_text):
            continue  # the manuscript says it already (an acknowledgement naming the tool counts)
        if kind == "data":
            heading = _cfg(style, "back_matter.data_heading") or data_cfg.get("heading", "Data availability statement")
        else:
            heading = _sentence_label(label)
        new[kind] = [f"{tag}{heading}", text]
        origins[kind] = origin
    if not _on(style, "back_matter.add_missing"):
        new = {k: v for k, v in new.items() if k in ("consent", "contributions", "data", "ack", "ethics", "ai")}
    if anchor is None:
        # No reference list found: add the sections after the last statement read here.
        seen = [p for p in nonempty if p.id in deleted or p.id in consumed or any(x.para is p for x in sections)]
        anchor = seen[-1] if seen else None
    if (not new or anchor is None) and moved_ack is None:
        return out
    # Each new section goes before the first section already there that comes after it in the journal's order
    # (consent, contributions, financial support, conflicts of interest), else with the others before the references.
    # A new section goes after the one before it in that order when that one is there (and nothing that comes
    # after it in the order stands between), else before the first one after it, else before the references.
    kept = [s for s in sections if s.para.id not in deleted]
    starts = {s.para.id for s in sections}
    at = {p.id: i for i, p in enumerate(nonempty)}

    def section_end(x: _Section) -> Paragraph:
        if x.end is not None:
            return x.end
        last = x.para
        for q in nonempty[at[x.para.id] + 1:at[x.para.id] + 31]:
            if q.id in starts or q.id in deleted or q.in_table or _CAPTION.match(q.text) or (q.id in after_refs) != (
                    x.para.id in after_refs) or _key(q.text) in ("references", "reference", "bibliography") or any(
                    r.match(q.text) for r in (_LABEL, _CONSENT_HEAD, _ETHICS, _DECLARATIONS, _DATA_HEAD)) \
                    or _section_contrib_label(q.text) or ai_label(q.text) \
                    or (x.kind != "contributions" and label_only(q)):
                break
            last = q
        return last

    targets: dict[tuple[str, str], list[str]] = {}
    target_kinds: dict[tuple[str, str], list[str]] = {}
    if moved_ack is not None:  # G148: the acknowledgement first, the statements after it
        above, lines, origin = moved_ack
        targets[("after", above.id)] = list(lines)
        target_kinds[("after", above.id)] = ["ack"]
        origins["ack"] = origin
        new.setdefault("ack_moved", lines)
    for kind in sorted(new, key=lambda k: rank.get(k, len(order))):
        if kind == "ack_moved":
            continue
        r = rank.get(kind, len(order))
        later = [s for s in kept if rank.get(s.kind, -1) > r]
        before = min(later, key=lambda s: pos[s.para.id]) if later else None
        earlier = [s for s in kept if rank.get(s.kind, len(order)) < r and s.para.id not in after_refs]
        prev_end = max((section_end(s) for s in earlier), key=lambda q: pos[q.id]) if earlier else None
        if prev_end is not None and pos[prev_end.id] < pos[anchor.id] and all(
                q.id in deleted for q in nonempty if pos[prev_end.id] < pos[q.id] <= pos[anchor.id]):
            prev_end = anchor  # only deleted statements between: with the others, after them
        if prev_end is not None and (before is None or pos[prev_end.id] < pos[before.para.id]):
            (targets.setdefault(("after", prev_end.id), []).extend(new[kind]), target_kinds.setdefault(("after", prev_end.id), []).append(kind))
        elif before is None or pos[before.para.id] > pos[anchor.id]:
            (targets.setdefault(("after", anchor.id), []).extend(new[kind]), target_kinds.setdefault(("after", anchor.id), []).append(kind))
        elif not before.heading and not any(c.para_id == before.para.id and c.start == 0 and c.end > 0 for c in out):
            (targets.setdefault(("before", before.para.id), []).extend(new[kind]), target_kinds.setdefault(("before", before.para.id), []).append(kind))
        else:
            # Before a heading line (heading_tags inserts its tag there): after the paragraph above it.
            j = ms.paragraphs.index(before.para)
            # (a stray "." line is a template leftover that goes whole, with whatever is put on it: jehp_1193)
            above = next((q for q in reversed(ms.paragraphs[:j]) if re.search(r"\w", q.text) and not q.in_table
                          and q.id not in deleted), None)
            if above is None:
                continue
            (targets.setdefault(("after", above.id), []).extend(new[kind]), target_kinds.setdefault(("after", above.id), []).append(kind))
    # a target paragraph this rule deletes ("after" the moved consent's old place): the nearest one above it that stays
    alive = [q for q in ms.paragraphs if re.search(r"\w", q.text) and not q.in_table]
    for key in [k for k in targets if k[0] == "after" and k[1] in deleted]:
        j = next(n for n, q in enumerate(alive) if q.id == key[1])
        keep = next((q for q in reversed(alive[:j]) if q.id not in deleted), None)
        if keep is None:
            continue
        lines, kinds = targets.pop(key), target_kinds.pop(key, [])
        targets.setdefault(("after", keep.id), []).extend(lines)
        target_kinds.setdefault(("after", keep.id), []).extend(kinds)
    by_id = {p.id: p for p in ms.paragraphs}
    fmt = {"bold": False, "italic": False, "superscript": False}
    consent_text = (new.get("consent") or [None, None])[1] if plan.get("consent") == "added" else None
    heads = {v[0] for v in new.values() if v}

    bold_leads = getattr(getattr(ms, "front_sources", None), "bold_leads", None) or {}

    def segments(lines: list[str], lead: bool) -> list | None:
        """Bold for the heading lines (head_bold), the statements as `fmt` with the bold lead they have in the title
        page ("Concept:" in "Concept: Mehmet ..."); `lead`: an empty first segment."""
        segs = [hfmt if ln in heads and head_bold else
                ([[len(bold_leads[ln]), {"bold": True}]] if bold_leads.get(ln) else None) for ln in lines]
        if not any(segs):
            return None
        return ([None] if lead else []) + segs

    for (where, pid), lines in targets.items():
        p = by_id[pid]
        prov = provenance.combine([origins.get(k) for k in target_kinds.get((where, pid), [])])
        # The standard consent statement asserts that consent was obtained: its own reason, so that autopilot can hold it.
        # Missing standard statements added with the journal's default text ("Nil."): their own reason too.
        # (a statement read from the submission is not the journal's default text: not ".added")
        reason = "sections_back_matter.consent_added" if consent_text and consent_text in lines else \
            "sections_back_matter.added" if any(k in std and new[k][1:] == [std[k].get("default", "")]
                                                for k in target_kinds.get((where, pid), [])) else "sections_back_matter"
        if where == "after":
            out.append(_change(p, len(p.text), len(p.text), BR + BR.join(lines), Layer.STRUCTURE, reason, fmt))
            out[-1].segment_format = segments(lines, True)
        else:
            # Merge with a heading this rule adds at the same point ("before" an unlabelled statement).
            mine = next((c for c in out if c.para_id == pid and c.start == c.end == 0), None)
            if mine is not None:
                out.remove(mine)
                out.append(_change(p, 0, 0, BR.join(lines) + BR + mine.new, Layer.STRUCTURE, reason, mine.format))
                if head_bold:
                    out[-1].segment_format = (segments(lines, False) or [None] * len(lines)) \
                        + list(mine.segment_format or [None])
            else:
                out.append(_change(p, 0, 0, BR.join(lines) + BR, Layer.STRUCTURE, reason, fmt))
                out[-1].segment_format = segments(lines, False)
        if prov:  # the change just added (a merged heading included)
            out[-1].provenance = prov
    return out


_FUNDING_WORDS = re.compile(r"\b(?:fund(?:ing|ed|s)?|grants?|sponsor\w*|financial\s+support|financially\s+supported)\b", re.I)


def _funding_sentence(text: str, start: int = 0) -> tuple[int, int] | None:
    """(start, end) of a sentence on funding inside a conflicts statement that has another sentence of its own: "The
    authors declare no conflict of interest. The study received no commercial funding." -> the second."""
    body = text[start:]
    spans = [(m.start() + start, m.end() + start) for m in re.finditer(r"[^.!?]+[.!?]*", body) if m.group().strip()]
    if len(spans) < 2:
        return None
    hits = [(s, e) for s, e in spans if _FUNDING_WORDS.search(text[s:e])
            and not re.search(r"(?i)conflict|competing|interests?\b", text[s:e])]
    if len(hits) != 1:
        return None
    s, e = hits[0]
    while s > start and text[s - 1] == " ":
        s -= 1  # the space before it goes with it
    if s == start:
        e = spans[1][0] if spans[0][0] == hits[0][0] else e
    return s, e


_STATEMENT_WORDS = re.compile(r"\b(?:registration|trial|availability|consent|ethic\w*|approval|AI|artificial|generative"
                              r"|funding|financial|finance|sponsor\w*|conflicts?|competing|acknowledge?ments?|contributions?"
                              r"|declarations?|disclosures?)\b", re.I)
_STATEMENT_LABEL = re.compile(r"([A-Z][^:.!?;\n]{1,70}?)(\s*:\s+)(?=\S)")
_DATA_LABEL = re.compile(r"^(?:data\s+availability|availability\s+of\s+data)", re.I)


def _split_run_in(style: JournalStyle, nonempty: list[Paragraph], std: dict, levels: dict, tag: str, deleted: set,
                  consumed: set, after_refs: set, satisfied: set, out: list[Change], standard_label) -> None:
    """Back-matter statements given as run-in labels ("Acknowledgement: The authors ...", "Consent to Publish: Not
    applicable. This study ...") -> the heading on its own line and the statement under it; several run on in one
    paragraph ("Trial registration: ... . Financial support and sponsorship: Nil. Conflicts of interest: ...") -> one
    section each. The required sections' labels take the journal's heading; a bare "Not applicable." ethics / consent
    / trial statement goes (back_matter.nil_ethics_bare_only)."""
    touched = {c.para_id for c in out}
    bare = _on(style, "back_matter.delete_nil_ethics") and _cfg(style, "back_matter.nil_ethics_bare_only")
    ack_heading = _cfg(style, "back_matter.acknowledgement_heading")
    contrib_heading = (_cfg(style, "back_matter.contributions") or {}).get("heading", "Author contributions")
    data_heading = _cfg(style, "back_matter.data_heading")
    for p in nonempty:
        text = p.text
        if p.id in deleted or p.id in consumed or p.id in after_refs or p.in_table or _CAPTION.match(text):
            continue
        starts = [0] + [m.end() for m in re.finditer(r"[.!?]\s+", text)]
        labels = []  # (segment start, label start, label end, statement start)
        for k in starts:
            m = _STATEMENT_LABEL.match(text, k)
            if not m or len(m.group(1).split()) > 8 or not _STATEMENT_WORDS.search(m.group(1)):
                continue
            km = _CONTRIB.match(m.group(1) + ":")
            if km and not km.group("author") and not km.group("label").lower().startswith(("credit", "contribution list", "contributorship")):
                continue  # "Contribution: This author helped ..." under an author's name
            if k == 0 and (p.id in touched or not text[m.end():].strip()):
                continue  # the rules above have dealt with the first label
            ws = k if k == 0 else len(text[:k].rstrip())
            labels.append((ws, m.start(1), m.end(1), m.end()))
        if not labels:
            continue
        for n, (ws, ls, le, cs) in enumerate(labels):
            seg_end = labels[n + 1][0] if n + 1 < len(labels) else len(text.rstrip())
            label = text[ls:le]
            content = text[cs:seg_end]
            if bare and _ETHICS.match(label + ": ") and _BARE_NIL.match(content):
                # "Consent for publication: Not applicable." goes
                if ws == 0:
                    nxt_start = labels[n + 1][1] if n + 1 < len(labels) else len(text)
                    out.append(_change(p, 0, nxt_start, "", Layer.STRUCTURE, "sections_back_matter"))
                else:
                    out.append(_change(p, ws, seg_end, "", Layer.STRUCTURE, "sections_back_matter"))
                continue
            lm = _LABEL.match(label + ":")
            kind = _kind(lm) if lm and lm.end("label") == len(label) and not ai_disclosure(text[ls:]) else None
            if kind in ("funding", "conflict") and kind in std:
                low = re.sub(r"\s+", " ", label).lower()
                kept = {k.lower() for k in std[kind].get("keep_labels") or []}
                name = label if low in kept else std[kind]["heading"]
                satisfied.add(kind)
            elif kind == "ack":
                name = (label[:-1] if label[-1:] in "sS" else label) if ack_heading == "singular" else (ack_heading or label)
                name = _sentence_label(name)
            elif _CONTRIB.match(label + ":"):
                name = contrib_heading
            elif data_heading and _DATA_LABEL.match(label):
                name = data_heading
            else:
                name = _sentence_label(label)
            level = levels.get(name.lower(), "H2")
            if name.isupper() and level != "H1" and style.get("headings.subheading_case", "keep") == "sentence":
                name = caps_sentence(name, " ".join(q.text for q in nonempty))  # S: "H2 | Sentence Case"
            head = (f"<{level}>" if tag else "") + name
            if ws == 0:
                out.append(_change(p, 0, cs, head + BR, Layer.STRUCTURE, "sections_back_matter"))
            else:
                out.append(_change(p, ws, cs, BR + head + BR, Layer.STRUCTURE, "sections_back_matter"))


def caps_sentence(text: str, body: str) -> str:
    """An all-capitals heading in sentence case: "COMPLIANCE WITH JEHP AI POLICY STATEMENT" -> "Compliance with JEHP AI
    policy statement"; a word the text writes in capitals elsewhere (an acronym, "JEHP", "AI") keeps them."""
    words = text.split(" ")
    out = []
    for i, w in enumerate(words):
        core = re.sub(r"\W", "", w)
        if core and re.search(rf"(?<![\w-]){re.escape(core)}(?![\w-])", body.replace(text, " ")):
            out.append(w)
        elif i == 0:
            out.append(w[:1] + w[1:].lower())
        else:
            out.append(w.lower())
    return " ".join(out)


def _move_consent(ms: Manuscript, style: JournalStyle, sections: list[_Section], nonempty: list[Paragraph],
                  deleted: set[str], after_refs: set[str], out: list[Change]) -> tuple[list[str], dict] | None:
    """S392 (IJO): the patient consent section's "Placement is before the "Financial support and sponsorship"
    section". A consent section after the reference list (or the figure legends), or after a funding / conflicts
    statement, is moved: its paragraphs are deleted there (tracked) and its heading and statement come back as a
    section placed in the journal's order with the others. Returns the lines to insert and their provenance."""
    cons = next((x for x in sections if x.kind == "consent" and x.para.id not in deleted), None)
    if cons is None:
        return None
    at = {p.id: i for i, p in enumerate(nonempty)}
    if cons.para.id not in at:
        return None
    i = at[cons.para.id]
    funding_before = any((m := _LABEL.match(q.text)) and _kind(m) in ("funding", "conflict", "both")
                         for q in nonempty[:i] if q.id not in after_refs)
    if cons.para.id not in after_refs and not funding_before:
        return None
    cfg = _cfg(style, "back_matter.consent") or {}
    standard = cfg.get("heading", "Declaration of patient consent")
    tag = "<H2>" if style.get("headings.tag") else ""
    p = cons.para
    paras = [p]
    statements: list[str] = []
    cm = _CONSENT_HEAD.match(p.text)
    if cm:
        label = cm.group("label")
        heading = standard if cfg.get("rename_heading") and re.search(r"patient|declaration", label, re.I) \
            else _sentence_label(label)
        content = p.text[cm.end():].strip()
        if content:
            statements.append(content)
    elif _BARE_CONSENT.match(p.text):
        heading = standard
    else:  # an unlabelled statement
        heading = standard
        statements.append(p.text.strip())
    if not statements or cons.heading:
        starts = {x.para.id for x in sections}
        for q in nonempty[i + 1:i + 4]:
            if q.id in starts or q.id in deleted or q.in_table or _CAPTION.match(q.text) or _LEGEND.match(q.text) \
                    or (q.id in after_refs) != (p.id in after_refs) or label_only_text(q.text) \
                    or _key(q.text) in ("references", "reference", "bibliography") \
                    or any(r.match(q.text) for r in (_LABEL, _CONSENT_HEAD, _CONTRIB, _ETHICS, _DECLARATIONS, _DATA_HEAD)):
                break
            paras.append(q)
            statements.append(q.text.strip())
    if not statements:
        return None
    said = " ".join(statements)
    if cm and cfg.get("rename_heading") and _CONSENT_PUBLICATION.search(said) \
            and re.search(r"(?i)\bpatient|\bimages?\b|photo", said):
        heading = standard  # "Informed consent" over the patient's consent to publication (S388)
    gone = {q.id for q in paras}
    out[:] = [c for c in out if c.para_id not in gone]
    for q in paras:
        out.append(_change(q, 0, len(q.text), "", Layer.STRUCTURE, "sections_consent"))
        deleted.add(q.id)
    origin = provenance.make("manuscript", "Manuscript", "the patient consent section, moved before Financial support "
                             "and sponsorship (stylesheet)", para_id=p.id)
    return [f"{tag}{heading}"] + statements, origin


def _move_acknowledgement(ms: Manuscript, style: JournalStyle, ack_at, sections: list[_Section], deleted: set[str],
                          after_refs: set[str], levels: dict, tag: str, out: list[Change]):
    """G148: the declarations (conflicts, contributions, funding, ...) go in the acknowledgement section, after it. An
    acknowledgement the author put after the statements (jehp_1588: "Acknowledgements: ..." last, the statements left
    under the Conclusion) is moved before the first of them: its paragraphs are deleted there (tracked) and its heading
    and statement come back after the paragraph above the statements. Returns (that paragraph, the lines, provenance)."""
    p, m, content_p = ack_at
    pos = {q.id: i for i, q in enumerate(ms.paragraphs)}
    kinds = {"consent", "contributions", "funding", "conflict", "ethics", "data", "ai"}
    first = min((x for x in sections if x.kind in kinds and x.para.id not in deleted and x.para.id not in after_refs
                 and pos[x.para.id] < pos[p.id]), key=lambda x: pos[x.para.id], default=None)
    if first is None:
        return None
    label = m.group("ack")
    ack_heading = _cfg(style, "back_matter.acknowledgement_heading")
    name = (label[:-1] if label[-1:] in "sS" else label) if ack_heading == "singular" else (ack_heading or label)
    name = _sentence_label(name)
    head = (f"<{levels.get(name.lower(), 'H2')}>" if tag else "") + name
    inline = p.text[m.end():].strip()
    paras = [p] + ([content_p] if content_p is not None and not inline else [])
    statements = [inline] if inline else [q.text.strip() for q in paras[1:]]
    if not statements:
        return None
    j = pos[first.para.id]
    above = next((q for q in reversed(ms.paragraphs[:j]) if q.text.strip() and not q.in_table and q.id not in deleted),
                 None)
    if above is None:
        return None
    gone = {q.id for q in paras}
    out[:] = [c for c in out if c.para_id not in gone]
    for q in paras:
        out.append(_change(q, 0, len(q.text), "", Layer.STRUCTURE, "sections_back_matter"))
        deleted.add(q.id)
    origin = provenance.make("manuscript", "Manuscript", "the acknowledgement, moved before the statements (G148)",
                             para_id=p.id)
    return above, [head] + statements, origin


def label_only_text(text: str) -> bool:
    return bool(text.strip()) and _words(text) <= 8 and not text.rstrip().endswith(".")


def _supplementary(ms: Manuscript, style: JournalStyle) -> list[Change]:
    """The stylesheets' "Supplementary Material" heading over supporting information (sections.back_matter.
    supplementary_heading). Off by default: no editor added it in the training files (0 of 6 articles)."""
    heading = _cfg(style, "back_matter.supplementary_heading")
    if not heading or not _on(style, "back_matter.enabled"):
        return []
    tag = "<H2>" if style.get("headings.tag") else ""
    region, _ = _back_region(ms)
    head = re.compile(r"^\s*(?:<H\d>)?\s*(?P<label>(?:supporting|supplementary|supplemental)\s+(?:information|materials?|data"
                      r"|files?|digital\s+content))\s*:?\s*$", re.I)
    for p in region:
        m = head.match(p.text)
        if m:
            if m.group("label") != heading:
                return [_change(p, m.start("label"), m.end("label"), heading, Layer.STRUCTURE, "sections_back_matter")]
            return []
    first = next((p for p in region if re.match(r"^\s*(?:supplementary|supplemental)\s+(?:table|figure|fig\.?|file|material"
                                                r"|video|appendix)\s*S?\d", p.text, re.I)), None)
    if first is None:
        return []
    return [_change(first, 0, 0, f"{tag}{heading}{BR}", Layer.STRUCTURE, "sections_back_matter",
                    {"bold": False, "italic": False})]


# --- author queries ------------------------------------------------------------------------

TEXT = {
    "abstract.citations": "Reference citations should not be present in the abstract. Kindly check.",
    "abstract.too_long": "As per journal style, kindly reduce the abstract to {n} words.",
    "abstract.structured_case_report": "Please note that in the Case Report there shouldn’t be any structured "
                                       "abstract, hence will have to remove the headings.",
    "keywords.too_many": "Kindly provide a maximum of {n} key words.",
    "abstract.heading": "The abstract heading “{head}” is not one of the journal’s headings ({allowed}). Kindly check "
                        "and revise the abstract under these headings.",
    "sections.patient_consent": "Kindly confirm whether consent form obtained or not?",
    "sections.consent_added": "Please ensure that patient consent form has been obtained.",
    "sections.contributions_missing": "Please provide authors’ contributions.",
    "sections.data_availability_missing": "Please provide the missing Data availability statement, if available.",
    "headings.duplicate": "The heading “{head}” is repeated. Kindly check and delete the duplicate.",
    "sections.statement_twice": "The {head} statement is given twice. Kindly check and retain one.",
}


def duplicate_headings(ms: Manuscript, style: JournalStyle) -> list[tuple[Paragraph, int, int, str]]:
    """(paragraph, start, end, heading) for a heading the text gives twice in a row: a heading line, then the same
    heading again as a heading of its own (a line, or a run-in label the heading rules split off its text:
    "Sleep Habits Questionnaire:" over "Sleep Habits Questionnaire: A three-item ...").

    A run-in label left in its paragraph under the same heading ("<H2>Ethical considerations" over "Ethical
    Considerations: The study ...") is not a second heading and is not queried: JEHP's and MJDRDYPU's editors leave
    it as written (jehp_1044, mjdrdypu_601: 2 of 2 kept, no query)."""
    from .rules_format import _tagged_headings

    levels = _tagged_headings(ms, style)
    ends = getattr(ms, "_heading_ends", {}) or {}
    paras = [p for p in ms.paragraphs if p.text.strip() and not p.in_table]
    out = []
    for p, q in zip(paras, paras[1:]):
        if p.id not in levels or ends.get(p.id):
            continue
        head = _key(p.text)
        if not head or len(head) < 4:
            continue
        if q.id not in levels:
            continue  # a run-in label kept in its paragraph: the editors leave it
        cut = ends.get(q.id)
        label = q.text[:cut] if cut else q.text
        if label and _key(label) == head:
            s = len(q.text) - len(q.text.lstrip())
            out.append((q, s, s + len(label.strip()), p.text.strip().rstrip(":").strip()))
    return out


_ABS_LABEL = re.compile(r"^\s*(?P<label>" + _SUBHEADS + r")\s*(?::|\s*[–—-]\s*|\s*$)", re.I)
_ABS_LABEL_ANY = re.compile(r"^\s*(?P<label>[A-Z][A-Za-z]+(?:\s+(?:and\s+|of\s+)?[A-Za-z]+){0,3})\s*:\s*(?=\S)")


def abstract_heading_deviations(lay: Layout, style: JournalStyle, allowed: list[str]) -> list[tuple]:
    """(paragraph, start, end, heading) for each structured-abstract sub-heading that is not one of the sheet's
    (after the sheet's renames: "Aim" -> "Purpose", "Materials and Methods" -> "Methods"): IJO_2733 "Design-",
    "Setting-" (S: "Raise query in case of any deviation from these headings"). Only for an abstract with two such
    headings at least (a structured one)."""
    rename = {k.lower(): v for k, v in (_cfg(style, "abstract.rename") or {}).items()}
    ok = {a.lower() for a in allowed} | {k for k, v in rename.items() if v.lower() in {a.lower() for a in allowed}}
    ok |= {a + "s" for a in ok}  # "Conclusions", "Results" / "Result"
    ok |= {a[:-1] for a in ok if a.endswith("s")}
    found = []
    for p in lay.abstract:
        m = _ABS_LABEL.match(p.text) or _ABS_LABEL_ANY.match(p.text)
        if m:
            found.append((p, m.start("label"), m.end("label"), m.group("label")))
    if len(found) < 2:
        return []
    return [(p, s, e, head) for p, s, e, head in found if re.sub(r"\s+", " ", head).strip().lower() not in ok]


def add_queries(ms: Manuscript, style: JournalStyle, qs) -> None:
    """Author queries for the abstract, key words and back matter (see queries.collect)."""
    from .rules import _ROUND, _SQUARE, enumeration_markers, superscript_citations

    if not _cfg(style, "enabled", True):
        return
    lay = layout(ms)
    if style.get("headings.duplicate_query", True):
        for q, s, e, head in duplicate_headings(ms, style):
            qs.add(q.id, "headings.duplicate", s, e, head=re.sub(r"<H\d>", "", head))
    if lay.abstract:
        if _on(style, "abstract.citation_query"):
            for p in lay.abstract:
                spans = [m.span() for m in _SQUARE.finditer(p.text)]
                listed = enumeration_markers(p.text)  # "(1) relational ...; and (4) boundary-spanning" (997)
                spans += [m.span() for m in _ROUND.finditer(p.text) if m.start() not in listed
                          and re.search(r"[A-Za-z]\s*$", p.text[:m.start()]) and int(re.findall(r"\d+", m.group(2))[0]) < 200]
                spans += [(s, e) for s, e, _ in superscript_citations(p) if s]
                if spans:
                    s, e = min(spans)
                    qs.add(p.id, "abstract.citations", s, e)
                    break
        if _on(style, "abstract.word_limit_query"):
            structured = _structured(lay)
            limit = _cfg(style, "abstract.word_limit.structured" if structured else "abstract.word_limit.unstructured")
            n = sum(_abstract_words(p.text) for p in lay.abstract)
            if lay.keywords_at:
                n -= _word_count(lay.keywords.text[lay.keywords_at:])  # the key words' line ends the last paragraph
            if limit and n > limit:
                qs.add(lay.abstract[0].id, "abstract.too_long", n=limit)
        allowed = _cfg(style, "abstract.headings") or []
        if allowed and _cfg(style, "abstract.heading_query", True):
            for p, s, e, head in abstract_heading_deviations(lay, style, allowed):
                qs.add(p.id, "abstract.heading", s, e, head=head, allowed=", ".join(allowed))
        if _on(style, "abstract.structured_case_report_query") and _structured(lay) and _is_case_report(ms):
            qs.add((lay.abstract_head or lay.abstract[0]).id, "abstract.structured_case_report")
    if lay.keywords is not None and _on(style, "keywords.max_query"):
        kws = listed_keywords(ms)
        hi = style.get("keywords.max")
        if hi and len(kws) > hi:
            qs.add(lay.keywords.id, "keywords.too_many", n=hi)
    elif lay.keywords is None and lay.abstract and _on(style, "keywords.add_missing") \
            and not (_on(style, "keywords.from_xml") and xml_keywords(ms)):
        qs.add(lay.abstract[-1].id, "keywords.missing")
    if _cfg(style, "back_matter.enabled", True):
        ref_i = references.find_heading(ms)
        at = ms.paragraphs[ref_i].id if ref_i is not None else None
        plan = back_matter_plan(ms, style, run=True)
        consent_cfg = _cfg(style, "back_matter.consent") or {}
        if plan.get("consent") == "added" and consent_cfg.get("query", True) and at:
            qs.add(at, "sections.consent_added")
        elif _on(style, "back_matter.patient_consent_query") and _is_case_report(ms) \
                and not has_consent_statement(ms) and at:
            qs.add(at, "sections.patient_consent")
        for pid, head in plan.get("twice") or []:
            qs.add(pid, "sections.statement_twice", head=head.lower())
        if plan.get("contributions") == "missing" and (_cfg(style, "back_matter.contributions") or {}).get(
                "query_missing") and at:
            qs.add(at, "sections.contributions_missing")
        if plan.get("data") == "missing" and (_cfg(style, "back_matter.data_availability") or {}).get(
                "query_missing") and at:
            qs.add(at, "sections.data_availability_missing")
